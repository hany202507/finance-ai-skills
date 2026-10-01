# -*- coding: utf-8 -*-
"""장부·규칙을 읽고 엑셀을 쓰는 공용 함수. 매입매출장검토의 review.py 와 같은 내용이다.

두 스킬은 서로 불러 쓰지 않는다. 하나만 설치해도 돌아야 한다. 그래서 여기 한 벌을 더 둔다.
고칠 일이 있으면 두 곳을 같이 고친다.
"""
import os
import io
import csv
import glob
from collections import defaultdict


def _rows_from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    head = next(it, None)
    if not head:
        return []
    head = [str(c).strip() if c is not None else "" for c in head]
    out = []
    for r in it:
        if all(c is None for c in r):
            continue
        out.append({head[k]: r[k] for k in range(min(len(head), len(r)))})
    wb.close()
    return out


def _rows_from_csv(path):
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load(folder):
    """폴더의 엑셀과 CSV 를 표 이름 -> 행목록 으로 올린다. 같은 이름이면 엑셀."""
    tables = {}
    for path in sorted(glob.glob(os.path.join(folder, "*.csv"))) + \
            sorted(glob.glob(os.path.join(folder, "**", "*.csv"), recursive=True)):
        tables[os.path.splitext(os.path.basename(path))[0]] = ("csv", path)
    for path in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
        if os.path.basename(path).startswith("~$"):
            continue
        tables[os.path.splitext(os.path.basename(path))[0]] = ("xlsx", path)
    out = {}
    for name, (kind, path) in tables.items():
        out[name] = (_rows_from_xlsx if kind == "xlsx" else _rows_from_csv)(path)
    return out


def pick(tables, prefix):
    for name in tables:
        if name.startswith(prefix):
            return tables[name]
    return []


def n(x):
    if x in (None, ""):
        return 0
    if isinstance(x, (int, float)):
        return int(x)
    try:
        return int(float(str(x).replace(",", "")))
    except ValueError:
        return 0


def s(x):
    return "" if x is None else str(x).strip()


def d3(r):
    return f"{n(r.get('년도')):04d}-{n(r.get('월')):02d}-{n(r.get('일')):02d}"


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    head = PatternFill("solid", fgColor="1F3864")
    neg = PatternFill("solid", fgColor="FCE4E4")
    thin = Side(style="thin", color="8EA9DB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    wb.remove(wb.active)
    for title, cols, rows, widths in sheets:
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF", size=9)
            cell.fill = head
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = box
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            has_neg = any(isinstance(c.value, (int, float)) and c.value < 0 for c in row)
            for c in row:
                c.border = box
                c.font = Font(size=9)
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = "#,##0"
                if has_neg:
                    c.fill = neg
    wb.save(path)


def coverage_md(findings, cat, skipped=None):
    """진단항목 전수에 대해 무엇이 돌았고 무엇이 안 돌았나. rules/진단항목.csv 가 없으면 비운다.
    skipped 는 자동 항목인데 원천이나 규칙이 없어 못 돈 것. 통과로 찍지 않는다."""
    if not cat:
        return []
    skipped = skipped or {}
    hit = defaultdict(int)
    for f in findings:
        hit[s(f["검사"])] += 1
    cnt = defaultdict(int)
    L = ["## 검토 범위", "",
         "검토항목 전수입니다. 자동은 이 스크립트가 대조한 것, 검토자는 회계사가 검토 방법대로 "
         "직접 보고 확인사항에 적는 것, 자료 미확보는 원천자료가 없어 지금은 못 보는 것입니다.", "",
         "| 번호 | 검토 항목 | 판정 | 수행 | 결과 |", "|---|---|---|---|---|"]
    for r in cat:
        no, how = s(r["번호"]), s(r["수행"])
        cnt[how] += 1
        if how == "자동" and no in skipped and not hit[no]:
            res = f"미실행 ({skipped[no]})"
        elif how == "자동":
            res = f"검출 {hit[no]}건" if hit[no] else "통과"
        elif how == "검토자":
            res = "검토자 수행 필요"
        else:
            res = "자료 확보 후"
        L.append(f"| {no} | {r['검토 항목']} | {r['판정']} | {how} | {res} |")
    L += ["", f"자동 {cnt['자동']} · 검토자 {cnt['검토자']} · 자료 미확보 {cnt['자료 미확보']}. "
          "검토자 수행 필요와 미실행은 통과가 아닙니다. 검토 방법은 rules/진단항목.csv 에 있습니다.", ""]
    return L

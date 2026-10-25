"""ctax ar 2회차 시험: AI 메모 검사(--memo) · 회계사 결론 이어받기 · 결손금 이력 점검 · 근거 조문 · 작은 조정 후보 접음.

고객 자료 없이 test_ar 의 합성 장부 · 전표 · 전기로 돌린다.
"""
from __future__ import annotations

import json

import openpyxl

from ctax.ar import cli, engine, excel, memo
from ctax.tests.test_ar import acc, by_title, ledger, prior, run, write_case

URL = "https://www.law.go.kr/법령/법인세법/제25조"


def law(grade="★★★", url=URL):
    x = {"ref": "법인세법 제25조 제4항", "grade": grade}
    if url:
        x["url"] = url
    return x


def good_memo(ar: dict) -> dict:
    """AR.json 이슈 전부에 대한 통과 메모. 금액은 이슈 amounts 값을 그대로 쓴다."""
    items = []
    for i in ar["issues"]:
        a = i.get("amounts") or {}
        amt = next((a[k] for k in ("current", "prior", "change") if a.get(k)), None)
        why = f"장부 금액 {abs(amt):,}원을 기준으로 봅니다" if amt else "장부 금액이 없어 회사 답을 보고 정합니다"
        items.append({"issue": i["no"], "judgment": "판단 필요", "why": why, "law": [law()],
                      "questions": ["관련 명세를 보내 주십시오"]})
    return {"items": items}


def read_ar(case_dir) -> dict:
    return json.loads((case_dir / "_색인" / "AR.json").read_text(encoding="utf-8"))


def issue_sheet(case_dir):
    wb = openpyxl.load_workbook(case_dir / "1 AR" / "AR_이슈.xlsx")
    ws = wb["이슈 목록"]
    head = next(r for r in range(1, 6) if ws.cell(r, 1).value == "번호")
    col = {ws.cell(head, c).value: c for c in range(1, ws.max_column + 1)}
    return wb, ws, head, col


def row_of(ws, head, no):
    return next(r for r in range(head + 1, ws.max_row + 1) if ws.cell(r, 1).value == no)


# ---------------------------------------------------------------- 메모 검사

def test_memo_validate_passes_and_numbers_come_from_ar():
    ar = run()
    assert memo.validate(good_memo(ar), ar) == []
    # 근거 글(basis) 속 금액도 AR.json 값이다
    ent = by_title(ar, "기업업무추진비 한도")[0]
    m = good_memo(ar)
    item = next(x for x in m["items"] if x["issue"] == ent["no"])
    item["why"] = "전기 15호에 1,000,000원 한도초과가 있었고 당기 30,000,000원입니다. 기본한도는 1천200만원입니다"
    assert memo.validate(m, ar) == []


def test_memo_rejects_numbers_not_in_ar():
    ar = run()
    m = good_memo(ar)
    m["items"][0]["why"] = "한도초과액은 1,234,567원으로 보입니다"
    m["items"][1]["questions"] = ["12345678원 지출의 명세를 보내 주십시오"]
    probs = memo.validate(m, ar)
    assert any("AR.json 에 없는 금액 1,234,567" in p for p in probs)
    assert any("AR.json 에 없는 금액 12345678" in p for p in probs)


def test_memo_rejects_missing_extra_and_duplicate_issues():
    ar = run()
    m = good_memo(ar)
    last = m["items"].pop()
    m["items"].append({**m["items"][0]})
    m["items"].append({**last, "issue": 99})
    probs = memo.validate(m, ar)
    assert len(probs) == 1 and "빠진 이슈 [" + str(last["issue"]) + "]" in probs[0]
    assert "AR.json 에 없는 이슈 [99]" in probs[0] and "겹친 이슈 [1]" in probs[0]
    assert memo.validate({"memo": []}, ar)


def test_memo_rejects_bad_grade_url_judgment_and_style():
    ar = run()
    m = good_memo(ar)
    m["items"][0]["law"] = [law(grade=None)]
    m["items"][1]["law"] = [law(url=None)]                                  # ★★★ 는 url 필수
    m["items"][2]["law"] = [law(grade="☆", url="https://example.com/x")]
    m["items"][3]["law"] = []
    m["items"][4]["judgment"] = "조정 대상"
    m["items"][5]["why"] = "한도를 넘나요?"
    m["items"][6]["why"] = "한도 안입니다 — 조정 없음"
    m["items"][7]["questions"] = ["명세를 보내 줄 수 있나요"]
    probs = "\n".join(memo.validate(m, ar))
    assert "근거 1: grade" in probs
    assert "url 이 있어야" in probs
    assert "law.go.kr 주소" in probs
    assert "한 줄 이상" in probs
    assert "judgment" in probs
    assert "물음표" in probs
    assert "줄표" in probs
    assert "주십시오" in probs
    # ☆ 는 url 없이 된다(조문을 찾지 못한 사실을 적는 줄)
    m = good_memo(ar)
    m["items"][0]["law"] = [{"ref": "korean-law 도구가 응답하지 않아 원문을 확인하지 못했습니다", "grade": "☆"}]
    assert memo.validate(m, ar) == []


def test_cli_memo_puts_or_rejects(case_dir, capsys):
    write_case(case_dir)
    assert cli.main([str(case_dir)]) == 0
    ar = read_ar(case_dir)
    bad = good_memo(ar)
    bad["items"][0]["why"] = "9,999,999원"
    path = case_dir.parent / "메모.json"
    path.write_text(json.dumps(bad, ensure_ascii=False), encoding="utf-8")
    assert cli.main([str(case_dir), "--memo", str(path)]) == 1
    assert "아무것도 넣지 않았습니다" in capsys.readouterr().out
    assert all(i["memo"] is None for i in read_ar(case_dir)["issues"])

    path.write_text(json.dumps(good_memo(ar), ensure_ascii=False), encoding="utf-8")
    assert cli.main([str(case_dir), "--memo", str(path)]) == 0
    ar2 = read_ar(case_dir)
    for i in ar2["issues"]:
        m = i["memo"]
        assert m["by"] == "AI" and m["confirmed"] is False and m["at"] and m["judgment"] == "판단 필요"
        assert m["law"][0]["grade"] == "★★★" and m["law"][0]["url"] == URL
    assert engine.validate(ar2) == []
    _wb, ws, head, col = issue_sheet(case_dir)
    r = row_of(ws, head, 1)
    assert ws.cell(r, col["AI 판단(확인 전)"]).value == "판단 필요"
    assert "회사에 물을 것" in ws.cell(r, col["AI 까닭 · 질문"]).value
    assert "★★★" in ws.cell(r, col["AI 근거 조문 · 등급"]).value
    step = json.loads((case_dir / "_색인" / "신고건.json").read_text(encoding="utf-8"))["steps"]["ar"]
    assert "AI 메모" in step["note"]
    # 메모 파일 · AR.json 이 없으면 끝 코드 2
    assert cli.main([str(case_dir), "--memo", str(case_dir / "없음.json")]) == 2


def test_cli_memo_without_ar(case_dir):
    write_case(case_dir)
    path = case_dir.parent / "메모.json"
    path.write_text('{"items": []}', encoding="utf-8")
    assert cli.main([str(case_dir), "--memo", str(path)]) == 2


# ---------------------------------------------------------------- 이어받기

def test_conclusion_carried_from_excel(case_dir):
    write_case(case_dir)
    assert cli.main([str(case_dir)]) == 0
    ar = read_ar(case_dir)
    ent = by_title(ar, "기업업무추진비 한도")[0]
    wb, ws, head, col = issue_sheet(case_dir)
    ws.cell(row_of(ws, head, ent["no"]), col["회계사 결론"]).value = "한도초과 손금불산입. 23호 작성"
    # 같은 key 가 둘인 전기 연속 줄(rule 로 가린다)에도 적는다
    res = by_title(ar, "전기 유보 미수수익")[0]
    ws.cell(row_of(ws, head, res["no"]), col["회계사 결론"]).value = "추인 대상"
    wb.save(case_dir / "1 AR" / "AR_이슈.xlsx")

    assert cli.main([str(case_dir)]) == 0
    ar2 = read_ar(case_dir)
    c = by_title(ar2, "기업업무추진비 한도")[0]["conclusion"]
    assert c["text"] == "한도초과 손금불산입. 23호 작성" and c["by"] == "회계사" and c["at"]
    assert by_title(ar2, "전기 유보 미수수익")[0]["conclusion"]["text"] == "추인 대상"
    others = [i for i in ar2["issues"] if i["conclusion"] and i["rule"] not in ("entertainment", "reserve:미수수익")]
    assert others == []
    # 다시 돌려도 엑셀에 그대로 있고 at 은 바뀌지 않는다
    assert cli.main([str(case_dir)]) == 0
    ar3 = read_ar(case_dir)
    assert by_title(ar3, "기업업무추진비 한도")[0]["conclusion"]["at"] == c["at"]
    _wb, ws, head, col = issue_sheet(case_dir)
    assert ws.cell(row_of(ws, head, ent["no"]), col["회계사 결론"]).value == "한도초과 손금불산입. 23호 작성"


def test_conclusion_survives_memo_and_memo_survives_rerun(case_dir):
    write_case(case_dir)
    assert cli.main([str(case_dir)]) == 0
    ar = read_ar(case_dir)
    wb, ws, head, col = issue_sheet(case_dir)
    ws.cell(row_of(ws, head, 1), col["회계사 결론"]).value = "특수관계인 아님(직원 대여)"
    wb.save(case_dir / "1 AR" / "AR_이슈.xlsx")
    path = case_dir.parent / "메모.json"
    path.write_text(json.dumps(good_memo(ar), ensure_ascii=False), encoding="utf-8")
    assert cli.main([str(case_dir), "--memo", str(path)]) == 0
    ar2 = read_ar(case_dir)
    assert ar2["issues"][0]["conclusion"]["text"] == "특수관계인 아님(직원 대여)"
    _wb, ws, head, col = issue_sheet(case_dir)
    assert ws.cell(row_of(ws, head, 1), col["회계사 결론"]).value == "특수관계인 아님(직원 대여)"
    # 장부가 그대로면 다시 돌려도 메모가 남는다
    assert cli.main([str(case_dir)]) == 0
    ar3 = read_ar(case_dir)
    assert all(i["memo"] and i["memo"]["by"] == "AI" for i in ar3["issues"])


def test_memo_dropped_when_amount_changes():
    def with_ent(amount):
        return ledger(accounts=[a if a["code"] != "813" else acc("813", "접대비(기업업무추진비)(판)", "비용", amount, 25_000_000, lines=40)
                                for a in ledger()["accounts"]])
    ar = run(with_ent(33_333_333))
    m = good_memo(ar)
    assert memo.validate(m, ar) == []
    memo.apply(m, ar)
    new = run(with_ent(34_444_444))                                       # 접대비 금액이 바뀐 새 실행
    notes = memo.carry(new, ar, None)
    ent = by_title(new, "기업업무추진비 한도")[0]
    assert ent["memo"] is None
    assert any(f"이슈 {ent['no']} AI 메모는 금액이 바뀌어" in n for n in notes)
    assert by_title(new, "법인세등을 손금불산입")[0]["memo"]["by"] == "AI"


def test_conclusion_rows_from_old_excel_without_key(tmp_path):
    """key · rule 열이 없는 1회차 엑셀은 종류 · 제목으로 잇는다."""
    ar = run()
    path = tmp_path / "AR_이슈.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "이슈 목록"
    ws.append(["분석적 검토 이슈 목록"])
    ws.append([])
    ws.append(["번호", "종류", "확실도", "가이드", "계정", "제목", "회계사 결론"])
    ent = by_title(ar, "기업업무추진비 한도")[0]
    ws.append([7, ent["kind"], ent["certainty"], "", "", ent["title"], "1회차 결론"])
    wb.save(path)
    rows = excel.read_conclusions(path)
    assert rows == [{"no": 7, "key": None, "rule": None, "kind": ent["kind"], "title": ent["title"], "text": "1회차 결론"}]
    memo.carry(ar, None, rows)
    assert by_title(ar, "기업업무추진비 한도")[0]["conclusion"]["text"] == "1회차 결론"
    assert excel.read_conclusions(tmp_path / "없음.xlsx") is None


# ---------------------------------------------------------------- 결손금 이력

def test_loss_missing_row_asks_for_older_returns():
    pr = prior(cells={"3": {"06": -7_000_000}}, losses=[])
    hit = by_title(run(pr=pr), "50호(갑) 이월결손금 줄에 없습니다")[0]
    assert hit["kind"] == "신고 점검" and hit["certainty"] == "확실" and hit["papers"] == ["loss"]
    assert any("전전기 이전 신고서로 결손금 이력" in b for b in hit["basis"])
    assert any("결손금이 처음 난 사업연도부터" in q for q in hit["questions"])
    # 5회차에 법인세법 제13조 제1항 제1호 원문을 확인해 ★★★ 로 올렸다
    assert hit["law"] and hit["law"][0]["grade"] == "★★★" and hit["law"][0]["ref"].startswith("법인세법 제13조 제1항 제1호")
    assert not by_title(run(pr=pr), "결손금이 난 해의 신고서")          # 같은 사건을 두 줄로 내지 않는다


def test_loss_history_signals():
    # 이월결손금 공제가 있었는데 표가 빔
    ar = run(pr=prior(cells={"3": {"06": 0, "07": 4_000_000}}, losses=[]))
    hit = by_title(ar, "결손금이 난 해의 신고서로 확인해야 합니다")
    assert hit and hit[0]["kind"] == "신고 점검" and hit[0]["certainty"] == "가능"
    assert any("07 이월결손금 공제 4,000,000원" in b for b in hit[0]["basis"])
    # 표 줄 잔액이 발생 − 공제와 다름
    pr = prior(losses=[{"year": "2022", "incurred": 10_000_000, "used": 3_000_000, "remaining": 9_000_000}])
    hit = by_title(run(pr=pr), "결손금이 난 해의 신고서로 확인해야 합니다")
    assert hit and any("서로 맞지 않습니다" in b for b in hit[0]["basis"])
    # 전기 재무상태표 미처리결손금이 표 잔액보다 큼
    fs = {"bs": [{"label": "Ⅲ. 이익잉여금(결손금)", "cur": -50_000_000, "prior": -40_000_000, "code": None, "kind": "subtotal"}], "is": []}
    pr = prior(losses=[{"year": "2023", "incurred": 10_000_000, "used": 0, "remaining": 10_000_000}])
    hit = by_title(run(ledger(fs=fs), pr=pr), "결손금이 난 해의 신고서로 확인해야 합니다")
    assert hit and any("이익잉여금 △40,000,000원" in b and "10,000,000원보다 큽니다" in b for b in hit[0]["basis"])
    assert hit[0]["amounts"] == {"prior": 10_000_000}
    # 신호가 없으면 서지 않는다
    pr = prior(losses=[{"year": "2023", "incurred": 50_000_000, "used": 0, "remaining": 50_000_000}])
    assert not by_title(run(ledger(fs=fs), pr=pr), "결손금이 난 해의 신고서")
    assert not by_title(run(), "결손금이 난 해의 신고서")


# ---------------------------------------------------------------- 근거 조문 · 접음 · key

def test_issue_law_key_rule_and_folded():
    ar = run()
    for i in ar["issues"]:
        assert isinstance(i["law"], list) and isinstance(i["folded"], bool)
        assert i["key"] == f"{i['kind']}|{','.join(i['guide'])}|{','.join(i['accounts'])}"
        assert i["rule"]
        for x in i["law"]:
            assert x["grade"] in engine.GRADES and x["ref"]
            if x["grade"] == "★★★":
                assert x["url"].startswith("https://www.law.go.kr/")
    assert by_title(ar, "기업업무추진비 한도")[0]["law"][0]["ref"].startswith("법인세법 제25조")
    assert by_title(ar, "중소기업 특별세액감면을 당기에도")[0]["law"][0]["ref"].startswith("조세특례제한법 제7조")
    rules = [i["rule"] for i in ar["issues"]]
    assert len(rules) == len(set(rules))


def test_small_candidates_fold_last_in_bucket(case_dir):
    accounts = ledger()["accounts"] + [acc("279", "판매보증충당부채", "부채", 1_200_000, 1_000_000),
                                       acc("971", "외화환산이익", "수익", 2_500_000, 0)]
    ar = run(ledger(accounts=accounts))
    small = [i for i in ar["issues"] if i["folded"]]
    assert {i["rule"] for i in small} == {"provision", "fx"}
    assert all(i["kind"] == "조정 후보" and i["certainty"] == "가능" for i in small)
    possible = [i for i in ar["issues"] if Builder_bucket(i) == 3]
    assert [i["folded"] for i in possible] == sorted(i["folded"] for i in possible)    # 접은 줄은 가능 묶음 끝
    assert by_title(ar, "감가상각 시부인")[0]["folded"] is False                      # 8,000,000원은 접지 않는다
    assert engine.validate(ar) == []
    # 엑셀: 머리 줄 하나와 숨긴 줄
    write_case(case_dir)
    (case_dir / "_색인" / "장부.json").write_text(json.dumps(ledger(accounts=accounts), ensure_ascii=False), encoding="utf-8")
    assert cli.main([str(case_dir)]) == 0
    _wb, ws, head, col = issue_sheet(case_dir)
    marks = [r for r in range(head + 1, ws.max_row + 1) if ws.cell(r, 1).value == excel.FOLD_MARK]
    assert len(marks) == 1 and "작은 조정 후보 2건" in ws.cell(marks[0], 6).value
    hidden = [r for r in range(head + 1, ws.max_row + 1) if ws.row_dimensions[r].hidden]
    assert hidden == [marks[0] + 1, marks[0] + 2]
    assert all(ws.row_dimensions[r].outline_level == 1 for r in hidden)
    from openpyxl.utils import get_column_letter
    assert ws.column_dimensions[get_column_letter(excel.COL["key"])].hidden and ws.cell(head, col["key"]).value == "key"


def Builder_bucket(i):
    return engine.Builder.bucket(i)


def test_prior_reserve_upgrades_candidate():
    """전기 50호(을) 유보가 당기 조정 후보에 붙으면 추인 조정이 따르므로 확실로 올리고 전기 처리에 적는다."""
    accounts = ledger()["accounts"] + [acc("109", "대손충당금", "자산", -5_000_000, -4_000_000)]
    pr = prior(reserves=[{"subject": "대손충당금", "opening": 1_000_000, "decrease": 0, "increase": 0, "closing": 1_000_000}])
    bd = by_title(run(ledger(accounts=accounts), pr=pr), "대손충당금 한도")[0]
    assert bd["certainty"] == "확실" and "전기 50호(을) 대손충당금 기말 유보 1,000,000원" in bd["prior_treatment"]
    assert bd["amounts"] == {"prior": -4_000_000, "current": -5_000_000, "change": -1_000_000}


def test_subject_keys_for_new_prior_lines():
    from ctax.ar.engine import subject_key
    assert subject_key("일시상각충당금", "당기지원금은감가상각비에대하여손금불산입") == "grant"
    assert subject_key("일반보조금") == "grant"
    assert subject_key("무형자산상각비") == "depreciation"
    assert subject_key("단기투자자산") == "securities"
    assert subject_key("단기금융상품평가이익") == "securities"
    assert subject_key("직원급여") == "accrued_expense"
    assert subject_key("잡이익", "법인세환급액익금불산입") == "tax_refund"
    assert subject_key("외상매출금", "매출중복과다계상") == "receivable"
    assert subject_key("퇴직급여충당부채") == "retirement_allowance"


# ---------------------------------------------------------------- 2회차 규칙 보강(열 곳 돌려 본 뒤)

def test_prior_local_withholding_blank():
    forms = prior()["forms"] + [{"code": "10갑", "name": "10갑", "page": 2, "pages": 1, "blank": False, "known": True},
                                {"code": "지방43의5갑", "name": "지방43의5갑", "page": 3, "pages": 1, "blank": True, "known": True}]
    hit = by_title(run(pr=prior(forms=forms)), "특별징수세액명세서가 비었습니다")
    assert hit and hit[0]["kind"] == "신고 점검" and hit[0]["certainty"] == "가능" and "지방43의5갑" in hit[0]["forms"]
    forms[-1]["blank"] = False
    assert not by_title(run(pr=prior(forms=forms)), "특별징수세액명세서가 비었습니다")


def test_rnd_credit_uses_prior_credit_once():
    accounts = ledger()["accounts"] + [acc("826", "경상연구개발비(판)", "비용", 50_000_000, 40_000_000)]
    pr = prior(credits=[{"name": "일반연구ㆍ인력개발비세액공제", "code": "16B", "amount": 6_000_000, "carryforward": 0}])
    ar = run(ledger(accounts=accounts), pr=pr)
    rnd = by_title(ar, "연구 · 인력개발비 세액공제 대상인지")
    assert len(rnd) == 1 and rnd[0]["certainty"] == "가능" and "1-4" in rnd[0]["guide"]
    assert any("전기에 일반연구ㆍ인력개발비세액공제 6,000,000원을 받았습니다" in b for b in rnd[0]["basis"])
    assert not by_title(ar, "전기 연구 · 인력개발비 세액공제를 당기에도")


def test_reserve_joins_living_prior_line_and_foreign_securities():
    accounts = ledger()["accounts"] + [acc("110", "받을어음", "자산", 9_000_000, 8_000_000),
                                       acc("178", "해외투자주식", "자산", 5_000_000, 0)]
    pr = prior(adjustments=[{"side": "차감", "subject": "외상매출금", "amount": 4_000_000, "disposition": "유보", "code": "400",
                             "note": "매출중복과다계상분익금불산입"}],
               reserves=[{"subject": "외상매출금", "opening": 0, "decrease": 0, "increase": -4_000_000, "closing": -4_000_000}])
    ar = run(ledger(accounts=accounts), pr=pr)
    rec = [i for i in ar["issues"] if i["rule"].startswith("prior:receivable")]
    assert len(rec) == 1 and "유보 △4,000,000원이 남아 있습니다" in rec[0]["title"]
    assert "reserve_50" in rec[0]["papers"] and "50을" in rec[0]["forms"]
    assert not [i for i in ar["issues"] if i["rule"].startswith("reserve:")]
    sec = by_title(ar, "유가증권")[0]
    assert "2-3-1-4" in sec["guide"] and "해외현지법인명세서" in sec["forms"]
    assert any(x["ref"].startswith("국제조세조정에 관한 법률") for x in sec["law"])


def test_memo_may_use_workpaper_amounts():
    """10-19: 조서 계산 금액(조정.json)을 메모 글에 쓸 수 있다."""
    ar = run()
    m = good_memo(ar)
    m["items"][0]["why"] = "조서가 계산한 한도초과 1,234,567원을 손금불산입합니다."
    assert any("1,234,567" in p for p in memo.validate(m, ar))
    assert not any("1,234,567" in p for p in memo.validate(m, ar, {1_234_567}))

"""ctax ar 시험. 고객 자료 없이 합성 장부 · 전표 · 전기로 돌린다."""
from __future__ import annotations

import csv
import json

import openpyxl
import pytest
import yaml

from ctax.ar import cli, engine
from ctax.ar.guide import HERE, load_guide

CLOSE = "손익계정에 대체"


def acc(code, name, cls, amount, prior=0, lines=1, debit=None, credit=None, closing=None):
    debit = debit if debit is not None else (amount if cls in ("자산", "비용") and amount > 0 else 0)
    credit = credit if credit is not None else (amount if cls in ("부채", "자본", "수익") and amount > 0 else 0)
    return {"code": code, "name": name, "class": cls, "section": "", "opening": 0, "debit": debit, "credit": credit,
            "closing": closing if closing is not None else debit - credit, "amount": amount, "prior_amount": prior,
            "fs_label": name, "lines": lines}


def ledger(**over) -> dict:
    accounts = [
        acc("103", "보통예금", "자산", 50_000_000, 40_000_000),
        acc("108", "외상매출금", "자산", 30_000_000, 20_000_000),
        acc("137", "주.임.종단기채권", "자산", 70_000_000, 70_000_000, lines=0),
        acc("208", "차량운반구", "자산", 40_000_000, 40_000_000, lines=0),
        acc("209", "감가상각누계액", "자산", -16_000_000, -8_000_000),
        acc("295", "퇴직급여충당부채", "부채", 12_000_000, 0),
        acc("331", "자본금", "자본", 100_000_000, 100_000_000, lines=0),
        acc("400", "손익", "손익대체", 0, 0, lines=3),
        acc("401", "상품매출", "수익", 900_000_000, 800_000_000),
        acc("802", "직원급여(판)", "비용", 120_000_000, 100_000_000),
        acc("806", "퇴직급여(판)", "비용", 12_000_000, 5_000_000),
        acc("813", "접대비(기업업무추진비)(판)", "비용", 30_000_000, 25_000_000, lines=40),
        acc("817", "세금과공과금(판)", "비용", 3_000_000, 2_000_000, lines=10),
        acc("818", "감가상각비(판)", "비용", 8_000_000, 8_000_000),
        acc("822", "차량유지비(판)", "비용", 9_000_000, 7_000_000, lines=12),
        acc("901", "이자수익", "수익", 200_000, 150_000),
        acc("933", "기부금", "비용", 2_000_000, 0, lines=2),
        acc("960", "잡손실", "비용", 500_000, 0, lines=2),
        acc("998", "법인세등", "비용", 4_000_000, 3_000_000),
    ]
    data = {
        "source": "wehago", "files": [], "accounts": accounts,
        "partners": [{"account": "137", "partner_code": "000001", "partner": "가상 임원", "opening": 70_000_000,
                      "debit": 0, "credit": 0, "closing": 70_000_000}],
        "fs": {"bs": [], "is": []},
        "totals": {"revenue": 900_000_000, "net_income": 40_000_000, "income_before_tax": 44_000_000,
                   "corporate_tax": 4_000_000, "assets": 174_000_000, "liabilities": 12_000_000, "equity": 162_000_000},
        "checks": [{"id": "balanced", "ok": True, "title": "차변합 = 대변합", "detail": "", "amounts": {"diff": 0}},
                   {"id": "closing_residual", "ok": True, "title": "손익(400) 계정 잔액 0", "detail": "",
                    "amounts": {"손익 400 잔액": 0}}],
        "journal_csv": "_색인/전표.csv",
    }
    data.update(over)
    return data


def journal() -> list[dict]:
    rows = [
        ("813", "접대비(기업업무추진비)(판)", 25_000_000, 0, "거래처 접대비/신용카드(법인)"),
        ("813", "접대비(기업업무추진비)(판)", 5_000_000, 0, "거래처 경조사비 지급"),
        ("822", "차량유지비(판)", 6_000_000, 0, "렌트카 대여료 12가3456"),
        ("822", "차량유지비(판)", 3_000_000, 0, "주유대"),
        ("817", "세금과공과금(판)", 2_800_000, 0, "재산세"),
        ("817", "세금과공과금(판)", 200_000, 0, "주정차 과태료"),
        ("960", "잡손실", 450_000, 0, "부가세 가산세"),
        ("960", "잡손실", 50_000, 0, "단수차이"),
        ("933", "기부금", 2_000_000, 0, "가상 재단"),
        ("813", "접대비(기업업무추진비)(판)", 0, 30_000_000, CLOSE),
    ]
    out = []
    for i, (code, name, dr, cr, memo) in enumerate(rows):
        out.append({"date": "2025-06-30", "no": f"{i:05d}", "kind": "차변", "account": code, "account_name": name,
                    "debit": dr, "credit": cr, "partner_code": "", "partner": "", "memo_code": "", "memo": memo,
                    "closing": memo == CLOSE})
    return out


def prior(**over) -> dict:
    data = {
        "pdf": "00 받은자료/전기/가상.pdf", "fy": "2024", "fy_end": "2024-12-31",
        "forms": [{"code": c, "name": c, "page": 1, "pages": 1, "blank": False, "known": True}
                  for c in ("1", "3", "15", "23갑", "23을", "29", "50갑", "세금공과금명세서", "8갑")],
        "company": {"name": "가상회사", "sme": True, "kind_code": "30"},
        "cells": {"3": {"01": 30_000_000, "06": 36_000_000}, "50갑": {}, "1": {}},
        "adjustments": [
            {"side": "가산", "subject": "접대비(기업업무추진비)", "amount": 1_000_000, "disposition": "기타사외유출", "code": "500",
             "note": "접대비한도초과액손금불산입"},
            {"side": "가산", "subject": "법인세등", "amount": 3_000_000, "disposition": "기타사외유출", "code": "500", "note": ""},
            {"side": "가산", "subject": "미수이자", "amount": 300_000, "disposition": "유보", "code": "400", "note": ""},
        ],
        "reserves": [{"subject": "미수수익", "opening": -300_000, "decrease": 0, "increase": -300_000, "closing": -300_000}],
        "losses": [],
        "credits": [{"name": "중소기업에 대한 특별세액감면", "code": "112", "amount": 500_000, "carryforward": 0}],
        "problems": [],
    }
    data.update(over)
    return data


def run(led=None, jr=None, pr="default"):
    return engine.build(led or ledger(), jr if jr is not None else journal(), prior() if pr == "default" else pr,
                        {"company": "가상회사", "fy": "2025"})


def by_title(ar, word):
    return [i for i in ar["issues"] if word in i["title"]]


# ---------------------------------------------------------------- 가이드

def test_guide_yaml_counts_and_text():
    g = load_guide()
    assert len(g) == 39
    assert sum(x["id"].startswith("1-") for x in g) == 4
    assert sum(x["id"] in ("2-1", "2-2") for x in g) == 2
    assert sum(x["id"].startswith("2-3-") for x in g) == 33
    row = next(x for x in g if x["id"] == "2-3-16-1")
    assert row["title"] == "16) 기업업무추진비"
    assert row["text"] == "원장을 통해 직부인대상확인 및 한도시부인"
    assert row["raw"] == ": 원장을 통해 직부인대상확인 및 한도시부인"
    raw = yaml.safe_load((HERE / "guide.yaml").read_text(encoding="utf-8"))
    assert raw["lines"][11]["raw"].endswith("확인 ")       # 원문 끝 빈칸까지 그대로


def test_every_guide_line_judged():
    ar = run()
    assert [g["id"] for g in ar["guide"]] == [g["id"] for g in load_guide()]
    for g in ar["guide"]:
        assert g["status"] in engine.STATUSES, g
        assert g["facts"] and all(g["facts"]), g
        assert g["title"] and g["text"] and g["group"]
        if g["status"] == "이슈":
            assert g["issues"], g
    assert engine.validate(ar) == []


def test_guide_without_prior_marks_data_needed():
    ar = run(pr=None)
    st = {g["id"]: g["status"] for g in ar["guide"]}
    assert st["1-3"] == "자료 필요" and st["1-4"] == "자료 필요"
    assert all(s for s in st.values())


# ---------------------------------------------------------------- 이슈 순서

def test_issue_order_and_numbers():
    led = ledger(checks=[{"id": "balanced", "ok": False, "title": "차변합 = 대변합", "detail": "차 100 · 대 90",
                          "amounts": {"debit": 100, "credit": 90, "diff": 10}}])
    ar = run(led)
    issues = ar["issues"]
    assert [i["no"] for i in issues] == list(range(1, len(issues) + 1))
    buckets = [engine.Builder.bucket(i) for i in issues]
    assert buckets == sorted(buckets)
    assert issues[0]["kind"] == "신고 점검"
    assert issues[-1]["kind"] == "서식 필요"
    for b in set(buckets):
        sizes = [engine.Builder.size(i) for i in issues if engine.Builder.bucket(i) == b]
        assert sizes == sorted(sizes, reverse=True)
    for i in issues:
        assert i["kind"] in engine.KINDS and i["certainty"] in engine.CERTAINTIES
        assert set(i) >= {"no", "kind", "certainty", "guide", "accounts", "title", "amounts", "basis", "prior_treatment",
                          "papers", "forms", "questions", "memo", "conclusion"}
        assert i["memo"] is None and i["conclusion"] is None


# ---------------------------------------------------------------- 신고 점검

def test_failed_ledger_check_is_issue():
    led = ledger(checks=[{"id": "bs_rollforward", "ok": False, "title": "재무상태표", "detail": "외상매출금 차이",
                          "amounts": {"외상매출금": 1_000}}])
    ar = run(led)
    hit = by_title(ar, "전기 재무상태표에 분개장 증감")
    assert hit and hit[0]["kind"] == "신고 점검"


def test_closing_residual_and_offset():
    led = ledger(checks=[{"id": "closing_residual", "ok": False, "title": "손익(400) 계정 잔액 0", "detail": "",
                          "amounts": {"손익 400 잔액": 1_234}}])
    assert by_title(run(led), "손익(400) 계정에 잔액이 남았습니다")
    led = ledger(checks=[{"id": "closing_residual", "ok": True, "title": "", "detail": "",
                          "amounts": {"손익 400 잔액": 0, "813 접대비": 5_000, "855 영업비": -5_000}}])
    hit = by_title(run(led), "마감 뒤에도 잔액이 남은")
    assert hit and "813" in hit[0]["accounts"]
    # checks 에 closing_residual 이 없으면 400 계정 잔액으로 본다
    accounts = [a for a in ledger()["accounts"] if a["code"] != "400"] + [acc("400", "손익", "손익대체", 0, closing=777)]
    assert by_title(run(ledger(accounts=accounts, checks=[])), "손익(400) 계정에 잔액이 남았습니다")


def test_profit_without_corporate_tax():
    accounts = [a for a in ledger()["accounts"] if a["code"] != "998"]
    led = ledger(accounts=accounts, totals={**ledger()["totals"], "corporate_tax": 0})
    hit = by_title(run(led), "법인세등이 0원")
    assert hit and hit[0]["kind"] == "신고 점검" and hit[0]["certainty"] == "확실"
    # 전기 결손이면 확실도는 가능으로 내린다
    pr = prior(cells={"3": {"06": -5_000_000}}, losses=[{"year": "2024", "incurred": 5_000_000, "used": 0,
                                                        "remaining": 5_000_000}])
    assert by_title(run(led, pr=pr), "법인세등이 0원")[0]["certainty"] == "가능"


def test_prior_loss_row_missing():
    pr = prior(cells={"3": {"06": -7_000_000}}, losses=[])
    hit = by_title(run(pr=pr), "50호(갑) 이월결손금 줄에 없습니다")
    assert hit and hit[0]["kind"] == "신고 점검" and hit[0]["amounts"]["prior"] == -7_000_000
    # 전기 읽기가 3호 06 으로 채운 줄(derived)은 신고서에 있던 줄이 아니다
    pr = prior(cells={"3": {"06": -7_000_000}},
               losses=[{"year": "2024", "incurred": 7_000_000, "used": 0, "remaining": 7_000_000, "derived": "3호 06"}])
    assert by_title(run(pr=pr), "50호(갑) 이월결손금 줄에 없습니다")
    pr = prior(cells={"3": {"06": -7_000_000}}, losses=[{"year": "2024", "incurred": 7_000_000, "used": 0, "remaining": 7_000_000}])
    assert not by_title(run(pr=pr), "50호(갑) 이월결손금 줄에 없습니다")


def test_loan_without_prior_interest():
    hit = by_title(run(), "인정이자 조정과 19호가 없습니다")
    assert hit and hit[0]["kind"] == "신고 점검" and hit[0]["forms"] == ["19갑", "19을"]
    pr = prior(adjustments=[{"side": "가산", "subject": "가지급금인정이자", "amount": 2_000_000, "disposition": "상여",
                             "code": "100", "note": ""}])
    ar = run(pr=pr)
    assert not by_title(ar, "인정이자 조정과 19호가 없습니다")
    hit = by_title(ar, "인정이자 계산이 필요합니다")
    assert hit and hit[0]["kind"] == "조정 후보" and "가지급금인정이자" in hit[0]["prior_treatment"]


# ---------------------------------------------------------------- 조정 후보 · 전기 연속

def test_adjustment_candidates():
    ar = run()
    corp = by_title(ar, "법인세등을 손금불산입")[0]
    assert corp["certainty"] == "확실" and corp["papers"] == ["corporate_tax"] and corp["amounts"]["current"] == 4_000_000
    ent = by_title(ar, "기업업무추진비 한도")[0]
    assert ent["certainty"] == "확실" and ent["forms"] == ["23갑", "23을"] and ent["papers"] == ["entertainment"]
    assert "전기 15호 가산 접대비(기업업무추진비) 1,000,000원 기타사외유출" in ent["prior_treatment"]
    assert any("경조사" in b for b in ent["basis"])
    car = by_title(ar, "업무용승용차")[0]
    assert car["certainty"] == "확실" and car["papers"] == ["business_car"]      # 전기에 29호를 냈다
    pen = by_title(ar, "가산세 · 과태료로 보이는")[0]
    assert pen["certainty"] == "확실" and pen["papers"] == ["taxes_dues"]
    assert any(b.startswith("세금과공과") and "200,000" in b for b in pen["basis"])    # 주정차 과태료
    assert any(b.startswith("잡손실 적요") and "450,000" in b for b in pen["basis"])   # 부가세 가산세(단수차이는 빠짐)
    ret = by_title(ar, "퇴직급여충당금 한도")[0]
    assert ret["papers"] == ["retirement"] and ret["forms"] == ["32"]
    don = by_title(ar, "기부금 한도")[0]
    assert don["papers"] == ["donation"] and don["forms"] == ["21", "22"]


def test_car_without_prior_is_possible():
    pr = prior(forms=[f for f in prior()["forms"] if f["code"] != "29"])
    assert by_title(run(pr=pr), "업무용승용차")[0]["certainty"] == "가능"


def test_prior_continuity():
    ar = run()
    # 미수이자: 당기 미수수익 계정이 없어 소멸, 50호(을) 유보는 추인 대상
    gone = by_title(ar, "전기 미수이자 조정 대상이 당기에는 없습니다")
    assert gone and gone[0]["kind"] == "전기 연속"
    res = by_title(ar, "전기 유보 미수수익")
    assert res and res[0]["kind"] == "전기 연속" and "추인" in res[0]["title"]
    g13 = next(g for g in ar["guide"] if g["id"] == "1-3")
    assert g13["status"] == "이슈" and len(g13["issues"]) >= 3
    # 결손이라 법인세등이 없으면 전기 법인세 조정은 소멸
    accounts = [a for a in ledger()["accounts"] if a["code"] != "998"]
    led = ledger(accounts=accounts, totals={**ledger()["totals"], "income_before_tax": -1, "corporate_tax": 0})
    assert by_title(run(led), "전기 법인세등 조정 대상이 당기에는 없습니다")


def test_credits_and_forms_list():
    ar = run()
    assert by_title(ar, "중소기업 특별세액감면을 당기에도 적용할지")
    codes = [f["code"] for f in ar["forms"]]
    assert codes[:len(engine.BASIC_FORMS)] == engine.BASIC_FORMS
    by = {f["code"]: f for f in ar["forms"]}
    assert all(by[c]["status"] == "해당" for c in engine.BASIC_FORMS)
    assert by["29"]["prior_filed"] is True and by["29"]["status"] == "해당" and by["29"]["why"]
    assert by["19갑"]["status"] == "확인" and by["19갑"]["prior_filed"] is False
    assert by["23갑"]["by"] == "엔진"
    assert by["8갑"]["prior_filed"] is True
    for f in ar["forms"]:
        assert f["status"] in engine.FORM_STATUSES and f["by"] in ("엔진", "사람") and f["name"]


def test_cost_accounts_with_unknown_prior():
    """원가 계정(prior_known false)은 증감 · 상태에서 뺀다. 이름으로 모으는 묶음은 판관비와 원가를 함께 센다."""
    accounts = ledger()["accounts"] + [
        {**acc("613", "접대비(기업업무추진비)(도)", "비용", 6_000_000, 0, lines=5), "prior_known": False, "transfer": 6_000_000},
        {**acc("603", "급여(도)", "비용", 500_000_000, 0), "prior_known": False, "transfer": 500_000_000},
    ]
    ar = run(ledger(accounts=accounts))
    ent = by_title(ar, "기업업무추진비 한도")[0]
    assert ent["accounts"] == ["813", "613"]
    # 전기 · 증감은 원가 계정 전기 금액을 몰라 비우고(당기 − 전기 = 증감이 늘 맞게), 아는 계정끼리 비교는 근거에 둔다(5회차)
    assert ent["amounts"] == {"prior": None, "current": 36_000_000, "change": None}
    assert any("전기 금액을 아는 계정만 보면 전기 25,000,000원, 당기 30,000,000원, 증감 5,000,000원" in b for b in ent["basis"])
    pay = by_title(ar, "통합고용세액공제")[0]                                           # 증가 판정은 802 만 센다
    assert pay["amounts"]["prior"] is None and any("증감 20,000,000원" in b for b in pay["basis"])
    for i in ar["issues"]:
        a = i["amounts"]
        if a.get("prior") is not None and a.get("current") is not None and a.get("change") is not None:
            assert a["current"] - a["prior"] == a["change"], i["title"]


def test_prior_only_income_statement_lines():
    """당기 0 이고 코드 없는 손익계산서 줄(전기에만 있는 줄)과 totals.prior 를 쓴다."""
    accounts = [a for a in ledger()["accounts"] if a["code"] != "998"]
    fs = {"bs": [], "is": [{"label": "법인세등", "cur": 0, "prior": 3_000_000, "code": None, "kind": "account", "section": "법인세등"},
                           {"label": "매출", "cur": 0, "prior": 2_500_000_000, "code": None, "kind": "account", "section": "매출액"}]}
    tot = {**ledger()["totals"], "income_before_tax": -5_000_000, "corporate_tax": 0,
           "prior": {"revenue": 3_300_000_000, "corporate_tax": 3_000_000}}
    ar = run(ledger(accounts=accounts, fs=fs, totals=tot))
    assert by_title(ar, "결손금 소급공제")                     # 전기 법인세등은 totals.prior 와 코드 없는 줄에서
    assert by_title(ar, "지출증명서류 합계표")                 # 전기 매출은 totals.prior.revenue
    for i in ar["issues"]:
        assert "" not in i["accounts"]


# ---------------------------------------------------------------- 명령

def write_case(case_dir, with_prior=True):
    idx = case_dir / "_색인"
    idx.mkdir(parents=True, exist_ok=True)
    (idx / "장부.json").write_text(json.dumps(ledger(), ensure_ascii=False), encoding="utf-8")
    if with_prior:
        (idx / "전기.json").write_text(json.dumps(prior(), ensure_ascii=False), encoding="utf-8")
    cols = ["date", "no", "kind", "account", "account_name", "debit", "credit", "partner_code", "partner", "memo_code",
            "memo", "closing"]
    with open(idx / "전표.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in journal():
            w.writerow({**r, "closing": 1 if r["closing"] else 0})


def test_cli_end_to_end(case_dir):
    write_case(case_dir)
    assert cli.main([str(case_dir)]) == 0
    ar = json.loads((case_dir / "_색인" / "AR.json").read_text(encoding="utf-8"))
    assert set(ar) == {"guide", "issues", "forms"}
    assert engine.validate(ar) == []
    wb = openpyxl.load_workbook(case_dir / "1 AR" / "AR_이슈.xlsx")
    assert wb.sheetnames == ["이슈 목록", "서식 목록", "가이드 대조"]
    ws = wb["이슈 목록"]
    head = next(r for r in range(1, 6) if ws.cell(r, 1).value == "번호")
    col = {ws.cell(head, c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(head + 1, col["회계사 결론"]).fill.fgColor.rgb.endswith("FFF2CC")         # 노란 입력 칸
    folded = sum(1 for i in ar["issues"] if i["folded"])
    assert ws.max_row - head == len(ar["issues"]) + (1 if folded else 0)                     # 접음 머리 줄
    g = wb["가이드 대조"]
    ghead = next(r for r in range(1, 6) if g.cell(r, 1).value == "번호")
    assert g.max_row - ghead == len(load_guide())
    assert all(g.cell(r, 5).value for r in range(ghead + 1, g.max_row + 1))
    step = json.loads((case_dir / "_색인" / "신고건.json").read_text(encoding="utf-8"))["steps"]["ar"]
    assert step["ok"] is True and "_색인/AR.json" in step["outputs"]


def test_cli_missing_prior(case_dir, capsys):
    write_case(case_dir, with_prior=False)
    assert cli.main([str(case_dir)]) == 2
    assert "ctax prior" in capsys.readouterr().out


def test_cli_missing_ledger(case_dir):
    assert cli.main([str(case_dir)]) == 2


def test_text_rules():
    ar = run()
    text = json.dumps(ar, ensure_ascii=False)
    assert "—" not in text
    for i in ar["issues"]:
        assert not i["title"].endswith("?")
        for q in i["questions"]:
            assert q.endswith("주십시오"), q


@pytest.mark.parametrize("name,expect", [("접대비(기업업무추진비)(판)", "접대비(기업업무추진비)"), ("세금과공과금(판)", "세금과공과금"),
                                         ("복리후생비(도)", "복리후생비")])
def test_clean_name(name, expect):
    from ctax.ar.topics import clean
    assert clean(name) == expect


def _first_year(prior_pay=0):
    accounts = [acc("103", "보통예금", "자산", 50_000_000), acc("212", "기계장치", "자산", 30_000_000),
                acc("262", "미지급비용", "부채", 1_000_000), acc("331", "자본금", "자본", 50_000_000, lines=1),
                acc("401", "상품매출", "수익", 200_000_000), acc("802", "직원급여(판)", "비용", 120_000_000, prior_pay)]
    return engine.build(ledger(accounts=accounts, partners=[]), [],
                        {"no_prior": True, "company": {"sme": True}, "forms": [], "adjustments": [], "reserves": [],
                         "losses": [], "credits": [], "problems": []}, {"company": "가상회사", "fy": "2025"})


def test_first_year_issues_do_not_compare_with_prior():
    ar = _first_year()
    titles = {i["title"] for i in ar["issues"]}
    head = "첫 사업연도라 전기와 견줄 수 없습니다. "
    assert head + "통합고용세액공제 대상인지 상시근로자 수로 확인합니다" in titles
    assert head + "통합투자세액공제 대상인지 확인합니다" in titles
    assert head + "주식등변동상황명세서(54호)에 설립 때 발행한 주식과 주주를 적습니다" in titles
    assert not any(w in t for t in titles for w in ("늘었습니다", "바뀌었습니다"))
    texts = [b for i in ar["issues"] for b in i["basis"]] + [f for g in ar["guide"] for f in g["facts"]]
    assert not any("(전기 0원)" in t or "증감 " in t and "유형자산" not in t or "신규" in t for t in texts)
    pay = by_title(ar, "통합고용")[0]
    assert pay["amounts"] == {"prior": 0, "current": 120_000_000, "change": 120_000_000}     # 금액은 그대로
    assert pay["questions"] == ["올해 월별 상시근로자 수(4대보험 가입자 명부)를 보내 주십시오"]


def test_first_year_wording_needs_all_prior_zero():
    ar = _first_year(prior_pay=100_000_000)                             # 전기 신고서는 없지만 전기 금액이 있다(첫 수임)
    titles = {i["title"] for i in ar["issues"]}
    assert "급여가 늘었습니다. 통합고용세액공제 대상인지 상시근로자 수로 확인합니다" in titles
    assert "자본금이 바뀌었습니다. 주식등변동상황명세서(54호)에 증자 · 감자를 반영합니다" in titles
    assert any("(전기 0원)" in b for i in ar["issues"] for b in i["basis"])
    assert not by_title(run(), "첫 사업연도")                            # 전기 신고서가 있는 회사는 그대로

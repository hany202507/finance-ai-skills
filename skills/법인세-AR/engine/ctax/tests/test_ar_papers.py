"""AR 이슈 → 조서 잇기(3회차 R3-5). 이슈의 papers 는 조서 등록부(ctax.wp.papers ORDER · AGGREGATES)에 있는 id 만 쓴다.

고객 자료 없이 test_ar 의 합성 장부에 계정을 더해 돌린다.
"""
from __future__ import annotations

import json

import openpyxl

from ctax.ar import cli, engine, excel
from ctax.tests.test_ar import acc, by_title, ledger, prior, run, write_case
from ctax.wp import papers as registry

KNOWN = {pid for pid, _ in registry.ORDER} | set(registry.AGGREGATES)


def _more_accounts():
    return ledger()["accounts"] + [
        acc("116", "미수수익", "자산", 1_500_000, 1_000_000),
        acc("133", "선급비용", "자산", 2_400_000, 1_200_000),
        acc("109", "대손충당금", "자산", -5_000_000, -4_000_000, debit=0, credit=1_000_000),
        acc("107", "단기매매증권", "자산", 30_000_000, 20_000_000),
        acc("905", "단기매매증권평가이익", "수익", 4_000_000, 0),
        acc("903", "배당금수익", "수익", 6_000_000, 5_000_000),
        acc("910", "외화환산이익", "수익", 12_000_000, 0),
        acc("262", "미지급비용", "부채", 8_000_000, 6_000_000),
        acc("925", "전기오류수정이익", "수익", 5_000_000, 0),
        acc("812", "급여(판)", "비용", 300_000_000, 200_000_000),
    ]


def test_every_paper_id_is_registered():
    ar = run(ledger(accounts=_more_accounts()))
    used = {p for i in ar["issues"] for p in i["papers"]}
    assert used <= KNOWN
    # 3회차 새 조서 열 개 중 이 장부로 서는 것
    want = {"미수이자 익금불산입": "accrued_income", "미확정 비용": "accrued_expense", "선급비용이 있습니다": "prepaid",
            "대손충당금 한도": "bad_debt", "유가증권 평가손익": "securities", "수입배당금": "dividend",
            "외화환산손익": "fx", "전기오류수정손익": "prior_error", "인정이자": "loan_interest",
            "중소기업 특별세액감면을 당기에도": "credits", "통합고용세액공제": "credits"}
    for word, pid in want.items():
        hit = by_title(ar, word)
        assert hit, word
        assert pid in hit[0]["papers"], (word, hit[0]["papers"])
    assert engine.validate(ar) == []


def test_unregistered_paper_id_is_dropped(monkeypatch):
    monkeypatch.setitem(engine.KEY_PAPERS, "fx", ["fx", "no_such_paper"])
    ar = run(ledger(accounts=_more_accounts()))
    assert by_title(ar, "외화환산손익")[0]["papers"] == ["fx"]


def test_paper_ids_come_from_registry():
    assert engine.paper_ids() == frozenset(KNOWN)
    assert engine.papers_for("credit:중소기업에 대한 특별세액감면") == ["credits"]
    assert engine.papers_for("loan_interest") == ["interest_denial"]   # 업무무관 지급이자(26호)는 인정이자 조서가 아니라 17 조서(10-21)
    assert engine.papers_for("prior:x", ["reserve_50", "accrued_income"]) == ["reserve_50", "accrued_income"]


def test_prior_reserve_links_paper():
    ar = run()
    # 전기 50호(을) 미수수익 유보가 계정이 없어 전액 추인 대상: 유보 집계와 미수수익 조서
    gone = by_title(ar, "전기 유보 미수수익")
    assert gone and gone[0]["papers"] == ["reserve_50", "accrued_income"]
    # 전기 15호 미수이자 유보 조정이 당기에 소멸: 유보라서 조서를 잇는다
    gone = by_title(ar, "전기 미수이자 조정 대상이 당기에는 없습니다")
    assert gone and gone[0]["papers"] == ["accrued_income"]


def test_excel_paper_names_from_registry(case_dir):
    names = excel.paper_names()
    for pid, no in registry.ORDER:
        assert names[pid].startswith(f"{no} ")
    assert names["loss"] == "결손금"
    write_case(case_dir)
    assert cli.main([str(case_dir)]) == 0
    ar = json.loads((case_dir / "_색인" / "AR.json").read_text(encoding="utf-8"))
    ws = openpyxl.load_workbook(case_dir / "1 AR" / "AR_이슈.xlsx")["이슈 목록"]
    head = next(r for r in range(1, 6) if ws.cell(r, 1).value == "번호")
    col = {ws.cell(head, c).value: c for c in range(1, ws.max_column + 1)}
    loan = next(i for i in ar["issues"] if i["rule"] == "loan")
    row = next(r for r in range(head + 1, ws.max_row + 1) if ws.cell(r, 1).value == loan["no"])
    assert ws.cell(row, col["조서"]).value == names["loan_interest"]


# ---------------------------------------------------------------- 결손금 이력 다음 할 일

def test_loss_issue_has_todo_for_manual_rows(case_dir):
    pr = prior(cells={"3": {"06": -7_000_000}}, losses=[])
    hit = by_title(run(pr=pr), "50호(갑) 이월결손금 줄에 없습니다")[0]
    assert hit["papers"] == ["loss"]
    assert any("결손 결손금" in t and "노란 줄" in t and "ctax wp" in t for t in hit["todo"])
    assert any("세무조정계산서" in q for q in hit["questions"])
    # 신호만 있는 loss_history 도 같은 안내
    pr = prior(cells={"3": {"07": 5_000_000}}, losses=[])
    hist = [i for i in run(pr=pr)["issues"] if i["rule"] == "loss_history"]
    assert hist and hist[0]["todo"] == [engine.LOSS_TODO]
    # 다른 이슈는 빈 목록
    assert all(i["todo"] == [] for i in run()["issues"])
    # 엑셀 「다음 할 일」 열
    write_case(case_dir)
    (case_dir / "_색인" / "전기.json").write_text(json.dumps(prior(cells={"3": {"06": -7_000_000}}, losses=[]),
                                                          ensure_ascii=False), encoding="utf-8")
    assert cli.main([str(case_dir)]) == 0
    ws = openpyxl.load_workbook(case_dir / "1 AR" / "AR_이슈.xlsx")["이슈 목록"]
    head = next(r for r in range(1, 6) if ws.cell(r, 1).value == "번호")
    col = {ws.cell(head, c).value: c for c in range(1, ws.max_column + 1)}
    todo = [ws.cell(r, col["다음 할 일"]).value for r in range(head + 1, ws.max_row + 1)]
    assert engine.LOSS_TODO in todo

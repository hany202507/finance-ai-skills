"""회계사 입력 되읽기(계약 14.2). `ctax wp` 가 조서.xlsx 를 다시 쓰기 전에 노란 칸 값을 읽는다.

1. 숨은 시트 `_입력` 의 key ↔ 칸 표로 칸마다 지금 값을 읽는다.
2. 엔진 기본값(그 표의 default)과 다른 칸만 회계사 값으로 본다. 같으면(손대지 않았거나 되돌렸으면) 회계사 값을 지운다.
3. `_색인/조서입력.json` 에 모은다: {"cells": [{"paper", "key", "value", "at", "from": "조서.xlsx"}]}.
   지난 실행에 모은 값 중 이번 엑셀 `_입력` 표에 없는 key(줄이 장부에서 사라진 경우 등)는 그대로 둔다.

칸을 가리는 것은 칸 주소가 아니라 key 다. 시트 모양이 바뀌어도 key 가 같으면 값이 따라간다.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from ..case import Case
from .model import normalize
from .sheet import INPUT_SHEET

INPUT_JSON = "조서입력.json"


def read_workbook(path: Path) -> tuple[list[dict], str]:
    """조서.xlsx 의 노란 칸 [{paper, key, value, default, kind}] 과 문제(읽지 못했으면 까닭)."""
    if not path.exists():
        return [], ""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path)
    except Exception as e:  # 열려 있거나 깨진 파일
        return [], f"{path.name} 을 읽지 못했습니다({type(e).__name__})"
    if INPUT_SHEET not in wb.sheetnames:
        return [], f"{path.name} 에 {INPUT_SHEET} 시트가 없습니다(1회차 조서). 노란 칸을 읽지 않고 새로 씁니다"
    out = []
    for row in wb[INPUT_SHEET].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        paper, key, sheet, cell, default, kind = (list(row) + [None] * 6)[:6]
        kind = kind or "text"
        if sheet not in wb.sheetnames:
            continue
        value = wb[sheet][cell].value
        if isinstance(value, str) and value.startswith("="):
            continue   # 회계사가 수식을 넣은 칸은 값을 알 수 없어 읽지 않는다
        out.append({"paper": str(paper), "key": str(key), "value": normalize(value, kind),
                    "default": normalize(default, kind), "kind": kind})
    return out, ""


def collect(case: Case, xlsx: Path) -> tuple[dict, list[str]]:
    """조서.xlsx 와 지난 조서입력.json 을 합쳐 조서입력.json 을 쓰고 {(paper, key): value} 를 돌려준다."""
    old = case.read_json(INPUT_JSON) if case.has(INPUT_JSON) else {"cells": []}
    cells = {(c["paper"], c["key"]): c for c in old.get("cells", []) if c.get("paper") and c.get("key")}
    rows, problem = read_workbook(xlsx)
    problems = [problem] if problem else []
    now = datetime.now().isoformat(timespec="seconds")
    for r in rows:
        k = (r["paper"], r["key"])
        if r["value"] == r["default"]:
            cells.pop(k, None)
            continue
        prev = cells.get(k)
        if prev and prev.get("value") == r["value"]:
            continue
        cells[k] = {"paper": r["paper"], "key": r["key"], "value": r["value"], "at": now, "from": "조서.xlsx"}
    data = {"cells": sorted(cells.values(), key=lambda c: (c["paper"], c["key"]))}
    if data["cells"] or case.has(INPUT_JSON):
        case.write_json(INPUT_JSON, data)
    return {(c["paper"], c["key"]): c["value"] for c in data["cells"]}, problems

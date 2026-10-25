"""조서 엑셀 수식의 작은 계산기. 시험에서 「수식 결과 = 조정.json 값」 을 확인할 때 쓴다.

조서 엑셀(`ctax.wp.excel`)이 거는 수식만 계산한다: 사칙연산, 비교, 셀 · 범위 참조(다른 시트 포함),
SUM · SUMIF · SUMIFS · COUNTIF · MIN · MAX · ROUNDDOWN · ROUND · IF · AND · OR · ABS · VLOOKUP(FALSE). 엑셀처럼 15자리에서 반올림한 뒤 자른다.
"""
from __future__ import annotations

import math
import re

from openpyxl.utils import column_index_from_string, get_column_letter

_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:'(?:[^']|'')+'!|[A-Za-z0-9_가-힣.]+!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?%?)
  | (?P<func>[A-Z][A-Z0-9.]*)\(
  | (?P<bool>TRUE|FALSE)
  | (?P<op><>|<=|>=|[-+*/^&=<>(),])
""", re.X)


class FormulaError(Exception):
    pass


def _tokens(text: str):
    pos, out = 0, []
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m:
            raise FormulaError(f"읽지 못한 수식 조각: {text[pos:pos + 20]!r}")
        pos = m.end()
        kind = m.lastgroup
        if kind == "ws":
            continue
        out.append((kind, m.group(kind)))
    return out


class Range:
    def __init__(self, values: list, cols: int = 1):
        self.values = values
        self.cols = cols


def _num(v):
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        return 1 if v else 0
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, Range):
        raise FormulaError("범위를 숫자로 쓸 수 없습니다")
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0


def _excel_round(x: float) -> float:
    """엑셀은 15자리 유효숫자로 계산한다. 0.7 같은 소수의 이진 오차를 지운다."""
    if x == 0:
        return 0.0
    digits = 15 - int(math.floor(math.log10(abs(x)))) - 1
    return round(x, max(0, min(digits, 12)))


def _match(value, criterion) -> bool:
    if isinstance(criterion, str):
        m = re.match(r"^(<>|<=|>=|=|<|>)(.*)$", criterion)
        if m:
            op, rhs = m.groups()
            try:
                a, b = _num(value), float(rhs)
            except ValueError:
                a, b = str(value or ""), rhs
            return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
        return str(value if value is not None else "").strip().lower() == criterion.strip().lower()
    if value is None or value == "":
        return criterion in (None, "")
    try:
        return _num(value) == _num(criterion)
    except FormulaError:
        return False


class Workbook:
    """openpyxl 통합 문서(수식 문자열 그대로 읽은 것)를 계산한다."""

    def __init__(self, wb):
        self.wb = wb
        self.cache: dict[tuple[str, str], object] = {}
        self.stack: set[tuple[str, str]] = set()

    def value(self, sheet: str, addr: str):
        addr = addr.replace("$", "")
        key = (sheet, addr)
        if key in self.cache:
            return self.cache[key]
        if key in self.stack:
            raise FormulaError(f"순환 참조: {sheet}!{addr}")
        self.stack.add(key)
        raw = self.wb[sheet][addr].value
        if isinstance(raw, str) and raw.startswith("="):
            v = _Parser(self, sheet, _tokens(raw[1:])).parse()
            if isinstance(v, Range):
                raise FormulaError(f"{sheet}!{addr} 수식이 범위를 돌려줍니다")
        else:
            v = raw
        self.stack.discard(key)
        self.cache[key] = v
        return v

    def ref(self, sheet: str, text: str):
        if "!" in text:
            sname, text = text.rsplit("!", 1)
            sheet = sname[1:-1].replace("''", "'") if sname.startswith("'") else sname
        text = text.replace("$", "")
        if ":" in text:
            a, b = text.split(":")
            ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
            cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
            vals = []
            c0, c1 = column_index_from_string(ca), column_index_from_string(cb)
            for r in range(int(ra), int(rb) + 1):
                for c in range(c0, c1 + 1):
                    vals.append(self.value(sheet, f"{get_column_letter(c)}{r}"))
            return Range(vals, c1 - c0 + 1)
        return self.value(sheet, text)


class _Parser:
    PREC = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}

    def __init__(self, book: Workbook, sheet: str, toks):
        self.book, self.sheet, self.toks, self.i = book, sheet, toks, 0

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def take(self):
        t = self.peek()
        self.i += 1
        return t

    def parse(self):
        v = self.expr(0)
        if self.i != len(self.toks):
            raise FormulaError(f"수식 끝이 남았습니다: {self.toks[self.i:]}")
        return v

    def expr(self, min_prec):
        left = self.unary()
        while True:
            kind, op = self.peek()
            if kind != "op" or op not in self.PREC or self.PREC[op] < min_prec:
                return left
            self.take()
            right = self.expr(self.PREC[op] + (0 if op == "^" else 1))
            left = self.apply(op, left, right)

    def apply(self, op, a, b):
        if op == "&":
            return f"{'' if a is None else a}{'' if b is None else b}"
        if op in ("=", "<>", "<", ">", "<=", ">="):
            if isinstance(a, str) or isinstance(b, str):
                a, b = str(a or "").lower(), str(b or "").lower()
            else:
                a, b = _num(a), _num(b)
            return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
        a, b = _num(a), _num(b)
        if op == "+":
            return a + b
        if op == "-":
            return a - b
        if op == "*":
            return a * b
        if op == "/":
            if b == 0:
                raise FormulaError("0 으로 나눔")
            return a / b
        return a ** b

    def unary(self):
        kind, tok = self.peek()
        if kind == "op" and tok in ("-", "+"):
            self.take()
            v = self.expr(4)
            return -_num(v) if tok == "-" else _num(v)
        return self.atom()

    def atom(self):
        kind, tok = self.take()
        if kind == "num":
            return float(tok[:-1]) / 100 if tok.endswith("%") else (float(tok) if "." in tok or "e" in tok.lower() else int(tok))
        if kind == "str":
            return tok[1:-1].replace('""', '"')
        if kind == "bool":
            return tok == "TRUE"
        if kind == "ref":
            return self.book.ref(self.sheet, tok)
        if kind == "op" and tok == "(":
            v = self.expr(0)
            self.expect(")")
            return v
        if kind == "func":
            args = []
            if self.peek() != ("op", ")"):
                while True:
                    args.append(self.expr(0))
                    if self.peek() == ("op", ","):
                        self.take()
                        continue
                    break
            self.expect(")")
            return self.call(tok, args)
        raise FormulaError(f"뜻밖의 조각: {tok!r}")

    def expect(self, tok):
        k, t = self.take()
        if t != tok:
            raise FormulaError(f"{tok} 가 있어야 합니다(받은 것 {t!r})")

    @staticmethod
    def flat(args):
        out = []
        for a in args:
            if isinstance(a, Range):
                out.extend(v for v in a.values if isinstance(v, (int, float)) and not isinstance(v, bool))
            else:
                out.append(_num(a))
        return out

    def call(self, name, args):
        if name == "SUM":
            return sum(self.flat(args))
        if name == "MIN":
            vals = self.flat(args)
            return min(vals) if vals else 0
        if name == "MAX":
            vals = self.flat(args)
            return max(vals) if vals else 0
        if name == "ABS":
            return abs(_num(args[0]))
        if name in ("ROUNDDOWN", "ROUND"):
            x = _excel_round(_num(args[0]))
            d = int(_num(args[1])) if len(args) > 1 else 0
            f = 10 ** d
            if name == "ROUND":
                return math.floor(abs(x) * f + 0.5) / f * (1 if x >= 0 else -1)
            return math.trunc(x * f) / f
        if name == "IF":
            cond = args[0]
            if isinstance(cond, str):
                cond = cond != ""
            return args[1] if _num(cond) else (args[2] if len(args) > 2 else False)
        if name == "AND":
            return all(_num(a) for a in args)
        if name == "OR":
            return any(_num(a) for a in args)
        if name == "SUMIF":
            rng, crit = args[0], args[1]
            srng = args[2] if len(args) > 2 else rng
            return sum(_num(s) for v, s in zip(rng.values, srng.values) if _match(v, crit))
        if name == "SUMIFS":
            srng, pairs = args[0], list(zip(args[1::2], args[2::2]))
            total = 0
            for idx, s in enumerate(srng.values):
                if all(_match(r.values[idx], c) for r, c in pairs):
                    total += _num(s)
            return total
        if name == "COUNTIF":
            return sum(1 for v in args[0].values if _match(v, args[1]))
        if name == "VLOOKUP":
            key, table, col = args[0], args[1], int(_num(args[2]))
            exact = len(args) > 3 and not _num(args[3])
            if not exact:
                raise FormulaError("VLOOKUP 은 정확히 일치(FALSE)만 계산합니다")
            rows = [table.values[i:i + table.cols] for i in range(0, len(table.values), table.cols)]
            for r in rows:
                if _match(r[0], key if isinstance(key, str) else key):
                    return r[col - 1]
            return None   # 엑셀은 #N/A. 조서 수식은 IF 로 감싸 이 값을 쓰지 않는다
        raise FormulaError(f"모르는 함수: {name}")


def evaluate(wb, sheet: str, addr: str):
    """한 셀 값. 같은 통합 문서를 여러 번 물으면 `Workbook` 을 만들어 다시 쓴다."""
    return Workbook(wb).value(sheet, addr)


def split_ref(ref: str) -> tuple[str, str]:
    """`'6 기업업무추진비'!$G$40` → ('6 기업업무추진비', 'G40')."""
    sname, addr = ref.rsplit("!", 1)
    sheet = sname[1:-1].replace("''", "'") if sname.startswith("'") else sname
    return sheet, addr.replace("$", "")

"""조서 엑셀 공통 도우미. 조서 파일(`papers/<id>.py`)의 `write` 가 이 함수들로 시트 한 장을 쓴다.

조서 한 장: 머리(번호 · 이름 · 이슈 · 작성/검토 란) → 1 금액 검증 → 2 원장 검토(줄마다 노란 칸) →
3 세무조정 계산 · 소득처분 → 4 이어지는 서식 란.

노란 칸은 `Sheet.inp` 로 쓴다. 칸마다 key 를 붙여 숨은 시트 `_입력` 에 key ↔ 칸 주소 표를 남기고,
다음 `ctax wp` 가 그 표로 회계사 값을 읽는다(`inputs.py`).
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .. import xlsx as X
from .books import Books
from .model import MISSING, Paper

PAPER_WIDTHS = [12, 9, 8, 22, 18, 40, 16, 22, 12, 12, 40]
BOLD = Font(name=X.FONT, bold=True, size=11)
INPUT_SHEET = "_입력"


def q(sheet: str, addr: str) -> str:
    """다른 시트에서 부르는 절대 주소."""
    col = "".join(ch for ch in addr if ch.isalpha())
    row = "".join(ch for ch in addr if ch.isdigit())
    return f"'{sheet}'!${col}${row}"


def local(ref: str) -> str:
    """`'시트'!$G$10` → `$G$10` (같은 시트 안에서 부를 때)."""
    return ref.split("!")[1]


class Book:
    """조서 통합 문서, 셀 지도, 노란 칸 표."""

    def __init__(self, books: Books):
        self.b = books
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.cells: dict[str, dict[str, str]] = {}   # paper id → values 키 → 주소
        self.ref: dict[str, str] = {}                # 기준 시트 이름 → 주소
        self.input_cells: list[dict] = []            # {"paper", "key", "sheet", "cell", "default", "kind"}

    def reg(self, pid: str, key: str, address: str) -> None:
        self.cells.setdefault(pid, {})[key] = address

    def sheet(self, title: str, widths: list[int]) -> "Sheet":
        return Sheet(self, title, widths)


class Sheet:
    def __init__(self, bk: Book, title: str, widths: list[int]):
        self.bk = bk
        self.title = title
        self.ws = bk.wb.create_sheet(title)
        X.widths(self.ws, widths)
        self.ws.sheet_view.showGridLines = False

    def put(self, row: int, col: int, value=None, *, kind: str = "body", money: bool = False, wrap: bool = False,
            fmt: str | None = None):
        c = X.cell(self.ws, row, col, value, money=money, kind=kind, wrap=wrap)
        if fmt:
            c.number_format = fmt
        return c

    def inp(self, row: int, col: int, value, *, pid: str, key: str, default=MISSING, kind: str = "text",
            money: bool = False, fmt: str | None = None):
        """노란 칸. `default` 는 엔진 기본값(주지 않으면 value). 빈 값(None · "")도 기본값이 될 수 있다."""
        c = self.put(row, col, value, kind="input", money=money, fmt=fmt)
        self.bk.input_cells.append({"paper": pid, "key": key, "sheet": self.title,
                                    "cell": f"{get_column_letter(col)}{row}",
                                    "default": value if default is MISSING else default, "kind": kind,
                                    "value": value})
        return c

    def addr(self, row: int, col: int) -> str:
        return q(self.title, f"{get_column_letter(col)}{row}")

    def section(self, row: int, text: str) -> int:
        self.ws.cell(row=row, column=1, value=text).font = BOLD
        return row + 1


# 조서 한 장의 틀 ----------------------------------------------------------------------------------

def head(bk: Book, p: Paper) -> Sheet:
    s = bk.sheet(p.sheet, PAPER_WIDTHS)
    ws = s.ws
    ws.cell(row=1, column=1, value=f"{p.no} {p.name}").font = X.TITLE_FONT
    issues = ", ".join(f"#{i}" for i in p.issues) or "없음"
    forms = ", ".join(p.forms)
    X.note(ws, 2, f"{bk.b.company} · {bk.b.fy} 사업연도 · AR 이슈 {issues} · 이어지는 서식 {forms} · 상태 {p.status}")
    for i, label in enumerate(["작성", "작성일", "검토", "검토일"]):
        s.put(3, 1 + i * 2, label)
        s.put(3, 2 + i * 2, None, kind="input")
    return s


def notes(s: Sheet, row: int, p: Paper) -> int:
    if p.overridden:
        X.note(s.ws, row, f"회계사 값 {len(p.overridden)}칸을 반영했습니다(노란 칸, 조정.json inputs)")
        row += 1
    for n in p.confirm:
        s.put(row, 1, f"확인 필요: {n}").font = Font(name=X.FONT, color="C00000", size=9)
        row += 1
    for n in p.notes:
        X.note(s.ws, row, n)
        row += 1
    return row


def verify_rows(p: Paper) -> int:
    return 2 + max(1, len(p.accounts)) + 1


def lines_dry(p: Paper, start: int) -> tuple[int, int, int]:
    """`ledger` 를 쓰기 전에 원장 검토 표의 (다음 행, 첫 줄, 끝 줄)을 미리 안다(금액 검증이 그 범위를 부른다)."""
    first = start + 2
    last = first + max(1, len(p.lines)) - 1
    return last + 2, first, last


def verify(bk: Book, s: Sheet, row: int, p: Paper, first: int, last: int) -> int:
    """1. 금액 검증: 재무제표 금액과 원장(아래 표) 합계."""
    row = s.section(row, "1. 금액 검증")
    for col, label in ((3, "코드"), (4, "계정"), (7, "재무제표"), (8, "원장 합계"), (9, "차이")):
        X.header(s.ws, row, [label], col=col)
    row += 1
    for a in p.accounts:
        s.put(row, 3, a["code"])
        s.put(row, 4, a["name"])
        s.put(row, 7, int(a.get("amount") or 0), money=True)
        s.put(row, 8, f'=SUMIF($C${first}:$C${last},"{a["code"]}",$G${first}:$G${last})', money=True)
        s.put(row, 9, f"=G{row}-H{row}", money=True)
        row += 1
    if not p.accounts:
        s.put(row, 4, "관련 계정이 없습니다")
        row += 1
    return row + 1


def ledger(bk: Book, s: Sheet, row: int, p: Paper, pick: str, pick2: str = "", title: str = "2. 원장 검토",
           kind: str = "pick") -> tuple[int, int, int]:
    """2. 원장 검토 표. (다음 행, 첫 줄, 끝 줄)을 돌려준다. 줄이 없어도 빈 줄 하나를 두어 범위가 깨지지 않게 한다.

    `pick` · `pick2` 가 있으면 그 열이 노란 칸이고 key 는 `<pick>/<줄 열쇠>` 다.
    """
    row = s.section(row, title)
    X.header(s.ws, row, ["일자", "전표", "코드", "계정", "거래처", "적요", "금액", "엔진 분류", pick, pick2, "근거"])
    row += 1
    first = row
    for lr in p.lines:
        ln = lr.line
        s.put(row, 1, ln.date)
        s.put(row, 2, ln.no)
        s.put(row, 3, ln.account)
        s.put(row, 4, bk.b.by_code.get(ln.account, {}).get("name") or ln.account_name)
        s.put(row, 5, ln.partner)
        s.put(row, 6, ln.memo)
        s.put(row, 7, ln.amount, money=True)
        s.put(row, 8, lr.label)
        if pick:
            s.inp(row, 9, lr.pick, pid=p.id, key=f"{pick}/{lr.key}", default=lr.default, kind=kind)
        else:
            s.put(row, 9, lr.pick)
        if pick2:
            s.inp(row, 10, lr.pick2, pid=p.id, key=f"{pick2}/{lr.key}", default=lr.default2)
        else:
            s.put(row, 10, lr.pick2)
        s.put(row, 11, lr.why)
        row += 1
    if not p.lines:
        s.put(row, 6, "해당 줄 없음")
        row += 1
    last = row - 1
    s.ws.freeze_panes = None
    return row + 1, first, last


def calc_row(bk: Book, s: Sheet, row: int, p: Paper, key: str | None, label: str, value, *, kind: str = "body",
             note: str = "", fmt: str | None = None, input_kind: str | None = None) -> int:
    """계산 줄 하나(G 열). `kind="input"` 이면 노란 칸이고 key 로 `_입력` 에 올린다(기본값은 p.defaults[key])."""
    s.put(row, 1, label)
    money = fmt is None and not isinstance(value, str) or (isinstance(value, str) and value.startswith("="))
    if kind == "input" and key:
        s.inp(row, 7, value, pid=p.id, key=key, default=p.defaults.get(key, value),
              kind=input_kind or ("int" if isinstance(value, int) else "text"), money=money and fmt is None, fmt=fmt)
    else:
        s.put(row, 7, value, kind=kind, money=money and fmt is None, fmt=fmt)
    if note:
        X.note(s.ws, row, note, col=8)
    if key:
        bk.reg(p.id, key, s.addr(row, 7))
    return row + 1


def adj_table(bk: Book, s: Sheet, row: int, p: Paper, title: str = "3-2. 세무조정 · 소득처분") -> int:
    row = s.section(row, title)
    for col, label in ((1, "구분"), (4, "과목"), (5, "소득처분"), (6, "조정내용"), (7, "금액")):
        X.header(s.ws, row, [label], col=col)
    row += 1
    cells = bk.cells.get(p.id, {})
    for a in p.adjustments:
        s.put(row, 1, a.side)
        s.put(row, 4, a.subject)
        s.put(row, 5, f"{a.disposition}({a.code})")
        s.put(row, 6, a.note)
        src = cells.get(a.key)
        s.put(row, 7, f"={local(src)}" if src else a.amount, money=True)
        bk.reg(p.id, f"adj:{a.kind}", s.addr(row, 7))
        row += 1
    return row + 1


def forms_table(bk: Book, s: Sheet, row: int, p: Paper, items: list[tuple[str, str, str]]) -> int:
    row = s.section(row, "4. 이어지는 서식 란")
    for col, label in ((1, "서식"), (4, "란"), (7, "값")):
        X.header(s.ws, row, [label], col=col)
    row += 1
    cells = bk.cells.get(p.id, {})
    for form, cell_label, key in items:
        s.put(row, 1, form)
        s.put(row, 4, cell_label)
        src = cells.get(key)
        s.put(row, 7, f"={local(src)}" if src else None, money=True)
        row += 1
    return row + 1


class Rows:
    """계산 줄을 차례로 쓰며 key → 행을 기억한다(`g(key)` 는 같은 시트의 G 열 주소)."""

    def __init__(self, bk: Book, s: Sheet, p: Paper, row: int):
        self.bk, self.s, self.p, self.row = bk, s, p, row
        self.at: dict[str, int] = {}

    def put(self, key, label, value, **kw):
        self.at[key] = self.row
        self.row = calc_row(self.bk, self.s, self.row, self.p, key, label, value, **kw)

    def g(self, key) -> str:
        return f"G{self.at[key]}"

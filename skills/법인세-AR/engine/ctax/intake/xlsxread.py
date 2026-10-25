"""받은자료 xlsx 를 행 목록으로 읽는다. 숫자 · 이름 정리 도우미도 여기 둔다.

Drive 에서 받은 xlsx 는 styles.xml 의 `<xf ... count="n">` 때문에 openpyxl 이 TypeError 를 내는 일이 있어
그 속성만 지운 사본으로 다시 연다(corptax-impl `xlsx_io` 에서 옮김).
"""
from __future__ import annotations

import io
import re
import warnings
import zipfile
import zlib
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


class Unreadable(Exception):
    """파일이 깨졌거나 xlsx 가 아니다."""


CONTENT_ERRORS = (zipfile.BadZipFile, KeyError, OSError, zlib.error, EOFError, SyntaxError)


def _repair_bytes(path: Path) -> bytes:
    with zipfile.ZipFile(path) as src:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for it in src.infolist():
                data = src.read(it.filename)
                if it.filename == "xl/styles.xml":
                    data = re.sub(rb'(<xf\b[^>]*?)\s+count="\d+"', rb"\1", data)
                out.writestr(it.filename, data)
    return buf.getvalue()


def _open(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except TypeError:
        pass
    except CONTENT_ERRORS as e:
        raise Unreadable(f"{path.name}: {e}") from e
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            return load_workbook(io.BytesIO(_repair_bytes(path)), read_only=True, data_only=True)
    except (*CONTENT_ERRORS, TypeError) as e:
        raise Unreadable(f"{path.name}: {e}") from e


def xls_rows(path) -> list[list]:
    """BIFF xls(더존 ERP 계정명세서 · 합계잔액시산표) 첫 시트의 행. 빈 칸은 None, 정수인 실수는 정수로."""
    path = Path(path)
    try:
        import xlrd
    except ImportError as e:        # pragma: no cover - requirements.txt 에 있다
        raise Unreadable(f"{path.name}: xls 를 읽으려면 xlrd 가 필요합니다") from e
    try:
        sh = xlrd.open_workbook(str(path)).sheet_by_index(0)
    except Exception as e:          # xlrd 는 깨진 파일마다 다른 예외를 낸다
        raise Unreadable(f"{path.name}: {e}") from e
    out = []
    for r in range(sh.nrows):
        row = []
        for c in range(sh.ncols):
            v = sh.cell_value(r, c)
            if v == "":
                v = None
            elif isinstance(v, float) and v.is_integer():
                v = int(v)
            row.append(v)
        out.append(row)
    return out


def rows(path, limit: int | None = None) -> list[list]:
    """첫 시트의 행을 목록으로. limit 이 있으면 그 행 수까지만. `.xls` 는 xlrd 로 읽는다."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        out = xls_rows(path)
        return out[:limit] if limit is not None else out
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = _open(path)
        try:
            ws = wb.worksheets[0]
            out = []
            for r in ws.iter_rows(values_only=True):
                out.append(list(r))
                if limit is not None and len(out) >= limit:
                    break
        except CONTENT_ERRORS as e:
            raise Unreadable(f"{path.name}: {e}") from e
        finally:
            wb.close()
    return out


# ---------------------------------------------------------------- 값 정리

def num(v) -> int | None:
    """셀 값을 정수(원)로. 쉼표 문자열, `△` · `▲` 음수, 괄호 음수, 정수 · 실수를 안다. 빈 값은 None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(round(v))
    s = str(v).strip().replace(",", "").replace(" ", "")
    if s in ("", "-"):
        return None
    neg = False
    if s[0] in "△▲":
        neg, s = True, s[1:]
    elif s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    try:
        x = int(s)
    except ValueError:
        try:
            x = int(round(float(s)))
        except ValueError:
            return None
    return -x if neg else x


def n0(v) -> int:
    x = num(v)
    return 0 if x is None else x


def compact(s) -> str:
    """공백(전각 포함)을 모두 지운 글."""
    return re.sub(r"[\s　]+", "", str(s or ""))


SUFFIX = re.compile(r"\((판|제|도|분|보|운|공|임|건|용|기|관)\)$")
ROMAN = re.compile(r"^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩⅪⅫ]+\.")
PAREN = re.compile(r"^\(\d+\)")


def norm(s) -> str:
    """계정 이름 비교용. 공백과 `(판)` · `(제)` · `(도)` 같은 끝 꼬리를 지운다."""
    return SUFFIX.sub("", compact(s))


def strip_no(s) -> str:
    """`Ⅳ. 판매비와 관리비` → `판매비와관리비`, `(1) 당좌자산` → `당좌자산`."""
    c = compact(s)
    c = ROMAN.sub("", c, count=1)
    c = PAREN.sub("", c, count=1)
    return c.lstrip(".")


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def date_text(v) -> str | None:
    """전표일자를 `YYYY-MM-DD` 로. 날짜가 아니면 None(합계 줄 등)."""
    if isinstance(v, (datetime, date)):
        return text(v)
    s = text(v)
    m = re.fullmatch(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return None

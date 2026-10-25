# -*- coding: utf-8 -*-
"""엑셀을 다시 쓰기 전에 「AI 추가 이슈」 시트의 회계사 결론을 옮겨 둔다(engine · memo · finish 가 부른다)."""
from __future__ import annotations

from datetime import datetime

from openpyxl import load_workbook

from bestar import paths

ADDED_SHEET = "AI 추가 이슈"
ADDED_CONCL = "추가결론.json"


def save_added_conclusions(c) -> int:
    """엑셀 「AI 추가 이슈」 시트의 결론 칸을 _색인/AR_단계/추가결론.json 에 옮긴다(엑셀을 다시 쓰기 전에 부른다)."""
    path = c.root / paths.XLSX
    if not path.exists():
        return 0
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:                                    # noqa: BLE001 (열 수 없으면 앞 기록을 그대로 둔다)
        return 0
    if ADDED_SHEET not in wb.sheetnames:
        wb.close()
        return 0
    rows = list(wb[ADDED_SHEET].iter_rows(values_only=True))
    wb.close()
    if not rows:
        return 0
    head = [str(x or "") for x in rows[0]]
    if "번호" not in head or "회계사 결론" not in head:
        return 0
    ki, ci = head.index("번호"), head.index("회계사 결론")
    saved = paths.read_json(paths.steps_dir(c) / ADDED_CONCL, {}) or {}
    n = 0
    for r in rows[1:]:
        key, text = r[ki], r[ci]
        if not key:
            continue
        text = str(text).strip() if text is not None else ""
        old = saved.get(str(key)) or {}
        if text and old.get("text") != text:
            saved[str(key)] = {"text": text, "by": "회계사", "at": datetime.now().isoformat(timespec="seconds")}
            n += 1
        elif not text and old:
            saved.pop(str(key))
    paths.write_json(paths.steps_dir(c) / ADDED_CONCL, saved)
    return n

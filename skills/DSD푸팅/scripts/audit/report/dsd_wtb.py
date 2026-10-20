"""DSD 푸팅: 정산표(8100B)와 회사 제시 dsd 본표의 대사(2026-10-13).

dsd 안에서 숫자가 서로 맞아도 감사한 숫자와 다르면 소용이 없다. 정산표의 「공시 재무제표」 시트(회사 재무제표 양식)와
「재무제표」 시트(표시과목)에서 과목 · 당기(최종) · 전기를 읽어 dsd 의 재무상태표 · 손익계산서 과목과 이름으로 짝짓는다.
정산표는 엑셀에서 계산된 값이 저장돼 있어야 한다(엔진이 방금 만든 파일은 엑셀로 한 번 열어 저장하거나 tools/excel_recalc.py 를 돌린다).

감사 법인 양식(2026-10-18, 실제 회사 자료 전체 테스트 B-1): 감사인이 쓰는 WTB 는 시트 BS · IS(공시 과목의 수정 후 당기 · 전기),
WBS · WIS(계정별 최초 회사제시 · 수정 차변 · 대변 · 수정후), AJE(감사 수정분개 목록)로 되어 있다. 과목의 수정 후 금액은 BS · IS 에서 읽고,
그 과목에 드는 계정의 수정(수정후 - 최초 회사제시)을 WBS · WIS 에서 모아 과목의 회사제시 금액을 거꾸로 구한다. 계정을 과목에 잇는 순서는
이름, 금액, 이웃 계정 금액의 합이다. AJE 의 수정분개 번호는 그 계정이 든 과목에 붙는다. 어느 양식으로도 읽지 못하면 WtbFormatError 한 줄로 끝낸다.
"""
from __future__ import annotations

import re
from pathlib import Path

from .base import norm
from .dsd_grid import Doc, Part, Table
from .dsd_links import Check, _isnum, _labels, _ref, period_col

WTB_ID = "조서:"                              # 정산표 표 id 앞머리(검증의 terms 와 엑셀 주소에서 dsd 표와 가른다)
SHEETS = ("공시 재무제표", "재무제표", "공시용BS", "공시용PL")   # 8100B 의 회사 양식 과목 · 표시과목, WTB 양식(wtbform)의 공시용 시트
# WTB 양식(2026-10-17): 8100A 수정분개는 8100B 가 아니라 WTB 양식 · 8800 에 들어간다. 감사한 최종 숫자와 견주려면 WTB 양식을 준다.
# 공시용BS · 공시용PL 은 과목이 B 열이고 머리 이름이 「회사 제시 · Final · 전기」 다
STATEMENTS = ("재무상태표", "손익계산서")
FIRM_FS = ("BS", "IS", "PL", "재무상태표", "손익계산서")      # 감사 법인 양식의 공시 과목 시트
FIRM_ACC = {"BS": "WBS", "IS": "WIS", "PL": "WPL", "재무상태표": "WBS", "손익계산서": "WIS"}   # 같은 재무제표의 계정별 시트
AJE_SHEET = re.compile(r"^\s*(AJE|RJE|수정분개)", re.I)
FIRM_HEAD = ["과목", "당기(최종)", "전기", "당기(회사제시)", "수정", "감사 수정분개"]
ADJ = "수정분개"                               # 수정분개 목록 표의 kind. 과목 대사에는 쓰지 않는다
MIN_AMOUNT = 1000                              # 금액만으로 계정을 과목에 이을 때 이보다 작은 금액은 쓰지 않는다


class WtbFormatError(ValueError):
    """정산표 엑셀을 어느 양식으로도 읽지 못했다. 명령줄은 역추적 없이 이 글 한 줄로 끝낸다."""


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _fmt(v) -> str:
    return "" if v is None else f"{v:,.0f}"


def read_wtb(path: Path) -> Doc:
    """정산표 엑셀에서 과목 · 당기 · 전기를 읽는다. 시트마다 표 하나(머리줄: 과목, 당기, 전기).
    8100B · WTB 양식 시트가 없으면 감사 법인 양식(BS · IS · WBS · WIS · AJE)으로 읽는다."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:                       # 엑셀이 아니거나 깨진 파일
        raise WtbFormatError(f"엑셀로 열지 못함({type(e).__name__}): {Path(path).name}") from None
    parts = _read_8100b(wb, path)
    stats = {"양식": "8100B · WTB 양식"} if parts else {}
    if not parts:
        parts, stats = _read_firm(wb)
    if not parts:
        raise WtbFormatError(_why_not(wb, path))
    return Doc(Path(path).name, "", parts, stats)


def _why_not(wb, path) -> str:
    """읽지 못한 까닭 한 줄: 시트 이름과 첫 시트의 머리줄."""
    names = wb.sheetnames
    ws = wb[names[0]]
    head = ""
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [str(ws.cell(r, j).value).strip() for j in range(1, min(ws.max_column, 10) + 1) if ws.cell(r, j).value not in (None, "")]
        if vals:
            head = " | ".join(vals)[:80]
            break
    more = f" 외 {len(names) - 8}개" if len(names) > 8 else ""
    return (f"시트 {' · '.join(names[:8])}{more} · 머리줄 「{head}」. 읽는 양식은 8100B(「공시 재무제표」 · 「재무제표」), WTB 양식(「공시용BS」 · 「공시용PL」), "
            f"감사 법인 양식(BS · IS 와 WBS · WIS, AJE)이다: {Path(path).name}")


def _read_8100b(wb, path) -> list[Part]:
    parts = []
    for name in SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        found = next(((r, j) for r in range(1, min(ws.max_row, 30) + 1) for j in range(1, min(ws.max_column, 4) + 1)
                      if norm(ws.cell(r, j).value) == "과목"), None)
        if found is None:
            continue
        head, lc = found
        cols = {j: norm(ws.cell(head, j).value) for j in range(lc + 1, ws.max_column + 1)}
        cur = (next((j for j, h in cols.items() if h.startswith("당기") and "최종" in h), None)
               or next((j for j, h in cols.items() if h.lower() == "final" or h == "최종"), None)
               or next((j for j, h in cols.items() if h.startswith("당기")), None))
        pri = next((j for j, h in cols.items() if h.startswith("전기")), None)
        sub = next((j for j, h in cols.items() if "회사제시" in h and not h.startswith("증감")), None)      # 수정분개 반영 전 금액
        if cur is None:
            continue
        rows, nums = [], []
        for r in range(head + 1, ws.max_row + 1):
            label = str(ws.cell(r, lc).value or "").strip()
            if not label:
                continue
            a, b, c = _num(ws.cell(r, cur).value), (_num(ws.cell(r, pri).value) if pri else None), (_num(ws.cell(r, sub).value) if sub else None)
            rows.append([label, _fmt(a), _fmt(b), _fmt(c)])
            nums.append([None, a, b, c])
        if not any(_isnum(n[1]) for n in nums):
            raise WtbFormatError(f"정산표 「{name}」 시트에 계산된 값이 없습니다. 엑셀로 열어 저장한 뒤 다시 주십시오: {Path(path).name}")
        if name == "공시용BS":                       # WTB 양식은 마감 전 시산표라 당기순손익을 이익잉여금 바로 아래 줄에 따로 둔다(wtbform). 공시 과목은 둘의 합이다
            for k in range(1, len(rows)):            # (2026-10-20 시나리오 B: 이익잉여금 430,100,000 + 당기순손익 171,980,000 = dsd 602,080,000)
                if norm(rows[k][0]).startswith("당기순손익") and re.search(r"잉여금|결손금", rows[k - 1][0]):
                    for j in (1, 2, 3):
                        if _isnum(nums[k][j]):
                            nums[k - 1][j] = (nums[k - 1][j] or 0) + nums[k][j]
                            rows[k - 1][j] = _fmt(nums[k - 1][j])
        key = f"정산표 {name}"
        parts.append(Part(key, key, [("t", Table(key, "", 1, ["과목", "당기(최종)", "전기", "당기(회사제시)"], rows, nums))]))
    return parts


# ---------------------------------------------------------------- 감사 법인 양식
def _clean(v) -> str:
    from . import dsd

    return dsd._label(str(v or "").strip()) if str(v or "").strip() else ""


def _find_head(ws, words, rows: int = 15, cols: int = 8):
    for r in range(1, min(ws.max_row, rows) + 1):
        for j in range(1, min(ws.max_column, cols) + 1):
            if norm(ws.cell(r, j).value).lower() in words:
                return r, j
    return None


def _firm_fs(ws) -> list[dict] | None:
    """BS · IS 시트: 과목 | 당기(내역 · 합계 두 칸) | 전기(두 칸). 합계 칸에 금액이 있는 줄은 합계 줄이다."""
    h = _find_head(ws, ("과목",))
    if h is None:
        return None
    r0, lc = h
    cols = {j: norm(ws.cell(r0, j).value) for j in range(lc + 1, ws.max_column + 1)}
    cur = next((j for j, x in cols.items() if x.startswith("당기") or "(당)" in x), None)
    pri = next((j for j, x in cols.items() if x.startswith("전기") or "(전)" in x), None)
    if cur is None or pri is None or pri <= cur:
        return None
    width = pri - cur
    out = []
    for r in range(r0 + 1, ws.max_row + 1):
        label = _clean(ws.cell(r, lc).value)
        if not label:
            continue
        a = next(((v, j - cur) for j in range(cur, pri) if (v := _num(ws.cell(r, j).value)) is not None), (None, 0))
        b = next((v for j in range(pri, pri + width) if (v := _num(ws.cell(r, j).value)) is not None), None)
        out.append({"name": label, "final": a[0], "prior": b, "sub": a[1] > 0})
    return out


def _firm_acc(ws) -> list[dict] | None:
    """WBS · WIS 시트: 계정과목 | 최초 회사제시 | 수정(차변 · 대변) | 수정후. 차변 · 대변 칸이 있는 줄이 계정 줄이다."""
    h = _find_head(ws, ("계정과목", "과목"))
    if h is None:
        return None
    r0, lc = h
    col, last = {}, r0
    for r in range(r0, min(r0 + 3, ws.max_row) + 1):
        for j in range(lc + 1, ws.max_column + 1):
            x = norm(ws.cell(r, j).value).lower()
            if not x:
                continue
            k = ("given" if "회사제시" in x else "final" if ("수정후" in x or x in ("final", "최종")) else
                 "dr" if x in ("차변", "dr", "dr.") else "cr" if x in ("대변", "cr", "cr.") else None)
            if k and k not in col:
                col[k], last = j, max(last, r)
    if "given" not in col or "final" not in col:
        return None
    out = []
    for r in range(last + 1, ws.max_row + 1):
        label = _clean(ws.cell(r, lc).value)
        if not label:
            continue
        g, f = _num(ws.cell(r, col["given"]).value), _num(ws.cell(r, col["final"]).value)
        leaf = any(ws.cell(r, col[k]).value is not None for k in ("dr", "cr") if k in col)
        out.append({"name": label, "given": g, "final": f, "delta": (f or 0) - (g or 0), "leaf": leaf, "line": None, "aje": []})
    return out


def _firm_aje(ws) -> list[dict]:
    """AJE 시트: No. | 내역 | 차변(계정 · 금액) | 대변(계정 · 금액). 번호 칸에 숫자가 있는 줄에서 수정분개 하나가 시작한다."""
    h = _find_head(ws, ("no.", "no", "번호"), cols=6)
    if h is None:
        return []
    r0, nc = h
    col = {}
    for j in range(nc + 1, ws.max_column + 1):
        x = norm(ws.cell(r0, j).value)
        k = "desc" if x in ("내역", "적요", "내용") else "dr" if x == "차변" else "cr" if x == "대변" else None
        if k and k not in col:
            col[k] = j
    if "dr" not in col or "cr" not in col:
        return []
    out, cur = [], None
    for r in range(r0 + 1, ws.max_row + 1):
        no = ws.cell(r, nc).value
        if _num(no) is not None:
            cur = {"no": int(no) if float(no).is_integer() else no, "desc": str(ws.cell(r, col["desc"]).value or "").strip() if "desc" in col else "",
                   "lines": []}
            out.append(cur)
        elif isinstance(no, str) and no.strip() and cur is not None:
            break                                    # 수정분개 아래의 대사 표(「당기순이익 대사」)
        if cur is None:
            continue
        for side, j in (("차변", col["dr"]), ("대변", col["cr"])):
            acc, amt = ws.cell(r, j).value, _num(ws.cell(r, j + 1).value)
            if isinstance(acc, str) and acc.strip() and amt:
                cur["lines"].append((side, _clean(acc), amt))
    return out


def _keys(s: str) -> set[str]:
    """이름 견줄 열쇠: 이름, 괄호 말을 뗀 것, 괄호 안 말(「대손충당금(외상매출금)」 · 「미처분이익잉여금(미처리결손금)」)."""
    n = norm(s)
    out = {n, norm(_alias(n))}
    m = re.match(r"^(.+?)\((.+)\)$", n)
    if m:
        out |= {m.group(1), m.group(2), norm(_alias(m.group(1))), norm(_alias(m.group(2)))}
    return {x for x in out if x}


def _same(a, b) -> bool:
    return a is not None and b is not None and abs(abs(a) - abs(b)) < 0.5


def _targets(lines: list[dict]) -> list[int]:
    """계정을 이을 수 있는 과목 줄: 내역 줄과, 바로 아래 내역 줄이 없는 합계 줄(「II. 이익잉여금」 한 줄로 적은 과목)."""
    out = []
    for i, x in enumerate(lines):
        if x["final"] is None:
            continue
        if not x["sub"]:
            out.append(i)
            continue
        nxt = next((y for y in lines[i + 1:] if y["final"] is not None), None)
        if nxt is None or nxt["sub"]:
            out.append(i)
    return out


def _link(lines: list[dict], accs: list[dict]) -> None:
    """계정 줄을 공시 과목(내역 줄)에 잇는다(acc["line"] = 과목 번호). 이름 → 금액 → 이웃 계정 금액의 합 순서.
    이웃 계정 합은 수정이 있는 계정에만 쓴다(그 계정이 든 과목을 알아야 과목의 회사제시 금액을 구한다)."""
    leaf = [i for i in _targets(lines)]
    used: dict[int, list[int]] = {}
    for k, a in enumerate(accs):                      # 1 이름
        if not a["leaf"]:
            continue
        hit = [i for i in leaf if _keys(lines[i]["name"]) & _keys(a["name"])]
        if len(hit) > 1:
            hit = [i for i in hit if _same(lines[i]["final"], a["final"])] or [i for i in hit if i not in used] or hit
        if hit:
            a["line"] = hit[0]
            used.setdefault(hit[0], []).append(k)
    for k, a in enumerate(accs):                      # 2 금액(과목 하나와만 같을 때)
        if not a["leaf"] or a["line"] is not None or a["final"] is None or abs(a["final"]) < MIN_AMOUNT:
            continue
        hit = [i for i in leaf if i not in used and _same(lines[i]["final"], a["final"])]
        if len(hit) == 1:
            a["line"] = hit[0]
            used[hit[0]] = [k]
    runs, run = [], []                                # 3 이웃 계정의 합. 합계 줄에서 끊긴 계정 묶음 안에서만 본다
    for k, a in enumerate(accs):
        if a["leaf"]:
            run.append(k)
        elif run:
            runs.append(run)
            run = []
    if run:
        runs.append(run)
    for run in runs:
        for k in run:
            if accs[k]["line"] is not None or not accs[k]["delta"]:
                continue
            p = run.index(k)
            done = False
            for size in range(2, min(len(run), 12) + 1):
                for s in range(max(0, p - size + 1), min(p, len(run) - size) + 1):
                    win = run[s:s + size]
                    if any(accs[x]["line"] is not None for x in win):
                        continue
                    tot = sum(accs[x]["final"] or 0 for x in win)
                    hit = [i for i in leaf if i not in used and abs(tot) >= MIN_AMOUNT and _same(lines[i]["final"], tot)]
                    if len(hit) == 1:
                        for x in win:
                            accs[x]["line"] = hit[0]
                        used[hit[0]] = list(win)
                        done = True
                        break
                if done:
                    break


def _refs(nos) -> str:
    return " · ".join(f"No.{v}" for v in nos)


def _read_firm(wb) -> tuple[list[Part], dict]:
    names = {norm(n).upper(): n for n in wb.sheetnames}
    aje_name = next((n for n in wb.sheetnames if AJE_SHEET.match(n)), None)
    entries = _firm_aje(wb[aje_name]) if aje_name else []
    fs, accs = [], {}
    for want in FIRM_FS:
        n = names.get(norm(want).upper())
        if n is None:
            continue
        lines = _firm_fs(wb[n])
        if lines is None or not any(x["final"] is not None for x in lines):
            continue
        an = names.get(FIRM_ACC[want])
        acc = accs.get(an) or (_firm_acc(wb[an]) if an else None)
        fs.append((n, lines, an if acc else None, acc))
        if acc:
            accs[an] = acc
    for want in sorted(set(FIRM_ACC.values())):        # 공시 과목 시트 없이 계정별 시트만 있는 양식
        n = names.get(want)
        if n and n not in accs and (acc := _firm_acc(wb[n])):
            accs[n] = acc
    if not fs and not accs:
        return [], {}
    by_name: dict[str, dict] = {}                    # 수정분개 번호를 계정에 붙인다(계정 이름으로 찾는다)
    for acc in accs.values():
        for a in acc:
            if a["leaf"]:
                by_name.setdefault(norm(a["name"]), a)
    lost_aje = set()
    for e in entries:
        for _, acc_name, _ in e["lines"]:
            a = by_name.get(norm(acc_name))
            if a is None:
                lost_aje.add(acc_name)
            elif e["no"] not in a["aje"]:
                a["aje"].append(e["no"])
    parts, unlinked, n_adj = [], [], 0
    for n, lines, an, acc in fs:
        delta = [0.0] * len(lines)
        linked = [0] * len(lines)
        aje: list[list] = [[] for _ in lines]
        if acc:
            _link(lines, acc)
            for a in acc:
                if a["line"] is None:
                    if a["leaf"] and a["delta"]:
                        unlinked.append(f"{an} {a['name']}")
                    continue
                delta[a["line"]] += a["delta"]
                linked[a["line"]] += 1
                aje[a["line"]] += [x for x in a["aje"] if x not in aje[a["line"]]]
        given = [None] * len(lines)
        for i, x in enumerate(lines):                      # 내역 줄: 수정 후 금액 - 그 줄에 든 계정의 수정
            if acc and x["final"] is not None and not x["sub"]:
                given[i] = x["final"] - delta[i]
        for i, x in enumerate(lines):                      # 합계 줄: 계정별 시트의 같은 이름(없으면 같은 금액) 합계 줄, 그것도 없으면 바로 아래 내역 줄의 합
            if not acc or x["final"] is None or not x["sub"]:
                continue
            subs = [a for a in acc if not a["leaf"] and a["given"] is not None and _same(a["final"], x["final"])]
            y = next((a for a in subs if _keys(a["name"]) & _keys(x["name"])), None) or (subs[0] if len(subs) == 1 and abs(x["final"]) >= MIN_AMOUNT else None)
            if y is not None:
                given[i] = (-1 if (y["final"] or 0) * x["final"] < 0 else 1) * y["given"]
                continue
            if linked[i]:                                  # 내역 줄 없이 계정을 이은 합계 줄
                given[i] = x["final"] - delta[i]
                continue
            block = []
            for j in range(i + 1, len(lines)):
                if lines[j]["sub"]:
                    break
                if lines[j]["final"] is not None:
                    block.append(j)
            if block and all(given[j] is not None for j in block) and _same(sum(lines[j]["final"] for j in block), x["final"]):
                given[i] = (-1 if sum(lines[j]["final"] for j in block) * x["final"] < 0 else 1) * sum(given[j] for j in block)
        rows, nums = [], []
        for i, x in enumerate(lines):
            d = None if given[i] is None else x["final"] - given[i]
            n_adj += 1 if d else 0
            rows.append([x["name"], _fmt(x["final"]), _fmt(x["prior"]), _fmt(given[i]), _fmt(d) if d else ("" if d is None else "-"), _refs(aje[i])])
            nums.append([None, x["final"], x["prior"], given[i], d, None])
        key = f"정산표 {n}"
        cap = f"감사 법인 WTB {n} 시트" + (f"(회사제시 · 수정은 {an} 시트에서 구함)" if an else "")
        parts.append(Part(key, key, [("t", Table(key, cap, 1, list(FIRM_HEAD), rows, nums))]))
    for an, acc in accs.items():                       # 계정별 시트: 과목 시트에 없는 이름을 찾을 때와 보기용
        rows = [[a["name"], _fmt(a["final"]), "", _fmt(a["given"]), _fmt(a["delta"]) if a["delta"] else "", _refs(a["aje"])] for a in acc]
        nums = [[None, a["final"], None, a["given"], a["delta"] or None, None] for a in acc]
        key = f"정산표 {an}"
        parts.append(Part(key, key, [("t", Table(key, f"감사 법인 WTB {an} 시트(계정별)", 1, list(FIRM_HEAD), rows, nums))]))
    if entries:
        rows, nums = [], []
        for e in entries:
            for side, acc_name, amt in e["lines"] or [("", "", None)]:
                rows.append([f"No.{e['no']}", e["desc"], side, acc_name, _fmt(amt)])
                nums.append([None, None, None, None, amt])
        key = f"정산표 {aje_name}"
        parts.append(Part(key, key, [("t", Table(key, f"감사 수정분개({aje_name} 시트)", 1, ["번호", "내역", "차대", "계정", "금액"], rows, nums, kind=ADJ))]))
    stats = {"양식": "감사 법인 WTB", "시트": [n for n, *_ in fs] + [a for a in accs] + ([aje_name] if aje_name else []),
             "수정분개": sum(1 for e in entries if e["lines"]), "금액 없는 수정분개": sum(1 for e in entries if not e["lines"]),
             "수정한 과목": n_adj, "과목에 잇지 못한 수정 계정": unlinked, "계정을 찾지 못한 수정분개 계정": sorted(lost_aje)}
    return parts, stats


def describe(wdoc: Doc) -> str:
    """정산표를 어떻게 읽었는지 한 줄."""
    s = wdoc.stats or {}
    if s.get("양식") != "감사 법인 WTB":
        return f"정산표 {wdoc.source}: {s.get('양식', '8100B · WTB 양식')}"
    out = (f"정산표 {wdoc.source}: 감사 법인 WTB(시트 {' · '.join(s['시트'])}) · 감사 수정분개 {s['수정분개']}건"
           + (f"(금액 없는 것 {s['금액 없는 수정분개']}건 뺌)" if s["금액 없는 수정분개"] else "") + f" · 수정이 든 과목 줄 {s['수정한 과목']}곳")
    if s["과목에 잇지 못한 수정 계정"]:
        out += f" · 과목에 잇지 못한 수정 계정 {len(s['과목에 잇지 못한 수정 계정'])}개({', '.join(s['과목에 잇지 못한 수정 계정'][:5])})"
    if s["계정을 찾지 못한 수정분개 계정"]:
        out += f" · 계정별 시트에 없는 수정분개 계정 {len(s['계정을 찾지 못한 수정분개 계정'])}개({', '.join(s['계정을 찾지 못한 수정분개 계정'][:5])})"
    return out


def no_adjustments(wdoc: Doc) -> bool:
    """정산표의 당기(회사제시)와 당기(최종)이 모든 과목에서 같으면 수정분개가 없는 정산표다. 회사제시 열이 없으면 알 수 없어 False.
    10단계 재평가: 수정분개를 넣기 전의 정산표와 대사하면 정산표 대사가 회사 장부와의 대사일 뿐인데 그 사실이 엑셀에 보이지 않았다."""
    seen = False
    for t in wdoc.tables():
        if t.kind:
            continue
        for n in t.nums:
            a, c = n[1], n[3] if len(n) > 3 else None
            if c is None or a is None:
                continue
            seen = True
            if abs(a - c) > 0.5:
                return False
    return seen


ALIAS = (("미처리결손금", "미처분이익잉여금"), ("결손금", "이익잉여금"), ("순손실", "순이익"), ("영업손실", "영업이익"), ("총손실", "총이익"))
MIN_HINT = 1000                                # 이보다 작은 금액은 우연히 같을 수 있어 같은 금액의 줄을 찾지 않는다
NAME_ONLY = "과목 이름 차이"                    # 정산표의 다른 줄이 같은 금액일 때 붙이는 유형 이름


def _alias(name: str) -> str:
    """손실 · 결손금으로 적은 과목을 정산표의 이익 · 잉여금 줄과 짝짓기 위한 이름. 금액의 부호는 절댓값 비교가 받는다."""
    for a, b in ALIAS:
        if a in name:
            return name.replace(a, b)
    return name


def _same_amount(wtb: Doc, v: float, b: int, name: str) -> str:
    """정산표에서 같은 금액(절댓값)인 다른 이름의 줄(셋까지). 과목 이름만 다른 경우를 알려 주는 참고다."""
    if abs(v) < MIN_HINT:
        return ""
    hits = []
    for q in wtb.tables():
        if q.kind:
            continue
        for k in range(len(q.rows)):
            x = q.rows[k][0].strip()
            if x != name.strip() and x not in hits and _isnum(q.nums[k][b]) and abs(abs(q.nums[k][b]) - abs(v)) < 0.5:
                hits.append(x)
    return " · ".join(f"「{x}」" for x in hits[:3])


def checks(doc: Doc, wtb: Doc, totals: set | None = None) -> tuple[list[Check], list[dict]]:
    """dsd 본표 과목 = 정산표 금액. (검증, 정산표에서 찾지 못한 과목과 금액).
    `totals` 는 세로 합계 검증이 걸린 줄 (표 id, 줄). 정산표에 없는 소계 줄(유동자산 · 비유동부채)은 아래 과목의 합으로 이미 검증되니 찾지 못함으로 알리지 않는다
    (2026-10-20, 10-17 시나리오 B 에서 소계 여섯 줄이 알림으로 나왔다).
    과목은 이름으로 짝짓는다(회사 양식 과목 시트 먼저, 없으면 표시과목 시트). 부호 표기가 서로 달라 절댓값이 같으면 같은 금액으로 본다.
    금액이 다른데 정산표의 다른 줄 하나가 같은 금액이면 FALSE 로 두지 않고 「확인」 으로 내리고 그 줄 이름을 적는다."""
    from .dsd_compare import _pairs

    out, lost = [], []
    for name in STATEMENTS:
        t = doc.table(name)
        if t is None:
            continue
        labels = _labels(t)
        cols = [(a, b, word) for a, b, word in ((period_col(t, "당기"), 1, "당기(최종)"), (period_col(t, "전기"), 2, "전기")) if a is not None]
        done, met = set(), set()
        for names in (labels, [_alias(x) for x in labels]):
            for q in [x for x in wtb.tables() if not x.kind]:
                open_rows = {i for i in range(len(labels)) if any((i, a) not in done for a, _, _ in cols)}
                pairs = _pairs([names[i] if i in open_rows else f"\x00{i}" for i in range(len(labels))], [r[0] for r in q.rows])
                for i, k in sorted(pairs.items()):
                    if i not in open_rows:
                        continue
                    met.add(i)
                    for a, b, word in cols:
                        if (i, a) in done or not _isnum(t.nums[i][a]) or not _isnum(q.nums[k][b]):
                            continue
                        v, w = t.nums[i][a] * t.unit, q.nums[k][b]
                        sign = -1 if v and w and abs(v + w) < abs(v - w) else 1      # 차감 과목은 두 곳의 부호 표기가 다를 수 있다
                        other = _same_amount(wtb, v, b, q.rows[k][0]) if abs(abs(v) - abs(w)) >= 0.5 else ""
                        c = Check(f"조서|{t.id}|{labels[i]}#{_ref(labels, i)[1]}|{word}", "조서-재무제표", "레퍼", t.id, i, a, [(WTB_ID + q.id, k, b, sign)], None,
                                  f"{q.id} {q.rows[k][0]} · {word}" + (" (부호 반대)" if sign < 0 else ""), "조서",
                                  "감사한 금액(8100B 정산표)과 같아야 한다" + (" (부호 표기는 서로 다르다)" if sign < 0 else ""),
                                  when=f"정산표의 {other} 줄이 같은 금액이다. 과목 이름만 다르면 차이가 없다" if other else "", usual=NAME_ONLY if other else "")
                        given = q.nums[k][3] if b == 1 and len(q.nums[k]) > 3 else None      # 수정분개 반영 확인: dsd 가 회사제시 금액 그대로인가
                        refs = q.rows[k][5] if b == 1 and len(q.rows[k]) > 5 else ""        # 감사 법인 WTB: 그 과목에 든 감사 수정분개 번호
                        if abs(abs(v) - abs(w)) >= 0.5 and _isnum(given) and abs(abs(v) - abs(given)) < 0.5 * max(t.unit, 1):
                            c.finding = (f"숫자 오류: dsd 금액 {v:,.0f}원이 정산표의 회사제시 금액과 같고 최종 금액 {w:,.0f}원과 다르다. "
                                         f"수정분개 {w - given:,.0f}원이 dsd 에 반영되지 않았다." + (f" 감사 수정분개 {refs}." if refs else ""))
                            c.when, c.usual = "", ""
                        elif abs(abs(v) - abs(w)) >= 0.5 and refs:
                            c.label += f" · 감사 수정분개 {refs}"
                        out.append(c)
                        done.add((i, a))
        for i in range(len(labels)):
            for a, b, word in cols:
                if (i, a) in done or not _isnum(t.nums[i][a]) or not t.nums[i][a]:
                    continue
                if i not in met and (a != cols[0][0] or (t.id, i) in (totals or ())):
                    continue                                 # 과목을 통째로 못 찾은 줄은 당기에서 한 번만 알린다. 합계 검증이 걸린 소계 줄은 알리지 않는다
                why = f"조서: 정산표에서 「{labels[i]}」 과목을 찾지 못함" if i not in met else f"조서: 정산표의 「{labels[i]}」 과목에 {word} 금액이 없음"
                other = _same_amount(wtb, t.nums[i][a] * t.unit, b, labels[i]) if i not in met else ""
                if other:                                    # 시나리오 B: dsd 「금융원가」 한 줄 = 정산표 「이자비용」(2026-10-20)
                    why += f". 정산표의 {other} 줄이 같은 금액이다"
                lost.append({"table": t.id, "kind": "조서", "why": why})
    return out, lost

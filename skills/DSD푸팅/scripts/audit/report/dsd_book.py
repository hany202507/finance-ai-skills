"""DSD 푸팅 엑셀(2026-10-12). 부분마다 시트 하나: 왼쪽에 원문(문단과 표), 표 바로 오른쪽에 합계 검증, 그 오른쪽에 표 사이 연결의 검증 줄,
그 오른쪽에 기존 판 대비 차이.

표 안 합계 검증은 합계 칸과 같은 줄, 표 오른쪽의 같은 차례 열에 TRUE · 단수 · FALSE 로 보인다. 재무상태표가 A~C 열이면
유동자산 줄의 D · E 열이 당기 · 전기 합계가 맞는지다. 칸의 수식이 합계 칸과 더한 칸들을 직접 가리킨다.
단수(차이가 표시 단위로 MINOR 이하, 천원 · 백만원 표만)는 노랑, FALSE 는 빨강이다.
검증 줄의 값 · 차이 · 결과와 기존 판 대비 차이는 모두 수식이다. 오버뷰는 시트마다 네 가지를 신호등으로 보인다:
합계(표 안 소계 · 합계), 레퍼(표 사이 연결과 계산), 전기(작년 보고서와 일치), 버전(기존 판 대비 바뀐 곳, 바뀐 곳이 있으면 노랑).
신호등은 조건부 서식 아이콘이고 칸의 값은 각 시트의 결과 열과 「바뀜」 열을 세는 수식이다. 바뀐 칸과 문단은 원문에서 빨간 글자가 된다.
두 시트에 걸친 연결은 검증 줄의 칸이 있는 시트에 한 번 적고, 대사 대상이 있는 시트에는 결과를 끌어온 줄을 둔다.
기존 판 원문은 「전_」, 작년 보고서 원문은 「작년_」, 정산표는 「조서_」, 함께 낸 다른 보고서는 「짝_」 시트에 값으로 둔다.

회계사가 O · X 를 적는 확인 칸은 「검증 목록」 시트 한 곳이다(9단계 재평가: 확인 칸이 두 군데라 잘못 적었다). 주석 시트의 확인 칸은 그 칸을 끌어온다.
수식 칸에는 엔진이 계산한 값을 함께 저장한다(엑셀 프로그램이 아닌 미리보기에서도 TRUE · FALSE 가 보이게). 엑셀로 열면 다시 계산한다.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink

from .. import wp
from . import dsd_compare
from .dsd_compare import Match
from .dsd_grid import Doc
from . import dsd_links
from .dsd_links import MINOR, PAIR_ID, PRIOR_ID, Check

COLS = ["묶음", "구분", "칸", "이 칸 값", "대사 대상", "대사 대상 값", "차이", "결과", "출처", "근거", "확인", "id", "단위"]
K_GROUP, K_KIND, K_CELL, K_LEFT, K_LABEL, K_RIGHT, K_DIFF, K_OK, K_SRC, K_BASIS, K_MARK, K_ID, K_UNIT = range(13)
NUM = '#,##0;(#,##0);"-"'
RED = "B3261E"
GREEN = "1E7B34"
LINK = "1B2A4A"
FILL_MINOR = PatternFill("solid", fgColor="FFE699")
SUM, MINOR_TEXT, ASK_TEXT = "합계", "단수", "확인"
FILL_ASK = PatternFill("solid", fgColor="F8CBAD")
OVERVIEW, LIST, UNRESOLVED = "오버뷰", "검증 목록", "풀지 못한 연결"
MIRROR = "끌어옴"
HEAD_ROW, FIRST_ROW = 2, 4
PREV, PRIOR = "전_", "작년_"            # 기존 판 · 작년 보고서 원문 시트 이름 앞머리
WTB, WTB_ID = "조서_", "조서:"          # 정산표 시트 이름 앞머리와 표 id 앞머리
PAIR = "짝_"                            # 함께 낸 다른 보고서(연결 ↔ 별도) 시트 이름 앞머리
CHANGES = "바뀐 곳"
PREV_ID = "전:"                         # 기존 판의 표 id 앞머리(주소를 가를 때만 쓴다)
OV_HEAD = ("시트", "제목", "합계", "레퍼", "전기", "버전", "바뀐 곳", "검증 범위")
TODO, SCOPE, COVER, WPS = "볼 것", "검증 범위", "표지", "조서 대사"
TODO_HEAD = ("상태", "원인", "판정", "검증", "시트", "칸", "이 칸", "이 칸 값", "대사 대상", "대사 대상 값", "차이", "단위",
             "차이 설명 후보", "고칠 칸 · 맞는 숫자", "근거 · 조건", "작년 보고서", "검토 메모", "확인", "늘 나는 차이", "묶음", "id")
TODO_COL = {h: j for j, h in enumerate(TODO_HEAD, 1)}      # 「볼 것」 열 번호(11단계에 검증 · 차이 설명 후보 · 고칠 칸 열을 더했다)
OV_WTB = "정산표"
FILL_NONE = PatternFill("solid", fgColor="E7E6E6")       # 어느 검증에도 들어가지 않은 숫자 칸
FILL_WP = PatternFill("solid", fgColor="DDEBF7")         # 보고서 안에 맞춰 볼 곳이 없어 조서와 맞출 숫자 칸
LIST_HEAD = ("No", "시트", "구분", "칸", "대사 대상", "이 칸 값", "대사 대상 값", "차이", "결과", "출처", "확인", "id", "묶음", "근거", "조건", "판정", "0 = 0", "단위")
OV_ROW = 5
UNIT_NAME = {1: "원", 1000: "천원", 10000: "만원", 1000000: "백만원", 100000000: "억원", 10 ** 12: "조원", 0.01: "%"}


def unit_name(u: float) -> str:
    return UNIT_NAME.get(u, f"×{u:g}")


def _num(x: float) -> str:
    """수식에 넣는 숫자(지수 꼴 없이)."""
    if float(x).is_integer():
        return str(int(x))
    return f"{x:.12f}".rstrip("0").rstrip(".")


@dataclass
class Layout:
    addr: dict = field(default_factory=dict)          # (표 id, 행, 열) → (시트, 줄, 칸). 머리줄은 행 -1
    check_row: dict = field(default_factory=dict)     # 검증 id → (시트, 줄). 표 사이 연결과 전기 대사만. 합계 검증은 표 옆 칸에 있다
    sum_at: dict = field(default_factory=dict)        # (표 id, 행, 열) → [그 칸의 합계 검증]
    list_row: dict = field(default_factory=dict)      # 검증 id → 검증 목록 시트의 줄
    mirrors: list = field(default_factory=list)       # (시트, 줄, 검증)
    cb: dict = field(default_factory=dict)            # 시트 → 검증 줄 첫 칸
    db: dict = field(default_factory=dict)            # 시트 → 기존 판 대비 첫 칸(기존 판이 있을 때)
    items: dict = field(default_factory=dict)         # 시트 → [(줄, 항목 번호, 항목)]. 표는 머리줄의 줄
    para_row: dict = field(default_factory=dict)      # (시트, 항목 번호) → 줄
    gone_row: dict = field(default_factory=dict)      # 시트 → 없어진 것 첫 줄
    last: dict = field(default_factory=dict)          # 시트 → 마지막 줄
    cache: dict = field(default_factory=dict)         # (시트, 줄, 칸) → 수식 칸에 함께 저장할 계산값
    changed: dict = field(default_factory=dict)       # 시트 → 「바뀜」 열의 합(기존 판 대비)

    def hb(self, sheet: str) -> int:
        """「바뀜」 열(줄마다 바뀐 곳 수)."""
        return self.db[sheet] + self.cb[sheet] - 1


def _put(lay: Layout, ws, r: int, c: int, value, cached=None, **style):
    """칸에 값이나 수식을 쓴다. 수식이고 cached 가 있으면 계산값을 함께 저장할 목록에 넣는다."""
    x = ws.cell(r, c, value)
    if isinstance(value, str) and value.startswith("=") and cached is not None:
        lay.cache[(ws.title, r, c)] = cached
    for k, v in style.items():
        setattr(x, k, v)
    return x


def _q(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'!"


def _place_plain(lay: Layout, doc: Doc, sheet_prefix: str, id_prefix: str) -> None:
    """원문만 두는 시트(기존 판 · 작년 보고서 · 정산표 · 다른 보고서)의 배치. 표마다 위에 표 번호 줄이 있다."""
    for p in doc.parts:
        sheet, row, placed = sheet_prefix + p.key, FIRST_ROW, []
        for k, it in enumerate(p.items):
            if it[0] == "p":
                placed.append((row, k, it))
                lay.para_row[(sheet, k)] = row
                row += 1
                continue
            row += 1                                  # 표 번호 줄
            placed.append((row, k, it))
            t = it[1]
            for j in range(len(t.header)):
                lay.addr[(id_prefix + t.id, -1, j)] = (sheet, row, j + 1)
                for i in range(len(t.rows)):
                    lay.addr[(id_prefix + t.id, i, j)] = (sheet, row + 1 + i, j + 1)
            row += len(t.rows) + 2
        lay.items[sheet], lay.last[sheet] = placed, row


def place(doc: Doc, checks: list[Check], prev: Doc | None = None, prior: Doc | None = None, m: Match | None = None, wtb: Doc | None = None,
          pair: Doc | None = None) -> Layout:
    lay = Layout()
    own: dict[str, list] = {}
    for c in checks:
        own.setdefault(c.table, []).append(c)
    sheet_of = {t.id: doc.part_of(t.id).key for t in doc.tables()}
    mirror: dict[str, list] = {}
    for c in checks:
        other = next((tid for tid, *_ in c.terms if tid in sheet_of and sheet_of[tid] != sheet_of[c.table]), None)
        if other:
            mirror.setdefault(other, []).append(c)
    summed = {c.table for c in checks if c.group == SUM}
    for n, c in enumerate(checks, 2):
        lay.list_row[c.id] = n
    for p in doc.parts:                           # 합계 검증이 있는 표는 오른쪽에 숫자 열만큼 자리를 더 쓴다
        width = max([2 * len(x[1].header) - 1 if x[1].id in summed else len(x[1].header) for x in p.items if x[0] == "t"] + [4])
        lay.cb[p.key] = width + 2
        if prev is not None:
            lay.db[p.key] = lay.cb[p.key] + len(COLS) + 1
        row, placed = FIRST_ROW, []
        for k, it in enumerate(p.items):
            if it[0] == "p":
                placed.append((row, k, it))
                lay.para_row[(p.key, k)] = row
                row += 1
                continue
            row += 1                                  # 표 번호 줄(「[주석5-1] 단위 천원」). 볼 것 · 검증 범위가 표 번호로 가리킨다
            placed.append((row, k, it))
            t = it[1]
            for j in range(len(t.header)):
                lay.addr[(t.id, -1, j)] = (p.key, row, j + 1)
            for i in range(len(t.rows)):
                for j in range(len(t.header)):
                    lay.addr[(t.id, i, j)] = (p.key, row + 1 + i, j + 1)
            k2 = row
            for c in own.get(t.id, []):
                if c.group == SUM:
                    lay.sum_at.setdefault((t.id, c.r, c.c), []).append(c)
                    continue
                lay.check_row[c.id] = (p.key, k2)
                k2 += 1
            for c in mirror.get(t.id, []):
                lay.mirrors.append((p.key, k2, c))
                k2 += 1
            row = max(row + 1 + len(t.rows), k2) + 1
        if m is not None and m.gone.get(p.key):
            lay.gone_row[p.key] = row + 1
            row += 2 + len(m.gone[p.key])
        lay.items[p.key], lay.last[p.key] = placed, row
    if prev is not None:
        _place_plain(lay, prev, PREV, PREV_ID)
    if prior is not None:
        _place_plain(lay, prior, PRIOR, PRIOR_ID)
    if wtb is not None:
        _place_plain(lay, wtb, WTB, WTB_ID)
    if pair is not None:
        _place_plain(lay, pair, PAIR, PAIR_ID)
    return lay


def _a1(lay: Layout, tid: str, r: int, c: int, sheet: str) -> str:
    s, row, col = lay.addr[(tid, r, c)]
    return ("" if s == sheet else _q(s)) + f"{L(col)}{row}"


def _unit_of(tabs: dict, tid: str, r: int, c: int) -> float:
    return tabs[tid].unit_at(r, c)


def _units(tabs: dict, c: Check) -> set:
    if c.plain:
        return {1}
    return {_unit_of(tabs, c.table, c.r, c.c)} | {_unit_of(tabs, tid, r, k) for tid, r, k, _ in c.terms}


def _run(terms: list) -> tuple | None:
    """항이 셋 이상이고 전부 더하며 한 표의 한 열(또는 한 줄)에서 빈틈없이 이어지면 (표, 첫 행, 첫 열, 끝 행, 끝 열)."""
    if len(terms) < 3 or any(s != 1 for *_, s in terms) or len({t[0] for t in terms}) != 1:
        return None
    rows, cols = [t[1] for t in terms], [t[2] for t in terms]
    if len(set(cols)) == 1 and rows == list(range(rows[0], rows[0] + len(rows))):
        return terms[0][0], rows[0], cols[0], rows[-1], cols[0]
    if len(set(rows)) == 1 and cols == list(range(cols[0], cols[0] + len(cols))):
        return terms[0][0], rows[0], cols[0], rows[0], cols[-1]
    return None


def _right_formula(lay: Layout, tabs: dict, c: Check, sheet: str) -> str:
    if c.expr:                                   # 계산 검증: 항을 원(수량 · 비율)으로 바꿔 식을 계산하고 왼쪽 칸의 단위로 되돌린다
        refs = []
        for tid, r, k, _ in c.terms:
            u = _unit_of(tabs, tid, r, k)
            refs.append(_a1(lay, tid, r, k, sheet) + ("" if u == 1 else f"*{_num(u)}"))
        lu = _unit_of(tabs, c.table, c.r, c.c)
        body = c.expr.format(*refs)
        return f"=({body})" if lu == 1 else f"=({body})/{_num(lu)}"
    units = _units(tabs, c)
    big = max(units)
    if c.span and len(units) == 1:
        tid, r0, c0, r1, c1 = c.span
        return f"=SUM({_a1(lay, tid, r0, c0, sheet)}:{_a1(lay, tid, r1, c1, sheet).split('!')[-1]})"
    run = _run(c.terms) if len(units) == 1 and c.left_div == 1 else None
    if run:                                      # 한 줄로 이어진 칸을 전부 더하면 SUM 으로 적는다(수식을 눌렀을 때 범위가 한눈에 보인다)
        tid, r0, c0, r1, c1 = run
        return f"=SUM({_a1(lay, tid, r0, c0, sheet)}:{_a1(lay, tid, r1, c1, sheet).split('!')[-1]})"
    parts = []
    for tid, r, k, s in c.terms:
        ref = _a1(lay, tid, r, k, sheet)
        u = 1 if c.plain else _unit_of(tabs, tid, r, k)
        if len(units) > 1 and u != big:
            ref = f"{ref}*{_num(u)}/{_num(big)}"
        if abs(s) != 1:
            ref = f"{ref}*{abs(s):.10g}"
        parts.append(("+" if s > 0 else "-") + ref)
    body = "".join(parts)
    scaled = len(units) > 1 or c.left_div != 1 or any(abs(s) != 1 for *_, s in c.terms)
    return f"=ROUND({body},0)" if scaled else f"={body}"


def _left_formula(lay: Layout, tabs: dict, c: Check, sheet: str) -> str:
    ref = _a1(lay, c.table, c.r, c.c, sheet)
    if c.zero_left:                              # 보고서에 없는 줄(처분계산서의 이익준비금 적립)은 0 으로 견준다
        return "=0"
    if c.expr:
        return f"={ref}"
    units = _units(tabs, c)
    if c.left_div != 1:
        return f"=ROUND({ref}/{c.left_div:.10g},0)"
    lu = 1 if c.plain else _unit_of(tabs, c.table, c.r, c.c)
    if len(units) > 1 and lu != max(units):
        return f"=ROUND({ref}*{_num(lu)}/{_num(max(units))},0)"
    return f"={ref}"


def _tol_text(tol: float) -> str:
    """허용 폭. 정수가 아니면(계산 검증) 소수 끝의 오차를 받으려고 아주 작은 수를 더한다."""
    return _num(tol) if float(tol).is_integer() else _num(tol + 1e-9)


def _state(d: str, tol, mark: str | None = None, when: bool = False, limit: float = MINOR, cmp: str = "") -> str:
    """차이 식 d 로 결과를 내는 식. 허용 폭 안이면 TRUE, 단수 한도(limit) 이하면 「단수」, 더 크면 FALSE. 확인 칸이 X 면 빈칸.
    조건부 연결(when)은 FALSE 자리에 「확인」 을 낸다. 원 단위끼리와 계산 검증은 단수 한도가 0 이라 「단수」 가 없다.
    부분 대사(cmp, 11단계)는 le 면 차이 ≤ 허용 폭, ge 면 차이 ≥ -허용 폭이면 TRUE 다."""
    bad = f'"{ASK_TEXT}"' if when else "FALSE"
    t = _tol_text(tol)
    if cmp:
        cond = f"{d}<={t}" if cmp == "le" else f"{d}>=-{t}"
        core = f'IF({cond},TRUE,{bad})' if when else cond
        return f'IF({mark}="X","",{core})' if mark else core
    if limit <= tol:
        core = f'IF(ABS({d})<={t},TRUE,{bad})' if when else f"ABS({d})<={t}"
    else:
        core = f'IF(ABS({d})<={t},TRUE,IF(ABS({d})<={_num(limit)},"{MINOR_TEXT}",{bad}))'
    return f'IF({mark}="X","",{core})' if mark else core


def state(c: Check):
    """엔진이 계산한 결과를 엑셀 칸의 값으로. TRUE · 「단수」 · 「확인」 · FALSE"""
    return MINOR_TEXT if c.minor else (ASK_TEXT if c.ok is False and c.when else c.ok)


def shown(c: Check):
    """확인 칸에 X 를 적은 검증과 판정으로 뺀 검증은 결과 칸이 빈칸이다."""
    return "" if c.excluded or c.dropped else state(c)


def _mark_of(c: Check) -> str:
    return "X" if c.excluded else ("O" if c.confirmed else "")


def _sum_formula(lay: Layout, tabs: dict, cs: list[Check], sheet: str) -> tuple[str, object]:
    """합계 칸 옆에 두는 식과 그 계산값. 합계 칸에서 더한 칸들을 뺀 차이로 결과를 낸다. 한 칸에 검증이 둘이면(세로합과 가로합) 나쁜 쪽을 보인다."""
    parts = []
    for c in cs:
        right = _right_formula(lay, tabs, c, sheet)[1:]
        d = f"{_left_formula(lay, tabs, c, sheet)[1:]}-" + (right if right.startswith("SUM(") else f"({right.lstrip('+')})")
        parts.append((d, c.tol, f"{_q(LIST)}$K${lay.list_row[c.id]}", c.minor_limit))
    if len(parts) == 1:
        d, tol, mark, limit = parts[0]
        if cs[0].dropped:
            return '=""', ""
        return "=" + _state(d, tol, mark, False, limit), shown(cs[0])
    codes = []
    for (d, tol, m, limit), c in zip(parts, cs):
        if c.dropped:                                # 판정으로 뺀 검증은 나쁜 쪽 고르기에 넣지 않는다
            codes.append("0")
            continue
        t = _tol_text(tol)
        inner = f"IF(ABS({d})<={_num(limit)},1,2)" if limit > tol else "2"
        codes.append(f'IF({m}="X",0,IF(ABS({d})<={t},0,{inner}))')
    worst = max(0 if c.excluded or c.dropped or c.ok else (1 if c.minor else 2) for c in cs)
    return f'=CHOOSE(MAX({",".join(codes)})+1,TRUE,"{MINOR_TEXT}",FALSE)', (True, MINOR_TEXT, False)[worst]


def _sum_block(ws, lay: Layout, tabs: dict, sheet: str, row: int, t) -> None:
    """표 바로 오른쪽에 합계 검증을 쓴다. 표의 j 번째 숫자 열의 검증은 표 너비 + j 번째 열, 합계 칸과 같은 줄이다."""
    w = len(t.header)
    cells = [(r, c) for (tid, r, c) in lay.sum_at if tid == t.id]
    if not cells:
        return
    for j in range(1, w):
        x = ws.cell(row, w + j, f"검증 {t.header[j]}".strip())
        x.font, x.alignment = wp.font(size=9, bold=True, color=wp.MUTED), Alignment(horizontal="center", wrap_text=True)
    for r, c in cells:
        f, v = _sum_formula(lay, tabs, lay.sum_at[(t.id, r, c)], sheet)
        _put(lay, ws, row + 1 + r, w + c, f, v, font=wp.font(size=9, color=GREEN), alignment=Alignment(horizontal="center"))
    rng = f"{L(w + 1)}{row + 1}:{L(2 * w - 1)}{row + len(t.rows)}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["FALSE"], font=Font(name=wp.FONT, size=9, bold=True, color=RED), fill=wp.FILL_FLAG))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{MINOR_TEXT}"'], font=Font(name=wp.FONT, size=9, bold=True), fill=FILL_MINOR))


def result_at(lay: Layout, tabs: dict, c: Check) -> tuple[str, int, int]:
    """검증의 결과가 보이는 칸 (시트, 줄, 열). 합계 검증은 표 옆 칸, 그 밖은 검증 줄의 결과 열."""
    if c.group == SUM:
        sheet, row, _ = lay.addr[(c.table, c.r, c.c)]
        return sheet, row, len(tabs[c.table].header) + c.c
    sheet, row = lay.check_row[c.id]
    return sheet, row, lay.cb[sheet] + K_OK


def _style_head(cell):
    cell.font = Font(name=wp.FONT, size=10, bold=True, color="FFFFFF")
    cell.fill = wp.FILL_H
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _link(x, location: str):
    x.hyperlink = Hyperlink(ref=x.coordinate, location=location)
    x.font = Font(name=wp.FONT, size=10, underline="single", color=LINK)


def _table_tag(t) -> str:
    """표 번호 줄의 글. 「[주석5-1] 단위 천원 · 당기 중 유형자산의 변동내역은 …」"""
    if t.kind:
        return f"[{t.id}] 주석 문단과 표 안 글자 칸에 적힌 금액을 모은 표(원문에는 없는 표, 금액은 적힌 단위)"
    tag = f"[{t.id}] 단위 {unit_name(t.unit)}"
    if t.cell_units:
        tag += " · 칸에 단위가 붙은 칸은 그 단위"
    return tag


def _original(ws, placed: list) -> None:
    """원문(문단과 표)을 쓴다. 표 위 줄에 표 번호와 단위를 적는다."""
    for row, _, it in placed:
        if it[0] == "p":
            ws.cell(row, 1, it[1]).font = wp.font()
            continue
        t = it[1]
        ws.cell(row - 1, 1, _table_tag(t)).font = wp.font(size=9, bold=True, color=wp.MUTED)
        for j, h in enumerate(t.header):
            x = ws.cell(row, j + 1, h)
            x.font, x.fill, x.border = wp.font(bold=True), wp.FILL_SUB, wp.BOX
            x.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, r in enumerate(t.rows):
            for j, text in enumerate(r):
                v = t.nums[i][j] if j else None
                x = ws.cell(row + 1 + i, j + 1, v if v is not None else (text or None))
                x.font, x.border = wp.font(), wp.BOX
                if v is not None:
                    u = (t.cell_units or {}).get((i, j))
                    if u == 0.01:
                        x.number_format = '0.0#####"%"'
                    elif u is not None and u != t.unit:      # 칸에 단위가 붙은 칸은 그 단위를 붙여 보인다(「346,820천원」)
                        w = unit_name(u)
                        x.number_format = f'#,##0"{w}";(#,##0"{w}");"-"' if isinstance(v, int) else f'General"{w}"'
                    else:
                        x.number_format = NUM if isinstance(v, int) else "General"


def _title(p) -> str:
    return f"주석 {p.key[2:]}. {p.title}" if p.key.startswith("주석") else p.title


def _plain_sheet(wb, lay: Layout, p, prefix: str, note: str):
    ws = wb.create_sheet(prefix + p.key)
    ws.cell(1, 1, f"{_title(p)} ({note})").font = Font(name=wp.FONT, size=12, bold=True, color=wp.MUTED)
    _original(ws, lay.items[prefix + p.key])
    ws.column_dimensions["A"].width = 26
    for j in range(2, 12):
        ws.column_dimensions[L(j)].width = 15
    return ws


def _diff_block(ws, lay: Layout, p, prev: Doc, m: Match) -> None:
    """기존 판 대비 차이. 표는 같은 모양으로, 칸마다 「새 판 칸 - 기존 판 같은 칸」. 문단은 같은지(0 · 1).
    「바뀜」 열이 줄마다 바뀐 곳 수를 센다. 새 줄은 1 을 더하고 없어진 것은 한 줄에 1 이다."""
    key = p.key
    db, hb, first = lay.db[key], lay.hb(key), lay.db[key]
    last = L(hb - 1)
    old_key = m.parts.get(key)
    old_tabs = {t.id: t for t in prev.tables()}
    old_items = next((q.items for q in prev.parts if q.key == old_key), [])
    _style_head(ws.cell(HEAD_ROW, db, "기존 판 대비"))
    _style_head(ws.cell(HEAD_ROW, hb, "바뀜"))
    total = 0

    def count(row, vals, extra=0):
        nonlocal total
        n = sum(1 for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool) and v != 0) + extra
        total += n
        rng = f"{L(first)}{row}:{last}{row}"
        _put(lay, ws, row, hb, f"=SUMPRODUCT(--ISNUMBER({rng}),--({rng}<>0))" + (f"+{extra}" if extra else ""), n, font=wp.font(color=wp.MUTED))

    for row, k, it in lay.items[key]:
        if it[0] == "p":
            j = m.paras.get((key, k))
            if j is None:
                ws.cell(row, db, 1).font = wp.font(color=wp.MUTED)
                ws.cell(row, db + 1, "새 문단").font = wp.font(color=RED)
                count(row, [1])
            else:
                orow = lay.para_row[(PREV + old_key, j)]
                v = 0 if old_items[j][1] == it[1] else 1
                _put(lay, ws, row, db, f"=IF(EXACT(A{row},{_q(PREV + old_key)}A{orow}),0,1)", v, font=wp.font(color=wp.MUTED))
                count(row, [v])
            continue
        t = it[1]
        mt = m.tables.get(t.id)
        q = old_tabs[mt[0]] if mt else None
        for j, h in enumerate(t.header):
            x = ws.cell(row, db + j, h if j else (h if q is not None else f"새 표: {h}"))
            x.font, x.alignment = wp.font(bold=True, color=wp.MUTED), Alignment(horizontal="center")
        for i in range(len(t.rows)):
            r = row + 1 + i
            new_row = q is None or i not in mt[1]
            ws.cell(r, db, ("새 줄: " if new_row else "") + t.rows[i][0]).font = wp.font(color=RED if new_row else wp.MUTED)
            vals = []
            for j in range(1, len(t.header)):
                here = f"{L(j + 1)}{r}"
                v = t.nums[i][j]
                if new_row or j not in mt[2]:
                    f, val = (f"={here}", v) if isinstance(v, (int, float)) else (None, None)
                else:
                    oi, oj = mt[1][i], mt[2][j]
                    old = _a1(lay, PREV_ID + q.id, oi, oj, key)
                    w = q.nums[oi][oj]
                    both = dsd_compare.numeric(t, i, j) and dsd_compare.numeric(q, oi, oj)
                    plain = isinstance(v, (int, float)) and isinstance(w, (int, float))
                    if both:
                        f, val = (f"={here}-{old}" if plain else f"=N({here})-N({old})"), (v or 0) - (w or 0)
                    else:
                        f, val = f'=IF(EXACT({here}&"",{old}&""),0,1)', 0 if dsd_compare.cell_text(t, i, j) == dsd_compare.cell_text(q, oi, oj) else 1
                if f:
                    _put(lay, ws, r, db + j, f, round(val, 6) if isinstance(val, float) else val, font=wp.font(color=wp.MUTED), number_format=NUM)
                    vals.append(val)
            count(r, vals, 1 if new_row else 0)
        if t.rows:                                   # 바뀐 칸은 원문에서 빨간 글자
            r0, r1 = row + 1, row + len(t.rows)
            ws.conditional_formatting.add(f"B{r0}:{L(len(t.header))}{r1}",
                                          FormulaRule(formula=[f"AND(ISNUMBER({L(db + 1)}{r0}),{L(db + 1)}{r0}<>0)"], font=Font(name=wp.FONT, color=RED, bold=True)))
    if key in lay.gone_row:
        r = lay.gone_row[key]
        ws.cell(r - 1, 1, "기존 판에 있다가 없어지거나 이름이 바뀐 것").font = wp.font(bold=True, color=RED)
        for text in m.gone[key]:
            ws.cell(r, 1, text).font = wp.font(color=RED)
            ws.cell(r, hb, 1).font = wp.font(color=wp.MUTED)
            total += 1
            r += 1
    lay.changed[key] = total
    end = max(lay.last[key], FIRST_ROW)
    ws.conditional_formatting.add(f"A{FIRST_ROW}:A{end}",       # 바뀐 문단. 표 줄은 이 칸이 글자라 걸리지 않는다
                                  FormulaRule(formula=[f"AND(ISNUMBER(${L(db)}{FIRST_ROW}),${L(db)}{FIRST_ROW}<>0)"], font=Font(name=wp.FONT, color=RED, bold=True)))
    for j in range(db, hb):
        ws.column_dimensions[L(j)].width = 14
    ws.column_dimensions[L(hb)].width = 7


def _money_cells(p) -> list:
    """그 부분의 숫자 칸(0 이 아닌 금액) [(표 id, 줄, 열)]. % 칸은 금액이 아니라 세지 않는다."""
    return [(x[1].id, i, j) for x in p.items if x[0] == "t" for i in range(len(x[1].rows)) for j in range(1, len(x[1].header))
            if isinstance(x[1].nums[i][j], (int, float)) and x[1].nums[i][j] != 0 and not x[1].rate(i, j)]


def wp_cells(doc: Doc, workpapers: list[dict] | None) -> dict:
    """조서와 맞출 숫자 칸 {(표 id, 줄, 열): 조서 대사 번호}. 줄을 적지 않았으면 표 전체다."""
    out = {}
    for n, w in enumerate(workpapers or [], 1):
        t = doc.table(w.get("table", ""))
        if t is None:
            continue
        rows = [dsd_links.find_row(t, r["row"], int(r.get("row_nth", 1))) if isinstance(r, dict) else dsd_links.find_row(t, r)
                for r in w.get("rows") or []] or list(range(len(t.rows)))
        for i in rows:
            if i is None:
                continue
            for j in range(1, len(t.header)):
                if isinstance(t.nums[i][j], (int, float)) and t.nums[i][j] != 0:
                    out.setdefault((t.id, i, j), n)
    return out


def _part_sheet(wb, lay: Layout, tabs: dict, sheet_of: dict, p, checks_by_id: dict, prev: Doc | None, m: Match | None, cover: set | None = None,
                wpc: dict | None = None):
    ws = wb.create_sheet(p.key)
    cb = lay.cb[p.key]
    ws.cell(1, 1, _title(p)).font = Font(name=wp.FONT, size=12, bold=True)
    for k, name in enumerate(COLS):
        _style_head(ws.cell(HEAD_ROW, cb + k, "확인(검증 목록)" if k == K_MARK else name))
    _original(ws, lay.items[p.key])
    for key in (_money_cells(p) if cover is not None else []):      # 검증에 들어가지 않은 숫자는 회색, 조서와 맞출 숫자는 하늘색
        if key not in cover:
            _, row, col = lay.addr[key]
            ws.cell(row, col).fill = FILL_WP if wpc and key in wpc else FILL_NONE
    for row, _, it in lay.items[p.key]:
        if it[0] == "t":
            _sum_block(ws, lay, tabs, p.key, row, it[1])
    for cid, (sheet, row) in lay.check_row.items():
        if sheet != p.key:
            continue
        c = checks_by_id[cid]
        a = lambda k: f"{L(cb + k)}{row}"       # noqa: E731
        n = lay.list_row[c.id]
        mark = f"{_q(LIST)}K{n}"
        vals = {K_GROUP: (c.group, None), K_KIND: (c.kind, None), K_CELL: (_a1(lay, c.table, c.r, c.c, sheet), None),
                K_LEFT: (_left_formula(lay, tabs, c, sheet), c.left), K_LABEL: (c.label, None), K_RIGHT: (_right_formula(lay, tabs, c, sheet), c.right),
                K_DIFF: (f"={a(K_LEFT)}-{a(K_RIGHT)}", None if c.right is None else round(c.left - c.right, 6)),
                K_OK: (('=""' if c.dropped else "=" + _state(a(K_DIFF), c.tol, mark, bool(c.when), c.minor_limit, c.cmp)), shown(c)),
                K_SRC: (c.source, None), K_BASIS: (c.basis + (f" (단수 차이 {c.tol} 허용)" if c.tol and not c.expr else "") + (f" (조건: {c.when})" if c.when else ""), None),
                K_MARK: (f'=IF({mark}="","",{mark})', _mark_of(c)), K_ID: (c.id, None), K_UNIT: (unit_name(c.unit), None)}
        for k, (v, cached) in vals.items():
            x = _put(lay, ws, row, cb + k, v, cached, font=wp.font())
            if k in (K_LEFT, K_RIGHT, K_DIFF):
                x.number_format = NUM if not c.expr else "#,##0.######"
            if k == K_BASIS:
                x.alignment = Alignment(wrap_text=False)
    for sheet, row, c in lay.mirrors:
        if sheet != p.key:
            continue
        tid, r, k0, _ = next(x for x in c.terms if sheet_of.get(x[0]) == sheet)
        rs, rr = lay.check_row[c.id]
        vals = {K_GROUP: (c.group, None), K_KIND: (c.kind, None), K_CELL: (_a1(lay, tid, r, k0, sheet), None),
                K_LEFT: ("=" + _a1(lay, tid, r, k0, sheet), tabs[tid].nums[r][k0]),
                K_LABEL: (f"{rs} 시트 {rr}줄의 검증", None), K_OK: (f"={_q(rs)}{L(lay.cb[rs] + K_OK)}{rr}", shown(c)), K_SRC: (MIRROR, None)}
        for k, (v, cached) in vals.items():
            x = _put(lay, ws, row, cb + k, v, cached, font=wp.font(color=wp.MUTED))
            if k == K_LEFT:
                x.number_format = NUM
    end = max(lay.last[p.key], ws.max_row, FIRST_ROW)
    ok = f"{L(cb + K_OK)}{FIRST_ROW}:{L(cb + K_OK)}{end}"
    ws.conditional_formatting.add(ok, CellIsRule(operator="equal", formula=["FALSE"], font=Font(name=wp.FONT, bold=True, color=RED), fill=wp.FILL_FLAG))
    ws.conditional_formatting.add(ok, CellIsRule(operator="equal", formula=[f'"{MINOR_TEXT}"'], font=Font(name=wp.FONT, bold=True), fill=FILL_MINOR))
    ws.conditional_formatting.add(ok, CellIsRule(operator="equal", formula=[f'"{ASK_TEXT}"'], font=Font(name=wp.FONT, bold=True), fill=FILL_ASK))
    if prev is not None:
        _diff_block(ws, lay, p, prev, m)
    ws.column_dimensions["A"].width = 26
    for j in range(2, cb):
        ws.column_dimensions[L(j)].width = 15
    for k, w in zip(range(len(COLS)), (6, 12, 8, 15, 30, 15, 12, 8, 8, 30, 6, 6, 7)):
        ws.column_dimensions[L(cb + k)].width = w
    for k in (K_GROUP, K_ID):
        ws.column_dimensions[L(cb + k)].hidden = True
    ws.freeze_panes = ws.cell(FIRST_ROW, 2)
    return ws


def _light(states: list) -> object:
    """신호등 값. 검증이 없으면 빈칸, FALSE 가 있으면 0, 단수 · 확인만 있으면 0.5, 전부 맞으면 1."""
    if not states:
        return ""
    if any(s is False for s in states):
        return 0
    if any(s in (MINOR_TEXT, ASK_TEXT) for s in states):
        return 0.5
    return 1


def _overview(wb, lay: Layout, doc: Doc, checks: list[Check], unresolved: list, prev: Doc | None, prior: Doc | None, m: Match | None,
              ratios: dict | None = None, wtb: bool = False, pair: Doc | None = None):
    ws = wb[OVERVIEW]
    ws.cell(1, 1, "DSD 푸팅 오버뷰").font = Font(name=wp.FONT, size=14, bold=True)
    src = (f"새 판 {doc.source}" + (f" · 기존 판 {prev.source}" if prev is not None else "") + (f" · 작년 보고서 {prior.source}" if prior is not None else "")
           + (f" · 함께 본 보고서 {pair.source}" if pair is not None else ""))
    ws.cell(1, 3, src).font = wp.font(color=wp.MUTED)
    end = OV_ROW + len(doc.parts) - 1
    states = [shown(c) for c in checks]
    _put(lay, ws, 2, 1, f'=IF(COUNTIF({_q(LIST)}I:I,FALSE)>0,0,IF(COUNTIF({_q(LIST)}I:I,"{MINOR_TEXT}")+COUNTIF({_q(LIST)}I:I,"{ASK_TEXT}")>0,0.5,1))',
         0 if False in states else (0.5 if MINOR_TEXT in states or ASK_TEXT in states else 1))
    gone = f" · 없어진 시트 {len(m.gone_parts)}" if m is not None and m.gone_parts else ""
    n_un = sum(1 for u in unresolved if u.get("kind") not in ("전기", "조서"))
    n_lost = sum(1 for u in unresolved if u.get("kind") == "전기")       # 작년 보고서에서 줄을 찾지 못한 전기 금액은 따로 센다
    n_wtb = sum(1 for u in unresolved if u.get("kind") == "조서")
    lost = (f" · 전기 대사에서 못 찾은 줄 {n_lost}" if n_lost else "") + (f" · 정산표에서 못 찾은 과목 {n_wtb}" if n_wtb else "")
    changed = f'&" · 바뀐 곳 "&SUM(G{OV_ROW}:G{end})' if prev is not None else ""
    n_changed = sum(lay.changed.values()) if prev is not None else 0
    text = (f"틀린 곳 {states.count(False)} · 확인 {states.count(ASK_TEXT)} · 단수 차이 {states.count(MINOR_TEXT)}"
            + (f" · 바뀐 곳 {n_changed}" if prev is not None else "") + f" · 제외 {sum(1 for c in checks if c.excluded)} · 풀지 못한 연결 {n_un}{lost}{gone}")
    _put(lay, ws, 2, 2, f'="틀린 곳 "&COUNTIF({_q(LIST)}I:I,FALSE)&" · 확인 "&COUNTIF({_q(LIST)}I:I,"{ASK_TEXT}")&" · 단수 차이 "&COUNTIF({_q(LIST)}I:I,"{MINOR_TEXT}"){changed}&" · 제외 "&COUNTIF({_q(LIST)}K:K,"X")&" · 풀지 못한 연결 {n_un}{lost}{gone}"',
         text, font=wp.font(bold=True))
    ws.cell(3, 2, f"합계: 표 안 소계 · 합계 / 레퍼: 표 사이 연결과 계산(곱셈 · 나눗셈) / 전기: 작년 보고서와 일치 / 버전: 기존 판 대비. 노랑은 확인(조건이 맞을 때만 같은 식의 차이)이나 단수 차이(천원 · 백만원 표에서 {MINOR} 이하)만 있음, 버전의 노랑은 바뀐 곳 있음. 검증 범위는 숫자 칸 가운데 검증에 들어간 칸의 비율(들어가지 않은 칸은 회색, 조서와 맞출 칸은 하늘색). 볼 곳은 「볼 것」 시트에 모았다. 결론과 서명은 「표지」 시트에 적는다").font = wp.font(size=9, color=wp.MUTED)
    for j, h in enumerate(OV_HEAD + ((OV_WTB,) if wtb else ()), 1):
        _style_head(ws.cell(OV_ROW - 1, j, h))
    sheet_of = {c.id: lay.addr[(c.table, c.r, c.c)][0] for c in checks}
    by_sheet: dict = {}                                  # 시트 → [(묶음, 결과)]: 그 시트의 검증 줄과 끌어온 줄
    for c in checks:
        if c.id in lay.check_row:
            by_sheet.setdefault(lay.check_row[c.id][0], []).append((c.group, shown(c), c.kind))
    for s, _, c in lay.mirrors:
        by_sheet.setdefault(s, []).append((c.group, shown(c), c.kind))
    row = OV_ROW
    for p in doc.parts:
        cb = lay.cb[p.key]
        g, ok = f"{_q(p.key)}${L(cb + K_GROUP)}:${L(cb + K_GROUP)}", f"{_q(p.key)}${L(cb + K_OK)}:${L(cb + K_OK)}"
        x = ws.cell(row, 1, p.key)
        _link(x, f"{_q(p.key)}A1")
        ws.cell(row, 2, p.title).font = wp.font()
        if ratios and p.key in ratios:
            y = ws.cell(row, 8, ratios[p.key])
            y.number_format, y.alignment, y.font = "0%", Alignment(horizontal="center"), wp.font(color=wp.MUTED if ratios[p.key] >= 0.999 else RED)
        lb, lm, li = (f"{_q(LIST)}${k}:${k}" for k in "BMI")       # 합계 검증은 표 옆 칸에 있어 검증 목록 시트에서 센다
        own = f'{lb},"{p.key}",{lm},"{SUM}"'
        sums = [shown(c) for c in checks if c.group == SUM and sheet_of[c.id] == p.key]
        v = "" if not sums else (0 if False in sums else (0.5 if MINOR_TEXT in sums else 1))
        _put(lay, ws, row, 3, f'=IF(COUNTIFS({own})=0,"",IF(COUNTIFS({own},{li},FALSE)>0,0,IF(COUNTIFS({own},{li},"{MINOR_TEXT}")>0,0.5,1)))', v,
             alignment=Alignment(horizontal="center"))
        rows_here = by_sheet.get(p.key, [])
        for j, name in ((4, "레퍼"), (5, "전기")):
            if name == "전기" and prior is None:
                continue
            if name == "레퍼" and wtb:
                continue
            f = (f'=IF(COUNTIF({g},"{name}")=0,"",IF(COUNTIFS({g},"{name}",{ok},FALSE)>0,0,'
                 f'IF(COUNTIFS({g},"{name}",{ok},"{MINOR_TEXT}")+COUNTIFS({g},"{name}",{ok},"{ASK_TEXT}")>0,0.5,1)))')
            _put(lay, ws, row, j, f, _light([s for grp, s, _ in rows_here if grp == name]), alignment=Alignment(horizontal="center"))
        if wtb:                                          # 정산표 대사는 검증 목록의 구분 열로 갈라 센다
            lc = f"{_q(LIST)}$C:$C"
            mine = [c for c in checks if sheet_of[c.id] == p.key]
            for j, cond, pick in ((4, f'{lb},"{p.key}",{lm},"레퍼",{lc},"<>조서-재무제표"', lambda c: c.group == "레퍼" and c.kind != "조서-재무제표"),
                                  (9, f'{lb},"{p.key}",{lc},"조서-재무제표"', lambda c: c.kind == "조서-재무제표")):
                f = (f'=IF(COUNTIFS({cond})=0,"",IF(COUNTIFS({cond},{li},FALSE)>0,0,'
                     f'IF(COUNTIFS({cond},{li},"{MINOR_TEXT}")+COUNTIFS({cond},{li},"{ASK_TEXT}")>0,0.5,1)))')
                _put(lay, ws, row, j, f, _light([shown(c) for c in mine if pick(c)]), alignment=Alignment(horizontal="center"))
        if prev is not None:
            h = f"{_q(p.key)}${L(lay.hb(p.key))}:${L(lay.hb(p.key))}"
            n = lay.changed.get(p.key, 0)
            _put(lay, ws, row, 6, f"=IF(SUM({h})=0,1,0.5)", 1 if n == 0 else 0.5, alignment=Alignment(horizontal="center"))
            _put(lay, ws, row, 7, f'=IF(SUM({h})=0,"",SUM({h}))', "" if n == 0 else n, font=wp.font(bold=True, color=RED), alignment=Alignment(horizontal="center"))
        row += 1
    ws.conditional_formatting.add(f"C{OV_ROW}:F{max(end, OV_ROW)}", IconSetRule("3TrafficLights1", "num", [0, 0.5, 1], showValue=False))
    ws.conditional_formatting.add("A2", IconSetRule("3TrafficLights1", "num", [0, 0.5, 1], showValue=False))
    if wtb:
        ws.conditional_formatting.add(f"I{OV_ROW}:I{max(end, OV_ROW)}", IconSetRule("3TrafficLights1", "num", [0, 0.5, 1], showValue=False))
    for col, w in zip("ABCDEFGHI", (16, 34, 8, 8, 8, 8, 9, 10, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{OV_ROW}"


def _list(wb, lay: Layout, tabs: dict, checks: list[Check]):
    """검증 전부를 한 줄씩. 표 사이 연결과 전기 대사는 그 시트의 검증 줄을 가리키고, 합계 검증은 여기서 직접 계산한다.
    확인 칸(K)은 회계사가 O · X 를 적는 유일한 칸이다. X 를 적으면 그 검증의 결과 칸이 빈칸이 된다(합계는 표 옆 칸, 연결은 그 시트의 검증 줄)."""
    ws = wb.create_sheet(LIST)
    for j, h in enumerate(LIST_HEAD, 1):
        _style_head(ws.cell(1, j, h))
    for c in checks:
        n = lay.list_row[c.id]
        cell = f"{L(lay.addr[(c.table, c.r, c.c)][2])}{lay.addr[(c.table, c.r, c.c)][1]}"
        diff = None if c.right is None else round(c.left - c.right, 6)
        if c.group == SUM:
            sheet = lay.addr[(c.table, c.r, c.c)][0]
            vals = [(n - 1, None), (sheet, None), (c.kind, None), (cell, None), (c.label, None), (_left_formula(lay, tabs, c, LIST), c.left),
                    (_right_formula(lay, tabs, c, LIST), c.right), (f"=F{n}-G{n}", diff),
                    (('=""' if c.dropped else "=" + _state(f"H{n}", c.tol, f"K{n}", False, c.minor_limit)), shown(c))]
        else:
            sheet, row = lay.check_row[c.id]
            cb = lay.cb[sheet]
            ref = lambda k: f"={_q(sheet)}{L(cb + k)}{row}"      # noqa: E731
            vals = [(n - 1, None), (sheet, None), (c.kind, None), (cell, None), (c.label, None), (ref(K_LEFT), c.left), (ref(K_RIGHT), c.right),
                    (ref(K_DIFF), diff), (ref(K_OK), shown(c))]
        vals += [(c.source, None), (_mark_of(c) or None, None), (c.id, None), (c.group, None), (c.basis, None), (c.when, None),
                 (c.finding + (" (판정으로 뺀 검증)" if c.dropped else ""), None),
                 ("0 = 0" if c.ok and not c.left and not c.right else None, None), (unit_name(c.unit), None)]
        for j, (v, cached) in enumerate(vals, 1):
            x = _put(lay, ws, n, j, v, cached, font=wp.font())
            if j in (6, 7, 8):
                x.number_format = NUM if not c.expr else "#,##0.######"
            if j == 11:
                x.fill = wp.FILL_IN
            if j in (14, 15, 16):
                x.alignment = Alignment(wrap_text=True, vertical="top")
        _link(ws.cell(n, 2), f"{_q(lay.addr[(c.table, c.r, c.c)][0])}{cell}")      # 시트 이름을 누르면 그 칸으로 간다
    end = max(len(checks) + 1, 2)
    ws.conditional_formatting.add(f"I2:I{end}", CellIsRule(operator="equal", formula=["FALSE"], font=Font(name=wp.FONT, bold=True, color=RED), fill=wp.FILL_FLAG))
    ws.conditional_formatting.add(f"I2:I{end}", CellIsRule(operator="equal", formula=[f'"{MINOR_TEXT}"'], font=Font(name=wp.FONT, bold=True), fill=FILL_MINOR))
    ws.conditional_formatting.add(f"I2:I{end}", CellIsRule(operator="equal", formula=[f'"{ASK_TEXT}"'], font=Font(name=wp.FONT, bold=True), fill=FILL_ASK))
    ws.auto_filter.ref = f"A1:{L(len(LIST_HEAD))}{end}"
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLMNOPQR", (6, 14, 12, 8, 34, 15, 15, 12, 8, 8, 6, 40, 8, 50, 40, 50, 7, 7)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["L"].hidden = True
    print_setup(ws, 1)


def _first_sentence(text: str) -> str:
    t = re.sub(r"^[^:]{1,20}:\s*", "", str(text or ""))      # 「보고서 불일치: 」 를 뗀다
    return re.split(r"(?<=다)\.\s*", t, maxsplit=1)[0].strip()


def causes(doc: Doc, checks: list[Check]) -> dict:
    """FALSE 를 원인으로 묶는다. {검증 id: (묶음 번호, 원인 칸 (표, 줄, 열) 또는 None, 그 묶음의 FALSE 수, 원인 글)}.
    1) 판정이 원인(cause)을 적었으면 같은 원인끼리. 2) FALSE 여럿에 함께 든 칸(한 칸이 틀리면 그 칸이 든 검증이 함께 FALSE 가 된다).
    3) 판정 글의 첫 문장이 같은 것끼리. 4) 나머지는 혼자 한 묶음.
    9단계 재평가는 FALSE 9줄을 원인 셋으로 묶는 일을 사람이 했고, 10단계 재평가는 같은 칸일 때만 묶여
    판정 글이 「같은 원인」 이라 한 줄이 따로 번호를 받는다고 지적했다."""
    tabs = {t.id for t in doc.tables()}
    bad = [c for c in checks if dsd_links.failed(c) and not c.excluded]
    order = sorted(bad, key=lambda c: -abs((c.left or 0) - (c.right or 0)))
    out, n, left = {}, 0, {c.id for c in bad}
    by_cause: dict = {}
    for c in order:
        if c.cause:
            by_cause.setdefault(c.cause, []).append(c.id)
    for text, ids in by_cause.items():
        n += 1
        for cid in ids:
            out[cid] = (n, None, len(ids), text)
        left -= set(ids)
    cells = {c.id: {(c.table, c.r, c.c)} | {(t, r, k) for t, r, k, _ in c.terms if t in tabs} for c in bad}
    while left:
        count: dict = {}
        for cid in left:
            for k in cells[cid]:
                count[k] = count.get(k, 0) + 1
        best = max(count.items(), key=lambda kv: (kv[1], -min(i for i, c in enumerate(order) if c.id in left and kv[0] in cells[c.id])), default=None)
        if best is None or best[1] < 2:
            break
        n += 1
        group = [cid for cid in left if best[0] in cells[cid]]
        for cid in group:
            out[cid] = (n, best[0], len(group), "")
        left -= set(group)
    by_text: dict = {}
    for c in order:
        if c.id in left and _first_sentence(c.finding):
            by_text.setdefault(_first_sentence(c.finding), []).append(c.id)
    for text, ids in by_text.items():
        if len(ids) < 2:
            continue
        n += 1
        for cid in ids:
            out[cid] = (n, None, len(ids), text)
        left -= set(ids)
    for c in order:
        if c.id in left:
            n += 1
            out[c.id] = (n, None, 1, "")
    return out


def last_year(doc: Doc, tabs: dict, checks: list[Check]) -> dict:
    """「볼 것」 의 작년 보고서 열. {검증 id: 글}. 전기 대사가 짝지은 칸(당기 보고서의 전기 칸 ↔ 작년 보고서의 당기 칸)을 써서
    FALSE · 확인 검증의 칸이 모두 작년 보고서에 있으면 같은 식을 작년 금액으로 다시 계산한다. 차이가 같으면 작년에서 넘어온 차이다."""
    cellmap = {}
    for c in checks:
        if c.group == "전기" and c.terms:
            tid, r, k, _ = c.terms[0]
            cellmap[(c.table, c.r, c.c)] = (tid, r, k)
    out = {}
    for c in checks:
        if c.group == "전기" or c.ok is not False or c.expr:
            continue
        here = [(c.table, c.r, c.c)] + [(t, r, k) for t, r, k, _ in c.terms]
        if not all(x in cellmap for x in here):
            if (c.table, c.r, c.c) in cellmap:
                t, r, k = cellmap[(c.table, c.r, c.c)]
                v = tabs[t].nums[r][k]
                if isinstance(v, (int, float)):
                    out[c.id] = f"작년 보고서의 이 칸 {v:,.0f}" if float(v).is_integer() else f"작년 보고서의 이 칸 {v:,}"
            continue

        def won(x):
            t, r, k = cellmap[x]
            v = tabs[t].nums[r][k]
            return (v if isinstance(v, (int, float)) else 0) * tabs[t].unit_at(r, k)

        diff = won(here[0]) - sum(s * won(x) for x, (_, _, _, s) in zip(here[1:], c.terms))
        now = (c.left or 0) - (c.right or 0)
        dshow = diff / c.unit if c.unit else diff
        if abs(dshow) <= max(c.tol, 0.5):
            out[c.id] = "작년 보고서에서는 맞았다"
        elif abs(dshow - now) <= max(c.tol, 1):
            out[c.id] = f"작년 보고서에도 같은 차이 {dshow:,.0f}(작년에서 넘어온 차이)"
        else:
            out[c.id] = f"작년 보고서의 차이 {dshow:,.0f}"
    return out


def _todo(wb, lay: Layout, doc: Doc, tabs: dict, checks: list[Check]) -> None:
    """「볼 것」 시트: TRUE 가 아닌 검증만 모은다. FALSE(원인으로 묶음), 확인, 작년 보고서의 오류, 늘 나는 차이, 근거가 약한 연결, 단수 순이고
    맨 아래에 판정으로 뺀 검증과 회계사가 X 를 적은 검증을 흐린 글씨로 둔다(무엇을 뺐는지 한 장에서 보이게).
    판정은 셋째 열이다(10단계 재평가: 열한째 열이라 노트북 화면에서 가로로 넘겨야 보였다).
    상태와 금액은 검증 목록 시트를 가리키는 수식이라 원문 칸을 고치면 따라 바뀐다. 줄 구성은 만든 때의 것이다."""
    ws = wb.create_sheet(TODO, 1)
    cause = causes(doc, checks)
    prior = last_year(doc, tabs, checks)
    gsize: dict = {}
    for c in checks:
        if c.id in cause:
            g = cause[c.id][0]
            gsize[g] = max(gsize.get(g, 0), abs((c.left or 0) - (c.right or 0)))

    def rank(c):
        if c.dropped or c.excluded:
            return (8, 0, 0)
        if dsd_links.failed(c):
            g = cause[c.id][0]
            return (0, -gsize[g], g)
        if dsd_links.asks(c):
            order = {dsd_links.DISPLAY: 1, dsd_links.EXPLAIN: 2, dsd_links.LAST_YEAR: 4}
            if c.when in order:
                return (order[c.when], 0, 0)
            return ((6 if c.weak else (5 if c.usual else 3)), 0, 0)
        return (7, 0, 0)

    rows = sorted((c for c in checks if c.ok is False), key=lambda c: (rank(c), -abs((c.left or 0) - (c.right or 0))))
    ws.cell(1, 1, "볼 것").font = Font(name=wp.FONT, size=14, bold=True)
    live = [c for c in rows if not (c.dropped or c.excluded)]
    states = [shown(c) for c in live]
    n_group = len({cause[c.id][0] for c in live if c.id in cause})
    kinds = [dsd_links.kind_of(c) for c in live]
    judged = [c for c in live if c.ok is False and not c.minor and not dsd_links.no_review(c)]      # 늘 나는 차이 · 근거 확인 숫자 오류는 독립 검증 대상이 아니다(2026-10-22)
    tail = (f" (FALSE 는 숫자 오류, 원인 {n_group} · 표시 오류 {kinds.count('표시 오류')} · 설명 부족 {kinds.count('설명 부족')}"
            f" · 독립 검증 {sum(1 for c in judged if c.verified)}/{len(judged)}(바꿈 {sum(1 for c in judged if c.verified == '바꿈')})"
            f" · 작년 보고서의 오류 {sum(1 for c in live if c.when == dsd_links.LAST_YEAR)}"
            f" · 늘 나는 차이 {sum(1 for c in live if dsd_links.asks(c) and c.usual)}(그중 독립 검증 대신 회계사 표본 {sum(1 for c in live if dsd_links.sampled(c))})"
            f" · 엔진이 근거를 확인한 숫자 오류 {sum(1 for c in live if dsd_links.engine_checked(c))}(독립 검증 안 함)"
            f" · 판정으로 뺀 것 {sum(1 for c in rows if c.dropped)} · 회계사가 뺀 것 {sum(1 for c in rows if c.excluded and not c.dropped)})")
    end_row = len(rows) + 4
    _put(lay, ws, 1, 4, f'="FALSE "&COUNTIF(A4:A{end_row},FALSE)&" · 확인 "&COUNTIF(A4:A{end_row},"{ASK_TEXT}")&" · 단수 "&COUNTIF(A4:A{end_row},"{MINOR_TEXT}")&"{tail}"',
         f"FALSE {states.count(False)} · 확인 {states.count(ASK_TEXT)} · 단수 {states.count(MINOR_TEXT)}{tail}", font=wp.font(bold=True))
    ws.cell(2, 1, "TRUE 가 아닌 검증만 모았다. FALSE 는 숫자 오류(같아야 하는 금액이 다르고 고칠 숫자가 있음)이고 원인이 같은 줄끼리 묶었다. 표시 오류는 금액은 맞고 이름 · 기간 · 분류 표시가 틀린 것, "
                  "설명 부족은 금액은 맞고 문장이 모자란 것이다. 확인은 조건이 맞을 때만 같은 식의 차이이거나 판정이 확인으로 둔 차이다. 작년 보고서의 오류는 올해 숫자는 맞고 작년 보고서가 틀린 것이다. "
                  "단수는 천원 · 백만원 표의 3 이하 차이다. 맨 아래 흐린 줄은 판정으로 뺀 검증(걸지 않음 · 엔진 판독 · 설명됨)과 회계사가 X 를 적은 검증이다. "
                  "검증 열은 판정을 쓴 사람과 다른 사람(에이전트)이 다시 본 결과(유지 · 바꿈)다. 판정이 확인인 늘 나는 차이는 독립 검증을 하지 않는다. "
                  "회계사가 아래쪽에 모인 늘 나는 차이에서 유형마다 표본을 골라 근거를 보고 확인 칸에 O 를 적는다. 차이 설명 후보는 엔진이 차이를 보고서의 다른 숫자로 설명해 본 것이고(자리바꿈 · 부호 · 단위 · 다른 열 · 차이와 같은 칸), "
                  "고칠 칸 · 맞는 숫자는 판정이 고치라고 한 칸과 그 숫자를 받치는 세 번째 근거를 엔진이 계산한 결과다. "
                  "시트 이름을 누르면 그 칸으로, 대사 대상을 누르면 견준 칸으로, 원인을 누르면 원인 칸으로 간다. "
                  "O · X 는 「검증 목록」 시트의 확인 칸에 적는다(확인 열을 누르면 그 줄로 간다).").font = wp.font(size=9, color=wp.MUTED)
    for j, h in enumerate(TODO_HEAD, 1):
        _style_head(ws.cell(3, j, h))
    for n, c in enumerate(rows, 4):
        t = tabs[c.table]
        sheet, row, col = lay.addr[(c.table, c.r, c.c)]
        k = lay.list_row[c.id]
        li = lambda col_: f"={_q(LIST)}{col_}{k}"      # noqa: E731
        name = f"{dsd_links._labels(t)[c.r]} · {t.header[c.c]}".strip(" ·")
        note = c.basis + (f" (조건: {c.when})" if c.when else "")
        out_ = bool(c.dropped or c.excluded)
        why = ""
        if c.dropped:
            why = f"판정으로 뺌({c.dropped})"
        elif c.excluded:
            why = "회계사가 X 로 뺌"
        elif c.id in cause:
            g, cell, cnt, text = cause[c.id]
            if cell is not None:
                ct = tabs[cell[0]]
                why = f"원인 {g}: {cell[0]} {dsd_links._labels(ct)[cell[1]]} · {ct.header[cell[2]]} (FALSE {cnt}건이 이 칸을 함께 쓴다)"
            elif text and cnt > 1:
                why = f"원인 {g}: {text[:80]} (FALSE {cnt}건)"
            elif text:
                why = f"원인 {g}: {text[:80]}"
            else:
                why = f"원인 {g}: 이 검증 하나"
        elif c.when == dsd_links.LAST_YEAR:
            why = "작년 보고서의 오류"
        elif c.when == dsd_links.DISPLAY:
            why = "표시 오류(금액은 맞음)"
        elif c.when == dsd_links.EXPLAIN:
            why = "설명 부족(금액은 맞음)"
        elif dsd_links.asks(c):
            why = "근거가 약한 연결" if c.weak else ("늘 나는 차이" if c.usual else "확인할 차이")
        diff = None if c.right is None else round(c.left - c.right, 6)
        first = ("뺌", None) if out_ else (li("I"), shown(c))
        check = "" if (out_ or c.minor) else (f"{c.verified}: {c.verify_why}"[:300] if c.verified else
                                                ("독립 검증 안 함(늘 나는 차이. 회계사가 표본으로 본다)" if dsd_links.sampled(c) else
                                                 ("독립 검증 안 함(엔진이 근거 칸으로 맞는 숫자를 확인했다)" if dsd_links.engine_checked(c) else "검증 전")))
        hint = "\n".join(f"[{h['kind']}] {h['text']}" for h in c.hints[:4])
        fix = ""
        if c.fix_cell is not None and c.fix_cell[0] in tabs:
            ft = tabs[c.fix_cell[0]]
            fix = f"{c.fix_cell[0]} {dsd_links._labels(ft)[c.fix_cell[1]]} · {ft.header[c.fix_cell[2]]} → {c.correct:,g}. {c.fix_note}"
        elif dsd_links.failed(c) and not out_:
            fix = c.fix_note or "고칠 쪽 미정(판정이 고칠 칸과 근거를 적지 않음)"
        vals = (first, (why, None), (c.finding, None), (check, None), (sheet, None), (f"{L(col)}{row}", None), (name, None), (li("F"), c.left), (c.label, None),
                (li("G"), c.right), (li("H"), diff), (unit_name(c.unit), None), (hint, None), (fix, None), (note, None), (prior.get(c.id, ""), None), (None, None),
                (f'=IF({_q(LIST)}K{k}="","→ 확인 칸",{_q(LIST)}K{k})', _mark_of(c) or "→ 확인 칸"), (c.usual, None), (f"{c.group} · {c.kind}", None), (c.id, None))
        wrap = {TODO_COL[h] for h in ("원인", "판정", "검증", "이 칸", "대사 대상", "차이 설명 후보", "고칠 칸 · 맞는 숫자", "근거 · 조건", "작년 보고서")}
        nums = {TODO_COL[h] for h in ("이 칸 값", "대사 대상 값", "차이")}
        for j, (v, cached) in enumerate(vals, 1):
            x = _put(lay, ws, n, j, v, cached, font=wp.font(color=wp.MUTED) if out_ else wp.font())
            x.alignment = Alignment(vertical="top", wrap_text=j in wrap)
            if j in nums:
                x.number_format = NUM if not c.expr else "#,##0.######"
        _link(ws.cell(n, TODO_COL["시트"]), f"{_q(sheet)}{L(col)}{row}")
        far = lay.addr.get(tuple(c.terms[0][:3])) if c.terms else None       # 대사 대상의 첫 칸으로 가는 바로가기(다른 표, 작년 보고서, 정산표)
        if far is not None:
            _link(ws.cell(n, TODO_COL["대사 대상"]), f"{_q(far[0])}{L(far[2])}{far[1]}")
        if c.id in cause and cause[c.id][1] is not None:
            at = lay.addr.get(cause[c.id][1])
            if at is not None:
                _link(ws.cell(n, TODO_COL["원인"]), f"{_q(at[0])}{L(at[2])}{at[1]}")
        _link(ws.cell(n, TODO_COL["확인"]), f"{_q(LIST)}K{k}")
        ws.cell(n, TODO_COL["검토 메모"]).fill = wp.FILL_IN
    end = max(len(rows) + 3, 4)
    rng = f"A4:A{end}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["FALSE"], font=Font(name=wp.FONT, bold=True, color=RED), fill=wp.FILL_FLAG))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{ASK_TEXT}"'], font=Font(name=wp.FONT, bold=True), fill=FILL_ASK))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{MINOR_TEXT}"'], font=Font(name=wp.FONT, bold=True), fill=FILL_MINOR))
    ws.auto_filter.ref = f"A3:{L(len(TODO_HEAD))}{end}"
    ws.freeze_panes = "E4"
    for j, w in enumerate((8, 24, 48, 22, 11, 7, 26, 15, 30, 15, 13, 7, 44, 34, 36, 26, 24, 10, 16, 14, 30), 1):
        ws.column_dimensions[L(j)].width = w
    for h in ("묶음", "id"):
        ws.column_dimensions[L(TODO_COL[h])].hidden = True
    print_setup(ws, 3)


def print_setup(ws, head_row: int) -> None:
    """가로 인쇄, 너비 한 장, 머리줄 되풀이(10단계 재평가: 인쇄 설정이 없어 PDF 로 편철하기 어려웠다)."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head_row}:{head_row}"


_AMOUNT = re.compile(r"\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?%")
_NUMS = re.compile(r"\d[\d,]*(?:\.\d+)?")


def _rich_diff(old: str, new: str, width: int = 70) -> tuple:
    """두 글의 다른 글자를 표시한 (기존 글, 새 글). 바뀐 글자는 빨간 굵은 글씨, 지운 글자는 빨간 취소선이다.
    글이 길면 바뀐 곳 앞뒤 width 자만 남기고 「…」 로 줄인다(9단계 재평가: 300자로 잘린 두 문단에서 다른 글자를 찾을 수 없었다)."""
    sm = SequenceMatcher(None, old, new, autojunk=False)
    ops = sm.get_opcodes()
    keep_o, keep_n = set(), set()
    for tag, a0, a1, b0, b1 in ops:
        if tag != "equal":
            keep_o |= set(range(max(0, a0 - width), min(len(old), a1 + width)))
            keep_n |= set(range(max(0, b0 - width), min(len(new), b1 + width)))
    plain = InlineFont(rFont=wp.FONT, sz=10)
    hot = InlineFont(rFont=wp.FONT, sz=10, b=True, color=RED)
    gone = InlineFont(rFont=wp.FONT, sz=10, strike=True, color=RED)

    def build(text, keep, spans):
        out: list = []                            # [(글꼴, 글)]. 같은 글꼴이 이어지면 한 덩이로 붙인다
        for a, b, f in spans:
            for i in range(a, b):
                piece, font = (text[i], f) if i in keep else (" … ", plain)
                if i not in keep and out and out[-1][1].endswith(" … "):
                    continue
                if out and out[-1][0] is font:
                    out[-1] = (font, out[-1][1] + piece)
                else:
                    out.append((font, piece))
        blocks = [TextBlock(f, t) for f, t in out if t]
        return CellRichText(*blocks) if blocks else ""

    so = [(a0, a1, plain if tag == "equal" else gone) for tag, a0, a1, b0, b1 in ops]
    sn = [(b0, b1, plain if tag == "equal" else hot) for tag, a0, a1, b0, b1 in ops]
    return build(old, keep_o, so), build(new, keep_n, sn)


def _changes_sheet(wb, lay: Layout, doc: Doc, prev: Doc, m: Match, compare: list | None = None) -> None:
    """「바뀐 곳」 시트: 기존 판 대비 바뀐 칸과 문단을 한 장에. 숫자는 기존 값 · 새 값 · 차이, 글자는 기존 글과 새 글(다른 글자에 표시).
    숫자가 든 글이 바뀌면(지분율, 문장 속 금액) 「숫자 든 글 바뀜」 으로 글 바뀜보다 앞에 둔다. 새 표는 줄마다 한 줄로 푼다.
    compare 를 주면 기존 판과 새 판의 검증 결과가 달라진 것을 맨 아래에 모은다(고쳐 달라고 한 곳이 고쳐졌는지)."""
    ws = wb.create_sheet(CHANGES, 3)
    ws.cell(1, 1, "바뀐 곳").font = Font(name=wp.FONT, size=14, bold=True)
    ws.cell(2, 1, "기존 판과 견줘 바뀐 숫자와 글이다. 숫자가 바뀐 칸, 숫자가 든 글, 그 밖의 글, 새로 생긴 것, 없어진 것 순이고 숫자는 차이가 큰 순이다. "
                  "글은 바뀐 글자를 빨갛게 적었고 긴 글은 바뀐 곳 앞뒤만 남겼다. 시트 이름을 누르면 그 칸으로 간다. 맨 아래에 두 판의 검증 결과가 달라진 것을 모았다. "
                  "오버뷰의 바뀐 곳 수는 칸마다 세어 이 시트의 줄 수보다 크다.").font = wp.font(size=9, color=wp.MUTED)
    for j, h in enumerate(("종류", "시트", "칸", "표 · 줄 · 열", "기존 값", "새 값", "차이"), 1):
        _style_head(ws.cell(3, j, h))
    olds = {p.key: p for p in prev.parts}
    rows = []

    def text_change(where, name, old, new, sheet):
        a, b = _NUMS.findall(old), _NUMS.findall(new)
        kind, rank = ("숫자 든 글 바뀜", 2) if a != b and (a or b) else ("글 바뀜", 3)
        rows.append((rank, kind, sheet, where, name, old, new, None))

    for p in doc.parts:
        op = olds.get(m.parts.get(p.key))
        if op is None:
            rows.append((5, "새 주석", p.key, "A1", p.title, "", "", None))
            continue
        ot = {x[1].id: x[1] for x in op.items if x[0] == "t"}
        for i, x in enumerate(p.items):
            if x[0] == "p":
                j = m.paras.get((p.key, i))
                where = f"A{lay.para_row[(p.key, i)]}"
                if j is None:
                    rows.append((4, "새 문단", p.key, where, "", "", x[1][:300], None))
                elif op.items[j][1] != x[1]:
                    text_change(where, "", op.items[j][1], x[1], p.key)
                continue
            t = x[1]
            mt = m.tables.get(t.id)
            if not mt:
                for r in range(len(t.rows)):
                    rows.append((4, "새 표", p.key, _a1(lay, t.id, r, 0, p.key), f"{t.id} {t.rows[r][0]}", "",
                                 " · ".join(f"{t.header[c]} {t.rows[r][c]}".strip() for c in range(1, len(t.header)) if t.rows[r][c])[:200], None))
                continue
            q = ot[mt[0]]
            for r in range(len(t.rows)):
                if r not in mt[1]:
                    rows.append((4, "새 줄", p.key, _a1(lay, t.id, r, 0, p.key), f"{t.id} {t.rows[r][0]}", "",
                                 " · ".join(f"{t.header[c]} {t.rows[r][c]}".strip() for c in range(1, len(t.header)) if t.rows[r][c])[:200], None))
                    continue
                for c in range(1, len(t.header)):
                    if c not in mt[2]:
                        continue
                    orow, ocol = mt[1][r], mt[2][c]
                    v, w = t.nums[r][c], q.nums[orow][ocol]
                    where, name = _a1(lay, t.id, r, c, p.key), f"{t.id} {t.rows[r][0]} · {t.header[c]}"
                    if dsd_compare.numeric(t, r, c) and dsd_compare.numeric(q, orow, ocol):
                        if (v or 0) != (w or 0):
                            rows.append((1, "숫자 바뀜", p.key, where, name, w or 0, v or 0, (v or 0) - (w or 0)))
                    elif dsd_compare.cell_text(t, r, c) != dsd_compare.cell_text(q, orow, ocol):
                        text_change(where, name, dsd_compare.cell_text(q, orow, ocol), dsd_compare.cell_text(t, r, c), p.key)
        for g in m.gone.get(p.key, []):
            rows.append((6, g.split(":")[0], p.key, "A1", g.split(": ", 1)[-1][:200], "", "", None))
    rows.sort(key=lambda r: (r[0], -abs(r[7] or 0)))
    n = 4
    for _, kind, sheet, where, name, old, new, diff in rows:
        if kind in ("글 바뀜", "숫자 든 글 바뀜"):
            old, new = _rich_diff(str(old), str(new))
        for j, v in enumerate((kind, sheet, where, name, old, new, diff), 1):
            x = ws.cell(n, j, v)
            x.font = wp.font(color=RED, bold=True) if kind == "숫자 든 글 바뀜" and j == 1 else wp.font()
            x.alignment = Alignment(vertical="top", wrap_text=j in (4, 5, 6))
            if isinstance(v, (int, float)):
                x.number_format = NUM
        _link(ws.cell(n, 2), f"{_q(sheet)}{where}")
        n += 1
    ws.auto_filter.ref = f"A3:G{max(n - 1, 4)}"
    if compare:
        n += 1
        ws.cell(n, 1, f"검증 결과가 달라진 것 {len(compare)}건").font = wp.font(bold=True)
        n += 1
        for j, h in enumerate(("달라진 것", "시트", "칸", "검증", "기존 판 결과", "새 판 결과", "차이(새 판)"), 1):
            _style_head(ws.cell(n, j, h))
        n += 1
        for x in compare:
            vals = (x["what"], x["sheet"], x["cell"], x["label"], _shown_text(x["old"]), _shown_text(x["new"]), x["diff"])
            for j, v in enumerate(vals, 1):
                y = ws.cell(n, j, v)
                y.font = wp.font(bold=j == 1, color=RED if x["what"] == "새로 틀림" else (GREEN if x["what"] == "고쳐짐" else None))
                y.alignment = Alignment(vertical="top", wrap_text=j == 4)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    y.number_format = NUM
            _link(ws.cell(n, 2), f"{_q(x['sheet'])}{x['cell']}")
            n += 1
    ws.freeze_panes = "A4"
    for col, w in zip("ABCDEFG", (14, 12, 7, 44, 40, 40, 16)):
        ws.column_dimensions[col].width = w
    print_setup(ws, 3)


def _shown_text(v) -> str:
    if v is True:
        return "TRUE"
    if v is False:
        return "FALSE"
    if v is None:
        return "없음(새 검증)"
    return "제외" if v == "" else str(v)


def _cell_key(c: Check, back: dict) -> tuple | None:
    """검증의 (묶음, 구분, 왼쪽 칸, 대사 대상 칸들). back 은 옛 표 id → (새 표 id, {옛 행: 새 행}, {옛 열: 새 열}) 로 옛 칸을 새 칸으로 옮긴다.
    back 이 없으면 칸을 그대로 쓴다(새 판)."""
    def cell(t, r, k):
        if back is None:
            return (t, r, k)
        if t not in back:
            return (t, r, k) if t.startswith((PRIOR_ID, WTB_ID, PAIR_ID)) else None
        nt, rows, cols = back[t]
        if r == -1:
            return (nt, -1, cols.get(k, k))
        if r not in rows or k not in cols:
            return None
        return (nt, rows[r], cols[k])
    left = cell(c.table, c.r, c.c)
    if left is None:
        return None
    terms = []
    for t, r, k, sg in c.terms:
        x = cell(t, r, k)
        if x is None:
            return None
        terms.append((x, sg))
    return (c.group, c.kind, left, tuple(sorted(terms)))


def compare_checks(lay: Layout, new: list[Check], old: list[Check], m: Match | None = None) -> list[dict]:
    """기존 판과 새 판의 검증 결과 비교. 고쳐짐 · 새로 틀림 · 여전히 틀림 · 새 검증(틀림) 을 낸다.
    m 을 주면 옛 검증의 칸을 표 · 줄 · 열 짝으로 새 칸에 옮겨 (묶음, 구분, 칸, 대사 대상 칸)으로 짝짓는다. 주석 번호가 밀린 판에서 검증 id(표 번호)로
    짝지으면 다른 검증끼리 견주게 된다(10단계 재평가). 칸으로 짝을 못 지은 것만 id 로 짝짓는다."""
    before = {c.id: c for c in old}
    by_cell: dict = {}
    if m is not None:
        back = {ot: (nt, {o: n for n, o in rows.items()}, {o: n for n, o in cols.items()}) for nt, (ot, rows, cols) in m.tables.items()}
        for o in old:
            k = _cell_key(o, back)
            if k is not None:
                by_cell.setdefault(k, o)
    out = []
    for c in new:
        o = by_cell.get(_cell_key(c, None)) if m is not None else None
        if o is None and m is None:
            o = before.get(c.id)
        a = state(o) if o is not None else None
        b = state(c)
        bad_b, bad_a = b is not True, a is not None and a is not True
        if o is None:
            if not bad_b:
                continue
            what = "새 검증에서 틀림"
        elif bad_a and not bad_b:
            what = "고쳐짐"
        elif bad_b and not bad_a:
            what = "새로 틀림"
        elif bad_a and bad_b:
            what = "여전히 틀림"
        else:
            continue
        sheet, row, col = lay.addr[(c.table, c.r, c.c)]
        out.append({"what": what, "sheet": sheet, "cell": f"{L(col)}{row}", "label": f"{c.kind} {c.table}: {c.label}", "old": a, "new": b,
                    "diff": None if c.right is None else round(c.left - c.right, 6)})
    order = {"새로 틀림": 0, "여전히 틀림": 1, "새 검증에서 틀림": 2, "고쳐짐": 3}
    return sorted(out, key=lambda x: order[x["what"]])


def _scope(wb, lay: Layout, doc: Doc, cover: set, extra: dict | None, cross: set | None = None, wpc: dict | None = None) -> dict:
    """「검증 범위」 시트: 시트마다 숫자 칸 가운데 어느 검증에든 들어간 칸의 비율, 조서와 맞출 칸, 검증 없는 칸(회색), 검증이 하나도 없는 표,
    주석을 가리키는데 대사가 없는 과목, 합계인데 검증이 없는 칸, 걸지 않은 연결, 원문 대조(표 수, 주석 번호 참조, 머리줄).
    TRUE 가 많은 것과 다 본 것은 다르다. 회색 칸과 이 시트가 보지 못한 곳을 알려 준다. {시트: 비율}"""
    ws = wb.create_sheet(SCOPE, 2)
    wpc = wpc or {}
    ws.cell(1, 1, "검증 범위").font = Font(name=wp.FONT, size=14, bold=True)
    ws.cell(2, 1, "숫자 칸(0 이 아닌 금액, % 칸 제외) 가운데 어느 검증에든 들어간 칸의 비율이다. 검증에 들어가지 않은 칸은 두 가지다. 보고서 안에 맞춰 볼 곳이 없어 조서와 맞출 칸(하늘색, "
                  "「조서 대사」 시트)과 검증이 없는 칸(회색, 눈으로 본다). 검증된 칸에는 표 안 합계에만 들어간 칸도 있어 다른 표 · 정산표와 맞춘 칸은 따로 센다. "
                  "문장 속 금액은 주석 문단에서 꺼내 「주석N-문장」 표로 모은 칸이다.").font = wp.font(size=9, color=wp.MUTED)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 45
    head = ("시트", "제목", "숫자 칸", "검증된 칸", "비율", "조서와 맞출 칸", "검증 없는 칸(회색)", "다른 표와 대사된 칸", "문장 속 금액", "그중 검증된 칸", "검증이 하나도 없는 표")
    for j, h in enumerate(head, 1):
        _style_head(ws.cell(3, j, h))
    ratios, n, bare_rows = {}, 4, []
    first = n
    for p in doc.parts:
        cells = _money_cells(p)
        if not cells:
            continue
        sent = [k for k in cells if (doc.table(k[0]) and doc.table(k[0]).kind)]
        hit = sum(1 for k in cells if k in cover)
        to_wp = sum(1 for k in cells if k not in cover and k in wpc)
        ratios[p.key] = hit / len(cells)
        bare = [x[1].id for x in p.items if x[0] == "t" and any(k[0] == x[1].id for k in cells) and not any(k in cover for k in cells if k[0] == x[1].id)]
        for x in p.items:
            if x[0] == "t" and x[1].id in bare:
                t = x[1]
                bare_rows.append((p.key, _a1(lay, t.id, -1, 0, p.key), t.id, (t.caption or " · ".join(h for h in t.header if h))[:60],
                                  sum(1 for k in cells if k[0] == t.id), "조서 대사" if all(k in wpc for k in cells if k[0] == t.id) else ""))
        _link(ws.cell(n, 1, p.key), f"{_q(p.key)}A1")
        vals = (p.title, len(cells), hit, hit / len(cells), to_wp, len(cells) - hit - to_wp, sum(1 for k in cells if cross is not None and k in cross),
                len(sent) or None, sum(1 for k in sent if k in cover) if sent else None, ", ".join(bare))
        for j, v in enumerate(vals, 2):
            y = ws.cell(n, j, v)
            y.font = wp.font()
            if j == 5:
                y.number_format = "0%"
        n += 1
    ws.cell(n, 1, "합계").font = wp.font(bold=True)
    tot = {}
    for j in (3, 4, 6, 7, 8, 9, 10):
        vals = [ws.cell(r, j).value for r in range(first, n)]
        tot[j] = sum(v for v in vals if isinstance(v, (int, float)))
        _put(lay, ws, n, j, f"=SUM({L(j)}{first}:{L(j)}{n - 1})", tot[j], font=wp.font(bold=True))
    _put(lay, ws, n, 5, f'=IF(C{n}=0,"",D{n}/C{n})', (tot[4] / tot[3]) if tot[3] else "", font=wp.font(bold=True), number_format="0%")
    n += 2
    ws.cell(n, 1, f"검증이 하나도 없는 표 {len(bare_rows)}개").font = wp.font(bold=True)
    n += 1
    if bare_rows:
        for j, h in enumerate(("시트", "칸", "표", "제목 · 머리줄", "숫자 칸", "조서 대사"), 1):
            _style_head(ws.cell(n, j, h))
        n += 1
        for sheet, where, tid, cap, cnt, w in bare_rows:
            for j, v in enumerate((sheet, where, tid, cap, cnt, w), 1):
                ws.cell(n, j, v).font = wp.font()
            _link(ws.cell(n, 1), f"{_q(sheet)}{where}")
            n += 1
    n += 1
    ex = extra or {}
    for title, head, rows in (
        ("주석을 가리키는데 대사가 없는 재무제표 과목", ("표", "과목", "당기", "전기", "가리킨 주석"),
         [(u["table"], u["row"], u.get("당기"), u.get("전기"), ", ".join(u.get("notes", []))) for u in ex.get("unlinked", [])]),
        ("합계 줄 · 열인데 합계 검증이 없는 칸", ("표", "줄 · 열", "구분", "검증이 없는 칸", ""),
         [(u.get("table"), u.get("row", u.get("col", "")), u.get("what", ""), ", ".join(u.get("cells", [])), "") for u in ex.get("unchecked_totals", [])]),
        ("AI 가 걸지 않기로 한 연결", ("연결", "이유", "", "", ""),
         [(u.get("what", ""), u.get("why", ""), "", "", "") for u in ex.get("skipped", [])]),
        ("원문 대조", ("항목", "수", "내용", "", ""),
         [(k, v, "", "", "") for k, v in (doc.stats or {}).items()] + [(x["what"], x.get("n", ""), x["why"], "", "") for x in ex.get("consistency", [])]),
    ):
        ws.cell(n, 1, f"{title} {len(rows)}{'곳' if title != '원문 대조' else '줄'}").font = wp.font(bold=True)
        n += 1
        if rows:
            for j, h in enumerate(head, 1):
                if h:
                    _style_head(ws.cell(n, j, h))
            n += 1
            for r in rows:
                for j, v in enumerate(r, 1):
                    y = ws.cell(n, j, v)
                    y.font = wp.font()
                    if j in (3, 4) and isinstance(v, (int, float)):
                        y.number_format = NUM
                n += 1
        n += 1
    ws.cell(n, 1, "원문 대조의 표 수: 원문 표는 테두리 있는 표와 테두리만 없는 숫자 표, 틀 표는 칸 안에 문단과 표를 담은 테두리 없는 표(안의 표를 읽었다), "
                  "문단으로 읽은 숫자 표는 금액이 든 테두리 없는 표인데 표로 보지 않은 것이다. 어느 부분에도 넣지 못한 표가 있으면 원문과 견준다.").font = wp.font(size=9, color=wp.MUTED)
    for col, w in zip("ABCDEFGHIJK", (16, 34, 12, 12, 9, 13, 14, 14, 12, 12, 40)):
        ws.column_dimensions[col].width = w
    print_setup(ws, 3)
    return ratios


def _wp_sheet(wb, lay: Layout, doc: Doc, workpapers: list[dict], wpc: dict, cover: set) -> None:
    """「조서 대사」 시트: 보고서 안에서는 맞춰 볼 곳이 없는 숫자(약정 한도, 담보, 특수관계자 거래, 보험)를 조서와 맞출 목록.
    AI 가 연결.json 의 workpapers 에 적은 것이다. 회계사가 조서 번호 · 결과 · 메모를 적는다."""
    ws = wb.create_sheet(WPS)
    ws.cell(1, 1, "조서 대사").font = Font(name=wp.FONT, size=14, bold=True)
    ws.cell(2, 1, "보고서 안에서는 맞춰 볼 곳이 없어 조서 · 증빙과 맞출 숫자다. 원문 시트에서 하늘색 칸이다. 조서 번호, 결과(O · X), 메모는 회계사가 적는다.").font = wp.font(size=9, color=wp.MUTED)
    head = ("시트", "표", "줄", "숫자 칸", "무엇과 맞추나", "조서(제안)", "까닭", "조서 번호", "결과", "메모")
    for j, h in enumerate(head, 1):
        _style_head(ws.cell(3, j, h))
    for n, w in enumerate(workpapers, 4):
        t = doc.table(w["table"])
        if t is None:
            continue
        sheet = doc.part_of(t.id).key
        cnt = sum(1 for k, v in wpc.items() if v == n - 3 and k not in cover)
        names = [r["row"] if isinstance(r, dict) else r for r in w.get("rows", [])]
        vals = (sheet, t.id, ", ".join(names) or "표 전체", cnt, w["what"], w.get("wp", ""), w["why"], None, None, None)
        for j, v in enumerate(vals, 1):
            x = ws.cell(n, j, v)
            x.font = wp.font()
            x.alignment = Alignment(vertical="top", wrap_text=j in (5, 7, 10))
            if j >= 8:
                x.fill = wp.FILL_IN
        at = lay.addr.get((t.id, -1, 0))
        if at is not None:
            _link(ws.cell(n, 1), f"{_q(at[0])}A{at[1]}")
    for col, w in zip("ABCDEFGHIJ", (12, 12, 24, 8, 34, 16, 40, 12, 7, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    print_setup(ws, 3)


def done_text(checks: list[Check], n_wp: int) -> str:
    """표지의 「한 일」. 이 보고서에서 실제로 건 검증만 센다(10단계 재평가: 정산표 · 작년 보고서가 없는 보고서도 같은 글이었다)."""
    live = [c for c in checks if not c.dropped]
    n = lambda f: sum(1 for c in live if f(c))      # noqa: E731
    parts = [(f"표 안 합계 {n(lambda c: c.group == SUM):,}건", n(lambda c: c.group == SUM)),
             (f"표 사이 연결 {n(lambda c: c.group == '레퍼' and not c.expr and c.kind not in ('조서-재무제표', '보고서 간')):,}건(재무제표 ↔ 주석, 주석 ↔ 주석)",
              n(lambda c: c.group == '레퍼' and not c.expr and c.kind not in ('조서-재무제표', '보고서 간'))),
             (f"곱셈 · 나눗셈 계산 {n(lambda c: bool(c.expr)):,}건", n(lambda c: bool(c.expr))),
             (f"작년 보고서의 같은 칸 {n(lambda c: c.group == '전기'):,}건", n(lambda c: c.group == '전기')),
             (f"정산표 금액 {n(lambda c: c.kind == '조서-재무제표'):,}건", n(lambda c: c.kind == '조서-재무제표')),
             (f"함께 낸 보고서(연결 ↔ 별도) {n(lambda c: c.kind == '보고서 간'):,}건", n(lambda c: c.kind == '보고서 간'))]
    did = [t for t, k in parts if k]
    text = " · ".join(did) + "을 엑셀 수식으로 견줬다. " if did else "건 검증이 없다. "
    text += "연결은 AI 가 주석을 읽고 정했고 금액은 엔진과 엑셀이 계산했다. 검증에 들어가지 않은 숫자는 원문 시트에서 회색(눈으로 볼 것)"
    text += f"과 하늘색(조서와 맞출 것, {n_wp}항목)으로 표시했다." if n_wp else "으로 표시했다."
    return text


def _cover(wb, lay: Layout, doc: Doc, checks: list[Check], prev: Doc | None, prior: Doc | None, wtb: Doc | None, pair: Doc | None,
           ratios: dict, n_wp: int, n_wpc: int = 0) -> None:
    """「표지」 시트: 조서로 편철할 때의 표지. 조서 번호 · 회사 · 결론 · 작성자 · 검토자 · 일자는 회계사가 적는다.
    결과 요약은 다른 시트를 세는 수식이다(9단계 재평가: 조서로 편철하려면 표지를 따로 만들어야 했다)."""
    ws = wb.create_sheet(COVER, 0)
    ws.cell(1, 1, "DSD 푸팅 조서(재무제표 · 주석 숫자 대사)").font = Font(name=wp.FONT, size=14, bold=True)
    live = [c for c in checks if not c.dropped]
    states = [shown(c) for c in live]
    total_cells = sum(len(_money_cells(p)) for p in doc.parts)
    hit = sum(r * len(_money_cells(p)) for p in doc.parts if (r := ratios.get(p.key)) is not None)
    rows = [
        ("조서 번호", None, True), ("회사", None, True), ("대상 파일", doc.source, False),
        ("견준 자료", " · ".join(x for x in ((f"작년 보고서 {prior.source}" if prior is not None else ""), (f"정산표 {wtb.source}" if wtb is not None else ""),
                                          (f"기존 판 {prev.source}" if prev is not None else ""), (f"함께 본 보고서 {pair.source}" if pair is not None else "")) if x) or "없음", False),
        ("한 일", done_text(checks, n_wp), False),
    ]
    r = 3
    for name, v, inp in rows:
        ws.cell(r, 1, name).font = wp.font(bold=True)
        x = ws.cell(r, 2, v)
        x.font, x.alignment = wp.font(), Alignment(wrap_text=True, vertical="top")
        if inp:
            x.fill = wp.FILL_IN
        r += 1
    ws.row_dimensions[r - 1].height = 60
    r += 1
    ws.cell(r, 1, "결과").font = wp.font(bold=True, size=11)
    r += 1
    li = f"{_q(LIST)}I:I"
    summary = [
        ("검증 수", f"=COUNTA({_q(LIST)}L:L)-1", len(checks)),
        ("판정으로 뺀 검증", None, len(checks) - len(live)),
        ("FALSE(숫자 오류)", f"=COUNTIF({li},FALSE)", states.count(False)),
        ("표시 오류(금액은 맞음)", None, sum(1 for c in live if dsd_links.kind_of(c) == "표시 오류")),
        ("설명 부족(금액은 맞음)", None, sum(1 for c in live if dsd_links.kind_of(c) == "설명 부족")),
        ("독립 검증을 거친 판정", None, f"{sum(1 for c in live if c.ok is False and not c.minor and not dsd_links.no_review(c) and c.verified)} / "
                                     f"{sum(1 for c in live if c.ok is False and not c.minor and not dsd_links.no_review(c))}"),
        ("늘 나는 차이(독립 검증 대신 회계사 표본)", None, sum(1 for c in live if dsd_links.sampled(c))),
        ("엔진이 근거를 확인한 숫자 오류(독립 검증 안 함)", None, sum(1 for c in live if dsd_links.engine_checked(c))),
        ("확인", f'=COUNTIF({li},"{ASK_TEXT}")', states.count(ASK_TEXT)),
        ("단수 차이", f'=COUNTIF({li},"{MINOR_TEXT}")', states.count(MINOR_TEXT)),
        ("회계사가 뺀 검증(X)", f'=COUNTIF({_q(LIST)}K:K,"X")', sum(1 for c in checks if c.excluded)),
        ("숫자 칸", None, total_cells),
        ("검증된 칸의 비율", None, round(hit / total_cells, 4) if total_cells else ""),
        ("조서와 맞출 칸을 뺀 비율", None, round(hit / (total_cells - n_wpc), 4) if total_cells - n_wpc > 0 else ""),
        ("조서와 맞출 항목", None, n_wp),
    ]
    for name, f, v in summary:
        ws.cell(r, 1, name).font = wp.font()
        x = _put(lay, ws, r, 2, f if f else v, v if f else None, font=wp.font(bold=True))
        if name.endswith("비율"):
            x.number_format = "0.0%"
        r += 1
    r += 1
    ws.cell(r, 1, "결론").font = wp.font(bold=True, size=11)
    r += 1
    x = ws.cell(r, 1, None)
    x.fill = wp.FILL_IN
    x.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 90
    r += 2
    for name in ("작성자", "작성일", "검토자", "검토일"):
        ws.cell(r, 1, name).font = wp.font(bold=True)
        ws.cell(r, 2).fill = wp.FILL_IN
        r += 1
    r += 1
    ws.cell(r, 1, "노란 칸은 회계사가 적는다. FALSE 와 확인의 판정은 「볼 것」 시트, 검증에 들어가지 않은 숫자는 「검증 범위」 시트, 조서와 맞출 숫자는 「조서 대사」 시트에 있다.").font = wp.font(size=9, color=wp.MUTED)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _inject(path: Path, wb, cache: dict) -> None:
    """저장한 xlsx 의 수식 칸에 계산값을 넣는다. openpyxl 은 수식만 쓰고 값 자리를 비워 둔다(<v></v>).
    엑셀은 열 때 다시 계산하고(fullCalcOnLoad), 엑셀이 아닌 미리보기는 이 값을 보여 준다."""
    by: dict = {}
    for (sheet, r, c), v in cache.items():
        by.setdefault(sheet, {})[f"{L(c)}{r}"] = v
    files = {f"xl/worksheets/sheet{i + 1}.xml": ws.title for i, ws in enumerate(wb.worksheets)}
    cell = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)
    tmp = Path(path).with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            sheet = files.get(item.filename)
            if sheet is not None:                     # 글자 서식 조각의 공백을 지킨다. 공백만 든 조각(「 」)이 지워지면 엑셀이 파일을 복구한다(10단계 재평가)
                data = data.replace(b"<t>", b'<t xml:space="preserve">')
            if sheet in by:
                vals = by[sheet]

                def fill(m):
                    v = vals.get(m.group(1))
                    attrs = re.sub(r'\s+t="[^"]*"', "", m.group(2))
                    if v is None:
                        return m.group(0)
                    if isinstance(v, bool):
                        return f'<c r="{m.group(1)}"{attrs} t="b"><f>{m.group(3)}</f><v>{int(v)}</v></c>'
                    if isinstance(v, (int, float)):
                        return f'<c r="{m.group(1)}"{attrs}><f>{m.group(3)}</f><v>{v!r}</v></c>'
                    return f'<c r="{m.group(1)}"{attrs} t="str"><f>{m.group(3)}</f><v>{escape(str(v))}</v></c>'

                data = cell.sub(fill, data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def write(path: Path, doc: Doc, checks: list[Check], unresolved: list[dict], prev: Doc | None = None, prior: Doc | None = None,
          m: Match | None = None, extra: dict | None = None, wtb: Doc | None = None, pair: Doc | None = None, old_checks: list[Check] | None = None) -> Path:
    """엑셀을 쓴다. prev(기존 판)를 주면 m = dsd_compare.match(doc, prev) 도 함께 준다. old_checks 는 기존 판에서 같은 연결로 계산한 검증이다.
    extra 는 「검증 범위」 시트에 적을 것이다({unlinked, unchecked_totals}: dsd_ai.request 의 결과, skipped: 걸지 않은 연결, workpapers: 조서와 맞출 숫자,
    consistency: 원문 대조)."""
    lay = place(doc, checks, prev, prior, m, wtb, pair)
    cover = dsd_links.covered(checks)
    tabs = {t.id: t for t in doc.tables()}
    for other, pre in ((prior, PRIOR_ID), (wtb, WTB_ID), (pair, PAIR_ID)):
        if other is not None:
            tabs.update({pre + t.id: t for t in other.tables()})
    sheet_of = {t.id: doc.part_of(t.id).key for t in doc.tables()}
    workpapers = list((extra or {}).get("workpapers", []))
    wpc = wp_cells(doc, workpapers)
    wb = Workbook()
    wb.active.title = OVERVIEW
    by_id = {c.id: c for c in checks}
    for p in doc.parts:
        _part_sheet(wb, lay, tabs, sheet_of, p, by_id, prev, m, cover, wpc)
    _list(wb, lay, tabs, checks)
    _todo(wb, lay, doc, tabs, checks)
    ratios = _scope(wb, lay, doc, cover, extra, dsd_links.covered([c for c in checks if c.group != SUM]), wpc)
    if prev is not None and m is not None:
        _changes_sheet(wb, lay, doc, prev, m, compare_checks(lay, checks, old_checks, m) if old_checks is not None else None)
    if workpapers:
        _wp_sheet(wb, lay, doc, workpapers, wpc, cover)
    if unresolved or (m is not None and m.gone_parts):
        ws = wb.create_sheet(UNRESOLVED)
        for j, h in enumerate(("표", "구분", "사유"), 1):
            _style_head(ws.cell(1, j, h))
        rows = list(unresolved) + [{"table": x, "kind": "버전", "why": "기존 판에 있다가 새 판에 없는 시트"} for x in (m.gone_parts if m is not None else [])]
        for n, u in enumerate(rows, 2):
            for j, k in enumerate(("table", "kind", "why"), 1):
                ws.cell(n, j, u[k]).font = wp.font()
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["C"].width = 70
    for other, prefix, note in ((prev, PREV, "기존 판"), (prior, PRIOR, "작년 보고서"), (wtb, WTB, "정산표"), (pair, PAIR, "함께 본 보고서")):
        if other is not None:
            for p in other.parts:
                _plain_sheet(wb, lay, p, prefix, note)
    _overview(wb, lay, doc, checks, unresolved, prev, prior, m, ratios, wtb is not None, pair)
    tres = (extra or {}).get("text")
    if tres is not None:                             # 문장 검토(2026-10-17, dsd_text): 시트와 오버뷰 한 줄
        from . import dsd_text

        dsd_text.add_sheet(wb, tres)
        r = OV_ROW + len(doc.parts) + 1
        c = wb[OVERVIEW].cell(r, 1, dsd_text.SHEET)
        _link(c, f"{_q(dsd_text.SHEET)}A1")
        wb[OVERVIEW].cell(r, 2, dsd_text.summary(tres)).font = wp.font(bold=True)
    n_wpc = sum(1 for p in doc.parts for k in _money_cells(p) if k not in cover and k in wpc)   # 숫자 칸만 센다(2026-10-20: 글자 칸까지 세어 비율이 100% 를 넘었다)
    _cover(wb, lay, doc, checks, prev, prior, wtb, pair, ratios, len(workpapers), n_wpc)
    wb.active = wb.sheetnames.index(OVERVIEW)        # 표지가 맨 앞이지만 열면 오버뷰가 보인다
    wb.calculation.fullCalcOnLoad = True
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _inject(path, wb, lay.cache)
    return path


def read_marks(path: Path) -> dict:
    """지난 엑셀의 확인 칸을 읽는다. X 는 쓰지 않음, O 는 회계사가 확인한 연결.
    지금은 「검증 목록」 시트의 확인 칸이 유일한 칸이다. 예전 엑셀(표 사이 연결의 확인을 주석 시트의 검증 줄에 적던 판)도 읽는다."""
    out = {}
    if not Path(path).exists():
        return out
    wb = load_workbook(path, read_only=True)          # 값만 읽는다. 한 상장사 반기 별도 엑셀 1.3초 → 0.5초(2026-10-22 실행 시간 P5)

    def at(row, j):
        return row[j] if j < len(row) else None

    try:
        if LIST in wb.sheetnames:
            rows = wb[LIST].iter_rows(values_only=True)
            heads = {v: j for j, v in enumerate(next(rows, ()))}
            if {"확인", "id"} <= set(heads):
                for row in rows:
                    cid, mark = at(row, heads["id"]), at(row, heads["확인"])
                    if cid and isinstance(mark, str) and mark.strip().upper() in ("X", "O"):
                        out[cid] = {"excluded": True} if mark.strip().upper() == "X" else {"confirmed": True}
        for ws in wb.worksheets:
            head = next(ws.iter_rows(min_row=HEAD_ROW, max_row=HEAD_ROW, values_only=True), ())
            heads = {v: j for j, v in enumerate(head)}
            if "확인" not in heads or "id" not in heads or heads.get("출처") is None:
                continue
            for row in ws.iter_rows(min_row=FIRST_ROW, values_only=True):
                cid, mark, src = at(row, heads["id"]), at(row, heads["확인"]), at(row, heads["출처"])
                if not cid or src == MIRROR or not isinstance(mark, str):
                    continue
                v = mark.strip().upper()
                if v == "X":
                    out[cid] = {"excluded": True}
                elif v == "O":
                    out[cid] = {"confirmed": True}
    finally:
        wb.close()
    return out

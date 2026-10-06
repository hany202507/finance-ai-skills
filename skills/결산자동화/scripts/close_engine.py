# -*- coding: utf-8 -*-
"""결산 결정적 엔진 — raw → 분개장 → 살아있는 수식 워크북 + 더존 업로드 2종.

설계서 §2(계산은 결정적 엔진) / §5(표준 처리 관례) / §8(산출물 규격).
숫자는 이 엔진이 정한다. LLM 암산 금지. 산출물의 재무제표·시산표는 전부 수식.

사용:
  python close_engine.py <raw_dir> <out_dir> <기간라벨> [profile_module]
예:
  python close_engine.py ".../2분기_당기" ".../output" 2026Q2 profile_sample

raw_dir 에 있어야 할 파일(부분 이름 매칭):
  전기이월분개장*.xlsx  기초BS*.xlsx  급여대장*.xlsx  은행거래내역*.xlsx
  카드승인내역*.xlsx    세금계산서*.xlsx  거래처마스터*.xlsx
"""
import sys, os, glob, importlib, math
from collections import defaultdict, OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------------- 유틸
def find(raw_dir, key):
    hits = glob.glob(os.path.join(raw_dir, f"*{key}*.xlsx"))
    if not hits:
        raise FileNotFoundError(f"raw에 '{key}' 파일이 없습니다: {raw_dir}")
    return sorted(hits)[0]

def rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [r for r in ws.iter_rows(values_only=True)]

def mm(dstr):            # 'YYYY-MM-DD' → 'MM'
    return str(dstr)[5:7]

def ymd(year, m, d):
    return f"{year}-{int(m):02d}-{int(d):02d}"

def to_int(x):
    return 0 if x in (None, "") else int(round(float(x)))

def _norm(name):
    """상호를 비교용으로 깎는다. 법인격 표기와 공백·기호만 버리고 글자는 남긴다.

    '네이버파이낸셜(주)' 와 '네이버파이낸셜' 을 같게 보려는 것이다. 여기서 더 깎으면
    '네이버파이낸셜(주)'(플랫폼) 과 '네이버(주) 검색광고'(광고) 가 섞이므로 더 깎지 않는다.
    """
    t = str(name or "")
    for junk in ("(주)", "주식회사", "(유)", "(사)", "(재)", "（주）"):
        t = t.replace(junk, "")
    return "".join(ch for ch in t if not ch.isspace() and ch not in "·.,-_()[]")

# ---------------------------------------------------------------- 엔진
class Engine:
    def __init__(self, raw_dir, period, P):
        self.raw = raw_dir; self.period = period; self.P = P
        self.year = int(str(period)[:4])
        self.J = []          # 분개장 라인
        self.ADJ = []        # 결산조정 라인
        self.no = 0
        self.master = {}     # 거래처명 → (코드, 사업자번호)
        self._mnorm = {}     # 정규화 상호 → (코드, 사업자번호)
        self.party_miss = {}  # 해소 못한 거래처라벨 → 등장 계정 set. 업로드 사전검증이 읽는다
        self._load()

    # --- 거래처 해소 (더존 업로드 거래처코드·사업자번호)
    def resolve_party(self, label, acct=None):
        """거래처라벨 → (거래처코드, 사업자번호). 못 찾으면 ("", 0) 이고 party_miss 에 남긴다.

        더존은 채권·채무 계정에 거래처코드가 없으면 거래처별 잔액을 못 잡는다. 그래서
        빈칸을 조용히 두지 않고 못 찾은 라벨을 모아 업로드 사전검증에서 드러낸다.

        찾는 순서는 좁은 것부터다. 추측으로 넓히지 않는다.
          1 프로파일 거래처별칭  — 전문가가 확정한 라벨→상호 대응. 이게 정답이다
          2 거래처마스터 상호 완전일치
          3 법인격 표기만 깎은 정규화 일치
          4 정규화 부분일치가 '단 하나'일 때만. 둘 이상이면 빈칸 + miss 등재
        프로파일 집합거래처에 적은 라벨은 한 거래처가 아니므로(임직원·다수(매입처) 등)
        빈칸이 정상이고 miss 에 넣지 않는다.
        """
        lab = str(label or "")
        if not lab:
            return "", 0
        if lab in getattr(self.P, "집합거래처", ()):
            return "", 0
        alias = getattr(self.P, "거래처별칭", {}).get(lab)
        if alias:
            if alias in self.master:
                return self.master[alias]
            hit = self._mnorm.get(_norm(alias))
            if hit:
                return hit
            # 별칭이 가리키는 상호가 마스터에 없다. 별칭 오타거나 거래처 미등록이다
            self.party_miss.setdefault(f"{lab}→{alias}(마스터 없음)", set()).add(acct)
            return "", 0
        if lab in self.master:
            return self.master[lab]
        hit = self._mnorm.get(_norm(lab))
        if hit:
            return hit
        n = _norm(lab)
        cand = {v for k, v in self._mnorm.items() if n and n in k}
        if len(cand) == 1:
            return cand.pop()
        self.party_miss.setdefault(lab, set()).add(acct)
        return "", 0

    def party_code(self, label, acct=None):
        return self.resolve_party(label, acct)[0]

    def party_biz(self, label, acct=None):
        return to_int(self.resolve_party(label, acct)[1])

    # --- raw 로드
    def _load(self):
        self.opening = rows(find(self.raw, "전기이월분개장"))
        self.tax = rows(find(self.raw, "세금계산서"))[1:]
        self.bank = rows(find(self.raw, "은행거래내역"))[1:]
        self.card = rows(find(self.raw, "카드승인내역"))[1:]
        self.pay = rows(find(self.raw, "급여대장"))[1:]
        for r in rows(find(self.raw, "거래처마스터"))[1:]:
            if r[0] is None: continue
            self.master[str(r[1])] = (str(r[0]), r[3])
            self._mnorm.setdefault(_norm(r[1]), (str(r[0]), r[3]))
        self.tax = [r for r in self.tax if r and r[0] is not None]
        self.bank = [r for r in self.bank if r and r[0] is not None]
        self.card = [r for r in self.card if r and r[0] is not None]
        self.pay = [r for r in self.pay if r and r[0] is not None]
        # 채널매출·정산차감은 선택이다. 세금계산서 매출만 있는 회사는 이 파일이 없다
        self.chsales = self._opt("채널매출")
        self.setl = self._opt("정산차감")

    def _opt(self, key):
        try:
            return [r for r in rows(find(self.raw, key))[1:] if r and r[0] is not None]
        except FileNotFoundError:
            return []

    def months(self):
        ms = {mm(r[0]) for r in self.tax}
        ms |= {str(r[0])[5:7] for r in self.chsales}
        return sorted(ms)

    # --- 분개장 기록 헬퍼
    def entry(self, date, gubun, lines, cls="확정"):
        """lines: [(code, name, debit, credit, cf, party, desc)]"""
        self.no += 1
        for code, name, dr, cr, cf, party, desc in lines:
            self.J.append(dict(no=self.no, date=date, gubun=gubun, code=code, acct=name,
                               debit=to_int(dr), credit=to_int(cr), cf=cf or "",
                               cls=cls, desc=desc, party=party or ""))

    # =========================================================== STEP: 전기이월
    def carry_forward(self):
        for r in self.opening[1:]:
            if r[0] is None: continue
            date, gubun, code, name, dr, cr, party, desc = r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]
            if gubun != "전기이월": continue
            d = str(date)[:10]
            self.entry(d, "전기이월",
                       [(int(code), name, dr, cr, "", party, desc)])

    # =========================================================== STEP: 매출(발생)
    def sales(self):
        agg = defaultdict(lambda: [0, 0])   # (month, 계정) → [공급, 세액]
        for r in self.tax:
            if r[1] != "매출": continue
            m = mm(r[0]); item = r[5]; code = self.P.매출품목_계정[item]
            agg[(m, code)][0] += to_int(r[6]); agg[(m, code)][1] += to_int(r[7])
        day = self.P.게시일["매출"]
        for m in self.months():
            for code in (401, 404):
                if (m, code) not in agg: continue
                sup, vat = agg[(m, code)]
                nm = self.P.매출계정명[code]
                d = ymd(self.year, m, day)
                self.entry(d, "매출", [
                    (108, "외상매출금", sup + vat, 0, "", "다수(매출처)", f"{int(m)}월 {nm} 외상매출"),
                    (code, nm, 0, sup, "", "다수(매출처)", f"{int(m)}월 {nm}"),
                    (255, "부가세예수금", 0, vat, "", "", f"{int(m)}월 매출부가세"),
                ])

    # =========================================================== STEP: 채널매출(총액 인식)
    def channel_sales(self):
        """채널 원천에서 매출을 총액으로 인식한다. 통장 입금에서 거꾸로 잡지 않는다.

        직판은 외상매출금으로 받고, 중개(입점사)는 판매대금 전액이 아니라 수수료만
        수익으로 잡는다. 판매대금은 통장에서 선수금으로 받아 여기서 수수료만 대체한다.
        """
        day = self.P.게시일["매출"]
        for x in self._prep_channel():
            u = x["acct"]; m = x["ym"][5:7]; d = ymd(self.year, m, day)
            tag = f" {x['dtype']}" if x["dtype"] not in ("NONE", None, "") else ""
            head = f"{int(m)}월 {x['chname']} {x['taxkind']}{tag} 매출({x['n']:,}건)"
            lines = []
            if x["gap"] and x["gap_kind"] == "비용":
                # 보전받지 못하는 차액. 대가만 채권으로 잡고 나머지는 판촉 성격의 비용이다
                lines.append((u["상대코드"], u["상대명"], x["paid"], 0, "", x["party"], head))
                lines.append((x["cost"][0], x["cost"][1], x["gap"], 0, "", x["party"],
                              f"{int(m)}월 {x['chname']} {x['dtype']} 시가과세 차액({x['근거']})"))
            else:
                # 조정 없음이거나 제3자가 보전하는 차액. 전액 채권이다
                lines.append((u["상대코드"], u["상대명"], x["base"], 0, "", x["party"], head))
            lines.append((u["수익코드"], u["수익명"], 0, x["sup"], "", x["party"], head))
            if x["vat"]:
                lines.append((255, "부가세예수금", 0, x["vat"], "", x["party"],
                              f"{int(m)}월 {x['chname']} 매출부가세"))
            self.entry(d, "매출", lines)

    def _prep_channel(self):
        """채널매출 원천에 매출조정 규칙을 적용해 과세표준을 확정한다.

        할인은 종류마다 과세표준이 다르다. 에누리와 자기적립마일리지는 대가에서 빠지고,
        제3자가 보전하는 쿠폰은 보전액을 더하며, 보전 없는 마일리지는 시가로 본다.
        규칙과 근거 조문은 회사 프로파일의 매출조정_규칙이 들고 있다.
        """
        if hasattr(self, "chrows"): return self.chrows
        adj = getattr(self.P, "매출조정_규칙", {})
        acct = getattr(self.P, "채널매출_계정", {})
        cost = getattr(self.P, "차액비용계정", (833, "광고선전비"))
        out = []
        for r in self.chsales:
            (ym, ch, chname, kind, taxkind, dtype, gift, n,
             paid, listamt, disc) = [r[i] for i in range(11)]
            party = r[11] if len(r) > 11 else ""
            paid, listamt = to_int(paid), to_int(listamt)
            if paid == 0 or kind not in acct: continue
            rule = adj.get(dtype, dict(기준="대가", 차액처리=None, 근거="규칙 없음 · 대가 기준"))
            base = paid if rule["기준"] == "대가" else listamt
            sup = base * 10 // 11 if taxkind == "TAXABLE" else base
            vat = base - sup if taxkind == "TAXABLE" else 0
            out.append(dict(ym=str(ym), ch=ch, chname=chname, kind=kind, taxkind=taxkind,
                            dtype=dtype, gift=gift, n=to_int(n), paid=paid, base=base,
                            sup=sup, vat=vat, gap=base - paid,
                            gap_kind=rule["차액처리"], 근거=rule["근거"],
                            acct=acct[kind], cost=cost, party=party))
        self.chrows = out
        return out

    # =========================================================== STEP: 정산차감
    def settlement_deductions(self):
        """정산에서 빠진 것으로 외상매출금을 닫는다. 항목별 처리는 프로파일이 정한다."""
        rules = getattr(self.P, "정산차감_규칙", {})
        day = self.P.게시일.get("정산", 25)
        for r in self.setl:
            ym, ch, item, counter, amt, party = r[:6]
            amt = to_int(amt)
            if amt == 0 or item not in rules: continue
            code, name, vat_kind = rules[item]
            m = str(ym)[5:7]; d = ymd(self.year, m, day)
            if vat_kind == "공제":
                sup = amt * 10 // 11; vat = amt - sup
                lines = [(code, name, sup, 0, "", counter, f"{int(m)}월 {item}({counter})"),
                         (135, "부가세대급금", vat, 0, "", counter, f"{int(m)}월 {item} 매입부가세({counter})")]
            else:
                lines = [(code, name, amt, 0, "", counter, f"{int(m)}월 {item}({counter})")]
            lines.append((108, "외상매출금", 0, amt, "", party,
                          f"{int(m)}월 {item} 정산상계({counter})"))
            self.entry(d, "정산", lines)

    # =========================================================== STEP: 은행 라벨 규칙
    def bank_rules(self):
        """프로파일이 정한 은행 계정 라벨을 월별로 묶어 기표한다."""
        rules = getattr(self.P, "은행라벨_규칙", {})
        if not rules: return
        by = OrderedDict()
        for r in self.bank:
            lab = r[6]
            if lab not in rules: continue
            amt = to_int(r[2]) if rules[lab]["side"] == "입금" else to_int(r[3])
            key = (mm(r[0]), lab)
            by.setdefault(key, [0, r[7]])
            by[key][0] += amt
        for (m, lab), (amt, party) in sorted(by.items()):
            if amt == 0: continue
            u = rules[lab]
            d = ymd(self.year, m, u.get("day", 15))
            self.entry(d, "이체", [
                (u["dr"], u["dr_name"], amt, 0, "" if u["dr"] != 103 else u.get("cf", ""),
                 party, f"{int(m)}월 {u['desc']}"),
                (u["cr"], u["cr_name"], 0, amt, u.get("cf", "") if u["cr"] == 103 else "",
                 party, f"{int(m)}월 {u['desc']}"),
            ])

    # =========================================================== STEP: 미분류(가계정)
    def unclassified(self):
        """규칙에 안 걸린 거래를 가계정에 두고 목록으로 남긴다.

        분개에서 빼면 통장 순증감과 대차가 안 맞고, 임의로 계정을 주면 어느 것이
        규칙이고 어느 것이 추측인지 구분이 없어진다. 가계정이 그 사이를 지킨다.
        """
        acct = getattr(self.P, "미분류_계정", None)
        self.unclassified_rows = []
        if not acct:
            return
        by = OrderedDict()
        for r in self.bank:
            if r[6] != "미분류":
                continue
            side = "입금" if to_int(r[2]) else "출금"
            amt = to_int(r[2]) or to_int(r[3])
            self.unclassified_rows.append(
                [str(r[0])[:10], r[1], side, amt, r[4],
                 "전기에 없던 거래이고 적요가 성격을 설명하지 않는다"])
            by.setdefault((mm(r[0]), side), 0)
            by[(mm(r[0]), side)] += amt
        for (m, side), amt in sorted(by.items()):
            if not amt:
                continue
            code, name = acct[side]
            d = ymd(self.year, m, 28)
            desc = f"{int(m)}월 미분류({side}) 가계정 대체 · 전문가 확인 필요"
            if side == "출금":
                lines = [(code, name, amt, 0, "", "미확인", desc),
                         (103, "보통예금", 0, amt, "영업-미분류", "미확인", desc)]
            else:
                lines = [(103, "보통예금", amt, 0, "영업-미분류", "미확인", desc),
                         (code, name, 0, amt, "", "미확인", desc)]
            self.entry(d, "미분류", lines)

    # =========================================================== STEP: 매입(발생)
    def purchases(self):
        # (month, 품목) → [공급, 세액, set(거래처)]
        agg = OrderedDict()
        order = list(self.P.매입품목_계정.keys())
        for r in self.tax:
            if r[1] != "매입": continue
            m = mm(r[0]); item = r[5]
            agg.setdefault((m, item), [0, 0, set()])
            agg[(m, item)][0] += to_int(r[6]); agg[(m, item)][1] += to_int(r[7])
            agg[(m, item)][2].add(r[3])
        day = self.P.게시일["매입"]
        for m in self.months():
            for item in order:
                if (m, item) not in agg: continue
                sup, vat, parties = agg[(m, item)]
                code, name, sang, cf = self.P.매입품목_계정[item]
                party = list(parties)[0] if len(parties) == 1 else "다수(매입처)"
                d = ymd(self.year, m, day)
                if sang == 251:   # 외상매입(상품)
                    self.entry(d, "매입", [
                        (code, name, sup, 0, "", party, f"{int(m)}월 {name.replace('상품','상품매입') if name=='상품' else name}"),
                        (135, "부가세대급금", vat, 0, "", party, f"{int(m)}월 매입부가세"),
                        (251, "외상매입금", 0, sup + vat, "", party, f"{int(m)}월 상품 외상매입"),
                    ])
                else:             # 경비(현금지급)
                    self.entry(d, "매입", [
                        (code, name, sup, 0, "", party, f"{int(m)}월 {name}"),
                        (135, "부가세대급금", vat, 0, "", party, f"{int(m)}월 매입부가세"),
                        (103, "보통예금", 0, sup + vat, cf, party, f"{int(m)}월 {name} 대금지급"),
                    ])

    # =========================================================== STEP: 상품매입대금 결제(외상 상환)
    def pay_ap(self):
        by = defaultdict(int)
        for r in self.bank:
            if r[6] == "매입" and "상품매입대금" in str(r[4]):
                by[mm(r[0])] += to_int(r[3])
        for m in sorted(by):
            amt = by[m]; d = ymd(self.year, m, 15)
            self.entry(d, "이체", [
                (251, "외상매입금", amt, 0, "", "다수(매입처)", f"{int(m)}월 상품매입대금 결제"),
                (103, "보통예금", 0, amt, "영업-상품매입", "다수(매입처)", f"{int(m)}월 상품매입대금 결제"),
            ])

    # =========================================================== STEP: 매출채권 회수
    def collect(self):
        by = OrderedDict()   # (month,거래처라벨) → amt
        buckets = []
        for r in self.bank:
            if r[6] == "매출채권회수":
                m = mm(r[0]); lab = r[7]
                by.setdefault((m, lab), 0); by[(m, lab)] += to_int(r[2])
                if lab not in buckets: buckets.append(lab)
        day = self.P.게시일["매출회수"]
        for m in self.months():
            for lab in buckets:
                if (m, lab) not in by: continue
                amt = by[(m, lab)]; d = ymd(self.year, m, day)
                self.entry(d, "이체", [
                    (103, "보통예금", amt, 0, "영업-매출회수", lab, f"{int(m)}월 매출대금 회수({lab})"),
                    (108, "외상매출금", 0, amt, "", lab, f"{int(m)}월 매출채권 회수({lab})"),
                ])

    # =========================================================== STEP: 카드(기중)
    def cards(self):
        # (month, 가맹점) → [승인, 취소]
        agg = OrderedDict()
        order = []
        for r in self.card:
            m = mm(r[0]); shop = r[3]
            if shop not in order: order.append(shop)
            agg.setdefault((m, shop), [0, 0])
            agg[(m, shop)][0] += to_int(r[6]); agg[(m, shop)][1] += to_int(r[7])
        cardparty = self.P.카드대금거래처
        for m in self.months():
            for shop in order:
                if (m, shop) not in agg: continue
                gross = agg[(m, shop)][0] - agg[(m, shop)][1]
                if gross == 0: continue
                code, name, vat_kind = self.P.카드가맹점_계정[shop]
                d = ymd(self.year, m, 14)
                if vat_kind == "공제":
                    sup = gross * 10 // 11
                    vat = gross - sup
                    self.entry(d, "카드", [
                        (code, name, sup, 0, "", shop, f"{int(m)}월 카드 {shop}(매입세액 공제, 가맹점 사업자번호 확보)"),
                        (135, "부가세대급금", vat, 0, "", shop, f"{int(m)}월 카드매입 부가세({shop})"),
                        (253, "미지급금", 0, gross, "", cardparty, f"{int(m)}월 카드 {shop}"),
                    ])
                elif vat_kind == "불공제":
                    # 불공제 사유는 가맹점마다 다르다(접대비·승용차·면세·사업무관).
                    # 사유를 한 문구로 박으면 산출물에서 근거가 사라진다. 프로파일이 정한다
                    reason = getattr(self.P, "카드불공제_사유", {}).get(shop, "")
                    # 「매입세액 불공제」는 항상 들어간다. 검증과 더존 사유코드 매핑이 이 말로 고른다
                    tail = f"매입세액 불공제 · {reason}" if reason else "매입세액 불공제"
                    self.entry(d, "카드", [
                        (code, name, gross, 0, "", shop, f"{int(m)}월 카드 {shop}({tail}, 전액비용)"),
                        (253, "미지급금", 0, gross, "", cardparty, f"{int(m)}월 카드 {shop} 매입세액 불공제"),
                    ])
                else:  # 해외
                    self.entry(d, "카드", [
                        (code, name, gross, 0, "", cardparty, f"{int(m)}월 카드 해외결제 {shop}(전액비용)"),
                        (253, "미지급금", 0, gross, "", cardparty, f"{int(m)}월 카드 해외결제 {shop}"),
                    ])

    # =========================================================== STEP: 카드대금 결제
    def pay_card(self):
        by = defaultdict(int)
        for r in self.bank:
            if r[6] == "신용카드대금": by[mm(r[0])] += to_int(r[3])
        for m in sorted(by):
            amt = by[m]; d = ymd(self.year, m, 5)
            self.entry(d, "이체", [
                (253, "미지급금", amt, 0, "", self.P.카드대금거래처, f"{int(m)}월 법인카드대금 결제"),
                (103, "보통예금", 0, amt, "영업-카드대금", self.P.카드대금거래처, f"{int(m)}월 법인카드대금 결제"),
            ])

    # =========================================================== STEP: 급여
    def payroll(self):
        by = OrderedDict()   # 귀속월 → [총지급, 공제합, 실지급]
        for r in self.pay:
            ym = str(r[0]); by.setdefault(ym, [0, 0, 0])
            by[ym][0] += to_int(r[6]); by[ym][1] += to_int(r[13]); by[ym][2] += to_int(r[14])
        self.payroll_yesu = []
        self.payroll_yesu_by_m = {}
        for ym in sorted(by):
            tot, ded, net = by[ym]
            m = ym[5:7]; d = ymd(self.year, m, 25)
            self.payroll_yesu.append(ded)
            self.payroll_yesu_by_m[m] = ded
            self.entry(d, "급여", [
                (801, "급여", tot, 0, "", "임직원", f"{int(m)}월 급여"),
                (254, "예수금", 0, ded, "", "임직원", f"{int(m)}월 급여 예수금(4대보험·원천세)"),
                (103, "보통예금", 0, net, "영업-급여", "임직원", f"{int(m)}월 급여 실지급"),
            ])

    # =========================================================== STEP: 4대보험
    def insurance4(self):
        # 기초 예수금(전기이월 254 대변)
        opening_yesu = 0
        for r in self.opening[1:]:
            if r and r[1] == "전기이월" and r[2] == 254:
                opening_yesu += to_int(r[5])
        by = OrderedDict()
        for r in self.bank:
            if r[6] == "4대보험": by.setdefault(mm(r[0]), 0); by[mm(r[0])] += to_int(r[3])
        prior = [opening_yesu] + self.payroll_yesu   # 상계 대상: 직전 예수금
        ybm = getattr(self, "payroll_yesu_by_m", {})
        for i, m in enumerate(sorted(by)):
            paid = by[m]
            # 납부월의 직전 급여월 예수금이 있으면 그것을 상계한다. 전기이월 예수금을
            # 첫 납부에 통째로 상계하면 회사부담분이 음수가 된다(1원도 안 맞는 신호).
            prev = "%02d" % (int(m) - 1)
            sang = ybm.get(prev, prior[i] if i < len(prior) else 0)
            corp = paid - sang
            d = ymd(self.year, m, 10)
            self.entry(d, "이체", [
                (254, "예수금", sang, 0, "", "공단/세무서", f"{int(m)}월 4대보험·원천세 예수금 상계"),
                (811, "복리후생비", corp, 0, "", "공단/세무서", f"{int(m)}월 4대보험 회사부담분"),
                (103, "보통예금", 0, paid, "영업-4대보험", "공단/세무서", f"{int(m)}월 4대보험·원천세 납부"),
            ])

    # =========================================================== STEP: 은행 단순 비용/수익
    def bank_simple(self):
        # 라벨 → (코드, 계정, CF, 게시일, 적요, 차변쪽=비용/자산)
        rules = [
            ("보험료",     821, "보험료",     "영업-보험료",   "화재/배상책임 보험료"),
            ("세금과공과", 817, "세금과공과금", "영업-세금과공과", "사업소분/면허세 등"),
            ("이자비용",   951, "이자비용",   "재무-이자지급", "차입금이자"),
        ]
        for lab, code, name, cf, desc in rules:
            by = OrderedDict()
            for r in self.bank:
                if r[6] == lab:
                    by.setdefault(mm(r[0]), [0, r[7]]); by[mm(r[0])][0] += to_int(r[3])
            for m in sorted(by):
                amt, party = by[m]
                d = ymd(self.year, m, 30 if lab == "이자비용" else 20)
                self.entry(d, "이체", [
                    (code, name, amt, 0, "", party, f"{int(m)}월 {desc}"),
                    (103, "보통예금", 0, amt, cf, party, f"{int(m)}월 {desc} 지급" if lab=="이자비용" else f"{int(m)}월 {name}"),
                ])
        # 이자수익(입금)
        by = OrderedDict()
        for r in self.bank:
            if r[6] == "이자수익":
                by.setdefault(mm(r[0]), [0, r[7]]); by[mm(r[0])][0] += to_int(r[2])
        for m in sorted(by):
            amt, party = by[m]; d = ymd(self.year, m, 30)
            self.entry(d, "이체", [
                (103, "보통예금", amt, 0, "영업-이자수취", party, f"{int(m)}월 예금이자"),
                (901, "이자수익", 0, amt, "", party, f"{int(m)}월 예금이자수익"),
            ])

    # =========================================================== STEP: 전기 부가세 납부
    def prior_vat(self):
        for r in self.bank:
            if r[6] == "부가세납부":
                amt = to_int(r[3]); d = str(r[0])[:10]
                self.entry(d, "이체", [
                    (261, "미지급세금", amt, 0, "", r[7], "1분기 부가세 예정신고 납부(전기 미지급세금 상계)"),
                    (103, "보통예금", 0, amt, "영업-부가세납부", r[7], "1분기 부가세 예정신고 납부"),
                ])

    # =========================================================== STEP: 결산조정
    def _by_period(self, value):
        """기간별 dict 면 이 결산의 기간 값을, 아니면 그대로 돌려준다. 월 단위로 두 번 돌리는 회사용."""
        return value[self.period] if isinstance(value, dict) else value

    def _sum_journal(self, code, side, gubun=None):
        s = 0
        for l in self.J:
            if l["code"] == code and (gubun is None or l["gubun"] == gubun):
                s += l[side]
        return s

    def adjust(self):
        last = str(self.year) + "-06-30"   # 분기말 — 기간에 맞춰 조정 필요시 프로파일화
        # 기간 말일 추정: 마지막 매출월 말일
        import calendar
        lm = int(self.months()[-1]); last = ymd(self.year, lm, calendar.monthrange(self.year, lm)[1])
        n = 0
        # 1) 감가/무형
        for code, name, accum, acqcode, acq, life, mths in self.P.상각자산:
            amt = acq * mths // life
            accname = "감가상각누계액" if code == 820 else "무형자산상각누계액"
            n += 1
            self.ADJ.append((n, last, code, name, amt, 0, f"{'비품 감가상각' if code==820 else '소프트웨어 무형자산상각'} (정액5년, {mths}개월분)"))
            self.ADJ.append((n, last, accum, accname, 0, amt, f"{'비품 감가상각' if code==820 else '소프트웨어 무형자산상각'} (정액5년, {mths}개월분)"))
        # 2) 상품매출원가 대체
        cur_purchase = self._sum_journal(146, "debit", gubun="매입")   # 당기 상품매입(이월 제외)
        inv_open = self._by_period(self.P.상품_기초재고)
        inv_close = self._by_period(self.P.상품_기말재고_잠정)
        cogs = inv_open + cur_purchase - inv_close
        n += 1
        self.ADJ.append((n, last, 451, "상품매출원가", cogs, 0, f"상품매출원가 대체(기말재고 잠정 {inv_close:,} - 확인 No.1)"))
        self.ADJ.append((n, last, 146, "상품", 0, cogs, f"상품매출원가 대체(기말재고 잠정 {inv_close:,} - 확인 No.1)"))
        # 3) 부가세 정리
        yesu = self._sum_journal(255, "credit")    # 부가세예수금 총
        dae = self._sum_journal(135, "debit")       # 부가세대급금 총
        mizigib = yesu - dae
        n += 1
        desc = f"부가세 정리(1기확정: 예수금{yesu:,}-대급금{dae:,}=미지급{mizigib:,}, 7/25납부)"
        self.ADJ.append((n, last, 255, "부가세예수금", yesu, 0, desc))
        self.ADJ.append((n, last, 135, "부가세대급금", 0, dae, desc))
        self.ADJ.append((n, last, 261, "미지급세금", 0, mizigib, desc))
        self.vat_payable = mizigib; self.cogs = cogs; self.card_deduct = self._card_deduct()

    def _card_deduct(self):
        s = 0
        for l in self.J:
            if l["gubun"] == "카드" and l["code"] == 135: s += l["debit"]
        return s

    def _tax_sales(self):
        return sum(to_int(r[6]) for r in self.tax if r[1] == "매출")

    def _channel_sales(self):
        """채널매출 공급가액 합계(매출조정 반영). 채널매출 파일이 없으면 0이다."""
        return sum(x["sup"] for x in self._prep_channel())

    def balance(self, code):
        """계정 기말 순차변(net-debit). 자산 +, 부채·자본·수익 −."""
        s = 0
        for l in self.J:
            if l["code"] == code: s += l["debit"] - l["credit"]
        for a in self.ADJ:
            if a[2] == code: s += to_int(a[4]) - to_int(a[5])
        return s

    def summary(self):
        b = self.balance
        sga_codes = sorted({l["code"] for l in self.J if 800 <= l["code"] < 900}
                           | {a[2] for a in self.ADJ if 800 <= a[2] < 900})
        sales = -(b(401) + b(404))
        cogs = b(451) + b(452)
        gp = sales - cogs
        sga = sum(b(c) for c in sga_codes)
        op = gp - sga
        oi = -b(901); oe = b(951); tax = b(998)
        ni = op + oi - oe - tax
        # 확인사항 본문이 기간을 따라가려면 금액만으로는 안 된다. 건수와 달도 같이 준다.
        # 여기를 빼면 프로파일이 숫자를 문자열에 박게 되고, 다른 달로 돌릴 때 틀린 값을 말한다
        p = str(self.period)
        month = int(p[5:7]) if len(p) == 7 and p[4] == "-" else 0
        cardno = {l["no"] for l in self.J if l["gubun"] == "카드"}
        carded = {l["no"] for l in self.J if l["code"] == 135}
        return dict(sales=sales, cogs=cogs, gp=gp, sga=sga, op=op, oi=oi, oe=oe, ni=ni,
                    ar=b(108), ap=-b(251), card_ap=-b(253), yesu=-b(254),
                    loan=-b(260), vat_payable=self.vat_payable,
                    card_deduct=self.card_deduct, cash_end=b(103),
                    period=p, month=month,
                    tax_n=len(self.tax),
                    tax_purchase_n=sum(1 for r in self.tax if str(r[1]) == "매입"),
                    tax_sales_n=sum(1 for r in self.tax if str(r[1]) == "매출"),
                    card_deny_n=len(cardno - carded),
                    corp_tax_paid=sum(l["debit"] for l in self.J
                                      if l["code"] == 261 and "법인세" in str(l["desc"])))

    # =========================================================== 전체 실행
    def run(self):
        self.carry_forward(); self.sales(); self.channel_sales()
        self.purchases(); self.pay_ap(); self.settlement_deductions()
        self.collect(); self.cards(); self.pay_card(); self.payroll()
        self.insurance4(); self.bank_simple(); self.bank_rules()
        self.unclassified()
        self.prior_vat(); self.adjust()
        return self

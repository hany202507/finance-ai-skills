# -*- coding: utf-8 -*-
"""확인사항 리포트(.docx) 생성 — 전문가 검토·회신용 문서(설계서 §7 STEP7 · 부록 B).

가장 중요한 산출물: 맨 앞에 「전문가 확정·회신 필요 항목」 요약(체크리스트) →
판단 5필드 표 → 전기 방식 승계·편차 표 → 채권채무/손익/검증 요약 → 꼬리말.

판단 항목(잠정/전기승계/편차)은 회사 프로파일 + 엔진 결과로 채우되,
회사별 특수 항목은 이 스크립트를 복제해 편집한다. 여기 기본값은 예시푸드마켓 예시.

사용:
  python build_docx.py <raw_dir> <out_dir> <기간라벨> [profile_module]
"""
import io, sys, os, importlib
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from close_engine import Engine

def shade(cell, hex_):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hex_)
    tcPr.append(sh)

def hdr_row(row):
    for c in row.cells:
        shade(c, "305496")
        for p in c.paragraphs:
            for r in p.runs: r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
            if not p.runs:
                run = p.add_run(""); run.font.bold = True

def add_table(doc, headers, data, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    hdr_row(t.rows[0])
    for rowd in data:
        cells = t.add_row().cells
        for i, v in enumerate(rowd): cells[i].text = str(v)
    return t

def won(n): return f"{int(round(n)):,}"


def 검증결과(out_dir, period):
    """검증 스크립트가 남긴 JSON 을 읽는다. 없으면 빈 목록이다(= 미실행)."""
    import json
    got = []
    for 이름, 파일 in (("업로드 사전검증", f"_검증_업로드_{period}.json"),
                     ("워크북 검산", f"_검증_워크북_{period}.json")):
        fp = os.path.join(out_dir, 파일)
        if os.path.exists(fp):
            try:
                got.append((이름, json.load(io.open(fp, encoding="utf-8"))))
            except Exception:
                pass
    return got

def build(eng, out_path, period):
    P = eng.P; S = eng.summary()
    doc = Document()
    doc.styles["Normal"].font.name = "맑은 고딕"; doc.styles["Normal"].font.size = Pt(10)

    h = doc.add_heading(f"결산 확인사항 — {P.회사명} {period}", level=0)
    sub = doc.add_paragraph()
    r = sub.add_run("전문가 검토 완료 전제 · 자동발송·업로드 금지 · 비식별 교육용 예시")
    r.italic = True; r.font.size = Pt(9)

    # 회사별 판단 항목은 프로파일이 준다. 없으면 아래 예시 기본값을 쓴다.
    # 값이 호출 가능하면 엔진 요약 S 를 넘겨 준다(금액을 본문에 끼워 넣을 때).
    def from_profile(attr, default):
        v = getattr(P, attr, None)
        if v is None: return default
        return v(S) if callable(v) else v

    # ── ① 전문가 확정·회신 필요 항목(요약)
    doc.add_heading("① 전문가 확정·회신 필요 항목 (요약)", level=1)
    for line in from_profile("확인사항_요약", [
        "☐ No.1 기말 상품재고 40,000,000 (실사 미도착) — 실사표 회신 필요 / 원가율 잠정",
        "☐ No.2 해외카드(META·OpenAI) 대리납부 부가세(§52) 해당 여부 — 확정 필요",
        "☑ No.3 카드 식대·포장 매입세액 공제(과세유형 57) — 전기 Q1 승계, 이견 시 회신",
        "☑ No.4 카드 택시(여객운송) 매입세액 불공제·전액비용 — 전기 Q1 플래그 확정",
        "☐ No.5 법인세 분기 미계상(연말 확정) — 중간예납 대상 여부 확인",
    ]):
        doc.add_paragraph(line, style="List Bullet")

    # ── ② 판단 항목 상세(5필드)
    doc.add_heading("② 판단 항목 상세 (5필드)", level=1)
    rows5 = from_profile("확인사항_5필드", [
        ["1","기말재고(원가대체)","실사표 미도착. 기초 40,000,000·당기매입 126,000,000",
         "①실사확정 ②원가율 잠정","잠정 40,000,000 유지(전기 수준) — 확정 시 매출원가 재계산",
         f"매출원가·순이익 변동(현 매출원가 상품 {won(S['cogs']-45000000)})"],
        ["2","해외카드 대리납부","META·OpenAI 해외결제 전액비용 처리",
         "①대리납부 대상 ②비대상","비대상 잠정(추가자료 시 재검토)","부가세 납부세액 변동 가능"],
        ["3","카드 식대·포장 공제","가맹점 사업자번호 확보 → 과세유형 57 공제",
         "①공제 ②불공제","공제(전기 Q1 No.5 승계)",f"매입세액 +{won(S['card_deduct'])}(미지급세금 감소)"],
    ])
    add_table(doc, ["No","항목","상황(근거)","선택지","추천안과 근거","재무제표 영향"], rows5)

    # ── ③ 전기 방식 승계·편차
    doc.add_heading("③ 전기 방식 승계·편차", level=1)
    seung = from_profile("전기승계_표", [
        ["1","카드 식대·포장 매입세액","공제(과세유형 57)","공제 동일","승계"],
        ["2","카드 택시(여객운송)","잠정 공제 → 재분류 플래그(Q1 No.8)","불공제·전액비용 확정","편차(플래그 확정)"],
        ["3","4대보험 회사부담분","복리후생비","복리후생비","승계"],
        ["4","감가·무형 상각","정액 5년","정액 5년","승계"],
        ["5","이자 현금흐름","수익=영업 / 비용=재무","동일","승계"],
        ["6","법인세","분기 미계상","분기 미계상","승계"],
    ])
    add_table(doc, ["No","항목","전기 방식","당기 처리","승계/편차"], seung)

    # ── ④ 채권·채무 기말잔액
    doc.add_heading("④ 채권·채무 기말잔액", level=1)
    debt = [
        ["외상매출금(108)", won(S["ar"]), ""],
        ["외상매입금(251)", won(S["ap"]), ""],
        ["미지급금-카드(253)", won(S["card_ap"]), ""],
        ["예수금(254)", won(S["yesu"]), "4대보험·원천세"],
        ["단기차입금(260)", won(S["loan"]), "국민은행 연6%"],
        ["미지급세금(261)", won(S["vat_payable"]), "1기 확정 부가세, 7/25 납부"],
    ]
    add_table(doc, ["계정","기말잔액","비고"], debt)

    # ── ⑤ 손익 요약
    doc.add_heading("⑤ 손익 요약", level=1)
    pnl = [
        ["매출액", won(S["sales"])],
        ["매출원가", won(S["cogs"])],
        ["매출총이익", won(S["gp"])],
        ["판매관리비", won(S["sga"])],
        ["영업이익", won(S["op"])],
        ["영업외수익(이자수익)", won(S["oi"])],
        ["영업외비용(이자비용)", won(S["oe"])],
        ["당기순이익", won(S["ni"])],
        ["법인세비용", "0 (분기 미계상, 연말 확정)"],
    ]
    add_table(doc, ["과목","금액"], pnl)

    # ── ⑥ 검증 요약 (실제 결과를 읽어 적는다. 지어내지 않는다)
    doc.add_heading("⑥ 검증 결과", level=1)
    검증 = 검증결과(os.path.dirname(os.path.abspath(out_path)), period)
    if not 검증:
        r = doc.add_paragraph().add_run(
            "검증 미실행. verify_upload.py 와 verify_close.py 를 돌리기 전에 만든 문서다. "
            "이 문서만 보고 확정하면 안 된다.")
        r.bold = True
    else:
        rows검 = []
        for 이름, d in 검증:
            for c in d.get("검사", []):
                rows검.append([이름, c["항목"], c["결과"], c.get("상세", "")])
            for m in d.get("fail", []):
                rows검.append([이름, m, "FAIL", ""])
            for m in d.get("확인", []):
                rows검.append([이름, m, "확인", ""])
            if not d.get("검사") and not d.get("fail"):
                rows검.append([이름, "형식·필수값·거래처코드·사업자번호",
                               "PASS" if d.get("pass") else "FAIL", ""])
        add_table(doc, ["검증", "항목", "결과", "상세"], rows검)
        나쁨 = [x for x in rows검 if x[2] == "FAIL"]
        p2 = doc.add_paragraph()
        if 나쁨:
            r = p2.add_run(f"FAIL {len(나쁨)}건. 이 상태로 더존에 올리거나 확정하면 안 된다.")
            r.bold = True
        else:
            p2.add_run("위 항목은 검증 스크립트가 실제로 실행한 결과다. "
                       "여기에 없는 항목은 검증하지 않은 것이다.")
    doc.add_paragraph(
        "검증하지 않는 것: 원천 자료 자체의 정확성, 판단 항목(공제·불공제, 계정 분류, 기말재고)의 타당성. "
        "이 둘은 전문가가 본다.", style="List Bullet")

    # ── 꼬리말
    doc.add_heading("산출물 4종 & 주의", level=1)
    for line in [
        f"결산확정_{period}.xlsx · 일반전표_업로드_더존_{period}.xlsx · 매입매출전표_업로드_더존_{period}.xlsx · 본 확인사항_{period}.docx",
        "더존 재업로드 시 기존 자동분개 삭제 후 올릴 것(중복 방지).",
        "잠정 항목(기말재고·해외카드) 확정 전 더존 업로드 금지.",
    ]:
        p = doc.add_paragraph(line, style="List Bullet")
    doc.save(out_path)

def main():
    raw, out, period = sys.argv[1], sys.argv[2], sys.argv[3]
    pmod = sys.argv[4] if len(sys.argv) > 4 else "profile_sample"
    P = importlib.import_module(pmod)
    os.makedirs(out, exist_ok=True)
    eng = Engine(raw, period, P).run()
    path = os.path.join(out, f"확인사항_{period}.docx")
    build(eng, path, period)
    print("확인사항 생성:", path)

if __name__ == "__main__":
    main()

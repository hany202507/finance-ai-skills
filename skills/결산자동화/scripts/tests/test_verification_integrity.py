# -*- coding: utf-8 -*-
"""검증기가 실제로 막는지 고정한다.

2026-09-13 피드백에서 나온 것: 검증기가 검사하지 않은 항목을 PASS 로 표시하고 있었다.
  · build_workbook 이 12개 중 6개를 ="PASS" 리터럴로 채웠다
  · 순이익 검사가 =IF(NI=NI,...) 로 자기 자신과 비교했다
  · build_docx 가 「전부 PASS」를 고정 기재하고, 검증보다 먼저 돌았다
  · verify_close 가 수식 0개·무관 상수 워크북을 통과시켰다
  · verify_upload 가 월13·일32·구분과 차대 반대·빈 파일을 통과시켰다

여기 있는 반례가 다시 통과하기 시작하면 같은 일이 되풀이된 것이다.
"""
import os
import subprocess
import sys

import openpyxl
import pytest

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)


def 돌리기(script, args):
    r = subprocess.run([sys.executable, os.path.join(HERE, script)] + args, cwd=HERE,
                       capture_output=True, text=True, encoding="utf-8", errors="replace",
                       env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    return r.returncode, (r.stdout or "") + (r.stderr or "")


# ── 업로드 사전검증 ────────────────────────────────────────
일반전표_컬럼 = ["월", "일", "구분", "계정과목코드", "계정과목명",
                "거래처코드", "거래처", "적요", "차변", "대변"]
매입매출_컬럼 = ["년도", "월", "일", "매입매출구분(1-매출/2-매입)", "과세유형", "불공제사유",
                "신용카드거래처코드", "신용카드사명", "신용카드(가맹점)번호", "거래처명",
                "사업자(주민)등록번호", "공급가액", "부가세", "품명", "전자세금(1.전자)",
                "기본계정", "상대계정", "현금영수증 승인번호"]
성한_일반 = [6, 25, 3, 811, "복리후생비", None, "임직원", "정상", 1000, None]
성한_대변 = [6, 25, 4, 103, "보통예금", None, "임직원", "정상", None, 1000]
성한_매출 = [2026, 6, 1, 2, 51, None, None, None, None, "(주)가나",
            1234567890, 100000, 10000, "원료", 1, 146, 251, None]


def 업로드쓰기(d, 일반행, 매출행, period="2026-06"):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "일반전표"
    ws.append(일반전표_컬럼)
    for r in 일반행:
        ws.append(r)
    wb.save(os.path.join(d, f"일반전표_업로드_더존_{period}.xlsx"))
    wb2 = openpyxl.Workbook(); ws2 = wb2.active; ws2.title = "매출자료 & 매입자료"
    ws2.append(매입매출_컬럼)
    for r in 매출행:
        ws2.append(r)
    wb2.save(os.path.join(d, f"매입매출전표_업로드_더존_{period}.xlsx"))


def 사전검증(d):
    return 돌리기("verify_upload.py", [str(d), "2026-06"])


def test_정상_업로드는_통과한다(tmp_path):
    업로드쓰기(tmp_path, [성한_일반, 성한_대변], [성한_매출])
    code, out = 사전검증(tmp_path)
    assert code == 0, out


@pytest.mark.parametrize("바꿀것, 기대문구", [
    ({0: 13}, "월=13"),
    ({1: 32}, "일=32"),
])
def test_말이_안_되는_날짜를_잡는다(tmp_path, 바꿀것, 기대문구):
    행 = list(성한_일반)
    for i, v in 바꿀것.items():
        행[i] = v
    업로드쓰기(tmp_path, [행, 성한_대변], [성한_매출])
    code, out = 사전검증(tmp_path)
    assert code == 1 and 기대문구 in out, out


def test_구분과_차대_방향이_반대면_잡는다(tmp_path):
    """구분 3(차변)인데 금액이 대변에 있으면 더존이 반대로 기표한다."""
    뒤집힘 = [6, 25, 3, 811, "복리후생비", None, "임직원", "반대", None, 1000]
    맞춤 = [6, 25, 4, 103, "보통예금", None, "임직원", "반대", 1000, None]
    업로드쓰기(tmp_path, [뒤집힘, 맞춤], [성한_매출])
    code, out = 사전검증(tmp_path)
    assert code == 1 and "구분=3(차변)인데 금액이 대변" in out, out


def test_빈_파일을_통과시키지_않는다(tmp_path):
    업로드쓰기(tmp_path, [], [])
    code, out = 사전검증(tmp_path)
    assert code == 1 and "한 줄도 없다" in out, out


def test_음수_금액을_잡는다(tmp_path):
    행 = list(성한_일반); 행[8] = -1000
    업로드쓰기(tmp_path, [행, 성한_대변], [성한_매출])
    code, out = 사전검증(tmp_path)
    assert code == 1 and "음수" in out, out


def test_매출입_방향과_과세유형이_어긋나면_잡는다(tmp_path):
    """매출(1)로 적고 매입 코드 51 을 쓰면 신고서가 반대편에 집계된다."""
    행 = list(성한_매출); 행[3] = 1
    업로드쓰기(tmp_path, [성한_일반, 성한_대변], [행])
    code, out = 사전검증(tmp_path)
    assert code == 1 and "매입 코드" in out, out


def test_사업자번호가_텍스트면_잡는다(tmp_path):
    행 = list(성한_매출); 행[10] = "1234567890"
    업로드쓰기(tmp_path, [성한_일반, 성한_대변], [행])
    code, out = 사전검증(tmp_path)
    assert code == 1 and "숫자형" in out, out


# ── 워크북 검산 ───────────────────────────────────────────
def 가짜워크북(path, 매출, 영업이익, 순이익, 시산표합=0):
    """수식이 하나도 없고 숫자만 적어 넣은 워크북. 결산확정 워크북이 아니다."""
    wb = openpyxl.Workbook()
    ts = wb.create_sheet("시산표")
    ts.append(["계정코드", "계정과목", "차변", "대변", "", "", "", "기말"])
    ts.append([103, "보통예금", 0, 0, "", "", "", 시산표합])
    is_ = wb.create_sheet("손익계산서"); is_.append(["과목", "금액"])
    for n, v in (("I. 매출액", 매출), ("V. 영업이익", 영업이익), ("X. 당기순이익", 순이익)):
        is_.append([n, v])
    bs = wb.create_sheet("재무상태표"); bs.append(["과목", "", "", "금액"])
    bs.append(["자산총계", "", "", 500]); bs.append(["부채와자본총계", "", "", 500])
    cf = wb.create_sheet("현금흐름표_직접법"); cf.append(["", "항목", "금액"])
    cf.append(["", "기말현금", 시산표합])
    del wb["Sheet"]
    wb.save(path)


def test_수식_없는_워크북을_통과시키지_않는다(tmp_path):
    """매출·영업이익·순이익이 서로 무관한 상수인데 통과하면 검산이 없는 것이다."""
    p = tmp_path / "결산확정_2026-06.xlsx"
    가짜워크북(p, 111, 222, 333)
    code, out = 돌리기("verify_close.py", [str(p)])
    assert code == 1, out
    assert "수식 연결" in out and "FAIL" in out, out


def test_손익_항등식이_깨지면_잡는다(tmp_path):
    p = tmp_path / "결산확정_2026-06.xlsx"
    가짜워크북(p, 111, 222, 333)
    _, out = 돌리기("verify_close.py", [str(p)])
    assert "당기순이익=매출-원가-판관비+영업외-법인세" in out, out
    표시 = [l for l in out.splitlines() if "당기순이익=매출" in l]
    assert 표시 and "FAIL" in 표시[0], 표시


def test_검증리포트에_가짜_PASS_가_없다():
    """build_workbook 이 ="PASS" 리터럴을 다시 넣으면 실패한다."""
    src = open(os.path.join(HERE, "build_workbook.py"), encoding="utf-8").read()
    i = src.index("검증리포트")
    블록 = src[i:i + 4000]
    assert "'=\"PASS\"'" not in 블록 and '"=\\"PASS\\""' not in 블록, \
        "검증리포트에 검사 없는 PASS 리터럴이 있다"
    assert "=IF({NI}={NI}" not in 블록, "순이익을 자기 자신과 비교하고 있다"


def test_확인사항_문서가_PASS_를_지어내지_않는다():
    """build_docx 에 '전부 PASS' 같은 고정 문구가 있으면 실패한다."""
    src = open(os.path.join(HERE, "build_docx.py"), encoding="utf-8").read()
    assert "전부 PASS" not in src, "문서에 검증 결과를 고정 기재하고 있다"
    assert "검증결과(" in src, "문서가 검증 결과 JSON 을 읽지 않는다"


def test_전표_하나가_어긋나면_전체합이_맞아도_잡는다(tmp_path):
    """더존은 전표 단위로 받는다. 전체 차=대 만 보면 이걸 놓친다.

    전표1: 차 1000 / 대 900  (모자람)
    전표2: 차 900  / 대 1000 (넘침)
    전체는 1900 = 1900 으로 맞지만 두 전표 다 틀렸다.
    """
    행 = [
        [6, 25, 3, 811, "복리후생비", None, "임직원", "전표1", 1000, None],
        [6, 25, 4, 103, "보통예금", None, "임직원", "전표1", None, 900],
        [6, 26, 3, 811, "복리후생비", None, "임직원", "전표2", 900, None],
        [6, 26, 4, 103, "보통예금", None, "임직원", "전표2", None, 1000],
    ]
    업로드쓰기(tmp_path, 행, [성한_매출])
    code, out = 사전검증(tmp_path)
    assert code == 1, out
    assert "차대가 맞지 않는다" in out, out


def test_정상_전표는_전표수를_센다(tmp_path):
    업로드쓰기(tmp_path, [성한_일반, 성한_대변], [성한_매출])
    code, out = 사전검증(tmp_path)
    assert code == 0 and "전표 1개" in out, out

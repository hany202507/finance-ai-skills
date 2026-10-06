# -*- coding: utf-8 -*-
"""더존 산출 테스트."""
import os, tempfile
from vat_engine import JudgmentInput, Purpose, judge, Routing
from vat_douzone import (
    Voucher, DENY_CODE, douzone_purchsale_rows, douzone_general_rows,
    build_douzone_xlsx, PURCHSALE_HEADER, GENERAL_HEADER,
)


def _v(inp, **ctx):
    return Voucher(년=2026, 월=7, 일=15, judgment=judge(inp), **ctx)


def test_purchsale_row_sales_taxinvoice():
    v = _v(JudgmentInput(방향="매출", 증빙=frozenset({"세금계산서"}), 과세구분="과세",
                         공급가=1000, 세액=100),
           거래처명="ABC", 사업자번호="123-45-67890", 품명="상품", 전자=True,
           기본계정=401, 상대계정=108)
    rows = douzone_purchsale_rows([v])
    assert len(rows) == 1
    r = rows[0]
    assert r[3] == 1              # 매입매출구분 매출
    assert r[4] == 11            # 과세유형
    assert r[10] == 1234567890   # 사업자번호 숫자형(하이픈 제거)
    assert r[11] == 1000 and r[12] == 100
    assert r[14] == 1            # 전자
    assert r[15] == 401 and r[16] == 108

def test_purchsale_purchase_gubun_2():
    v = _v(JudgmentInput(방향="매입", 증빙=frozenset({"세금계산서"}), 과세구분="과세",
                         공급가=1000, 세액=100, 지출목적=Purpose.일반경비))
    r = douzone_purchsale_rows([v])[0]
    assert r[3] == 2 and r[4] == 51

def test_purchsale_ungong_deny_code():
    v = _v(JudgmentInput(방향="매입", 증빙=frozenset({"세금계산서"}), 과세구분="과세",
                         공급가=1000, 세액=100, 지출목적=Purpose.토지_자본적지출, 자본적지출=True))
    r = douzone_purchsale_rows([v])[0]
    assert r[4] == 54            # 불공
    assert r[5] == "5"          # 토지관련 → 더존 코드 5

def test_deny_code_map_covers_all_reasons():
    from vat_engine import DenyReason
    for reason in DenyReason:
        assert reason in DENY_CODE

def test_general_rows_balanced_debit_credit():
    # 여객운송(택시) → 일반전표 전액비용
    v = _v(JudgmentInput(방향="매입", 증빙=frozenset({"카드"}), 총액=11000,
                         지출목적=Purpose.여객운송_개인서비스),
           기본계정=812, 기본계정명="여비교통비", 상대계정=253, 상대계정명="미지급금",
           적요="택시")
    rows = douzone_general_rows([v])
    assert len(rows) == 2
    차변, 대변 = rows
    assert 차변[2] == 3 and 차변[8] == 11000 and 차변[9] == 0   # 구분3 차변
    assert 대변[2] == 4 and 대변[8] == 0 and 대변[9] == 11000   # 구분4 대변
    assert 차변[3] == 812 and 대변[3] == 253

def test_purchsale_excludes_general_routing():
    v = _v(JudgmentInput(방향="매입", 증빙=frozenset({"카드"}), 총액=11000,
                         지출목적=Purpose.여객운송_개인서비스))
    assert douzone_purchsale_rows([v]) == []   # 일반전표로 갔으므로 매입매출 없음

def test_build_douzone_xlsx_smoke():
    vs = [
        _v(JudgmentInput(방향="매출", 증빙=frozenset({"세금계산서"}), 과세구분="과세",
                         공급가=1000, 세액=100), 거래처명="ABC", 사업자번호="1234567890",
           기본계정=401, 상대계정=108),
        _v(JudgmentInput(방향="매입", 증빙=frozenset({"카드"}), 총액=11000,
                         지출목적=Purpose.여객운송_개인서비스),
           기본계정=812, 상대계정=253, 거래처코드="00138"),
    ]
    d = tempfile.mkdtemp()
    ps, gen = os.path.join(d, "ps.xlsx"), os.path.join(d, "gen.xlsx")
    build_douzone_xlsx(vs, ps, gen)
    import openpyxl
    w1 = openpyxl.load_workbook(ps).active
    w2 = openpyxl.load_workbook(gen).active
    assert [c.value for c in w1[1]] == PURCHSALE_HEADER
    assert [c.value for c in w2[1]] == GENERAL_HEADER
    assert w1.max_row == 2       # header + 1 매입매출(매출)
    assert w2.max_row == 3       # header + 2 일반전표행
    # 거래처코드 앞자리0 텍스트 보존
    assert w2.cell(2, 6).value == "00138"

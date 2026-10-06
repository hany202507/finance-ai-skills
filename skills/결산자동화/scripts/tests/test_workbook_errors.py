# -*- coding: utf-8 -*-
"""정상 생성물의 수식 참조와 오류 셀 차단을 실제 재계산/CLI로 검증한다.

외부 실습파일이나 개발자 개인 경로에 의존하지 않는 작은 결산 fixture다.
원본 생성 워크북을 한 번 만들고 각 반례는 tmp_path 복사본만 바꾼다.
"""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from types import SimpleNamespace

import openpyxl
from openpyxl.formula import Tokenizer
import pytest

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))


def _line(no, code, name, debit=0, credit=0, opening=False, cf=""):
    return dict(no=no, date="2026-05-31" if opening else "2026-06-20",
                gubun="전기이월" if opening else "거래", code=code, acct=name,
                debit=debit, credit=credit, cf=cf, cls="", desc="시험 거래", party="")


@pytest.fixture(scope="module")
def generated_workbook(tmp_path_factory):
    """현금 1,090 / 부채 10 / 자본 1,080 / 순이익 80인 작은 정상 결산."""
    from build_workbook import build

    eng = SimpleNamespace(
        P=SimpleNamespace(),
        J=[_line(1, 103, "보통예금", debit=1000, opening=True),
           _line(1, 331, "자본금", credit=1000, opening=True),
           _line(2, 103, "보통예금", debit=110, cf="영업-매출"),
           _line(2, 401, "상품매출", credit=100),
           _line(2, 255, "부가세예수금", credit=10),
           _line(3, 811, "복리후생비", debit=20),
           _line(3, 103, "보통예금", credit=20, cf="영업-경비")],
        ADJ=[(1, "2026-06-30", 255, "부가세예수금", 10, 0, "부가세 정리"),
             (1, "2026-06-30", 135, "부가세대급금", 0, 0, "부가세 정리"),
             (1, "2026-06-30", 261, "미지급세금", 0, 10, "부가세 정리")],
        vat_payable=10, card_deduct=0, _tax_sales=lambda: 100, _channel_sales=lambda: 0)
    path = tmp_path_factory.mktemp("generated_close") / "결산확정_2026-06.xlsx"
    build(eng, str(path), "2026-06")
    return path


@pytest.fixture
def workbook_copy(generated_workbook, tmp_path):
    path = tmp_path / generated_workbook.name
    shutil.copy2(generated_workbook, path)
    return path


def _verify(path):
    result = subprocess.run(
        [sys.executable, "-X", "utf8", "-B", str(HERE / "verify_close.py"), str(path)],
        cwd=HERE, capture_output=True, text=True, encoding="utf-8", errors="replace",
        env=dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1"))
    report_path = path.with_name("_검증_워크북_2026-06.json")
    assert report_path.exists(), result.stdout + result.stderr
    return result, json.loads(report_path.read_text(encoding="utf-8"))


def _assert_error(path, sheet, coordinate, token):
    result, report = _verify(path)
    assert result.returncode == 1, result.stdout + result.stderr
    assert report["pass"] is False, report
    assert "ALL PASS" not in result.stdout, result.stdout
    assert report.get("오류셀"), report
    detail = json.dumps(report["오류셀"], ensure_ascii=False)
    assert sheet in detail and coordinate in detail and token in detail, detail


def test_generated_report_references_income_statement_and_recalculates_pass(workbook_copy):
    from verify_close import recalc_values

    original = openpyxl.load_workbook(workbook_copy)
    formula = original["검증리포트"]["C6"].value
    refs = [t.value for t in Tokenizer(formula).items if t.subtype == "RANGE"]
    assert len(refs) == 7 and all("손익계산서!" in r for r in refs), formula
    original.close()
    values = recalc_values(str(workbook_copy))
    assert values["검증리포트"]["C6"].value == "PASS"
    values.close()

    result, report = _verify(workbook_copy)
    assert result.returncode == 0, result.stdout + result.stderr
    assert report["pass"] is True and not report.get("오류셀"), report


def test_report_formula_with_old_unqualified_references_is_rejected(workbook_copy):
    wb = openpyxl.load_workbook(workbook_copy)
    income = wb["손익계산서"]
    labels = {row[0].value: row[0].row for row in income if row[0].value}
    sales, cogs, sga, oi, oe, tax, ni = [labels[name] for name in (
        "I. 매출액", "II. 매출원가", "IV. 판매비와관리비", "VI. 영업외수익",
        "VII. 영업외비용", "IX. 법인세비용", "X. 당기순이익")]
    wb["검증리포트"]["C6"] = (
        f'=IF(ROUND(손익계산서!$B${ni}-(B{sales}-B{cogs}-B{sga}+B{oi}-B{oe}-B{tax}),0)=0,"PASS","FAIL")')
    wb.save(workbook_copy); wb.close()
    _assert_error(workbook_copy, "검증리포트", "C6", "#VALUE!")


@pytest.mark.parametrize("value, expected", [
    ("=1/0", "#DIV/0!"),
    ("=#REF!", "#REF!"),
    ("#REF!", "#REF!"),
])
def test_unrelated_hidden_sheet_formula_and_static_errors_are_rejected(workbook_copy, value, expected):
    wb = openpyxl.load_workbook(workbook_copy)
    hidden = wb.create_sheet("숨김근거")
    hidden.sheet_state = "hidden"
    hidden["Z9"] = value
    wb.save(workbook_copy); wb.close()
    _assert_error(workbook_copy, "숨김근거", "Z9", expected)


def test_error_in_main_numeric_cell_writes_fail_instead_of_crashing(workbook_copy):
    wb = openpyxl.load_workbook(workbook_copy)
    wb["시산표"]["H2"] = "=1/0"
    wb.save(workbook_copy); wb.close()
    _assert_error(workbook_copy, "시산표", "H2", "#DIV/0!")


def test_recalculation_exception_replaces_previous_pass_json(workbook_copy):
    wb = openpyxl.load_workbook(workbook_copy)
    wb.create_sheet("수식오류")["A1"] = "=SUM("
    wb.save(workbook_copy); wb.close()
    report_path = workbook_copy.with_name("_검증_워크북_2026-06.json")
    report_path.write_text(json.dumps({"pass": True, "오래된성공": True}), encoding="utf-8")

    result, report = _verify(workbook_copy)
    assert result.returncode == 1, result.stdout + result.stderr
    assert report["pass"] is False and "오래된성공" not in report, report
    assert "ALL PASS" not in result.stdout, result.stdout
    assert "재계산" in json.dumps(report, ensure_ascii=False), report

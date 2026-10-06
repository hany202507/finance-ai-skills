# -*- coding: utf-8 -*-
"""더존 업로드 사전검증 — 올리기 전에 반려 사유를 먼저 잡는다 (설계서 §9 체크 11).

더존은 반려 메시지가 거칠다. 사업자번호가 텍스트라서 막힌 것과 거래처가 미등록이라
막힌 것이 같은 문구로 나온다. 그래서 올리기 전에 양식 쪽을 여기서 전부 본다.
양식 기준은 references/더존업로드양식.md 다. 이 스크립트가 그 문서의 집행부다.

사용:
  python verify_upload.py <산출폴더> <기간> [profile_module] [거래처마스터.xlsx]

거래처마스터를 안 주면 <산출폴더>/raw/거래처마스터.xlsx 를 찾는다. 있으면 전표의 거래처코드가
그 마스터에 실제로 있는지까지 본다. 더존이 제일 자주 반려하는 것이 미등록 거래처다.
"""
import sys, os, io, json, importlib
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

일반전표_컬럼 = ["월", "일", "구분", "계정과목코드", "계정과목명",
                "거래처코드", "거래처", "적요", "차변", "대변"]
매입매출_컬럼 = ["년도", "월", "일", "매입매출구분(1-매출/2-매입)", "과세유형", "불공제사유",
                "신용카드거래처코드", "신용카드사명", "신용카드(가맹점)번호", "거래처명",
                "사업자(주민)등록번호", "공급가액", "부가세", "품명", "전자세금(1.전자)",
                "기본계정", "상대계정", "현금영수증 승인번호"]

# 거래처별 잔액을 더존이 따로 잡는 계정. 여기에 거래처코드가 없으면 보조부가 안 선다
채권채무계정 = {108, 110, 120, 131, 133, 232, 251, 253, 254, 259, 260, 262, 293}
구분유효 = {1, 2, 3, 4, 5, 6}
# 1 출금·3 차변·5 결산차변은 차변에, 2 입금·4 대변·6 결산대변은 대변에 금액이 온다
차변구분 = {1, 3, 5}
대변구분 = {2, 4, 6}
# 매출유형은 매입매출구분 1, 매입유형은 2 여야 한다. 뒤집히면 신고서가 반대로 간다
매출유형 = {11, 12, 13, 14, 16, 17, 22, 23}
매입유형 = {51, 52, 53, 54, 55, 57, 58, 61}
과세유형유효 = {11, 12, 13, 14, 16, 17, 22, 23, 51, 52, 53, 54, 55, 57, 58, 61}


def data_rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] not in (None, "")]


def main():
    out, period = sys.argv[1], sys.argv[2]
    P = importlib.import_module(sys.argv[3]) if len(sys.argv) > 3 else None
    집합 = set(getattr(P, "집합거래처", ()) or ())
    별칭 = getattr(P, "거래처별칭", {}) or {}
    # 거래처 해소를 설정한 프로파일은 남은 빈칸을 반려로 본다. 설정이 아예 없는 프로파일은
    # 빈칸이 의도인지 누락인지 판단할 근거가 없으므로 FAIL 로 단정하지 않고 미설정으로 알린다.
    해소설정 = bool(집합 or 별칭)

    F = []                                  # 반려로 이어지는 것
    W = []                                  # 확인사항으로 넘길 것
    def fail(m): F.append(m)
    def warn(m): W.append(m)

    # ── 1) 일반전표
    gp = os.path.join(out, f"일반전표_업로드_더존_{period}.xlsx")
    wb = openpyxl.load_workbook(gp)
    ws = wb.worksheets[0]
    if ws.title != "일반전표":
        fail(f"일반전표 시트명이 '{ws.title}' 이다. '일반전표' 여야 한다")
    hdr = [c.value for c in ws[1]]
    if hdr != 일반전표_컬럼:
        fail(f"일반전표 컬럼이 양식과 다르다\n      기대 {일반전표_컬럼}\n      실제 {hdr}")
    G = data_rows(ws)

    if not G:
        fail("일반전표에 데이터가 한 줄도 없다 — 빈 파일은 올릴 것이 없다")

    for i, r in enumerate(G, start=2):
        for c, nm in ((0, "월"), (1, "일"), (2, "구분"), (3, "계정과목코드")):
            if not isinstance(r[c], int):
                fail(f"일반전표 {i}행 {nm}={r[c]!r} 숫자형이 아니다")
        # 숫자형인 것과 말이 되는 값인 것은 다르다. 13월 32일도 숫자형이다
        if isinstance(r[0], int) and not 1 <= r[0] <= 12:
            fail(f"일반전표 {i}행 월={r[0]} 1~12 가 아니다")
        if isinstance(r[1], int) and not 1 <= r[1] <= 31:
            fail(f"일반전표 {i}행 일={r[1]} 1~31 이 아니다")
        if r[2] not in 구분유효:
            fail(f"일반전표 {i}행 구분={r[2]} 유효하지 않다(1,2,3,4,5,6)")
        for c, nm in ((8, "차변"), (9, "대변")):
            if r[c] not in (None, "") and not isinstance(r[c], (int, float)):
                fail(f"일반전표 {i}행 {nm}={r[c]!r} 숫자형이 아니다(콤마·통화기호 금지)")
        # 구분이 가리키는 쪽에 금액이 있어야 한다. 반대쪽에 있으면 더존이 반대로 기표한다
        차, 대 = r[8] or 0, r[9] or 0
        if 차 and 대:
            fail(f"일반전표 {i}행 차변({차:,})과 대변({대:,})에 모두 금액이 있다")
        elif not 차 and not 대:
            fail(f"일반전표 {i}행 차변·대변이 모두 비었다")
        elif r[2] in 차변구분 and not 차:
            fail(f"일반전표 {i}행 구분={r[2]}(차변)인데 금액이 대변에 있다")
        elif r[2] in 대변구분 and not 대:
            fail(f"일반전표 {i}행 구분={r[2]}(대변)인데 금액이 차변에 있다")
        if (차 or 대) < 0:
            fail(f"일반전표 {i}행 금액이 음수다({차 or 대:,}) — 반대 구분으로 적어야 한다")
        if r[5] not in (None, ""):
            if not isinstance(r[5], str):
                fail(f"일반전표 {i}행 거래처코드={r[5]!r} 텍스트여야 한다(앞자리 0 유실)")
            elif not r[5].isdigit():
                fail(f"일반전표 {i}행 거래처코드={r[5]!r} 숫자만이어야 한다")

    # 거래처코드가 마스터에 실제로 있나. 더존은 미등록 거래처 한 곳에 파일 전체를 반려한다
    master = None
    cand = sys.argv[4] if len(sys.argv) > 4 else os.path.join(out, "raw", "거래처마스터.xlsx")
    if os.path.isdir(cand):
        for n in ("거래처마스터.xlsx", "partners.xlsx"):
            if os.path.exists(os.path.join(cand, n)):
                cand = os.path.join(cand, n)
                break
    if os.path.exists(cand):
        mws = openpyxl.load_workbook(cand, read_only=True, data_only=True).active
        mit = mws.iter_rows(values_only=True)
        mh = [str(c).strip() for c in next(mit)]
        mi = mh.index("더존거래처코드") if "더존거래처코드" in mh else mh.index("거래처코드")
        master = {str(r[mi]).strip() for r in mit if r[mi] not in (None, "")}
    if master:
        본적없는코드 = {}
        for i, r in enumerate(G, start=2):
            if r[5] not in (None, "") and str(r[5]).strip() not in master:
                본적없는코드.setdefault(str(r[5]).strip(), (i, r[6]))
        for code, (i, nm) in sorted(본적없는코드.items()):
            fail(f"일반전표 {i}행 거래처코드={code} '{nm}' 가 거래처마스터에 없다 "
                 f"— 더존에 먼저 등록해야 파일이 올라간다")
    else:
        warn("거래처마스터를 못 찾아 미등록 거래처 검사를 건너뛰었다 "
             "— 네 번째 인자로 거래처마스터.xlsx 를 주면 본다")

    빈코드 = {}
    for r in G:
        if r[3] in 채권채무계정 and not r[5] and r[6]:
            빈코드.setdefault(str(r[6]), set()).add(r[3])
    미선언 = {k: v for k, v in 빈코드.items() if k not in 집합}
    for k, v in sorted(미선언.items()):
        msg = (f"채권채무 계정 {sorted(v)} 의 거래처 '{k}' 에 거래처코드가 없다 "
               f"— 거래처별칭에 상호를 넣거나 집합거래처로 선언해야 한다")
        (fail if 해소설정 else warn)(msg)
    for k, v in sorted({k: v for k, v in 빈코드.items() if k in 집합}.items()):
        warn(f"'{k}' 는 집합거래처로 선언돼 거래처코드를 비웠다(계정 {sorted(v)}) "
             f"— 거래처별 잔액 관리가 필요하면 더존에 기타거래처 등록")

    # 전표 차대 균형 (파일 전체)
    dtot = sum(r[8] or 0 for r in G); ctot = sum(r[9] or 0 for r in G)
    if dtot != ctot:
        fail(f"일반전표 차변합 {dtot:,} ≠ 대변합 {ctot:,}")

    # 전표별 차대 균형. 전체 합이 맞아도 한 전표가 모자라고 다음이 넘치면
    # 더존은 그 전표를 반려한다. 전체 합계로는 상쇄돼서 안 잡힌다.
    #
    # 양식에 전표번호 열이 없으므로 두 가지를 같이 쓴다.
    #   · 전표는 날짜를 넘지 않는다. 날짜가 바뀌는 지점에서 누적은 0 이어야 한다
    #   · 같은 날짜 안에서는 누적이 0 으로 돌아오는 지점이 전표 경계다
    # 날짜 신호가 없으면 앞뒤 전표가 서로 상쇄될 때 한 전표로 보여 놓치게 된다
    한전표최대 = 40
    잔, 시작행, 경계 = 0, 0, 0
    날 = None
    for i, r in enumerate(G):
        key = (r[0], r[1])
        if 날 is not None and key != 날 and 잔 != 0:
            fail(f"일반전표 {날[0]}월 {날[1]}일분 전표의 차대가 맞지 않는다"
                 f"(누적 {잔:,}) — {시작행+2}행부터 {i+1}행")
            잔, 시작행 = 0, i
        날 = key
        잔 += (r[8] or 0) - (r[9] or 0)
        if 잔 == 0:
            경계 += 1
            시작행 = i + 1
        elif i - 시작행 + 1 > 한전표최대:
            fail(f"일반전표 {시작행+2}행부터 {한전표최대}행이 넘도록 차대가 맞지 않는다"
                 f"(누적 {잔:,}) — 전표 하나의 차변합과 대변합이 어긋났다")
            잔, 시작행 = 0, i + 1
    if 잔 != 0:
        fail(f"일반전표 마지막 전표({시작행+2}행부터)의 차대가 맞지 않는다(누적 {잔:,})")
    else:
        print(f"  일반전표 전표 {경계}개 · 전표별 차대 균형")

    # ── 2) 매입매출전표
    sp = os.path.join(out, f"매입매출전표_업로드_더존_{period}.xlsx")
    wb2 = openpyxl.load_workbook(sp)
    ws2 = wb2.worksheets[0]
    if ws2.title != "매출자료 & 매입자료":
        fail(f"매입매출전표 시트명이 '{ws2.title}' 이다. '매출자료 & 매입자료' 여야 한다")
    hdr2 = [c.value for c in ws2[1]]
    if hdr2 != 매입매출_컬럼:
        fail(f"매입매출전표 컬럼이 양식과 다르다\n      기대 {매입매출_컬럼}\n      실제 {hdr2}")
    S = data_rows(ws2)

    if not S:
        fail("매입매출전표에 데이터가 한 줄도 없다 — 빈 파일은 올릴 것이 없다")

    for i, r in enumerate(S, start=2):
        for c, nm in ((0, "년도"), (1, "월"), (2, "일"), (3, "매입매출구분"), (4, "과세유형"),
                      (11, "공급가액"), (12, "부가세"), (15, "기본계정"), (16, "상대계정")):
            if r[c] in (None, ""):
                if c in (11, 12):      # 공급가액 0·부가세 0 은 있을 수 있다
                    continue
                fail(f"매입매출 {i}행 {nm} 이 비었다")
            elif not isinstance(r[c], int):
                fail(f"매입매출 {i}행 {nm}={r[c]!r} 숫자형이 아니다")
        if isinstance(r[1], int) and not 1 <= r[1] <= 12:
            fail(f"매입매출 {i}행 월={r[1]} 1~12 가 아니다")
        if isinstance(r[2], int) and not 1 <= r[2] <= 31:
            fail(f"매입매출 {i}행 일={r[2]} 1~31 이 아니다")
        if isinstance(r[0], int) and not 2000 <= r[0] <= 2100:
            fail(f"매입매출 {i}행 년도={r[0]} 가 이상하다")
        if r[4] not in 과세유형유효:
            fail(f"매입매출 {i}행 과세유형={r[4]} 양식 코드표에 없다")
        if r[3] not in (1, 2):
            fail(f"매입매출 {i}행 매입매출구분={r[3]} 1(매출)·2(매입) 여야 한다")
        # 유형과 매출입 방향이 어긋나면 부가세 신고서가 반대편에 집계된다
        elif r[3] == 1 and r[4] in 매입유형:
            fail(f"매입매출 {i}행 매출(1)로 적었는데 과세유형 {r[4]} 는 매입 코드다")
        elif r[3] == 2 and r[4] in 매출유형:
            fail(f"매입매출 {i}행 매입(2)로 적었는데 과세유형 {r[4]} 는 매출 코드다")
        # 11열 사업자번호 — 텍스트면 더존이 받지 않는다
        biz = r[10]
        if biz not in (None, ""):
            if not isinstance(biz, int):
                fail(f"매입매출 {i}행 사업자번호={biz!r} 숫자형이어야 한다(텍스트면 실패)")
            elif len(str(biz)) != 10:
                warn(f"매입매출 {i}행 사업자번호 {biz} 가 10자리가 아니다 — 주민번호·앞자리0 예외는 더존 개별등록")
        else:
            if r[4] == 12:
                warn(f"매입매출 {i}행 영세(12) '{r[9]}' 사업자번호 없음 — 국외 거래처는 정상. 더존 거래처 등록 확인")
            elif r[4] in (57, 58, 61):
                fail(f"매입매출 {i}행 카드·현금({r[4]}) '{r[9]}' 사업자번호가 없다 "
                     f"— 사업자번호 없는 카드매입은 공제 대상이 아니다. 일반전표로 보내야 한다")
            else:
                fail(f"매입매출 {i}행 과세유형 {r[4]} '{r[9]}' 사업자번호가 비었다 — 업로드가 막힌다")
        if r[6] not in (None, "") and not isinstance(r[6], str):
            fail(f"매입매출 {i}행 신용카드거래처코드={r[6]!r} 텍스트여야 한다(앞자리 0 유실)")
        if r[4] == 54 and r[5] in (None, ""):
            fail(f"매입매출 {i}행 불공(54) 인데 6열 불공제사유가 비었다")

    # ── 3) 엔진이 남긴 미해결 거래처
    miss = os.path.join(out, f"_거래처미해소_{period}.json")
    if os.path.exists(miss):
        m = json.load(io.open(miss, encoding="utf-8"))
        for k, v in m.items():
            msg = (f"거래처라벨 '{k}' 을 거래처마스터에서 찾지 못했다(계정 {v}) "
                   f"— 프로파일 거래처별칭·집합거래처를 채워야 한다")
            (fail if 해소설정 else warn)(msg)

    # ── 리포트
    print("=== 더존 업로드 사전검증 ===")
    print(f"  일반전표 {len(G)}행 · 매입매출전표 {len(S)}행")
    ccfill = sum(1 for r in G if r[3] in 채권채무계정 and r[5])
    cctot = sum(1 for r in G if r[3] in 채권채무계정)
    bizfill = sum(1 for r in S if r[10])
    print(f"  채권채무 거래처코드 {ccfill}/{cctot} · 매입매출 사업자번호 {bizfill}/{len(S)}")
    if not 해소설정:
        print("  [미설정] 이 프로파일에 거래처별칭·집합거래처가 없다. 거래처코드 빈칸을 "
              "반려로 판정하지 않았다 — 두 맵을 채우면 이 검증이 엄격해진다")
    for m in W:
        print(f"  [확인] {m}")
    for m in F:
        print(f"  [FAIL] {m}")
    # 결과를 남긴다. 확인사항 문서가 이것을 읽어 실제 결과를 적는다.
    # 안 남기면 문서가 「전부 PASS」를 지어내게 된다
    json.dump({"스크립트": "verify_upload.py", "기간": period,
               "pass": not F, "fail": F, "확인": W,
               "일반전표행": len(G), "매입매출행": len(S), "전표수": 경계,
               "채권채무_거래처코드": [ccfill, cctot],
               "사업자번호": [bizfill, len(S)],
               "미등록거래처_검사": bool(master)},
              io.open(os.path.join(out, f"_검증_업로드_{period}.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    if F:
        print(f"업로드 사전검증 FAIL {len(F)}건 — 올리면 반려된다")
        sys.exit(1)
    print("업로드 사전검증 ALL PASS" + (f" (확인사항 {len(W)}건)" if W else ""))


if __name__ == "__main__":
    main()

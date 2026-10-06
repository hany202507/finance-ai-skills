# -*- coding: utf-8 -*-
"""부문손익·CC 분석 모듈 (결산 스킬 확장, additive — close_engine과 독립).

원자료(은행 거래내역·카드 승인내역·세금계산서)를 규칙으로 태깅해
수익부문별 매출 × 코스트센터(CC)별 비용의 월별 부문손익을 산출한다.

원칙
- 현금주의: 매출=은행 입금, 비용=은행 출금+카드 확정금액. 세금계산서는 태깅 근거로만
  (발생주의 이중집계 방지). 현금영수증은 정산 입금과 중복이라 제외.
- 완전 분할(partition): 모든 행이 정확히 한 분류에 속한다(미분류 포함) →
  분류 합계 = 원본 합계가 구조적으로 보장(1원 tie-out을 스크립트가 검증).
- 카드대금·해외카드 은행 출금은 '대체'(비용 아님) — 카드 지출은 카드내역에서 잡는다.
- 판단성(개인명 이체·용도 미상)은 FLAG 시트로 — 단정하지 않고 전문가 확인 큐.
- 산출 워크북의 부문손익 표는 SUMIFS 살아있는 수식(태그 시트를 고치면 재계산).

사용:
  python segment_pnl.py <raw_dir> <out_xlsx> [rules_module]
rules_module 생략 시 내장 RULES(가상 예시회사 GF, 식품 제조·유통) 사용.
새 회사는 RULES와 같은 구조의 dict를 가진 .py를 만들어 모듈명을 넘긴다.
"""
import sys, os, re, glob, datetime, importlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ===== 내장 규칙 — 가상 예시회사(GF, 데모용 더미: 식품 제조·유통) =====
# 각 규칙: (정규식, 계정, 분류) — 분류는 매출이면 수익부문, 지출이면 CC.
RULES = {
    '수익부문': ['스마트스토어', '오픈마켓', 'B2B도매', '기타·미분류'],
    'CC': ['생산', '마케팅', '물류', '본사공통', '미분류'],
    '매출': [  # 은행 입금 내용 매칭
        (r'스마트스토어정산', '상품매출', '스마트스토어'),
        (r'지마켓|쿠팡페이|옥션|차액정산', '상품매출', '오픈마켓'),
        (r'후치꼬치|영풍|네모난오렌|넥솔위즈빌|현대그린푸드|볼프강|느티나무|정밀|예스에프앤(?!에스)|제이피푸드', '상품매출', 'B2B도매'),
        (r'이자', '이자수익', '기타·미분류'),
    ],
    '지출': [  # 은행 출금 내용 + 카드 가맹점 공용 매칭 (위에서부터 첫 매칭)
        # 대체·비비용성 (비용 집계 제외)
        (r'KB카드출금|카드출금|마스타해외승인출금', '카드대금(대체)', '대체'),
        (r'^국세|국민연금|건강보험|고용보험|산재', '제세공과(예수금성)', '대체'),
        # 원재료·생산
        (r'예스에프앤|제이피(푸드)?|햇살푸드|태원식품|하나태원|신화케이푸드|리안|보감|농부누리|영농|축산|웰스토리|푸드|식품|에프에스', '원재료매입', '생산'),
        # 마케팅 (광고·판매수수료·외주)
        (r'FACEBK|META|오픈엑스|OPENX|광고', '광고선전비', '마케팅'),
        (r'알파브라더스', '지급수수료(외주·대행)', '마케팅'),
        (r'네이버페이|이니시스|NICE|네이버파이낸셜|지마켓 ?수수료', '지급수수료(판매)', '마케팅'),
        # 물류
        (r'우정사업본부|우체국|대한통운|한진|택배|로지스', '운반비', '물류'),
        # 본사공통
        (r'퀸즈에비뉴|관리단|관리비', '임차·관리비', '본사공통'),
        (r'세무사|법무사|노무사', '지급수수료(자문)', '본사공통'),
        (r'Adobe|SLACK|구글플레이|유튜브|네이버플러스|카카오(?!T)|GOOGLE|Google', '지급수수료(SaaS)', '본사공통'),
        (r'카카오 ?T|택시|주차|고속버스|KOBUS', '여비교통비', '본사공통'),
        (r'캐피탈', '리스료(확인)', '본사공통'),
        (r'ＳＫＢ|SKB|브로드밴드|통신', '통신비', '본사공통'),
        (r'하나카드기업|송금수수료|법원행정처', '지급수수료(금융·행정)', '본사공통'),
        (r'컴포즈|메가 ?MGC|버거킹|순대국|오봉집|부뚜막|써브웨이|롯데리아|돈까스|닭갈비|네이처볼|백년옥|쿠팡이츠|배달|코페이|키오스크|국수', '복리후생비(식대)', '본사공통'),
        (r'상자|포장', '소모품비(포장)', '물류'),
        (r'씨유|CU|지에스25|GS25|이마트|컬리|쿠팡|마트|편의점|다이소', '소모품비', '본사공통'),
    ],
}

E_HDR = Font(bold=True, color='FFFFFF')
E_FILL = PatternFill('solid', fgColor='1E3A6B')
E_SUB = PatternFill('solid', fgColor='F4F8FD')


def find_file(raw_dir, kw):
    hits = [f for f in glob.glob(os.path.join(raw_dir, '*.xlsx')) if kw in os.path.basename(f)]
    return hits[0] if hits else None


def month_of(v):
    if isinstance(v, datetime.datetime) or isinstance(v, datetime.date):
        return f'{v.year:04d}-{v.month:02d}'
    s = str(v)
    return s[:7].replace('.', '-').replace('/', '-')


def classify(text, rules):
    for pat, acct, seg in rules:
        if re.search(pat, text):
            return acct, seg
    return None, None


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return list(wb.active.iter_rows(min_row=2, values_only=True))


def run(raw_dir, out_path, rules=RULES):
    bank_f = find_file(raw_dir, '은행')
    card_f = find_file(raw_dir, '카드')
    inv_f = find_file(raw_dir, '세금계산서')
    assert bank_f and card_f, f'은행/카드 파일을 {raw_dir}에서 찾지 못함'

    months = set()
    tag_bank, tag_card, flags = [], [], []

    # --- 은행: 입금=매출(수익부문) / 출금=비용(CC) 또는 대체 ---
    for r in load_rows(bank_f):
        if r[0] is None:
            continue
        m = month_of(r[0]); desc = str(r[8] or '').strip()
        inn = int(r[5] or 0); out = int(r[6] or 0)
        months.add(m)
        if inn:
            acct, seg = classify(desc, rules['매출'])
            conf = '높음'
            if seg is None:
                acct, seg, conf = '기타수입(확인)', '기타·미분류', '낮음'
                flags.append(('은행입금', m, desc, inn, '매출 부문 미분류 — 거래처·성격 확인'))
            tag_bank.append((m, '매출', seg, acct, desc, inn, conf))
        if out:
            acct, seg = classify(desc, rules['지출'])
            conf = '높음'
            if seg is None:
                if re.match(r'^[가-힣]{2,4}$', desc):
                    acct, seg, conf = '인건비 추정(확인)', '미분류', '낮음'
                    flags.append(('은행출금', m, desc, out, '개인명 이체 — 급여/외주/가지급 확인 필요'))
                else:
                    acct, seg, conf = '미분류', '미분류', '낮음'
                    flags.append(('은행출금', m, desc, out, '용도 미상 — 확인 필요'))
            kind = '대체' if seg == '대체' else '비용'
            tag_bank.append((m, kind, seg, acct, desc, out, conf))

    # --- 카드: 전건 비용(확정금액 기준) ---
    for r in load_rows(card_f):
        if r[0] is None:
            continue
        m = month_of(r[0]); mer = str(r[6] or '').strip()
        amt = int(r[12] or 0)
        if amt == 0:
            continue
        months.add(m)
        acct, seg = classify(mer, rules['지출'])
        conf = '높음'
        if seg is None:
            acct, seg, conf = '미분류', '미분류', '낮음'
            flags.append(('카드', m, mer, amt, '가맹점 미분류 — 계정·CC 확인'))
        if seg == '대체':
            seg, conf = '미분류', '낮음'
        tag_card.append((m, '비용', seg, acct, mer, amt, conf))

    # --- 세금계산서: 참고 태깅(발생주의 참고 — 부문손익 집계에는 미포함) ---
    tag_inv = []
    if inv_f:
        for r in load_rows(inv_f):
            if r[0] is None:
                continue
            m = month_of(r[0]); kind = str(r[2] or ''); party = str(r[4] or ''); item = str(r[6] or '')
            supply = int(r[7] or 0)
            if kind == '매출':
                acct, seg = classify(party, rules['매출'])
                tag_inv.append((m, '매출(발생)', seg or 'B2B도매', '상품매출', f'{party} | {item}', supply))
            else:
                acct, seg = classify(party + ' ' + item, rules['지출'])
                tag_inv.append((m, '매입(발생)', seg or '미분류', acct or '미분류', f'{party} | {item}', supply))

    months = sorted(months)
    wb = openpyxl.Workbook()

    def tag_sheet(name, rows, cols):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = E_HDR; c.fill = E_FILL
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = 'A2'
        for i, wd in enumerate([9, 8, 12, 18, 34, 14, 8][:len(cols)]):
            ws.column_dimensions[get_column_letter(i + 1)].width = wd
        return ws

    tag_sheet('태그_은행', tag_bank, ['월', '구분', '부문·CC', '계정', '적요', '금액', '신뢰도'])
    tag_sheet('태그_카드', tag_card, ['월', '구분', '부문·CC', '계정', '가맹점', '금액', '신뢰도'])
    if tag_inv:
        tag_sheet('참고_세금계산서', tag_inv, ['월', '구분', '부문·CC', '계정', '거래처|품명', '공급가액'])
    tag_sheet('FLAG_확인필요', flags, ['출처', '월', '적요·가맹점', '금액', '확인 요청'])

    # --- 부문손익 (SUMIFS 살아있는 수식) ---
    ws = wb.active; ws.title = '부문손익'
    ws.sheet_view.showGridLines = False
    ws['B2'] = '부문손익 — 수익부문 × 코스트센터 (현금주의: 입금·출금·카드 기준)'
    ws['B2'].font = Font(bold=True, size=14, color='16264A')
    ws['B3'] = '태그 시트의 분류를 고치면 이 표가 다시 계산됩니다 (SUMIFS). 세금계산서는 참고(발생주의) 시트.'
    ws['B3'].font = Font(size=10, color='6B7A90')

    def block(r0, title, kind, segs, sheets):
        ws.cell(r0, 2, title).font = Font(bold=True, size=11, color='2F6FB0')
        ws.cell(r0 + 1, 2, '부문').font = Font(bold=True)
        for j, m in enumerate(months):
            c = ws.cell(r0 + 1, 3 + j, m); c.font = Font(bold=True); c.fill = E_SUB
        c = ws.cell(r0 + 1, 3 + len(months), '합계'); c.font = Font(bold=True); c.fill = E_SUB
        for i, seg in enumerate(segs):
            ws.cell(r0 + 2 + i, 2, seg)
            for j, m in enumerate(months):
                parts = [f"SUMIFS('{sh}'!F:F,'{sh}'!A:A,\"{m}\",'{sh}'!B:B,\"{kind}\",'{sh}'!C:C,\"{seg}\")" for sh in sheets]
                ws.cell(r0 + 2 + i, 3 + j, '=' + '+'.join(parts)).number_format = '#,##0'
            a = get_column_letter(3); b = get_column_letter(2 + len(months))
            rr = r0 + 2 + i
            ws.cell(rr, 3 + len(months), f'=SUM({a}{rr}:{b}{rr})').number_format = '#,##0'
        rt = r0 + 2 + len(segs)
        ws.cell(rt, 2, '계').font = Font(bold=True)
        for j in range(len(months) + 1):
            col = get_column_letter(3 + j)
            ws.cell(rt, 3 + j, f'=SUM({col}{r0 + 2}:{col}{rt - 1})').number_format = '#,##0'
            ws.cell(rt, 3 + j).font = Font(bold=True)
        return rt

    r = block(5, '매출 — 수익부문별 (은행 입금)', '매출', rules['수익부문'], ['태그_은행'])
    r2 = block(r + 2, '비용 — 코스트센터별 (은행 출금 + 카드)', '비용', [s for s in rules['CC']], ['태그_은행', '태그_카드'])
    r3 = block(r2 + 2, '(참고) 대체성 출금 — 비용 아님 (카드대금·제세 등)', '대체', ['대체'], ['태그_은행'])
    rr = r3 + 2
    ws.cell(rr, 2, '순영업현금(매출계−비용계)').font = Font(bold=True, color='16264A')
    for j in range(len(months) + 1):
        col = get_column_letter(3 + j)
        ws.cell(rr, 3 + j, f'={col}{r}-{col}{r2}').number_format = '#,##0'
        ws.cell(rr, 3 + j).font = Font(bold=True)
    ws.column_dimensions['B'].width = 34
    for j in range(len(months) + 1):
        ws.column_dimensions[get_column_letter(3 + j)].width = 14

    # --- 검증(1원 tie-out): 원본 합계 = 태그 완전분할 합계 ---
    bank_in_raw = sum(int(x[5] or 0) for x in load_rows(bank_f) if x[0] is not None)
    bank_out_raw = sum(int(x[6] or 0) for x in load_rows(bank_f) if x[0] is not None)
    card_raw = sum(int(x[12] or 0) for x in load_rows(card_f) if x[0] is not None)
    t_in = sum(x[5] for x in tag_bank if x[1] == '매출')
    t_out = sum(x[5] for x in tag_bank if x[1] in ('비용', '대체'))
    t_card = sum(x[5] for x in tag_card)
    assert t_in == bank_in_raw, f'입금 tie-out 실패 {t_in} != {bank_in_raw}'
    assert t_out == bank_out_raw, f'출금 tie-out 실패 {t_out} != {bank_out_raw}'
    assert t_card == card_raw, f'카드 tie-out 실패 {t_card} != {card_raw}'
    ws = wb.create_sheet('검증')
    ws.append(['체크', '원본 합계', '태그 합계(수식)', '판정(수식)'])
    for c in ws[1]:
        c.font = E_HDR; c.fill = E_FILL
    ws.append(['은행 입금 = 매출 분할합', bank_in_raw, "=SUMIFS(태그_은행!F:F,태그_은행!B:B,\"매출\")", '=IF(B2=C2,"PASS","FAIL")'])
    ws.append(['은행 출금 = 비용+대체 분할합', bank_out_raw, "=SUMIFS(태그_은행!F:F,태그_은행!B:B,\"비용\")+SUMIFS(태그_은행!F:F,태그_은행!B:B,\"대체\")", '=IF(B3=C3,"PASS","FAIL")'])
    ws.append(['카드 확정금액 = 카드 분할합', card_raw, '=SUM(태그_카드!F:F)', '=IF(B4=C4,"PASS","FAIL")'])
    for i in (2, 3, 4):
        ws.cell(i, 2).number_format = '#,##0'; ws.cell(i, 3).number_format = '#,##0'
    ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 18; ws.column_dimensions['C'].width = 18

    wb.save(out_path)
    return dict(months=months, bank_in=bank_in_raw, bank_out=bank_out_raw, card=card_raw,
                flags=len(flags), bank_rows=len(tag_bank), card_rows=len(tag_card))


if __name__ == '__main__':
    raw_dir, out_path = sys.argv[1], sys.argv[2]
    rules = RULES
    if len(sys.argv) > 3:
        sys.path.insert(0, os.path.dirname(sys.argv[3]) or '.')
        rules = importlib.import_module(os.path.splitext(os.path.basename(sys.argv[3]))[0]).RULES
    info = run(raw_dir, out_path, rules)
    print('OK', info)

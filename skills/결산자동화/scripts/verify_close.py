# -*- coding: utf-8 -*-
"""독립 검증 — 결산확정 워크북을 `formulas`로 실제 재계산하여 셀값을 대조.

설계서 §9 검증 12체크 중 '수식 재계산' 축(하드코딩 아님 + 대차/현금 대사)을 담당.
raw 재집계 축은 별도 서브에이전트(설계서 §9)가 수행한다.

사용:
  python verify_close.py <결산확정_xlsx> [expected_json]
expected_json 없으면 내부 정합성(대차·현금·순이익 연결)만 검증.
"""
import sys, os, json, tempfile, glob
import formulas
import openpyxl


def count_formulas(path):
    """원본 파일의 수식 개수. 재계산본이 아니라 원본을 봐야 한다.

    재무제표를 수식으로 엮는 것이 이 산출물의 요건이다(설계서 §4-7).
    상수로 적어 넣으면 분개를 고쳐도 안 따라오고, 그건 결산확정 워크북이 아니다.
    """
    wb = openpyxl.load_workbook(path, data_only=False)
    n = 0
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
    return n

def recalc_values(path):
    """formulas로 재계산한 워크북을 임시로 써서 값만 읽어온다."""
    xl = formulas.ExcelModel().loads(path).finish()
    xl.calculate()
    outdir = tempfile.mkdtemp()
    xl.write(dirpath=outdir)
    f = glob.glob(os.path.join(outdir, "*.xlsx"))[0]
    wb = openpyxl.load_workbook(f, data_only=True)
    return wb


# Excel 표준 오류와 동적 배열/외부 데이터 오류. 설명문에 포함된 문자열은 제외한다.
EXCEL_ERRORS = frozenset({
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!",
    "#UNKNOWN!", "#CONNECT!", "#BUSY!", "#PYTHON!",
})


def workbook_errors(wb):
    """숨김 시트를 포함한 모든 셀의 실제 오류값을 수집한다."""
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "e" or (isinstance(value, str) and value in EXCEL_ERRORS):
                    errors.append({"시트": ws.title, "셀": cell.coordinate, "오류": str(value)})
    return errors


def finish(path, checks, nf, errors, summary=None):
    """성공·실패 모두 새 JSON을 남겨 이전 PASS 결과가 재사용되지 않게 한다."""
    print("=== 재계산 검증 ===")
    for name, result, detail in checks:
        print(f"  [{result}] {name}  {detail}")
    if summary:
        print("\n" + summary)
    allpass = all(result == "PASS" for _, result, _ in checks)
    directory = os.path.dirname(os.path.abspath(path))
    period = os.path.basename(path).replace("결산확정_", "").replace(".xlsx", "")
    result_path = os.path.join(directory, f"_검증_워크북_{period}.json")
    result = {"스크립트": "verify_close.py", "기간": period, "pass": allpass,
              "수식개수": nf, "오류셀": errors,
              "검사": [{"항목": name, "결과": res, "상세": detail}
                       for name, res, detail in checks]}
    # 결과 저장에 실패하면 성공으로 종료하지 않는다.
    with open(result_path, "w", encoding="utf-8") as stream:
        json.dump(result, stream, ensure_ascii=False, indent=1)
    print("ALL PASS" if allpass else "FAIL 있음")
    raise SystemExit(0 if allpass else 1)


def find_row(ws, col, needle):
    for r in range(1, ws.max_row+1):
        if str(ws.cell(r, col).value).strip() == needle:
            return r
    return None

def cellval(ws, label_col, label, val_col):
    r = find_row(ws, label_col, label)
    return ws.cell(r, val_col).value if r else None

def main():
    path = sys.argv[1]
    exp = {}
    if len(sys.argv) > 2 and os.path.exists(sys.argv[2]):
        exp = json.load(open(sys.argv[2], encoding="utf-8"))
    nf = count_formulas(path)
    source = openpyxl.load_workbook(path, data_only=False)
    try:
        errors = workbook_errors(source)
    finally:
        source.close()
    try:
        wb = recalc_values(path)
    except Exception as exc:
        checks = [("수식 재계산", "FAIL", f"{type(exc).__name__}: {exc}")]
        checks.extend(("워크북 오류 셀", "FAIL", f"{e['시트']}!{e['셀']} = {e['오류']}")
                      for e in errors)
        finish(path, checks, nf, errors)
    # 원본 상수 오류도 보존하고, 재계산 후 새로 생긴 오류를 전 시트에서 검사한다.
    unique_errors = {(e["시트"], e["셀"], e["오류"]): e
                     for e in errors + workbook_errors(wb)}
    errors = list(unique_errors.values())
    if errors:
        checks = [("워크북 오류 셀", "FAIL", f"{e['시트']}!{e['셀']} = {e['오류']}")
                  for e in errors]
        wb.close()
        finish(path, checks, nf, errors)
    ts = wb["시산표"]; is_ = wb["손익계산서"]; bs = wb["재무상태표"]; cf = wb["현금흐름표_직접법"]

    # 시산표 기말합(마지막 행 H)
    ts_tot_h = ts.cell(ts.max_row, 8).value
    ni = cellval(is_, 1, "X. 당기순이익", 2)
    sales = cellval(is_, 1, "I. 매출액", 2)
    op = cellval(is_, 1, "V. 영업이익", 2)
    at = cellval(bs, 1, "자산총계", 4)
    le = cellval(bs, 1, "부채와자본총계", 4)
    endcash = cellval(cf, 2, "기말현금", 3)

    R = lambda x: None if x is None else round(float(x))
    checks = [("워크북 오류 셀 없음", "PASS", "전 시트 오류 셀 0개")]
    def chk(name, cond, detail=""):
        checks.append((name, "PASS" if cond else "FAIL", detail))

    chk("시산표 기말합=0(대차균형)", R(ts_tot_h) == 0, f"H합계={R(ts_tot_h)}")
    chk("재무상태표 자산=부채+자본", R(at) == R(le), f"자산={R(at):,} / 부채+자본={R(le):,}")
    # 현금흐름 기말현금 = BS 보통예금(103) 기말
    r103 = find_row(ts, 1, "103")
    bs103 = ts.cell(r103, 8).value if r103 else None
    chk("현금흐름 기말현금=BS 보통예금", R(endcash) == R(bs103), f"CF={R(endcash):,} / BS={R(bs103):,}")
    # 수식이 없으면 재무제표가 아니라 숫자를 적어 넣은 표다
    chk("수식 연결(하드코딩 아님)", nf > 0, f"수식 {nf:,}개")

    # 손익 항등식. 셀이 비어 있지 않은지가 아니라 숫자가 맞는지 본다
    cogs = cellval(is_, 1, "II. 매출원가", 2)
    sga = cellval(is_, 1, "IV. 판매비와관리비", 2)
    oi = cellval(is_, 1, "VI. 영업외수익", 2)
    oe = cellval(is_, 1, "VII. 영업외비용", 2)
    tax = cellval(is_, 1, "IX. 법인세비용", 2)
    Z = lambda x: 0 if x is None else round(float(x))
    계산 = Z(sales) - Z(cogs) - Z(sga) + Z(oi) - Z(oe) - Z(tax)
    chk("당기순이익=매출-원가-판관비+영업외-법인세", ni is not None and R(ni) == 계산,
        f"손익계산서 {R(ni) if ni is not None else None} / 구성요소 합 {계산:,}")
    chk("영업이익=매출-원가-판관비", op is not None and R(op) == Z(sales) - Z(cogs) - Z(sga),
        f"영업이익 {R(op) if op is not None else None} / {Z(sales)-Z(cogs)-Z(sga):,}")

    # 시산표에서 손익계정만 다시 더해 순이익과 맞춘다. 손익계산서를 안 거치는 대사다
    pl = 0
    found = 0
    for r in range(2, ts.max_row + 1):
        v = ts.cell(r, 1).value
        try:
            code = int(str(v).strip())
        except (TypeError, ValueError):
            continue
        if 400 <= code < 1000:
            pl += Z(ts.cell(r, 8).value)
            found += 1
    chk("순이익=-(시산표 손익계정 기말 합)", found > 0 and ni is not None and R(ni) == -pl,
        f"손익계정 {found}개 합 {pl:,} / 순이익 {R(ni) if ni is not None else None}")

    if exp:
        for k, label in [("순이익","당기순이익"),("기말현금","기말현금"),
                         ("자산총계","자산총계"),("매출액","매출액"),("영업이익","영업이익")]:
            if k in exp:
                got = {"순이익":R(ni),"기말현금":R(endcash),"자산총계":R(at),
                       "매출액":R(sales),"영업이익":R(op)}[k]
                chk(f"[기대값] {label}={exp[k]:,}", got == exp[k], f"산출={got:,}")

    summary = (f"순이익={R(ni):,} 기말현금={R(endcash):,} 자산총계={R(at):,} "
               f"매출액={R(sales):,} 영업이익={R(op):,}")
    wb.close()
    finish(path, checks, nf, errors, summary)


if __name__ == "__main__":
    main()

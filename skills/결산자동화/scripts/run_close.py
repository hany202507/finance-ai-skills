# -*- coding: utf-8 -*-
"""결산 엔진 드라이버 — 워크북 + 더존 업로드 2종 생성 + 앵커 출력.

사용:
  python run_close.py <raw_dir> <out_dir> <기간라벨> [profile_module]
"""
import sys, os, importlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from close_engine import Engine, to_int
import build_workbook as BW

WON = '#,##0'
HDR = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="305496")

def hdr(ws, heads):
    ws.append(heads)
    for c in range(1, len(heads)+1):
        ws.cell(1, c).font = HDR; ws.cell(1, c).fill = HFILL
        ws.cell(1, c).alignment = Alignment(horizontal="center")

# ---------------- 더존 일반전표
def upload_general(eng, path):
    P = eng.P
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "일반전표"
    hdr(ws, ["월","일","구분","계정과목코드","계정과목명","거래처코드","거래처","적요","차변","대변"])
    def code_of(party, acct=None):
        return eng.party_code(party, acct)
    # 전표 그룹핑
    from collections import OrderedDict
    groups = OrderedDict()
    for l in eng.J:
        groups.setdefault(l["no"], []).append(l)
    for no, lines in groups.items():
        g = lines[0]["gubun"]
        if g in ("전기이월", "매출", "매입"): continue
        if g == "카드" and any(x["code"] == 135 for x in lines): continue  # 공제→매입매출
        for l in lines:
            mmm, dd = int(l["date"][5:7]), int(l["date"][8:10])
            차, 대 = to_int(l["debit"]), to_int(l["credit"])
            if 차 < 0:
                차, 대 = 0, 대 - 차
            if 대 < 0:
                차, 대 = 차 - 대, 0
            gubun = 3 if 차 else 4
            dcode = P.더존코드변환.get(l["code"], l["code"])
            ws.append([mmm, dd, gubun, dcode, l["acct"], code_of(l["party"], dcode), l["party"],
                       l["desc"], 차 or None, 대 or None])
    # 결산조정 → 일반전표(5결산차/6결산대)
    # 더존은 음수 금액을 받지 않는다. 음수 차변은 대변으로 뒤집어야 같은 분개가 된다
    for a in eng.ADJ:
        mmm, dd = int(str(a[1])[5:7]), int(str(a[1])[8:10])
        차, 대 = to_int(a[4]), to_int(a[5])
        if 차 < 0:
            차, 대 = 0, 대 - 차
        if 대 < 0:
            차, 대 = 차 - 대, 0
        gubun = 5 if 차 else 6
        dcode = P.더존코드변환.get(a[2], a[2])
        ws.append([mmm, dd, gubun, dcode, a[3], "", "", a[6], 차 or None, 대 or None])
    for r in range(2, ws.max_row+1):
        ws.cell(r, 9).number_format = WON; ws.cell(r, 10).number_format = WON
        cc = ws.cell(r, 6)
        if cc.value: cc.value = str(cc.value); cc.number_format = "@"   # 거래처코드 텍스트(앞자리 0)
    dtot = sum(to_int(ws.cell(r,9).value) for r in range(2, ws.max_row+1))
    ctot = sum(to_int(ws.cell(r,10).value) for r in range(2, ws.max_row+1))
    wb.save(path)
    return dtot, ctot

# ---------------- 더존 매입매출
def upload_salespur(eng, path):
    P = eng.P
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "매출자료 & 매입자료"
    hdr(ws, ["년도","월","일","매입매출구분(1-매출/2-매입)","과세유형","불공제사유","신용카드거래처코드",
             "신용카드사명","신용카드(가맹점)번호","거래처명","사업자(주민)등록번호","공급가액","부가세",
             "품명","전자세금(1.전자)","기본계정","상대계정","현금영수증 승인번호"])
    txt_cols = set()  # 신용카드거래처코드는 텍스트
    suptot = vattot = 0
    # 세금계산서 per-invoice
    for r in eng.tax:
        d = str(r[0]); y, mo, dd = int(d[:4]), int(d[5:7]), int(d[8:10])
        mi = r[1]; item = r[5]; sup = to_int(r[6]); vat = to_int(r[7]); biz = to_int(r[4])
        if mi == "매출":
            gubun, gtype = 1, 11
            base = P.매출품목_계정[item]; sang = 108
        else:
            gubun, gtype = 2, 51
            code, name, s, cf = P.매입품목_계정[item]; base = code; sang = s
        if not biz:   # 원천에 사업자번호가 없으면 마스터에서 보충한다
            biz = eng.party_biz(r[3], base)
        ws.append([y, mo, dd, gubun, gtype, "", "", "", "", r[3], biz or None,
                   sup, vat, item, 1, base, sang, ""])
        suptot += sup; vattot += vat
    # 채널매출 — 신용카드 매출이라 세금계산서가 아니다. 과세유형은 프로파일이 정한다.
    # 여기를 빼면 시산표는 맞는데 부가세 신고서 (3)란이 통째로 빈다.
    ch_gtype = getattr(P, "채널매출_과세유형", {})
    from collections import OrderedDict as _OD
    chagg = _OD()
    for x in eng._prep_channel():
        if x["sup"] == 0 or x["taxkind"] not in ch_gtype: continue
        k = (x["ym"], x["ch"], x["taxkind"])
        a = chagg.setdefault(k, [0, 0, x["chname"], x["acct"], x["party"]])
        a[0] += x["sup"]; a[1] += x["vat"]
    for (ym, ch, taxkind), (sup, vat, chname, u, party) in chagg.items():
        y, mo = int(str(ym)[:4]), int(str(ym)[5:7])
        dd = P.게시일["매출"]
        # 11열을 비워 두면 더존이 거래처를 못 붙인다. 플랫폼은 마스터에 상호·사업자번호가 있다.
        # 수출(해외바이어)은 국내 사업자번호가 없는 것이 정상이라 빈칸으로 두고 확인사항에 올린다
        pbiz = eng.party_biz(party, u["상대코드"])
        ws.append([y, mo, dd, 1, ch_gtype[taxkind], "", "", "", "", party, pbiz or None,
                   sup, vat, f"{chname} {taxkind} 매출", "", u["수익코드"], u["상대코드"], ""])
        suptot += sup; vattot += vat
    # 카드 공제(식대·포장) per month per shop  → 과세유형 57
    from collections import OrderedDict
    cardg = OrderedDict()
    for l in eng.J:
        if l["gubun"] == "카드":
            cardg.setdefault(l["no"], []).append(l)
    cardsa = eng.card[0][1] if eng.card else "하나법인"
    gaeng_biz = {r[3]: to_int(r[5]) for r in eng.card if r[5] not in (None, "")}
    for no, lines in cardg.items():
        if not any(x["code"] == 135 for x in lines): continue  # 공제만
        exp = [x for x in lines if x["code"] not in (135, 253)][0]
        vatl = [x for x in lines if x["code"] == 135][0]
        shop = exp["party"]
        d = exp["date"]; y, mo, dd = int(d[:4]), int(d[5:7]), int(d[8:10])
        ccode, mbiz = eng.resolve_party(shop, exp["code"])
        biz = gaeng_biz.get(shop) or to_int(mbiz)
        ws.append([y, mo, dd, 2, 57, "", ccode, cardsa, biz or None, shop, to_int(biz) or None,
                   exp["debit"], vatl["debit"], shop, "", exp["code"], 253, ""])
        suptot += exp["debit"]; vattot += vatl["debit"]
    # 숫자형/텍스트 지정
    for r in range(2, ws.max_row+1):
        for cc in (1,2,3,4,5,11,12,13,16,17):
            v = ws.cell(r, cc).value
            if v not in (None, ""): ws.cell(r, cc).value = int(v)
        ws.cell(r,12).number_format = WON; ws.cell(r,13).number_format = WON
        cc7 = ws.cell(r,7)
        if cc7.value: cc7.value = str(cc7.value); cc7.number_format = "@"
    wb.save(path)
    return suptot, vattot

# ---------------- 미분류 목록
def unclassified_list(eng, path):
    """전문가가 볼 목록. 규칙에 안 걸린 거래를 건별로 낸다.

    가계정에 넣어 두었으므로 금액은 맞지만 계정이 정해진 것이 아니다.
    이 목록이 강의안 3차·4차 실습(유형화·확인사항)의 입력이다.
    """
    rows = getattr(eng, "unclassified_rows", [])
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "미분류"
    hdr(ws, ["거래일자", "계좌", "입출금", "금액", "적요", "왜 미분류인가", "전문가 판단"])
    for r in rows:
        ws.append(list(r) + [""])
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 4).number_format = WON
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 46
    ws.column_dimensions["G"].width = 22
    wb.save(path)
    return len(rows)


# ================================================================= main
def main():
    raw, out, period = sys.argv[1], sys.argv[2], sys.argv[3]
    pmod = sys.argv[4] if len(sys.argv) > 4 else "profile_sample"
    P = importlib.import_module(pmod)
    os.makedirs(out, exist_ok=True)
    eng = Engine(raw, period, P).run()
    wbpath = os.path.join(out, f"결산확정_{period}.xlsx")
    BW.build(eng, wbpath, period)
    g = upload_general(eng, os.path.join(out, f"일반전표_업로드_더존_{period}.xlsx"))
    s = upload_salespur(eng, os.path.join(out, f"매입매출전표_업로드_더존_{period}.xlsx"))
    n미 = unclassified_list(eng, os.path.join(out, f"미분류목록_{period}.xlsx"))
    # 거래처마스터에서 못 찾은 라벨을 업로드 사전검증에 넘긴다. 조용히 빈칸으로 두지 않는다
    import json
    mp = os.path.join(out, f"_거래처미해소_{period}.json")
    if eng.party_miss:
        json.dump({k: sorted(x for x in v if x) for k, v in eng.party_miss.items()},
                  open(mp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    elif os.path.exists(mp):
        os.remove(mp)
    print("=== 결산 엔진 앵커 ===")
    print(f"분개장 라인 {len(eng.J)} / 전표 {eng.no} / 조정 {len(eng.ADJ)}")
    sm = eng.summary()
    print(f"매출액(계)    {sm['sales']:>15,}")
    print(f"  세금계산서  {eng._tax_sales():>15,}")
    print(f"  채널매출    {eng._channel_sales():>15,}")
    print(f"매출원가      {eng.cogs:>15,}")
    print(f"미지급세금    {eng.vat_payable:>15,}")
    print(f"카드공제세액  {eng.card_deduct:>15,}")
    print(f"일반전표 차={g[0]:,} 대={g[1]:,} (일치 {g[0]==g[1]})")
    print(f"매입매출 공급={s[0]:,} 세액={s[1]:,}")
    # 강의안 2-8: 분개 건수와 미분류 건수의 합이 전체와 맞아야 한다
    n통장 = len(eng.bank)
    n규칙 = sum(1 for r in eng.bank if r[6] and r[6] != "미분류")
    if n규칙 + n미 != n통장:
        raise SystemExit(f"통장 {n통장:,}건 ≠ 규칙 {n규칙:,} + 미분류 {n미:,} — 빠진 거래가 있다")
    print(f"통장 {n통장:,}건 = 규칙 {n규칙:,} + 미분류 {n미:,} (합계 일치)")
    if eng.party_miss:
        print(f"거래처 미해소 {len(eng.party_miss)}건: {', '.join(sorted(eng.party_miss))}")
    print(f"산출물: {out}")

if __name__ == "__main__":
    main()

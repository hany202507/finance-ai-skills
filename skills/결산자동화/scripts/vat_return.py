# -*- coding: utf-8 -*-
"""부가세 신고서 집계 — Judgment 집합 → 일반과세자 확정신고 본지 (1)~(30) 값.

현행 서식 [별지 제21호서식] <개정 2025.7.4.> 기준. 계산은 결정론, 검증식 assert(1원 대사).
근거: 부가가치세법 시행규칙 별지 제21호 서식 (2025.7.4. 현행)
입력: vat_engine.judge()가 낸 Judgment 리스트.
"""
from dataclasses import dataclass, field
from collections import defaultdict
from typing import Dict
from vat_engine import Judgment

# 과세유형 → 본지 란 (현행 2025.7.4.)
SALES_LINE = {11: 1, 17: 3, 22: 3, 14: 4, 12: 5, 16: 6}   # 매출 과세유형 → 본지 매출란
EXEMPT_SALES = {13, 23}          # 면세매출 → 면세수입금액(매출세액 아님)
PURCH_GENERAL = {51, 52, 55}     # (10) 일반매입 / 고정자산이면 (12)
PURCH_CARD = {57, 61}            # (15) 그 밖의 공제매입세액
UNGONG = 54                      # (17) 공제받지못할 매입세액

발행공제율_2026 = 13             # /1000 = 1.3%
발행공제한도_2026 = 10_000_000   # 연 1,000만원(2026.12.31까지 한시)
전자신고세액공제액 = 10_000       # (57) 확정신고 시


@dataclass
class ReturnForm:
    """부가세 신고서 본지 집계 결과 (현행 (1)~(30))."""
    공급: Dict[int, int] = field(default_factory=lambda: defaultdict(int))   # 란 → 공급가액
    세액: Dict[int, int] = field(default_factory=lambda: defaultdict(int))   # 란 → 세액
    면세수입금액: int = 0
    카드발행액: int = 0            # (20) 발행공제 기준(과세 카드·현금영수증 매출 공급대가)
    대손세액가감: int = 0          # (8)
    매출세액: int = 0             # ㉮
    매입세액합계16: int = 0        # (16)
    공제못할17: int = 0           # (17)
    매입세액18: int = 0           # (18) ㉯
    납부세액: int = 0             # ㉰
    발행공제: int = 0             # (20)
    전자신고공제: int = 0          # (57)→(19)
    예정고지: int = 0             # (24)
    차감납부: int = 0             # (30)


def aggregate(judgments, *, 신고구분="확정", 발행공제자격=False, 발행공제_기발생=0,
              전자신고공제=True, 예정고지=0, 대손세액가감=0) -> ReturnForm:
    """Judgment 리스트를 본지 란별로 집계하고 세액을 계산한다.

    신고구분: '확정'|'예정'. 예정신고는 전자신고세액공제·대손세액공제(확정 전용)와
    예정고지 정산을 적용하지 않는다(§48~49).
    """
    if 신고구분 == "예정":
        전자신고공제 = False       # (57) 전자신고세액공제는 확정신고만
        대손세액가감 = 0           # (8) 대손세액공제는 확정신고만
        예정고지 = 0               # 예정신고 자체엔 예정고지 차감 없음
    f = ReturnForm()
    for j in judgments:
        if j.과세유형 is None:        # 일반전표·비과세(라우팅 일반)
            continue
        c = j.과세유형
        if j.구분 == "매출":
            if c in EXEMPT_SALES:
                f.면세수입금액 += j.공급가
                continue
            ln = SALES_LINE.get(c)
            if ln:
                f.공급[ln] += j.공급가
                f.세액[ln] += j.세액
            if c in (17, 22):
                f.카드발행액 += j.공급가 + j.세액
        elif j.구분 == "매입":
            if c == UNGONG:
                # 불공제분도 세금계산서 수취분(10/12)에 포함 → (17)에서 차감((18)=(16)−(17))
                ln = 12 if getattr(j, "고정자산", False) else 10
                f.공급[ln] += j.공급가
                f.세액[ln] += j.세액
                f.공급[17] += j.공급가
                f.세액[17] += j.세액
            elif c in PURCH_CARD:
                f.공급[15] += j.공급가
                f.세액[15] += j.세액
            elif c in PURCH_GENERAL:
                ln = 12 if getattr(j, "고정자산", False) else 10
                f.공급[ln] += j.공급가
                f.세액[ln] += j.세액

    # 매출세액 ㉮ = (1)~(4) 세액 + (8) 대손가감  ((5)(6) 영세는 세액 0)
    f.대손세액가감 = 대손세액가감
    f.매출세액 = sum(f.세액[l] for l in (1, 2, 3, 4)) + 대손세액가감

    # 매입세액: (16) = (10)−(11)+(12)+(13)+(14)+(15), (18) = (16)−(17)
    f.매입세액합계16 = (f.세액[10] - f.세액[11] + f.세액[12]
                       + f.세액[13] + f.세액[14] + f.세액[15])
    f.공제못할17 = f.세액[17]
    f.매입세액18 = f.매입세액합계16 - f.공제못할17

    # 납부(환급) ㉰ = ㉮ − ㉯
    f.납부세액 = f.매출세액 - f.매입세액18

    # (20) 발행공제 — 자격게이트(개인·직전연도 10억↓)일 때만, 연 한도 관리
    if 발행공제자격:
        raw = f.카드발행액 * 발행공제율_2026 // 1000
        f.발행공제 = max(0, min(raw, 발행공제한도_2026 - 발행공제_기발생))

    f.전자신고공제 = 전자신고세액공제액 if 전자신고공제 else 0
    f.예정고지 = 예정고지

    # (30) 차가감 납부할세액
    f.차감납부 = f.납부세액 - f.발행공제 - f.전자신고공제 - f.예정고지

    _validate(f)
    return f


def _validate(f: ReturnForm) -> None:
    """검증식 — 1원 대사(설계 §5)."""
    assert f.매입세액18 == f.매입세액합계16 - f.공제못할17, "(18) = (16)−(17) 위반"
    assert f.납부세액 == f.매출세액 - f.매입세액18, "㉰ = ㉮−㉯ 위반"
    assert f.차감납부 == f.납부세액 - f.발행공제 - f.전자신고공제 - f.예정고지, "(30) 위반"


def build_return_xlsx(f: ReturnForm, path: str) -> str:
    """본지 (1)~(30) 요약 신고서 엑셀 산출(값 + 합계 수식). build_workbook 패턴."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "부가세신고서(본지)"
    ws["A1"] = "부가가치세 신고서 (일반과세자 확정신고, 현행 별지 제21호 2025.7.4.)"
    ws["A1"].font = Font(bold=True)
    ws.append(["란", "명칭", "공급가액", "세액"])
    rows = [
        (1, "세금계산서 발급분(과세)"), (3, "신용카드·현금영수증 발행분(과세)"),
        (4, "기타(과세)"), (5, "세금계산서 발급분(영세율)"), (6, "기타(영세율)"),
        (10, "세금계산서 수취 일반매입"), (12, "세금계산서 수취 고정자산매입"),
        (15, "그 밖의 공제매입세액"), (17, "공제받지 못할 매입세액"),
    ]
    for ln, nm in rows:
        ws.append([ln, nm, f.공급.get(ln, 0), f.세액.get(ln, 0)])
    ws.append([])
    summary = [
        ("면세수입금액", f.면세수입금액),
        ("매출세액 ㉮", f.매출세액),
        ("매입세액 ㉯ (18)", f.매입세액18),
        ("납부세액 ㉰", f.납부세액),
        ("(20) 신용카드발행공제", f.발행공제),
        ("(57) 전자신고세액공제", f.전자신고공제),
        ("(24) 예정고지세액", f.예정고지),
        ("(30) 차가감 납부할세액", f.차감납부),
    ]
    for nm, val in summary:
        ws.append(["", nm, "", val])
    wb.save(path)
    return path

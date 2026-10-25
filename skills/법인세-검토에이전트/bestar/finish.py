# -*- coding: utf-8 -*-
"""마무리: 빠진 것 검사 → 엔진 다시(AR · 조서, 회계사 결론 이어받기) → 세 묶음 요약 → 엑셀 시트 더함 → 한 파일 화면(목업 「3. 분석적 검토」, 읽기 전용). 판단을 저장하며 보려면 ar.py serve.

  ar.py finish <신고건> [--draft]

--draft 는 메모 · 제안이 덜 끝났어도 지금 있는 것으로 요약 · 엑셀 · 화면을 만든다(검사 실패는 화면 머리에 보인다).
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ctax.ar import cli as ar_cli
from ctax.ar import engine as ar_engine
from ctax.ar import excel as ar_excel
from ctax.ar.guide import form_names

from bestar import facts as facts_mod
from bestar import memos, paths, rank, steps
from bestar.keep import ADDED_CONCL, ADDED_SHEET, save_added_conclusions

YELLOW = PatternFill("solid", fgColor="FFF2CC")
ALIAS_FILE = "설정.json"          # {"alias": "(주)스타일"}: 보고서 · 화면 · 엑셀에 실제 회사 이름 대신 쓸 이름
HEAD = PatternFill("solid", fgColor="17365D")


# ---------------------------------------------------------------- 자료 모으기
def added_issues(c, f1: dict) -> list[dict]:
    """F1 AI 추가 이슈 + 메모 「added」 줄(근거 조문 · 질문지 · 세액 영향 · 숫자 · 받은 앞 단계 항목)."""
    concl = paths.read_json(paths.steps_dir(c) / ADDED_CONCL, {}) or {}
    ex = memos.added_extras(c)
    out = []
    for a in f1.get("accounts") or []:
        if a.get("issues"):
            continue
        e = ex.get(a["id"]) or {}
        out.append({"id": a["id"], "account": a.get("account"), "title": a.get("title"), "current": a.get("current"), "prior": a.get("prior"),
                    "finding": a.get("finding"), "why": a.get("why"), "risk": a.get("risk"),
                    "judgment": a.get("judgment") or "판단 필요", "ask": a.get("ask") or [], "conclusion": concl.get(a["id"]),
                    "law": e.get("law") or [], "choices": e.get("choices") or a.get("choices"),
                    "tax_effect": bool(e.get("tax_effect") or a.get("tax_effect")), "facts": e.get("facts") or a.get("facts") or [],
                    "answers": e.get("answers") or [], "stake": e.get("stake")})
    return out


def received(c, ar: dict, extras: dict, credits: dict, added: list[dict]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for i in ar.get("issues", []):
        for k in extras.get(i["no"], {}).get("answers") or []:
            out.setdefault(k, []).append(f"이슈 {i['no']}")
    for x in (credits or {}).get("items") or []:
        for k in x.get("answers") or []:
            out.setdefault(k, []).append(f"공제 제안 {x.get('name')}")
    for a in added:
        out.setdefault(a["id"], []).append("AI 추가 이슈")
        for k in a.get("answers") or []:
            out.setdefault(k, []).append(f"AI 추가 이슈 {a['id']}")
    for k, _name, step, item in steps.handed(c):
        if item.get("outlook") in ("안 생길 듯", "종료"):
            out.setdefault(k, []).append(f"{step} 에서 닫음({item.get('outlook')})")
    return out


def gaps(c, ar: dict, rec: dict, credits: dict | None) -> list[str]:
    out = []
    miss = [i["no"] for i in ar.get("issues", []) if not i.get("memo")]
    if miss:
        out.append(f"메모가 없는 이슈 {miss}: ar.py memo 로 넣습니다")
    if credits is None:
        out.append(f"_색인/{paths.CREDITS} 이 없습니다: 제안할 것이 없어도 items 를 빈 목록으로 ar.py credits 를 돌립니다")
    left = steps.outside_unreviewed(paths.read_json(paths.index(c) / "장부.json", {}) or {}, ar, steps.load_step(c, "F1")["carry"])
    if left:
        out.append(f"2-(1) 에서 훑지 않은 엔진 규칙에 안 걸린 계정 {len(left)}개({', '.join(left[:8])}{' …' if len(left) > 8 else ''}): "
                   "F1 의 others_reviewed 에 적고 step done F1 --redo")
    for k, name, step, _item in steps.handed(c):
        if k not in rec:
            out.append(f"{step} 에서 넘긴 {k} {name}: 받은 곳이 없습니다(이슈 메모 answers 또는 공제 제안 answers)")
    return out


def collect(c, draft_gaps: list[str] | None = None) -> dict:
    ar = paths.read_ar(c)
    st = {s: steps.load_step(c, s) for s in steps.ORDER}
    ex = memos.extras(c)
    credits = paths.read_json(paths.index(c) / paths.CREDITS)
    added = added_issues(c, st["F1"]["carry"])
    for i in ar.get("issues", []):
        e = ex.get(i["no"], {})
        i["answers"], i["risk"] = e.get("answers") or [], e.get("risk")
        i["facts"], i["same_as"] = e.get("facts") or [], e.get("same_as")
    rec = received(c, ar, ex, credits, added)
    adj_sum = (paths.read_json(paths.index(c) / "조정.json", {}) or {}).get("summary") or {}
    income = adj_sum.get("income")
    summary = rank.build(ar, ex, st, added, credits or {}, [g["id"] for g in ar.get("guide", [])],
                         taxable=None if income is None else income > 0, tax_summary=adj_sum)
    ledger = paths.read_json(paths.index(c) / "장부.json", {}) or {}
    info = c.info()
    alias = (paths.read_json(paths.steps_dir(c) / ALIAS_FILE, {}) or {}).get("alias")
    if alias:
        info = {**info, "company": alias}
    kinds = Counter(i["kind"] for i in ar.get("issues", []))
    f2 = st["F2"]["carry"]
    return {
        "company": info.get("company"), "fy": info.get("fy"), "fy_start": info.get("fy_start"), "fy_end": info.get("fy_end"),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "premise": {k: f2.get(k) for k in ("verdict", "status", "premise", "point")},
        "progress": {"steps_done": len(steps.read_log(c)), "steps_total": len(steps.ORDER),
                     "memos": sum(1 for i in ar.get("issues", []) if i.get("memo")), "issues": len(ar.get("issues", [])),
                     "conclusions": sum(1 for i in ar.get("issues", []) if i.get("conclusion")) + sum(1 for a in added if a.get("conclusion")),
                     "credits": len((credits or {}).get("items") or [])},
        "gaps": draft_gaps or [],
        "summary": summary,
        "ar": {**ar, "kinds": [{"kind": k, "count": kinds[k]} for k in ar_engine.KINDS if kinds.get(k)]},
        "account_names": {a["code"]: a["name"] for a in ledger.get("accounts") or [] if a.get("code")},
        "form_names": form_names(),
        "paper_names": ar_excel.paper_names(),
        "added": added,
        "facts": facts_mod.load(c),
        "credits": credits or {"point": "", "items": []},
        "steps": [{"id": s, "step": steps.LABEL[s], "title": steps.TITLE[s], "finding": st[s].get("finding"), "note": st[s].get("note"),
                   "carry": st[s]["carry"], "ask": st[s].get("ask") or st[s]["carry"].get("ask") or [],
                   "received": {k: rec.get(k, []) for k in [x.get("id") for key in ("forecast", "accounts") for x in st[s]["carry"].get(key) or []]}}
                  for s in steps.ORDER],
    }


# ---------------------------------------------------------------- 엑셀
def _sheet(wb, name: str, head: list[str], rows: list[list], widths: list[int], index: int | None = None):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, index) if index is not None else wb.create_sheet(name)
    ws.append(head)
    for r in rows:
        ws.append(r)
    for n, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(n)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, Font(bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def write_excel(c, data: dict, ar_sheet: dict | None = None) -> None:
    path = c.root / paths.XLSX
    wb = load_workbook(path)
    s = data["summary"]
    rows = []
    for tier, items in (("가장 중요한 3가지", s["top"]), ("확인사항", s["check"]), ("확인하면 좋은 세부사항", s["detail"])):
        for n, u in enumerate(items, 1):
            rows.append([tier, n, u["title"], u.get("judgment") or "메모 없음", u.get("risk") or "", u.get("size") or None,
                         u["where"], u.get("line") or "", "\n".join(u.get("asks")[:3]), u["uid"]])
    ws = _sheet(wb, "요약", ["묶음", "순서", "사안", "AI 판단", "위험도", "크기(원)", "어디서 잡혔나", "한 줄 결론", "받을 자료", "가리키는 곳"],
                rows, [16, 6, 44, 11, 7, 15, 26, 60, 46, 12], index=0)
    for r in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        r[0].number_format = "#,##0"
    JD = {"조정 필요": "조정 예상"}
    rows = []
    for g in (ar_sheet or {}).get("groups") or []:
        for r in g["rows"]:
            rows.append([g["label"], "이슈" if r.get("issue") else "", r["title"], r.get("relation"), JD.get(r["judgment"], r["judgment"]), r.get("judgment_text"),
                         r.get("amount"), chr(10).join(r.get("asks") or []), ", ".join(r.get("guide") or [])])
    if rows:
        ws = _sheet(wb, "AR 한 장", ["묶음", "이슈", "항목", "관련(장부 · 전기)", "추정 판단", "판단 내용", "금액", "추가 요청", "가이드 줄"],
                    rows, [14, 7, 30, 46, 11, 56, 15, 50, 16], index=1)
        ws.auto_filter.ref = ws.dimensions          # 「이슈」 열로 이슈만 거른다
        for r in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            r[0].number_format = "#,##0"
    acc = ((ar_sheet or {}).get("coverage") or {}).get("accounts") or {}
    if acc.get("rows"):
        ws = _sheet(wb, "계정 전수", ["코드", "계정", "구분", "금액", "다룬 곳"],
                    [[r["code"], r["name"], r["class"], r["amount"], r["how"] or "엔진 규칙에 안 걸림 · 2-(1) 에서 훑음"] for r in acc["rows"]],
                    [10, 28, 8, 18, 40], index=2)
        ws.auto_filter.ref = ws.dimensions
        for r in ws.iter_rows(min_row=2, min_col=4, max_col=4):
            r[0].number_format = "#,##0"
    added = data["added"]
    ws = _sheet(wb, ADDED_SHEET, ["번호", "계정", "당기", "확인한 사실", "왜 보나", "위험도", "AI 판단", "회계사 결론"],
                [[a["id"], a.get("account"), a.get("current"), a.get("finding"), a.get("why"), a.get("risk"), a.get("judgment"),
                  (a.get("conclusion") or {}).get("text")] for a in added], [7, 18, 15, 60, 40, 7, 11, 40])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).fill = YELLOW
        ws.cell(r, 3).number_format = "#,##0"
    cr = data["credits"]
    rows = []
    for x in cr.get("items") or []:
        cands = "\n".join(f"전표 {k.get('line')}줄 {k.get('what')} {k.get('amount'):,}원" if isinstance(k.get("amount"), (int, float)) else str(k)
                          for k in x.get("candidates") or [])
        rows.append([x.get("state"), x.get("name"), x.get("basis"), x.get("purpose"), cands, "\n".join(x.get("judge") or []),
                     "\n".join(x.get("ask") or []), ", ".join(str(n) for n in x.get("issues") or [])])
    _sheet(wb, "공제 제안", ["묶음", "제도", "근거", "입법 취지", "취지에 맞는 거래", "회계사가 판단할 점", "받을 자료", "관련 이슈"],
           rows, [11, 26, 24, 50, 56, 40, 36, 10])
    rows = []
    for f in data.get("facts") or []:
        if f["type"] == "합계":
            rows.append([f["id"], "엔진 합계", f["what"], f["rule"], f["count"], f["debit"], f["credit"], f["net"], ""])
        else:
            rows.append([f["id"], "전기 PDF(AI 판독, 확인 전)", f["what"], f"{f['form']}호 {f['page']}쪽", None, None, None, f["amount"], ""])
    ws = _sheet(wb, "합계 · PDF 값", ["번호", "종류", "무엇", "고른 줄 · 쪽", "줄 수", "차변", "대변", "순액 · 금액", "확인"],
                rows, [7, 22, 34, 40, 7, 15, 15, 15, 10])
    for r in ws.iter_rows(min_row=2, min_col=6, max_col=8):
        for cell in r:
            cell.number_format = "#,##0"
    from bestar import decide
    dec = decide.load(c)
    rows = [[k, "답", v.get("values") or v.get("value"), v.get("other") or "", v.get("by"), v.get("at")] for k, v in dec["decisions"].items()]
    rows += [[k, "반영", "", "", v.get("by"), v.get("at")] for k, v in dec["accepted"].items()]
    from bestar import link
    names = link.paper_names(c)
    for v in ((paths.read_json(link.state_path(c), {}) or {}).get("linked") or {}).values():
        rows.append([link.describe(v, names), "조서 칸", v.get("value"), f"넣기 전 {v.get('prev')}", v.get("source"), ""])
    _sheet(wb, "화면 판단", ["항목", "종류", "고른 답", "직접 적은 답", "누가", "언제"], [[str(x) if isinstance(x, list) else x for x in r] for r in rows],
           [50, 8, 30, 40, 10, 20])
    rows = []
    for st in data["steps"]:
        got = "\n".join(f"{k} → {', '.join(v) or '받은 곳 없음'}" for k, v in st["received"].items())
        rows.append([st["step"], st["title"], st.get("finding"), st.get("note"), got])
    _sheet(wb, "순서와 전제", ["순서", "단계", "확인한 사실", "다음 단계로 넘긴 말", "넘긴 항목 → 받은 곳"], rows, [8, 26, 70, 50, 50])
    alias = (paths.read_json(paths.steps_dir(c) / ALIAS_FILE, {}) or {}).get("alias")
    real = (paths.case(c.root).info() or {}).get("company")
    if alias and real and real != alias:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and real in cell.value:
                        cell.value = cell.value.replace(real, alias)
    wb.active = 0
    wb.save(path)
    wb.close()


# ---------------------------------------------------------------- 명령
def main(argv: list[str]) -> int:
    p = argparse.ArgumentParser(prog="ar.py finish")
    p.add_argument("folder")
    p.add_argument("--draft", action="store_true")
    p.add_argument("--alias", help="보고서 · 화면에 실제 회사 이름 대신 쓸 이름(한 번 주면 남는다, 지우려면 --alias \"\")")
    a = p.parse_args(argv)
    c = paths.case(a.folder)
    if a.alias is not None:
        paths.write_json(paths.steps_dir(c) / ALIAS_FILE, {"alias": a.alias.strip()} if a.alias.strip() else {})
    if not steps.finished(c):
        print("1-(1) ~ 2-(2) 여섯 단계를 먼저 끝냅니다(ar.py step status).")
        return 2
    save_added_conclusions(c)
    code = ar_cli.build(c)                 # 엑셀 결론 칸 이어받기(메모도 이어받음)
    if code == 2:
        return 2
    from bestar import link                # 화면 답 · 2-(2) 전제를 조서 노란 칸에 맞춘 뒤
    if link.xlsx_path(c).exists():
        r = link.sync(c, rerun=False)
        names = link.paper_names(c)
        for w in r["writes"]:
            print(f"조서 칸: {link.describe(w, names)} (전 {w['current']}) · {w['source']}")
        for w in r["kept"] + r["unknown"]:
            print(f"조서 칸 넣지 않음: {link.describe(w, names)} · {w['why']}")
        if r.get("problem"):
            print(f"조서 칸: {r['problem']}")
    from ctax.wp import cli as wp_cli      # AR.json 이 바뀌었으니 조서도 다시(조서 엑셀 노란 칸 이어받음)
    if wp_cli.main([str(c.root)]) == 2:
        return 2
    ar = paths.read_ar(c)
    st_f1 = steps.load_step(c, "F1")["carry"]
    credits = paths.read_json(paths.index(c) / paths.CREDITS)
    rec = received(c, ar, memos.extras(c), credits, added_issues(c, st_f1))
    found = gaps(c, ar, rec, credits)
    if found and not a.draft:
        print("마무리하지 못했습니다:\n- " + "\n- ".join(found) + "\n(지금 것으로 보려면 --draft)")
        return 1
    data = collect(c, found)
    from bestar import review, review_view
    rd = review.build(c, found)
    paths.write_json(paths.index(c) / paths.SUMMARY, {k: data[k] for k in ("company", "fy", "generated_at", "premise", "progress", "gaps", "summary")})
    try:
        write_excel(c, data, rd.get("ar_sheet"))
    except PermissionError:
        print(f"{paths.XLSX} 를 쓸 수 없습니다(엑셀에서 열려 있으면 닫고 다시 돌리세요)")
        return 1
    from bestar import report
    report.write(rd, c.root / paths.REPORT)
    s = data["summary"]
    print(f"가장 중요한 3가지 {len(s['top'])} · 확인사항 {len(s['check'])} · 세부사항 {len(s['detail'])}")
    for n, u in enumerate(s["top"], 1):
        print(f"  {n}. [{u['where']}] {u['title']} ({u.get('judgment') or '메모 없음'} · 위험 {u.get('risk')})")
    print(f"보고서: {paths.REPORT}\n엑셀: {paths.XLSX}\n화면이 필요하면: python scripts/ar.py view \"<신고건>\" (판단 저장은 ar.py serve)")
    return 1 if found else 0

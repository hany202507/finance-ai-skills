"""조서 · 서식 · AR 엑셀의 공통 꼴(CONTRACT.md 11절). openpyxl 만 쓴다."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONEY = "#,##0;[Red]-#,##0;-"
FONT = "맑은 고딕"

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=13)
NOTE_FONT = Font(name=FONT, color="666666", size=9)
BODY_FONT = Font(name=FONT, size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # 노란 바탕 = 사람이 넣는 칸
INPUT_FONT = Font(name=FONT, color="1F4E9E", size=10)  # 파란 글자
LINK_FONT = Font(name=FONT, color="2E7D32", size=10)   # 초록 글자 = 다른 시트 값
TOTAL_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def title(ws, text: str, sub: str = "", row: int = 1) -> int:
    """시트 제목과 부제. 다음에 쓸 행 번호를 돌려준다."""
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    if sub:
        ws.cell(row=row + 1, column=1, value=sub).font = NOTE_FONT
        return row + 3
    return row + 2


def header(ws, row: int, labels: list[str], col: int = 1) -> None:
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def cell(ws, row: int, col: int, value=None, *, money: bool = False, kind: str = "body", wrap: bool = False):
    """kind: body · input(노란 칸) · link(다른 시트 값) · total(합계 줄)."""
    c = ws.cell(row=row, column=col, value=value)
    c.border = BOX
    c.font = {"input": INPUT_FONT, "link": LINK_FONT}.get(kind, BODY_FONT)
    if kind == "input":
        c.fill = INPUT_FILL
    elif kind == "total":
        c.fill = TOTAL_FILL
        c.font = Font(name=FONT, bold=True, size=10)
    if money:
        c.number_format = MONEY
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def widths(ws, values: list[int]) -> None:
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def note(ws, row: int, text: str, col: int = 1) -> None:
    ws.cell(row=row, column=col, value=text).font = NOTE_FONT

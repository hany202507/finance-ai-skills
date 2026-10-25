"""`1 AR/AR_이슈.xlsx`. 시트 셋: 이슈 목록 · 서식 목록 · 가이드 대조(CONTRACT.md 11절 꼴, ctax.xlsx 도우미)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ctax import xlsx as X
from ctax.ar.guide import load_guide

# 집계 시트 이름(조서 등록부에 없고 wp/summary.py 가 만든다)
AGG_NAMES = {"summary_15": "소득금액조정", "tax_3": "과세표준 · 세액", "reserve_50": "유보 · 자본금", "loss": "결손금"}


def paper_names() -> dict[str, str]:
    """조서 id → 「<번호> <이름>」(조서.xlsx 시트 이름). 조서 등록부(ctax.wp.papers ORDER)와 조서 파일의 NAME 에서 만든다.
    조서 파일 하나를 읽지 못하면 그 조서는 id 를 그대로 보인다."""
    from importlib import import_module
    out = dict(AGG_NAMES)
    try:
        from ctax.wp import papers as registry
    except Exception:                                    # noqa: BLE001
        return out
    for pid, no in registry.ORDER:
        try:
            name = import_module(f"{registry.__name__}.{pid}").NAME
        except Exception:                                # noqa: BLE001 (다른 에이전트가 고치는 중인 조서)
            continue
        out[pid] = f"{no} {name}"
    return out


COLS = ["번호", "종류", "확실도", "가이드", "계정", "제목", "전기", "당기", "증감", "근거", "전기 처리", "조서", "서식",
        "회사에 물을 것", "다음 할 일", "근거 조문(엔진)", "AI 판단(확인 전)", "AI 까닭 · 질문", "AI 근거 조문 · 등급", "회계사 결론", "key", "rule"]
COL = {name: i + 1 for i, name in enumerate(COLS)}
FOLD_MARK = "접음"


def _lines(items) -> str:
    return "\n".join(str(x) for x in items if x)


def _laws(law) -> str:
    return _lines(f"{x.get('ref')} {x.get('grade') or ''}".strip() for x in (law or []) if isinstance(x, dict))


def _fold_head(ws, row: int, folded: list[dict]) -> None:
    """작은 조정 후보 묶음 머리 줄. 아래 줄은 접어 둔다(엑셀 왼쪽 + 로 펼침)."""
    from ctax.ar.engine import FOLD_LIMIT
    nos = ", ".join(str(i["no"]) for i in folded)
    X.cell(ws, row, 1, FOLD_MARK)
    X.cell(ws, row, 2, "조정 후보")
    X.cell(ws, row, 3, "가능")
    X.cell(ws, row, 6, f"작은 조정 후보 {len(folded)}건(이슈 {nos}). 당기 · 전기 · 증감이 모두 {FOLD_LIMIT:,}원 미만이라 접었습니다. "
                       "왼쪽 + 를 누르면 펼칩니다", wrap=True)
    for col in range(1, len(COLS) + 1):
        c = ws.cell(row=row, column=col)
        c.fill = X.TOTAL_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _issue_row(ws, row: int, iss: dict, names: dict, papers: dict | None = None) -> None:
    papers = papers or {}
    a = iss.get("amounts") or {}
    codes = iss.get("accounts", [])
    acc = _lines([f"{c} {names.get(c, '')}".strip() for c in codes[:6]] + ([f"외 {len(codes) - 6}개"] if len(codes) > 6 else []))
    vals = [iss["no"], iss["kind"], iss["certainty"], _lines(iss.get("guide", [])), acc, iss["title"]]
    for col, v in enumerate(vals, start=1):
        X.cell(ws, row, col, v, wrap=col in (4, 5, 6))
    X.cell(ws, row, 7, a.get("prior"), money=True)
    X.cell(ws, row, 8, a.get("current"), money=True)
    if a.get("prior") is not None and a.get("current") is not None and a.get("change", a["current"] - a["prior"]) == a["current"] - a["prior"]:
        X.cell(ws, row, 9, f"=H{row}-G{row}", money=True)
    else:
        X.cell(ws, row, 9, a.get("change"), money=True)
    X.cell(ws, row, 10, _lines(iss.get("basis", [])), wrap=True)
    X.cell(ws, row, 11, iss.get("prior_treatment"), wrap=True)
    X.cell(ws, row, 12, _lines(papers.get(p, p) for p in iss.get("papers", [])), wrap=True)
    X.cell(ws, row, 13, " · ".join(iss.get("forms", [])), wrap=True)
    X.cell(ws, row, 14, _lines(iss.get("questions", [])), wrap=True)
    X.cell(ws, row, COL["다음 할 일"], _lines(iss.get("todo", [])), wrap=True)
    X.cell(ws, row, COL["근거 조문(엔진)"], _laws(iss.get("law")), wrap=True)
    m = iss.get("memo") if isinstance(iss.get("memo"), dict) else None
    if m:
        why = m.get("why") or ""
        if m.get("questions"):
            why += "\n회사에 물을 것: " + " / ".join(m["questions"])
        X.cell(ws, row, COL["AI 판단(확인 전)"], m.get("judgment"), wrap=True)
        X.cell(ws, row, COL["AI 까닭 · 질문"], why, wrap=True)
        X.cell(ws, row, COL["AI 근거 조문 · 등급"], _laws(m.get("law")), wrap=True)
    else:
        for col in (COL["AI 판단(확인 전)"], COL["AI 까닭 · 질문"], COL["AI 근거 조문 · 등급"]):
            X.cell(ws, row, col, None)
    c = iss.get("conclusion")
    X.cell(ws, row, COL["회계사 결론"], c.get("text") if isinstance(c, dict) else c, kind="input", wrap=True)
    X.cell(ws, row, COL["key"], iss.get("key"))
    X.cell(ws, row, COL["rule"], iss.get("rule"))
    for col in range(1, len(COLS) + 1):
        ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")


def read_conclusions(path: Path) -> list[dict] | None:
    """「이슈 목록」 의 회계사 결론 칸을 읽는다 [{no, key, rule, kind, title, text}]. 파일이 없거나 읽지 못하면 None.

    key · rule 열이 없는 앞 꼴 엑셀은 종류 · 제목으로 잇는다(memo.match).
    """
    path = Path(path)
    if not path.exists():
        return None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if "이슈 목록" not in wb.sheetnames:
            return None
        rows = list(wb["이슈 목록"].iter_rows(values_only=True))
    finally:
        wb.close()
    head = next((i for i, r in enumerate(rows[:10]) if r and r[0] == "번호"), None)
    if head is None:
        return None
    idx = {str(v).strip(): j for j, v in enumerate(rows[head]) if v is not None}
    if "회계사 결론" not in idx:
        return []

    def get(r, name):
        j = idx.get(name)
        return r[j] if j is not None and j < len(r) else None

    out = []
    for r in rows[head + 1:]:
        if not r or not isinstance(get(r, "번호"), (int, float)):
            continue
        text = get(r, "회계사 결론")
        out.append({"no": int(get(r, "번호")), "key": get(r, "key"), "rule": get(r, "rule"), "kind": get(r, "종류"),
                    "title": get(r, "제목"), "text": "" if text is None else str(text).strip()})
    return out


def write(ar: dict, path: Path, ledger: dict | None = None, info: dict | None = None) -> Path:
    info = info or {}
    names = {str(a.get("code")): str(a.get("name") or "") for a in (ledger or {}).get("accounts", [])}
    head = f"{info.get('company', '')} {info.get('fy', '')} 사업연도".strip()
    wb = Workbook()

    # 이슈 목록
    ws = wb.active
    ws.title = "이슈 목록"
    r = X.title(ws, "분석적 검토 이슈 목록",
                f"{head} · 엔진 판정이라 확인 전입니다. 순서는 신고 점검, 확실, 전기 연속, 가능, 자료 필요, 서식 필요이고 같은 묶음 안에서는 금액이 큰 순입니다. "
                "AI 메모는 회계사 확인 전이고, 회계사 결론은 노란 칸에 적으면 다음 ctax ar 가 이어받습니다")
    X.header(ws, r, COLS)
    ws.freeze_panes = ws.cell(row=r + 1, column=7)
    ws.sheet_properties.outlinePr.summaryBelow = False
    row = r
    issues = ar.get("issues", [])
    folded = [i for i in issues if i.get("folded")]
    pnames = paper_names()
    for iss in issues:
        row += 1
        if folded and iss is folded[0]:
            _fold_head(ws, row, folded)
            row += 1
        _issue_row(ws, row, iss, names, pnames)
        if iss.get("folded"):
            ws.row_dimensions[row].outline_level = 1
            ws.row_dimensions[row].hidden = True
    if issues:
        ws.auto_filter.ref = f"A{r}:{get_column_letter(len(COLS))}{row}"
    X.widths(ws, [6, 10, 9, 11, 22, 40, 14, 14, 14, 50, 26, 16, 18, 44, 36, 30, 12, 50, 34, 30, 20, 20])
    for col in (COL["key"], COL["rule"]):
        ws.column_dimensions[get_column_letter(col)].hidden = True

    # 서식 목록
    ws = wb.create_sheet("서식 목록")
    r = X.title(ws, "서식 목록(총괄표)",
                f"{head} · 기본 서식, 이슈가 가리키는 서식, 전기에 냈던 서식. 채움이 엔진이면 ctax 가 채우고 사람이면 회계사가 넣습니다")
    X.header(ws, r, ["서식", "서식 이름", "전기 제출", "당기 판정", "까닭 이슈", "채움", "회계사 확인"])
    for i, f in enumerate(ar.get("forms", [])):
        row = r + 1 + i
        X.cell(ws, row, 1, f["code"])
        X.cell(ws, row, 2, f["name"])
        X.cell(ws, row, 3, "제출" if f.get("prior_filed") else "")
        X.cell(ws, row, 4, f["status"])
        X.cell(ws, row, 5, ", ".join(str(n) for n in f.get("why", [])))
        X.cell(ws, row, 6, f["by"])
        X.cell(ws, row, 7, None, kind="input")
    ws.freeze_panes = ws.cell(row=r + 1, column=3)
    X.widths(ws, [14, 40, 10, 10, 16, 8, 30])

    # 가이드 대조
    ws = wb.create_sheet("가이드 대조")
    r = X.title(ws, "가이드 대조(이정아 회계사 세무조정 AR 가이드)",
                f"{head} · 가이드 줄마다 엔진 판정(이슈 · 해당 없음 · 자료 필요)과 근거 사실. 가이드 원문은 한 글자도 고치지 않았습니다")
    X.header(ws, r, ["번호", "구분", "항목", "가이드 원문", "판정", "근거 사실", "이슈"])
    raw = {g["id"]: g["raw"] for g in load_guide()}
    for i, g in enumerate(ar.get("guide", [])):
        row = r + 1 + i
        X.cell(ws, row, 1, g["id"])
        X.cell(ws, row, 2, g["group"], wrap=True)
        X.cell(ws, row, 3, g["title"], wrap=True)
        X.cell(ws, row, 4, raw.get(g["id"], g["text"]), wrap=True)
        X.cell(ws, row, 5, g["status"])
        X.cell(ws, row, 6, _lines(g.get("facts", [])), wrap=True)
        X.cell(ws, row, 7, ", ".join(str(n) for n in g.get("issues", [])), wrap=True)
        for col in range(1, 8):
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=r + 1, column=4)
    X.widths(ws, [10, 22, 26, 50, 10, 60, 10])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

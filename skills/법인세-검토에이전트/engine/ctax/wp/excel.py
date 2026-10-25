"""`2 조서/조서.xlsx` 를 쓴다. 조서 시트는 등록부 순서대로 각 조서 파일의 `write` 가 쓰고, 여기서는
기준 시트 · 집계 시트(15호 · 3호 · 결손금 · 50호) · 총괄 · 숨은 `_입력` 시트를 쓴다.

`write()` 는 values 키 → 셀 주소 지도를 돌려준다(시험이 수식 결과와 JSON 값을 견줄 때 쓴다).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl.styles import Alignment

from .. import xlsx as X
from . import prior_compare
from .books import Books
from .sheet import INPUT_SHEET, Book, Sheet, local, notes
from .summary import CARRYBACK, DIRECT_FIELDS, DIRECT_ROWS, LOSS_MANUAL_ROWS, Result


def write(result: Result, books: Books, path: Path) -> dict:
    bk = Book(books)
    overview = bk.sheet("총괄", [10, 26, 12, 16, 16, 14, 40])
    _basis(bk, result)
    for p in result.papers:
        result.mods[p.id].write(bk, p)
    _summary15(bk, result)
    _tax3(bk, result)
    _loss(bk, result)
    for p in result.papers:
        after = getattr(result.mods[p.id], "write_after", None)
        if after:
            after(bk, p)
    t3 = result.tax3.extra.pop("sheet")
    don = bk.cells.get("donation", {}).get("한도초과")
    t3.put(result.tax3.extra.pop("row05"), 3, f"={don}" if don else 0, kind="link" if don else "body", money=True)
    _reserve50(bk, result)
    _overview(bk, overview, result)
    _inputs_sheet(bk)
    path.parent.mkdir(parents=True, exist_ok=True)
    bk.wb.save(path)
    result.input_cells = [dict(c) for c in bk.input_cells]
    return bk.cells


# 기준 2025 ---------------------------------------------------------------------------------------

def _basis(bk: Book, res: Result) -> None:
    s = bk.sheet("기준 2025", [34, 18, 70, 8])
    r = X.title(s.ws, "기준 2025: 세율 · 한도", "rates_2025.yaml 값입니다. 노란 칸은 회사마다 넣는 값입니다")
    X.header(s.ws, r, ["항목", "값", "근거", "등급"])
    r += 1
    rates = bk.b.rates

    def row(name, value, basis="", grade="", fmt=None):
        nonlocal r
        s.put(r, 1, name)
        s.put(r, 2, value, money=isinstance(value, int) and fmt is None, fmt=fmt)
        s.put(r, 3, basis, wrap=False)
        s.put(r, 4, grade)
        bk.ref[name] = s.addr(r, 2)
        r += 1

    t3 = res.tax3
    s.put(r, 1, "중소기업 여부")
    s.inp(r, 2, t3.inputs["중소기업 여부"], pid="tax_3", key="중소기업 여부", default=t3.defaults["중소기업 여부"])
    s.put(r, 3, "전기 신고서(51호) 판정. 바뀌었으면 고치세요(여 · 부)")
    bk.ref["중소기업 여부"] = s.addr(r, 2)
    r += 1
    row("사업연도 월수", bk.b.months, "신고건.json 사업연도", "", fmt="0")
    br = rates["corporate_tax"]["brackets"]
    for i, (cap, rt) in enumerate(br["value"], start=1):
        if cap is not None:
            row(f"세율 구간{i} 상한", cap, br["basis"], br["grade"])
        row(f"세율 구간{i}", rt, br["basis"], br["grade"], fmt="0%")
    for p in res.papers:
        fn = getattr(res.mods[p.id], "basis", None)
        if fn:
            for name, value, basis, grade, fmt in fn(rates):
                row(name, value, basis, grade, fmt)
    ls = rates["loss"]
    row("이월결손금 공제율(중소기업)", ls["sme_ratio"]["value"], ls["sme_ratio"]["basis"], ls["sme_ratio"]["grade"], fmt="0%")
    row("이월결손금 공제율(일반)", ls["general_ratio"]["value"], ls["general_ratio"]["basis"], ls["general_ratio"]["grade"], fmt="0%")
    r += 1
    for p in res.papers:
        fn = getattr(res.mods[p.id], "basis_table", None)
        if fn:
            r = fn(bk, s, r) + 1
    X.note(s.ws, r, f"원문 확인 {rates.get('checked', '')} · {rates.get('source', '')}")


# 집계 ----------------------------------------------------------------------------------------------

def _summary15(bk: Book, res: Result) -> None:
    p = res.summary15
    s = bk.sheet(p.sheet, [8, 8, 4, 24, 16, 44, 16, 22])
    s.ws.cell(row=1, column=1, value="15 소득금액조정합계표").font = X.TITLE_FONT
    X.note(s.ws, 2, f"{bk.b.company} · {bk.b.fy} 사업연도 · 조서의 조정을 모읍니다(초록 글자 = 조서 시트 값, 노란 줄 = 직접 조정)")
    row = 4
    X.header(s.ws, row, ["구분", "코드", "", "과목", "소득처분", "조정내용", "금액", "조서"])
    row += 1
    first = row
    for paper, a in res.adjustments():
        src = bk.cells[paper.id][f"adj:{a.kind}"]
        s.put(row, 1, a.side)
        s.put(row, 2, a.code)
        s.put(row, 4, a.subject)
        s.put(row, 5, a.disposition)
        s.put(row, 6, a.note)
        s.put(row, 7, f"={src}", kind="link", money=True)
        s.put(row, 8, paper.sheet)
        row += 1
    row = prior_compare.write_applied(bk, s, row, res)   # 「전기 조정 대조」 에서 회계사가 O 로 넣은 줄
    by_row = {getattr(a, "row", 0): a for a in res.direct}
    cols = {"구분": 1, "과목": 4, "소득처분": 5, "조정내용": 6, "금액": 7}
    for n in range(1, DIRECT_ROWS + 1):
        a = by_row.get(n)
        for name, kind in DIRECT_FIELDS:
            key = f"직접 조정 {n}/{name}"
            s.inp(row, cols[name], p.inputs.get(key) or None, pid=p.id, key=key,
                  default=p.defaults.get(key) or None, kind=kind, money=name == "금액")
        s.put(row, 2, a.code if a else None)
        s.put(row, 8, f"직접 조정 {n}")
        if a:
            bk.reg(p.id, f"direct:{n}", s.addr(row, 7))
        row += 1
    X.note(s.ws, row, "직접 조정: 구분(가산 · 차감) · 과목 · 소득처분(유보 · 기타사외유출 · 상여 · 배당 · 기타소득 · 기타) · "
                      "조정내용 · 금액을 넣고 ctax wp 를 다시 돌리면 3호 · 50호(을)에 반영됩니다")
    last = row - 1
    row += 2
    s.put(row, 1, "익금산입 · 손금불산입 합계")
    s.put(row, 7, f'=SUMIF(A{first}:A{last},"가산",G{first}:G{last})', kind="total", money=True)
    bk.reg(p.id, "익금산입 손금불산입", s.addr(row, 7))
    row += 1
    s.put(row, 1, "손금산입 · 익금불산입 합계")
    s.put(row, 7, f'=SUMIF(A{first}:A{last},"차감",G{first}:G{last})', kind="total", money=True)
    bk.reg(p.id, "손금산입 익금불산입", s.addr(row, 7))
    row += 1
    s.put(row, 1, "줄 수")
    s.put(row, 7, f'=COUNTIF(A{first}:A{last},"가산")+COUNTIF(A{first}:A{last},"차감")', fmt="0")
    bk.reg(p.id, "줄 수", s.addr(row, 7))
    row += 2
    row = notes(s, row, p)
    prior_compare.write(bk, s, row + 1, res)


TAX3_ROWS = [("01", "결산서상 당기순손익"), ("02", "익금산입"), ("03", "손금산입"), ("04", "차가감소득금액(01 + 02 - 03)"),
             ("05", "기부금한도초과액"), ("54", "기부금한도초과이월액 손금산입"), ("06", "각사업연도소득금액(04 + 05 - 54)"),
             ("07", "이월결손금"), ("08", "비과세소득"), ("09", "소득공제"), ("10", "과세표준(06 - 07 - 08 - 09)"),
             ("11", "세율(%)"), ("12", "산출세액"), ("17", "최저한세 적용대상 공제감면세액"), ("18", "차감세액(12 - 17)"),
             ("19", "최저한세 적용제외 공제감면세액"), ("20", "가산세액"), ("21", "가감계(18 - 19 + 20)"),
             ("22", "중간예납세액"), ("24", "원천납부세액"), ("28", "기납부세액 합계(22 + 24)"),
             ("29", "감면분 추가납부세액"), ("30", "차감납부할세액(21 - 28 + 29)")]
TAX3_INPUTS = ("54", "08", "09", "17", "19", "20", "22", "24", "29")


def _tax3(bk: Book, res: Result) -> None:
    p = res.tax3
    s = bk.sheet(p.sheet, [8, 44, 18, 50])
    s.ws.cell(row=1, column=1, value="3 법인세 과세표준 및 세액조정계산서").font = X.TITLE_FONT
    X.note(s.ws, 2, f"{bk.b.company} · {bk.b.fy} 사업연도 · 노란 칸은 사람이 넣는 값입니다")
    X.header(s.ws, 4, ["란", "항목", "금액", "출처"])
    at = {code: 5 + i for i, code in enumerate(c for c, _ in TAX3_ROWS)}
    R = bk.ref
    c15 = bk.cells["summary_15"]

    def G(code):
        return f"C{at[code]}"

    m = R["사업연도 월수"]
    x = f"MAX(0,{G('10')})*12/{m}"
    c1, c2, c3 = R["세율 구간1 상한"], R["세율 구간2 상한"], R["세율 구간3 상한"]
    r1, r2, r3, r4 = R["세율 구간1"], R["세율 구간2"], R["세율 구간3"], R["세율 구간4"]
    tax = (f"=ROUNDDOWN(IF({G('10')}<=0,0,IF({x}<={c1},{x}*{r1},IF({x}<={c2},{c1}*{r1}+({x}-{c1})*{r2},"
           f"IF({x}<={c3},{c1}*{r1}+({c2}-{c1})*{r2}+({x}-{c2})*{r3},"
           f"{c1}*{r1}+({c2}-{c1})*{r2}+({c3}-{c2})*{r3}+({x}-{c3})*{r4})))*{m}/12),0)")
    rate_f = (f"=IF({G('10')}<=0,0,IF({x}<={c1},{r1},IF({x}<={c2},{r2},IF({x}<={c3},{r3},{r4}))))*100")
    spec = {
        "01": (p.values["01"], "body", "장부.json 당기순이익"),
        "02": (f"={c15['익금산입 손금불산입']}", "link", "15호 합계"),
        "03": (f"={c15['손금산입 익금불산입']}", "link", "15호 합계"),
        "04": (f"={G('01')}+{G('02')}-{G('03')}", "body", ""),
        "05": (None, "link", "8 기부금 한도초과"),
        "54": (None, "input", "전기 이월 기부금 손금산입(21호)"),
        "06": (f"={G('04')}+{G('05')}-{G('54')}", "body", ""),
        "07": (None, "link", "결손금 시트"),
        "08": (None, "input", ""),
        "09": (None, "input", ""),
        "10": (f"={G('06')}-{G('07')}-{G('08')}-{G('09')}", "body", ""),
        "11": (rate_f, "body", "과세표준 구간의 최고 세율"),
        "12": (tax, "body", "법인세법 제55조 누진세율(기준 2025 시트)"),
        "17": (None, "input", _credits_src(bk, res, "최저한세 적용대상", "8호(갑) 150")),
        "18": (f"={G('12')}-{G('17')}", "body", ""),
        "19": (None, "input", _credits_src(bk, res, "최저한세 적용제외", "8호(갑) 110")),
        "20": (None, "input", "가산세(신고건 가산세가 있으면 9호 (21) 합계)"),
        "21": (f"={G('18')}-{G('19')}+{G('20')}", "body", ""),
        "22": (None, "input", "기납부세액 명세의 중간예납(적요 「중간예납」, 법인세등 · 미지급세금 차변, 7~10월 예금 출금 추정)"),
        "24": (None, "input", "기납부세액 명세의 원천납부세액(법인세분, 지방소득세 뺌)"),
        "28": (f"={G('22')}+{G('24')}", "body", ""),
        "29": (None, "input", _credits_src(bk, res, "추가납부/합계", "8호(을) 5절 추가납부세액")),
        "30": (f"={G('21')}-{G('28')}+{G('29')}", "body", ""),
    }
    for code, label in TAX3_ROWS:
        r = at[code]
        val, kind, src = spec[code]
        s.put(r, 1, code)
        s.put(r, 2, label)
        if code in TAX3_INPUTS:
            s.inp(r, 3, p.inputs[code], pid=p.id, key=code, default=p.defaults[code], kind="int", money=True)
        else:
            s.put(r, 3, val, kind=kind, money=code != "11", fmt="0" if code == "11" else None)
        s.put(r, 4, src)
        bk.reg(p.id, code, s.addr(r, 3))
    r = at["30"] + 2
    s.put(r, 1, "참고")
    s.put(r, 2, "지방소득세 특별징수(선납세금 중, 3호에 넣지 않음)")
    s.inp(r, 3, p.inputs["지방소득세 특별징수"], pid=p.id, key="지방소득세 특별징수",
          default=p.defaults["지방소득세 특별징수"], kind="int", money=True)
    s.put(r, 4, "적요 지방소득세 · 특별징수 또는 거래처가 지자체인 줄. 법인지방소득세 신고의 기납부세액")
    bk.reg(p.id, "지방소득세 특별징수", s.addr(r, 3))
    p.extra["sheet"] = s
    p.extra["row07"] = at["07"]
    p.extra["row05"] = at["05"]
    notes(s, r + 2, p)


def _credits_src(bk: Book, res: Result, key: str, form: str) -> str:
    """3호 17 · 19 출처 글. 공제감면 조서(credits)의 그 줄 칸을 부른다. 노란 칸 기본값이 그 조서 값이다."""
    p = res.paper("credits")
    addr = bk.cells.get("credits", {}).get(key)
    if p is None or not addr:
        return f"공제감면 조서 {key}({form})"
    return f"조서 「{p.sheet}」 {key} {local(addr).replace('$', '')}({form}). 고치면 그 값이 이깁니다"


def _loss(bk: Book, res: Result) -> None:
    p = res.loss
    s = bk.sheet(p.sheet, [10, 30, 4, 16, 16, 16, 16, 16, 16])
    s.ws.cell(row=1, column=1, value="결손금(50호(갑) 이월결손금 표)").font = X.TITLE_FONT
    X.note(s.ws, 2, f"{bk.b.company} · {bk.b.fy} 사업연도 · 오래된 결손금부터 공제합니다. 노란 줄은 전기 신고서에 없는 결손금(연도 · 발생액 · 기공제액)")
    R = bk.ref
    row = 4
    s.put(row, 1, "공제율")
    s.put(row, 7, f'=IF({R["중소기업 여부"]}="여",{R["이월결손금 공제율(중소기업)"]},{R["이월결손금 공제율(일반)"]})', fmt="0%")
    bk.reg(p.id, "공제율", s.addr(row, 7))
    rate_at = row
    row += 1
    s.put(row, 1, "각사업연도소득금액(3호 06)")
    s.put(row, 7, f"={bk.cells['tax_3']['06']}", kind="link", money=True)
    inc_at = row
    row += 1
    s.put(row, 1, "공제 한도")
    s.put(row, 7, f"=IF(G{inc_at}>0,ROUNDDOWN(G{inc_at}*G{rate_at},0),0)", money=True)
    bk.reg(p.id, "공제한도", s.addr(row, 7))
    cap_at = row
    row += 2
    X.header(s.ws, row, ["연도", "출처", "", "발생액", "기공제액", "", "기초 잔액", "당기 공제", "기말 잔액"])
    row += 1
    first = row
    filled = {r.get("manual") for r in p.extra["rows"] if r.get("manual")}
    blanks = [{"manual": n, "year": "", "incurred": 0, "used_before": 0, "source": "회계사 입력(비어 있음)"}
              for n in range(1, LOSS_MANUAL_ROWS + 1) if n not in filled]
    for e in blanks + p.extra["rows"]:
        n = e.get("manual")
        if n:
            key = f"결손금 수동 {n}"
            s.inp(row, 1, e["year"] or None, pid=p.id, key=f"{key}/연도", default=p.defaults.get(f"{key}/연도") or None, kind="year")
            s.inp(row, 4, e["incurred"], pid=p.id, key=f"{key}/발생액", default=p.defaults.get(f"{key}/발생액", 0), kind="int", money=True)
            s.inp(row, 5, e["used_before"], pid=p.id, key=f"{key}/기공제액", default=p.defaults.get(f"{key}/기공제액", 0), kind="int", money=True)
            s.put(row, 7, f"=D{row}-E{row}", money=True)
        else:
            s.put(row, 1, e["year"])
            s.put(row, 4, e["incurred"], money=True)
            s.put(row, 7, e["opening"], money=True)
        s.put(row, 2, e["source"])
        if row == first:
            s.put(row, 8, f"=MAX(0,MIN(G{row},$G${cap_at}))", money=True)
        else:
            s.put(row, 8, f"=MAX(0,MIN(G{row},$G${cap_at}-SUM(H${first}:H{row - 1})))", money=True)
        s.put(row, 9, f"=G{row}-H{row}", money=True)
        row += 1
    last = row - 1
    row += 1
    at = {}
    for key, label, f in [("이월결손금기초", "이월결손금 기초 잔액", f"=SUM(G{first}:G{last})"),
                          ("당기공제", "당기 공제(3호 07)", f"=SUM(H{first}:H{last})"),
                          ("당기발생", "당기 발생 결손금", f"=MAX(0,-G{inc_at})")]:
        s.put(row, 1, label)
        s.put(row, 7, f, money=True)
        bk.reg(p.id, key, s.addr(row, 7))
        at[key] = row
        row += 1
    back = CARRYBACK in p.inputs
    if back:
        s.put(row, 1, "소급공제 결손금(68호, 신청할 때만)")
        s.inp(row, 7, p.inputs[CARRYBACK], pid=p.id, key=CARRYBACK, default=p.defaults.get(CARRYBACK, 0), kind="int", money=True)
        bk.reg(p.id, CARRYBACK, s.addr(row, 7))
        at[CARRYBACK] = row
        row += 1
    s.put(row, 1, "이월결손금 기말 잔액")
    s.put(row, 7, f"=G{at['이월결손금기초']}-G{at['당기공제']}+G{at['당기발생']}" + (f"-G{at[CARRYBACK]}" if back else ""),
          kind="total", money=True)
    bk.reg(p.id, "이월결손금기말", s.addr(row, 7))
    row += 2
    if back:
        row = _carryback_rows(bk, s, p, row, at)
    notes(s, row, p)
    # 3호 07 을 잇는다.
    t3: Sheet = res.tax3.extra["sheet"]
    t3.put(res.tax3.extra.pop("row07"), 3, f"={bk.cells['loss']['당기공제']}", kind="link", money=True)


def _carryback_rows(bk: Book, s: Sheet, p, row: int, at: dict) -> int:
    """결손금 소급공제 환급세액(법인세법 제72조 제1항, 시행령 제110조 제1항). 직전 사업연도 12개월, 세율은 기준 시트."""
    R, v = bk.ref, p.values
    row = s.section(row, "결손금 소급공제 환급세액(68호)")
    put = {}
    for key, label in (("직전 과세표준", "직전 사업연도 과세표준(전기 3호 10)"), ("직전 산출세액", "직전 사업연도 산출세액(전기 3호 12)"),
                       ("직전 공제감면세액", "직전 사업연도 공제 · 감면세액(전기 3호 17 + 19)")):
        s.put(row, 1, label)
        s.put(row, 7, v[key], money=True)
        bk.reg(p.id, key, s.addr(row, 7))
        put[key] = f"G{row}"
        row += 1
    s.put(row, 1, "직전 사업연도 법인세액(환급 한도 = 산출세액 - 공제 · 감면세액)")
    s.put(row, 7, f"=MAX(0,{put['직전 산출세액']}-{put['직전 공제감면세액']})", money=True)
    bk.reg(p.id, "직전 법인세액", s.addr(row, 7))
    cap = f"G{row}"
    row += 1
    s.put(row, 1, "소급공제할 수 있는 결손금(당기 결손금과 직전 과세표준 중 작은 것)")
    s.put(row, 7, f"=MIN(G{at['당기발생']},{put['직전 과세표준']})", money=True)
    bk.reg(p.id, "소급공제 최대 결손금", s.addr(row, 7))
    most = f"G{row}"
    row += 1
    c1, c2, c3 = R["세율 구간1 상한"], R["세율 구간2 상한"], R["세율 구간3 상한"]
    r1, r2, r3, r4 = R["세율 구간1"], R["세율 구간2"], R["세율 구간3"], R["세율 구간4"]

    def tax(x):
        return (f"ROUNDDOWN(IF({x}<=0,0,IF({x}<={c1},{x}*{r1},IF({x}<={c2},{c1}*{r1}+({x}-{c1})*{r2},"
                f"IF({x}<={c3},{c1}*{r1}+({c2}-{c1})*{r2}+({x}-{c2})*{r3},"
                f"{c1}*{r1}+({c2}-{c1})*{r2}+({c3}-{c2})*{r3}+({x}-{c3})*{r4})))),0)")

    for key, label, loss in (("소급공제 환급세액", "환급세액(노란 칸 결손금으로)", f"G{at[CARRYBACK]}"),
                             ("소급공제 최대 환급세액", "최대 환급세액(소급공제할 수 있는 결손금 전부)", most)):
        x = f"({put['직전 과세표준']}-{loss})"
        s.put(row, 1, label)
        s.put(row, 7, f"=MAX(0,MIN({put['직전 산출세액']}-{tax(x)},{cap}))", money=True)
        bk.reg(p.id, key, s.addr(row, 7))
        row += 1
    return row + 1


def _reserve50(bk: Book, res: Result) -> None:
    from .summary import reserve_move
    p = res.reserve50
    s = bk.sheet(p.sheet, [6, 30, 18, 18, 18, 18, 30])
    s.ws.cell(row=1, column=1, value="50 자본금과 적립금 조정명세서").font = X.TITLE_FONT
    X.note(s.ws, 2, f"{bk.b.company} · {bk.b.fy} 사업연도 · 기초는 전기 50호(을) 기말, 증감은 조서의 유보 조정")
    row = 4
    row = s.section(row, "50호(을) 유보")
    X.header(s.ws, row, ["", "과목", "기초", "감소", "증가", "기말", "출처"])
    row += 1
    first = row
    links: dict[tuple[str, str], list[str]] = {}
    openings = {r["subject"]: r["opening"] for r in p.extra["rows"]}
    sources = [(bk.cells[paper.id][f"adj:{a.kind}"], a) for paper, a in res.adjustments()]
    sources += [(bk.cells["summary_15"][f"direct:{a.row}"], a) for a in res.direct if a.amount]
    for ref, a in sources:
        if not a.reserve:
            continue
        subj = a.reserve_subject or a.subject
        col, amt = reserve_move({"opening": openings.get(subj, 0)}, a)
        sign = 1 if amt == a.amount else -1
        links.setdefault((subj, col), []).append(("-" if sign < 0 else "+") + ref)
    for r in p.extra["rows"]:
        subj = r["subject"]
        s.put(row, 2, subj)
        s.put(row, 3, r["opening"], money=True)
        for col, name in ((4, "감소"), (5, "증가")):
            parts = links.get((subj, name), [])
            f = "=" + "".join(parts).lstrip("+") if parts else 0
            s.put(row, col, f, kind="link" if parts else "body", money=True)
        s.put(row, 6, f"=C{row}-D{row}+E{row}", money=True)
        s.put(row, 7, "전기 50호(을)" if r["opening"] else "당기 조정")
        row += 1
    if not p.extra["rows"]:
        s.put(row, 2, "유보 없음")
        row += 1
    last = row - 1
    s.put(row, 2, "합계")
    for col, letter in ((3, "C"), (6, "F")):
        s.put(row, col, f"=SUM({letter}{first}:{letter}{last})", kind="total", money=True)
    bk.reg(p.id, "유보기초", s.addr(row, 3))
    bk.reg(p.id, "유보기말", s.addr(row, 6))
    tot_at = row
    row += 2
    row = s.section(row, "50호(갑) 자본금과 적립금")
    X.header(s.ws, row, ["", "과목", "기초(전기 기말)", "", "", "기말", "출처"])
    row += 1
    prior = p.extra["prior"]
    eq_first = row
    for name in ("자본금", "자본잉여금", "자본조정", "기타포괄손익누계액", "이익잉여금"):
        s.put(row, 2, name)
        s.put(row, 3, prior.get(name), money=True)
        s.put(row, 6, p.values[f"기말 {name}"], money=True)
        s.put(row, 7, "재무상태표")
        bk.reg(p.id, f"기말 {name}", s.addr(row, 6))
        row += 1
    s.put(row, 2, "계")
    s.put(row, 3, prior.get("계"), money=True)
    s.put(row, 6, f"=SUM(F{eq_first}:F{row - 1})", kind="total", money=True)
    bk.reg(p.id, "기말 계", s.addr(row, 6))
    sum_at = row
    row += 1
    s.put(row, 2, "유보(50호(을) 합계)")
    s.put(row, 6, f"=F{tot_at}", money=True)
    row += 1
    s.put(row, 2, "차가감계")
    s.put(row, 3, prior.get("차가감계"), money=True)
    s.put(row, 6, f"=F{sum_at}+F{row - 1}", kind="total", money=True)
    bk.reg(p.id, "차가감계", s.addr(row, 6))
    row += 2
    notes(s, row, p)


def _overview(bk: Book, s: Sheet, res: Result) -> None:
    ws = s.ws
    ws.cell(row=1, column=1, value=f"{bk.b.company} {bk.b.fy} 법인세 조서").font = X.TITLE_FONT
    X.note(ws, 2, "Gridie 법인세 조서 · 계산은 엔진, 노란 칸은 회계사가 고치는 칸입니다. 고친 뒤 ctax wp 를 다시 돌리면 그 값을 읽습니다")
    X.header(ws, 4, ["번호", "조서", "상태", "가산", "차감", "AR 이슈", "이어지는 서식"])
    row = 5
    for p in res.all_papers:
        ws.cell(row=row, column=1, value=p.no)
        link = ws.cell(row=row, column=2, value=p.name)
        link.hyperlink = f"#'{p.sheet}'!A1"
        ws.cell(row=row, column=3, value=p.status)
        add = sum(a.amount for a in p.live_adjustments() if a.side == "가산")
        ded = sum(a.amount for a in p.live_adjustments() if a.side == "차감")
        if p.id == "summary_15":
            add = sum(a.amount for a in res.direct if a.side == "가산")
            ded = sum(a.amount for a in res.direct if a.side == "차감")
        X.cell(ws, row, 4, add, money=True)
        X.cell(ws, row, 5, ded, money=True)
        ws.cell(row=row, column=6, value=", ".join(f"#{i}" for i in p.issues))
        ws.cell(row=row, column=7, value=", ".join(p.forms))
        row += 1
    row += 1
    X.header(ws, row, ["란", "3호", "금액"])
    row += 1
    for code, label in (("01", "당기순손익"), ("06", "각사업연도소득금액"), ("10", "과세표준"), ("12", "산출세액"),
                        ("30", "차감납부할세액")):
        ws.cell(row=row, column=1, value=code)
        ws.cell(row=row, column=2, value=label)
        X.cell(ws, row, 3, f"={bk.cells['tax_3'][code]}", kind="link", money=True)
        row += 1
    for c in ws[4]:
        c.alignment = Alignment(horizontal="center")


def _inputs_sheet(bk: Book) -> None:
    """숨은 시트 `_입력`: 노란 칸마다 조서 · key · 시트 · 칸 · 엔진 기본값 · 종류. 다음 실행이 이 표로 회계사 값을 읽는다."""
    ws = bk.wb.create_sheet(INPUT_SHEET)
    ws.append(["paper", "key", "sheet", "cell", "default", "kind"])
    for c in bk.input_cells:
        ws.append([c["paper"], c["key"], c["sheet"], c["cell"], c["default"], c["kind"]])
    ws.sheet_state = "hidden"


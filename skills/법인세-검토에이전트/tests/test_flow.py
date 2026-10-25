"""스킬 흐름 시험: 단계 검사 → 메모 → 공제 제안 → 마무리. 엔진 시험의 합성 장부 · 전표 · 전기만 쓴다(고객 자료 없음)."""
from __future__ import annotations

import csv
import json
import re

import openpyxl
import pytest

from ctax.ar import cli as ar_cli
from ctax.tests.test_ar import journal, ledger, prior

from bestar import finish, paths, rank, steps
from bestar import credits as credits_mod
from bestar import memos

SME_TESTS = ("업종", "규모", "자산총액", "독립성", "유예기간", "소기업")


@pytest.fixture
def case(tmp_path):
    root = tmp_path / "가상회사" / "2025 법인세"
    idx = root / "_색인"
    idx.mkdir(parents=True)
    (idx / "장부.json").write_text(json.dumps(ledger(), ensure_ascii=False), encoding="utf-8")
    (idx / "전기.json").write_text(json.dumps(prior(), ensure_ascii=False), encoding="utf-8")
    rows = journal()
    with (idx / "전표.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    assert ar_cli.main([str(root)]) == 0
    return root


def issues(root):
    return paths.read_ar(paths.case(root))["issues"]


def find(root, word):
    return next(i["no"] for i in issues(root) if word in i["title"])


def write_step(root, sid, carry, **extra):
    data = {"id": sid, "step": steps.LABEL[sid], "finding": f"{sid} 자료상 확인했습니다.", "note": "다음 단계로 넘깁니다.",
            "carry": carry, **extra}
    d = root / "_색인" / "AR_단계"
    d.mkdir(parents=True, exist_ok=True)
    (d / f"{sid}.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def good_steps(root, f2_tests=None):
    ent = find(root, "기업업무추진비")
    write_step(root, "P1", {"class": "중소기업", "small": "중기업", "industry": "도매", "grace": "없음", "effects": []})
    write_step(root, "P2", {"stock_change": "미확인", "stock_form": "미확인", "rental_main": "해당 아님", "effects": []})
    write_step(root, "P3", {"forecast": [
        {"id": "f1", "src": "adj:0", "item": "접대비 한도초과", "kind": "반복", "outlook": "다시 생길 듯", "why": "전기 1,000,000원", "issues": [ent]},
        {"id": "f2", "src": "adj:1", "item": "법인세등", "kind": "반복", "outlook": "다시 생길 듯", "why": "법인세등 있음", "issues": [find(root, "법인세등")]},
        {"id": "f3", "src": "adj:2", "item": "미수이자", "kind": "유보 추인", "outlook": "안 생길 듯", "why": "미수수익 계정 없음", "issues": []},
        {"id": "f4", "src": "res:0", "item": "미수수익 유보", "kind": "유보 추인", "outlook": "안 생길 듯", "why": "계정 없음", "issues": []}]})
    write_step(root, "P4", {"forecast": [
        {"id": "c1", "src": "credit:0", "name": "중소기업에 대한 특별세액감면", "outlook": "계속", "why": "업종 같음", "issues": []}]})
    write_step(root, "F1", {"forecast_check": [{"forecast": "f1", "status": "계정 있음"}], "accounts": [
        {"id": "a1", "account": "주.임.종단기채권", "why": "임원 채권", "issues": [find(root, "대여금")]},
        {"id": "a2", "account": "차량유지비", "issues": [], "finding": "자료상 렌트 차량 대여료가 있습니다.", "why": "업무용승용차",
         "risk": "중", "judgment": "판단 필요", "ask": ["차량 계약서를 보내 주십시오"]}]})
    f1 = json.loads((root / "_색인/AR_단계/F1.json").read_text(encoding="utf-8"))
    out = steps.outside_accounts(json.loads((root / "_색인/장부.json").read_text(encoding="utf-8")), paths.read_ar(paths.case(root)), f1["carry"])
    f1["carry"]["others_reviewed"] = [{"accounts": [o["name"] for o in out], "note": "일반 계정입니다. 조정 사유가 보이지 않습니다."}]
    (root / "_색인/AR_단계/F1.json").write_text(json.dumps(f1, ensure_ascii=False), encoding="utf-8")
    tests = f2_tests or [{"name": n, "result": "충족"} for n in SME_TESTS if n != "독립성"] + [{"name": "독립성", "result": "미확인"}]
    write_step(root, "F2", {"verdict": "중소기업", "status": "확인 필요", "premise": "중소기업", "compare_p1": "같음",
                            "point": "독립성만 미확인", "tests": tests, "effects": "", "ask": ["주주명부를 보내 주십시오"]})


def run_steps(root):
    c = paths.case(root)
    for sid in steps.ORDER:
        assert steps.cmd_done(c, sid, False) == 0, steps.check_step(c, sid)


# ---------------------------------------------------------------- 단계 검사

def test_order_is_enforced(case):
    good_steps(case)
    c = paths.case(case)
    assert steps.cmd_done(c, "P2", False) == 1
    assert steps.cmd_done(c, "P1", False) == 0
    assert steps.cmd_done(c, "P2", False) == 0
    assert steps.cmd_done(c, "P2", True) == 0                 # 고치면 뒤 단계 무효
    assert [e["id"] for e in steps.read_log(c)] == ["P1", "P2"]


def test_p3_must_cover_every_prior_line(case):
    good_steps(case)
    p3 = json.loads((case / "_색인/AR_단계/P3.json").read_text(encoding="utf-8"))
    p3["carry"]["forecast"] = p3["carry"]["forecast"][:2]
    (case / "_색인/AR_단계/P3.json").write_text(json.dumps(p3, ensure_ascii=False), encoding="utf-8")
    probs = steps.check_step(paths.case(case), "P3")
    assert any("빠짐 ['adj:2', 'res:0']" in p for p in probs)


def test_amount_not_in_engine_is_rejected(case):
    good_steps(case)
    write_step(case, "P1", {"class": "중소기업", "small": "중기업", "industry": "도매", "grace": "없음",
                            "effects": ["기본한도 12,345,678원"]})
    probs = steps.check_step(paths.case(case), "P1")
    assert any("12,345,678" in p for p in probs)


def test_question_mark_and_class_mismatch(case):
    good_steps(case)
    write_step(case, "P1", {"class": "중견기업", "small": "해당 없음", "industry": "도매", "grace": "없음", "effects": ["맞나?"]})
    probs = steps.check_step(paths.case(case), "P1")
    assert any("differs_reason" in p for p in probs)
    assert any("물음표" in p for p in probs)


def test_f1_added_issue_needs_finding(case):
    good_steps(case)
    c = paths.case(case)
    for sid in ("P1", "P2", "P3", "P4"):
        assert steps.cmd_done(c, sid, False) == 0
    write_step(case, "F1", {"forecast_check": [], "accounts": [{"id": "a1", "account": "용역비", "issues": []}]})
    probs = steps.check_step(c, "F1")
    assert any("finding" in p for p in probs) and any("risk" in p for p in probs)


# ---------------------------------------------------------------- 메모 · 제안 · 마무리

def memo_file(root, tmp_path, answers=True):
    items = []
    for i in issues(root):
        it = {"issue": i["no"], "judgment": "판단 필요", "why": "자료를 받아 판단합니다.", "law": [{"ref": "시험", "grade": "☆"}],
              "questions": ["자료를 보내 주십시오"], "risk": "하", "answers": []}
        if "기업업무추진비" in i["title"]:
            it.update(judgment="조정 필요", risk="중", answers=["f1"] if answers else [])
        if "법인세등" in i["title"] and answers:
            it["answers"] = ["f2"]
        if "대여금" in i["title"]:
            it.update(risk="상", answers=["a1"] if answers else [])
        items.append(it)
    p = tmp_path / "memo.json"
    p.write_text(json.dumps({"items": items}, ensure_ascii=False), encoding="utf-8")
    return p


def credit_file(tmp_path, line=1, amount=25_000_000):
    p = tmp_path / "credits.json"
    p.write_text(json.dumps({"point": "전기 감면을 이어서 봅니다.", "items": [
        {"name": "중소기업에 대한 특별세액감면", "basis": "조세특례제한법 제7조", "state": "진행 중", "purpose": "감면 업종 중소기업 지원",
         "candidates": [{"what": "접대비", "line": line, "amount": amount}], "judge": ["감면 업종인지"],
         "ask": ["매출 구성을 보내 주십시오"], "issues": [], "answers": ["c1"]}]}, ensure_ascii=False), encoding="utf-8")
    return p


def test_memo_needs_steps_and_valid_answers(case, tmp_path):
    good_steps(case)
    assert memos.main([str(case), str(memo_file(case, tmp_path))]) == 2      # 단계 전
    run_steps(case)
    bad = memo_file(case, tmp_path)
    data = json.loads(bad.read_text(encoding="utf-8"))
    data["items"][0]["answers"] = ["f9"]
    bad.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert memos.main([str(case), str(bad)]) == 1
    assert memos.main([str(case), str(memo_file(case, tmp_path))]) == 0
    assert all(i.get("memo") for i in issues(case))


def test_credit_line_amount_checked(case, tmp_path):
    good_steps(case)
    run_steps(case)
    c = paths.case(case)
    probs = credits_mod.problems(c, json.loads(credit_file(tmp_path, amount=24_000_000).read_text(encoding="utf-8")))
    assert any("다릅니다" in p for p in probs)
    assert credits_mod.main([str(case), str(credit_file(tmp_path))]) == 0


def test_finish_full(case, tmp_path):
    good_steps(case)
    run_steps(case)
    assert memos.main([str(case), str(memo_file(case, tmp_path))]) == 0
    assert finish.main([str(case)]) == 1                     # 공제 제안 없음
    assert credits_mod.main([str(case), str(credit_file(tmp_path))]) == 0
    assert finish.main([str(case)]) == 0
    s = paths.read_json(case / "_색인" / paths.SUMMARY)["summary"]
    top = [u["title"] for u in s["top"]]
    assert len(top) == 3
    assert not any("법인세등" in t for t in top)              # 상시 항목은 셋에서 뺀다
    assert "기업업무추진비" in top[0]                          # 조정 필요 · 중(3) 과 판단 필요 · 상(3) 이 같으면 조정 필요 먼저
    assert "대여금" in top[1]
    assert any(u["uid"] == "sme" for u in s["check"])         # 독립성만 미확인이면 확인사항
    assert any(u["routine"] for u in s["check"])
    wb = openpyxl.load_workbook(case / paths.XLSX)
    assert wb.sheetnames[0] == "요약"
    assert {"이슈 목록", "AI 추가 이슈", "공제 제안", "순서와 전제"} <= set(wb.sheetnames)
    md = (case / paths.REPORT).read_text(encoding="utf-8")
    assert "## 훑은 범위" in md and "## 1. 전사 기준" in md and "## 3. 세액공제" in md and "엔진 규칙에 안 걸린 계정" in md
    from bestar import view_cmd
    assert view_cmd.main([str(case)]) == 0
    html = (case / paths.VIEW).read_text(encoding="utf-8")
    assert "window.DATA" in html and "/*DATA*/" not in html
    art = case.parent / "art.html"
    assert view_cmd.main([str(case), "--artifact", str(art)]) == 0
    t = art.read_text(encoding="utf-8")
    assert t.startswith("<title>") and "<!doctype" not in t.lower() and "<body" not in t.lower()


def test_finish_reports_unanswered(case, tmp_path):
    good_steps(case)
    run_steps(case)
    assert memos.main([str(case), str(memo_file(case, tmp_path, answers=False))]) == 0
    assert credits_mod.main([str(case), str(credit_file(tmp_path))]) == 0
    c = paths.case(case)
    ar = paths.read_ar(c)
    rec = finish.received(c, ar, memos.extras(c), paths.read_json(case / "_색인" / paths.CREDITS), finish.added_issues(c, steps.load_step(c, "F1")["carry"]))
    gaps = finish.gaps(c, ar, rec, {})
    assert any("f1" in g for g in gaps) and any("a1" in g for g in gaps)
    assert not any("f3" in g for g in gaps)                   # 안 생길 듯 으로 닫은 예측
    assert not any(" a2 " in g for g in gaps)                 # AI 추가 이슈는 그 자체로 받음


def test_added_conclusion_survives_rerun(case, tmp_path):
    good_steps(case)
    run_steps(case)
    assert memos.main([str(case), str(memo_file(case, tmp_path))]) == 0
    assert credits_mod.main([str(case), str(credit_file(tmp_path))]) == 0
    assert finish.main([str(case)]) == 0
    wb = openpyxl.load_workbook(case / paths.XLSX)
    ws = wb["AI 추가 이슈"]
    ws.cell(2, 8).value = "렌트 3대, 29호 작성"
    wb.save(case / paths.XLSX)
    assert finish.main([str(case)]) == 0
    wb = openpyxl.load_workbook(case / paths.XLSX)
    assert wb["AI 추가 이슈"].cell(2, 8).value == "렌트 3대, 29호 작성"


def test_rank_independence_and_failed_test():
    base = {"P3": {"carry": {"forecast": []}}, "P4": {"carry": {"forecast": []}}, "F1": {"carry": {"accounts": []}}}
    ar = {"issues": [{"no": 1, "kind": "조정 후보", "certainty": "확실", "title": "큰 것", "rule": "x", "amounts": {"current": 900}, "memo": {"judgment": "판단 필요"}, "guide": []},
                     {"no": 2, "kind": "조정 후보", "certainty": "확실", "title": "작은 것", "rule": "y", "amounts": {"current": 10}, "memo": {"judgment": "판단 필요"}, "guide": []}]}
    f2 = {"status": "확인 필요", "compare_p1": "같음", "verdict": "중소기업", "tests": [{"name": "규모", "result": "미충족"}]}
    s = rank.build(ar, {2: {"risk": "상"}}, {**base, "F2": {"carry": f2}}, [], {}, [])
    assert [u["uid"] for u in s["top"]] == ["issue:2", "sme", "issue:1"]     # 위험도 상이 크기보다 먼저, 같은 위험도는 크기 순


def test_facts_sum_and_pdf_make_numbers_usable(case, tmp_path):
    from bestar import facts, numbers
    good_steps(case)
    c = paths.case(case)
    assert facts.main("sum", [str(case), "--what", "접대비 두 줄", "--lines", "1-2"]) == 0
    assert facts.main("pdf", [str(case), "--page", "3", "--form", "29", "--what", "x", "--amount", "1"]) == 2   # PDF 없음
    (case / "00 받은자료" / "전기").mkdir(parents=True)
    (case / "00 받은자료" / "전기" / "가상.pdf").write_bytes(b"%PDF-1.4")
    assert facts.main("pdf", [str(case), "--page", "3", "--form", "29", "--what", "업무외사용금액", "--amount", "5930000"]) == 0
    got = facts.load(c)
    assert got[0]["id"] == "s1" and got[0]["debit"] == 30_000_000 and got[1]["id"] == "p1"
    known = numbers.known(c)
    assert 30_000_000 in known and 5_930_000 in known
    write_step(case, "P1", {"class": "중소기업", "small": "중기업", "industry": "도매", "grace": "없음",
                            "effects": ["전표 10234줄 확인, 전기 29호 업무외사용금액 5,930,000원"]})
    assert steps.check_step(c, "P1") == []


def test_same_as_keeps_top_distinct():
    base = {"P3": {"carry": {"forecast": []}}, "P4": {"carry": {"forecast": []}}, "F1": {"carry": {"accounts": []}},
            "F2": {"carry": {"status": "확정", "compare_p1": "같음"}}}
    mk = lambda n, size: {"no": n, "kind": "조정 후보", "certainty": "확실", "title": f"이슈{n}", "rule": f"r{n}",
                          "amounts": {"current": size}, "memo": {"judgment": "조정 필요"}, "guide": []}
    ar = {"issues": [mk(1, 900), mk(2, 800), mk(3, 10), mk(4, 5)]}
    s = rank.build(ar, {1: {"risk": "상"}, 2: {"risk": "상", "same_as": 1}}, base, [], {}, [])
    assert [u["ref"] for u in s["top"]] == [1, 3, 4]


def test_tax_effect_issue_reaches_top():
    """10-19: 전기 결손금 표에 빠진 해(세액이 달라짐)는 위험 중 · 판단 필요여도 셋에 오른다."""
    base = {"P3": {"carry": {"forecast": []}}, "P4": {"carry": {"forecast": []}}, "F1": {"carry": {"accounts": []}},
            "F2": {"carry": {"status": "확정", "compare_p1": "같음"}}}
    mk = lambda n, rule, judge, size, kind="조정 후보": {"no": n, "kind": kind, "certainty": "확실", "title": f"이슈{n}", "rule": rule,
                                                        "amounts": {"current": size}, "memo": {"judgment": judge}, "guide": []}
    ar = {"issues": [mk(1, "loan", "조정 필요", 900), mk(2, "a", "판단 필요", 800), mk(3, "b", "판단 필요", 700),
                     mk(4, "prior_loss_row", "판단 필요", 10, "신고 점검"), mk(5, "c", "판단 필요", 5)]}
    ex = {1: {"risk": "상"}, 2: {"risk": "상"}, 3: {"risk": "상"}, 4: {"risk": "중"}, 5: {"risk": "중", "tax_effect": True}}
    s = rank.build(ar, ex, base, [], {}, [])
    assert [u["ref"] for u in s["top"]] == [1, 4, 5]
    s = rank.build(ar, ex, base, [], {}, [], taxable=False)              # 당기 결손이면 결손금 이력 규칙은 더하지 않는다
    assert [u["ref"] for u in s["top"]] == [1, 5, 2]


def test_same_as_list_in_memo(case, tmp_path):
    good_steps(case)
    run_steps(case)
    p = memo_file(case, tmp_path)
    data = json.loads(p.read_text(encoding="utf-8"))
    nos = [it["issue"] for it in data["items"]]
    data["items"][0]["same_as"] = [nos[1]]
    data["items"][1]["same_as"] = "x"
    p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert memos.main([str(case), str(p)]) == 1                       # 글자는 거절(TypeError 가 아니라 검사 문구)
    data["items"][1]["same_as"] = None
    p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert memos.main([str(case), str(p)]) == 0
    ex = memos.extras(paths.case(case))
    assert ex[nos[0]]["same_as"] == [nos[1]]


# ---------------------------------------------------------------- 목업 화면 자료 · 판단 저장

CHOICES = {"title": "대여 상대방은 누구인가요", "description": "특수관계인이면 인정이자를 계산해요.", "multiple": False, "allow_other": True,
           "options": [{"value": "rel", "label": "임원 · 주주예요", "effect": "인정이자 익금산입"},
                       {"value": "none", "label": "특수관계가 없어요", "effect": "조정 없음"}]}


def finished_case(case, tmp_path, with_choices=True):
    good_steps(case)
    run_steps(case)
    p = memo_file(case, tmp_path)
    if with_choices:
        data = json.loads(p.read_text(encoding="utf-8"))
        for it in data["items"]:
            if it["risk"] == "상":
                it["choices"] = CHOICES
        p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert memos.main([str(case), str(p)]) == 0
    assert credits_mod.main([str(case), str(credit_file(tmp_path))]) == 0
    assert finish.main([str(case)]) == 0
    return paths.case(case)


def test_choices_checked(case, tmp_path):
    good_steps(case)
    run_steps(case)
    p = memo_file(case, tmp_path)
    data = json.loads(p.read_text(encoding="utf-8"))
    data["items"][0]["choices"] = {"title": "x", "options": [{"value": "a", "label": "하나뿐 12,345,678원"}]}
    p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert memos.main([str(case), str(p)]) == 1


def test_review_data_and_decisions(case, tmp_path):
    from bestar import decide, review
    c = finished_case(case, tmp_path)
    d = review.build(c)
    assert d["adjustments"], "조서 조정이 화면에 있어야 한다"
    assert all(a["uid"] in d["cards"] for a in d["adjustments"])
    q = d["questions"][0]
    assert q["title"] == CHOICES["title"] and q["answer"] is None
    assert any(j["uid"] == q["uid"] and j["state"] == "답변 필요" for j in d["judgments"])
    assert d["income"]["net_income"] is not None and d["stages"][2]["note"] == "1건 남음"
    decide.decide(c, d["keys"][q["uid"]], value="rel")
    decide.accept(c, [d["keys"][a["uid"]] for a in d["adjustments"]])
    d2 = review.build(c)
    assert d2["questions"][0]["answer"]["value"] == "rel"
    assert d2["cards"][q["uid"]]["result"] == "임원 · 주주예요" and d2["cards"][q["uid"]]["badge"] == "반영됨"
    assert all(a["accepted"] for a in d2["adjustments"]) and d2["stages"][2]["note"] == "완료"
    assert finish.main([str(case)]) == 0                       # 다시 돌려도 판단은 남는다
    assert review.build(c)["questions"][0]["answer"]["value"] == "rel"
    wb = openpyxl.load_workbook(case / paths.XLSX)
    assert wb["화면 판단"].max_row >= 2


def test_serve_round_trip(case, tmp_path):
    import threading
    import urllib.request
    from http.server import ThreadingHTTPServer
    from bestar import serve
    c = finished_case(case, tmp_path)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), serve.make_handler(c))
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    base = f"http://127.0.0.1:{srv.server_address[1]}/"
    try:
        html = urllib.request.urlopen(base).read().decode("utf-8")
        assert '"live"' in html and "/*DATA*/" not in html
        d = json.loads(urllib.request.urlopen(base + "api/data").read().decode("utf-8"))
        uid = d["questions"][0]["uid"]
        req = urllib.request.Request(base + "api/decide", data=json.dumps({"uid": uid, "value": "none"}).encode("utf-8"), method="POST",
                                     headers={"Content-Type": "application/json"})
        d2 = json.loads(urllib.request.urlopen(req).read().decode("utf-8"))
        assert d2["questions"][0]["answer"]["value"] == "none"
        req = urllib.request.Request(base + "api/undo", data=json.dumps({"uid": uid}).encode("utf-8"), method="POST")
        assert json.loads(urllib.request.urlopen(req).read().decode("utf-8"))["questions"][0]["answer"] is None
    finally:
        srv.shutdown()


def test_ar_sheet_three_groups(case, tmp_path):
    from bestar import review
    c = finished_case(case, tmp_path)
    groups = review.build(c)["ar_sheet"]["groups"]
    assert [g["label"] for g in groups] == ["전사 기준", "재무제표 계정별", "세액공제"]
    keys = [r["key"] for r in groups[0]["rows"]]
    assert keys[:5] == ["P1", "P2", "P3", "F1", "F2"]
    acct = {r["title"]: r for r in groups[1]["rows"]}
    ent = next(r for t, r in acct.items() if "기업업무추진비" in t)
    assert ent["judgment"] in ("조정 필요", "판단 필요") and ent["uids"]
    assert any(r["key"].startswith("added:") for r in groups[1]["rows"])          # AI 추가 이슈는 계정별
    assert groups[2]["rows"][0]["key"] == "P4"
    assert any(r["key"].startswith("credit:") for r in groups[2]["rows"])
    for g in groups:
        for r in g["rows"]:
            assert {"title", "relation", "judgment", "judgment_text", "asks"} <= set(r)
    wb = openpyxl.load_workbook(case / paths.XLSX)
    assert wb.sheetnames[1] == "AR 한 장" and wb["AR 한 장"].max_row > 10
    assert wb["AR 한 장"].auto_filter.ref and wb["AR 한 장"].cell(1, 2).value == "이슈"
    sheet = review.build(c)["ar_sheet"]
    assert all("issue" in r for g in sheet["groups"] for r in g["rows"])
    assert sheet["coverage"]["guide_lines"] == 39 and sheet["coverage"]["account_groups"] == 22
    titles = [r["title"] for g in sheet["groups"][1:] for r in g["rows"]]
    assert sum(1 for t in titles if re.match(r"\d+\)", t)) == 22             # 전체 보기에 계정 22묶음이 다 있다



def test_f1_must_review_outside_accounts(case):
    good_steps(case)
    c = paths.case(case)
    for sid in ("P1", "P2", "P3", "P4"):
        assert steps.cmd_done(c, sid, False) == 0
    f1 = json.loads((case / "_색인/AR_단계/F1.json").read_text(encoding="utf-8"))
    f1["carry"]["others_reviewed"] = []
    (case / "_색인/AR_단계/F1.json").write_text(json.dumps(f1, ensure_ascii=False), encoding="utf-8")
    probs = steps.check_step(c, "F1")
    assert any("others_reviewed" in p for p in probs)


# ---------------------------------------------------------------- 10-19 나참유통 · 조미코리아 실행에서 고친 것

def test_step_rejects_not_but_sentence(case):
    good_steps(case)
    write_step(case, "P1", {"class": "중소기업", "small": "소기업", "industry": "도매", "grace": "없음",
                            "effects": ["감면 대상이 아니라 공제 대상입니다"]})
    probs = steps.check_step(paths.case(case), "P1")
    assert any("앞을 부정하고 뒤를 말하는" in p for p in probs)


def test_credit_row_has_no_candidate_amount_and_judge_not_in_asks(case, tmp_path):
    from bestar import review
    c = finished_case(case, tmp_path)
    rows = review.build(c)["ar_sheet"]["groups"][2]["rows"]
    cr = next(r for r in rows if r["key"].startswith("credit:"))
    assert cr["amount"] is None and "공제 · 감면액이 아님" in cr["judgment_text"] and "판단할 점: 감면 업종인지" in cr["judgment_text"]
    assert "감면 업종인지" not in cr["asks"] and "매출 구성을 보내 주십시오" in cr["asks"]


def test_added_title_is_short():
    long_why = "가" * 90 + "입니다. 둘째 문장"
    assert rank.added_title({"account": "전도금", "why": long_why}).endswith("…")
    assert len(rank.added_title({"account": "전도금", "why": long_why})) <= len("전도금: ") + rank.TITLE_MAX
    assert rank.added_title({"account": "전도금", "title": "임원 가지급금으로 보임", "why": long_why}) == "전도금: 임원 가지급금으로 보임"


def test_page_command_renders_prior_pdf(case):
    import pymupdf
    from bestar import page
    d = case / "00 받은자료" / "전기"
    d.mkdir(parents=True)
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "29")
    doc.save(str(d / "가상.pdf"))
    assert page.main([str(case), "2"]) == 2
    assert page.main([str(case), "1"]) == 0
    assert (paths.steps_dir(paths.case(case)) / "전기_1쪽.png").exists()

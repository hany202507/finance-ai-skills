# -*- coding: utf-8 -*-
"""현금흐름표 실행: 입력 → 엔진 → 워크북 → 재계산 → 대조.

    python run.py --분개장 J.xlsx [--재무상태표 BS.xlsx] [--계정 accounts.xlsx]
                  [--정책 이자지급=재무 ...] [--회사 이름] [--출력 폴더]

대조는 엑셀이 계산한 값을 다시 읽어 엔진이 따로 낸 숫자와 맞춰 본다. 수식이 엔진과 다른 걸
계산하고 있으면 여기서 걸린다. 결과는 같은 폴더의 검증_<기간>.json 에 남는다.
"""
import argparse, json, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from collections import defaultdict
import 입력, engine, build_workbook, 재계산


def 대조(경로, r, 셀):
    import openpyxl
    wb = openpyxl.load_workbook(경로, data_only=True)
    직, 간, 요약 = wb["현금흐름표_직접법"], wb["현금흐름표_간접법"], wb["요약"]
    틀림 = []
    def 비교(이름, 엑셀값, 엔진값):
        if 엑셀값 is None or round(엑셀값) != 엔진값:
            틀림.append(f"{이름}: 엑셀 {엑셀값} / 엔진 {엔진값}")
    for (활, 항), v in r.직접.items():
        비교(f"직접법 {활} {항}", 직[셀["직접법"][(활, 항)]].value, v)
    for 활 in engine.활동순서:
        비교(f"직접법 {활} 합계", 직[셀["직접법"][활]].value, r.직접합[활])
        비교(f"간접법 {활} 합계", 간[셀["간접법"][활]].value, r.직접합[활])
    비교("간접법 영업(엔진 최종)", 간[셀["간접법"][engine.영업]].value, r.간접최종영업합)
    정 = wb["간접법_정산표"]
    for 활 in engine.활동순서:
        비교(f"정산표 {활}", 정[셀["정산"][활]].value, r.직접합[활])
    비교("정산표 행 검산", 정[셀["정산"]["행검산"]].value, 0)
    비교("현금 증감", 직[셀["직접법"]["증감"]].value, r.현금증감)
    비교("기초 현금", 직[셀["직접법"]["기초"]].value, r.기초현금)
    독엑셀 = 요약[셀["독종합"]].value
    if (독엑셀 == "PASS") != r.독립통과:
        틀림.append(f"독립 검산 종합: 엑셀 {독엑셀} / 엔진 {'PASS' if r.독립통과 else 'FAIL'}")
    a, b = 셀["요약검증행"]
    판정 = [(요약.cell(row=i, column=1).value, 요약.cell(row=i, column=3).value) for i in range(a, b + 1)]
    for 라벨, p in 판정:
        if p != "PASS":
            틀림.append(f"요약 검증 {라벨}: {p}")
    return 틀림, 판정


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--분개장", required=True, nargs="+", help="여러 파일(분기별)이면 모두")
    ap.add_argument("--분개장시트")
    ap.add_argument("--재무상태표")
    ap.add_argument("--계정")
    ap.add_argument("--손익계산서", help="공시 양식 손익계산서(당기순이익 대조)")
    ap.add_argument("--손익계산서시트")
    ap.add_argument("--재무상태표시트")
    ap.add_argument("--시산표", help="기말 합계잔액시산표. 주면 기초를 시산표 - 분개장 증감으로 만든다")
    ap.add_argument("--현금계정", nargs="*", default=[], help="현금성자산으로 볼 계정코드(예: 만기 3개월 이내 정기예금)")
    ap.add_argument("--투자계정", nargs="*", default=[], help="현금성자산이 아니라 투자(금융상품)로 볼 계정코드(예: 만기 3개월 넘는 정기예금)")
    ap.add_argument("--세부유지", action="store_true", help="7자리 세부계정을 상위계정으로 합치지 않는다")
    ap.add_argument("--정책", nargs="*", default=[])
    ap.add_argument("--기준", default="일반기업", choices=["일반기업", "중소기업", "K-IFRS", "K-IFRS1118"],
                    help="K-IFRS1118: 2027년 이후. 이자·배당 분류 고정, 간접법은 영업이익에서 출발")
    ap.add_argument("--판정", nargs="*", default=[], help="회계사 판정 파일(판정후보_<기간>.csv 를 채운 것). 여러 달 것을 이어 써도 된다")
    ap.add_argument("--상세", nargs="*", default=[], help="가계정을 발라낼 은행거래내역·미분류목록")
    ap.add_argument("--상세필터", help="상세 파일에서 이 값이 든 행만(예: 미분류)")
    ap.add_argument("--명세서", help="전기말 추적 계정의 거래처별 잔액(계정명세서)")
    ap.add_argument("--회사", default="")
    ap.add_argument("--출력")
    a = ap.parse_args(argv)

    줄, 전기, 기말, 마스터, 메모, 기간, 부가 = 입력.준비(a.분개장, a.재무상태표, a.계정, a.분개장시트,
                                              세부통합=not a.세부유지, 손익계산서경로=a.손익계산서,
                                              시산표경로=a.시산표, 재무상태표시트=a.재무상태표시트,
                                              손익계산서시트=a.손익계산서시트)
    손익NI = 부가.get("손익계산서", {}).get("당기")
    정책 = dict(x.split("=", 1) for x in a.정책)
    상세 = [d for p in a.상세 for d in 입력.상세(p, a.상세필터)]
    명세 = 입력.계정명세서(a.명세서) if a.명세서 else []
    if not 명세 and a.재무상태표:
        try:                       # 기초잔액 파일에 거래처·원천계정 열이 있으면 그 자체가 명세서다
            명세 = 입력.계정명세서(a.재무상태표)
        except Exception:
            명세 = []
    이름코드 = {l["계정"]: l["코드"] for l in 줄}
    for m in 명세:
        if isinstance(m["코드"], str) or m["코드"] is None:
            m["코드"] = 이름코드.get(m["계정"] or m["코드"], m["코드"])
    현금추가 = {int(x) if x.isdigit() else x for x in a.현금계정}
    투자지정 = {int(x) if x.isdigit() else x for x in a.투자계정}
    판정 = [p for f in a.판정 for p in 입력.판정파일(f)]
    r = engine.실행(줄, 전기, 마스터, 정책, 기말, 기준=a.기준, 상세=상세, 명세서=명세, 손익NI=손익NI, 현금추가=현금추가,
                   투자지정=투자지정, 잔액열어둠=부가.get("잔액열어둠", False), 판정=판정)
    메모.append(f"기준: {r.기준}" + (f", 상세 {len(상세)}건" if 상세 else "") + (f", 계정명세서 {len(명세)}건" if 명세 else "")
              + (f", 회계사 판정 {len(r.판정)}줄" if r.판정 else ""))

    out = Path(a.출력 or Path(a.분개장[0]).parent)
    out.mkdir(parents=True, exist_ok=True)
    이름 = f"{기간[0]}_{기간[1]}"
    xl = out / f"현금흐름표_{이름}.xlsx"
    메모 = 메모 + [f"정책: " + ", ".join(f"{k}={v}" for k, v in r.정책.items())]
    셀 = build_workbook.만들기(r, xl, a.회사, 기간, 메모, 기말)
    도구 = 재계산.심기(xl)
    틀림, 판정 = (대조(xl, r, 셀) if 도구 else (["재계산 못 함: 엑셀에서 열어 저장한 뒤 다시 대조"], []))

    요약 = {
        "파일": str(xl), "기간": 기간, "메모": 메모,
        "기초현금": r.기초현금, "기말현금": r.기말현금, "현금증감": r.현금증감, "당기순이익": r.당기순이익,
        "직접법": {f"{k[0]}|{k[1]}": v for k, v in sorted(r.직접.items())},
        "활동별": r.직접합, "간접법초안": r.간접초안합,
        "재분류": [{"상대계정": k[1], "간접법": k[2], "직접법": k[3], "항목": k[4], "금액": v} for k, v in r.재분류.items()],
        "판정방법": {m: sum(1 for t in r.직접행 if t["방법"] == m) for m in dict.fromkeys(t["방법"] for t in r.직접행)},
        "엔진검증": [{"검증": k, "값": v} for k, v in r.검증], "엔진통과": r.통과,
        "독립검산": r.독립, "독립통과": r.독립통과,
        "엑셀검증": [{"검증": k, "판정": p} for k, p in 판정], "엑셀대조_틀림": 틀림,
        "확인사항": [{"분류": x[0], "내용": x[1], "금액": x[2]} for x in r.확인],
    }
    (out / f"검증_{이름}.json").write_text(json.dumps(요약, ensure_ascii=False, indent=1), encoding="utf-8")

    # 판정후보: 근거가 약한 행을 계정·거래처로 묶는다. 활동·항목·근거·판정자·판정일을 채워 --판정 으로 다시 준다
    후보 = defaultdict(lambda: [0, "", ""])
    for t in r.직접행:
        if t["방법"] in engine.근거방법 or t["방법"].endswith("+회계사 판정"):
            continue
        k = (t["상대코드"], t["상대계정"], t["거래처"])
        후보[k][0] += t["금액"]; 후보[k][1] = t["활동"]; 후보[k][2] = t["항목"]
    import csv
    with open(out / f"판정후보_{이름}.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["계정코드", "계정과목", "거래처", "현금영향", "지금 활동", "지금 항목", "활동", "항목", "근거", "판정자", "판정일"])
        for (c, n, 처), (v, 활, 항) in sorted(후보.items(), key=lambda kv: -abs(kv[1][0])):
            if v:
                w.writerow([c, n, 처, v, 활, 항, "", "", "", "", ""])

    print(f"{xl.name}")
    print(f"  기초현금 {r.기초현금:,}  기말현금 {r.기말현금:,}  증감 {r.현금증감:,}")
    for 활 in engine.활동순서:
        if 활 == engine.환율 and not r.직접합[활]:
            continue
        print(f"  {활}  {r.직접합[활]:,}  (계정 기본활동으로만 보면 {r.간접초안합[활]:,})")
    print(f"  재분류 {len(r.재분류)}건  확인사항 {len(r.확인)}건")
    print(f"  엔진 검증 {'PASS' if r.통과 else 'FAIL'}  엑셀 대조 {'PASS' if not 틀림 else 'FAIL'}"
          f"  독립 검산 {'PASS' if r.독립통과 else 'FAIL'}"
          f" (A {sum(x['통과'] is False for x in r.독립['A'])} · B {sum(x['통과'] is False for x in r.독립['B'])}"
          f" · C {sum(x['통과'] is False for x in r.독립['C'])} FAIL"
          + (", B 는 잔액 열어둠이라 판정 안 함" if r.잔액열어둠 else "") + ")")
    for t in 틀림:
        print("   ", t)
    return 0 if (r.통과 and not 틀림) else 1


if __name__ == "__main__":
    sys.exit(main())

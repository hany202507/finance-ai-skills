# -*- coding: utf-8 -*-
"""엔진 결과 → 현금흐름표 워크북. 표와 검증은 전부 살아있는 수식이다.

구조
  분개장            원 분개 줄마다 현금흐름 열(현금영향·활동·항목·판정방법·경로)을 붙인다.
                    한 줄의 현금이 여러 항목으로 나뉘면 바로 밑에 「분할」 행을 두고 현금영향은 분할 행에만 둔다.
                    회계사는 수정활동·수정항목 칸을 채운다.
  현금흐름표_직접법  항목별 SUMIFS(분개장 현금영향, 최종활동, 최종항목)
  간접법_정산표      감사인 정산표 방식. 재무상태표 계정 한 행의 증감을 칸으로 모두 설명하고 행마다 검산 0.
                    투자·재무·환율 칸과 비영업 계정의 영업 칸은 분개장 분류를 SUMIFS 로 가져온다.
                    비용가산·수익차감·비현금거래 칸은 엔진이 비현금 분개를 짝지어 나눈 값이다.
  현금흐름표_간접법  정산표에서 당기순이익 → 가산 → 차감 → 자산부채 변동. 투자·재무는 직접법 항목
"""
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
import 계정규칙 as R

NAVY = PatternFill("solid", fgColor="1E2750")
SUB = PatternFill("solid", fgColor="F1F3FB")
INPUT = PatternFill("solid", fgColor="FFF6D5")
OKF = PatternFill("solid", fgColor="E6F4EA")
SPLIT = PatternFill("solid", fgColor="F7F7F7")
thin = Side(style="thin", color="D0D5DD")
HF = Font(bold=True, color="FFFFFF", size=10)
NUM = "#,##0;[Red]-#,##0"
활동순서 = [R.영업, R.투자, R.재무, R.환율]
활동명 = {R.영업: "Ⅰ. 영업활동 현금흐름", R.투자: "Ⅱ. 투자활동 현금흐름", R.재무: "Ⅲ. 재무활동 현금흐름",
         R.환율: "Ⅳ. 현금및현금성자산의 환율변동효과"}


def _머리(ws, heads, widths, 줄=1):
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=줄, column=i, value=h)
        c.fill, c.font = NAVY, HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = f"A{줄 + 1}"


def _숫자(ws, cols, r0, r1):
    for col in cols:
        for r in range(r0, r1 + 1):
            ws[f"{col}{r}"].number_format = NUM


def _줄(ws, r, kind, 열=2):
    a, b = ws.cell(row=r, column=1), ws.cell(row=r, column=열)
    b.number_format = NUM
    if kind == "sec":
        for c in (a, b): c.font = Font(bold=True, size=11); c.fill = SUB
    elif kind == "grp":
        a.font = Font(bold=True, size=10, color="44506A"); a.alignment = Alignment(indent=1)
    elif kind == "item":
        a.alignment = Alignment(indent=2)
    elif kind == "total":
        for c in (a, b): c.font = Font(bold=True, size=11, color="FFFFFF"); c.fill = NAVY
    elif kind == "chk":
        for c in (a, b): c.fill = OKF; c.font = Font(size=10, color="1E6B34")
    for c in (a, b):
        c.border = Border(bottom=thin)


def _식값(코드):
    if isinstance(코드, int):
        return str(코드)
    return '"' + str(코드).replace('"', '""') + '"'


def 만들기(r, 경로, 회사="", 기간=("", ""), 메모=(), 기말BS=None):
    wb = Workbook()
    요약 = wb.active; 요약.title = "요약"
    기준 = wb.create_sheet("분류기준")
    직 = wb.create_sheet("현금흐름표_직접법")
    간 = wb.create_sheet("현금흐름표_간접법")
    정 = wb.create_sheet("간접법_정산표")
    비 = wb.create_sheet("비교")
    분 = wb.create_sheet("분개장")
    비현 = wb.create_sheet("비현금거래")
    확 = wb.create_sheet("확인사항")
    규 = wb.create_sheet("계정규칙")
    전 = wb.create_sheet("전기재무상태표")
    말 = wb.create_sheet("기말재무상태표")
    독 = wb.create_sheet("독립검산", 1)
    목 = wb.create_sheet("목록")
    계정 = r.계정
    셀 = {}

    # ════════ 분개장
    # A전표 B일자 C전표구분 D계정코드 E계정과목 F계정구분 G차변 H대변 I거래처 J적요 K세부계정 L현금전표 M행
    # N현금영향 O현금영향합계(분할) P기본활동 Q자동활동 R자동항목 S수정활동 T수정항목 U최종활동 V최종항목 W재분류
    # X판정방법 Y추적경로 Z원천전표 AA확인 AB적요판정
    _머리(분, ["전표", "일자", "전표구분", "계정코드", "계정과목", "계정구분", "차변", "대변", "거래처", "적요", "세부계정",
               "현금전표", "행", "현금영향\n(유입+ 유출-)", "현금영향 합계\n(분할 전)", "기본활동\n(계정)", "자동활동", "자동항목",
               "수정활동", "수정항목", "최종활동", "최종항목", "재분류", "판정 방법", "추적 경로", "원천 전표", "확인", "적요 판정",
               "분류 대상·이유", "정산칸", "정산라벨", "정산금액\n(현금영향 부호)", "투자·재무 관련\n비현금(상한용)", "관련 활동"],
          [9, 11, 8, 9, 14, 7, 15, 15, 14, 30, 18, 6, 8, 15, 15, 8, 8, 18, 8, 16, 8, 18, 7, 16, 26, 14, 34, 24, 46, 8, 16, 14, 14, 9])
    칸이름 = {"손익": "비용가산·수익차감", "비현금": "비현금거래", "영업잔차": "영업자산부채 증감"}

    def 정산(x):
        """정산표 칸으로 가는 비현금 몫: [(칸, 라벨, 금액)]. 손익·비현금만(영업잔차는 정산표에서 잔차로 계산)."""
        return [(k, lab, v) for (k, lab), v in sorted(r.줄정산.get(id(x), {}).items()) if k in ("손익", "비현금") and v]

    def 비현금이유(x, 현금전표):
        a = 계정[x["코드"]]
        간 = [k for k, v in r.줄정산.get(id(x), {}).items() if v]
        앞 = ("현금 전표지만 현금과 같은 편(함께 차변 또는 함께 대변)이라 이 줄로 현금이 오가지 않았다"
             if 현금전표 else "현금이 오가지 않은 전표")
        if not 간:
            return 앞 + ". 현금흐름표에 영향 없음"
        if ("순이익", "") in 간:
            return 앞 + ". 손익계정이라 간접법 당기순이익에 들어 있다"
        곳 = sorted({f"{칸이름[k]}" + (f"({lab})" if lab else "") for k, lab in 간})
        return 앞 + f". 간접법_정산표 {a['계정']} 행 " + ", ".join(곳)

    def 분류이유(t):
        뜻, 신 = R.방법풀이(t["방법"])
        return f"분류함(신뢰 {신}): {뜻}"
    행별 = defaultdict(list)
    for t in r.직접행:
        행별[t["원id"]].append(t)
    n = 1
    분석행 = []      # 현금영향이 있는 행 번호(수식 넣을 곳)

    def _정산분할(x, a, k, 정몫, n):
        """비현금 몫이 둘 이상의 정산표 칸으로 가면 줄 밑에 정산분할 행을 둔다. 차변·대변은 비워 잔액에 안 섞인다."""
        if len(정몫) < 2:
            return n
        for 칸_, lab, v in 정몫:
            n += 1
            분.append([k, x["일자"], "정산분할", x["코드"], x["계정"], a["구분"], None, None, x["거래처"], x["적요"],
                      x.get("세부", ""), "", "정산분할"] + [None] * 15 + ["비현금 몫이 여러 정산표 칸으로 나뉜다"] + [칸_, lab, v, None, None])
            for c in range(1, 35):
                분.cell(row=n, column=c).fill = SPLIT
        return n
    for k, v in r.전표.items():
        cash = any(계정[x["코드"]]["성격"] == "현금" for x in v)
        for x in v:
            n += 1          # max_row 는 호출마다 셀 전체를 훑어 18만 줄에서 멈춘다. 직접 센다
            a = 계정[x["코드"]]
            base = [k, x["일자"], x["구분"], x["코드"], x["계정"], a["구분"], x["차변"] or None, x["대변"] or None,
                    x["거래처"], x["적요"], x.get("세부", ""), "Y" if cash else ""]
            rows = 행별.get(id(x), [])
            정몫 = 정산(x)
            정열 = list(정몫[0]) if len(정몫) == 1 else (["정산 분할", f"{len(정몫)}개", None] if 정몫 else [None, None, None])
            상 = r.상한대체.get(id(x))
            정열 = 정열 + ([상[1], 상[0]] if 상 else [None, None])
            if a["성격"] == "현금":
                분.append(base + ["현금"] + [None] * 15 + ["현금 계정 줄. 현금흐름 그 자체라 분류하지 않고 같은 전표의 상대 줄을 분류한다"]
                         + [None] * 5)
                continue
            if not rows:
                분.append(base + [""] + [None] * 15 + [비현금이유(x, cash)] + 정열)
                n = _정산분할(x, a, k, 정몫, n)
                continue
            if len(rows) == 1:
                t = rows[0]
                분.append(base + ["", t["금액"], None, t["기본활동"], t["활동"], t["항목"], None, None, None, None, None,
                                 t["방법"], t["경로"], t["원천"], t["확인"], t.get("적요판정", ""), 분류이유(t)] + 정열)
                분석행.append(n)
                n = _정산분할(x, a, k, 정몫, n)
                continue
            분.append(base + ["분할 합계", None, sum(t["금액"] for t in rows), a["활동"], "분할", f"{len(rows)}개로 나뉨",
                             None, None, None, None, None, "", "", "", "현금영향은 아래 분할 행에 있다. 수정도 분할 행에서", "",
                             "현금이 여러 항목으로 나뉘어 바로 아래 분할 행에서 분류한다"] + 정열)
            for t in rows:
                n += 1
                분.append([k, t["일자"], "분할", x["코드"], x["계정"], a["구분"], None, None, t["거래처"], t["적요"],
                          x.get("세부", ""), "Y", "분할", t["금액"], None, t["기본활동"], t["활동"], t["항목"], None, None,
                          None, None, None, t["방법"], t["경로"], t["원천"], t["확인"], t.get("적요판정", ""), 분류이유(t)])
                for c in range(1, 30):
                    분.cell(row=n, column=c).fill = SPLIT
                분석행.append(n)
            n = _정산분할(x, a, k, 정몫, n)
    NJ = max(n, 2)
    assert 분.max_row == n, (분.max_row, n)    # 수식 행 번호가 실제 행과 같은지
    for i in 분석행:
        분[f"U{i}"] = f'=IF(S{i}="",Q{i},S{i})'
        분[f"V{i}"] = f'=IF(T{i}="",R{i},T{i})'
        분[f"W{i}"] = f'=IF(U{i}<>P{i},"재분류","")'
        분[f"S{i}"].fill = INPUT; 분[f"T{i}"].fill = INPUT
    _숫자(분, ["G", "H", "N", "O", "AF", "AG"], 2, NJ)
    분.auto_filter.ref = f"A1:AH{NJ}"
    J = {c: f"분개장!${c}$2:${c}${NJ}" for c in ("D", "F", "G", "H", "L", "M", "N", "P", "U", "V", "X", "AD", "AE", "AF", "AG", "AH")}
    항목전체 = list(dict.fromkeys(h for 활 in 활동순서 for h in R.항목순서[활]))
    for i, h in enumerate(항목전체, 1):
        목.cell(row=i, column=1, value=h)
    목.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1='"영업,투자,재무,환율"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1=f"=목록!$A$1:$A${len(항목전체)}", allow_blank=True,
                         showErrorMessage=True, errorTitle="항목", error="목록에 있는 항목만 쓸 수 있다")
    분.add_data_validation(dv); 분.add_data_validation(dv2)
    dv.add(f"S2:S{NJ}"); dv2.add(f"T2:T{NJ}")

    # ════════ 전기재무상태표
    _머리(전, ["계정코드", "계정과목", "구분", "잔액"], [12, 20, 8, 18])
    for b in r.전기BS:
        전.append([b["코드"], b["계정"], b["구분"], b["잔액"]])
    _숫자(전, "D", 2, max(전.max_row, 2))
    _머리(말, ["계정코드", "계정과목", "구분", "잔액"], [12, 20, 8, 18])
    for b in (기말BS or []):
        말.append([b["코드"], b["계정"], b.get("구분"), b["잔액"]])
    _숫자(말, "D", 2, max(말.max_row, 2))

    # ════════ 계정규칙
    _머리(규, ["계정코드", "계정과목", "구분", "성격", "기본활동", "유입항목", "유출항목", "판정 근거"],
          [12, 18, 8, 8, 9, 22, 22, 24])
    쓰인 = {x["코드"] for v in r.전표.values() for x in v} | {c for c, v in r.기초.items() if v}
    for c in sorted(쓰인, key=str):
        a = 계정[c]
        규.append([c, a["계정"], a["구분"], a["성격"], a["활동"], a["유입"], a["유출"], a["근거"]])

    # ════════ 현금흐름표_직접법
    직.column_dimensions["A"].width = 42; 직.column_dimensions["B"].width = 22
    직.sheet_view.showGridLines = False
    직["A1"] = f"{회사} 현금흐름표 (직접법)"; 직["A1"].font = Font(bold=True, size=13)
    직["A2"] = f"{기간[0]} ~ {기간[1]}  (단위: 원)  분개장 현금영향을 최종활동·최종항목으로 SUMIFS"; 직["A2"].font = Font(size=9, color="6B7688")
    항목들 = {활: [] for 활 in 활동순서}
    for (활, 항) in r.직접:
        if 활 in 항목들 and 항 not in 항목들[활]:
            항목들[활].append(항)
    항목들[R.환율] = [R.환율항목]
    for 활 in 활동순서:
        순 = R.항목순서[활]
        항목들[활].sort(key=lambda h: 순.index(h) if h in 순 else 99)
    직셀 = {}
    row = 4
    for 활 in 활동순서:
        s = row
        직.cell(row=row, column=1, value=활동명[활]); _줄(직, row, "sec"); row += 1
        for 항 in 항목들[활]:
            직.cell(row=row, column=1, value=항)
            직.cell(row=row, column=2, value=f'=SUMIFS({J["N"]},{J["U"]},"{활}",{J["V"]},A{row})')
            직셀[(활, 항)] = f"B{row}"; _줄(직, row, "item"); row += 1
        직.cell(row=row, column=1, value="그 밖의 항목(수정 입력으로 생긴 것)")
        직.cell(row=row, column=2, value=f'=SUMIFS({J["N"]},{J["U"]},"{활}")-SUM(B{s + 1}:B{row - 1})' if row > s + 1
                else f'=SUMIFS({J["N"]},{J["U"]},"{활}")')
        _줄(직, row, "item"); 직셀[(활, "그 밖의")] = f"B{row}"; row += 1
        직.cell(row=s, column=2, value=f"=SUM(B{s + 1}:B{row - 1})")
        직셀[활] = f"B{s}"; row += 1

    # ════════ 간접법_정산표
    #  A코드 B과목 C기말 D기초 E증감 F검산 G당기순이익 H/I비용가산 J/K수익차감 L/M영업자산부채 N/O투자유입 P/Q투자유출
    #  R/S재무유입 T/U재무유출 V/W비현금차변 X/Y비현금대변 Z/AA환율 AB주어진기말 AC차이
    heads = ["계정코드", "과목", "기말", "기초", "증감", "검산\n(0 이어야)", "당기순이익", "비용가산", "금액", "수익차감", "금액",
             "영업자산부채의 증감", "금액", "투자 유입", "금액", "투자 유출", "금액", "재무 유입", "금액", "재무 유출", "금액",
             "비현금 차변", "금액", "비현금 대변", "금액", "환율", "금액", "주어진 기말", "기말 차이", "손익 라벨\n(수식 기준)", "비현금 라벨\n(수식 기준)"]
    _머리(정, heads, [11, 22, 16, 16, 16, 12, 15, 16, 14, 16, 14, 22, 15, 16, 14, 16, 14, 16, 14, 16, 14, 18, 14, 18, 14, 10, 12, 16, 12, 16, 16], 줄=3)
    정["A1"] = f"{회사} 현금흐름 정산표  {기간[0]} ~ {기간[1]}"; 정["A1"].font = Font(bold=True, size=13)
    정["A2"] = ("행마다 증감을 칸으로 모두 설명한다(검산 0). 자산: 증감 + (비용가산 - 수익차감 + 영업 + 투자유입 - 투자유출 + 재무유입 - 재무유출"
               " - 비현금차변 + 비현금대변 + 환율) = 0, 부채·자본은 증감에서 뺀다. 투자·재무·환율 칸은 분개장 분류에서 온다")
    정["A2"].font = Font(size=9, color="6B7688")
    주어진 = defaultdict(int)
    for b in (기말BS or []):
        주어진[b["코드"]] += b["잔액"]
    구분순 = ["자산", "부채", "자본"]
    BS = [c for c in 쓰인 | set(주어진) if c in 계정 and 계정[c]["구분"] in 구분순]
    BS.sort(key=lambda c: (구분순.index(계정[c]["구분"]), 계정[c]["성격"] != "현금", str(c)))
    라벨표 = defaultdict(dict)    # 계정 → (활동, 부호) → 엔진이 붙인 항목 이름(여러 개면 「외」)
    모음 = defaultdict(lambda: defaultdict(list))
    for t in r.직접행:
        if t["활동"] in (R.투자, R.재무):
            모음[t["상대코드"]][(t["활동"], ">0" if t["금액"] > 0 else "<0")].append(t["항목"])
    for c, d in 모음.items():
        for k, v in d.items():
            u = list(dict.fromkeys(v))
            라벨표[c][k] = u[0] + (f" 외 {len(u) - 1}" if len(u) > 1 else "")
    row = 4
    잉여행 = []
    정행 = {}          # 코드 → (첫 행, 마지막 행)
    현금행 = []
    for 구분 in 구분순:
        정.cell(row=row, column=2, value=f"[{구분}]").font = Font(bold=True); row += 1
        for c in [c for c in BS if 계정[c]["구분"] == 구분]:
            a = 계정[c]
            칸 = r.정산칸.get(c, {})
            손익라벨 = sorted(set(칸.get("I", {})) | set(칸.get("K", {})))
            비현라벨 = sorted(set(칸.get("W", {})) | set(칸.get("Y", {})))
            높이 = max(1, len(손익라벨), len(비현라벨))
            s, e = row, row + 높이 - 1
            정행[c] = (s, e)
            정.cell(row=s, column=1, value=c); 정.cell(row=s, column=2, value=a["계정"])
            자산 = 구분 == "자산"
            차대 = (f"SUMIFS({J['G']},{J['D']},$A{s})-SUMIFS({J['H']},{J['D']},$A{s})" if 자산
                   else f"SUMIFS({J['H']},{J['D']},$A{s})-SUMIFS({J['G']},{J['D']},$A{s})")
            정[f"D{s}"] = f"=SUMIFS(전기재무상태표!$D:$D,전기재무상태표!$A:$A,$A{s})"
            정[f"C{s}"] = f"=D{s}+{차대}"
            정[f"E{s}"] = f"=C{s}-D{s}"
            if 기말BS and a["성격"] != "잉여금":
                정[f"AB{s}"] = f"=SUMIFS(기말재무상태표!$D:$D,기말재무상태표!$A:$A,$A{s})"
                정[f"AC{s}"] = f"=C{s}-AB{s}"
            if a["성격"] == "잉여금":
                잉여행.append(s)
            if a["성격"] == "현금":
                정[f"F{s}"] = 0
                정.cell(row=s, column=2, value=f"{a['계정']} (현금)")
                현금행.append(s)
                row = e + 1
                continue
            # 비용가산·수익차감·비현금 칸도 분개장 정산 열에서 SUMIFS 로 가져온다(라벨별 순액, 부호로 칸을 가른다)
            for i, lab in enumerate(손익라벨):
                q = s + i
                정[f"AD{q}"] = lab
                순 = f'SUMIFS({J["AF"]},{J["D"]},$A${s},{J["AD"]},"손익",{J["AE"]},$AD{q})'
                정[f"I{q}"] = f"=MAX(0,{순})"; 정[f"K{q}"] = f"=MAX(0,-{순})"
                정[f"H{q}"] = f'=IF(I{q}>0,AD{q},"")'; 정[f"J{q}"] = f'=IF(K{q}>0,AD{q},"")'
            for i, lab in enumerate(비현라벨):
                q = s + i
                정[f"AE{q}"] = lab
                순 = f'SUMIFS({J["AF"]},{J["D"]},$A${s},{J["AD"]},"비현금",{J["AE"]},$AE{q})'
                정[f"Y{q}"] = f"=MAX(0,{순})"; 정[f"W{q}"] = f"=MAX(0,-{순})"
                정[f"X{q}"] = f'=IF(Y{q}>0,AE{q},"")'; 정[f"V{q}"] = f'=IF(W{q}>0,AE{q},"")'
            ID = f"$A{s}"
            이름표 = 라벨표.get(c, {})
            for 라벨열, 금액열, 활, 조건, 부호, 기본 in (("N", "O", "투자", ">0", "", "투자 유입"), ("P", "Q", "투자", "<0", "-", "투자 유출"),
                                                   ("R", "S", "재무", ">0", "", "재무 유입"), ("T", "U", "재무", "<0", "-", "재무 유출")):
                정[f"{금액열}{s}"] = f'={부호}SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"{활}",{J["N"]},"{조건}")'
                정[f"{라벨열}{s}"] = f'=IF({금액열}{s}<>0,"{이름표.get((활, 조건), 기본)}","")'
            정[f"AA{s}"] = f'=SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"환율")'
            정[f"Z{s}"] = f'=IF(AA{s}<>0,"환율","")'
            합 = (f"SUM(I{s}:I{e})-SUM(K{s}:K{e})+SUM(M{s}:M{e})+SUM(O{s}:O{e})-SUM(Q{s}:Q{e})+SUM(S{s}:S{e})"
                 f"-SUM(U{s}:U{e})-SUM(W{s}:W{e})+SUM(Y{s}:Y{e})+SUM(AA{s}:AA{e})")
            if a["활동"] == R.영업:
                정[f"L{s}"] = f"{a['계정']}의 {'감소(증가)' if 자산 else '증가(감소)'}"
                나머지 = (f"SUM(I{s}:I{e})-SUM(K{s}:K{e})+SUM(O{s}:O{e})-SUM(Q{s}:Q{e})+SUM(S{s}:S{e})-SUM(U{s}:U{e})"
                        f"-SUM(W{s}:W{e})+SUM(Y{s}:Y{e})+SUM(AA{s}:AA{e})")
                정[f"M{s}"] = f"=-E{s}-({나머지})" if 자산 else f"=E{s}-({나머지})"
            else:
                정[f"M{s}"] = f'=SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"영업")'
                정[f"L{s}"] = f'=IF(M{s}<>0,"영업 분류분","")'
            정[f"F{s}"] = f"=E{s}+({합})" if 자산 else f"=E{s}-({합})"
            row = e + 1
    # 당기순이익 행
    BS끝 = row - 1
    당기순 = row
    정.cell(row=row, column=2, value="당기순이익(이익잉여금으로)").font = Font(bold=True)
    NI식 = (f'SUMIFS({J["H"]},{J["F"]},"수익")-SUMIFS({J["G"]},{J["F"]},"수익")'
           f'+SUMIFS({J["H"]},{J["F"]},"비용")-SUMIFS({J["G"]},{J["F"]},"비용")')
    정[f"C{row}"] = f"={NI식}"; 정[f"D{row}"] = 0; 정[f"E{row}"] = f"=C{row}-D{row}"
    정[f"G{row}"] = f"=C{row}"; 정[f"F{row}"] = f"=E{row}-G{row}"
    if 기말BS and 잉여행:
        # 분개장 이익잉여금에는 마감분개를 뺐으므로 당기순이익이 없다. 이익잉여금 합계 + 당기순이익으로 대조한다
        정[f"AB{row}"] = "=" + "+".join(f"SUMIFS(기말재무상태표!$D:$D,기말재무상태표!$A:$A,{_식값(c)})"
                                        for c in 계정 if 계정[c]["성격"] == "잉여금")
        정[f"AC{row}"] = "=" + "+".join(f"C{i}" for i in 잉여행) + f"+C{row}-AB{row}"
        정.cell(row=row, column=28).number_format = NUM
    row += 2
    # 손익계정 중 현금분이 영업 외 활동으로 간 것
    정.cell(row=row, column=2, value="[손익계정 중 투자·재무·환율로 간 현금분]").font = Font(bold=True); row += 1
    손익시작 = row
    for c in r.정산손익:
        a = 계정[c]
        정.cell(row=row, column=1, value=c); 정.cell(row=row, column=2, value=a["계정"])
        ID = f"$A{row}"
        정[f"H{row}"] = a["계정"]; 정[f"I{row}"] = f'=-SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"<>영업",{J["N"]},"<0")'
        정[f"J{row}"] = a["계정"]; 정[f"K{row}"] = f'=SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"<>영업",{J["N"]},">0")'
        정[f"O{row}"] = f'=SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"투자",{J["N"]},">0")'
        정[f"Q{row}"] = f'=-SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"투자",{J["N"]},"<0")'
        정[f"S{row}"] = f'=SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"재무",{J["N"]},">0")'
        정[f"U{row}"] = f'=-SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"재무",{J["N"]},"<0")'
        정[f"AA{row}"] = f'=SUMIFS({J["N"]},{J["D"]},{ID},{J["U"]},"환율")'
        정[f"F{row}"] = f"=I{row}-K{row}+O{row}-Q{row}+S{row}-U{row}+AA{row}"
        row += 1
    if not r.정산손익:
        정.cell(row=row, column=2, value="없음"); row += 1
    끝 = row - 1
    row += 1
    합행 = row
    정.cell(row=row, column=2, value="합계").font = Font(bold=True)
    for col in ("G", "I", "K", "M", "O", "Q", "S", "U", "W", "Y", "AA"):
        정[f"{col}{row}"] = f"=SUM({col}4:{col}{끝})"
        정[f"{col}{row}"].font = Font(bold=True)
    _숫자(정, ["C", "D", "E", "F", "G", "I", "K", "M", "O", "Q", "S", "U", "W", "Y", "AA", "AB", "AC"], 4, row)
    row += 2
    # 요약 칸
    정.cell(row=row, column=2, value="정산표 요약").font = Font(bold=True, size=11); row += 1
    Σ = lambda col: f"{col}{합행}"
    현금E = "+".join(f"E{i}" for i in 현금행) or "0"
    요약칸 = [
        ("영업활동 = 당기순이익 + 비용가산 - 수익차감 + 영업자산부채 증감", f"={Σ('G')}+{Σ('I')}-{Σ('K')}+{Σ('M')}", R.영업),
        ("투자활동 = 유입 - 유출", f"={Σ('O')}-{Σ('Q')}", R.투자),
        ("재무활동 = 유입 - 유출", f"={Σ('S')}-{Σ('U')}", R.재무),
        ("환율변동효과", f"={Σ('AA')}", R.환율),
    ]
    정요약 = {}
    for 라벨, 식, 활 in 요약칸:
        정.cell(row=row, column=2, value=라벨); 정[f"E{row}"] = 식; 정[f"E{row}"].number_format = NUM
        정[f"F{row}"] = f"=E{row}-현금흐름표_직접법!{직셀[활]}"; 정[f"F{row}"].number_format = NUM
        정[f"G{row}"] = "← 직접법과 차이"
        정요약[활] = f"E{row}"; row += 1
    정.cell(row=row, column=2, value="합계"); 정[f"E{row}"] = f"=SUM(E{row - 4}:E{row - 1})"; 정[f"E{row}"].number_format = NUM
    정요약["합계"] = f"E{row}"; row += 1
    정.cell(row=row, column=2, value="현금 증감(현금 행)"); 정[f"E{row}"] = f"={현금E}"; 정[f"E{row}"].number_format = NUM
    정요약["현금"] = f"E{row}"; row += 1
    정.cell(row=row, column=2, value="비현금 차변 - 대변(0 이어야)"); 정[f"E{row}"] = f"={Σ('W')}-{Σ('Y')}"; 정[f"E{row}"].number_format = NUM
    정요약["비현금"] = f"E{row}"; row += 1
    정.cell(row=row, column=2, value="행 검산 절댓값 합(0 이어야)"); 정[f"E{row}"] = f"=SUMPRODUCT(ABS(F4:F{끝}))"; 정[f"E{row}"].number_format = NUM
    정요약["행검산"] = f"E{row}"; row += 1
    정.cell(row=row, column=2, value="비용가산 - 수익차감 = 분개장 정산 손익 합(차이)")
    정[f"E{row}"] = f'=SUM(I4:I{BS끝})-SUM(K4:K{BS끝})-SUMIFS({J["AF"]},{J["AD"]},"손익")'; 정[f"E{row}"].number_format = NUM
    정요약["손익칸"] = f"E{row}"; row += 1
    정.cell(row=row, column=2, value="비현금 대변 - 차변 = 분개장 정산 비현금 합(차이)")
    정[f"E{row}"] = f'=SUM(Y4:Y{BS끝})-SUM(W4:W{BS끝})-SUMIFS({J["AF"]},{J["AD"]},"비현금")'; 정[f"E{row}"].number_format = NUM
    정요약["비현금칸"] = f"E{row}"; row += 2
    # 손익계산서 대사: 비용가산·수익차감 라벨별
    정.cell(row=row, column=2, value="손익계산서 대사(비용가산·수익차감 라벨별)").font = Font(bold=True, size=11); row += 1
    for j, h in enumerate(["계정코드", "손익계정", "손익계산서(분개장)", "정산표 가감", "차이", "비고"], 1):
        c = 정.cell(row=row, column=j + 1, value=h); c.fill, c.font = NAVY, HF
    row += 1
    라벨들 = sorted({lab for c in r.정산칸 for k in ("I", "K") for lab in r.정산칸[c].get(k, {})} |
                  {계정[c]["계정"] for c in r.정산손익})
    for lab in 라벨들:
        코드 = r.정산라벨코드.get(lab) or next((c for c in r.정산손익 if 계정[c]["계정"] == lab), None)
        if 코드 is None:
            continue
        비용 = 계정[코드]["구분"] == "비용"
        정.cell(row=row, column=2, value=코드); 정.cell(row=row, column=3, value=lab)
        ID = f"$B{row}"
        정[f"D{row}"] = (f"=SUMIFS({J['G']},{J['D']},{ID})-SUMIFS({J['H']},{J['D']},{ID})" if 비용
                        else f"=SUMIFS({J['H']},{J['D']},{ID})-SUMIFS({J['G']},{J['D']},{ID})")
        정[f"E{row}"] = f"=SUMIF($H$4:$H${끝},C{row},$I$4:$I${끝})-SUMIF($J$4:$J${끝},C{row},$K$4:$K${끝})" if 비용 \
            else f"=SUMIF($J$4:$J${끝},C{row},$K$4:$K${끝})-SUMIF($H$4:$H${끝},C{row},$I$4:$I${끝})"
        정[f"F{row}"] = f"=D{row}-E{row}"
        정[f"G{row}"] = "차이는 현금으로 주고받았거나 영업자산부채 증감에 남은 몫"
        for col in "DEF":
            정[f"{col}{row}"].number_format = NUM
        row += 1
    셀["정산"] = 정요약

    # ════════ 현금흐름표_간접법 (정산표에서)
    간.column_dimensions["A"].width = 46; 간.column_dimensions["B"].width = 22
    간.sheet_view.showGridLines = False
    간["A1"] = f"{회사} 현금흐름표 (간접법)"; 간["A1"].font = Font(bold=True, size=13)
    간["A2"] = f"{기간[0]} ~ {기간[1]}  (단위: 원)  영업은 간접법_정산표, 투자·재무·환율은 직접법 항목"; 간["A2"].font = Font(size=9, color="6B7688")
    row = 4
    간셀 = {}
    s영 = row
    간.cell(row=row, column=1, value=활동명[R.영업]); _줄(간, row, "sec"); row += 1
    간.cell(row=row, column=1, value="1. 당기순이익(손실)"); 간.cell(row=row, column=2, value=f"=간접법_정산표!G{당기순}")
    _줄(간, row, "grp"); 소계 = [row]; row += 1
    범 = f"간접법_정산표!$H$4:$H${끝}", f"간접법_정산표!$I$4:$I${끝}", f"간접법_정산표!$J$4:$J${끝}", f"간접법_정산표!$K$4:$K${끝}"
    가산 = sorted({lab for c in r.정산칸 for lab in r.정산칸[c].get("I", {})} | {계정[c]["계정"] for c in r.정산손익})
    차감 = sorted({lab for c in r.정산칸 for lab in r.정산칸[c].get("K", {})} | {계정[c]["계정"] for c in r.정산손익})
    for 제목, 라벨목록, 라벨범, 금액범, 부호, 총칸 in (
            ("2. 현금의 유출이 없는 비용 등의 가산", 가산, 범[0], 범[1], "", "I"),
            ("3. 현금의 유입이 없는 수익 등의 차감", 차감, 범[2], 범[3], "-", "K")):
        h = row
        간.cell(row=row, column=1, value=제목); _줄(간, row, "grp"); row += 1
        for lab in 라벨목록:
            간.cell(row=row, column=1, value=lab)
            간.cell(row=row, column=2, value=f"={부호}SUMIF({라벨범},A{row},{금액범})"); _줄(간, row, "item"); row += 1
        간.cell(row=row, column=1, value="그 밖의 것")
        간.cell(row=row, column=2, value=f"={부호}간접법_정산표!{총칸}{합행}-SUM(B{h + 1}:B{row - 1})" if row > h + 1
                else f"={부호}간접법_정산표!{총칸}{합행}")
        _줄(간, row, "item"); row += 1
        간.cell(row=h, column=2, value=f"=SUM(B{h + 1}:B{row - 1})"); 소계.append(h)
    h = row
    간.cell(row=row, column=1, value="4. 영업활동으로 인한 자산·부채의 변동"); _줄(간, row, "grp"); row += 1
    for c, (s, e) in 정행.items():
        if 계정[c]["성격"] == "현금" or 계정[c]["활동"] != R.영업:
            continue
        간.cell(row=row, column=1, value=f"=간접법_정산표!L{s}")
        간.cell(row=row, column=2, value=f"=간접법_정산표!M{s}"); _줄(간, row, "item"); row += 1
    간.cell(row=row, column=1, value="그 밖의 것(비영업 계정 중 영업으로 분류된 현금)")
    간.cell(row=row, column=2, value=f"=간접법_정산표!M{합행}-SUM(B{h + 1}:B{row - 1})" if row > h + 1 else f"=간접법_정산표!M{합행}")
    _줄(간, row, "item"); row += 1
    간.cell(row=h, column=2, value=f"=SUM(B{h + 1}:B{row - 1})"); 소계.append(h)
    간.cell(row=s영, column=2, value="=" + "+".join(f"B{x}" for x in 소계))
    간셀[R.영업] = f"B{s영}"
    row += 1
    for 활 in (R.투자, R.재무, R.환율):
        s = row
        간.cell(row=row, column=1, value=활동명[활]); _줄(간, row, "sec"); row += 1
        for 항 in 항목들[활] + ["그 밖의"]:
            간.cell(row=row, column=1, value="그 밖의 항목(수정 입력으로 생긴 것)" if 항 == "그 밖의" else 항)
            간.cell(row=row, column=2, value=f"=현금흐름표_직접법!{직셀[(활, 항)]}"); _줄(간, row, "item"); row += 1
        간.cell(row=s, column=2, value=f"=SUM(B{s + 1}:B{row - 1})"); 간셀[활] = f"B{s}"; row += 1
    for 표, 표셀 in ((직, 직셀), (간, 간셀)):
        rr = 표.max_row + 1
        표.cell(row=rr, column=1, value="Ⅴ. 현금의 증감 (Ⅰ+Ⅱ+Ⅲ+Ⅳ)")
        표.cell(row=rr, column=2, value="=" + "+".join(표셀[a] for a in 활동순서)); _줄(표, rr, "sec"); 표셀["증감"] = f"B{rr}"
        표.cell(row=rr + 1, column=1, value="Ⅵ. 기초의 현금" + (" (잔액 열어둠: 0 으로 두었다)" if r.잔액열어둠 else ""))
        표.cell(row=rr + 1, column=2, value="=" + ("+".join(f"간접법_정산표!D{i}" for i in 현금행) or "0")); _줄(표, rr + 1, "sec")
        표셀["기초"] = f"B{rr + 1}"
        표.cell(row=rr + 2, column=1, value="Ⅶ. 기말의 현금 (Ⅴ+Ⅵ)" + (" = 당기 변동만" if r.잔액열어둠 else ""))
        표.cell(row=rr + 2, column=2, value=f"=B{rr}+B{rr + 1}"); _줄(표, rr + 2, "total"); 표셀["기말"] = f"B{rr + 2}"
        표.cell(row=rr + 4, column=1, value="검증: 재무상태표 기말 현금")
        표.cell(row=rr + 4, column=2, value="=" + ("+".join(f"간접법_정산표!C{i}" for i in 현금행) or "0")); _줄(표, rr + 4, "chk")
        표.cell(row=rr + 5, column=1, value="검증: 차이 (0 이어야 한다)")
        표.cell(row=rr + 5, column=2, value=f"=B{rr + 2}-B{rr + 4}"); _줄(표, rr + 5, "chk"); 표셀["차이"] = f"B{rr + 5}"
    rr = 간.max_row + 1
    간.cell(row=rr, column=1, value="검증: 영업활동이 직접법과 같은가 (차이)")
    간.cell(row=rr, column=2, value=f"={간셀[R.영업]}-현금흐름표_직접법!{직셀[R.영업]}"); _줄(간, rr, "chk"); 간셀["영업차이"] = f"B{rr}"
    DJ = lambda k: f"현금흐름표_직접법!{직셀[k]}"
    GJ = lambda k: f"현금흐름표_간접법!{간셀[k]}"

    # ════════ 비교
    _머리(비, ["활동", "직접법\n(분개장 분류)", "간접법\n(정산표)", "차이"], [12, 20, 20, 16])
    for i, 활 in enumerate(활동순서, 2):
        비.append([활, f"={DJ(활)}", f"=간접법_정산표!{정요약[활]}", f"=B{i}-C{i}"])
    비.append(["합계", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)"])
    _숫자(비, "BCD", 2, 6)
    비["A8"] = "재분류 내역: 계정의 기본활동과 분개장 최종활동이 다른 현금영향(미지급금 결제분이 투자활동 등)"
    비["A8"].font = Font(bold=True, size=10)
    hr = 10
    for j, h in enumerate(["계정코드", "계정", "기본활동", "최종활동", "최종항목", "금액"], 1):
        c = 비.cell(row=hr, column=j, value=h); c.fill, c.font = NAVY, HF
    rr = hr + 1
    for (코드, 계, 기본, 활, 항), v in sorted(r.재분류.items(), key=lambda kv: str(kv[0])):
        for j, val in enumerate([코드, 계, 기본, 활, 항], 1):
            비.cell(row=rr, column=j, value=val)
        비.cell(row=rr, column=6, value=f'=SUMIFS({J["N"]},{J["D"]},A{rr},{J["P"]},C{rr},{J["U"]},D{rr},{J["V"]},E{rr})').number_format = NUM
        rr += 1
    if rr == hr + 1:
        비.cell(row=rr, column=1, value="재분류 없음")

    # ════════ 비현금거래
    _머리(비현, ["전표", "일자", "종류", "계정코드", "계정과목", "기본활동", "차변", "대변", "간접법 제외액", "적요"],
          [12, 11, 20, 11, 16, 9, 16, 16, 14, 34])
    for x in r.비현금:
        for y in x["줄"]:
            비현.append([x["전표"], y["일자"], x["종류"], y["코드"], y["계정"], 계정[y["코드"]]["활동"],
                        y["차변"] or None, y["대변"] or None, x["제외"].get(id(y)), y["적요"]])
    _숫자(비현, "GHI", 2, max(비현.max_row, 2))
    if not r.비현금:
        비현.append(["없음"])

    # ════════ 확인사항
    _머리(확, ["번호", "분류", "내용", "금액", "회계사 판단"], [6, 10, 96, 18, 30])
    for i, (분류, 내용, 금액) in enumerate(r.확인, 1):
        확.append([i, 분류, 내용, 금액, None])
        확[f"E{i + 1}"].fill = INPUT
    _숫자(확, "D", 2, max(확.max_row, 2))

    # ════════ 분류기준
    기준.column_dimensions["A"].width = 30; 기준.column_dimensions["B"].width = 92
    for c, w in zip("CDE", (10, 12, 20)):
        기준.column_dimensions[c].width = w
    기준.sheet_view.showGridLines = False
    기준["A1"] = "현금흐름 분류 기준"; 기준["A1"].font = Font(bold=True, size=13)
    기준["A2"] = "분개장 「분류 대상·이유」 열이 줄마다 이 기준의 어디에 해당하는지 적는다"; 기준["A2"].font = Font(size=9, color="6B7688")
    row = 4
    def 제목(t):
        nonlocal row
        기준.cell(row=row, column=1, value=t).font = Font(bold=True, size=11); row += 1
    def 표머리(hs):
        nonlocal row
        for j, h in enumerate(hs, 1):
            c = 기준.cell(row=row, column=j, value=h); c.fill, c.font = NAVY, HF
        row += 1
    제목("1. 활동의 정의")
    표머리(["활동", "정의와 근거"])
    for a, b in (("영업", "주된 수익창출 활동, 그리고 투자·재무가 아닌 모든 활동(K-IFRS 1007 문단 14). 모르는 것의 기본값이 영업이 되는 이유"),
                 ("투자", "장기성 자산과 현금성자산이 아닌 투자자산의 취득·처분(문단 16). 유형·무형자산, 금융상품, 대여금, 보증금"),
                 ("재무", "자본과 차입금의 크기·구성을 바꾸는 활동(문단 17). 차입·상환, 증자, 자기주식, 배당 지급"),
                 ("환율", "외화 현금및현금성자산의 환율변동효과. 세 활동과 따로 표시(문단 28)")):
        기준.append([a, b]); row += 1
    row += 1
    제목("2. 분개장 줄이 분류 대상인지")
    표머리(["줄의 상태", "왜 활동이 없거나 있는가", "", "줄 수", "현금영향"])
    상태 = [
        ("분류한 줄(영업)", "현금 전표의 상대 줄 중 현금이 실제로 오간 몫", f'=COUNTIFS({J["N"]},"<>",{J["U"]},"영업")', f'=SUMIFS({J["N"]},{J["U"]},"영업")'),
        ("분류한 줄(투자)", "", f'=COUNTIFS({J["N"]},"<>",{J["U"]},"투자")', f'=SUMIFS({J["N"]},{J["U"]},"투자")'),
        ("분류한 줄(재무)", "", f'=COUNTIFS({J["N"]},"<>",{J["U"]},"재무")', f'=SUMIFS({J["N"]},{J["U"]},"재무")'),
        ("분류한 줄(환율)", "", f'=COUNTIFS({J["N"]},"<>",{J["U"]},"환율")', f'=SUMIFS({J["N"]},{J["U"]},"환율")'),
        ("현금 계정 줄", "현금흐름 그 자체. 분류하지 않고 같은 전표의 상대 줄을 분류한다", f'=COUNTIF({J["M"]},"현금")', None),
        ("분할 합계 줄", "현금이 여러 항목으로 나뉘어 바로 아래 분할 줄에서 분류한다", f'=COUNTIF({J["M"]},"분할 합계")', None),
        ("현금이 없는 전표의 줄", "현금이 오가지 않았다. 직접법에는 없고 간접법 정산표(비용가산·수익차감·비현금거래·영업자산부채 증감)나 당기순이익으로 설명된다",
         f'=COUNTIFS({J["L"]},"",{J["M"]},"")', None),
        ("현금 전표지만 현금과 같은 편인 줄", "현금과 함께 차변(또는 함께 대변)이라 이 줄로는 현금이 오가지 않았다. 예: 이자 받을 때 떼인 원천세(선납세금), "
         "현금 300 + 미지급금 700 으로 산 비품의 미지급금 700", f'=COUNTIFS({J["L"]},"Y",{J["M"]},"",{J["N"]},"")', None),
    ]
    for a, b, 식1, 식2 in 상태:
        기준.cell(row=row, column=1, value=a); 기준.cell(row=row, column=2, value=b)
        기준.cell(row=row, column=4, value=식1).number_format = NUM
        if 식2:
            기준.cell(row=row, column=5, value=식2).number_format = NUM
        기준.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    if r.제외전표:
        기준.cell(row=row, column=1, value="마감분개(분개장 시트에 없음)")
        기준.cell(row=row, column=2, value=f"손익 → 이익잉여금 대체 {len(r.제외전표)}건. 현금흐름과 무관하고 당기순이익을 0 으로 만들어 뺐다"); row += 1
    row += 1
    제목("3. 분류 순서(현금이 오간 줄)")
    for t in ("① 현금 전표를 현금분과 비현금분으로 나눈다. 현금과 반대편 줄에만 현금액을 금액 비율로 붙인다",
              "② 상대계정이 명확한 계정(매출채권·매입채무·재고·유형자산·차입금·자본·손익)이면 계정규칙 시트의 활동·항목",
              "③ 부가세·예수금·외환차손익은 같은 전표의 주 계정을 따른다",
              "④ 추적 계정(미지급금·미수금·가지급금·가수금·선급금·선급비용·미지급비용·미지급세금·규칙 없는 계정)은 원천 전표를 찾는다: "
              "건별 일치 → 기초잔액 → 거래처 순차 → 다른 거래처 기초잔액 → 계정 순차 → 원천 없음. 원천이 또 추적 계정이면 3단계까지",
              "⑤ 원천으로 못 정한 것: 현금 반환 상계 → 상세 내역 발라내기 → 적요 → 거래처 이력 → 기본값(기타 영업활동)",
              "⑥ 잡이익·잡손실은 강한 적요가 가리키는 활동을 따른다",
              "⑦ 계정으로 정한 줄도 강한 적요가 다른 활동을 가리키면 확인 열에 적는다(값은 바꾸지 않는다)",
              "⑧ 회계사가 수정활동·수정항목을 채우면 그것이 최종이다"):
        기준.cell(row=row, column=1, value=t); row += 1
    row += 1
    제목("4. 판정 방법별 뜻과 신뢰도")
    표머리(["판정 방법", "뜻", "신뢰도", "줄 수", "현금영향"])
    방법들 = sorted({t["방법"] for t in r.직접행}, key=lambda m: ({"높음": 0, "중간": 1, "낮음": 2}.get(R.방법풀이(m)[1], 3), m))
    for m in 방법들:
        뜻, 신 = R.방법풀이(m)
        기준.cell(row=row, column=1, value=m); 기준.cell(row=row, column=2, value=뜻); 기준.cell(row=row, column=3, value=신)
        기준.cell(row=row, column=4, value=f'=COUNTIF({J["X"]},A{row})').number_format = NUM
        기준.cell(row=row, column=5, value=f'=SUMIFS({J["N"]},{J["X"]},A{row})').number_format = NUM
        if 신 == "낮음":
            for c in range(1, 6):
                기준.cell(row=row, column=c).fill = INPUT
        row += 1
    row += 1
    제목("5. 정책과 적요 사용 원칙")
    기준.cell(row=row, column=1, value="기준"); 기준.cell(row=row, column=2, value=f"{r.기준}. 이자·배당·법인세 정책: "
                                                                            + ", ".join(f"{k}={v}" for k, v in r.정책.items())); row += 1
    기준.cell(row=row, column=1, value="적요(강)"); 기준.cell(row=row, column=2, value="그 말만으로 활동이 정해지는 것(법인세, 급여, 대출 실행·상환, 설비 구입, 보증금, 해외송금 등). "
                                                                           "적요 글자만 보고, 적합성 점검에도 쓴다"); row += 1
    기준.cell(row=row, column=1, value="적요(약)"); 기준.cell(row=row, column=2, value="대개 그렇지만 다른 뜻일 수 있는 것(대금, 정산, 수수료, 카드대금). 거래처 이름까지 보고, 원천을 못 찾은 줄에만 쓴다"); row += 1
    기준.cell(row=row, column=1, value="계정별 기본 활동"); 기준.cell(row=row, column=2, value="계정규칙 시트. 판정 근거 열이 어떤 키워드·재무상태표 분류로 정했는지 적는다"); row += 1

    # ════════ 독립검산
    독.column_dimensions["A"].width = 12; 독.column_dimensions["B"].width = 22
    for c, w in zip("CDEFGHIJKLMNOPQRS", (10, 8, 16, 16, 16, 16, 16, 16, 9, 52, 16, 16, 16, 16, 16, 16, 12)):
        독.column_dimensions[c].width = w
    독.sheet_view.showGridLines = False
    독["A1"] = "독립 검산"; 독["A1"].font = Font(bold=True, size=13)
    독["A2"] = ("직접법과 간접법(정산표)의 일치는 같은 분개장 분류를 공유해 구조상 늘 맞는다. 분류가 맞는지는 아래처럼 분류와 무관한 사실"
               "(손익계정 성격, 잔액, 비현금 대체액, 손익계산서)로 확인한다")
    독["A2"].font = Font(size=9, color="6B7688")
    독["A3"] = ("B 는 분개장만으로는 범위가 넓을 수 있다(기말 미지급금 중 자산 관련 몫을 모르므로). 기말 계정명세서의 투자 관련 금액을 "
               "노란 칸에 넣으면 그 범위가 한 점이 되어 강한 검산이 된다")
    독["A3"].font = Font(size=9, color="6B7688")
    row = 5
    독결과 = []
    def 독제목(t, hs):
        nonlocal row
        독.cell(row=row, column=1, value=t).font = Font(bold=True, size=11); row += 1
        for j, h in enumerate(hs, 1):
            c = 독.cell(row=row, column=j, value=h); c.fill, c.font = NAVY, HF
            c.alignment = Alignment(wrap_text=True)
        row += 1
    독제목("A. 비현금 성격 손익(감가상각비·상각비·대손상각비·평가·손상·처분손익)은 현금분이 0 이어야 한다",
          ["계정코드", "계정", "", "", "현금분", "", "", "", "", "", "판정", ""])
    a시작 = row
    for x in r.독립["A"]:
        독.cell(row=row, column=1, value=x["코드"]); 독.cell(row=row, column=2, value=x["계정"])
        독.cell(row=row, column=5, value=f"=SUMIFS({J['N']},{J['D']},A{row})").number_format = NUM
        독.cell(row=row, column=11, value=f'=IF(ROUND(E{row},0)=0,"PASS","FAIL")')
        row += 1
    if row == a시작:
        독.cell(row=row, column=2, value="해당 계정 없음"); row += 1
    독결과.append(("A. 비현금 성격 손익의 현금분", a시작, row - 1))
    row += 1
    독제목("B. 미지급금·선급금 등에서 투자·재무로 분류한 현금은 그 계정의 잔액과 비현금 대체액이 허락하는 범위 안이어야 한다",
          ["계정코드", "계정", "활동", "방향", "비현금 대체액\n(짝, 하한용)", "분류한 현금", "기초", "기말", "하한", "상한", "판정",
           "범위 계산", "관련 비현금\n(상한용)", "현금 없이 반대로\n움직인 금액(하한 완화)", "범위 폭\n(좁을수록 강한 검산)",
           "기말에 남아야 할\n관련 잔액 하한", "기말에 남아야 할\n관련 잔액 상한", "기말 명세서 확인값\n(입력)", "명세서 대조"])
    b시작 = row
    정행코드 = {c: se for c, se in 정행.items()}
    for x in r.독립["B"]:
        c = x["코드"]
        독.cell(row=row, column=1, value=c); 독.cell(row=row, column=2, value=x["계정"])
        독.cell(row=row, column=3, value=x["활동"]); 독.cell(row=row, column=4, value=x["방향"])
        이름들 = sorted({lab for (k, lab) in (kk for d in r.줄정산.values() for kk in d)
                       if k == "비현금" and r.정산라벨활동.get(lab) == x["활동"]})
        조건 = '">0"' if x["방향"] == "유출" else '"<0"'
        부호 = "" if x["방향"] == "유출" else "-"
        대체식 = "+".join(f'SUMIFS({J["AF"]},{J["D"]},A{row},{J["AD"]},"비현금",{J["AE"]},{_식값(n_)},{J["AF"]},{조건})'
                        for n_ in 이름들) or "0"
        독.cell(row=row, column=5, value=f"={부호}({대체식})")
        독.cell(row=row, column=6, value=(f'=-SUMIFS({J["N"]},{J["D"]},A{row},{J["U"]},"{x["활동"]}",{J["N"]},"<0")' if x["방향"] == "유출"
                                          else f'=SUMIFS({J["N"]},{J["D"]},A{row},{J["U"]},"{x["활동"]}",{J["N"]},">0")'))
        상조건 = '">0"' if x["방향"] == "유출" else '"<0"'
        # 다른 추적 계정을 거치는 몫은 그 계정들의 실제 투자·재무 대체액까지만
        경유들 = sorted(r.경유대상.get((c, x["활동"], x["방향"]), set()), key=str)
        이름들전 = sorted({lab for (k, lab) in (kk for d in r.줄정산.values() for kk in d)
                        if k == "비현금" and r.정산라벨활동.get(lab) == x["활동"]})
        경유상한 = "+".join(f'ABS(SUMIFS({J["AF"]},{J["D"]},{_식값(u)},{J["AD"]},"비현금",{J["AE"]},{_식값(n_)}))'
                          for u in 경유들 for n_ in 이름들전) or "0"
        독.cell(row=row, column=13, value=(f'=MAX(E{row},{부호}SUMIFS({J["AG"]},{J["D"]},A{row},{J["AH"]},"{x["활동"]}",{J["AG"]},{상조건})'
                                           f'+MIN({부호}SUMIFS({J["AG"]},{J["D"]},A{row},{J["AH"]},"추적경유",{J["AG"]},{상조건}),{경유상한}))'))
        독.cell(row=row, column=13).number_format = NUM
        독.cell(row=row, column=14, value=(f"=MAX(0,SUMIFS({J['G']},{J['D']},A{row})+SUMIFS({J['N']},{J['D']},A{row},{J['N']},\"<0\"))"
                                           if x["방향"] == "유출" else
                                           f"=MAX(0,SUMIFS({J['H']},{J['D']},A{row})-SUMIFS({J['N']},{J['D']},A{row},{J['N']},\">0\"))"))
        독.cell(row=row, column=14).number_format = NUM
        s_ = 정행코드.get(c, (None,))[0]
        독.cell(row=row, column=7, value=f"=MAX(0,간접법_정산표!D{s_})" if s_ else 0)
        독.cell(row=row, column=8, value=f"=MAX(0,간접법_정산표!C{s_})" if s_ else 0)
        if x["발생먼저"]:
            독.cell(row=row, column=9, value=f"=MAX(0,E{row}-H{row}-N{row})"); 독.cell(row=row, column=10, value=f"=M{row}+G{row}")
            독.cell(row=row, column=12, value="발생이 먼저(미지급금 결제·미수금 회수): 하한 = 짝 대체 - 기말 - 현금 없이 반대로 움직인 금액, 상한 = 관련 비현금 + 기초")
        else:
            독.cell(row=row, column=9, value=f"=MAX(0,E{row}-G{row}-N{row})"); 독.cell(row=row, column=10, value=f"=M{row}+H{row}")
            독.cell(row=row, column=12, value="현금이 먼저(선급금·가지급금·가수금): 하한 = 짝 대체 - 기초 - 현금 없이 반대로 움직인 금액, 상한 = 관련 비현금 + 기말")
        독.cell(row=row, column=11, value="잔액 열어둠" if r.잔액열어둠 else f'=IF(AND(F{row}>=I{row}-0.5,F{row}<=J{row}+0.5),"PASS","FAIL")')
        독.cell(row=row, column=15, value=f"=J{row}-I{row}").number_format = NUM
        # 분류가 맞다면 기말에 이 계정에 남아 있어야 할 투자·재무 관련 잔액. 기말 계정명세서의 해당 금액을 R 에 넣으면 대조된다
        if x["발생먼저"]:
            독.cell(row=row, column=16, value=f"=MIN(H{row},MAX(0,E{row}-F{row}))")
            독.cell(row=row, column=17, value=f"=MIN(H{row},MAX(0,M{row}+G{row}-F{row}))")
        else:
            독.cell(row=row, column=16, value=f"=MIN(H{row},MAX(0,F{row}-M{row}))")
            독.cell(row=row, column=17, value=f"=MIN(H{row},MAX(0,G{row}+F{row}-E{row}))")
        독.cell(row=row, column=18).fill = INPUT
        독.cell(row=row, column=19, value=f'=IF(R{row}="","명세서 없음",IF(AND(R{row}>=P{row}-0.5,R{row}<=Q{row}+0.5),"PASS","FAIL"))')
        for col in (16, 17, 18):
            독.cell(row=row, column=col).number_format = NUM
        for col in range(5, 11):
            독.cell(row=row, column=col).number_format = NUM
        row += 1
    if row == b시작:
        독.cell(row=row, column=2, value="추적 계정에서 투자·재무로 간 현금 없음"); row += 1
    독결과.append(("B. 추적 계정 → 투자·재무 현금의 범위", b시작, row - 1))
    B명세 = (b시작, row - 1)
    row += 1
    독제목("C. 비현금 성격 손익은 손익계산서 = 정산표 가감 + 현금분(부호 맞춰)",
          ["계정코드", "계정", "구분", "", "손익계산서(분개장)", "정산표 가감", "현금분", "차이", "", "", "판정", ""])
    c시작 = row
    끝정 = BS끝
    for x in r.독립["C"]:
        c = x["코드"]; 비용 = 계정[c]["구분"] == "비용"
        독.cell(row=row, column=1, value=c); 독.cell(row=row, column=2, value=x["계정"]); 독.cell(row=row, column=3, value=계정[c]["구분"])
        독.cell(row=row, column=5, value=(f"=SUMIFS({J['G']},{J['D']},A{row})-SUMIFS({J['H']},{J['D']},A{row})" if 비용
                                          else f"=SUMIFS({J['H']},{J['D']},A{row})-SUMIFS({J['G']},{J['D']},A{row})"))
        가 = (f"SUMIF(간접법_정산표!$H$4:$H${끝정},B{row},간접법_정산표!$I$4:$I${끝정})",
             f"SUMIF(간접법_정산표!$J$4:$J${끝정},B{row},간접법_정산표!$K$4:$K${끝정})")
        독.cell(row=row, column=6, value=f"={가[0]}-{가[1]}" if 비용 else f"={가[1]}-{가[0]}")
        독.cell(row=row, column=7, value=f"=SUMIFS({J['N']},{J['D']},A{row})")
        독.cell(row=row, column=8, value=f"=E{row}-F{row}+G{row}" if 비용 else f"=E{row}-F{row}-G{row}")
        독.cell(row=row, column=11, value=f'=IF(ROUND(H{row},0)=0,"PASS","FAIL")')
        for col in range(5, 9):
            독.cell(row=row, column=col).number_format = NUM
        row += 1
    if row == c시작:
        독.cell(row=row, column=2, value="해당 계정 없음"); row += 1
    독결과.append(("C. 비현금 성격 손익의 손익계산서 대사", c시작, row - 1))
    row += 1
    # D. 참고: 분석적 대사(판정 없음)
    독.cell(row=row, column=1, value="D. 참고: 분석적 대사(판정 없이 크기만 본다)").font = Font(bold=True, size=11); row += 1
    for j, h in enumerate(["항목", "", "", "", "직접법", "비교 기준", "차이", "", "", "", "", "비교 기준"], 1):
        c = 독.cell(row=row, column=j, value=h); c.fill, c.font = NAVY, HF
    row += 1
    이자비용 = [c for c in 계정 if 계정[c]["구분"] == "비용" and 계정[c].get("정책키") == "이자지급"]
    이자수익 = [c for c in 계정 if 계정[c]["구분"] == "수익" and 계정[c].get("정책키") == "이자수취"]
    법인세비용 = [c for c in 계정 if 계정[c]["구분"] == "비용" and 계정[c].get("정책키") == "법인세"]
    def 합식(코드들, 비용):
        if not 코드들:
            return "0"
        return "+".join((f"(SUMIFS({J['G']},{J['D']},{_식값(c)})-SUMIFS({J['H']},{J['D']},{_식값(c)}))" if 비용
                         else f"(SUMIFS({J['H']},{J['D']},{_식값(c)})-SUMIFS({J['G']},{J['D']},{_식값(c)}))") for c in 코드들)
    참고 = [
        ("이자의 지급 vs 이자비용", f'=-SUMIFS({J["N"]},{J["V"]},"이자의 지급")', f"={합식(이자비용, True)}", "이자비용(미지급이자 증감만큼 다를 수 있다)"),
        ("이자의 수취 vs 이자수익", f'=SUMIFS({J["N"]},{J["V"]},"이자의 수취")', f"={합식(이자수익, False)}", "이자수익(미수수익·원천징수만큼 다를 수 있다)"),
        ("법인세의 납부 vs 법인세비용", f'=-SUMIFS({J["N"]},{J["V"]},"법인세의 납부")-SUMIFS({J["N"]},{J["V"]},"법인세의 환급")',
         f"={합식(법인세비용, True)}", "법인세비용(선납·미지급 법인세 증감만큼 다를 수 있다)"),
        ("고객으로부터의 유입(0 이상이어야)", f'=SUMIFS({J["N"]},{J["V"]},"고객으로부터의 유입")', "=0", "부호 점검"),
        ("공급자·종업원에 대한 유출(0 이하여야)", f'=SUMIFS({J["N"]},{J["V"]},"공급자에 대한 유출")+SUMIFS({J["N"]},{J["V"]},"종업원에 대한 유출")', "=0", "부호 점검"),
    ]
    for 이름, 식1, 식2, 설명 in 참고:
        독.cell(row=row, column=1, value=이름)
        독.cell(row=row, column=5, value=식1).number_format = NUM
        독.cell(row=row, column=6, value=식2).number_format = NUM
        독.cell(row=row, column=7, value=f"=E{row}-F{row}").number_format = NUM
        독.cell(row=row, column=12, value=설명)
        row += 1
    독요약 = []
    for 이름, a_, b_ in 독결과:
        독요약.append((이름, f'=COUNTIF(독립검산!K{a_}:K{b_},"FAIL")'))
    독요약.insert(2, ("B'. 기말 계정명세서와 대조(명세서를 넣은 행만)", f'=COUNTIF(독립검산!S{B명세[0]}:S{B명세[1]},"FAIL")'))

    # ════════ 요약·검증
    요약.column_dimensions["A"].width = 52; 요약.column_dimensions["B"].width = 20; 요약.column_dimensions["C"].width = 10
    요약.sheet_view.showGridLines = False
    요약["A1"] = f"{회사} 현금흐름표 검증"; 요약["A1"].font = Font(bold=True, size=13)
    요약["A2"] = f"{기간[0]} ~ {기간[1]}"; 요약["A2"].font = Font(size=9, color="6B7688")
    row = 4
    for m in 메모:
        요약.cell(row=row, column=1, value=m).font = Font(size=9, color="6B7688"); row += 1
    row += 1
    for j, h in enumerate(["검증", "값 (0 이어야 한다)", "판정"], 1):
        c = 요약.cell(row=row, column=j, value=h); c.fill, c.font = NAVY, HF
    row += 1
    현금순 = f'SUMIFS({J["G"]},{J["M"]},"현금")-SUMIFS({J["H"]},{J["M"]},"현금")'
    검증식 = [
        ("분개장 차변 합계 = 대변 합계", f"=SUM({J['G']})-SUM({J['H']})"),
        ("분개장 현금영향 합계 = 현금 계정 순증감", f"=SUM({J['N']})-({현금순})"),
        ("분개장 미분류 현금 줄 수",
         f'=COUNTIFS({J["N"]},"<>",{J["U"]},"<>영업",{J["U"]},"<>투자",{J["U"]},"<>재무",{J["U"]},"<>환율")'),
        ("직접법 기말현금 = 재무상태표 기말현금", f"={DJ('차이')}"),
        ("정산표 행 검산(절댓값 합)", f"=간접법_정산표!{정요약['행검산']}"),
        ("정산표 비현금 차변 = 대변", f"=간접법_정산표!{정요약['비현금']}"),
        ("정산표 비용가산·수익차감 = 분개장 정산 손익", f"=간접법_정산표!{정요약['손익칸']}"),
        ("정산표 비현금 칸 = 분개장 정산 비현금", f"=간접법_정산표!{정요약['비현금칸']}"),
        ("정산표 활동 합계 = 현금 증감", f"=간접법_정산표!{정요약['합계']}-간접법_정산표!{정요약['현금']}"),
        ("간접법(정산표) 영업활동 = 직접법 영업활동", f"={GJ('영업차이')}"),
    ]
    if getattr(r, "손익NI", None) is not None:
        검증식.append((f"분개장 당기순이익 = 손익계산서 당기순이익 ({r.손익NI:,})", f"=간접법_정산표!G{당기순}-({r.손익NI})"))
    if 기말BS:
        주현 = sum(b["잔액"] for b in 기말BS if b["코드"] in r.현금계정)
        검증식.append((f"분개장 기말현금 = 주어진 기말 재무상태표 현금 ({주현:,})",
                      "=" + ("+".join(f"간접법_정산표!C{i}" for i in 현금행) or "0") + f"-{주현}"))
    검시작 = row
    for 라벨, 식 in 검증식:
        요약.cell(row=row, column=1, value=라벨)
        요약.cell(row=row, column=2, value=식).number_format = NUM
        요약.cell(row=row, column=3, value=f'=IF(ROUND(B{row},0)=0,"PASS","FAIL")')
        for c in range(1, 4):
            요약.cell(row=row, column=c).border = Border(bottom=thin)
        row += 1
    요약.cell(row=row, column=1, value="종합(내부 일관성)").font = Font(bold=True)
    요약.cell(row=row, column=3, value=f'=IF(COUNTIF(C{검시작}:C{row - 1},"PASS")={len(검증식)},"PASS","FAIL")').font = Font(bold=True)
    row += 1
    요약.cell(row=row, column=1, value="위는 산술·대사·구조가 맞는지다. 분류가 맞는지는 아래 독립 검산이 본다").font = Font(size=9, color="6B7688")
    row += 2
    for j, h in enumerate(["독립 검산(분류와 무관한 사실로 확인)", "FAIL 수", "판정"], 1):
        c = 요약.cell(row=row, column=j, value=h); c.fill, c.font = NAVY, HF
    row += 1
    독시작 = row
    for 이름, 식 in 독요약:
        요약.cell(row=row, column=1, value=이름)
        요약.cell(row=row, column=2, value=식)
        요약.cell(row=row, column=3, value=f'=IF(B{row}=0,"PASS","FAIL")')
        row += 1
    요약.cell(row=row, column=1, value="종합(독립 검산)").font = Font(bold=True)
    요약.cell(row=row, column=3, value=f'=IF(COUNTIF(C{독시작}:C{row - 1},"PASS")={len(독요약)},"PASS","FAIL")').font = Font(bold=True)
    독종합 = f"C{row}"
    row += 2
    if 기말BS:
        요약.cell(row=row, column=1, value="주어진 기말 재무상태표와 다른 계정 수(정산표 기말 차이)")
        요약.cell(row=row, column=2, value=f'=COUNTIF(간접법_정산표!AC4:AC{끝},">0.5")+COUNTIF(간접법_정산표!AC4:AC{끝},"<-0.5")')
        row += 1
    요약.cell(row=row, column=1, value="확인사항 (회계사 판단 필요)")
    요약.cell(row=row, column=2, value=f"=COUNTA(확인사항!C2:C{max(확.max_row, 2)})"); row += 1
    요약.cell(row=row, column=1, value="재분류된 분개장 행")
    요약.cell(row=row, column=2, value=f'=COUNTIF(분개장!$W$2:$W${NJ},"재분류")'); row += 2
    for 활 in 활동순서:
        요약.cell(row=row, column=1, value=활동명[활])
        요약.cell(row=row, column=2, value=f"={DJ(활)}").number_format = NUM
        row += 1
    요약.cell(row=row, column=1, value="현금의 증감")
    요약.cell(row=row, column=2, value=f"={DJ('증감')}").number_format = NUM

    wb.save(경로)
    return {"직접법": 직셀, "간접법": 간셀, "정산": 정요약, "요약검증행": (검시작, 검시작 + len(검증식) - 1), "독종합": 독종합}

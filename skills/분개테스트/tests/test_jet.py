"""모의 분개장(가상 회사 가온정밀(가상))으로 분개 테스트 조서를 만든다. 2026-10-21 시험 열세 가지를 더하며 합성 분개장 시험을 붙였다."""
import os
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parents[1]
EX = HERE / "examples" / "모의자료"
sys.path.insert(0, str(HERE / "scripts"))
ENV = {**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1", "AUDIT_OCR": "off"}


def _run(args):
    r = subprocess.run([sys.executable, str(HERE / "scripts" / "run.py"), *map(str, args)],
                       env=ENV, capture_output=True, text=True, encoding="utf-8", errors="replace")
    return r.returncode, r.stdout + r.stderr


def test_jet_example(tmp_path):
    code, log = _run([EX / "분개장_2025.xlsx", "--after", EX / "분개장_2026_기말후.xlsx", "--tb", EX / "재무제표_시산표.xlsx",
                      "--fy-end", "2025-12-31", "--pm", "60000000", "--company", "가온정밀(가상)", "--out", tmp_path / "out"])
    assert code == 0, log
    assert "분개장 1,751행" in log, log
    for line in ("주말 · 공휴일 전표: 1건", "기말 후 취소 · 역분개: 6건", "결산일 근처 큰 금액: 32건", "드문 계정 조합: 10건",
                 "끝자리 0 반복: 16건", "적요 없음: 0건", "수기 조정 전표: 4건",
                 # 2026-10-21 추가 시험
                 "전표 형식 이상: 0건", "전표번호 중복 · 순서 역전: 0건", "한 번만 나온 거래처의 큰 금액: 4건", "결산일 근처 이익을 늘린 전표: 5건",
                 "중복 의심 전표: 0건", "대표 · 주주 · 임원 관련 계정: 0건", "자본 계정 분개: 0건", "적요 주의 낱말: 3건",
                 "계정별로 평소보다 큰 금액: 4건", "작성자 = 승인자 · 승인 없음: 5건", "드물게 분개한 작성자: 5건", "결산 후 입력 · 늦은 입력: 6건",
                 "업무 시간이 아닌 입력: 6건"):
        assert line in log, line
    assert "여러 시험에 걸린 전표: " in log and "202512-0028(10개)" in log, log
    p = tmp_path / "out" / "3600-JET_분개테스트.xlsx"
    wb = load_workbook(p, read_only=True)
    for name in ("요약", "결과 목록", "여러 시험에 걸린 전표", "전표번호 연속성", "분포 분석"):
        assert name in wb.sheetnames, name


# ---------------------------------------------------------------- 합성 분개장(새 시험마다 걸릴 전표를 하나씩 심었다)
HEAD = ["일자", "전표번호", "계정코드", "계정명", "차변", "대변", "적요", "거래처", "작성자", "승인자", "입력일시"]


def _workday(d: date) -> datetime:
    """평소 전표는 다음 업무일 오전 10시에 입력한 것으로 둔다(주말 · 공휴일 입력 시험에 걸리지 않게)."""
    import yaml

    hol = {date.fromisoformat(str(x)) for x in yaml.safe_load((HERE / "scripts" / "definitions" / "jet.yaml").read_text(encoding="utf-8"))["holidays"]}
    while d.weekday() >= 5 or d in hol:
        d = date.fromordinal(d.toordinal() + 1)
    return datetime(d.year, d.month, d.day, 10, 0)


def _synthetic(path: Path) -> None:
    rows = []

    def v(d, no, lines, user="김회계", appr="이팀장", when=None):
        when = when if when is not None else _workday(d)
        for code, name, dr, cr, memo, party in lines:
            rows.append([d, no, code, name, dr, cr, memo, party, user, appr, when])

    for m in range(1, 13):
        for k in range(1, 11):                     # 달마다 열 장, 일자는 번호 순서대로
            if (m, k) == (9, 7):
                continue                           # 빠진 번호 202509-0007
            d = date(2025, m, min(3 * k, 28))
            amt = 1_000_000 * (k + m) + 12_345 * k
            party = f"거래처{k}"
            v(d, f"2025{m:02d}-{k:04d}", [("108", "외상매출금", amt, 0, "제품 매출", party), ("404", "제품매출", 0, amt, "제품 매출", party)])
    # 차대 불일치
    v(date(2025, 3, 29), "202503-0011", [("103", "보통예금", 1_000_000, 0, "입금", "거래처1"), ("108", "외상매출금", 0, 900_000, "입금", "거래처1")])
    # 음수 금액(마이너스 역분개)
    v(date(2025, 4, 29), "202504-0011", [("108", "외상매출금", -2_000_000, 0, "매출 정정", "거래처2"), ("404", "제품매출", 0, -2_000_000, "매출 정정", "거래처2")])
    # 중복 지급 둘
    for no, dd in (("202505-0011", 20), ("202505-0012", 23)):
        v(date(2025, 5, dd), no, [("251", "외상매입금", 3_300_000, 0, "대금 지급", "가나상사"), ("103", "보통예금", 0, 3_300_000, "대금 지급", "가나상사")])
    # 대표이사 가지급금
    v(date(2025, 6, 29), "202506-0011", [("134", "가지급금", 5_000_000, 0, "대표 가지급", "대표이사 홍길동"),
                                         ("103", "보통예금", 0, 5_000_000, "대표 가지급", "대표이사 홍길동")])
    # 번호는 뒤인데 일자가 앞선 전표(끼워 넣음)
    v(date(2025, 7, 2), "202507-0011", [("103", "보통예금", 7_000_000, 0, "회수", "거래처3"), ("108", "외상매출금", 0, 7_000_000, "회수", "거래처3")])
    # 같은 전표번호가 다른 날짜에도
    v(date(2025, 8, 26), "202508-0005", [("103", "보통예금", 1_500_000, 0, "회수", "거래처4"), ("108", "외상매출금", 0, 1_500_000, "회수", "거래처4")])
    # 자본 계정 · 결산 후 입력 · 밤 입력 · 자기 승인 · 드문 작성자(입력일시는 글자로)
    v(date(2025, 12, 31), "202512-0011", [("375", "이월이익잉여금", 9_000_000, 0, "배당 결의", None), ("265", "미지급배당금", 0, 9_000_000, "배당 결의", None)],
      user="대표", appr="대표", when="2026-01-20 23:10")
    wb = Workbook()
    ws = wb.active
    ws.title = "분개장"
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_jet_synthetic(tmp_path):
    src = tmp_path / "합성_분개장.xlsx"
    _synthetic(src)
    code, log = _run([src, "--fy-end", "2025-12-31", "--pm", "4000000", "--company", "합성", "--out", tmp_path / "out"])
    assert code == 0, log
    wb = load_workbook(tmp_path / "out" / "3600-JET_분개테스트.xlsx")
    ws = wb["결과 목록"]
    hits: dict[str, list] = {}
    for r in ws.iter_rows(min_row=8, values_only=True):
        if r[0] and r[3]:
            hits.setdefault(r[0], []).append((r[3], r[10]))
    nos = {k: sorted(n for n, _ in v) for k, v in hits.items()}
    assert nos["전표 형식 이상"] == ["202503-0011", "202504-0011"], nos
    assert any("음수 금액 줄 2개" in w for _, w in hits["전표 형식 이상"])
    assert nos["중복 의심 전표"] == ["202505-0011", "202505-0012"], nos
    assert nos["대표 · 주주 · 임원 관련 계정"] == ["202506-0011"], nos
    assert nos["자본 계정 분개"] == ["202512-0011"], nos
    seq = dict(hits["전표번호 중복 · 순서 역전"])
    assert "202507-0011" in seq and "일자가" in seq["202507-0011"], seq
    assert "202508-0005" in seq and "같은 전표번호" in seq["202508-0005"], seq
    assert nos["작성자 = 승인자 · 승인 없음"] == ["202512-0011"], nos
    assert nos["드물게 분개한 작성자"] == ["202512-0011"], nos
    assert nos["결산 후 입력 · 늦은 입력"] == ["202512-0011"], nos
    assert nos["업무 시간이 아닌 입력"] == ["202512-0011"], nos
    assert "202512-0011" not in nos.get("결산일 근처 이익을 늘린 전표", []), nos      # 자본 계정은 손익이 아니다
    gaps = wb["전표번호 연속성"]
    assert any(r[0] == "202509-" and r[4] == 1 and str(r[5]) == "7" for r in gaps.iter_rows(values_only=True)), "빠진 번호 202509-0007"


# ---------------------------------------------------------------- 도구 함수
def test_helpers():
    from audit.workpapers import common, jet

    spec = common.load_def("jet.yaml")
    acc = spec["accounts"]

    def line(name, code=None, dr=0.0, cr=0.0):
        return common.Line(date=date(2025, 1, 1), voucher="1", row=None, code=code, name=name, dr=dr, cr=cr, memo="", party=None)

    # 코드 없는 분개장: 외상매출금 · 매출원가는 매출이 아니다. 선급비용 · 이익잉여금은 손익이 아니다
    assert jet._in(line("제품매출"), acc["revenue"]) and not jet._in(line("외상매출금"), acc["revenue"])
    assert not jet._in(line("매출원가"), acc["revenue"])
    assert jet._in(line("급여"), acc["pnl"]) and not jet._in(line("선급비용"), acc["pnl"]) and not jet._in(line("이익잉여금"), acc["pnl"])
    # 코드가 있으면 코드로: 제조원가 급여(504)는 손익 효과에 넣지 않는다
    assert not jet._in(line("급여", code="504"), acc["pnl"]) and jet._in(line("급여", code="801"), acc["pnl"])
    # 특수관계 묶음은 코드 체계가 달라도 계정명으로 잡는다
    assert jet._in(line("주주임원종업원단기채권", code="1370"), acc["related"])

    # 입력일시 읽기
    assert jet._when(datetime(2026, 1, 20, 23, 10)) == (date(2026, 1, 20), 23, 10)
    assert jet._when(datetime(2026, 1, 20)) == (date(2026, 1, 20), None, None)
    assert jet._when("2026-01-20 23:10") == (date(2026, 1, 20), 23, 10)
    assert jet._when("20260120231000") == (date(2026, 1, 20), 23, 10)
    assert jet._when(46042.5) == (date(2026, 1, 20), 12, 0)
    assert jet._when("2026.01.20") == (date(2026, 1, 20), None, None)

    # 줄지 않는 가장 긴 부분열: 끼워 넣은 하나와 잘못 친 하나만 빠진다
    assert jet._lnds([1, 2, 3, 0, 4, 5]) == {0, 1, 2, 4, 5}
    assert jet._lnds([1, 2, 99, 3, 4, 5]) == {0, 1, 3, 4, 5}

    # 2026-10-21 검토로 더한 꼴: 오전/오후 · 「년 월 일」 · 붙여 쓴 날짜 · 시각만 · 0시 0분 0초(날짜만) · 일련번호 반올림
    from datetime import time as dtime
    assert jet._when("2026-01-20 오후 11:10:05") == (date(2026, 1, 20), 23, 10)
    assert jet._when("2026-01-20T23:10:05.123Z") == (date(2026, 1, 20), 23, 10)
    assert jet._when("2026년 1월 20일 23:10") == (date(2026, 1, 20), 23, 10)
    assert jet._when("20260120 2310") == (date(2026, 1, 20), 23, 10)
    assert jet._when("2026-01-20 00:00:00") == (date(2026, 1, 20), None, None)
    assert jet._when(dtime(23, 10)) == (None, 23, 10) and jet._when("오후 11:10") == (None, 23, 10)
    assert jet._when(46042.999999) == (date(2026, 1, 21), None, None)
    # 입력일자 열과 입력시각 열이 따로면 합친다
    v = common.Voucher(date=date(2025, 12, 31), no="1", lines=[common.Line(date=date(2025, 12, 31), voucher="1", row=None, code="1", name="a", dr=1, cr=0,
                                                                          memo="", party=None, extra={"입력일자": date(2026, 1, 9), "입력시각": "23:40"})], file="a")
    e = jet.entry_of(v, {"entered": "입력일자", "entered_time": "입력시각"})
    assert (e["entered"], e["hour"], e["minute"]) == (date(2026, 1, 9), 23, 40), e
    # 손익 · 자본 · 특수관계 계정명(코드 없는 분개장)
    for n in ("복리후생비", "지급수수료", "경상연구개발비", "대손충당금환입", "법인세비용"):
        assert jet._in(line(n), acc["pnl"]), n
    for n in ("자기주식처분이익", "퇴직급여충당부채", "미지급법인세", "기타포괄손익-공정가치측정금융자산평가이익"):
        assert not jet._in(line(n), acc["pnl"]), n
    assert not jet._in(line("기타포괄손익-공정가치측정금융자산", code="178"), acc["equity"])
    assert jet._in(line("기타포괄손익-공정가치측정금융자산평가이익", code="394"), acc["equity"])
    assert jet._in(line("임원장기대여금"), acc["related"]) and not jet._in(line("임원급여"), acc["related"])

    t = next(t for t in spec["tests"] if t["id"] == "voucher_seq")

    def seq(items):
        return jet.voucher_seq([common.Voucher(date=d, no=no, lines=[], file="a") for d, no in items], t)

    # 일자마다 1번부터 붙이는 번호(더존): 번호가 날짜마다 겹쳐도 중복이 아니고 일자별 빠진 번호만 센다
    items = [(date(2025, 3, dd), f"{n:05d}") for dd in range(1, 6) for n in ((1, 2, 3) if dd != 4 else (1, 3))]
    r = seq(items)
    assert r["scheme"].startswith("일자마다") and r["hits"] == [] and r["gaps"] == 1, r
    # 전표가 몰린 하루가 있어도 일자별 번호로 본다(번호 개수가 아니라 전표 수로 센다)
    items = [(date(2025, 1, 1) + timedelta(days=k), f"{n:05d}") for k in range(0, 300, 3) for n in range(1, 6)]
    items += [(date(2025, 12, 31), f"{n:05d}") for n in range(1, 201)]
    r = seq(items)
    assert r["scheme"].startswith("일자마다") and r["hits"] == [] and r["gaps"] == 0, r["note"]
    # 앞부분 없이 달마다 0001 부터: 달 단위로 보고 가짜 빠진 번호를 만들지 않는다
    items = [(date(2025, m, 1 + n // 3), f"{n + 1:04d}") for m in range(1, 13) for n in range(60)]
    r = seq(items)
    assert r["scheme"].startswith("달마다") and r["gaps"] == 0 and r["hits"] == [], r["note"]
    # 번호대 둘(일반 00001~ · 매입매출 50001~): 그 사이 5만 개를 빠진 번호로 세지 않는다
    items = [(date(2025, 3, 5), f"{n:05d}") for n in list(range(1, 11)) + list(range(50001, 50006))]
    r = seq(items)
    assert r["gaps"] == 0, r["note"]
    # 번호가 10억 단위여도 펼치지 않는다
    items = [(date(2025, 1, 1) + timedelta(days=k), f"{2025000000000 + k * 7}") for k in range(50)]
    r = seq(items)
    assert r["gaps"] == 49 * 6, r["note"]


def test_no_pm_and_blank_approver(tmp_path):
    """수행중요성 없이 돌아가고, 승인자 칸이 대부분 비면 승인 없음을 전표마다 올리지 않는다(2026-10-21 검토)."""
    src = tmp_path / "합성_분개장.xlsx"
    _synthetic(src)
    wb = load_workbook(src)
    ws = wb.active
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 10).value = None                       # 승인자 열을 모두 비운다
    wb.save(src)
    code, log = _run([src, "--fy-end", "2025-12-31", "--company", "합성", "--out", tmp_path / "out"])
    assert code == 0, log
    assert "결산일 근처 큰 금액: 0건 (수행중요성이 없어" in log, log
    assert "작성자 = 승인자 · 승인 없음: 0건" in log and "목록에 올리지 않았습니다" in log, log


def test_example_without_codes(tmp_path):
    """계정코드 열을 지운 예시 분개장: 기말 후 취소는 계정명으로 짝을 찾고, 이익을 늘린 전표는 시산표 계정명으로 손익을 가린다(2026-10-21 검토)."""
    src = {}
    for name in ("분개장_2025.xlsx", "분개장_2026_기말후.xlsx"):
        wb = load_workbook(EX / name)
        ws = wb.worksheets[0]
        hi = next(r for r in range(1, 11) if "계정코드" in [str(c.value) for c in ws[r]])
        col = [str(c.value) for c in ws[hi]].index("계정코드") + 1
        ws.delete_cols(col)
        src[name] = tmp_path / name
        wb.save(src[name])
    code, log = _run([src["분개장_2025.xlsx"], "--after", src["분개장_2026_기말후.xlsx"], "--tb", EX / "재무제표_시산표.xlsx",
                      "--fy-end", "2025-12-31", "--pm", "60000000", "--company", "가온정밀(가상)", "--out", tmp_path / "out"])
    assert code == 0, log
    assert "기말 후 취소 · 역분개: 6건" in log and "결산일 근처 이익을 늘린 전표: 5건" in log, log

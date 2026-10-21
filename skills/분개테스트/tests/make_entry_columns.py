# -*- coding: utf-8 -*-
"""모의 분개장에 작성자 · 승인자 · 입력일시 열을 붙인다(2026-10-21).

행 · 금액 · 전표번호는 건드리지 않고 거래처 오른쪽에 세 열만 더한다. 그래서 기존 시험 일곱 가지의 건수와 시산표 대사는 그대로다.
이미 열이 있으면 값을 다시 채운다(여러 번 돌려도 같은 결과).

    python tests/make_entry_columns.py examples/모의자료/분개장_2025.xlsx examples/모의자료/분개장_2026_기말후.xlsx

심은 전표(가상 회사 가온정밀(가상), 이름은 모두 지어냈다)
- 평소: 매출 · 회수 · 매출 취소는 김민지, 매입 · 지급 · 결산 · 자금은 이수진, 급여 · 예수금 · 퇴직연금은 박정훈이 쓰고 재무팀장 정태호가 승인한다.
  입력은 전표 일자 업무 시간(09~17시)에 한다. 주말 · 공휴일 일자인 월말 결산 전표는 다음 업무일에 입력한다.
- 기말 후 취소된 매출 셋(202512-0028 · 0029 · 0058): 대표이사 최가온이 전표 일자 밤 10~11시에 입력했고 승인자가 없다.
- 기말재고 실사 조정 · 개발비 자산화(202512-0064 · 0066): 최가온이 2026-02-11 밤에 입력하고 스스로 승인했다(결산 후 입력).
- 12-31 일자 결산 전표(매출원가 대체 · 감가상각 · 법인세 계상 · 성과급)는 2026-01-07~08 에 입력했다(결산 후 입력, 정상 결산 절차).
- 주말 전표 202503-0028(기계장치 취득)은 그 토요일 오후에 입력했다.
"""
from __future__ import annotations

import sys
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

HEAD = ("작성자", "승인자", "입력일시")
APPROVER = "정태호"
CEO = "최가온"
HOLIDAYS = {date(2025, 1, 1), date(2025, 1, 27), date(2025, 1, 28), date(2025, 1, 29), date(2025, 1, 30), date(2025, 3, 3),
            date(2025, 5, 5), date(2025, 5, 6), date(2025, 6, 3), date(2025, 6, 6), date(2025, 8, 15), date(2025, 10, 3),
            date(2025, 10, 6), date(2025, 10, 7), date(2025, 10, 8), date(2025, 10, 9), date(2025, 12, 25), date(2026, 1, 1),
            date(2026, 2, 16), date(2026, 2, 17), date(2026, 2, 18), date(2026, 3, 2)}

# 전표번호 → (작성자, 승인자, 입력일시). 승인자 None 은 빈 칸
SPECIAL = {
    "202512-0028": (CEO, None, datetime(2025, 12, 18, 22, 47)),
    "202512-0029": (CEO, None, datetime(2025, 12, 22, 23, 15)),
    "202512-0058": (CEO, None, datetime(2025, 12, 29, 22, 31)),
    "202512-0064": (CEO, CEO, datetime(2026, 2, 11, 23, 42)),
    "202512-0066": (CEO, CEO, datetime(2026, 2, 11, 23, 51)),
    "202512-0059": ("이수진", APPROVER, datetime(2026, 1, 7, 10, 12)),
    "202512-0061": ("이수진", APPROVER, datetime(2026, 1, 7, 10, 25)),
    "202512-0063": ("이수진", APPROVER, datetime(2026, 1, 7, 16, 40)),
    "202512-0065": ("박정훈", APPROVER, datetime(2026, 1, 8, 11, 20)),
    "202503-0028": ("이수진", APPROVER, datetime(2025, 3, 15, 15, 20)),
}


def _workday(d: date) -> date:
    while d.weekday() >= 5 or d in HOLIDAYS:
        d += timedelta(days=1)
    return d


def _user(memo: str) -> str:
    if any(w in memo for w in ("급여", "예수금 납부", "퇴직", "성과급")):
        return "박정훈"
    if any(w in memo for w in ("제품 매출", "외상매출금 회수", "매출 취소")):
        return "김민지"
    return "이수진"


def fill(path: Path) -> int:
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    hi = next(r for r in range(1, 11) if "전표번호" in [str(c.value) for c in ws[r]])
    head = [str(c.value) if c.value is not None else "" for c in ws[hi]]
    col = {h: i + 1 for i, h in enumerate(head)}
    base = len([h for h in head if h and h not in HEAD])
    for k, h in enumerate(HEAD):
        c = ws.cell(hi, base + 1 + k, h)
        src = ws.cell(hi, base)
        c.font, c.fill, c.border, c.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
        ws.column_dimensions[c.column_letter].width = 18 if h == "입력일시" else 10
    # 전표마다 줄을 모은다
    rows: dict[str, list[int]] = {}
    for r in range(hi + 1, ws.max_row + 1):
        no = ws.cell(r, col["전표번호"]).value
        if no:
            rows.setdefault(str(no), []).append(r)
    n = 0
    for k, (no, rs) in enumerate(rows.items()):
        d = ws.cell(rs[0], col["일자"]).value
        d = d.date() if isinstance(d, datetime) else d
        memo = " ".join(str(ws.cell(r, col["적요"]).value or "") for r in rs)
        if no in SPECIAL:
            user, appr, when = SPECIAL[no]
        else:
            user, appr = _user(memo), APPROVER
            day = _workday(d)
            mins = (k * 37) % (8 * 60)
            when = datetime(day.year, day.month, day.day, 9, 10) + timedelta(minutes=mins)
        for r in rs:
            ws.cell(r, base + 1, user)
            ws.cell(r, base + 2, appr)
            c = ws.cell(r, base + 3, when)
            c.number_format = "yyyy-mm-dd hh:mm"
            n += 1
    wb.save(path)
    return n


if __name__ == "__main__":
    for p in sys.argv[1:]:
        print(p, fill(Path(p)), "줄")

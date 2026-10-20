"""전기 자료: 원천 판정 · 수집 · 계정 단위 전기 잔액 (R7, 2026-10-06).

전기 자료가 어떤 상태든 전기 숫자와 전기 판단을 이 모듈 한 곳에서 정한다. 감사범위가 당기만이어도 분석적 검토는
늘 전기 · 당기 두 해로 하므로, 계정마다 전기 잔액을 아래 우선순위로 정하고 출처와 감사받은 숫자인지를 남긴다.

  situation(d)                 "작년 폴더" · "전임 전기조서" · "전기조서 없음"
  prior_year_dir(d)            같은 회사 폴더의 '<결산 연도 − 1> 기말감사'(감사건이 있을 때만)
  prior_engagement(d)          읽을 작년 폴더. 보관본(_보관/<보고서일>)이 있으면 가장 늦은 보관본(해시로 고정된 최종본)
  collect(d, gathered=None)    _색인/시작자료.json 의 prior_source(전기 판단 정보)
  info(d)                      시작자료.json 의 prior_source. 없으면 collect 로 계산만 한다(쓰지 않음)
  sources(d)                   계정 단위 전기 잔액 원천 1~3 (작년 8800 · 전임 8100B · 회사제시 전기 시산표 / 비교식 전기 열)
  account_prior(d, tb, book)   계정마다 전기 잔액 · 출처 · 감사 여부와 원천 사이 차이

계정 단위 전기 잔액 우선순위
  1 작년 감사 폴더 8800 수정 후 시산표(보관본 우선)          감사받은 숫자
  2 전임감사인 전기조서 8100B 계정별 최종                    감사받은 숫자
  3 회사제시 전기 시산표(계정별) · 당기 시산표 파일의 전기 열  회사제시
  4 재무상태표: 당기 시산표 − 당기 분개장 / 손익: 전기 분개장   장부(분석이 book 으로 넘긴다)
  5 없음 → 분석은 표시과목 단위로 전기 재무제표를 쓴다

원천 1~3 이 마감 전 시산표(합계 0, 손익 계정 있음)면 이익잉여금에 전기 손익 계정 합을 더해 마감 후로 맞춘다.
원천 4 의 재무상태표 값(당기 기초)이 마감 후 금액이라 같은 기준으로 견주려는 것이다(start.current_statement 와 같은 규칙).
마감 후 시산표(재무상태표 계정만으로 합계 0)는 손익 정보가 없으므로 손익 계정의 원천으로 쓰지 않는다.

엔진은 판단을 정하지 않는다. 전기 미수정 왜곡의 당기 영향, 전기 이슈를 당기 이슈로 볼지, 대체절차가 충분한지는 조서 기입란이다.
숫자는 엔진이 파일에서 읽은 값만 쓴다. 작년 폴더는 읽기만 한다.
"""
from __future__ import annotations

import json
import os
import re
from datetime import date
from pathlib import Path

from . import figures, folder, reader

SITUATIONS = ("작년 폴더", "전임 전기조서", "전기조서 없음")
PRIOR_TB = "전기 시산표"
PRIOR_WP = "전기조서"
BASIS = {
    "last_year": "작년 감사 폴더 8800 수정 후 시산표",
    "predecessor": "전임 전기조서 8100B 최종",
    "prior_tb": "회사제시 전기 시산표",
    "comparative": "당기 시산표 파일의 전기 열",
    "book_bs": "당기 시산표 − 당기 분개장",
    "book_pl": "전기 분개장",
}
AUDITED = {"last_year": True, "predecessor": True, "prior_tb": False, "comparative": False, "book": False}
NONE = "없음"
BS_SECTIONS, PL_SECTIONS = ("자산", "부채", "자본"), ("수익", "비용")


# ---------------------------------------------------------------- 공통
def _read(p: Path, default=None):
    try:
        return json.loads(Path(p).read_text(encoding="utf-8")) if Path(p).exists() else default
    except (OSError, json.JSONDecodeError, UnicodeDecodeError):
        return default


def _case(d: Path) -> dict:
    return _read(Path(d) / "_색인" / "감사건.json", {}) or {}


def _rel(d: Path, p: Path) -> str:
    """감사 폴더 d 기준 상대경로(작년 폴더면 ../2024 기말감사/...). 항상 / 로 쓴다."""
    return Path(os.path.relpath(Path(p).resolve(), Path(d).resolve())).as_posix()


def _code(v) -> str | None:
    """계정코드 숫자 그대로(R8: 보조계정을 앞 3자리로 합치지 않는다). 시산표와 자리 수가 다르면 account_prior 가 맞춘다."""
    c = reader._code(v)
    return None if c is None else str(c)


def _c3(c: str) -> int:
    return int(str(c)[:3])


def _year_before(x: date) -> date:
    try:
        return x.replace(year=x.year - 1)
    except ValueError:                     # 2월 29일
        return x.replace(year=x.year - 1, day=28)


def _prior_end(d: Path) -> date | None:
    """전기말 = 사업연도 시작일 전날(반기 · 짧은 사업연도는 감사건 fiscal_year_start, R8)."""
    try:
        return date.fromordinal(folder.fiscal_start(_case(d)).toordinal() - 1)
    except (KeyError, ValueError, TypeError):
        return None


# ---------------------------------------------------------------- 작년 폴더
def prior_year_dir(d: Path) -> Path | None:
    """같은 회사 폴더의 작년 감사 폴더. 결산일 연도 − 1 의 '<연도> 기말감사' 이고 감사건 결산일이 전기 결산일과 같아야 한다."""
    d = Path(d).resolve()
    pe = _prior_end(d)
    if pe is None:
        return None
    pattern = folder.load_standard().get("engagement_folder", "{연도} 기말감사")
    p = d.parent / pattern.format(연도=pe.year)
    if p.resolve() == d or not (p / "_색인" / "감사건.json").is_file():
        return None
    fye = _case(p).get("fiscal_year_end")
    if fye and fye != pe.isoformat():
        return None
    return p


def _archive_key(name: str) -> tuple:
    m = re.fullmatch(r"(\d{4}-\d{2}-\d{2})(?:_(\d+))?", name)
    return (m.group(1), int(m.group(2) or 1)) if m else ("", 0)


def latest_archive(base: Path) -> Path | None:
    """가장 늦은 보관본. 보관.json 기록을 먼저 보고, 없으면 _보관 아래 날짜 이름 폴더를 본다."""
    cands = []
    for r in (_read(base / "_색인" / "보관.json", {}) or {}).get("archives", []):
        p = base / r.get("path", "")
        if (p / "_색인" / "감사건.json").is_file():
            cands.append(p)
    if not cands and (base / "_보관").is_dir():
        cands = [p for p in (base / "_보관").iterdir() if p.is_dir() and (p / "_색인" / "감사건.json").is_file()]
    cands = [p for p in cands if _archive_key(p.name)[0]]
    return max(cands, key=lambda p: _archive_key(p.name)) if cands else None


def prior_engagement(d: Path) -> Path | None:
    base = prior_year_dir(d)
    if base is None:
        return None
    return latest_archive(base) or base


def _received(d: Path, kind: str) -> list[str]:
    return sorted(figures.received(d, kind))


def situation(d: Path) -> str:
    if prior_engagement(d) is not None:
        return SITUATIONS[0]
    if _received(d, PRIOR_WP):
        return SITUATIONS[1]
    return SITUATIONS[2]


# ---------------------------------------------------------------- 시산표 원천 맞추기
def _normalize(values: dict[str, int], names: dict[str, str]) -> dict:
    """원천 하나를 마감 후 기준으로 맞춘다. {"values", "pl", "complete", "note"}."""
    vals = dict(values)
    total = sum(vals.values())
    pl_sum = sum(v for c, v in vals.items() if _c3(c) >= 400)
    bs_sum = total - pl_sum
    re_code = next((c for c in sorted(vals) if 370 <= _c3(c) <= 379 or "이익잉여금" in (names.get(c) or "")), None)
    note = None
    if abs(total) < 1 and any(vals.get(c) for c in vals if _c3(c) >= 400):
        pl, complete = True, True
        if re_code:
            vals[re_code] += pl_sum
            note = "마감 전 시산표라 이익잉여금에 전기 손익 계정 합을 더해 마감 후로 맞춤"
    elif abs(bs_sum) < 1:
        pl, complete = False, True
        note = "마감 후 시산표라 손익 계정 원천으로 쓰지 않음"
    else:
        pl, complete = True, False
        note = f"시산표 합계가 0 이 아님({total:,}). 마감 전후를 정하지 못해 이익잉여금을 그대로 둠"
    return {"values": vals, "pl": pl, "complete": complete, "note": note}


def _source(key: str, path: str, rows: list[tuple[str, str, int]], sheet: str | None = None) -> dict | None:
    """rows: [(코드, 계정명, 금액 차변 +)]. 같은 코드는 더한다."""
    vals: dict[str, int] = {}
    names: dict[str, str] = {}
    for code, name, amt in rows:
        if code is None or amt is None:
            continue
        vals[code] = vals.get(code, 0) + reader._int(float(amt))
        names.setdefault(code, name or "")
    if not vals:
        return None
    n = _normalize(vals, names)
    return {"key": key, "basis": BASIS[key], "audited": AUDITED[key], "path": path, "sheet": sheet, "names": names,
            "raw": vals, **n}


def _read_8800(d: Path, prev: Path) -> dict | None:
    from openpyxl import load_workbook

    from .report.base import wp_file

    p = wp_file(prev, "8800")
    if p is None:
        return None
    wb = load_workbook(p, data_only=True, read_only=True)
    try:
        ws = wb["시산표"] if "시산표" in wb.sheetnames else wb.worksheets[0]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    for i, r in enumerate(rows[:15]):
        h = [reader._norm(v) for v in r]
        if "코드" in h and "수정후" in h:
            c = {k: h.index(k) for k in ("코드", "계정명", "회사제시", "수정차변", "수정대변", "수정후") if k in h}
            out = []
            for row in rows[i + 1:]:
                code = _code(reader._cell(row, c["코드"]))
                if code is None:
                    continue
                v = reader._num(reader._cell(row, c["수정후"]))
                if v is None:      # 엑셀이 계산하지 않은 수식: 회사제시 + 수정 차변 − 수정 대변
                    v = sum((reader._num(reader._cell(row, c.get(k))) or 0) * s
                            for k, s in (("회사제시", 1), ("수정차변", 1), ("수정대변", -1)))
                out.append((code, reader._norm(reader._cell(row, c.get("계정명"))), v))
            return _source("last_year", _rel(d, p), out, ws.title)
    return None


def _read_8100b(d: Path) -> dict | None:
    from .start import read_prior_workpapers

    paths = _received(d, PRIOR_WP)
    if not paths:
        return None
    pw = read_prior_workpapers(d, paths)
    accs = pw["wtb"]["accounts"]
    if not accs:
        return None
    return _source("predecessor", pw["wtb"]["path"], [(_code(a["code"]), a["name"], a["final"]) for a in accs])


def _read_prior_tb(d: Path) -> dict | None:
    for rel in _received(d, PRIOR_TB):
        if Path(rel).suffix.lower() not in reader.EXCEL:
            continue
        got = reader.tb_rows(Path(d) / rel)
        if got:
            s = _source("prior_tb", rel, [(_code(r["code"]), r["name"], r["amount"]) for r in got["rows"]], got["sheet"])
            if s:
                s["column"] = got["column"]
                return s
    return None


def _read_comparative(d: Path) -> dict | None:
    found = figures._tb_file(d)
    if not found:
        return None
    path = found[0]
    got = reader.prior_columns(path)
    if not got:
        return None
    s = _source("comparative", Path(path).relative_to(d).as_posix(),
                [(_code(r["code"]), r["name"], r["amount"]) for r in got["rows"]], got["sheet"])
    if s:
        s["column"] = got["column"]
    return s


def sources(d: Path) -> list[dict]:
    """계정 단위 원천 1~3. 순서가 우선순위다."""
    d = Path(d)
    out = []
    prev = prior_engagement(d)
    for fn in ((lambda: _read_8800(d, prev)) if prev else None, lambda: _read_8100b(d), lambda: _read_prior_tb(d),
               lambda: _read_comparative(d)):
        if fn is None:
            continue
        s = fn()
        if s:
            out.append(s)
    return out


# ---------------------------------------------------------------- 장부(원천 4)
def book(d: Path, tb: list[dict]) -> dict[str, tuple[int, str]]:
    """분석(analysis.data)이 없을 때 쓰는 원천 4. {코드: (금액, 출처)}. 분개장은 계정코드 열이 있는 것만 읽는다."""
    d = Path(d)
    case = _case(d)
    try:
        fy_end = date.fromisoformat(case["fiscal_year_end"])
    except (KeyError, ValueError, TypeError):
        return {}
    pe = _prior_end(d)
    fy_start = date.fromordinal(pe.toordinal() + 1)
    p_start = date.fromordinal(_year_before(pe).toordinal() + 1)
    from .analysis import defs as adefs

    keys = {a["code"] for a in tb}
    cur, pri, months = {}, {}, set()
    plan = figures.journal_plan(d)                  # 겹치는 분개장은 빼거나 같은 줄을 거른다(2026-10-18)
    dd = figures.JournalDedupe(plan)
    for f in plan:
        if not f["use"]:
            continue
        for _, dt, dr, cr, _, vno, code in reader._journal(d / f["path"])[1]:
            k = figures.journal_line_key(dt, vno, code, dr, cr)
            if not dd.keep(f["path"], k, bool(f["dedupe"])):
                continue
            dd.seen(k)
            c = adefs.align_code(_code(code), keys)
            if c is None:
                continue
            if fy_start <= dt <= fy_end:
                cur[c] = cur.get(c, 0) + dr - cr
            elif p_start <= dt <= pe:
                pri[c] = pri.get(c, 0) + dr - cr
                months.add((dt.year, dt.month))
    want = set()
    y, m = p_start.year, p_start.month
    for _ in range(12):
        want.add((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    full = want <= months
    out = {}
    for a in tb:
        c, sec = a["code"], a.get("section")
        if sec in BS_SECTIONS and cur:
            out[c] = (reader._int(a["cur"] - cur.get(c, 0)), BASIS["book_bs"])
        elif sec in PL_SECTIONS and pri and full:
            out[c] = (reader._int(pri.get(c, 0)), BASIS["book_pl"])
    return out


def current_accounts(d: Path) -> list[dict]:
    """지금 회사제시 시산표 계정(3자리 코드)과 구분. 분석 없이 원천 분포를 셀 때 쓴다."""
    from .analysis import defs

    out: dict[str, dict] = {}
    for r in figures.current_tb(d):
        code = defs.code_key(r["code"]) or str(r["code"])
        if code in out:
            out[code]["cur"] += r["amount"]
            continue
        m = defs.map_account(code, r["name"])
        out[code] = {"code": code, "name": r["name"], "line": r["line"], "cur": r["amount"], "section": m["section"]}
    return list(out.values())


def _aligned(s: dict, keys: set[str]) -> dict:
    """원천의 계정코드를 시산표 열쇠에 맞춘 사본(자리 수가 다른 자료, R8)."""
    from .analysis import defs

    vals, names = {}, {}
    for c, v in s["values"].items():
        k = defs.align_code(c, keys)
        vals[k] = vals.get(k, 0) + v
        names.setdefault(k, s["names"].get(c, ""))
    return dict(s, values=vals, names=names)


# ---------------------------------------------------------------- 계정 단위 전기 잔액
def account_prior(d: Path, tb: list[dict] | None = None, book_values: dict | None = None, srcs: list[dict] | None = None) -> dict:
    """계정마다 전기 잔액을 원천 우선순위로 정한다.

    tb: [{code, name, section, cur}] (없으면 지금 시산표), book_values: 원천 4 {코드: (금액, 출처)} (없으면 여기서 계산).
    돌려주는 것
      accounts      {코드: {prior, basis, audited, sources: [{basis, amount, audited, absent}]}}
      extra         원천 1~3 에만 있는 계정(당기 시산표 · 분개장에 없음) [{code, name, prior, basis, audited}]
      basis_counts  {출처: 계정 수} (당기 시산표 계정 기준, 원천이 없으면 '없음')
      missing_pl    손익 계정 가운데 계정 단위 전기 원천이 없는 것이 있음
      sources       원천 1~3 요약 [{basis, path, audited, accounts, complete, pl, note}]
      conflicts     원천 사이 차이 [{code, chosen, chosen_basis, other, other_basis, diff}]
    """
    d = Path(d)
    tb = tb if tb is not None else current_accounts(d)
    srcs = sources(d) if srcs is None else srcs
    bk = book(d, tb) if book_values is None else book_values
    accounts, counts, conflicts = {}, {}, []
    missing_pl = False
    have = set()
    keys = {a["code"] for a in tb}
    srcs = [_aligned(s, keys) for s in srcs]
    for a in tb:
        code, sec = a["code"], a.get("section")
        have.add(code)
        cands = []
        for s in srcs:
            if sec in PL_SECTIONS and not s["pl"]:
                continue
            if code in s["values"]:
                cands.append({"basis": s["basis"], "amount": s["values"][code], "audited": s["audited"], "absent": False})
            elif s["complete"]:
                # 합계가 맞는 시산표에 계정이 없으면 그 원천에서는 잔액 0 이다
                cands.append({"basis": s["basis"], "amount": 0, "audited": s["audited"], "absent": True})
        if code in bk and bk[code][0] is not None:
            cands.append({"basis": bk[code][1], "amount": bk[code][0], "audited": False, "absent": False})
        if sec not in BS_SECTIONS + PL_SECTIONS:
            cands = []
        top = cands[0] if cands else None
        accounts[code] = {"prior": top["amount"] if top else None, "basis": top["basis"] if top else None,
                          "audited": top["audited"] if top else None, "sources": cands}
        counts[top["basis"] if top else NONE] = counts.get(top["basis"] if top else NONE, 0) + 1
        if sec in PL_SECTIONS and top is None:
            missing_pl = True
        for c in cands[1:]:
            if c["amount"] != top["amount"]:
                conflicts.append({"code": code, "name": a.get("name"), "chosen": top["amount"], "chosen_basis": top["basis"],
                                  "other": c["amount"], "other_basis": c["basis"] + (" (계정 없음)" if c["absent"] else ""),
                                  "diff": c["amount"] - top["amount"]})
    extra = []
    if srcs:
        s = srcs[0]
        for code, v in s["values"].items():
            if code in have or not v:
                continue
            extra.append({"code": code, "name": s["names"].get(code) or code, "prior": v, "basis": s["basis"], "audited": s["audited"]})
    return {"accounts": accounts, "extra": extra, "basis_counts": counts, "missing_pl": missing_pl,
            "sources": [{"basis": s["basis"], "path": s["path"], "audited": s["audited"], "accounts": len(s["values"]),
                         "complete": s["complete"], "pl": s["pl"], "note": s["note"], **({"column": s["column"]} if s.get("column") else {})}
                        for s in srcs],
            "conflicts": conflicts}


# ---------------------------------------------------------------- 전기 판단 정보
def _prev_unadjusted(d: Path, prev: Path) -> list[dict]:
    from .report import misstatement
    from .report.base import issues as prev_issues, wp_file

    p = wp_file(prev, "8100A")
    if p is None:
        return []
    try:
        st = misstatement.read(prev)
        pl = misstatement.plan(prev, figures.current_tb(prev), st)
    except Exception:  # noqa: BLE001  깨진 작년 조서면 미수정 목록을 비운다(경고는 collect 가 남긴다)
        return []
    titles = {i["no"]: i for i in prev_issues(prev)}
    by: dict = {}
    for e in pl["unadjusted"]:
        by.setdefault(e["issue"], []).append(e)
    eff = misstatement.effects(pl["unadjusted"])
    out = []
    for n, es in sorted(by.items(), key=lambda x: (x[0] is None, x[0])):
        i = titles.get(n, {})
        why = st["decisions"].get(n, (None, None))[1]
        e = eff.get(n, {})
        out.append({"issue": n, "text": i.get("title") or f"이슈 {n}", "account": i.get("account"),
                    "amount": sum(x["dr"] for x in es), "pl_effect": e.get("손익·법인세 전", 0) + e.get("법인세비용", 0),
                    "reason": why, "entries": [{"code": x["code"], "name": x.get("name"), "dr": x["dr"], "cr": x["cr"]} for x in es],
                    "source": _rel(d, p)})
    return out


def _prev_opinion(d: Path, prev: Path) -> dict | None:
    from .report import opinion as op_mod

    try:
        o = op_mod.read(prev)
    except Exception:  # noqa: BLE001
        return None
    rd = None
    try:
        from .close import read_8700_dates

        x = read_8700_dates(prev).get("report_date")
        rd = x.isoformat() if x else None
    except Exception:  # noqa: BLE001
        rd = None
    if rd is None:
        arch = re.match(r"(\d{4}-\d{2}-\d{2})", prev.name) if prev.parent.name == "_보관" else None
        rd = arch.group(1) if arch else None
    if o.get("kind") is None and not o.get("path"):
        return None
    return {"kind": o.get("kind"), "auditor": "작년 우리 감사", "report_date": rd, "emphasis": bool(o.get("emphasis")),
            "going_concern": bool(o.get("going_concern")), "emphasis_text": o.get("emphasis"), "going_concern_text": o.get("going_concern"),
            "other_text": o.get("other"), "source": _rel(d, prev / o["path"]) if o.get("path") else None}


def _prev_materiality(d: Path, prev: Path) -> dict | None:
    from .report import misstatement
    from .report.base import prev_inputs, wp_file

    m = misstatement.materiality(prev, defaults=False)        # 전기 확정값은 감사인 값만(기본값 제외)
    if m.get("overall") is None and m.get("performance") is None:
        return None
    p = wp_file(prev, "2700A")
    keep = prev_inputs(p) if p else {}
    basis = {}
    for k, name in (("overall", "재무제표 전체 중요성"), ("performance", "수행중요성"), ("trivial", "명백히 사소한 금액")):
        v = next((val for (s, lab), val in keep.items() if lab == f"{name} 기준·백분율"), None)
        basis[k] = str(v).strip() if v not in (None, "") else None
    return {k: ({"amount": int(m[k]), "basis": basis[k]} if m.get(k) is not None else None) for k in ("overall", "performance", "trivial")} | {
        "source": _rel(d, prev / m["source"]) if m.get("source") else None}


def _wp_materiality(pw: dict | None) -> dict | None:
    mat = (pw or {}).get("materiality") or {}
    if not mat:
        return None
    out = {k: ({"amount": v["amount"], "basis": v.get("basis")} if (v := mat.get(k)) else None) for k in ("overall", "performance", "trivial")}
    out["source"] = next((v["path"] for v in mat.values() if v.get("path")), None)
    return out


def _report_opinion(pr: dict | None, dp: dict | None) -> dict | None:
    if pr:
        return {"kind": pr.get("opinion"), "auditor": pr.get("auditor"), "report_date": pr.get("report_date"),
                "emphasis": pr.get("emphasis"), "going_concern": pr.get("going_concern"), "source": pr.get("path")}
    if dp:
        return {"kind": dp.get("opinion"), "auditor": dp.get("auditor"), "report_date": dp.get("report_date"),
                "emphasis": dp.get("emphasis"), "going_concern": dp.get("going_concern"), "source": dp.get("path")}
    return None


def _carry_from_opinion(op: dict | None) -> list[dict]:
    if not op:
        return []
    src = op.get("source") or "전기 감사보고서"
    out = []
    if op.get("kind") and op["kind"] != "적정":
        out.append({"what": "전기 감사의견 변형", "text": f"전기 감사의견 {op['kind']}. 변형을 초래한 사항이 당기 위험 평가에 미치는 영향을 봅니다(감사기준서 510 문단 9)",
                    "source": src})
    if op.get("going_concern"):
        out.append({"what": "계속기업 관련 중요한 불확실성", "text": op.get("going_concern_text") or "전기 감사보고서에 계속기업 관련 중요한 불확실성 단락이 있음",
                    "source": src})
    if op.get("emphasis"):
        out.append({"what": "강조사항", "text": op.get("emphasis_text") or "전기 감사보고서에 강조사항 단락이 있음", "source": src})
    return out


def _gathered(d: Path) -> dict:
    """시작자료.json 이 있으면 그 판독(전기 감사보고서 · DART · 전기조서)을 쓰고, 없으면 전기조서만 다시 읽는다."""
    data = _read(Path(d) / "_색인" / "시작자료.json", {}) or {}
    if data.get("prior_workpapers") is None and _received(d, PRIOR_WP):
        from .start import read_prior_workpapers

        data = dict(data, prior_workpapers=read_prior_workpapers(d, _received(d, PRIOR_WP)))
    return data


def collect(d: Path, gathered: dict | None = None, book_values: dict | None = None, accounts: bool = True) -> dict:
    """전기 판단 정보(시작자료.json prior_source). gathered 는 start.gather 결과(prior_report · dart · prior_workpapers).
    accounts=False 면 계정 단위 원천 분포(분개장을 읽는다)를 건너뛴다(account_prior 는 None)."""
    d = Path(d)
    g = gathered if gathered is not None else _gathered(d)
    sit = situation(d)
    prev = prior_engagement(d)
    base = prior_year_dir(d)
    warnings = []
    pr, dp, pw = g.get("prior_report"), g.get("dart"), g.get("prior_workpapers")
    out = {"situation": sit, "prior_audited": None, "audited_basis": None,
           "engagement_dir": _rel(d, prev) if prev else None, "archived": bool(prev and base and prev != base),
           "materiality": None, "unadjusted": [], "opinion": None, "carry": [], "inputs": {}, "account_prior": None,
           "workpapers": _received(d, PRIOR_WP), "warnings": warnings}
    if prev:
        has_8820 = (prev / "8000 감사완결").is_dir() and any((prev / "8000 감사완결").glob("8820_*"))
        made = out["archived"] or has_8820
        out["prior_audited"] = True if made else None
        out["audited_basis"] = "작년 감사 폴더(보관본)" if out["archived"] else ("작년 감사 폴더 8820" if made else None)
        if not made:
            warnings.append(f"작년 감사 폴더에 감사보고서(8820)와 보관본이 없습니다: {out['engagement_dir']}")
        out["materiality"] = _prev_materiality(d, prev)
        out["unadjusted"] = _prev_unadjusted(d, prev)
        out["opinion"] = _prev_opinion(d, prev)
        for i in (_read(prev / "_색인" / "이슈.json", {"issues": []}) or {}).get("issues", []):
            src = _rel(d, prev / "_색인" / "이슈.json")
            if i.get("close_method") == "미수정 왜곡표시":
                out["carry"].append({"what": "미수정 왜곡표시", "text": f"작년 이슈 {i['no']}: {i['title']}", "source": src})
            elif i.get("status") == "열림":
                out["carry"].append({"what": "열린 채 종결된 이슈", "text": f"작년 이슈 {i['no']}: {i['title']}", "source": src})
        out["carry"] += _carry_from_opinion(out["opinion"])
        people = prev / "1000 감사계약" / "감사반인원.xlsx"
        if people.is_file():
            out["inputs"]["1300 감사반 인원"] = _rel(d, people)
    else:
        out["materiality"] = _wp_materiality(pw)
        out["unadjusted"] = [{"issue": None, "text": u["text"], "account": u.get("account"), "amount": u["amount"], "pl_effect": None,
                              "reason": u.get("reason"), "entries": [], "source": u["path"]} for u in (pw or {}).get("unadjusted", [])]
        out["opinion"] = _report_opinion(pr, dp)
        out["carry"] = [{"what": "미수정 왜곡표시", "text": f"전기 미수정 왜곡: {u['text']}", "source": u["source"]} for u in out["unadjusted"]]
        out["carry"] += _carry_from_opinion(out["opinion"])
        if pr or dp:
            out["prior_audited"] = True
            out["audited_basis"] = "전기 감사보고서" if pr and pr.get("format") != "DART" else "DART 감사보고서"
    if base and prev is None:
        warnings.append(f"작년 감사 폴더를 읽지 못했습니다: {_rel(d, base)}")
    if accounts:
        ap = account_prior(d, book_values=book_values)
        out["account_prior"] = {"basis_counts": ap["basis_counts"], "missing_pl": ap["missing_pl"], "sources": ap["sources"],
                                "conflicts": len(ap["conflicts"])}
    return out


def info(d: Path, accounts: bool = False) -> dict:
    """시작자료.json 의 prior_source. 없으면 계산만 한다(쓰지 않는다). 조서가 이어받을 값(중요성 · 미수정 · 이어받을 사항)을 읽을 때 쓴다."""
    data = _read(Path(d) / "_색인" / "시작자료.json", {}) or {}
    return data.get("prior_source") or collect(d, accounts=accounts)


def label(ps: dict | None) -> str:
    """현황 한 줄: '전기 자료: 작년 감사 폴더(../2024 기말감사/_보관/2025-03-14)'."""
    if not ps:
        return "전기 자료: 판독 전(python -m audit start)"
    sit = ps.get("situation")
    if sit == SITUATIONS[0]:
        return f"전기 자료: 작년 감사 폴더({ps.get('engagement_dir')})"
    if sit == SITUATIONS[1]:
        return f"전기 자료: 전임 전기조서({', '.join(Path(p).name for p in ps.get('workpapers') or []) or '받은자료 전기/'})"
    op = ps.get("opinion") or {}
    return f"전기 자료: 전기조서 없음({'전기 감사보고서 · DART' if op.get('source') else '전기 감사보고서도 없음'})"

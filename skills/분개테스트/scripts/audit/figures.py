"""감사 폴더의 숫자 원천을 한 곳에서 읽는다(R2~R5 공통, 2026-10-05).

여러 단계(시작 자료 · 분석적 검토 · 조서 · 보고)가 같은 숫자를 각자 다르게 읽지 않게 이 모듈만 쓴다.
숫자는 엔진이 파일에서 읽은 값이다. AI 판독 값은 여기로 들어오지 않는다.

  current_tb(d)        지금 회사제시 판 시산표 행 [{code, name, line, amount, version, path}] (차변 +, 대변 −)
  current_lines(d)     표시과목별 합 {표시과목: 금액} (차변 +, 대변 −). 표시과목 열이 없으면 계정명
  prior(d)             전기 재무제표 {"as_of", "bs": {과목: 금액}, "is": {과목: 금액}, "path"} (표시 금액, 부호는 재무제표 그대로)
                       작년 감사 폴더(보관본 우선) 8810 날인재무제표 → 받은자료 엑셀 → _색인/시작자료.json 의 prior_fs
                       (R2 시작 자료가 전기조서 · DART · 감사보고서에서 읽은 것) 순(R7, 2026-10-06)
  received(d, kind)    문서목록에서 그 종류의 받은자료 경로들(상대경로)
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from . import folder, reader

PRIOR_FS = "전기 재무제표"


def _docs(d: Path) -> dict:
    p = Path(d) / "_색인" / "문서목록.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {"files": []}


def received(d: Path, kind: str) -> list[str]:
    return [f["path"] for f in _docs(d)["files"] if f.get("role") == "received" and f.get("kind") == kind]


def _tb_files(d: Path) -> list[tuple[Path, str | None]]:
    """회사제시 엑셀 받은자료. 판이 높은 것부터."""
    found = []
    for f in _docs(d)["files"]:
        if f.get("role") == "received" and f.get("kind") == reader.TB and Path(f["path"]).suffix.lower() in reader.EXCEL:
            found.append((folder._version_key((f.get("version") or "v0")[1:]), f))
    found.sort(key=lambda x: x[0], reverse=True)
    return [(Path(d) / f["path"], f.get("version")) for _, f in found]


def _tb_file(d: Path) -> tuple[Path, str | None] | None:
    """지금 회사제시 판 시산표 파일. 시산표로 읽히는 파일이 있으면 그중 판이 가장 높은 것, 없으면 판이 가장 높은 파일."""
    files = _tb_files(d)
    for p, v in files:
        if reader.tb_table(p):
            return p, v
    return files[0] if files else None


_COA_CACHE: dict = {}


def coa(d: Path) -> tuple[dict | None, str | None]:
    """받은자료 가운데 계정과목표로 읽히는 첫 파일(reader.coa_table 결과, 상대 경로). 없으면 (None, None).
    이름에 「계정과목」 · 「계정코드」 · 「chart」 가 든 파일을 먼저 본다. 분개장은 보지 않는다(크고 계정과목표가 아니다).
    문서목록이 바뀌지 않았으면 다시 읽지 않는다."""
    d = Path(d)
    p = d / "_색인" / "문서목록.json"
    key = (str(d.resolve()), p.stat().st_mtime_ns if p.exists() else None)
    if key in _COA_CACHE:
        return _COA_CACHE[key]
    files = [f for f in _docs(d)["files"] if f.get("role") == "received" and f.get("kind") != reader.JOURNAL
             and Path(f["path"]).suffix.lower() in reader.EXCEL and (d / f["path"]).exists()]
    named = lambda f: any(k in Path(f["path"]).name.lower() for k in ("계정과목", "계정코드", "chart"))  # noqa: E731
    out = (None, None)
    for f in sorted(files, key=lambda f: not named(f)):
        got = reader.coa_table(d / f["path"])
        if got:
            out = (got, f["path"])
            break
    _COA_CACHE.clear()
    _COA_CACHE[key] = out
    return out


def current_tb(d: Path) -> list[dict]:
    """지금 회사제시 판 시산표 행. 시산표가 없고 계정 단위 재무제표(재무상태표 · 손익계산서)만 있으면 그것으로 만든다(R8).
    코드 없는 더존 합계잔액시산표는 받은자료 계정과목표의 이름으로 코드를 붙인다(없으면 계정 사전, 2026-10-18).
    그렇게 붙인 코드의 근거는 code_basis 에 남는다."""
    d = Path(d)
    files = _tb_files(d)
    chart = None
    for path, ver in files:
        got = reader.tb_table(path)
        if not got:
            continue
        if got["form"] == reader.TB_NO_CODE:
            chart = chart or coa(d)
            if chart[0]:
                got = reader.tb_table(path, coa=chart[0])
        rows = [{"code": str(r["code"]), "name": r["name"], "line": r["line"], "amount": r["amount"],
                 "version": ver, "path": path.relative_to(d).as_posix(), **({"subgroup": r["subgroup"]} if r.get("subgroup") else {}),
                 **({"code_basis": r["code_basis"] + (f" {chart[1]}" if r["code_basis"] == "계정과목표" and chart else "")}
                    if r.get("code_basis") else {})}
                for r in got["rows"]]
        return confirmed_lines(d, _with_lines(rows))
    return confirmed_lines(d, _with_lines(fs_tb(d, files)))


CONFIRMED = "3650A 감사인 확정"
_CONF_CACHE: dict = {}


def confirmations(d: Path) -> dict:
    """3650A 「계정 매핑」 시트의 감사인 확정 {"line": {코드: 표시과목}, "area": {코드: 영역}}(2026-10-07).
    파일이 바뀌지 않았으면 다시 읽지 않는다."""
    from . import wp

    p = wp.path_for(Path(d), "3650A")
    if not p.exists():
        return {"line": {}, "area": {}}
    key = (str(p), p.stat().st_mtime_ns)
    if key not in _CONF_CACHE:
        out = {"line": {}, "area": {}}
        try:
            for x in wp.readback(p):
                if x["sheet"] != "계정 매핑" or wp._blank(x["value"]):
                    continue
                for kind, head in (("line", "표시과목 확정: "), ("area", "영역 확정: ")):
                    if str(x["label"]).startswith(head):
                        code = str(x["label"])[len(head):].split(" ", 1)[0]
                        out[kind][code] = str(x["value"]).strip()
        except Exception:  # noqa: BLE001  회계사가 다른 꼴로 저장한 파일은 확정 없음으로 본다
            pass
        _CONF_CACHE.clear()
        _CONF_CACHE[key] = out
    return _CONF_CACHE[key]


def confirmed_lines(d: Path, rows: list[dict]) -> list[dict]:
    """감사인이 3650A 에서 확정한 표시과목으로 바꾼다. 엔진 제안은 engine_line 에 남긴다."""
    lines = confirmations(d)["line"]
    for r in rows:
        v = lines.get(str(r["code"]))
        if v and v != r.get("line"):
            r["engine_line"], r["line"], r["line_basis"] = r.get("line"), v, CONFIRMED
        elif v:
            r["line_basis"] = CONFIRMED
    return rows


def _with_lines(rows: list[dict]) -> list[dict]:
    """표시과목. 시산표의 표시과목 열이 먼저, 없으면 accounts.yaml display_lines(일반기업회계기준 표시 과목), 그래도 없으면 계정명.
    line_basis 에 근거를 남긴다(R8, 진단 #5)."""
    from .analysis import defs

    names = {defs.norm(r["name"]) for r in rows}
    for r in rows:
        if r.get("line"):
            r["line_basis"] = "표시과목 열"
            continue
        line, basis = defs.display_line(r["name"], names)
        r["line"], r["line_basis"] = (line, basis) if line else (r["name"], "계정명")
    return rows


# 계정 단위 재무제표에서 코드를 정하지 못한 계정에 붙이는 범위 코드(report_fs.yaml 묶음 범위의 끝 3자리 + 순번 2자리).
# 코드 없는 더존 시산표도 같은 표를 쓴다(reader.assign_codes)
FS_PLACEHOLDER = reader.FS_PLACEHOLDER
FS_RANGES = reader.FS_RANGES
FS_SOURCE = "계정 단위 재무제표"


JOURNAL_FOLDER = "분개장"


def journal_plan(d: Path, exts: tuple[str, ...] | None = None) -> list[dict]:
    """분개장 받은자료를 어떤 순서로 어떻게 합칠지(2026-10-18, 분개장 겹침).
    [{"path", "range": (첫 날, 끝 날) | None, "use": bool, "dedupe": [앞 파일들], "why": 문장 | None}] 읽을 순서대로.

    순서: 받은자료 「분개장」 폴더의 파일이 먼저, 그다음 덮는 기간이 긴 파일, 그다음 경로.
    앞 파일들의 기간 안에 기간이 다 들어가는 파일(1년 분개장과 함께 받은 권한 검토용 분기 분개장)은 합치지 않는다(use False).
    기간이 일부만 겹치면 합치되, 앞 파일에 이미 있는 줄(같은 날짜 · 전표번호 · 계정 · 차변 · 대변)은 뺀다(dedupe).
    같은 내용(sha256)의 사본은 하나만 쓴다. 기간은 색인 엔진 사실(index read)에서 읽고, 없으면 파일을 읽는다."""
    d = Path(d)
    exts = exts or reader.EXCEL
    files, seen = [], {}
    for f in _docs(d)["files"]:
        if f.get("role") != "received" or f.get("kind") != reader.JOURNAL or Path(f["path"]).suffix.lower() not in exts:
            continue
        rng = reader.journal_range(f.get("facts") or [])
        if rng is None and (d / f["path"]).exists() and Path(f["path"]).suffix.lower() in reader.READABLE:
            lines = reader.journal_rows(d / f["path"])
            rng = (min(x["date"] for x in lines), max(x["date"] for x in lines)) if lines else None
        x = {"path": f["path"], "range": rng, "use": True, "dedupe": [], "why": None,
             "folder": f.get("folder_kind") == JOURNAL_FOLDER, "sha": f.get("sha256")}
        if x["sha"] and x["sha"] in seen:
            x.update(use=False, why=f"{seen[x['sha']]} 와 같은 파일이라 한 번만 셈")
        elif x["sha"]:
            seen[x["sha"]] = f["path"]
        files.append(x)
    span = lambda r: (r[1] - r[0]).days if r else -1  # noqa: E731
    files.sort(key=lambda x: (not x["folder"], -span(x["range"]), x["path"]))
    used: list[dict] = []
    for x in files:
        if not x["use"]:
            continue
        r = x["range"]
        over = [u for u in used if r and u["range"] and u["range"][0] <= r[1] and r[0] <= u["range"][1]]
        if over:
            names = " · ".join(u["path"] for u in over)
            if _covered(r, [u["range"] for u in over]):
                x.update(use=False, why=f"기간 {r[0]} ~ {r[1]} 이 {names} 의 기간 안에 들어 있어 합치지 않음(부분 분개장)")
                continue
            x.update(dedupe=[u["path"] for u in over],
                     why=f"기간 {r[0]} ~ {r[1]} 이 {names} 와 일부 겹쳐, 앞 파일에 있는 같은 줄(날짜 · 전표번호 · 계정 · 금액)은 빼고 합침")
        used.append(x)
    for x in files:
        x.pop("sha", None)
    return files


def _covered(r, ranges) -> bool:
    """기간 r 이 ranges 의 합집합 안에 다 들어가는지."""
    from datetime import timedelta

    start = r[0]
    for a, b in sorted(ranges):
        if a > start:
            return False
        if b >= start:
            start = b + timedelta(days=1)
        if start > r[1]:
            return True
    return start > r[1]


def journal_line_key(dt, voucher, code, dr, cr) -> tuple:
    """분개장 줄 겹침 판정 열쇠(journal_plan dedupe)."""
    return (dt, str(voucher or ""), str(code or ""), int(round(dr or 0)), int(round(cr or 0)))


class JournalDedupe:
    """journal_plan 의 dedupe 파일에서 앞 파일에 이미 있는 줄을 거른다. 앞 파일 줄은 seen 으로 넣고, 뒤 파일 줄은 keep 으로 묻는다.
    같은 열쇠가 여러 번이면 그 횟수만큼만 뺀다."""

    def __init__(self, plan: list[dict]):
        self.need = any(x["dedupe"] for x in plan)
        self.count: dict[tuple, int] = {}
        self.dropped: dict[str, int] = {}

    def seen(self, key: tuple) -> None:
        if self.need:
            self.count[key] = self.count.get(key, 0) + 1

    def keep(self, path: str, key: tuple, dedupe: bool) -> bool:
        if not dedupe:
            return True
        n = self.count.get(key, 0)
        if n:
            self.count[key] = n - 1
            self.dropped[path] = self.dropped.get(path, 0) + 1
            return False
        return True


def journal_codes(d: Path) -> dict[str, str]:
    """받은자료 분개장의 계정명 → 계정코드(가장 많이 쓰인 코드)."""
    count: dict[str, dict[str, int]] = {}
    for rel in received(d, reader.JOURNAL):
        p = Path(d) / rel
        if p.suffix.lower() not in reader.EXCEL:
            continue
        for x in reader.journal_rows(p):
            if x["code"] and x["name"]:
                c = count.setdefault(x["name"], {})
                c[x["code"]] = c.get(x["code"], 0) + 1
    return {n: max(c, key=c.get) for n, c in count.items()}


def fs_tb(d: Path, files: list | None = None) -> list[dict]:
    """계정 단위 재무제표(코드 없음)로 만든 시산표 행. 같은 판의 재무상태표 · 손익계산서 파일을 합친다.
    코드: 1 분개장에서 같은 계정명이 쓴 코드, 2 계정 사전(accounts.yaml)의 코드 가운데 재무제표 위치 범위 안의 것,
    3 둘 다 없으면 재무제표 위치 범위 코드(FS_PLACEHOLDER + 순번, code_basis 에 '재무제표 위치' 로 남겨 회계사 확인을 받는다).
    재무상태표 이익잉여금(결손금)에는 당기순이익이 들어 있어 손익계산서와 합치면 두 번 셈이 된다. 손익계산서가 있으면
    이익잉여금 계정에서 당기순이익을 빼 마감 전 시산표로 맞춘다."""
    from .analysis import defs

    d = Path(d)
    files = _tb_files(d) if files is None else files
    if not files:
        return []
    ver = files[0][1]
    same = [(p, v) for p, v in files if v == ver]
    sheets = []
    for p, _ in same:
        got = reader.fs_accounts(p)
        if got:
            sheets += [dict(s, path=p) for s in got["sheets"]]
    if not sheets:
        return []
    by_name = journal_codes(d)
    accs = defs.load("accounts")["accounts"]
    seq: dict[str, int] = {}
    used: dict[str, str] = {}
    out = []
    for s in sheets:
        for it in s["items"]:
            lo, hi = FS_RANGES[it["group"]]
            code, basis = by_name.get(it["name"]), "분개장 계정코드"
            if code and not lo <= int(code[:3]) <= hi:
                code = None
            if code is None:
                m = defs.map_account(None, it["name"])
                spec = next((a for a in accs if a["id"] == m["id"]), None) if m["id"] else None
                code = next((c for c in (spec or {}).get("codes", []) if lo <= int(c) <= hi), None)
                basis = "계정 사전"
            if code is None:
                base = FS_PLACEHOLDER[it["group"]]
                seq[base] = seq.get(base, 0) + 1
                code, basis = f"{base}{seq[base]:02d}", "재무제표 위치(회계사 확인 필요)"
            elif basis == "계정 사전" and used.get(code, it["name"]) != it["name"]:
                # 사전 코드가 다른 계정에 이미 쓰였으면 앞 3자리는 두고 순번을 붙여 계정을 따로 둔다
                seq[code] = seq.get(code, 0) + 1
                code = f"{code}{seq[code]:02d}"
            used.setdefault(code, it["name"])
            out.append({"code": code, "name": it["name"], "line": "", "amount": it["amount"], "version": ver,
                        "path": s["path"].relative_to(d).as_posix(), "source": FS_SOURCE, "code_basis": basis,
                        "fs_group": it["group"]})
    pl = [r for r in out if r["fs_group"] in ("매출액", "매출원가", "판매비와관리비", "영업외수익", "영업외비용", "법인세비용")]
    if pl:
        ni = -sum(r["amount"] for r in pl)
        eq = [r for r in out if r["fs_group"] == "자본" and ("잉여금" in r["name"] or "결손금" in r["name"])]
        re_row = next((r for r in eq if any(k in r["name"] for k in ("미처분", "미처리", "이월"))), eq[-1] if eq else None)
        if re_row is not None:
            re_row["amount"] = reader._int(re_row["amount"] + ni)
            re_row["closing_adjusted"] = True
    return out


def tb_absent_reason(d: Path) -> str:
    """시산표 행이 없을 때 까닭 한 줄. 파일이 없는 것과 파일은 있는데 판독하지 못한 것을 가른다(R8)."""
    found = _tb_file(d)
    if not found:
        files = [f for f in _docs(d)["files"] if f.get("role") == "received" and f.get("kind") == reader.TB]
        bad = next((f for f in files if f.get("read_error")), None)
        if bad:
            return f"회사제시 시산표 {bad['path']} {bad['read_error']}"
        return "회사제시 시산표(엑셀)가 받은자료에 없습니다"
    path, _ = found
    why = reader.read_failure(path, reader.TB) or reader.READ_FAIL
    return f"회사제시 시산표 {path.relative_to(d).as_posix()} {why}"


def current_lines(d: Path) -> dict[str, int]:
    out: dict[str, int] = {}
    for r in current_tb(d):
        out[r["line"]] = out.get(r["line"], 0) + r["amount"]
    return out


def _prior_sheet(ws) -> tuple[dict, str | None]:
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows[:10]):
        cells = [reader._norm(v) for v in r]
        if "과목" in cells:
            col = cells.index("과목")
            amt_col = next((j for j in range(col + 1, len(cells)) if "전기" in cells[j]), None) or                 next((j for j in range(col + 1, len(cells)) if cells[j]), None)
            if amt_col is None:
                return {}, None
            out = {}
            for row in rows[i + 1:]:
                name = reader._norm(reader._cell(row, col))
                v = reader._num(reader._cell(row, amt_col))
                if name and v is not None:
                    out[name] = reader._int(out.get(name, 0) + v)     # 같은 이름(감가상각누계액 · 대손충당금 둘)은 더한다(재무제표 꼴과 같게)
            return out, cells[amt_col]
    return {}, None


def prior(d: Path) -> dict | None:
    """전기 재무제표 과목별 금액. 작년 감사 폴더의 8810 날인재무제표(우리가 감사한 최종본)가 먼저, 다음은 받은자료
    '전기 재무제표' 엑셀, 없으면 시작 자료(R2)가 전기조서 8100B · DART · 전기 감사보고서에서 읽어 `_색인/시작자료.json` 에
    둔 전기 재무제표를 쓴다. 모두 없으면 None."""
    return drop_group_lines(prior_last_year(d) or prior_excel(d) or _prior_from_start(d))


# 전기 재무제표의 묶음 줄 후보(일반기업회계기준 제2장 중분류 · 손익 구분 이름). 이름만으로 정하지 않고, 바로 아래 과목들의 합과
# 같을 때만 묶음 줄로 본다(합성 회사처럼 '유형자산' 이 한 과목인 재무제표도 있기 때문이다)
GROUP_CANDIDATES = ("당좌자산", "재고자산", "투자자산", "유형자산", "무형자산", "기타비유동자산", "자본금", "자본잉여금", "자본조정",
                    "기타포괄손익누계액", "이익잉여금", "결손금", "매출액", "매출원가", "판매비와관리비", "판매관리비", "영업수익",
                    "영업비용", "영업외수익", "영업외비용", "법인세비용", "법인세등")
CONTRA_WORDS = ("충당금", "누계액", "할인차금")


def group_lines(items: list[tuple[str, int]], totals: set[str]) -> set[str]:
    """[(과목, 금액)] 순서에서 묶음 줄(소계)인 과목. 후보 이름이고, 다음 후보 · 합계 줄 전까지의 과목 합(또는 차감 계정을 뺀 합)과
    금액이 같으면 묶음 줄이다(R8, 진단 #8)."""
    from .analysis import defs

    import re

    # 전임 감사인 WTB 의 「II.자본잉여금」 · 「V.이익잉여금(결손금)」 처럼 번호(로마 숫자 · 숫자)가 앞에 붙은 묶음 줄도 본다(2026-10-18)
    key = lambda x: re.sub(r"^(?:[IVX]+|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+|\d+)[.．]", "", defs.norm(x))  # noqa: E731
    cand = {defs.norm(x) for x in GROUP_CANDIDATES}
    totals = {defs.norm(x) for x in totals}
    out = set()
    for i, (lab, v) in enumerate(items):
        if key(lab) not in cand or v is None:
            continue
        follow = []
        for lab2, v2 in items[i + 1:]:
            k = key(lab2)
            if k in cand or k in totals:
                break
            follow.append((k, v2 or 0))
        if not follow:
            continue
        plain = sum(x for _, x in follow)
        net = sum(-x if any(w in k for w in CONTRA_WORDS) else x for k, x in follow)
        if abs(plain - v) < 1 or abs(net - v) < 1 or abs(abs(net) - abs(v)) < 1:
            out.add(lab)
    return out


# 시산표 묶음 줄 이름: 더존 합계잔액시산표의 「<< 자 산 >>」 「[ 유동자산 ]」 「< 재고자산 >」 꼴. 계정이 아니라 아래 계정의 합(합계 열 포함)이다
GROUP_ROW_RE = re.compile(r"^\s*(?:<<?|\[|【|《|〈)[^<>\[\]]*(?:>>?|\]|】|》|〉)\s*$")


def is_group_row(name) -> bool:
    """시산표 행 이름이 묶음(소계) 줄인가. 이런 줄이 계정으로 들어오면 영역을 여는 기준 금액이 두 번 세어지거나
    합계 열 금액이 잔액처럼 쓰인다(2026-10-18 실제 회사 자료: 「<재고자산>」 대변 합계로 재고 없는 회사에 4000E 가 열림)."""
    return bool(GROUP_ROW_RE.match(str(name or "")))


def drop_group_lines(pf: dict | None) -> dict | None:
    """전기 재무제표에서 묶음 줄을 빼고 pf["groups"] 에 남긴다. 묶음 줄을 과목으로 비교하면 당기 0 인 과목과 전기 합계를 두 번 더하는
    결함이 생긴다(진단 #8)."""
    if not pf:
        return pf
    from .report.base import defs as fs_defs

    totals = set(fs_defs().get("prior_totals", [])) | {"자산총계", "부채총계", "자본총계", "부채와자본총계", "부채및자본총계"}
    groups = {}
    for k in ("bs", "is"):
        items = list((pf.get(k) or {}).items())
        g = group_lines(items, totals)
        if g:
            pf[k] = {lab: v for lab, v in items if lab not in g}
            groups[k] = sorted(g)
    if groups:
        pf["groups"] = groups
    return pf


def prior_last_year(d: Path) -> dict | None:
    """작년 감사 폴더(audit.prior.prior_engagement, 보관본 우선)의 8810 날인재무제표 워드 당기 열. 없으면 None."""
    from . import prior as prior_mod

    try:
        prev = prior_mod.prior_engagement(d)
    except Exception:  # noqa: BLE001  작년 폴더를 못 읽으면 다음 원천으로
        return None
    if prev is None:
        return None
    from .report.base import read_docx, wp_file
    from .report.footing import fs_doc_values

    p = wp_file(prev, "8810", "8810_*.docx")
    if p is None:
        return None
    try:
        v = fs_doc_values(read_docx(p))
    except Exception:  # noqa: BLE001
        return None
    # 8810 표에는 묶음 줄(유동자산 · 영업외비용 등)도 있다. 작년 8800 「표시과목」 시트의 과목과 합계 줄만 쓴다
    keep = _units_8800(prev)
    if keep is None:
        return None
    from .report.base import defs as fs_defs, norm

    keep |= {norm(x) for x in fs_defs().get("prior_totals", [])}
    out = {"as_of": (json.loads((prev / "_색인" / "감사건.json").read_text(encoding="utf-8")) or {}).get("fiscal_year_end"),
           "bs": {k: x[0] for k, x in v["bs"].items() if x and x[0] is not None and k in keep},
           "is": {k: x[0] for k, x in v["is"].items() if x and x[0] is not None and k in keep},
           "path": prior_mod._rel(d, p)}
    return out if out["bs"] or out["is"] else None


def _units_8800(prev: Path) -> set[str] | None:
    """작년 8800 「표시과목」 시트의 과목 이름(공백 없앰). 시트가 없으면 None."""
    import openpyxl

    from .report.base import norm, wp_file

    p = wp_file(prev, "8800")
    if p is None:
        return None
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    try:
        if "표시과목" not in wb.sheetnames:
            return None
        rows = [tuple(r) for r in wb["표시과목"].iter_rows(values_only=True)]
    finally:
        wb.close()
    head = next((i for i, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].startswith("표시 과목")), None)
    if head is None:
        return None
    return {norm(r[0]) for r in rows[head + 1:] if r and isinstance(r[0], str) and r[0].strip() and (len(r) < 2 or r[1] is not None)}


def _prior_from_start(d: Path) -> dict | None:
    p = Path(d) / "_색인" / "시작자료.json"
    if not p.exists():
        return None
    fs = (json.loads(p.read_text(encoding="utf-8")) or {}).get("prior_fs")
    if not fs or not (fs.get("bs") or fs.get("is")):
        return None
    return {"as_of": fs.get("as_of"), "bs": dict(fs.get("bs") or {}), "is": dict(fs.get("is") or {}), "path": fs.get("source")}


def _prior_column(rows) -> tuple | None:
    """비교식 재무제표에서 전기 열. (머리줄 행, 과목 열, 전기 열들, 근거) 또는 None. 「(전)」 · 「전기」 열이 먼저, 없으면
    「제 n 기」 가운데 작은 n, 기수 표시가 없으면 첫 금액 열(전기 재무제표만 따로 받은 파일)."""
    h = reader._fs_header(rows)
    if h is None:
        return None
    hi, lab, cur_cols, pri_cols = h
    head = rows[hi]
    if pri_cols:
        return hi, lab, pri_cols, f"'{str(head[pri_cols[0]]).strip()}' 열"
    return hi, lab, cur_cols, f"'{str(head[cur_cols[0]]).strip()}' 열(기수 표시 없음)"


def _prior_fs_sheet(rows) -> dict | None:
    """재무제표 꼴 시트(R8, 진단 #10)의 전기 열. {"stmt", "values": {과목: 금액}, "basis"}.
    묶음 줄(Ⅰ. · (1) 제목)은 빼고 계정 줄 · 합계 줄만 둔다. 금액은 엔진 8810 과 같은 부호다: 자산의 차감 계정(대손충당금 ·
    감가상각누계액)과 자본의 차감 항목(결손금 · 자기주식)은 음수, 손실 계산 줄(영업손실 · 당기순손실)은 음수."""
    got = _prior_column(rows)
    if got is None:
        return None
    hi, lab, cols, basis = got
    seq = reader._fs_lines(rows[hi + 1:], lab, cols)
    if not seq:
        return None
    keys = {x[1] for x in seq}
    stmt = "bs" if ("자산총계" in keys or "자산" in keys) else "is"
    vals, pos = {}, {}
    for it in reader._fs_items(seq, stmt):
        g, sign = reader._fs_group(stmt, it)
        v = it["value"]
        if stmt == "bs" and sign < 0 and g in ("유동자산", "비유동자산"):
            v = -abs(v)
        if stmt == "bs" and g == "자본" and sign > 0:
            v = -abs(v)
        vals[it["name"]] = reader._int(vals.get(it["name"], 0) + v)     # 같은 이름(감가상각누계액 · 대손충당금 둘)은 더한다
        pos.setdefault(it["name"], g)
    for kind, key, v, _ in seq:
        if v is None:
            continue
        if kind == "total" or (kind == "h1" and key in ("유동자산", "비유동자산", "유동부채", "비유동부채")):
            vals.setdefault(key, reader._int(-abs(v) if key.endswith("손실") else v))
    return {"stmt": stmt, "values": vals, "basis": basis, "positions": pos}


def prior_excel(d: Path) -> dict | None:
    """받은자료 '전기 재무제표' 엑셀만 읽는다. 없으면 None.
    재무제표 꼴(과목 · 제 n (당)기 · 제 n-1 (전)기)이면 전기 열을 읽고 묶음 줄을 뺀다(R8). 시트 이름 BS · PL · IS 도 알아본다."""
    paths = [p for p in received(d, PRIOR_FS) if Path(p).suffix.lower() in reader.EXCEL]
    if not paths:
        return None
    import openpyxl

    p = Path(d) / paths[0]
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = {"as_of": None, "bs": {}, "is": {}, "path": paths[0]}
    try:
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            got = _prior_fs_sheet(rows) if _is_fs_layout(rows) else None
            if got:
                out[got["stmt"]].update(got["values"])
                out.setdefault("positions", {}).update(got["positions"])     # 과목의 재무제표 묶음(유동자산 · 판매비와관리비 등)
                out.setdefault("column", {})[got["stmt"]] = got["basis"]
                if got["stmt"] == "bs" and out["as_of"] is None:
                    out["as_of"] = got["basis"]
                continue
            vals, head = _prior_sheet(ws)
            if not vals:
                continue
            title = ws.title.strip().upper()
            pl = title in ("PL", "IS", "P&L") or "손익" in ws.title or "매출액" in vals or "영업수익" in vals
            key = "is" if pl and "자산총계" not in vals else "bs"
            out[key].update(vals)
            if key == "bs" and head:
                out["as_of"] = head
    finally:
        wb.close()
    return out


def _is_fs_layout(rows) -> bool:
    """제목 줄(Ⅰ. · (1)) 이나 기수 머리(제 n 기 · (당) · (전))가 있는 재무제표 꼴인지."""
    h = reader._fs_header(rows)
    if h is None:
        return False
    hi = h[0]
    head = " ".join(str(v) for v in rows[hi] if v is not None)
    if "(당)" in head.replace(" ", "") or "(전)" in head.replace(" ", "") or "기" in head and "제" in head:
        return True
    return any(isinstance(r[h[1]] if h[1] < len(r) else None, str) and reader.ROMAN_PREFIX.match(r[h[1]]) for r in rows[hi + 1:hi + 40])


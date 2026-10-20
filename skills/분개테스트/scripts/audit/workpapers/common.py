"""R4 실행 단계 공통: 정의 · 감사건 · 중요성 · 시산표 표시 · 영역 대응 · 분개장 전수 판독 · 저장(2026-10-05).

숫자는 figures(시산표 · 전기 재무제표 · 받은자료 경로)와 reader(엑셀 판독 도우미)에서만 읽는다.
조서 엑셀은 wp.Book 으로만 만들고, 다시 만들 때 회계사가 적은 기입란 값을 (시트, 항목) 으로 옮겨 둔다.
"""
from __future__ import annotations

import json
import math
import re
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import yaml

from .. import figures, folder, reader, wp

DEFS = folder.DEFS


def load_def(name: str) -> dict:
    from .. import params

    return params.overlay(name.removesuffix(".yaml"), folder.read_yaml(DEFS / name))   # 2700B 값(2026-10-07)


def idx(d: Path, name: str, default=None):
    p = Path(d) / "_색인" / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def put_idx(d: Path, name: str, data) -> None:
    folder._write(Path(d) / "_색인" / name, data)


def case(d: Path) -> dict:
    return idx(d, "감사건.json", {}) or {}


def fy_end(d: Path) -> date:
    return date.fromisoformat(case(d).get("fiscal_year_end"))


def fy_start(d: Path) -> date:
    """사업연도 시작일(감사건 fiscal_year_start 가 있으면 그것, R8)."""
    return folder.fiscal_start(case(d))


def won(x) -> str:
    if x is None:
        return "-"
    x = reader._int(float(x))
    return f"{x:,}" if isinstance(x, int) else f"{x:,.2f}"


# ---------------------------------------------------------------- 계정 · 부호 · 영역
def code3(code) -> int | None:
    c = reader._code(code)
    if c is None:
        return None
    return c if c < 1000 else int(str(c)[:3])


def in_ranges(code, ranges) -> bool:
    k = code3(code)
    return k is not None and any(a <= k <= b for a, b in ranges)


def sign(code, defs: dict | None = None) -> int:
    """표시 부호. 대변 성격 계정(부채 · 자본 · 수익)은 -1."""
    defs = defs or load_def("workpapers.yaml")
    return -1 if in_ranges(code, defs["credit_ranges"]) else 1


def is_pl(code) -> bool:
    k = code3(code)
    return k is not None and k >= 401


def area_of(name: str, line: str, defs: dict | None = None, code=None) -> tuple[str | None, str]:
    """계정 → 4000 영역 코드와 근거. 계정명 낱말이 표시과목보다 먼저다. 둘 다 없으면 분석과 같은 계정 매핑
    (accounts.yaml 이름 · 코드 · 코드 범위)의 영역을 쓴다(R8, 진단 #5).
    낱말 규칙에 except 가 있으면 그 낱말이 든 계정명은 그 규칙에서 뺀다(「손실충당금」 규칙에서 「재고자산평가손실충당금」 을 빼는 데 쓴다)."""
    defs = defs or load_def("workpapers.yaml")
    n = reader._norm(name)
    for rule in defs.get("account_keywords", []):
        if any(k in n for k in rule["keys"]) and not any(x in n for x in rule.get("except", [])):
            return rule["area"], f"계정명 낱말 '{next(k for k in rule['keys'] if k in n)}'"
    a = (defs.get("line_areas") or {}).get(reader._norm(line))
    if a:
        return a, f"표시과목 '{line}'"
    if code is not None:
        from ..analysis import defs as adefs

        c3 = code3(code)
        m = adefs.map_account(str(c3) if c3 else None, name)
        if m["area"]:
            return m["area"], f"계정 매핑({m['basis']}) accounts.yaml '{m['id']}'"
    return None, "대응 없음(3650A 계정 매핑에서 정해 주십시오)"


def tb_rows(d: Path, defs: dict | None = None) -> list[dict]:
    """현재 시산표 행에 표시 금액(disp) · 영역(area) · 근거를 더한다."""
    defs = defs or load_def("workpapers.yaml")
    out = []
    areas = figures.confirmations(d)["area"]
    for r in figures.current_tb(d):
        a, basis = area_of(r["name"], r["line"], defs, r["code"])
        if areas.get(str(r["code"])):
            a, basis = areas[str(r["code"])].upper().removeprefix("4000"), figures.CONFIRMED
        s = sign(r["code"], defs)
        out.append({**r, "sign": s, "disp": s * r["amount"], "area": a, "area_basis": basis, "pl": is_pl(r["code"])})
    return out


def prior_for(line: str, prior: dict | None, defs: dict | None = None) -> tuple[int | None, str]:
    """표시과목의 전기 금액(재무제표 표시 부호)과 출처. 못 찾으면 (None, 까닭)."""
    if not prior:
        return None, "전기 재무제표 없음"
    defs = defs or load_def("workpapers.yaml")
    allp = {reader._norm(k): v for k, v in {**prior.get("bs", {}), **prior.get("is", {})}.items()}
    key = reader._norm(line)
    for k in [key] + [reader._norm(x) for x in (defs.get("prior_aliases") or {}).get(line, [])]:
        if k in allp:
            return allp[k], f"전기 재무제표 '{k}'"
    for k in allp:
        if key in k.split("·") and "·" in k:
            return None, f"전기 과목 '{k}' 에 합쳐 있음"
    return None, "전기 재무제표에 같은 과목 없음"


# ---------------------------------------------------------------- 중요성
def materiality(d: Path) -> dict:
    """수행중요성 · 명백히 사소한 금액. 2700A 기입란 → 2700A 표 → 엔진 후보 순."""
    d = Path(d)
    wps = {w["no"]: w for w in (idx(d, "조서목록.json", {"workpapers": []})["workpapers"])}
    w = wps.get("2700A")
    out = {"performance": None, "trivial": None, "basis": None, "candidate": False, "candidates": []}
    if w and w.get("file") and (d / w["file"]).suffix == ".xlsx" and (d / w["file"]).exists():
        p = d / w["file"]
        side = ("기준", "근거", "백분율", "÷", "비율")
        for x in wp.readback(p):
            v, lab = x["value"], str(x["label"])
            if not isinstance(v, (int, float)) or isinstance(v, bool) or v <= 0 or any(k in lab for k in side):
                continue
            if lab.startswith("수행중요성") and out["performance"] is None:
                out.update(performance=v, basis=f"2700A 기입란 {x['sheet']}!{x['ref']}")
            elif lab.startswith("명백히 사소") and out["trivial"] is None:
                out["trivial"] = v
        if out["performance"] is None:
            from openpyxl import load_workbook

            wb = load_workbook(p, data_only=True, read_only=True)
            try:
                for ws in wb.worksheets:
                    if ws.title.startswith("_"):
                        continue
                    for row in ws.iter_rows(values_only=True):
                        if not row:
                            continue
                        lab = reader._norm(row[0])
                        v = row[1] if len(row) > 1 else None
                        if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
                            if lab == "수행중요성" and out["performance"] is None:
                                out.update(performance=v, basis=f"2700A '{ws.title}' 시트 수행중요성 행(기입란 등록 없는 파일)")
                            elif lab.startswith("명백히사소") and out["trivial"] is None:
                                out["trivial"] = v
            finally:
                wb.close()
    if out["performance"] is None or out["trivial"] is None:
        # 분석이 2700A 빈 칸을 방법론 기본값으로 채운 적용값(2026-10-07, audit/methodology.py)
        from .. import methodology

        ap = methodology.applied(d)
        perf, triv = ap.get("materiality.performance") or {}, ap.get("materiality.trivial") or {}
        if out["performance"] is None and perf.get("value"):
            out.update(performance=perf["value"], candidate=bool(perf.get("default")),
                       basis=("방법론 " if perf.get("default") else "") + str(perf.get("from")))
        if out["trivial"] is None and triv.get("value"):
            out["trivial"] = triv["value"]
    if out["performance"] is None:
        cands = []
        rows = figures.current_tb(d)
        sales = -sum(r["amount"] for r in rows if 401 <= (code3(r["code"]) or 0) <= 430)
        assets = sum(r["amount"] for r in rows if (code3(r["code"]) or 0) <= 250)
        base = {"매출액": sales, "자산총계": assets}
        for c in load_def("sampling.yaml").get("materiality_candidates", []):
            b = base.get(c["base"])
            if b and b > 0:
                cands.append({"base": c["base"], "amount": b, "rate": c["rate"], "value": int(round(b * c["rate"]))})
        out["candidates"] = cands
        if cands:
            m = min(cands, key=lambda x: x["value"])
            out.update(performance=m["value"], candidate=True,
                       basis=f"엔진 후보({m['base']} {won(m['amount'])} × {m['rate']:.1%}). 2700A 에 수행중요성을 넣기 전")
    return out


# ---------------------------------------------------------------- 읽은 받은자료 기록 (2026-10-12)
_READING: ContextVar = ContextVar("audit_wp_reading", default=None)


@contextmanager
def reading():
    """조서 한 파일을 만드는 동안 엔진이 숫자를 가져온 받은자료(감사 폴더 기준 상대경로)를 모은다. lead.make 가 쓴다.
    받은자료에서 표를 읽어 값을 돌려주는 곳(시산표 · 분개장 · 명세서 · 대장 · 잔액증명)이 note_read 로 적는다.
    이 안이 아닌 곳(3320 · 3600-JET · 분석)에서 부른 note_read 는 아무것도 하지 않는다."""
    got: set[str] = set()
    token = _READING.set(got)
    try:
        yield got
    finally:
        _READING.reset(token)


def note_read(*rels) -> None:
    """지금 만들고 있는 조서가 읽은 받은자료로 적는다. 열어 보았으나 쓸 표가 없던 파일은 적지 않는다."""
    got = _READING.get()
    if got is not None:
        got.update(str(r).replace("\\", "/") for r in rels if r)


class Journals(dict):
    """journals 의 결과. cur · post · prior 행을 꺼낼 때 그 행이 나온 분개장 파일을 읽은 받은자료로 적는다(reading 안에서만).
    분개장은 한 번 읽어 여러 조서가 함께 쓰므로(ctx), 읽은 때가 아닌 꺼내 쓴 때에 적는다."""

    def __getitem__(self, key):
        if key in ("cur", "post", "prior"):
            note_read(*dict.get(self, "part_files", {}).get(key, ()))
        return dict.__getitem__(self, key)


# ---------------------------------------------------------------- 분개장 전수
@dataclass
class Line:
    date: date
    voucher: str
    row: int | None
    code: str | None
    name: str
    dr: float
    cr: float
    memo: str
    party: str | None
    extra: dict = field(default_factory=dict)
    file: str = ""


@dataclass
class Voucher:
    date: date
    no: str
    lines: list
    file: str

    @property
    def amount(self) -> float:
        return sum(x.dr for x in self.lines)

    def codes(self, side: str) -> list[str]:
        return sorted({x.code for x in self.lines if x.code and (x.dr if side == "dr" else x.cr)})

    @property
    def parties(self) -> list[str]:
        return sorted({x.party for x in self.lines if x.party})

    def net(self) -> dict:
        out: dict[tuple, float] = {}
        for x in self.lines:
            k = (x.code, x.party)
            out[k] = out.get(k, 0.0) + x.dr - x.cr
        return {k: v for k, v in out.items() if abs(v) > 0.5}

    def label(self, side: str) -> str:
        names = []
        for x in self.lines:
            if (x.dr if side == "dr" else x.cr) and x.name not in names:
                names.append(x.name or str(x.code))
        return ", ".join(names)

    @property
    def memo(self) -> str:
        ms = []
        for x in self.lines:
            if x.memo and x.memo not in ms:
                ms.append(x.memo)
        return " / ".join(ms)


EXTRA_COLS = reader.JOURNAL_EXTRA


def read_journal(path: Path, rel: str = "") -> list[Line]:
    """분개장 한 파일의 행 전부(reader.journal_rows 의 꼴 그대로). 합계 · 소계 행은 뺀다. 전표번호가 없으면 일자를 전표로 본다."""
    out = []
    for x in reader.journal_rows(Path(path)):
        dt = x["date"]
        out.append(Line(date=dt, voucher=x["voucher"] or dt.isoformat(),
                        row=int(x["row"]) if x["row"] is not None else None, code=x["code"], name=x["name"],
                        dr=x["dr"], cr=x["cr"], memo=x["memo"] or "", party=x["party"], extra=dict(x["extra"]), file=rel))
    return out


def journals(d: Path) -> dict:
    """분개장 받은자료 전부를 읽어 당기 · 기말 후 · 전기로 나눈다. {"cur", "post", "prior", "files", "extra_cols", "failed", "part_files",
    "excluded"}
    failed: 받은자료에 있으나 판독하지 못한 분개장 [{path, why}]. part_files: 당기 · 기말 후 · 전기 행이 나온 파일 {"cur": [..], ..}
    excluded: 다른 분개장과 기간이 겹쳐 합치지 않았거나 같은 줄을 뺀 분개장 [{path, why, dropped}](figures.journal_plan, 2026-10-18)"""
    d = Path(d)
    start, end = fy_start(d), fy_end(d)
    cur, post, prior, files, extra, failed, excluded = [], [], [], [], set(), [], []
    parts = {"cur": (cur, set()), "post": (post, set()), "prior": (prior, set())}
    plan = figures.journal_plan(d, reader.READABLE)
    dd = figures.JournalDedupe(plan)
    for f in plan:
        rel = f["path"]
        if not f["use"]:
            excluded.append({"path": rel, "why": f["why"], "dropped": None})
            continue
        lines = read_journal(d / rel, rel)
        if not lines:
            # 파일은 있는데 판독하지 못했다. 「분개장 없음」 과 가른다(R8)
            failed.append({"path": rel, "why": reader.read_failure(d / rel, reader.JOURNAL) or reader.READ_FAIL})
            continue
        files.append(rel)
        for x in lines:
            k = figures.journal_line_key(x.date, x.voucher, x.code or x.name, x.dr, x.cr)
            if not dd.keep(rel, k, bool(f["dedupe"])):
                continue
            dd.seen(k)
            extra |= set(x.extra)
            rows, used = parts["cur" if start <= x.date <= end else "post" if x.date > end else "prior"]
            rows.append(x)
            used.add(rel)
        if f["dedupe"]:
            excluded.append({"path": rel, "why": f["why"], "dropped": dd.dropped.get(rel, 0)})
    return Journals({"cur": cur, "post": post, "prior": prior, "files": files, "extra_cols": sorted(extra), "failed": failed,
                     "part_files": {k: sorted(used) for k, (_, used) in parts.items()}, "excluded": excluded})


def vouchers(lines: list[Line]) -> list[Voucher]:
    by: dict[tuple, Voucher] = {}
    for x in lines:
        k = (x.file, x.date, x.voucher)
        if k not in by:
            by[k] = Voucher(date=x.date, no=x.voucher, lines=[], file=x.file)
        by[k].lines.append(x)
    return sorted(by.values(), key=lambda v: (v.date, v.no))


def prior_parties(j: dict) -> set[str] | None:
    """전기 분개장의 거래처. 전기 분개장이 없으면 None(신규 거래처를 판정하지 못함)."""
    if not j["prior"]:
        return None
    return {x.party for x in j["prior"] if x.party}


# ---------------------------------------------------------------- 매출채권 거래처별
def ar_by_party(d: Path) -> dict:
    """기말 매출채권 거래처별 잔액. 연령분석표 → 계정명세서(외상매출금 · 매출채권 시트) 순. {"rows": [(거래처, 금액)], "source", "excluded"}
    excluded: 거래처 이름 칸에 결산 · 조정 · 선충전 같은 낱말이 있는 줄(workpapers.yaml confirmations.non_party_keywords)."""
    d = Path(d)
    skip = load_def("workpapers.yaml").get("confirmations", {}).get("non_party_keywords", [])
    for kind, sheet_keys in ((reader.AGING, None), ("계정명세서", ("외상매출금", "매출채권"))):
        for rel in figures.received(d, kind):
            if Path(rel).suffix.lower() not in reader.READABLE:
                continue
            tabs = reader._tables(d / rel, [["party"], ["ar_total"]], [])
            if kind == reader.AGING:                # 분기 연령표(Roll-Rate)가 함께 있어도 결산일 기준표만(2026-10-16)
                tabs = reader.asof_tabs(tabs, case(d).get("fiscal_year_end"))
            for sheet, c, rows, _, _ in tabs:
                if sheet_keys and not any(k in sheet for k in sheet_keys):
                    continue
                out, excluded = [], []
                for r in rows:
                    if reader._is_total_row(r):
                        continue
                    p = reader._cell(r, c["party"])
                    v = reader._num(reader._cell(r, c["ar_total"]))
                    if p in (None, "") or v is None:
                        continue
                    name = str(p).strip()
                    # 결산조정분개 · 예치금 선충전처럼 거래처가 아닌 줄은 조회 · 표본 모집단에서 뺀다(R8, 진단 #14)
                    hit = next((k for k in skip if k in reader._norm(name)), None)
                    if hit:
                        excluded.append({"party": name, "amount": reader._int(v), "why": f"거래처 아닌 줄(낱말 '{hit}')"})
                        continue
                    out.append((name, reader._int(v)))
                if out:
                    src = f"{rel} '{sheet}' 시트" + (f". 거래처 아닌 줄 {len(excluded)}개는 뺌(회계사 확인)" if excluded else "")
                    note_read(rel)
                    return {"rows": out, "source": src, "excluded": excluded}
    return {"rows": [], "source": None, "excluded": []}


def schedule_balances(d: Path, heads: tuple[str, ...] = ("기말",)) -> dict:
    """계정명세서의 계정코드별 기말 잔액(차변 양수). {코드: 금액}. 없으면 빈 dict.
    heads 는 금액 열 머리 이름의 앞 글자다. 기초 잔액은 schedule_openings 가 ("기초", "전기말") 로 부른다."""
    d = Path(d)
    for rel in figures.received(d, "계정명세서"):
        if Path(rel).suffix.lower() not in reader.READABLE:
            continue
        wb = reader._open(d / rel)
        try:
            for ws in wb.worksheets:
                rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
                for hi, row in enumerate(rows[:reader.HEADER_SCAN]):
                    hs = [reader._norm(v) for v in row]
                    ci = next((i for i, h in enumerate(hs) if h in ("계정코드", "코드")), None)
                    ei = next((i for i, h in enumerate(hs) if h.startswith(heads)), None)
                    if ci is None or ei is None:
                        continue
                    out = {}
                    for r in rows[hi + 1:]:
                        code = reader._code(reader._cell(r, ci))
                        v = reader._num(reader._cell(r, ei))
                        if code is not None and v is not None:
                            out[str(code)] = reader._int(v)
                    if out:
                        note_read(rel)
                        return {"balances": out, "source": f"{rel} '{ws.title}' 시트"}
        finally:
            wb.close()
    return {"balances": {}, "source": None}


def schedule_openings(d: Path) -> dict:
    """계정명세서의 계정코드별 기초 잔액(차변 양수). 머리 이름이 「기초」 나 「전기말」 로 시작하는 열을 읽는다. 없으면 빈 dict."""
    return schedule_balances(d, ("기초", "전기말"))


# ---------------------------------------------------------------- 기입란 · 저장
def ask(b: wp.Book, ws, row: int, col: int, label: str, rule: str = "required", **k):
    """같은 시트 안에서 항목 이름이 겹치지 않게 번호를 붙여 기입란을 등록한다(다시 만들 때 값을 옮기는 열쇠)."""
    used = {lab for sh, _, lab, _ in b.inputs if sh == ws.title}
    lab, n = label, 2
    while lab in used:
        lab, n = f"{label} ({n})", n + 1
    return b.ask(ws, row, col, lab, rule=rule, **k)


class Kept(Exception):
    """이전 기입값을 옮길 칸이 없어 저장하지 않았다."""

    def __init__(self, path: Path, lost: list[str], legacy: bool = False):
        if legacy:
            msg = f"{path.name}: 엔진이 만들지 않은 조서(기입란 목록 없음)라 덮어쓰지 않았습니다. 다시 만들려면 --force 를 주십시오"
        else:
            msg = f"{path.name}: 이전 기입값 {len(lost)}개를 옮길 칸이 없어 저장하지 않았습니다. --force 로 덮어쓸 수 있습니다: " + "; ".join(lost[:5])
        super().__init__(msg)
        self.path, self.lost, self.legacy = path, lost, legacy


# 4000 조서 골격 재편(2026-10-12) 전 파일의 시트 → 새 시트. 「리드」 한 장이 「Lead」 와 「표지」 로 갈렸다.
# 옛 「해당 여부: …」 는 Part 2 칸뿐이었고 새 이름은 「추가 절차 해당 여부: …」 다(Part 1 의 「해당 여부: …」 와 겹치지 않게)
OLD_SHEETS = {"리드": ("Lead", "표지")}
OLD_LABELS = {"리드": {"해당 여부: ": "추가 절차 해당 여부: "}}


def _carry_target(sheet: str, label: str, new_keys: dict, by_base: dict):
    """옛 기입란 (시트, 항목) 이 옮겨 갈 새 칸 (시트, 셀). 같은 이름 → 번호만 다른 같은 시트(.02 표본 ↔ .03 표본 ↔ 표본)
    → 골격 재편 전 시트 이름 순으로 찾는다. 없으면 None."""
    if (sheet, label) in new_keys:
        return new_keys[(sheet, label)]
    hit = by_base.get((wp.base_name(sheet), label))
    if hit:
        return hit
    for old, new in (OLD_LABELS.get(sheet) or {}).items():
        if label.startswith(old):
            label = new + label[len(old):]
            break
    for sh in OLD_SHEETS.get(sheet, ()):
        if (sh, label) in new_keys:
            return new_keys[(sh, label)]
    return None


def prev_version_file(path: Path) -> Path | None:
    """같은 조서의 앞 회사제시 판 파일(이름의 _v<판> 만 다른 엑셀). 여럿이면 판이 가장 높은 것.
    회사가 새 판을 내면 조서 이름이 바뀌어 새 파일로 만들어진다. 그때 앞 판 조서의 기입값을 옮기려고 쓴다(2026-10-12)."""
    stem = folder.VERSION_RE.sub("", path.stem)
    cands = [p for p in path.parent.glob("*" + path.suffix)
             if p != path and not p.name.startswith("~$") and folder.VERSION_RE.sub("", p.stem) == stem]
    if not cands:
        return None

    def key(p: Path):
        m = folder.VERSION_RE.search(p.stem)
        return folder._version_key(m.group(1) if m else None)

    return max(cands, key=key)


def save(b: wp.Book, force: bool = False) -> dict:
    """저장 전에 같은 파일의 기입란 값을 (시트, 항목) 으로 옮긴다. 옮길 칸이 없는 값이 있으면 force 없이는 저장하지 않는다.
    절차 시트 번호가 달라져도(.02 표본 → .03 표본) 번호를 뗀 시트 이름이 같으면 옮긴다.
    지금 판 이름의 파일이 없으면 앞 회사제시 판 조서에서 옮긴다. 앞 판 파일은 그대로 두므로 옮길 칸이 없는 값이 있어도 저장한다."""
    path = wp.path_for(b.d, b.no)
    carried, moved, lost = 0, 0, []
    src = path if path.exists() else prev_version_file(path)
    if src is not None and not _engine_made(src):
        # 회계사가 손으로 만든(또는 엔진 밖에서 만든) 조서는 기입란 목록이 없어 값을 옮길 수 없다. 덮어쓰지 않는다
        if src == path and not force:
            raise Kept(path, [], legacy=True)
    elif src is not None:
        old = {(x["sheet"], x["label"]): x["value"] for x in wp.readback(src) if not wp._blank(x["value"])}
        new_keys = {(sh, lab): (sh, ref) for sh, ref, lab, _ in b.inputs}
        by_base = {(wp.base_name(sh), lab): (sh, ref) for sh, ref, lab, _ in b.inputs}
        for k, v in old.items():
            target = _carry_target(k[0], k[1], new_keys, by_base)
            if target:
                sh, ref = target
                cell = b.wb[sh][ref]
                # carried 는 남은 값 수(엔진이 이미 같은 값을 넣은 칸 포함), moved 는 빈 칸에 실제로 옮긴 값 수(2026-10-16 시연:
                # 회계사가 아무것도 적지 않았는데 표지 조서번호 기본값 · 주석 단위 때문에 「기입값 7개 옮김」 으로 보였다)
                if wp._blank(cell.value):
                    cell.value = v
                    moved += 1
                carried += 1
            elif k[1] != "검토자":
                lost.append(f"{k[0]}!{k[1]}")
        if lost and not force and src == path:
            raise Kept(path, lost)
    p = b.save(path)
    return {"no": b.no, "path": p, "carried": carried, "moved": moved, "lost": lost}


def _engine_made(path: Path) -> bool:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return wp.META_SHEET in wb.sheetnames
    finally:
        wb.close()


def old_inputs(d: Path, no: str) -> dict:
    """이미 있는 조서의 기입란 값 {(시트, 항목): 값}. 지금 판 이름의 파일이 없으면 앞 회사제시 판 조서에서 읽는다."""
    path = wp.path_for(d, no)
    if not path.exists():
        path = prev_version_file(path)
    if path is None or not _engine_made(path):
        return {}
    return {(x["sheet"], x["label"]): x["value"] for x in wp.readback(path) if not wp._blank(x["value"])}


def ceil(x: float) -> int:
    return int(math.ceil(x - 1e-9))


def short(s: str, n: int = 40) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[:n - 1] + "…"


# ---------------------------------------------------------------- 표본 ↔ 증빙 (2026-10-09)
EVIDENCE_KIND = "표본 증빙"          # folder_standard.yaml 받은자료 「증빙」 폴더의 종류


def evidence_files(d: Path) -> list[str]:
    """받은자료 증빙 폴더의 파일(하위 폴더 포함). 표본 행과 이름으로 잇는다."""
    return [f["path"] for f in figures._docs(d)["files"] if f.get("role") == "received" and f.get("kind") == EVIDENCE_KIND]


def _ev_norm(s: str) -> str:
    return re.sub(r"[\s_\-·.()\[\]/]", "", str(s or "").replace("\\", ""))


def match_evidence(files: list[str], party: str | None = None, voucher: str | None = None, day: str | None = None) -> list[str]:
    """표본 한 행에 맞는 증빙 파일. 이름(경로 포함)에 전표번호가 있으면 그 파일, 없으면 거래처 이름이 있는 파일.
    day 를 주면(매출 분개 행) 거래처만으로는 잇지 않고 같은 날짜도 경로에 있어야 한다. 한 거래처의 증빙이 여러 거래에 걸리는 것을 막는다.
    파일 안의 글자는 보지 않는다(PDF · 스캔본 판독은 감사-맥락색인 몫)."""
    v = _ev_norm(voucher) if voucher else ""
    pt = _ev_norm(party) if party else ""
    dy = _ev_norm(day) if day else ""
    out = []
    for p in files:
        s = _ev_norm(p)
        if v and v in s:
            out.append(p)
        elif pt and pt in s and (not dy or dy in s):
            out.append(p)
    return out

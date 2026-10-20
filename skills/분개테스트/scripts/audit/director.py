"""지휘: 빠진 자료 · 다음 할 일 · 단계 (R1 T3, 2026-10-05).

_색인 을 읽어 결정적 규칙으로 정한다. 규칙은 정의 파일 두 개에 있다.
  definitions/pbc_match.yaml      요청자료(PBC) 항목과 이슈 대응 조서의 필요 자료 → 받은자료 종류
  definitions/director_rules.yaml 단계, 조서별 스킬과 필요 자료, 다음 할 일 규칙과 우선순위

  missing(dir)            빠진 자료 목록(문서목록.missing 형식). 앞서 적힌 requested_at 은 이어받는다.
                          분석.json requests 가운데 회사 몫(to 회사)도 잇는다(2026-10-18). 같은 요청자료 줄이 빠진 자료면 그 줄에 구체 문구
                          (detail)를 붙이고, 아니면 한 줄(source 분석)로 낸다. need 조건을 채운 받은자료가 들어오면 사라진다
  next_actions(dir)       다음 할 일 목록(감사건.next_actions 형식), 우선순위대로 max_actions 까지
  stage(dir)              조서 상태로 정한 단계(준비 · 계획 · 실행 · 보고 · 종결)
  run(dir, now)           셋을 _색인 에 쓰고 기록을 남긴 뒤 validate_index 로 검사한다

받은자료 판독(종류 · 사실)은 맥락 색인(reader · index apply) 몫이고, 여기서는 색인에 적힌 것만 본다.
"""
import calendar
import re
from datetime import date, timedelta
from pathlib import Path

from . import folder
from .errors import DefinitionError

STAGES = ("준비", "계획", "실행", "보고", "종결")
RULE_IDS = ("확인-판독", "요청-빠진자료", "다시확인", "기본값확인", "수정반영", "보고준비", "스킬실행", "전기이어받기", "초도기초잔액",
            "검토요청", "작성", "확인-분석값")
AREA_DEFAULT = "4000*"
# 손익 계정 단위 전기 원천이 없을 때 빠진 자료로 세는 항목(R7, 2026-10-06)
PRIOR_PL_ITEM = "전기 시산표(계정별) 또는 전기 분개장"
PERIODS = ("annual", "events", "fy_end", "after_end")
NO_BALANCE = "해당 없음 후보"          # 시산표 · 재무제표에 영역 잔액이 당기 · 전기 모두 0 인 요청 줄(2026-10-18)


# ---------------------------------------------------------------- 정의
def _yaml(name: str) -> dict:
    p = folder.DEFS / name
    if not p.exists():
        raise DefinitionError(f"정의 파일이 없습니다: {p}")
    return folder.read_yaml(p)


def _known_kinds(std: dict) -> set[str]:
    out = {str(v) for v in std["permanent"]["items"].values()}
    for f in std["folders"]:
        for c in f.get("children", []):
            for k in c["kinds"]:
                out.add(k["name"] if isinstance(k, dict) else k)
    return out


def _spec_kinds(spec: dict) -> list[str]:
    out = list(spec.get("kinds") or [])
    for s in spec.get("any_of") or []:
        out += _spec_kinds(s)
    return out


def load_definitions() -> tuple[dict, dict, dict]:
    """(폴더 표준, 대응 표, 지휘 규칙). 표에 없는 종류 이름이나 규칙 이름이 있으면 멈춘다."""
    std = folder.load_standard()
    match = _yaml("pbc_match.yaml")
    rules = _yaml("director_rules.yaml")
    known = _known_kinds(std)
    used = []
    for e in match.get("pbc", []):
        used += [(f"pbc {e.get('row')}", k) for k in _spec_kinds(e)]
        if e.get("from_stage") and e["from_stage"] not in STAGES:
            raise DefinitionError(f"pbc_match.yaml pbc {e.get('row')}: 모르는 단계 {e['from_stage']}")
        _check_period(f"pbc {e.get('row')}", e)
    for wp, needs in (match.get("issue_needs") or {}).items():
        for n in needs:
            used += [(f"issue_needs {wp}", k) for k in _spec_kinds(n)]
            _check_period(f"issue_needs {wp}", n)
    for wp, cfg in (rules.get("skills") or {}).items():
        for n in cfg.get("needs", []):
            used += [(f"skills {wp}", k) for k in (n if isinstance(n, list) else [n])]
    bad = [f"{where}: {k}" for where, k in used if k not in known]
    if bad:
        raise DefinitionError("폴더 표준에 없는 받은자료 종류입니다: " + "; ".join(bad))
    names = [s["name"] for s in rules["stages"]]
    if tuple(names) != STAGES:
        raise DefinitionError(f"director_rules.yaml stages 는 {' · '.join(STAGES)} 순서여야 합니다: {names}")
    for r in rules["rules"]:
        if r["id"] not in RULE_IDS:
            raise DefinitionError(f"director_rules.yaml 에 모르는 규칙이 있습니다: {r['id']}")
    return std, match, rules


def _period_list(v) -> list[str]:
    return [v] if isinstance(v, str) else list(v or [])


def _check_period(where: str, spec: dict) -> None:
    per = spec.get("period")
    if per is None:
        return
    items = per.items() if isinstance(per, dict) else [(None, per)]
    for kind, v in items:
        bad = [x for x in _period_list(v) if x not in PERIODS]
        if bad or (kind is not None and kind not in (spec.get("kinds") or [])):
            raise DefinitionError(f"pbc_match.yaml {where}: period 는 {' · '.join(PERIODS)} 이고 종류는 kinds 에 있어야 합니다: {per}")


# ---------------------------------------------------------------- 자료가 덮는 기간(2026-10-18)
def _month_end(y: int, m: int) -> date:
    return date(y, m, calendar.monthrange(y, m)[1])


def _point(s: str) -> tuple[date | None, bool]:
    """'2024' · '2024-10' · '2024-10-31' → (그 기간의 끝 날, 연도만인가)."""
    m = re.fullmatch(r"(\d{4})(?:-(\d{2})(?:-(\d{2}))?)?", s.strip())
    if not m:
        return None, False
    y = int(m.group(1))
    try:
        if m.group(3):
            return date(y, int(m.group(2)), int(m.group(3))), False
        if m.group(2):
            return _month_end(y, int(m.group(2))), False
    except ValueError:
        return None, False
    return date(y, 12, 31), True


def _year_of_month(m: int, fe: date) -> int:
    """연도 없이 월만 적힌 자료(「9월귀속」 · 「3Q」 · 「12월」)의 연도. 사업연도 안의 그 달로 본다."""
    return fe.year if m <= fe.month else fe.year - 1


def _month_rule(fe):
    return lambda m: _month_end(_year_of_month(int(m[1]), fe), int(m[1])) if 1 <= int(m[1]) <= 12 else None


def _name_end(name: str, fs: date, fe: date, folder_only: bool = False) -> date | None:
    """파일 · 폴더 이름의 연월이 덮는 마지막 날. 귀속월 → 분기 → 날짜 → 연월 → 월 순서로 처음 맞는 꼴만 쓴다
    (「9월귀속10월지급」 의 지급월 10 을 기간으로 읽지 않게). 사업연도 앞뒤 한 해를 넘는 날짜는 계정코드 같은 숫자로 보고 버린다.
    폴더 이름은 귀속월 · 분기만 본다(「2024년 1월~현재까지」 같은 폴더 이름을 기간으로 읽지 않게)."""
    yy_ok = {(fs.year - 1) % 100, fs.year % 100, fe.year % 100, (fe.year + 1) % 100}
    lo, hi = fs - timedelta(days=400), fe + timedelta(days=400)

    def quarter(m):
        q = int(m[1]) * 3
        return _month_end(_year_of_month(q, fe), q)

    groups = [
        [(_month_rule(fe), r"(\d{1,2})\s*월\s*귀속")],
        [(lambda m: _month_end(int(m[1]), int(m[2]) * 3), r"(20\d\d)\s*[.\-_ ]?\s*([1-4])\s*[Qq](?![A-Za-z])"),
         (quarter, r"(?<![0-9A-Za-z])([1-4])\s*[Qq](?![A-Za-z])"),
         (quarter, r"([1-4])\s*분기")],
    ]
    if not folder_only:
        groups += [
            [(lambda m: date(int(m[1]), int(m[2]), int(m[3])), r"(20\d\d)[.\-_]?(0[1-9]|1[0-2])[.\-_]?(0[1-9]|[12]\d|3[01])(?!\d)"),
             (lambda m: date(2000 + int(m[1]), int(m[2]), int(m[3])) if int(m[1]) in yy_ok else None,
              r"(?<!\d)(\d\d)(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)")],
            [(lambda m: _month_end(int(m[1]), int(m[2])) if 1 <= int(m[2]) <= 12 else None, r"(20\d\d)\s*년\s*(\d{1,2})\s*월"),
             (lambda m: _month_end(int(m[1]), int(m[2])), r"(20\d\d)[.\-_](0[1-9]|1[0-2])(?!\d)"),
             (lambda m: _month_end(2000 + int(m[1]), int(m[2])) if int(m[1]) in yy_ok else None, r"(?<!\d)(\d\d)(0[1-9]|1[0-2])(?!\d)")],
            [(_month_rule(fe), r"(?<![\d~])(\d{1,2})\s*월(?!\s*(귀속|~))")],
        ]
    for g in groups:
        got = []
        for rule, pat in g:
            for m in re.finditer(pat, name):
                try:
                    dt = rule(m)
                except ValueError:
                    dt = None
                if dt and lo <= dt <= hi:
                    got.append(dt)
        if got:
            return max(got)
    return None


def file_end(f: dict, fs: date, fe: date) -> date | None:
    """받은자료 한 파일이 덮는 마지막 날. 분개장 기간 사실 → 문서목록 period(엔진 · AI) → 파일 이름 → 바로 위 폴더 이름(분기 · 귀속월만).
    period 가 연도뿐이면 파일 이름에서 더 잘게 찾고, 없으면 그 해 전체로 본다. 알 수 없으면 None."""
    from . import reader

    rng = reader.journal_range(f.get("facts") or [])
    if rng:
        return rng[1]
    end, coarse = (None, False)
    if f.get("period"):
        end, coarse = _point(str(f["period"]).split("~")[-1])
        if end and not coarse:
            return end
    p = Path(f["path"])
    got = _name_end(p.stem, fs, fe) or _name_end(p.parent.name, fs, fe, folder_only=True)
    if got:
        return got
    if end and coarse:
        return fe if end.year == fe.year else end
    return None


def file_start(f: dict, fs: date, fe: date) -> date | None:
    """받은자료 한 파일의 날짜(기간이면 시작일). 문서목록 period 앞쪽 → 파일 이름의 날짜 · 연월. 알 수 없으면 None."""
    if f.get("period"):
        first = str(f["period"]).split("~")[0]
        m = re.fullmatch(r"(\d{4})-(\d{2})(?:-(\d{2}))?", first)
        if m:
            try:
                return date(int(m[1]), int(m[2]), int(m[3] or 1))
            except ValueError:
                return None
    end = _name_end(Path(f["path"]).stem, fs, fe)
    return end


def _gap_text(per: str, end: date, fe: date) -> str | None:
    """덮은 마지막 날 end 로 본 빠진 기간 글. 다 덮었으면 None."""
    if per == "after_end":
        return None if end > fe else "결산일 후분"
    if end >= fe:
        return None
    if per == "fy_end":
        return f"{fe.month}월말 기준"
    if per == "events":
        # 열린 날이 기간인 자료(의사록). 회의가 없던 달은 빠진 것이 아니므로 사업연도 마지막 3개월 안에 하나도 없을 때만 묻는다
        q = fe.month - 2 if fe.month > 2 else fe.month + 10
        last_q = date(fe.year if fe.month > 2 else fe.year - 1, q, 1)
        if end >= last_q:
            return None
    nxt = date(end.year + (end.month == 12), end.month % 12 + 1, 1)
    if nxt > fe:
        return None
    return f"{fe.month}월분" if (nxt.year, nxt.month) == (fe.year, fe.month) else f"{nxt.month}~{fe.month}월분"


# ---------------------------------------------------------------- 읽기
class _Case:
    """한 감사 폴더의 _색인 과 정의를 한 번 읽어 둔다."""

    def __init__(self, engagement_dir: Path):
        self.d = Path(engagement_dir)
        idx = self.d / "_색인"
        if not (idx / "감사건.json").exists():
            raise DefinitionError(f"_색인/감사건.json 이 없습니다. 먼저 folder new 나 folder scan 을 돌려 주십시오: {self.d}")
        self.std, self.match, self.rules = load_definitions()
        self.case = folder._read(idx / "감사건.json", None)
        self.docs = folder._read(idx / "문서목록.json", {"format_version": 1, "files": [], "missing": [], "warnings": []})
        self.wps = folder._read(idx / "조서목록.json", {"workpapers": []})["workpapers"]
        self.issues = folder._read(idx / "이슈.json", {"issues": []})["issues"]
        bases = set(self.rules.get("counted_bases", []))
        self.files = [f for f in self.docs["files"] if f["role"] in ("received", "permanent") and (f["confirmed"] or f["kind_basis"] in bases)]
        self.done = set(self.rules.get("done_statuses", folder.DONE))
        self.wp_folder = folder._wp_folder_map(self.std)
        self.names = {str(k): v for k, v in self.std.get("workpaper_names", {}).items()}
        self.folder_order = [f["path"] for f in self.std["folders"]]
        self.pbc_path = None
        self.warnings: list[str] = []
        self._stage = None
        self.assign = (folder._read(idx / "배정.json", {}) or {}).get("assignments") or {}      # 스태프 배정(2026-10-12)
        try:
            self.fy = (folder.fiscal_start(self.case), date.fromisoformat(self.case["fiscal_year_end"]))
        except (KeyError, ValueError, TypeError):
            self.fy = None
        self._no_balance = None

    def judge(self, spec: dict) -> tuple[bool, list[str]]:
        return _judge(spec, self.files, self.fy)

    def no_balance(self) -> set[str]:
        """시산표 · 재무제표에 계정 잔액이 당기 · 전기 모두 없는 4000 영역(2026-10-18). 당기 시산표를 못 읽으면 빈 집합(아무것도 빼지 않는다)."""
        if self._no_balance is None:
            from . import pbcform

            has = area_balances(self.d)
            areas = {str(e["area"]) for e in self.match.get("pbc", []) if e.get("area")}
            areas |= {str(r["area"]) for r in pbcform.load()["rows"] if r.get("area")}
            self._no_balance = set() if has is None else {a for a in areas if not has.get(a)}
        return self._no_balance

    def dropped(self) -> set[str]:
        """잔액이 없어 요청 줄도 내지 않는 영역(pbc_match.yaml no_balance.drop)."""
        return self.no_balance() & {str(x) for x in (self.match.get("no_balance") or {}).get("drop", [])}

    # 조서: 조서목록 + 열린 이슈의 대응 조서 가운데 아직 목록에 없는 것(상태 없음)
    def workpapers(self) -> list[dict]:
        out = list(self.wps)
        have = {w["no"] for w in out}
        for i in self.open_issues():
            no = i.get("response_wp")
            if no and no not in have:
                have.add(no)
                out.append({"no": no, "name": self.names.get(no, no), "folder": self._folder_of(no), "file": None,
                            "status": "없음", "needs_recheck": False})
        order = {p: n for n, p in enumerate(self.folder_order)}
        return [w for _, _, w in sorted((order.get(w["folder"], len(order)), n, w) for n, w in enumerate(out))]

    def _folder_of(self, no: str) -> str:
        return self.wp_folder.get(no) or ("4000 계정별" if folder._is_area(no) else "")

    def open_issues(self) -> list[dict]:
        return [i for i in self.issues if i["status"] == "열림"]

    def prior_source(self) -> dict | None:
        """시작 자료가 쓴 전기 판단 정보(_색인/시작자료.json prior_source). 시작 자료 전이면 None."""
        return (folder._read(self.d / "_색인" / "시작자료.json", {}) or {}).get("prior_source")

    def missing_prior_pl(self) -> bool:
        """손익 계정 단위 전기 원천이 없는지. 분석.json(나중에 돈 분석)이 시작자료.json 보다 앞선다.
        그 뒤에 전기 시산표나 전기 열두 달을 덮는 분개장이 들어왔으면 없는 것으로 보지 않는다."""
        an = (folder._read(self.d / "_색인" / "분석.json", {}) or {}).get("prior")
        ps = self.prior_source() or {}
        flag = an.get("missing_pl") if an else (ps.get("account_prior") or {}).get("missing_pl")
        if not flag:
            return False
        from datetime import date

        from . import reader

        try:
            from .folder import fiscal_start

            pe = date.fromordinal(fiscal_start(self.case).toordinal() - 1)     # 전기말(반기 · 짧은 사업연도 반영, R8)
        except (KeyError, ValueError, TypeError):
            return True
        p_start = pe.replace(year=pe.year - 1) if not (pe.month == 2 and pe.day == 29) else pe.replace(year=pe.year - 1, day=28)
        for f in self.files:
            if f["kind"] == reader.PRIOR_TB:
                return False
            if f["kind"] == reader.JOURNAL:
                rng = reader.journal_range(f.get("facts") or [])
                if rng and rng[0] <= date.fromordinal(p_start.toordinal() + 1) and rng[1] >= pe:
                    return False
        return True


def _hay(f: dict) -> str:
    return " ".join([f["path"], f["kind"]] + [x.get("text", "") for x in f.get("facts") or []])


def _read_ok(spec: dict, f: dict) -> bool:
    """read_as: 엔진 판독 사실이 그 글로 시작하는 파일만 센다. 엔진 사실이 없으면 unread_name 낱말이 파일 이름에 있어야 한다."""
    want = spec.get("read_as")
    if not want:
        return True
    eng = [x.get("text", "") for x in f.get("facts") or [] if x.get("basis") == "엔진"]
    if eng:
        return any(t.startswith(w) for t in eng for w in want)
    return any(w in Path(f["path"]).name for w in spec.get("unread_name") or [])


def _judge(spec: dict, files: list[dict], fy: tuple[date, date] | None = None) -> tuple[bool, list[str]]:
    """(받았는가, 빠진 것). 빠진 것은 need: all 일 때 빠진 종류와 period 로 본 모자란 기간 글(「10~12월분」 · 「12월말 기준」).
    fy(사업연도 시작 · 결산일)가 없으면 기간은 보지 않는다. 형식은 pbc_match.yaml 머리 설명."""
    if spec.get("any_of"):
        return any(_judge(s, files, fy)[0] for s in spec["any_of"]), []
    kinds = spec.get("kinds") or []
    kws = spec.get("keywords") or []

    def in_year(f):
        """dated_in_year: 날짜가 사업연도 안인 자료만 센다(당기 신규 약정). 날짜를 모르는 자료는 세지 않는다."""
        if not spec.get("dated_in_year") or not fy:
            return True
        st = file_start(f, *fy)
        return bool(st and fy[0] <= st <= fy[1])

    def after(f):
        """dated_after: 날짜가 그날 뒤인 자료만 센다(주주명부 기준일 뒤 의사록, 2026-10-18). 날짜를 모르는 자료는 세지 않는다."""
        if not spec.get("dated_after") or not fy:
            return True
        st = file_start(f, *fy)
        return bool(st and st > date.fromisoformat(spec["dated_after"][:10]))

    def hit(f, kind=None):
        return ((kind is None or f["kind"] == kind) and (not kws or any(k in _hay(f) for k in kws)) and _read_ok(spec, f) and in_year(f)
                and after(f))

    if not kinds:
        return any(hit(f) for f in files), []
    have = [k for k in kinds if any(hit(f, k) for f in files)]
    if not have:
        return False, []
    lack = [k for k in kinds if k not in have] if spec.get("need", "any") == "all" else []
    gaps = _period_gaps(spec, [f for f in files if hit(f) and f["kind"] in kinds], fy) if fy else []
    return not lack and not gaps, lack + gaps


def _period_gaps(spec: dict, hits: list[dict], fy: tuple[date, date]) -> list[str]:
    """period 로 본 모자란 기간. 덮는 기간을 하나도 모르면 받음으로 본다(빈 목록)."""
    per = spec.get("period")
    if not per:
        return []
    fs, fe = fy
    words = spec.get("period_words") or []
    groups = per.items() if isinstance(per, dict) else [(None, per)]
    out = []
    for kind, v in groups:
        fl = [f for f in hits if kind is None or f["kind"] == kind]
        def said(f):          # 종류 이름 · 폴더 이름(「잔액증명」)은 빼고 파일 이름 · 사실 문장만 본다
            return any(w in " ".join([Path(f["path"]).name] + [x.get("text", "") for x in f.get("facts") or []]) for w in words)

        if words and any(said(f) for f in fl):
            fl = [f for f in fl if said(f)]     # 잔액증명은 거래내역 · 화면 캡처가 아닌 증명서로 본다
        ends = [file_end(f, fs, fe) for f in fl]
        if not ends or None in ends:          # 덮는 기간을 모르는 파일이 하나라도 있으면 그 파일이 남은 기간일 수 있어 받음으로 본다
            continue
        for p in _period_list(v):
            t = _gap_text(p, max(ends), fe)
            if t:
                out.append(f"{kind} {t}" if kind else t)
    return out


GAP_RE = re.compile(r"(월분|월말 기준|결산일 후분)$")


def _item_text(base: str, lack: list[str]) -> str:
    """빠진 종류는 「(빠진 것: …)」, 모자란 기간은 「(10~12월분)」 처럼 붙인다."""
    kinds = [x for x in lack if not GAP_RE.search(x)]
    gaps = [x for x in lack if GAP_RE.search(x)]
    out = base + (f" (빠진 것: {' · '.join(kinds)})" if kinds else "")
    return out + (f" ({' · '.join(gaps)})" if gaps else "")


def base_item(text: str) -> str:
    """빠진 자료 글에서 엔진이 붙인 꼬리(빠진 것 · 모자란 기간)를 뗀 요청 항목 글."""
    return re.sub(r"( \((빠진 것: [^()]*|[^()]*(월분|월말 기준|결산일 후분))\))+$", "", text)


# ---------------------------------------------------------------- 영역 잔액(2026-10-18)
_BAL_CACHE: dict = {}


def area_balances(d: Path) -> dict[str, bool] | None:
    """{4000 영역: 잔액 있음}. 당기 회사제시 시산표(없으면 계정 단위 재무제표)와 전기 재무제표(시작자료.json prior_fs ·
    전기 감사보고서 요약)에서 계정 잔액이 0 이 아닌 영역. 당기 시산표를 못 읽으면 None(잔액을 모른다)."""
    from . import figures
    from .analysis import defs as adefs

    d = Path(d)
    idx = d / "_색인"
    key = (str(d),) + tuple((p.name, p.stat().st_mtime_ns) for p in (idx / "문서목록.json", idx / "시작자료.json") if p.exists())
    if key in _BAL_CACHE:
        return _BAL_CACHE[key]
    try:
        rows = figures.current_tb(d)
    except Exception:  # noqa: BLE001  시산표를 못 읽으면 잔액을 모르는 것으로 둔다
        rows = []
    has: dict[str, bool] | None = None
    if rows:
        has = {}
        for r in rows:
            name = str(r.get("name") or "").strip()
            if re.fullmatch(r"[<\[【(（].*[>\]】)）]", name):      # 「<재고자산>」 같은 묶음 머리 줄은 계정이 아니다
                continue
            a = adefs.map_account(r.get("code"), name).get("area")
            if a and r.get("amount"):
                has[a] = True
        st = folder._read(idx / "시작자료.json", {}) or {}
        pf = st.get("prior_fs") or {}
        for name, v in {**(pf.get("bs") or {}), **(pf.get("is") or {})}.items():
            a = adefs.map_account(None, str(name)).get("area")
            if a and v:
                has[a] = True
        for name, v in ((st.get("prior_report") or {}).get("figures") or {}).items():
            a = adefs.map_account(None, str(name)).get("area")
            if a and v and name not in ("자산총계", "부채총계", "자본총계"):
                has[a] = True
    _BAL_CACHE[key] = has
    return has


# ---------------------------------------------------------------- 요청자료 목록
def _pbc_file(c: _Case) -> Path | None:
    w = next((w for w in c.wps if w["no"] == "PBC" and w.get("file")), None)
    if w and (c.d / w["file"]).exists():
        return c.d / w["file"]
    cands = sorted(p for p in (c.d / "1000 감사계약").glob("PBC_*.xlsx") if not p.name.startswith("~$"))
    return cands[0] if cands else None


def read_pbc(path: Path) -> list[dict]:
    """PBC 엑셀의 요청 행. 옛 양식은 머리줄에 '요청 자료' 칸이 있고 # 칸이 숫자인 줄만 요청 항목이다.
    실제 요청자료 양식(2026-10-13, audit/pbcform.py)은 머리줄에 '요청자료명' 과 '순번' 이 있고 요청세부내역이 있으면 그것이 항목 문구다.
    그 양식의 '엔진 번호' 칸은 pbc_match.yaml row 번호다(없으면 낱말로 찾는다)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            head = next((n for n, r in enumerate(rows) if any(isinstance(v, str) and v.strip() == "요청 자료" for v in r)), None)
            new_form = False
            if head is None:
                head = next((n for n, r in enumerate(rows) if any(isinstance(v, str) and v.strip() == "요청자료명" for v in r)
                             and any(isinstance(v, str) and v.strip() == "순번" for v in r)), None)
                new_form = head is not None
            if head is None:
                continue
            col = {str(v).strip(): n for n, v in enumerate(rows[head]) if v is not None}
            out = []
            seq = 0
            for r in rows[head + 1:]:
                def cell(name):
                    n = col.get(name)
                    return r[n] if n is not None and n < len(r) else None
                if new_form:
                    name, detail = cell("요청자료명"), cell("요청세부내역")
                    if not (isinstance(name, str) and name.strip()):
                        continue
                    seq += 1
                    no = cell("순번")
                    eng = cell("엔진 번호")
                    item = (detail.strip() if isinstance(detail, str) and detail.strip() else name.strip())
                    out.append({"no": int(no) if isinstance(no, (int, float)) and not isinstance(no, bool) else seq,
                                "category": cell("구분"), "item": item, "name": name.strip(), "workpaper": None, "form": "new",
                                "engine": int(eng) if isinstance(eng, (int, float)) and not isinstance(eng, bool) else None})
                    continue
                no, item = cell("#"), cell("요청 자료")
                if isinstance(no, (int, float)) and not isinstance(no, bool) and isinstance(item, str) and item.strip():
                    wp = cell("관련 조서")
                    out.append({"no": int(no), "category": cell("분류"), "item": item.strip(), "workpaper": str(wp).strip() if wp else None})
            return out
    finally:
        wb.close()
    raise DefinitionError(f"요청자료 목록에서 '요청 자료' 머리줄을 찾지 못했습니다: {path}")


def _pbc_rows(c: _Case) -> list[tuple[dict, dict | None]]:
    """(요청 행, 대응 표 항목). PBC 파일이 없으면 대응 표의 표준 18건을 쓴다.
    새 양식은 한 대응 항목에 여러 줄이 닿을 수 있어 첫 줄만 판정에 쓰고, 대응 표에 없는 줄은 판정하지 않는다."""
    table = c.match.get("pbc", [])
    p = _pbc_file(c)
    if p is None:
        return [({"no": e["row"], "item": e["item"], "workpaper": e.get("workpaper")}, e) for e in table]
    c.pbc_path = p
    out, used = [], set()
    for row in read_pbc(p):
        if row.get("form") == "new":
            e = next((e for e in table if e["row"] == row["engine"]), None) if row.get("engine") is not None else                 next((e for e in table if all(w in row["name"] for w in e["match"])), None) or                 next((e for e in table if all(w in row["item"] for w in e["match"])), None)
            if e is not None and e["row"] in used:
                e = None                              # 같은 자료를 가리키는 뒷 줄은 판정에 다시 쓰지 않는다
            elif e is not None:
                used.add(e["row"])
                row = {**row, "item": e["item"]}      # 빠진 자료 문구는 표준 18건 문구를 그대로 쓴다(PBC 파일이 있든 없든 같다)
            out.append((row, e))                      # e 가 없는 줄은 엔진이 받았는지 보지 않는다(양식의 「엔진 판정」 칸에 그렇게 적는다)
            continue
        e = next((e for e in table if all(w in row["item"] for w in e["match"])), None)
        if e is None:
            c.warnings.append(f"요청자료 대응 표(pbc_match.yaml)에 없는 항목이라 받았는지 보지 않았습니다: {row['no']}. {row['item']}")
        out.append((row, e))
    return out


# ---------------------------------------------------------------- 단계
def _stage_of(c: _Case) -> str:
    if c._stage:
        return c._stage
    wps = c.workpapers()
    status = {w["no"]: w["status"] for w in wps}
    result = STAGES[-1]
    for st in c.rules["stages"]:
        lw = st.get("leave_when")
        if not lw:
            result = st["name"]
            break
        targets = [no for no in lw.get("workpapers", []) if no in status]
        for fo in lw.get("folders", []):
            targets += [w["no"] for w in wps if w["folder"] == fo]
        if not all(status[no] in c.done for no in targets):
            result = st["name"]
            break
    c._stage = result
    return result


def _wp_stage(c: _Case, w: dict) -> int:
    """조서가 속한 단계 번호. stages 의 leave_when 에 이름이 있거나 폴더가 맞는 첫 단계, 없으면 마지막(종결)."""
    for i, st in enumerate(c.rules["stages"]):
        lw = st.get("leave_when") or {}
        if w["no"] in [str(x) for x in lw.get("workpapers", [])] or w.get("folder") in lw.get("folders", []):
            return i
    return len(c.rules["stages"]) - 1


def stage(engagement_dir: Path) -> str:
    return _stage_of(_Case(engagement_dir))


# ---------------------------------------------------------------- 빠진 자료
def _missing_of(c: _Case) -> list[dict]:
    st = STAGES.index(_stage_of(c))
    prev = {}
    for m in c.docs.get("missing", []):
        base = base_item(m["item"])
        if m.get("requested_at"):
            for k in (m["item"], base):        # 기간 문구를 바꾼 글(period_words)도 원래 문구로 찾게 둘 다 둔다
                prev.setdefault(k, m["requested_at"])
                prev.setdefault(_unperiod(k, c.fy), m["requested_at"])

    out, seen, superseded, pbc_text, pbc_line = [], set(), set(), {}, {}
    rows = _pbc_rows(c)
    for row, e in rows:
        if e is not None:
            pbc_text[e["row"]] = row["item"]
    needs = c.match.get("issue_needs") or {}
    done_need = set()
    for i in c.open_issues():
        wp = i.get("response_wp")
        for n in needs.get(wp, []) if wp else []:
            if (wp, n["item"]) in done_need:
                continue
            done_need.add((wp, n["item"]))
            parties = issue_parties(c.d, i["no"]) if n.get("issue_parties") else []
            if parties:                 # 이슈를 낳은 거래의 상대방(분석.json 발견 사실)이 있으면 그 이름이 든 자료만 센다
                n = {**n, "keywords": parties}
            ok, lack = c.judge(n)
            if ok:
                continue
            refs = n.get("pbc") or []
            superseded.update(refs)
            item = issue_item(n["item"], i, wp, parties)
            asked = prev.get(item) or next((prev[pbc_text[r]] for r in refs if pbc_text.get(r) in prev), None)
            out.append({"item": _item_text(item, lack), "workpaper": wp, "issue": i["no"], "requested_at": asked})
            seen.add(item)
    nobal = c.no_balance()
    for row, e in rows:
        if e is None or e["row"] in superseded or row["item"] in seen:
            continue
        if e.get("from_stage") and st < STAGES.index(e["from_stage"]):
            continue
        if e.get("area") and str(e["area"]) in nobal:      # 잔액 · 거래가 없는 영역의 요청은 빠진 자료로 세지 않는다(2026-10-18)
            continue
        ok, lack = c.judge(e)
        if ok:
            continue
        seen.add(row["item"])
        m = {"item": _item_text(row["item"], lack), "issue": None, "requested_at": prev.get(row["item"])}
        wp = row.get("workpaper") or e.get("workpaper")
        if wp:
            m["workpaper"] = str(wp)
        out.append(m)
        pbc_line[e["row"]] = m
    # 분석적 검토가 대사하지 못해 남긴 자료 요청(분석.json requests, 회사 몫)을 잇는다(2026-10-18)
    for rq in analysis_requests(c, to="회사"):
        met = request_met(c, rq)
        if met is True:
            continue
        text = rq.get("ask") or rq["what"]
        if isinstance(met, list) and met:
            text += f" ({WAIT}: {', '.join(met[:3])}{' 외' if len(met) > 3 else ''})"
        m = pbc_line.get(rq.get("pbc"))
        if m is not None:
            if text not in m.setdefault("detail", []):
                m["detail"].append(text)
            m["rules"] = list(dict.fromkeys((m.get("rules") or []) + rq.get("rules", [])))
            continue
        if text in seen:
            continue
        seen.add(text)
        out.append({"item": text, "workpaper": str(rq.get("workpaper") or "2120A"), "issue": None, "requested_at": prev.get(text),
                    "source": "분석", "rules": list(rq.get("rules") or [])})
    # 분석적 검토는 감사범위와 상관없이 전기 · 당기 두 해로 한다. 손익을 계정 단위로 견줄 원천이 없으면 요청한다(R7)
    if c.missing_prior_pl() and PRIOR_PL_ITEM not in seen:
        out.append({"item": PRIOR_PL_ITEM, "workpaper": "2120A", "issue": None, "requested_at": prev.get(PRIOR_PL_ITEM)})
    for m in out:
        m["item"] = period_words(m["item"], c.fy)
    return out


def _unperiod(text: str, fy) -> str:
    if not fy:
        return text
    s, e = fy
    return text.replace(f"{s.month}~{e.month}월", "1~12월") if (e.year - s.year) * 12 + e.month - s.month + 1 < 12 else text


def period_words(text: str, fy) -> str:
    """요청 문구의 「(1~12월)」 을 감사 기간의 달로 바꾼다(1 ~ 9월 중간감사 · 짧은 사업연도, 2026-10-19). 12개월이면 그대로."""
    if not fy or "1~12월" not in text:
        return text
    s, e = fy
    months = (e.year - s.year) * 12 + e.month - s.month + 1
    if months >= 12:
        return text
    return text.replace("1~12월", f"{s.month}~{e.month}월")


def analysis_requests(c: "_Case", to: str | None = None) -> list[dict]:
    """분석.json requests. to 를 주면 그 몫만(회사 · 회계사). to 가 없는 옛 결과는 회사 몫으로 본다."""
    an = folder._read(c.d / "_색인" / "분석.json", {}) or {}
    reqs = an.get("requests") or []
    return [r for r in reqs if to is None or (r.get("to") or "회사") == to]


def analysis_checks(c: "_Case") -> list[str]:
    """회계사가 확인할 분석 값: OCR(kordoc) 값 확인 필요(분석.json ocr_checks)와 회계사 몫 요청."""
    an = folder._read(c.d / "_색인" / "분석.json", {}) or {}
    def page(p):
        return f" ({p}쪽)" if isinstance(p, int) else (f" ({p})" if p else "")

    out = [f"OCR 값 확인 필요: {Path(x['path']).name} " + x["what"] + page(x.get("page")) for x in an.get("ocr_checks") or []]
    out += [r["what"] for r in analysis_requests(c, to="회계사")]
    return out


def _need_spec(c: "_Case", n: dict) -> dict | None:
    """요청이 풀리는 조건 하나를 pbc_match 판정 형식으로. {'pbc': 행} 은 그 요청자료 줄 판정을 쓰고, kinds 를 주면 그 종류로 좁힌다."""
    spec: dict = {}
    if n.get("pbc") is not None:
        e = next((x for x in c.match.get("pbc", []) if x["row"] == n["pbc"]), None)
        if e is None:
            return None
        spec = {k: v for k, v in e.items() if k in ("kinds", "keywords", "need", "any_of", "read_as", "unread_name", "period",
                                                    "period_words", "dated_in_year")}
    if n.get("kinds"):
        spec["kinds"] = list(n["kinds"])
        spec["need"] = "any"
        per = spec.get("period")
        if isinstance(per, dict):
            per = {k: v for k, v in per.items() if k in spec["kinds"]}
            spec["period"] = per or None
        spec.pop("any_of", None)
    for k in ("keywords", "dated_after", "period"):
        if n.get(k) is not None:
            spec[k] = n[k]
    return spec


CONTRACT_RE = re.compile(r"(여신\s*거래\s*약정|대출\s*약정|차입\s*약정|한도\s*약정|금전\s*소비\s*대차)")
WAIT = "받음, 판독 대기"


def read_fact(f: dict, what: str, fe: date | None) -> bool:
    """그 파일에 요청 종류에 맞는 판독 사실이 있는지(2026-10-18). 엔진 · AI 사실을 모두 본다.
      잔액증명 결산일 · 주주명부 결산일   「기준일 …」 사실이나 기준일 · 현재 가 붙은 문장의 날짜가 결산일
      약정서                         엔진 사실 「약정서 종류 약정」 이나 여신 · 대출 · 차입 약정 낱말이 든 AI 사실"""
    facts = f.get("facts") or []
    if what == "약정서":
        return any((x.get("basis") == "엔진" and x["text"].startswith("약정서 종류 약정")) or
                   (x.get("basis") == "AI" and CONTRACT_RE.search(x["text"])) for x in facts)
    if fe is None:
        return False
    from .analysis import inst

    for x in facts:
        t = x.get("text") or ""
        if t.startswith("기준일 ") or (x.get("basis") == "AI" and ("기준일" in t or "현재" in t)):
            if any(dt == fe for dt in inst.dates(t)):
                return True
    return False


def read_any(f: dict, what: str) -> bool:
    """그 종류의 판독 사실이 하나라도 있는지(맞든 틀리든). 없으면 판독 전이다."""
    facts = f.get("facts") or []
    if what == "약정서":
        return any(x["text"].startswith("약정서 종류 ") or x.get("basis") == "AI" for x in facts)
    return any(x["text"].startswith("기준일 ") or (x.get("basis") == "AI" and ("기준일" in x["text"] or "현재" in x["text"])) for x in facts)


def request_met(c: "_Case", rq: dict) -> bool | list[str]:
    """분석 요청이 풀렸는지. need 의 조건마다 받은자료가 맞아야 한다. new_since 는 분석 때 없던 파일만, institution 은 그 금융기관 파일만 센다.
    read 가 있으면 그 파일에 맞는 판독 사실이 있어야 풀린다. 파일은 맞는데 판독 사실이 없으면 그 파일 이름 목록(받음, 판독 대기)을 돌려준다.
    need 가 없는 요청은 분석을 다시 돌릴 때까지 남는다."""
    need = rq.get("need")
    if not need:
        return False
    seen = set(rq.get("seen") or [])
    waiting: list[str] = []
    for n in need:
        spec = _need_spec(c, n)
        if spec is None:
            return False
        files = c.files
        if n.get("new_since"):
            files = [f for f in files if (f.get("sha256") or "")[:16] not in seen]
        if n.get("institution"):
            from .analysis import inst

            k = inst.key(n["institution"])
            files = [f for f in files if inst.key(inst.name(Path(f["path"]).stem, file=True)) == k]
        if not files:
            return False
        if n.get("read"):
            ok = [f for f in files if read_fact(f, n["read"], c.fy[1] if c.fy else None)]
            if ok and _judge(spec, ok, c.fy)[0]:
                continue
            kinds = set(spec.get("kinds") or [])
            # 판독 대기: 분석 뒤 새로 들어온 그 종류 파일 가운데 판독 사실이 아직 없는 것
            cand = [f for f in files if f not in ok and seen and (f.get("sha256") or "")[:16] not in seen
                    and (not kinds or f["kind"] in kinds) and not read_any(f, n["read"])]
            if cand and _judge(spec, cand, None)[0]:
                waiting += [Path(f["path"]).name for f in cand]
                continue
            return False
        if not _judge(spec, files, c.fy)[0]:
            return False
    return sorted(set(waiting)) or True


def item_text(m: dict) -> str:
    """빠진 자료 한 줄 글. 분석이 붙인 구체 문구(detail)를 뒤에 붙인다(지휘 · 요청 메일 · PBC 가 같이 쓴다)."""
    return m["item"] + (": " + "; ".join(m["detail"]) if m.get("detail") else "")


def issue_parties(d: Path, no: int) -> list[str]:
    """이슈를 낳은 발견 사실(_색인/분석.json findings 의 metric counterparty)의 거래 상대방 이름. 없으면 빈 목록.
    정의 파일에 거래처 · 은행 이름을 두지 않고 이슈의 실제 값으로 요청 글과 판정 낱말을 채운다(2026-10-18).
    의사록 · 약정 판독의 상대방(counterparty)만 쓴다. 분개장 거래처 칸(party)은 「(기업013)외화_USD」 같은 내부 계좌 이름이 섞여
    회사에 가는 메일에 넣지 않는다(실제 회사 자료 에서 확인). 세 곳까지만 쓴다."""
    an = folder._read(Path(d) / "_색인" / "분석.json", {}) or {}
    ids = next((x.get("findings") or [] for x in an.get("issues") or [] if x.get("issue") == no), [])
    out = []
    for f in an.get("findings") or []:
        if f.get("id") in ids:
            m = f.get("metric") or {}
            v = m.get("counterparty")
            if isinstance(v, str) and v.strip() and v.strip() not in out:
                out.append(v.strip())
    return out[:3]


def issue_item(template: str, issue: dict, wp: str | None, parties: list[str] | None = None) -> str:
    """이슈 대응 자료 글. {account} 를 이슈 계정 이름으로(없으면 대응 조서 영역 이름), {party} 를 이슈 거래 상대방으로 채운다.
    거래처 · 은행 이름은 정의에 두지 않는다."""
    if "{" not in template:
        return template
    account = issue.get("account")
    if not account:
        from .analysis import defs as adefs

        w = str(wp or "")
        account = adefs.area_name(w[4:]) if w.startswith("4000") and w[4:] else "해당 계정"
    party = f", {' · '.join(parties)}" if parties else ""
    try:
        return template.format(account=account, party=party)
    except (KeyError, IndexError, ValueError):
        return template


def missing(engagement_dir: Path) -> list[dict]:
    return _missing_of(_Case(engagement_dir))


# ---------------------------------------------------------------- 다음 할 일
def _more(items: list[str], n: int = 3) -> str:
    s = " · ".join(items[:n])
    return s + (f" 외 {len(items) - n}건" if len(items) > n else "")


def _skill_cfg(c: _Case, no: str) -> dict | None:
    skills = c.rules.get("skills") or {}
    if no in skills:
        return skills[no]
    return skills.get(AREA_DEFAULT) if folder._is_area(no) else None


def _needs_met(c: _Case, cfg: dict) -> list[str] | None:
    """필요 자료가 모두 있으면 맞은 종류 이름들, 하나라도 없으면 None."""
    kinds = {f["kind"] for f in c.files}
    got = []
    for n in cfg.get("needs", []):
        alts = n if isinstance(n, list) else [n]
        hit = next((k for k in alts if k in kinds), None)
        if hit is None:
            return None
        got.append(hit)
    return got


def _actions(c: _Case, miss: list[dict], full: bool = False, through: str | None = None) -> list[dict]:
    rules = sorted(c.rules["rules"], key=lambda r: r["priority"])
    wps = c.workpapers()
    out = []

    def add(r, workpaper=None, title_key="title", who_override=None, detail_key="detail", **kw):
        a = {"who": who_override or r["who"], "title": r[title_key].format(**kw)}
        if r.get(detail_key):
            a["detail"] = r[detail_key].format(**kw).strip()
        if (kw.get("skill") or r.get("skill")) and not who_override:
            a["skill"] = kw.get("skill") or r["skill"]
        if workpaper:
            a["workpaper"] = workpaper
            who = (c.assign.get(workpaper) or {}).get("reviewer" if r["id"] == "검토요청" else "to")
            if who:
                a["assignee"] = who            # 관련 조서의 담당자. 검토 요청 줄은 검토자
        out.append(a)

    def label(w):
        return {"no": w["no"], "name": w.get("name") or c.names.get(w["no"], w["no"])}

    per_rule = int(c.rules.get("max_per_rule", 3))

    def fold(start):
        """조서마다 한 줄인 규칙이 목록을 다 차지하지 않게 per_rule 줄까지만 두고 나머지는 마지막 줄에 묶는다."""
        extra = out[start + per_rule:]
        if extra:
            del out[start + per_rule:]
            last = out[-1]
            nos = [n for x in extra for n in (x.get("_nos") or [x.get("workpaper", "")])]
            more = f"같은 종류의 할 일이 {len(extra)}건 더 있습니다: " + _more(nos, 8)
            last["detail"] = f"{last['detail'].rstrip('.')}. {more}" if last.get("detail") else more

    for r in rules:
        rid = r["id"]
        start = len(out)
        if rid == "확인-판독":
            unc = [f for f in c.docs["files"] if f["role"] == "received" and not f["confirmed"]]
            if not unc:
                continue
            # 종류를 아직 못 정한 파일과, AI 가 종류 · 기간 · 사실을 넣어 확인을 기다리는 파일을 나눠 센다
            undecided = sum(1 for f in unc if "판독 필요" in f["kind"] or f["kind"] == "알 수 없음")
            failed = sum(1 for f in unc if f.get("read_error"))      # 종류는 알지만 엔진이 판독하지 못한 파일(R8)
            ai = len(unc) - undecided - failed
            parts = ([f"AI 판독 {ai}건"] if ai else []) + ([f"종류 판독 필요 {undecided}건"] if undecided else []) +                     ([f"엔진 판독 실패 {failed}건"] if failed else [])
            linked = {x for f in unc for x in f.get("linked_workpapers") or []}
            add(r, workpaper=next(iter(linked)) if len(linked) == 1 else None, count=len(unc),
                breakdown=" · ".join(parts), examples="예) " + _more([Path(f["path"]).name for f in unc]))
        elif rid == "요청-빠진자료":
            if not miss:
                continue
            todo = [m for m in miss if not m.get("requested_at")]
            if _stage_of(c) == "종결" and r.get("title_closed"):
                add(r, title_key="title_closed", who_override="회계사", count=len(miss), items=_more([m["item"] for m in miss]))
            elif todo:
                add(r, count=len(todo), items=_more([item_text(m) for m in todo]))
            else:
                add(r, title_key="title_again", count=len(miss), items=_more([item_text(m) for m in miss]))
        elif rid == "확인-분석값":
            # 2026-10-18. OCR(kordoc) 글에서 읽은 값과 회계사 몫 분석 요청(약정서 금액을 못 읽은 차입처 · 기준일 확인 등). 회사에 묻지 않는다
            items = analysis_checks(c)
            if items:
                add(r, workpaper="2120A", count=len(items), items=_more(items, 3))
        elif rid == "다시확인":
            for w in wps:
                if w["status"] == "작성중" and w.get("needs_recheck"):
                    add(r, workpaper=w["no"], basis=w.get("status_basis") or "다시 확인 표시", **label(w))
        elif rid == "기본값확인":
            an = folder._read(c.d / "_색인" / "분석.json", {}) or {}
            keys = an.get("default_inputs") or []
            if not keys:
                continue
            names = {"materiality.overall": "전체 중요성", "materiality.performance": "수행중요성", "materiality.trivial": "명백히 사소한 금액",
                     "analytics.sig_amount": "유의성 기준 금액", "analytics.sig_ratio": "유의성 기준 비율",
                     "analytics.tolerable_diff": "설명 후 허용 차이"}
            vals = []
            for k in keys:
                v = ((an.get("inputs") or {}).get(k) or {}).get("value")
                vals.append(f"{names.get(k, k)} {v:,.0f}" if isinstance(v, (int, float)) and v >= 1 else
                            f"{names.get(k, k)} {v:.0%}" if isinstance(v, float) else names.get(k, k))
            add(r, workpaper="2700A", items=" · ".join(vals))
        elif rid == "수정반영":
            # 회계사가 적은 수정(8100A · 8100B)이나 새 회사제시 판을 엔진 문서가 아직 담지 않았다(2026-10-12, audit/apply.py)
            from . import apply

            try:
                trig = apply.pending(c.d)["triggers"]
            except Exception:  # noqa: BLE001  8100A 를 다른 꼴로 저장했거나 시산표를 못 읽으면 이 줄을 내지 않는다
                trig = []
            if trig:
                add(r, workpaper="8100A", items=_more(trig, 2))
        elif rid == "보고준비":
            # DART 제출 준비(2026-10-06). 보고.json 의 dart_ready 를 report make · check · dart 가 쓴다
            rr = (folder._read(c.d / "_색인" / "보고.json", {}) or {}).get("dart_ready")
            if not rr:
                continue
            probs = rr.get("opinion_problems") or []
            if rr.get("opinion") is None and probs == ["감사의견 종류를 고르지 않았습니다"]:
                add(r, title_key="title_opinion", detail_key="detail_opinion", workpaper="8600", items="")
            elif probs:
                add(r, title_key="title_opinion_fix", detail_key="detail_opinion", workpaper="8600", items="; ".join(probs))
            if rr.get("notes_open"):
                add(r, title_key="title_notes", detail_key="detail_notes", workpaper="8400",
                    items=", ".join(str(n) for n in rr["notes_open"]))
            if rr.get("ready") and rr.get("html_warning"):
                add(r, title_key="title_rebuild", detail_key="detail_rebuild", workpaper="DART")
        elif rid == "스킬실행":
            # 스킬 하나가 조서 여럿을 한 번에 만든다(감사-시작자료 → PBC · 1100 · 1200 · 2100). 스킬마다 한 줄로 묶는다
            # 지금 단계와 다음 단계의 조서만 낸다(준비 단계에 종결 점검표를 권하지 않는다)
            now_i = STAGES.index(_stage_of(c))
            last_i = max(now_i + 1, STAGES.index(through)) if through else now_i + 1
            groups: dict = {}
            for w in wps:
                if _wp_stage(c, w) > last_i:
                    continue
                cfg = _skill_cfg(c, w["no"]) if w["status"] == "없음" else None
                got = _needs_met(c, cfg) if cfg else None
                if got is not None:
                    g = groups.setdefault(cfg["skill"], {"wps": [], "needs": []})
                    g["wps"].append(label(w))
                    g["needs"] += [x for x in got if x not in g["needs"]]
            for skill, g in groups.items():
                first = g["wps"][0]
                if len(g["wps"]) == 1:
                    kw = first
                else:
                    kw = {"no": " · ".join(x["no"] for x in g["wps"]), "name": f"({len(g['wps'])}개)"}
                needs = (f"필요한 자료가 모두 있습니다({' · '.join(g['needs'])})" if g["needs"] else "받은자료 없이 만들 수 있습니다")
                add(r, workpaper=first["no"], skill=skill, needs=needs, **kw)
                out[-1]["_nos"] = [x["no"] for x in g["wps"]]
        elif rid == "전기이어받기":
            # 2100 「전기에서 이어진 사항」 가운데 당기 이슈로 볼지 판단 칸이 빈 것(2100 이 없으면 전부)
            carry = (c.prior_source() or {}).get("carry") or []
            if not carry:
                continue
            open_ = _carry_open(c, len(carry))
            if open_:
                add(r, workpaper="2100", count=len(open_), items=_more([f"{carry[k - 1]['what']}: {carry[k - 1]['text']}" for k in open_], 2))
        elif rid == "초도기초잔액":
            why = _first_year_alt(c)
            if why:
                made = any(w["no"] == "3800" and w.get("file") for w in c.wps)
                add(r, workpaper="3800", why=why, make="" if made else ". 3800 이 없으면 python -m audit wp make <감사 폴더> 3800 으로 먼저 만듭니다")
        elif rid == "검토요청":
            for w in wps:
                if w["status"] == "검사 통과" and not w.get("needs_recheck"):
                    add(r, workpaper=w["no"], **label(w))
        elif rid == "작성":
            for w in wps:
                if w["status"] in ("초안", "작성중") and not w.get("needs_recheck"):
                    add(r, workpaper=w["no"], status=w["status"], **label(w))
        if not full and rid in ("다시확인", "스킬실행", "검토요청", "작성"):
            fold(start)
    if full:                                 # 실행기(runner)가 쓴다. 묶지 않고 줄이지 않으며 스킬 줄의 조서번호를 남긴다
        for a in out:
            if "_nos" in a:
                a["nos"] = a.pop("_nos")
        return out
    for a in out:
        a.pop("_nos", None)
    return out[: int(c.rules.get("max_actions", 8))]


def _readback(c: _Case, no: str) -> dict:
    """조서목록 파일의 기입란 값 {(시트, 항목): 값}. 없거나 못 읽으면 빈 dict."""
    from . import wp

    w = next((w for w in c.wps if w["no"] == no and w.get("file")), None)
    p = c.d / w["file"] if w else None
    if not p or not p.exists() or p.suffix.lower() != ".xlsx":
        return {}
    try:
        return {(x["sheet"], x["label"]): x["value"] for x in wp.readback(p)}
    except Exception:  # noqa: BLE001  회계사가 다른 꼴로 저장한 파일
        return {}


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _carry_open(c: _Case, n: int) -> list[int]:
    """2100 「전기에서 이어진 사항」에서 판단 칸(당기 이슈 여부)이 빈 번호들."""
    vals = _readback(c, "2100")
    return [k for k in range(1, n + 1) if _blank(vals.get(("전기에서 이어진 사항", f"이어진 사항 {k} 당기 이슈 여부")))]


def _first_year_alt(c: _Case) -> str | None:
    """초도감사이고 기초잔액 대체절차가 필요한데(전기조서 없음 · 전기 미감사) 3800 대체절차 칸이 비었으면 그 까닭."""
    if c.case.get("engagement_kind") != "초도감사":
        return None
    ps = c.prior_source()
    if not ps:
        return None
    vals = _readback(c, "3800")
    audited = ps.get("prior_audited")
    if audited is None and str(vals.get(("3800", "전기 재무제표 감사 여부")) or "").strip() == "아니오":
        audited = False
    if ps.get("situation") == "전기조서 없음":
        why = "초도감사인데 전기조서가 없습니다"
    elif audited is False:
        why = "초도감사인데 전기 재무제표가 감사받지 않았습니다"
    else:
        return None
    return why if _blank(vals.get(("3800", "기초잔액 대체절차"))) else None


def next_actions(engagement_dir: Path) -> list[dict]:
    c = _Case(engagement_dir)
    return _actions(c, _missing_of(c))


def plan(engagement_dir: Path, through: str | None = None) -> list[dict]:
    """다음 할 일 전부(묶지 않고 줄이지 않음). 스킬 줄에는 그 스킬이 만들 조서번호 nos 가 붙는다. _색인 에 쓰지 않는다.
    through 를 주면 그 단계의 조서까지 스킬 줄을 낸다(지휘는 지금 단계와 다음 단계만 낸다). 초안을 앞당겨 만들 때 실행기가 쓴다."""
    if through is not None and through not in STAGES:
        raise DefinitionError(f"모르는 단계입니다: {through} ({' · '.join(STAGES)})")
    c = _Case(engagement_dir)
    return _actions(c, _missing_of(c), full=True, through=through)


# ---------------------------------------------------------------- 이슈 닫기
CLOSE_METHOD = {"수정": "수정분개", "미수정": "미수정 왜곡표시", "왜곡 아님": "설명 종결"}


def _close_issues(c: "_Case", d: Path) -> list[int]:
    """8100A 에 회계사 처리(수정 · 미수정 · 왜곡 아님)가 적혀 있고, 대응 조서와 8100A 가 모두 검토 끝인 열린 이슈를 닫는다.
    닫는 방법은 처리에서 정하고 결론(conclusion)은 회계사 몫이라 비워 둔다(2026-10-05 통합)."""
    rep = folder._read(d / "_색인" / "보고.json", {})
    dec = {x["issue"]: x.get("decision") for x in rep.get("decisions", []) if x.get("decision")}
    # 조서목록에 실제로 있는 조서만 본다(c.workpapers() 는 아직 없는 이슈 대응 조서를 '없음' 으로 더한다)
    listed = folder._read(d / "_색인" / "조서목록.json", {"workpapers": []})["workpapers"]
    status = {w["no"]: w["status"] for w in listed}
    if status.get("8100A") != "검토 끝":
        return []
    closed = []
    for i in c.issues:
        if i["status"] != "열림" or i["no"] not in dec or dec[i["no"]] not in CLOSE_METHOD:
            continue
        resp = i.get("response_wp")
        if resp and resp in status and status[resp] != "검토 끝":
            continue
        i["status"] = "닫힘"
        i["close_method"] = CLOSE_METHOD[dec[i["no"]]]
        closed.append(i["no"])
    if closed:
        data = folder._read(d / "_색인" / "이슈.json", {"format_version": 1, "issues": []})
        data["issues"] = c.issues
        folder._write(d / "_색인" / "이슈.json", data)
    return closed


# ---------------------------------------------------------------- 쓰기
def run(engagement_dir: Path, now: str | None = None) -> dict:
    """빠진 자료 · 다음 할 일 · 단계를 _색인 에 쓰고 기록을 남긴다."""
    c = _Case(engagement_dir)
    at = now or folder._now()
    closed = _close_issues(c, c.d)
    miss = _missing_of(c)
    acts = _actions(c, miss)
    st = _stage_of(c)
    idx = c.d / "_색인"

    docs = c.docs
    docs["missing"] = miss
    warns = docs.setdefault("warnings", [])
    for w in c.warnings:
        if w not in warns:
            warns.append(w)
    folder._write(idx / "문서목록.json", docs)

    case = c.case
    before = case.get("stage")
    case["stage"] = st
    case["next_actions"] = acts
    what = f"단계 {st}. 빠진 자료 {len(miss)}건, 다음 할 일 {len(acts)}건을 정했습니다"
    if closed:
        what += f". 8100A 처리와 검토가 끝난 이슈 {len(closed)}건을 닫았습니다({', '.join(map(str, closed))})"
    if before and before != st:
        what += f". 단계가 {before} 에서 {st} 가 됐습니다"
    read = ["_색인/문서목록.json", "_색인/조서목록.json", "_색인/이슈.json"] + (["_색인/보고.json"] if closed else [])
    if c.pbc_path:
        read.append(c.pbc_path.relative_to(c.d).as_posix())
    case.setdefault("log", []).insert(0, {"at": at, "who": "지휘", "what": what, "read": read,
                                          "wrote": ["_색인/감사건.json", "_색인/문서목록.json"] + (["_색인/이슈.json"] if closed else [])})
    folder._write(idx / "감사건.json", case)
    folder.validate_index(c.d)
    return {"stage": st, "missing": miss, "next_actions": acts, "warnings": c.warnings}

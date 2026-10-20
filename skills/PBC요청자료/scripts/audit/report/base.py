"""보고 단계 공통: 정의 · 색인 읽기 · 숫자 표기 · 기입란 이어받기 · 워드와 HTML 블록 (R4 report, 2026-10-05).

워드 · HTML 문서는 블록 목록으로 다룬다. 블록은 ("h", 단계, 글자) · ("p", 글자) · ("cap", 글자) · ("table", 행 목록) 이다.
엔진이 만든 문서에는 글자 해시를 붙여 둔다(워드는 문서 속성 메모, HTML 은 meta). 해시가 맞으면 엔진이 만든 그대로라
다시 만들어도 되고, 다르면 회계사가 고친 문서이므로 덮어쓰지 않는다.
"""
from __future__ import annotations

import hashlib
import json
import re
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

import yaml

from .. import wp

DEFS = Path(__file__).resolve().parents[2] / "definitions"
MARK = "【기입"
STAMP = "audit-engine:"


@lru_cache(maxsize=1)
def defs() -> dict:
    return yaml.safe_load((DEFS / "report_fs.yaml").read_text(encoding="utf-8"))


def read_json(p: Path, default=None):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def case(d: Path) -> dict:
    return read_json(Path(d) / "_색인" / "감사건.json", {}) or {}


def company_name(d: Path) -> str:
    c = case(d).get("company") or {}
    return c.get("name") if isinstance(c, dict) else str(c or "")


def issues(d: Path) -> list[dict]:
    return (read_json(Path(d) / "_색인" / "이슈.json", {"issues": []}) or {}).get("issues", [])


def norm(s) -> str:
    return re.sub(r"\s+", "", str(s or ""))


# ---------------------------------------------------------------- 숫자 표기
def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    v = int(round(v))
    if v == 0:
        return "-"
    return f"({-v:,})" if v < 0 else f"{v:,}"


def parse_num(s):
    """문서 칸 글자를 숫자로. '-' 는 0, 괄호는 음수, 숫자가 아니면 None."""
    if s is None:
        return None
    if isinstance(s, bool):
        return None
    if isinstance(s, (int, float)):
        return int(round(s))
    t = str(s).strip().replace(",", "").replace("원", "").replace(" ", "")
    if t in ("-", "–", "−"):
        return 0
    if not t:
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    if t.startswith("−"):
        t = "-" + t[1:]
    try:
        x = float(t)
    except ValueError:
        return None
    x = -x if neg else x
    return int(round(x))


# ---------------------------------------------------------------- 기입란 이어받기
def prev_inputs(path: Path | None) -> dict:
    """앞서 만든 조서의 기입란 값 {(시트, 항목): 값}. 다시 만들 때 회계사 기입을 지키려고 쓴다."""
    if not path or not Path(path).exists():
        return {}
    try:
        return {(x["sheet"], x["label"]): x["value"] for x in wp.readback(Path(path))}
    except Exception:  # noqa: BLE001  깨진 파일이면 이어받지 않는다
        return {}


class RBook(wp.Book):
    """wp.Book 에 앞선 기입값 이어받기를 더한 책. 값을 주지 않은 기입란은 같은 시트 · 항목의 앞선 값으로 채운다."""

    def __init__(self, *a, keep: dict | None = None, **k):
        super().__init__(*a, **k)
        self.keep = keep or {}

    def ask(self, ws, row, col, label, rule="required", value=None, fmt=None, wrap=True, keep=True):
        if value is None and keep:
            value = self.keep.get((ws.title, label))
        return super().ask(ws, row, col, label, rule=rule, value=value, fmt=fmt, wrap=wrap)


def wp_file(d: Path, no: str, pattern: str | None = None) -> Path | None:
    """감사 폴더에서 조서 파일을 찾는다(판이 높은 것). pattern 은 glob."""
    from .. import folder

    try:
        base = wp.path_for(d, no)
    except ValueError:
        return None
    cands = sorted(base.parent.glob(pattern or f"{no}_*.xlsx"))
    if not cands:
        return None

    def key(p: Path):
        m = re.search(r"_v(\d+(?:\.\d+)+)$", p.stem)
        return folder._version_key(m.group(1) if m else None)

    return max(cands, key=key)


# ---------------------------------------------------------------- 블록 해시
def blocks_text(blocks: list) -> str:
    out = []
    for b in blocks:
        if b[0] == "table":
            out.append("t|" + "\x1f".join("\x1e".join(norm(c) for c in r) for r in b[1]))
        elif b[0] == "h":
            out.append(f"h{b[1]}|{norm(b[2])}")
        else:
            out.append(f"{b[0]}|{norm(b[1])}")
    return "\n".join(out)


def blocks_hash(blocks: list) -> str:
    return hashlib.sha256(blocks_text(blocks).encode("utf-8")).hexdigest()


def count_marks(blocks: list) -> int:
    n = 0
    for b in blocks:
        if b[0] == "table":
            n += sum(str(c).count(MARK) for r in b[1] for c in r)
        else:
            n += str(b[-1]).count(MARK)
    return n


# ---------------------------------------------------------------- 워드
def _font(run, size=None, bold=None, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = wp.FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), wp.FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _runs(par, text: str, size=10, bold=False):
    """【기입 …】 부분은 노란 바탕으로 칠한다."""
    from docx.enum.text import WD_COLOR_INDEX

    for i, part in enumerate(re.split(r"(【기입[^】]*】)", text)):
        if not part:
            continue
        r = par.add_run(part)
        _font(r, size=size, bold=bold)
        if part.startswith(MARK):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def write_docx(path: Path, blocks: list, title: str = "") -> Path:
    """블록을 워드로 쓴다. 표의 첫 행은 머리줄, 숫자처럼 보이는 칸은 오른쪽 정렬."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = wp.FONT
    st.font.size = Pt(10)
    for b in blocks:
        if b[0] == "h":
            par = doc.add_heading(level=min(max(b[1], 1), 4))
            _runs(par, b[2], size={1: 15, 2: 12, 3: 11}.get(b[1], 10), bold=True)
            for r in par.runs:
                _font(r, color=wp.NAVY)
        elif b[0] in ("p", "cap"):
            par = doc.add_paragraph()
            _runs(par, b[1], size=9 if b[0] == "cap" else 10, bold=b[0] == "cap")
        elif b[0] == "table":
            rows = b[1]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j in range(ncol):
                    v = r[j] if j < len(r) else ""
                    cell = t.cell(i, j)
                    par = cell.paragraphs[0]
                    _runs(par, str(v), size=9, bold=i == 0)
                    if i > 0 and j > 0 and parse_num(v) is not None:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = title
    doc.core_properties.author = "audit 엔진"
    doc.save(path)
    doc.core_properties.comments = STAMP + blocks_hash(read_docx(path))
    doc.save(path)
    return path


def read_docx(path: Path) -> list:
    """워드를 블록으로 읽는다. 빈 문단은 뺀다."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            t = p.text.strip()
            if not t:
                continue
            sn = p.style.name if p.style is not None else ""
            m = re.match(r"Heading (\d)", sn or "")
            if m:
                out.append(("h", int(m.group(1)), t))
            elif re.match(r"^표\s*\d+-\d+\s", t):
                out.append(("cap", t))
            else:
                out.append(("p", t))
        elif tag == "tbl":
            t = Table(el, doc)
            out.append(("table", [[c.text.strip() for c in r.cells] for r in t.rows]))
    return out


def docx_untouched(path: Path) -> bool:
    """엔진이 만든 그대로인 워드인지(해시 비교). 파일이 없으면 참."""
    from docx import Document

    p = Path(path)
    if not p.exists():
        return True
    try:
        c = Document(str(p)).core_properties.comments or ""
        return c.startswith(STAMP) and c[len(STAMP):] == blocks_hash(read_docx(p))
    except Exception:  # noqa: BLE001
        return False


# ---------------------------------------------------------------- HTML
class _Blocks(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.buf, self.tag, self.rows, self.row, self.cell, self.cls = [], [], None, None, None, None, ""

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("h1", "h2", "h3", "h4", "p") and self.rows is None:
            self.tag, self.buf, self.cls = tag, [], a.get("class") or ""
        elif tag == "table":
            self.rows = []
        elif tag == "tr" and self.rows is not None:
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            self.row.append(re.sub(r"\s+", " ", "".join(self.cell)).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag == "table" and self.rows is not None:
            self.out.append(("table", self.rows))
            self.rows = None
        elif tag == self.tag:
            t = re.sub(r"\s+", " ", "".join(self.buf)).strip()
            if t:
                if tag.startswith("h"):
                    self.out.append(("h", int(tag[1]), t))
                elif "cap" in self.cls or re.match(r"^표\s*\d+-\d+\s", t):
                    self.out.append(("cap", t))
                else:
                    self.out.append(("p", t))
            self.tag = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)
        elif self.tag:
            self.buf.append(data)


def read_html(path: Path) -> list:
    p = _Blocks()
    p.feed(Path(path).read_text(encoding="utf-8"))
    return p.out


def html_untouched(path: Path) -> bool:
    p = Path(path)
    if not p.exists():
        return True
    text = p.read_text(encoding="utf-8")
    m = re.search(r'<meta name="audit-engine" content="([0-9a-f]{64})">', text)
    return bool(m) and m.group(1) == blocks_hash(read_html(p))


# ---------------------------------------------------------------- 표 찾기
FS_TITLES = ("재무상태표", "손익계산서", "자본변동표", "현금흐름표")


def doc_tables(blocks: list) -> dict:
    """블록에서 표를 찾는다. 재무제표 표는 바로 앞 제목(재무상태표 등)으로, 주석 표는 '표 6-1 …' 설명으로 이름을 붙인다.
    {이름: 행 목록}. 같은 이름이 또 나오면 뒤의 것을 버린다."""
    out, last_h, last_cap = {}, None, None
    for b in blocks:
        if b[0] == "h":
            last_h = next((t for t in FS_TITLES if t in norm(b[2])), None)
            last_cap = None
        elif b[0] == "cap":
            m = re.match(r"^표\s*(\d+-\d+)", b[1])
            last_cap = m.group(1) if m else None
        elif b[0] == "table":
            name = last_cap or last_h
            if name and name not in out:
                out[name] = b[1]
            last_cap = None
            if not name:
                continue
            if name == last_h:
                last_h = None
    return out


def table_values(rows: list[list[str]]) -> dict:
    """표 행을 {첫 칸 글자: [숫자 칸...]} 로. 머리줄(첫 행)은 뺀다. 숫자가 아닌 칸은 None."""
    out = {}
    for r in rows[1:]:
        if not r:
            continue
        lab = norm(r[0])
        if not lab or lab in out:
            continue
        out[lab] = [parse_num(c) for c in r[1:]]
    return out

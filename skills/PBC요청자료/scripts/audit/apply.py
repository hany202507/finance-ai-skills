"""수정 반영과 판 관리 (2026-10-12).

회계사가 8100A 에 수정분개를 적거나(처리 「수정」), 8100B 「수정 제안」 반영 여부를 바꾸거나, 회사가 새 판을 내면
엔진이 앞서 만든 문서는 옛 숫자로 남는다. 이 모듈은 그 차이를 찾아 문서를 다시 만들고, 앞뒤 상태를 감사 폴더 안
`_판/<번호>_<시각>/` 에 남겨 되돌릴 수 있게 한다. 정의는 definitions/apply.yaml.

  pending(d)                 반영 안 된 변경 {triggers, report, areas, all_wp, signature, base, notes}. 쓰지 않는다
  apply(d, note, now, dry)   변경이 있으면 판을 남기고 다시 만든 뒤 바뀐 숫자를 _색인/수정반영.json 에 적는다
  status(d)                  판 목록 · 지금 판 · 반영 안 된 변경 · 판 뒤에 바뀐 파일
  undo(d, to, force, now)    그 판으로 되돌린다. 번호가 없으면 지금 판의 바로 앞 판
  numbers(d)                 엔진 문서의 숫자 {문서: {항목: 값}} (바뀐 숫자 목록의 원천)

엔진은 수정분개를 짓지 않는다. 무엇을 수정할지는 8100A 의 회계사 기입이 정하고, 여기서는 그 기입을 문서에 옮긴다.

무엇이 반영돼 있는지는 두 곳에서 읽는다. 보고 문서는 report.make 가 쓰는 _색인/보고.json(adjustments · company_version),
계정별 조서는 _색인/판.json 의 지금 판 signature 다. 판이 하나도 없으면 계정별 조서는 수정을 담지 않은 것으로 본다.

다시 만든 파일의 내용이 앞과 같으면 앞 파일을 그대로 둔다(openpyxl 저장은 만든 시각 때문에 해시가 늘 달라진다).
내용을 견줄 때 만든 시각 · 「엔진 초안 <날짜>」 · 표지 작성 일자는 뺀다. 다른 날 다시 만들어도 같은 내용으로 본다.
그래야 숫자가 바뀌지 않은 조서의 검사 통과 · 검토 끝과 서명이 남고, 바뀐 문서 목록이 실제로 바뀐 것만 보인다.

동시에 돌리지 않는다. 한 폴더에서 엔진 명령 둘이 함께 돌면 _색인 을 서로 덮어쓴다(runner 와 같다).
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from . import folder, wp
from .errors import DefinitionError

SUGGEST_SHEET, SUGGEST_HEAD, YES = "수정 제안", "반영 여부: ", "예"
DRAFT_RE = re.compile(r"엔진 초안 \d{4}-\d{2}-\d{2}")
KINDS = ("처음", "반영", "되돌리기 전")


class ApplyError(Exception):
    """반영 · 되돌리기 조건이 맞지 않는다. problems 에 파일 목록이나 이유가 있다."""

    def __init__(self, message: str, problems: list[str] | None = None):
        super().__init__(message)
        self.problems = problems or []


def defs() -> dict:
    p = folder.DEFS / "apply.yaml"
    if not p.exists():
        raise DefinitionError(f"정의 파일이 없습니다: {p}")
    return folder.read_yaml(p)


def _idx(d: Path, name: str, default=None):
    return folder._read(Path(d) / "_색인" / name, default)


def _put(d: Path, name: str, data) -> None:
    folder._write(Path(d) / "_색인" / name, data)


def _load(d: Path) -> dict:
    return _idx(d, defs()["index_file"], None) or {"format_version": 1, "current": None, "versions": []}


def _ver(idx: dict, no) -> dict | None:
    return next((v for v in idx["versions"] if v["no"] == no), None)


def _stamp(at: str) -> str:
    s = re.sub(r"\D", "", at)[:12].ljust(12, "0")
    return f"{s[:8]}-{s[8:]}"


def _log(d: Path, at: str, what: str, read=None, wrote=None) -> None:
    case = _idx(d, "감사건.json")
    e = {"at": at, "who": "수정반영", "what": what}
    if read:
        e["read"] = list(read)
    if wrote:
        e["wrote"] = list(wrote)
    case.setdefault("log", []).insert(0, e)
    _put(d, "감사건.json", case)


# ---------------------------------------------------------------- 지금 기입과 반영된 것
def _entry(e: dict) -> list:
    return [e["issue"], str(e["code"]), e["dr"] or 0, e["cr"] or 0]


def _sorted(rows: list) -> list:
    return sorted(rows, key=lambda r: json.dumps(r, ensure_ascii=False))


def signature(d: Path) -> dict:
    """지금 회계사 기입과 회사제시 판. adjust 는 8100A 처리 「수정」 분개(misstatement.plan 의 adjust),
    wtb 는 8100B 「수정 제안」 에서 반영 여부가 「예」 인 줄."""
    from . import figures
    from .report import misstatement
    from .report.base import prev_inputs, wp_file

    d = Path(d)
    adjust = []
    if wp_file(d, "8100A") is not None:
        tb = figures.current_tb(d)
        if tb:
            pl = misstatement.plan(d, tb, misstatement.read(d))
            adjust = _sorted([_entry(e) for e in pl["adjust"]])
    wtb = sorted(lab[len(SUGGEST_HEAD):] for (sh, lab), v in prev_inputs(wp_file(d, "8100B")).items()
                 if sh == SUGGEST_SHEET and str(lab).startswith(SUGGEST_HEAD) and str(v or "").strip() == YES)
    return {"adjust": adjust, "wtb": wtb, "company_version": (_idx(d, "감사건.json", {}) or {}).get("company_version")}


def _made(d: Path) -> list[str]:
    return [w["no"] for w in (_idx(d, "조서목록.json", {"workpapers": []}) or {"workpapers": []})["workpapers"] if w.get("file")]


def targets(d: Path, all_wp: bool) -> list[str]:
    """다시 만들 실행 단계 조서. 이미 있는 4000 영역 조서 전부, 회사제시 판이 바뀌었으면 3650A · 3320 · 3600-JET · 3800 도.
    순서는 workpapers.plan 과 같다(3320 이 4000 보다 먼저)."""
    have = _made(d)
    fixed = [str(n) for n in defs()["version_only"]]
    areas = [n for n in have if folder._is_area(n)]
    if not all_wp:
        return areas
    head = [n for n in fixed if n in have and n in ("3650A", "3320")]
    tail = [n for n in fixed if n in have and n not in head]
    return head + areas + tail


def _won(v) -> str:
    return f"{v:,.0f}" if isinstance(v, (int, float)) else str(v)


def _entry_text(e: list) -> str:
    side = f"차변 {_won(e[2])}" if e[2] else f"대변 {_won(e[3])}"
    return f"이슈 {e[0]} 계정 {e[1]} {side}"


def _delta(old: list, new: list, text) -> str:
    added = [x for x in new if x not in old]
    removed = [x for x in old if x not in new]
    parts = ([f"더해짐 {len(added)}줄({', '.join(text(x) for x in added[:3])}{' 외' if len(added) > 3 else ''})"] if added else []) + \
            ([f"빠짐 {len(removed)}줄({', '.join(text(x) for x in removed[:3])}{' 외' if len(removed) > 3 else ''})"] if removed else [])
    return " · ".join(parts)


def pending(d: Path) -> dict:
    """반영 안 된 변경. triggers 가 비면 반영할 것이 없다.
    report: 보고 문서를 다시 만들 것. areas: 4000 영역 조서를 다시 만들 것. all_wp: 회사제시 판이 바뀌어 실행 단계 조서 전부."""
    d = Path(d)
    idx = _load(d)
    rep = _idx(d, "보고.json")
    areas_made = any(folder._is_area(n) for n in _made(d))
    form = _wtbform_record(d)
    out = {"triggers": [], "report": False, "areas": False, "all_wp": False, "wtbform": False, "notes": [], "signature": None, "base": None}
    if rep is None and not areas_made and not idx["versions"] and form is None:
        return out                                 # 보고 문서도 계정별 조서도 WTB 양식도 없다. 반영할 곳이 없다
    sig = signature(d)
    out["signature"] = sig
    cv = sig["company_version"]
    cur = _ver(idx, idx.get("current"))
    rep_adj = _sorted([[a["issue"], str(a["code"]), a["dr"] or 0, a["cr"] or 0] for a in rep.get("adjustments", [])]) if rep else None
    rep_cv = rep.get("company_version") if rep else None
    base = dict(cur["signature"]) if cur else {"adjust": [], "wtb": [], "company_version": rep_cv if rep else cv}
    out["base"] = base

    if rep is not None and rep_adj != sig["adjust"]:
        out["report"] = True
        out["triggers"].append("8100A 수정분개가 보고 문서에 반영되지 않았습니다: " + _delta(rep_adj, sig["adjust"], _entry_text))
    ver_from = next((v for v in ((rep_cv if rep else None), base.get("company_version")) if v and cv and v != cv), None)
    if ver_from:
        out["report"] = rep is not None
        out["all_wp"] = True
        out["triggers"].append(f"회사제시 판이 {ver_from} 에서 {cv} 가 됐습니다")
        out["notes"].append("회사제시 판이 바뀌었습니다. 분석(python -m audit analyze)은 이 명령이 다시 돌리지 않습니다. "
                            "2120A · 8100B 와 이슈는 analyze 를 따로 돌려 새 판 숫자로 맞춰 주십시오")
    if areas_made:
        if base["adjust"] != sig["adjust"]:
            out["areas"] = True
            if not any(t.startswith("8100A") for t in out["triggers"]):
                out["triggers"].append("계정별 조서(4000)의 수정 열이 지금 8100A 수정분개와 다릅니다: "
                                       + _delta(base["adjust"], sig["adjust"], _entry_text))
        if base.get("wtb", []) != sig["wtb"]:
            out["areas"] = True
            out["triggers"].append("8100B 수정 제안의 반영 여부가 달라졌습니다: " + _delta(base.get("wtb", []), sig["wtb"], str))
        if out["all_wp"]:
            out["areas"] = True
    if form is not None:
        if form.get("adjust", []) != sig["adjust"]:
            out["wtbform"] = True
            out["triggers"].append("WTB 양식의 수정사항이 지금 8100A 수정분개와 다릅니다: " + _delta(form.get("adjust", []), sig["adjust"], _entry_text))
        elif form.get("wtb", []) != sig["wtb"]:
            out["wtbform"] = True
            out["triggers"].append("WTB 양식의 엔진 후보 반영 여부가 지금 8100B 와 다릅니다: " + _delta(form.get("wtb", []), sig["wtb"], str))
        if form.get("manual"):
            out["notes"].append(f"WTB 양식 수정사항 시트에 직접 적은 줄이 {form['manual']}개 있습니다. 재무제표 · DART 에 넣으려면 8100A 에 옮겨 주십시오")
    return out


def _wtbform_record(d: Path) -> dict | None:
    """감사 폴더 안에 만든 WTB 양식(wtbform)의 기록. 파일이 없어졌으면 없는 것으로 본다."""
    from .wtbform import RECORD

    rec = _idx(d, RECORD)
    if not rec or not (Path(d) / rec.get("path", "")).is_file():
        return None
    return rec


def _rebuild_wtbform(d: Path) -> str:
    """WTB 양식을 같은 자리에 다시 만든다. 분석 결과(_색인/분석.json)는 다시 돌리지 않는다."""
    from . import wtbform

    rec = _wtbform_record(d)
    res = _idx(d, "분석.json")
    if rec is None or res is None:
        raise ApplyError("WTB 양식을 다시 만들 분석 결과가 없습니다. python -m audit analyze 를 먼저 돌려 주십시오")
    wtbform.build(d, res, Path(d) / rec["path"])
    return rec["path"]


# ---------------------------------------------------------------- 판에 담는 파일
def _files(d: Path) -> list[Path]:
    std = folder.load_standard()
    out = []
    for top in defs()["folders"]:
        base = Path(d) / top
        if base.is_dir():
            out += [p for p in sorted(base.rglob("*")) if p.is_file() and not folder._ignored(p.name, std)]
    return out


def _rel(d: Path, p: Path) -> str:
    return Path(p).relative_to(d).as_posix()


def _hashes(d: Path) -> dict[str, str]:
    return {_rel(d, p): folder._sha(p) for p in _files(d)}


def _locked(paths: list[Path]) -> list[str]:
    """다른 프로그램(엑셀 · 워드)이 열어 둔 파일. 쓰기 전에 걸러 중간에 멈추지 않게 한다."""
    out = []
    for p in paths:
        try:
            with open(p, "r+b"):
                pass
        except PermissionError:
            out.append(p.name)
        except OSError:
            continue
    return out


def _snapshot(d: Path, idx: dict, kind: str, sig: dict, at: str, note: str = "", parent: int | None = None) -> dict:
    """지금 상태를 판 하나로 남긴다. 조서 폴더 셋과 _색인 의 json 사본, 해시목록."""
    cfg = defs()
    no = max((v["no"] for v in idx["versions"]), default=-1) + 1
    target = Path(d) / cfg["dir"] / f"{no:03d}_{_stamp(at)}"
    n = 2
    while target.exists():
        target = Path(d) / cfg["dir"] / f"{no:03d}_{_stamp(at)}_{n}"
        n += 1
    files = []
    for src in _files(d):
        rel = _rel(d, src)
        dst = target / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        h = folder._sha(dst)
        if h != folder._sha(src):
            raise ApplyError(f"판에 복사한 파일의 해시가 원본과 다릅니다: {rel}")
        files.append({"path": rel, "sha256": h, "size": dst.stat().st_size})
    names = []
    (target / "_색인").mkdir(parents=True, exist_ok=True)
    for p in sorted((Path(d) / "_색인").glob("*.json")):
        if p.name != cfg["index_file"]:
            shutil.copy2(p, target / "_색인" / p.name)
            names.append(p.name)
    body = json.dumps({"format_version": 1, "no": no, "kind": kind, "created_at": at, "count": len(files), "files": files, "index": names},
                      ensure_ascii=False, indent=1).encode("utf-8")
    (target / cfg["hash_list"]).write_bytes(body)
    rec = {"no": no, "at": at, "kind": kind, "path": _rel(d, target), "parent": parent, "note": note or "", "signature": sig,
           "files": len(files), "hash_list_sha256": hashlib.sha256(body).hexdigest()}
    idx["versions"].append(rec)
    return rec


def _listed(d: Path, rec: dict) -> dict[str, dict]:
    """판의 해시목록 {경로: {sha256, size}}. 해시목록이 없거나 판.json 기록과 다르면 멈춘다."""
    lp = Path(d) / rec["path"] / defs()["hash_list"]
    if not lp.exists():
        raise ApplyError(f"판 {rec['no']} 의 해시목록이 없습니다: {rec['path']}")
    body = lp.read_bytes()
    if hashlib.sha256(body).hexdigest() != rec["hash_list_sha256"]:
        raise ApplyError(f"판 {rec['no']} 의 해시목록이 판을 남긴 때와 다릅니다: {rec['path']}")
    return {f["path"]: f for f in json.loads(body.decode("utf-8"))["files"]}


def edited_since(d: Path, rec: dict) -> dict:
    """판 rec 을 남긴 뒤 조서 폴더에서 달라진 파일. {changed, new, missing}"""
    listed = _listed(d, rec)
    now = _hashes(d)
    return {"changed": sorted(p for p, h in now.items() if p in listed and listed[p]["sha256"] != h),
            "new": sorted(p for p in now if p not in listed),
            "missing": sorted(p for p in listed if p not in now)}


# ---------------------------------------------------------------- 숫자
def _doc_name(p: Path) -> str:
    return folder.VERSION_RE.sub("", p.stem)


def numbers(d: Path) -> dict[str, dict[str, float]]:
    """엔진 문서의 숫자. 엔진 조서의 `_검산` 기대값(8800 계정별 수정 후 · Lead 확정 · 8100A 합계 등), 8400 주석 표 값,
    8810 날인재무제표와 DART 제출 파일의 표 값, 보고.json 의 주요 금액과 푸팅 · DART 검토 불일치 건수."""
    from .report import dart_file, notes
    from .report.base import doc_tables, read_docx, read_html, table_values, wp_file

    d = Path(d)
    out: dict[str, dict[str, float]] = {}

    def put(doc: str, item: str, v) -> None:
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            return
        tab = out.setdefault(doc, {})
        key, n = item, 2
        while key in tab:
            key, n = f"{item} ({n})", n + 1
        tab[key] = v

    for p in _files(d):
        if p.suffix.lower() != ".xlsx" or len(p.relative_to(d).parts) != 2:
            continue
        try:
            wb = load_workbook(p, read_only=True)
        except Exception:  # noqa: BLE001  회계사가 다른 꼴로 저장한 파일은 건너뛴다
            continue
        try:
            if wp.CHECK_SHEET in wb.sheetnames:
                for row in wb[wp.CHECK_SHEET].iter_rows(min_row=2, values_only=True):
                    if row and row[0] and len(row) > 2:
                        put(_doc_name(p), str(row[0]), row[2])
        finally:
            wb.close()
    x = wp_file(d, "8400", "8400_공시사항점검표*.xlsx")
    if x is not None:
        try:
            for t in notes.read_xlsx(x):
                heads = [str(c) for c in t["cols"]]
                for row in t["rows"] + ([t["total"]] if t.get("total") else []):
                    for h, v in zip(heads[1:], row[1:]):
                        put("8400 주석", f"표 {t['id']} {row[0]} · {h}", v)
        except Exception:  # noqa: BLE001
            pass
    docs = [("8810 날인재무제표", wp.path_for(d, "8810", ".docx"), read_docx), ("DART 제출파일", dart_file.paths(d)[0], read_html)]
    for label, path, read in docs:
        if not path.exists():
            continue
        try:
            for name, rows in doc_tables(read(path)).items():
                heads = [str(c) for c in rows[0]] if rows else []
                for lab, nums in table_values(rows).items():
                    for i, v in enumerate(nums):
                        put(label, f"{name} {lab} · {heads[i + 1] if i + 1 < len(heads) and heads[i + 1] else i + 1}", v)
        except Exception:  # noqa: BLE001
            continue
    rep = _idx(d, "보고.json")
    if rep:
        for k, v in (rep.get("statements") or {}).items():
            put("재무제표 주요 금액", k, v)
        put("푸팅", "불일치 건수", (rep.get("footing") or {}).get("불일치"))
        if rep.get("dart_review"):
            put("DART 검토", "불일치 건수", rep["dart_review"].get("불일치"))
    return out


def changed_numbers(before: dict, after: dict) -> list[dict]:
    """바뀐 숫자 목록. 열은 조서 · 항목 · 전 · 후 · 차이. 한쪽에만 있는 항목은 없는 쪽을 null 로 둔다."""
    out = []
    for doc in list(dict.fromkeys(list(before) + list(after))):
        b, a = before.get(doc, {}), after.get(doc, {})
        for item in list(dict.fromkeys(list(b) + list(a))):
            x, y = b.get(item), a.get(item)
            if x is not None and y is not None and abs(x - y) <= 0.5:
                continue
            out.append({"workpaper": doc, "item": item, "before": x, "after": y,
                        "diff": (y - x) if x is not None and y is not None else None})
    return out


# ---------------------------------------------------------------- 내용이 같은 파일 가리기
COVER_DATE = ("표지", 4, 5)          # 4000 조서 표지의 작성 일자 칸(lead._cover_head). 만든 날짜라 다시 만들면 달라진다


def _xlsx_key(p: Path) -> tuple:
    wb = load_workbook(p, read_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            cells = []
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if ws.title == wp.META_SHEET and row and row[0] == "만든 시각":
                    continue
                for j, v in enumerate(row, 1):
                    if v is None or (ws.title, i, j) == COVER_DATE:
                        continue
                    if isinstance(v, str):
                        v = DRAFT_RE.sub("엔진 초안", v)
                    cells.append((i, j, repr(v)))
            out.append((ws.title, ws.sheet_state, tuple(cells)))
        return tuple(out)
    finally:
        wb.close()


def content_key(p: Path):
    """파일 내용의 열쇠. 만든 시각처럼 다시 만들 때마다 달라지는 것은 뺀다. 읽지 못하면 None(바이트로만 견준다)."""
    from .report import dsd
    from .report.base import blocks_text, read_docx, read_html

    suf = Path(p).suffix.lower()
    try:
        if suf == ".xlsx":
            return _xlsx_key(p)
        if suf == ".docx":
            return blocks_text(read_docx(p))
        if suf == ".html":
            return blocks_text(read_html(p))
        if suf == ".dsd":
            return blocks_text(dsd.read_dsd(p))
    except Exception:  # noqa: BLE001
        return None
    return None


def _keep_same(d: Path, stash: Path, before: dict[str, str]) -> dict:
    """다시 만든 뒤 내용이 앞과 같은 파일은 앞 파일을 되돌려 놓는다. {changed, same, new}"""
    out = {"changed": [], "same": [], "new": []}
    for rel, h in _hashes(d).items():
        if rel not in before:
            out["new"].append(rel)
        elif h != before[rel]:
            old = content_key(stash / rel)
            if old is not None and old == content_key(Path(d) / rel):
                shutil.copy2(stash / rel, Path(d) / rel)
                out["same"].append(rel)
            else:
                out["changed"].append(rel)
    return out


WP_KEEP = ("status", "status_basis", "signs", "needs_recheck", "sha256")


def _restore_status(d: Path, before: dict[str, dict]) -> list[str]:
    """파일이 앞과 바이트까지 같은 조서는 앞 상태(검사 통과 · 검토 끝 · 서명)를 되살린다.
    다시 만드는 동안 엔진 명령의 scan 이 바뀐 해시를 보고 작성중으로 내렸기 때문이다."""
    data = _idx(d, "조서목록.json")
    tops = tuple(defs()["folders"])
    back = []
    for w in data["workpapers"]:
        old = before.get(w["no"])
        f = w.get("file")
        if not old or not f or old.get("file") != f or not f.startswith(tops) or not (Path(d) / f).exists():
            continue
        if old.get("sha256") and folder._sha(Path(d) / f) == old["sha256"]:
            if any(w.get(k) != old.get(k) for k in WP_KEEP):
                back.append(w["no"])
            for k in WP_KEEP:
                if k in old:
                    w[k] = old[k]
    _put(d, "조서목록.json", data)
    return back


def _wps(d: Path) -> dict[str, dict]:
    return {w["no"]: dict(w) for w in (_idx(d, "조서목록.json", {"workpapers": []}) or {"workpapers": []})["workpapers"]}


def _mismatches(d: Path) -> list[dict]:
    """반영 뒤에도 남은 문서 사이 불일치(푸팅 · DART 검토). 아직 수정되지 않은 숫자를 찾는 줄이다."""
    rep = _idx(d, "보고.json") or {}
    out = []
    for src, key in (("푸팅", "footing"), ("DART 검토", "dart_review")):
        for x in ((rep.get(key) or {}).get("items") or []):
            if x.get("result") == "불일치":
                out.append({"source": src, "area": x.get("area"), "where": x.get("where"), "item": x.get("item"),
                            "doc": x.get("doc"), "ref": x.get("ref"), "diff": x.get("diff")})
    return out


def _is_listed(d: Path) -> bool:
    """앞서 만든 8820 이 주권상장법인 꼴(핵심감사사항 기입란)인지. report dart 와 같은 방법으로 본다."""
    from .report.base import read_docx

    p = wp.path_for(d, "8820", ".docx")
    if not p.exists():
        return False
    try:
        return any("【기입: 핵심감사사항" in str(b[-1]) for b in read_docx(p))
    except Exception:  # noqa: BLE001
        return False


def _old_version_edits(d: Path, old_ver: str | None) -> list[str]:
    """앞 판 이름의 워드 · HTML 가운데 회계사가 고친 것. 새 판 문서는 새 이름으로 새로 만들어지므로 손 수정이 옮겨지지 않는다."""
    from .report.base import docx_untouched, html_untouched

    if not old_ver:
        return []
    out = []
    for p in _files(d):
        if f"_{old_ver}" not in p.stem:
            continue
        if (p.suffix.lower() == ".docx" and not docx_untouched(p)) or (p.suffix.lower() == ".html" and not html_untouched(p)):
            out.append(_rel(d, p))
    return out


# ---------------------------------------------------------------- 반영
def apply(d: Path, note: str | None = None, now: str | None = None, dry: bool = False) -> dict:
    """반영 안 된 변경이 있으면 판을 남기고 다시 만든다. dry 면 무엇을 할지만 돌려준다(파일을 바꾸지 않는다)."""
    from . import report, workpapers

    d = Path(d)
    at = now or folder._now()
    if not dry:
        folder.scan(d, now=now, quiet_log=True)
    pend = pending(d)
    nos = targets(d, pend["all_wp"]) if pend["areas"] else []
    res = {"changed": bool(pend["triggers"]), "applied": False, "triggers": pend["triggers"], "notes": list(pend["notes"]),
           "plan": {"report": pend["report"], "workpapers": nos, "wtbform": pend["wtbform"]}}
    if dry or not pend["triggers"]:
        return res
    locked = _locked(_files(d))
    if locked:
        raise ApplyError("열려 있는 파일이 있어 반영하지 않았습니다. 엑셀 · 워드에서 닫은 뒤 다시 돌려 주십시오", locked)

    idx = _load(d)
    if not idx["versions"]:
        _snapshot(d, idx, "처음", pend["base"], at, note="반영 전 상태")
        idx["current"] = 0
    parent = idx["current"]
    num_before = numbers(d)
    wps_before = _wps(d)
    hash_before = _hashes(d)
    old_edits = _old_version_edits(d, pend["base"].get("company_version")) if pend["all_wp"] else []
    not_rebuilt, errors = [], []
    with tempfile.TemporaryDirectory(prefix="audit-apply-") as tmp:
        stash = Path(tmp)
        for rel in hash_before:
            (stash / rel).parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(d / rel, stash / rel)
        try:
            if pend["report"]:
                r = report.make(d, now=now, listed=_is_listed(d))
                not_rebuilt += [{"path": p, "why": "회계사가 고친 문서라 다시 만들지 않았습니다. 숫자가 옛 값일 수 있습니다"} for p in r["kept"]]
                errors += [f"8100A: {p}" for p in r["problems"]]
            if pend["wtbform"]:                 # WTB 가 먼저다. 계정별 조서 Lead 는 WTB 판과 같은지 본다(2026-10-16)
                _rebuild_wtbform(d)
            if nos:
                r = workpapers.make(d, nos=nos, now=now)
                not_rebuilt += [{"path": k["no"], "why": k["message"]} for k in r["kept"]]
                errors += [f"{e['no']}: {e['message']}" for e in r["errors"]]
        except Exception as e:  # noqa: BLE001  다시 만들다 멈추면 반영 전으로 돌려놓는다
            for rel in hash_before:
                if not (d / rel).exists() or folder._sha(d / rel) != hash_before[rel]:
                    shutil.copy2(stash / rel, d / rel)
            folder.scan(d, now=now, quiet_log=True)
            _restore_status(d, wps_before)
            _put(d, defs()["index_file"], idx)
            folder.validate_index(d)
            raise ApplyError(f"다시 만들다 멈췄습니다. 파일은 반영 전으로 되돌렸습니다({type(e).__name__}: {e})") from e
        files = _keep_same(d, stash, hash_before)
    folder.scan(d, now=now, quiet_log=True)
    kept_status = _restore_status(d, wps_before)
    num_after = numbers(d)
    rec = _snapshot(d, idx, "반영", pend["signature"], at, note=note or "", parent=parent)
    idx["current"] = rec["no"]
    _put(d, defs()["index_file"], idx)

    notes = list(pend["notes"])
    if old_edits:
        notes.append("앞 판 이름의 문서에 손으로 고친 내용은 새 판 문서로 옮기지 않았습니다: " + ", ".join(old_edits))
    changes = changed_numbers(num_before, num_after)
    result = {"format_version": 1, "at": at, "kind": "반영", "version": rec["no"], "from_version": parent, "note": note or "",
              "triggers": pend["triggers"], "changed_numbers": changes, "changed_files": files["changed"],
              "unchanged_files": files["same"], "new_files": files["new"], "not_rebuilt": not_rebuilt,
              "mismatches": _mismatches(d), "errors": errors, "notes": notes}
    _put(d, defs()["result_file"], result)
    what = (f"수정을 반영해 판 {rec['no']} 을 남겼습니다. 바뀐 문서 {len(files['changed'])}개 · 바뀐 숫자 {len(changes)}개"
            + (f" · 다시 만들지 않은 문서 {len(not_rebuilt)}개" if not_rebuilt else "")
            + (f" · 남은 불일치 {len(result['mismatches'])}건" if result["mismatches"] else "")
            + (f". 사유: {note}" if note else ""))
    _log(d, at, what, read=["_색인/보고.json", "_색인/조서목록.json"],
         wrote=files["changed"][:50] + [f"_색인/{defs()['result_file']}", f"_색인/{defs()['index_file']}", rec["path"]])
    folder.validate_index(d)
    return dict(res, applied=True, version=rec["no"], from_version=parent, result=result, kept_status=kept_status)


# ---------------------------------------------------------------- 현황
def status(d: Path) -> dict:
    d = Path(d)
    idx = _load(d)
    cur = _ver(idx, idx.get("current"))
    pend = pending(d)
    return {"versions": idx["versions"], "current": idx.get("current"), "triggers": pend["triggers"], "notes": pend["notes"],
            "edited": edited_since(d, cur) if cur else {"changed": [], "new": [], "missing": []}}


# ---------------------------------------------------------------- 되돌리기
def undo_preview(d: Path, to: int | None = None) -> dict:
    """되돌리기 전에 보일 것(2026-10-16, 화면). 파일을 바꾸지 않는다.
    {to, from, restore: 판 to 내용으로 바뀔 파일, extra: 판 to 뒤에 생겨 그대로 둘 파일, edited: 지금 판 뒤에 바뀐 파일,
     need_force, problems, signature: 판 to 의 수정분개 줄 수 · 회사제시 판}"""
    d = Path(d)
    idx = _load(d)
    cur = _ver(idx, idx.get("current"))
    if cur is None:
        raise ApplyError("판이 없습니다. 먼저 수정을 반영해 주십시오")
    if to is None:
        to = cur.get("parent")
        if to is None:
            raise ApplyError(f"판 {cur['no']} 앞에는 판이 없습니다. 되돌릴 판 번호를 골라 주십시오")
    target = _ver(idx, to)
    if target is None:
        raise ApplyError(f"없는 판 번호입니다: {to}")
    listed = _listed(d, target)
    tdir = d / target["path"]
    bad = [p for p, f in listed.items() if not (tdir / p).exists() or folder._sha(tdir / p) != f["sha256"]]
    ed = edited_since(d, cur)
    now = _hashes(d)
    restore = sorted(p for p, f in listed.items() if now.get(p) != f["sha256"])
    sig = target.get("signature") or {}
    return {"to": to, "from": cur["no"], "same": to == cur["no"] and not (ed["changed"] or ed["missing"]),
            "restore": restore, "extra": sorted(p for p in now if p not in listed), "edited": ed,
            "need_force": bool(ed["changed"] or ed["missing"]), "problems": bad,
            "locked": _locked([d / p for p in listed if (d / p).exists()]),
            "signature": {"adjust": len(sig.get("adjust") or []), "wtb": len(sig.get("wtb") or []), "company_version": sig.get("company_version")}}


def undo(d: Path, to: int | None = None, force: bool = False, now: str | None = None) -> dict:
    """판 to 로 되돌린다. to 가 없으면 지금 판의 바로 앞 판. 되돌리기 직전 상태가 지금 판과 다르면 그것도 판으로 남긴다.
    지금 판을 남긴 뒤 바뀐 파일이 있으면 force 없이는 되돌리지 않는다. 판 뒤에 생긴 파일은 지우지 않고 알린다."""
    d = Path(d)
    at = now or folder._now()
    cfg = defs()
    idx = _load(d)
    cur = _ver(idx, idx.get("current"))
    if cur is None:
        raise ApplyError("판이 없습니다. 먼저 python -m audit apply 로 수정을 반영해 주십시오")
    if to is None:
        to = cur.get("parent")
        if to is None:
            raise ApplyError(f"판 {cur['no']} 앞에는 판이 없습니다. 되돌릴 판 번호를 주십시오(--status 로 목록을 봅니다)")
    target = _ver(idx, to)
    if target is None:
        raise ApplyError(f"없는 판 번호입니다: {to}")
    listed = _listed(d, target)
    tdir = d / target["path"]
    bad = [p for p, f in listed.items() if not (tdir / p).exists() or folder._sha(tdir / p) != f["sha256"]]
    if bad:
        raise ApplyError(f"판 {to} 의 파일이 해시목록과 다릅니다. 되돌리지 않았습니다", bad)
    ed = edited_since(d, cur)
    if to == cur["no"] and not (ed["changed"] or ed["missing"]):
        raise ApplyError(f"이미 판 {to} 상태입니다")
    if (ed["changed"] or ed["missing"]) and not force:
        raise ApplyError(f"판 {cur['no']} 을 남긴 뒤 바뀐 파일이 {len(ed['changed']) + len(ed['missing'])}개 있습니다. "
                         f"되돌리면 이 파일들이 판 {to} 때의 내용이 됩니다. 지금 상태는 판으로 남기니 --force 를 주면 되돌립니다",
                         ed["changed"] + [f"{p} (없어짐)" for p in ed["missing"]])
    locked = _locked([d / p for p in listed if (d / p).exists()])
    if locked:
        raise ApplyError("열려 있는 파일이 있어 되돌리지 않았습니다. 엑셀 · 워드에서 닫은 뒤 다시 돌려 주십시오", locked)

    num_before = numbers(d)
    pre = None
    if ed["changed"] or ed["new"] or ed["missing"]:
        pre = _snapshot(d, idx, "되돌리기 전", cur["signature"], at, note=f"판 {to} 로 되돌리기 직전", parent=cur["no"])
    restored = []
    for rel, f in listed.items():
        dst = d / rel
        if not dst.exists() or folder._sha(dst) != f["sha256"]:
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(tdir / rel, dst)
            restored.append(rel)
    extra = sorted(p for p in _hashes(d) if p not in listed)
    for name in cfg["index_restore"]:
        src = tdir / "_색인" / name
        if src.exists():
            shutil.copy2(src, d / "_색인" / name)
    folder.scan(d, now=now, quiet_log=True)
    snap_wps = {w["no"]: w for w in (folder._read(tdir / "_색인" / "조서목록.json", {"workpapers": []}) or {"workpapers": []})["workpapers"]}
    _restore_status(d, snap_wps)
    idx["current"] = to
    _put(d, cfg["index_file"], idx)
    changes = changed_numbers(num_before, numbers(d))
    pend = pending(d)
    result = {"format_version": 1, "at": at, "kind": "되돌림", "version": to, "from_version": cur["no"], "note": "",
              "triggers": pend["triggers"], "changed_numbers": changes, "changed_files": restored, "unchanged_files": [],
              "new_files": extra, "not_rebuilt": [], "mismatches": _mismatches(d), "errors": [], "notes":
              ([f"되돌리기 직전 상태를 판 {pre['no']} 으로 남겼습니다"] if pre else [])
              + ([f"판 {to} 뒤에 생긴 파일 {len(extra)}개는 지우지 않았습니다"] if extra else [])}
    _put(d, cfg["result_file"], result)
    _log(d, at, f"판 {cur['no']} 에서 판 {to} 로 되돌렸습니다. 되살린 파일 {len(restored)}개"
               + (f", 되돌리기 직전 상태는 판 {pre['no']}" if pre else "")
               + (f", 판 {to} 뒤에 생긴 파일 {len(extra)}개는 그대로 둠" if extra else ""),
         read=[f"{target['path']}/{cfg['hash_list']}"], wrote=restored[:50] + [f"_색인/{cfg['index_file']}", f"_색인/{cfg['result_file']}"])
    folder.validate_index(d)
    return {"to": to, "from": cur["no"], "restored": restored, "extra": extra, "pre": pre["no"] if pre else None,
            "pending": pend["triggers"], "result": result}

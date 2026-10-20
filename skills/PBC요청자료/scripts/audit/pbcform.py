"""요청자료(PBC) 양식 (2026-10-13). 실제 초도감사 요청자료 목록(샘플문서)의 구조를 따른다.

  1. 초도감사및중간감사   머리(회사 · 제목 · 일자 · 안내 두 줄) 아래 표: 순번 · 구분 · 요청자료명 · 요청세부내역 · 요청형태 · 요청일자 ·
                         제공일자 · File명 · 보완 필요 여부(Y/N) · 보완요청 내역 · 비고. 오른쪽에 엔진 열(엔진 번호 · 판정 · 찾은 파일)을 더한다
  2. 질의서              부정에 대한 경영진 질문서 답변란
  3. 전산환경            IT 일반통제 · 응용통제 질의서(Y · N · N/A · 서술내용 · 샘플 문서명)
  4. 보완                데이터 백업 · 패스워드 관리 현황표

요청 항목과 문구는 definitions/pbc_form.yaml 이다. 받음 · 빠짐 판정은 엔진이 하는 18건(pbc_match.yaml)에만 한다. 양식 줄의 「엔진 번호」 가
그 18건의 번호이고, 번호가 없는 줄은 엔진이 받았는지 보지 않는다(회계사가 제공일자로 관리한다). 판단은 엔진이 하지 않는다.

원 양식은 다음 해 6월에 하는 초도감사 기준이라 다음 해 반기 · 6월 · 2분기 · 6개월을 묻는 줄이 30여 줄 있다(2026-10-18 실제 회사 자료).
감사건에 반기 기준일(half_year_end)이 없으면 그 줄은 빼거나(half_only) 결산 연도 기준 문구(fy_detail)로 쓴다.
시산표 · 재무제표에 영역 잔액이 당기 · 전기 모두 없으면 그 영역 줄(area)은 빼거나(pbc_match.yaml no_balance.drop) 「해당 없음 후보」 로 적는다.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import yaml
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from . import folder, wp

DEF = folder.DEFS / "pbc_form.yaml"
HEAD_ROW = 7
COLS = ["순번", "구분", "요청자료명", "요청세부내역", "요청형태", "요청일자", "제공일자", "File명", "보완 필요 여부", "보완요청 내역", "비고"]
ENGINE_COLS = ["엔진 번호", "엔진 판정", "엔진이 찾은 파일"]
WIDTHS = [2, 7, 14, 20, 62, 14, 12, 12, 36, 12, 34, 24, 9, 14, 40]


def load() -> dict:
    return yaml.safe_load(DEF.read_text(encoding="utf-8"))


def years(fy_end: str, half_year_end: str | None = None, fy_start: str | None = None) -> dict:
    """양식 글자 자리 값. 반기 기준일이 있으면 {ny} 는 그 기준일의 연도다(원 양식은 결산 다음 해 6월 초도감사)."""
    fy = int(str(fy_end)[:4]) if fy_end else date.today().year
    ny = int(str(half_year_end)[:4]) if half_year_end else fy + 1
    try:
        e = date.fromisoformat(str(fy_end)[:10])
    except ValueError:
        e = date(fy, 12, 31)
    try:
        st = date.fromisoformat(str(fy_start)[:10]) if fy_start else None
    except ValueError:
        st = None
    if st is None:
        st = date(e.year, 1, 1) if (e.month, e.day) == (12, 31) else date(e.year - 1, e.month % 12 + 1, 1)
    after = (e.year + 1, 1) if e.month == 12 else (e.year, e.month + 1)
    return {"fy": fy, "py": fy - 1, "ny": ny, "fy2": f"{fy % 100:02d}", "py2": f"{(fy - 1) % 100:02d}", "ny2": f"{ny % 100:02d}",
            "fs_k": f"{st.year}년 {st.month}월 {st.day}일", "fem": e.month, "after1": f"{after[0]}년 {after[1]}월"}


def rows_for(cfg: dict, first_year: bool, y: dict, half: bool = False, drop_areas: set | frozenset = frozenset()) -> list[dict]:
    """이 감사 건에 보낼 요청 줄. 계속감사는 초도감사에서만 점검하는 줄을 빼고, 남은 초도감사(1차) 구분은 감사계획으로 바꾼다.
    half(반기 기준일 있음)가 아니면 half_only 줄을 빼고 fy_detail 이 있는 줄은 그 문구를 쓴다. drop_areas 영역 줄은 뺀다."""
    out = []
    for r in cfg["rows"]:
        if r.get("first_year_only") and not first_year:
            continue
        if r.get("half_only") and not half:
            continue
        if r.get("area") and str(r["area"]) in drop_areas:
            continue
        cat = r["cat"] if first_year or r["cat"] != "초도감사(1차)" else "감사계획"
        detail = r.get("detail") if half or not r.get("fy_detail") else r["fy_detail"]
        out.append({**r, "cat": cat, "name": r["name"].format(**y), "detail": (detail or "").format(**y)})
    return out


def engine_status(rows: list[dict], std: list[dict]) -> list[dict]:
    """줄마다 엔진 번호의 받음 · 빠짐과 찾은 파일. std 는 start.pbc_status 결과."""
    by = {x["no"]: x for x in std}
    return [by.get(r.get("engine")) for r in rows]


def write(b: wp.Book, d: Path, std: list[dict]) -> int:
    from . import director

    cfg = load()
    case = wp._case(d)
    first = case.get("engagement_kind") == "초도감사"
    half = case.get("half_year_end")
    y = years(b.fy_end, half, case.get("fiscal_year_start"))
    c = director._Case(d)
    nobal = c.no_balance()
    rows = rows_for(cfg, first, y, half=bool(half), drop_areas=c.dropped())
    stat = engine_status(rows, std)
    key = "first_year" if first else "continuing"

    ws, _ = b.sheet(cfg["sheet_names"][key], WIDTHS, head=False)
    ws["B1"] = b.company or ""
    ws["B1"].font = wp.font(size=14, bold=True, color=wp.NAVY)
    ws["B2"] = cfg["titles"][key]
    ws["B2"].font = wp.font(size=12, bold=True)
    ws["B3"] = date.today()
    ws["B3"].number_format = "yyyy-mm-dd"
    for i, t in enumerate(cfg["notes"]):
        ws.cell(5 + i, 2, t).font = wp.font(size=10)
    b.head(ws, HEAD_ROW, COLS, col=2)
    b.head(ws, HEAD_ROW, ENGINE_COLS, col=2 + len(COLS))
    row = HEAD_ROW + 1
    first_row = row
    for n, r in enumerate(rows, 1):
        s = stat[n - 1]
        b.put(ws, row, 2, n)
        b.put(ws, row, 3, r["cat"], wrap=True)
        b.put(ws, row, 4, r["name"], wrap=True)
        b.put(ws, row, 5, r["detail"] or None, wrap=True)
        b.put(ws, row, 6, r.get("form"))
        b.ask(ws, row, 7, f"{n} 요청일자", rule="optional", value=date.today(), fmt="yyyy-mm-dd")
        b.ask(ws, row, 8, f"{n} 제공일자", rule="optional", fmt="yyyy-mm-dd")
        b.ask(ws, row, 9, f"{n} File명", rule="optional")
        b.ask(ws, row, 10, f"{n} 보완 필요 여부", rule="optional")
        b.ask(ws, row, 11, f"{n} 보완요청 내역", rule="optional")
        b.ask(ws, row, 12, f"{n} 비고", rule="optional")
        b.put(ws, row, 13, r.get("engine"))
        status = s["status"] if s else (director.NO_BALANCE if r.get("area") and str(r["area"]) in nobal else "엔진 미판정")
        b.put(ws, row, 14, status, flag=status == "빠짐", bold=status == "빠짐", color=None if s else wp.MUTED)
        files = (s or {}).get("files") or []
        b.put(ws, row, 15, "\n".join(files[:3]) + (f"\n외 {len(files) - 3}개" if len(files) > 3 else "") if files else None, wrap=True, color=wp.MUTED)
        row += 1
    last = row - 1
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J{first_row}:J{last}")
    ws.freeze_panes = ws.cell(first_row, 5)
    ws.auto_filter.ref = f"B{HEAD_ROW}:{wp.get_column_letter(2 + len(COLS) - 1)}{last}"
    got = sum(1 for s in stat if s and s["status"] == "받음")
    judged = sum(1 for s in stat if s)
    ws.cell(row + 1, 4, "엔진이 판정하는 자료").font = wp.font(bold=True)
    ws.cell(row + 1, 5, f'=COUNTIF(N{first_row}:N{last},"받음")&" 받음 / "&COUNTIF(M{first_row}:M{last},">0")&" 줄(엔진 번호가 있는 줄)"')
    b.verify("PBC 받음 줄 수", f'=COUNTIF(\'{ws.title}\'!N{first_row}:N{last},"받음")', got)
    b.verify("PBC 엔진 판정 줄 수", f"=COUNTIF('{ws.title}'!M{first_row}:M{last},\">0\")", judged)

    _questionnaire_sheet(b)
    _it_sheet(b, cfg)
    _supplement_sheet(b, cfg)
    return len(rows)


FRAUD_QUESTIONS = [
    "경영진은 부정으로 인한 재무제표 중요왜곡표시위험을 평가하는 절차와 그 빈도 · 결과를 어떻게 정하고 있습니까.",
    "부정위험(경영진 부정 포함)을 식별하고 대응하는 회사의 절차와 통제를 설명해 주십시오.",
    "경영진이 알고 있거나 의심하는 부정 사례가 있으면 내용을 설명해 주십시오(회사에 영향을 주는 직원 · 경영진 · 외부자 부정 포함).",
    "부정 또는 부정 의심에 대해 종업원 · 투자자 · 규제기관 등에서 제기된 주장이 있습니까.",
    "부정위험 평가와 대응 절차를 이사회 · 감사(위원회)에 어떻게 전달합니까. 최근 전달한 사례가 있으면 설명해 주십시오.",
    "경영진이 종업원에게 윤리적 행동과 정직한 업무 수행 기준을 전달하는 방법(윤리규정 · 서약서 · 교육)을 설명해 주십시오.",
    "재무제표에 영향을 주는 법령 · 규정 위반 사실이나 의심 사례가 있습니까.",
    "특수관계자의 신원과 특수관계자와의 중요한 거래(정상 영업과정 밖의 거래 포함)를 설명해 주십시오.",
]


def _questionnaire_sheet(b: wp.Book) -> None:
    ws, _ = b.sheet("2. 질의서", [2, 7, 80, 60, 24], head=False)
    ws["B1"] = "부정에 대한 경영진 질문서"
    ws["B1"].font = wp.font(size=14, bold=True, color=wp.NAVY)
    ws["B2"] = "감사기준서 240 의 경영진 질문 사항입니다. 답변란에 서술해 주시고 근거 문서가 있으면 비고에 문서명을 적어 주십시오(별도 질의서 PDF 로 대신해도 됩니다)."
    ws["B2"].font = wp.font(size=10, color=wp.MUTED)
    b.head(ws, 4, ["번호", "질문", "답변", "비고"], col=2)
    row = 5
    for n, q in enumerate(FRAUD_QUESTIONS, 1):
        b.put(ws, row, 2, n)
        b.put(ws, row, 3, q, wrap=True)
        b.ask(ws, row, 4, f"질의서 {n} 답변", rule="optional")
        b.ask(ws, row, 5, f"질의서 {n} 비고", rule="optional")
        ws.row_dimensions[row].height = 44
        row += 1


def _it_sheet(b: wp.Book, cfg: dict) -> None:
    ws, _ = b.sheet("3. 전산환경", [64, 2, 20, 6, 6, 6, 2, 36, 22], head=False)
    for i, t in enumerate(["질 문 사 항", None, None, "Y", "N", "N/A", None, "서술내용", "샘플 문서명"], 1):
        if t:
            c = b.head(ws, 1, [t], col=i)
    row = 2
    for q in cfg["it_questionnaire"]:
        kind = q["kind"]
        c = ws.cell(row, 1, q["text"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = wp.font(bold=kind in ("major", "group"), size=11 if kind == "major" else 10)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        if kind == "item":
            for col, lab in ((4, "Y"), (5, "N"), (6, "N/A")):
                b.ask(ws, row, col, f"전산 {row} {lab}", rule="optional")
            b.ask(ws, row, 8, f"전산 {row} 서술내용", rule="optional")
            b.ask(ws, row, 9, f"전산 {row} 샘플 문서명", rule="optional")
        row += 1
    ws.freeze_panes = "A2"


def _supplement_sheet(b: wp.Book, cfg: dict) -> None:
    ws, _ = b.sheet("4. 보완", [2, 18, 44, 28, 28], head=False)
    row = 2
    for blk in cfg["backup"].values():
        ws.cell(row, 2, blk["title"]).font = wp.font(size=12, bold=True)
        row += 2
        b.head(ws, row, blk["heads"], col=2)
        row += 1
        for s in blk["systems"]:
            b.put(ws, row, 2, s)
            for k in range(3):
                b.ask(ws, row, 3 + k, f"보완 {blk['title']} {s} {k}", rule="optional")
            row += 1
        row += 2

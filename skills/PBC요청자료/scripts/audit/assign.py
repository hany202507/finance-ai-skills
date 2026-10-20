"""스태프 배정 (2026-10-12).

분석적 검토 뒤에 영역과 이슈를 감사반 인원에게 배정하고, 그 배정이 다음 할 일 · 빠진 자료 · 요청 메일 · 화면에 쓰이게 한다.
누가 무엇을 맡는지는 회계사가 정한다. 엔진은 배정표의 줄(영역 · 이슈 수 · 금액 · 유의적 여부)만 채우고 담당자 칸은 비워 둔다.

  rows(d)                      배정표 줄 [{no, name, issues, amount, significant, why}]. 조서목록의 조서와 만들 4000 영역
  make(d, now)                 배정표 조서 2110-배정 을 쓴다. 담당자 · 검토자 · 기한은 기입란이고 _색인/배정.json 값으로 채운다
  assign(d, to, reviewers, dues)   배정표를 되읽고 명령줄 값을 덮어 _색인/배정.json 에 쓴다. 조서목록에는 folder.scan 이 옮긴다
  of(d)                        {조서번호: {to, reviewer, due}}
  by_person(d)                 {이름: {"workpapers": [...], "missing": [...]}} (빠진 자료는 관련 조서의 담당자로 묶는다)

감사반 인원은 `1000 감사계약/감사반인원.xlsx` 다(checklists.people). 인원이 적혀 있으면 그 밖의 이름은 받지 않는다.
혼자 하는 감사는 배정하지 않아도 된다. 배정표는 감사반 인원이 둘 이상이거나 assign 을 돌렸을 때만 생긴다.
4000 조서 표지의 수행자 칸은 미리 채우지 않는다. 그 칸은 절차를 실제로 수행한 사람이 적는 기록이다.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl.worksheet.datavalidation import DataValidation

from . import checklists, folder, wp

NO = "2110-배정"
SHEET = "배정"
INDEX = "배정.json"
FIELDS = (("to", "담당자"), ("reviewer", "검토자"), ("due", "기한"))


class AssignError(Exception):
    def __init__(self, message: str, problems: list[str] | None = None):
        super().__init__(message)
        self.problems = problems or []


def _idx(d: Path, name: str, default=None):
    return folder._read(Path(d) / "_색인" / name, default)


def load(d: Path) -> dict:
    return _idx(d, INDEX, None) or {"format_version": 1, "updated_at": None, "assignments": {}}


def of(d: Path) -> dict[str, dict]:
    return load(d).get("assignments") or {}


def people_names(d: Path) -> list[str]:
    return [p["name"] for p in checklists.people(d)[0]]


def rows(d: Path) -> list[dict]:
    """배정표 줄. 조서목록의 조서 전부와, 아직 만들지 않았지만 만들 4000 영역(이슈 대응 · 수행중요성 초과)."""
    d = Path(d)
    std = folder.load_standard()
    names = {str(k): v for k, v in (std.get("workpaper_names") or {}).items()}
    wps = (_idx(d, "조서목록.json", {"workpapers": []}) or {"workpapers": []})["workpapers"]
    issues = [i for i in (_idx(d, "이슈.json", {"issues": []}) or {"issues": []})["issues"] if i.get("status") != "닫힘"]
    out = {w["no"]: {"no": w["no"], "name": w.get("name") or names.get(w["no"], w["no"]), "issues": 0, "amount": None,
                     "significant": "", "why": w.get("why") or ""} for w in wps if w["no"] != NO}
    try:                                         # 시산표가 있으면 만들 4000 영역과 영역 금액을 더한다
        from .workpapers import lead

        ctx = lead.context(d)
        for a in lead.select(d, ctx):
            r = out.setdefault(a["no"], {"no": a["no"], "name": names.get(a["no"], a["no"]), "issues": 0, "amount": None,
                                         "significant": "", "why": ""})
            r["why"] = " · ".join(a["why"])
            r["significant"] = "예"
        for no, r in out.items():
            if folder._is_area(no):
                mine = [x for x in ctx["tb"] if x["area"] == no[4:]]
                if mine:
                    r["amount"] = int(sum(abs(x["disp"]) for x in mine))
    except Exception:  # noqa: BLE001  시산표가 없거나 읽지 못하면 조서목록 줄만 낸다
        pass
    for i in issues:
        no = i.get("response_wp") or (f"4000{i['area']}" if i.get("area") else None)
        if no in out:
            out[no]["issues"] += 1
    order = {str(n): k for k, n in enumerate(n for f in std["folders"] for n in f.get("workpapers", []))}
    return sorted(out.values(), key=lambda r: (order.get(r["no"], len(order)), r["no"]))


def _due(v) -> str | None:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip().replace(".", "-").replace("/", "-")
    if not s:
        return None
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        raise AssignError(f"기한은 2026-02-10 꼴로 적어 주십시오: {v}")
    return s


def make(d: Path, now: str | None = None) -> Path:
    """배정표 조서. 줄은 엔진이, 담당자 · 검토자 · 기한은 회계사가 채운다. 기입란은 모두 선택이라 비워도 조서 검사를 통과한다."""
    d = Path(d)
    cur = of(d)
    ppl = people_names(d)
    b = wp.Book(d, NO, standard="감사기준서 300 문단 11(업무팀원의 지휘 · 감독 · 검토) · 감사기준서 220")
    ws, row = b.sheet(SHEET, widths=[14, 24, 9, 18, 10, 40, 14, 14, 14])
    ws.cell(6, 1, "담당자 · 검토자 · 기한을 적고 python -m audit assign <감사 폴더> 를 돌리면 다음 할 일과 빠진 자료가 담당자별로 묶입니다. "
                  "엔진은 담당자를 정하지 않습니다").font = wp.font(size=9, color=wp.MUTED)
    row = b.head(ws, row, ["조서번호", "영역 · 조서", "이슈 수", "금액", "유의적", "생긴 이유", "담당자", "검토자", "기한"])
    dv = None
    if ppl:
        dv = DataValidation(type="list", formula1='"' + ",".join(ppl) + '"', allow_blank=True)
        ws.add_data_validation(dv)
    data = rows(d)
    first = row
    for r in data:
        a = cur.get(r["no"]) or {}
        b.put(ws, row, 1, r["no"], bold=True)
        b.put(ws, row, 2, r["name"])
        b.put(ws, row, 3, r["issues"] or None, fmt=wp.WON, flag=bool(r["issues"]))
        b.put(ws, row, 4, r["amount"], fmt=wp.WON)
        b.put(ws, row, 5, r["significant"] or None, flag=bool(r["significant"]))
        b.put(ws, row, 6, r["why"], wrap=True, color=wp.MUTED)
        for col, (key, label) in enumerate(FIELDS, 7):
            c = b.ask(ws, row, col, f"{label}: {r['no']}", rule="optional", value=a.get(key), wrap=False)
            if dv is not None and key != "due":
                dv.add(c.coordinate)
        row += 1
    b.put(ws, row, 1, "합계", bold=True)
    b.put(ws, row, 3, f"=SUM(C{first}:C{row - 1})" if data else 0, fmt=wp.WON, bold=True)
    b.verify("배정표 이슈 수 합계", f"='{SHEET}'!C{row}", sum(r["issues"] for r in data))
    row += 2
    src = "1000 감사계약/감사반인원.xlsx" if ppl else "감사반인원.xlsx 가 없거나 비어 있습니다. 이름을 직접 적습니다"
    ws.cell(row, 1, f"감사반 인원: {' · '.join(ppl) if ppl else '-'} ({src})").font = wp.font(size=9, color=wp.MUTED)
    ws.cell(row + 1, 1, "금액은 그 영역 계정의 시산표 잔액 절대값 합입니다. 유의적은 이슈 대응 영역이거나 표시과목이 수행중요성을 넘는 영역입니다").font = \
        wp.font(size=9, color=wp.MUTED)
    ws.freeze_panes = f"C{first}"
    return b.save()


def read_sheet(d: Path) -> dict[str, dict]:
    """배정표에 회계사가 적은 값 {조서번호: {to, reviewer, due}}. 배정표가 없으면 빈 dict."""
    p = wp.path_for(Path(d), NO)
    if not p.exists():
        return {}
    out: dict[str, dict] = {}
    names = {label: key for key, label in FIELDS}
    for x in wp.readback(p):
        m = re.fullmatch(r"(담당자|검토자|기한): (.+)", str(x["label"] or ""))
        if not m or wp._blank(x["value"]):
            continue
        key = names[m.group(1)]
        out.setdefault(m.group(2), {})[key] = _due(x["value"]) if key == "due" else str(x["value"]).strip()
    return out


def assign(d: Path, to: dict | None = None, reviewers: dict | None = None, dues: dict | None = None, now: str | None = None) -> dict:
    """배정을 정한다. 배정표에 적힌 값을 먼저 읽고 명령줄 값(to · reviewers · dues, {조서번호: 값})을 덮는다. 빈 글자는 그 칸을 지운다."""
    d = Path(d)
    at = now or folder._now()
    folder.scan(d, now=now, quiet_log=True)
    known = {r["no"] for r in rows(d)}
    ppl = people_names(d)
    data = load(d)
    cur = {no: dict(v) for no, v in (data.get("assignments") or {}).items()}
    sheet = read_sheet(d)
    if wp.path_for(d, NO).exists():
        cur = {no: dict(v) for no, v in sheet.items()}          # 배정표가 있으면 표가 정본이다(지운 칸은 지워진다)
    problems = []
    for key, given in (("to", to), ("reviewer", reviewers), ("due", dues)):
        for no, v in (given or {}).items():
            if no not in known:
                problems.append(f"조서목록과 만들 영역에 없는 조서번호입니다: {no}")
                continue
            v = str(v or "").strip()
            if key == "due" and v:
                try:
                    v = _due(v)
                except AssignError as e:
                    problems.append(str(e))
                    continue
            if v:
                cur.setdefault(no, {})[key] = v
            else:
                cur.get(no, {}).pop(key, None)
    if ppl:
        for no, a in cur.items():
            for key in ("to", "reviewer"):
                if a.get(key) and a[key] not in ppl:
                    problems.append(f"{no} {dict(FIELDS)[key]} {a[key]} 은 감사반 인원에 없습니다({' · '.join(ppl)})")
    for no, a in cur.items():
        if a.get("to") and a.get("to") == a.get("reviewer"):
            problems.append(f"{no} 담당자와 검토자가 같습니다({a['to']})")
    if problems:
        raise AssignError("배정을 쓰지 않았습니다", problems)
    cur = {no: a for no, a in cur.items() if a}
    before = data.get("assignments") or {}
    folder._write(d / "_색인" / INDEX, {"format_version": 1, "updated_at": at, "assignments": cur})
    path = make(d, now=now)
    folder.scan(d, now=now, quiet_log=True)
    changed = sorted(no for no in set(before) | set(cur) if before.get(no) != cur.get(no))
    case = _idx(d, "감사건.json")
    case.setdefault("log", []).insert(0, {"at": at, "who": "배정", "what": f"조서 {len(cur)}개에 담당자를 적었습니다. 바뀐 조서 {len(changed)}개"
                                          + (": " + ", ".join(changed[:8]) if changed else ""),
                                          "read": [path.relative_to(d).as_posix()], "wrote": [f"_색인/{INDEX}", "_색인/조서목록.json"]})
    folder._write(d / "_색인" / "감사건.json", case)
    folder.validate_index(d)
    return {"assignments": cur, "changed": changed, "path": path, "people": ppl, "unassigned": sorted(known - set(cur))}


def by_person(d: Path) -> dict[str, dict]:
    """담당자별 조서와 빠진 자료. 빠진 자료는 관련 조서의 담당자에게 묶고, 담당자가 없는 것은 「미배정」 에 둔다."""
    d = Path(d)
    cur = of(d)
    out: dict[str, dict] = {}
    for no, a in sorted(cur.items()):
        if a.get("to"):
            out.setdefault(a["to"], {"workpapers": [], "missing": []})["workpapers"].append(no)
    docs = _idx(d, "문서목록.json", {"missing": []}) or {"missing": []}
    for m in docs.get("missing", []):
        who = (cur.get(str(m.get("workpaper") or "")) or {}).get("to") or "미배정"
        out.setdefault(who, {"workpapers": [], "missing": []})["missing"].append(m["item"])
    return out


def ensure_sheet(d: Path, now: str | None = None) -> Path | None:
    """분석 뒤에 부른다. 감사반 인원이 둘 이상이면 배정표를 만들거나 줄을 새로 채운다(담당자 칸은 적힌 값 그대로). 혼자면 만들지 않는다."""
    d = Path(d)
    if len(people_names(d)) < 2 and not wp.path_for(d, NO).exists():
        return None
    path = wp.path_for(d, NO)
    sheet = read_sheet(d)
    if path.exists() and sheet != of(d):
        folder._write(d / "_색인" / INDEX, {"format_version": 1, "updated_at": now or folder._now(), "assignments": sheet})
    old = path.read_bytes() if path.exists() else None
    make(d, now=now)
    if old is not None:
        # 줄이 달라지지 않았으면 앞 파일을 그대로 둔다. 다시 저장하면 해시가 바뀌어 검사 통과 · 검토 끝이 풀린다(apply 와 같은 까닭)
        import tempfile

        from . import apply

        with tempfile.TemporaryDirectory(prefix="audit-assign-") as tmp:
            prev = Path(tmp) / path.name
            prev.write_bytes(old)
            if apply.content_key(prev) == apply.content_key(path):
                path.write_bytes(old)
    return path

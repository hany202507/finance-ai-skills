"""WTB 양식 맞춤 출력 (2026-10-13).

회계사 샘플 정산표(오하임 · 홈앤쇼핑)와 같은 양식으로 WTB 를 쓴다. 숫자는 8100B 와 같은 분석 결과에서 오고, 계산은 전부 엑셀 수식이다.

  재무상태표 · 손익계산서   회사 제시 재무제표(공시 양식)의 계층 아래에 시산표 계정을 줄로 둔다. 열은 회사 계정 · 공시용 계정과목 ·
                          중간감사 · 회사 제시 · 수정 차변 · 수정 대변 · Final · 전기 · 전전기 · 담당자 · 시점별 증감 판정(Materiality 기준)
  수정사항                분개 줄(계정코드 · 차변 · 대변 · 반영 여부). 정산표의 수정 열이 계정코드로 SUMIFS 한다. 자산 · 부채 · 손익 · 자본 영향과 검증
  공시용BS · 공시용PL      정산표 Final 열을 공시용 계정과목으로 SUMIF 해서 공시 양식으로 보인다(회사 제시 · Final · 전기)
  공시용자본변동표         전기말 · 당기말 잔액과 회계사가 적는 변동 줄. 당기순이익은 공시용PL 에서 가져오고 공시용BS 와 대사한다

원칙: 엔진은 숫자를 적고 합계 · 대사는 수식이 한다. 양식에 자리가 없는 계정은 지우지 않고 가까운 줄에 붙이거나 「양식에 없는 계정」 줄로 보인다.
수정사항 반영 여부(예)가 있는 줄만 정산표에 들어간다(8100B 와 같은 약속).

수정사항 줄은 세 곳에서 온다(2026-10-15). 출처 열에 적는다.
  8100A            처리가 「수정」 인 이슈의 수정분개. 재무제표 · 주석 · DART 와 같은 분개다. 반영 여부는 「예」 로 고정하고 8100A 에서 고친다
  8100B 수정 제안   엔진 후보 가운데 8100A 에서 아직 처리를 정하지 않은 이슈의 것. 반영 여부는 앞 WTB 파일 값, 없으면 8100B 값
  (빈 칸)           회계사가 이 시트에 직접 적은 줄. 다시 만들 때 그대로 옮긴다. 재무제표 · DART 에는 들어가지 않으니 8100A 에 옮기라고 알린다
무엇을 담았는지는 _색인/WTB양식.json 에 남긴다. 수정 반영(audit.apply)이 이것을 지금 8100A 와 견주어 다르면 이 파일을 다시 만든다.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from .analysis import fslayout

WON = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GRAY = PatternFill("solid", fgColor="D9D9D9")
HEAD = PatternFill("solid", fgColor="EDEDED")
RED = Font(color="B3261E", size=9)
MUTED = Font(color="707070", size=9)
THIN = Side(style="thin", color="BBBBBB")

BS, IS = "재무상태표", "손익계산서"
ADJ = "수정사항"
Q = lambda s: "'" + s + "'"          # noqa: E731  시트 이름 따옴표(한글 · 공백 안전)

SECTIONS_SIGN = {"자산": 1, "비용": 1, "부채": -1, "자본": -1, "수익": -1}
EXTRA_LABEL = "양식에 없는 계정"


def _norm(s) -> str:
    return re.sub(r"[\s·ㆍ]", "", str(s or ""))


def _sec_group(section: str) -> str:
    return "BS" if section in ("자산", "부채", "자본") else "IS"


# ------------------------------------------------------------------ 줄 모형
class Line:
    """정산표의 물리 줄 하나. kind: section · head · acct · total · extra · note."""

    def __init__(self, kind, label="", level=0, **kw):
        self.kind, self.label, self.level = kind, label, level
        self.code = kw.get("code")
        self.name = kw.get("name")
        self.key = kw.get("key")              # 공시용 계정과목(SUMIF 열쇠)
        self.cur = kw.get("cur")
        self.prior = kw.get("prior")
        self.sec = kw.get("sec")              # 자산 · 부채 · 자본 · 수익 · 비용(계정 줄의 부호 규칙)
        self.children: list[Line] = []
        self.row = None
        self.link = kw.get("link")            # 줄 값을 다른 시트 수식으로 정할 때(당기순손익 줄)
        self.formula = kw.get("formula")      # total 줄의 계산식 만드는 함수 (rows 를 받아 문자열)
        self.cat = kw.get("cat")              # 손익 분류(rev · cogs · sga · nonop_in · nonop_out · tax)
        self.assign = kw.get("assign")
        self.area = kw.get("area")            # 4000 영역(담당자 배정에 쓴다)


def _display_prior(a):
    return a.get("prior_disp") if a.get("prior_disp") is not None else None


def build_lines(d: Path, res: dict) -> dict:
    """공시 양식 + 시산표 계정으로 BS · IS 줄 목록을 세운다. {"BS": [Line], "IS": [Line], "notes": [...]}"""
    accounts = {a["code"]: a for a in res["accounts"]}
    lay = fslayout.layout(d, res["accounts"], res.get("lines"))
    rows = lay["rows"]
    notes = []
    if not rows:
        return {"BS": [], "IS": [], "notes": [lay["basis"]], "accounts": accounts}
    used = {c for r in rows for c in r["codes"]}
    unplaced = [a for a in res["accounts"] if a["code"] not in used and (abs(a.get("cur") or 0) > 0.5 or a.get("basis") == ADJ_ONLY)]
    seen: dict = {}
    for r in rows:                                  # 공시용 계정과목 열쇠는 한 양식 안에서 유일해야 SUMIF 가 맞는다
        if r["kind"] != "acct":
            continue
        k = (r["stmt"], r["label"])
        n = seen.get(k, 0)
        seen[k] = n + 1
        r["_key"] = r["label"] if not n and not any(x is not r and x["kind"] == "acct" and x["stmt"] == r["stmt"] and x["label"] == r["label"] for x in rows) \
            else (f"{r['label']}({r['contra_of']})" if r["contra_of"] else f"{r['label']}({n + 1})")

    out = {"BS": [], "IS": []}
    equity_free_row = None
    head_cat = None
    for r in rows:
        stmt = "BS" if r["stmt"] == "bs" else "IS"
        lines = out[stmt]
        k = r["kind"]
        if k == "section":
            lines.append(Line("section", r["label"]))
        elif k in ("h1", "h2"):
            hc = _is_cat(r["label"]) if stmt == "IS" else None
            if k == "h1":
                head_cat = hc                      # 소제목이 없는 새 묶음이면 앞 묶음의 분류를 물려받지 않는다
            elif hc:
                head_cat = hc
            lines.append(Line("head", r["label"], level=1 if k == "h1" else 2))
        elif k == "total":
            head_cat = None
            lines.append(Line("total", r["label"], prior=r["prior_signed"]))
        elif k == "acct":
            codes = [c for c in r["codes"] if c in accounts]
            if not codes:
                ln = Line("acct", r["label"], key=r["_key"], name=f"({r['label']}) 이번 시산표에 계정 없음", prior=r["prior_signed"],
                          sec=_section_of(stmt, r, accounts), cur=0, cat=(head_cat or _is_cat(r["label"])) if stmt == "IS" else None)
                if stmt == "BS" and _is_equity_free(r["label"]):
                    equity_free_row = (stmt, ln)
                lines.append(ln)
                continue
            first = True
            for c in codes:
                a = accounts[c]
                ln = Line("acct", r["label"], code=c, name=a["name"], key=r["_key"], cur=a["cur_disp"], sec=a["section"], area=a.get("area"),
                          prior=(_display_prior(a) if _display_prior(a) is not None else (r["prior_signed"] if first else None)),
                          cat=(head_cat or _acct_cat(a)) if stmt == "IS" else None)
                lines.append(ln)
                first = False
    # 양식에 자리가 없는 계정: 자본의 이월결손금 같은 것은 미처리결손금 줄에 붙이고, 나머지는 양식에 없는 계정 줄로 모은다
    extra = {"BS": [], "IS": []}
    for a in unplaced:
        stmt = _sec_group(a["section"])
        if a["section"] == "자본" and equity_free_row and re.search(r"이월|결손|이익잉여|미처분", a["name"]):
            host = equity_free_row[1]
            ln = Line("acct", host.label, code=a["code"], name=a["name"], key=host.key, cur=a["cur_disp"], sec="자본")
            idx = out["BS"].index(host)
            if host.code is None and host.cur == 0 and not host.link:
                out["BS"][idx] = ln                          # 비어 있던 자리표시 줄을 실제 계정 줄로 바꾼다
                ln.prior = host.prior
            else:
                out["BS"].insert(idx + 1, ln)
            notes.append(f"{a['code']} {a['name']} 은 공시 양식에 같은 이름이 없어 「{host.label}」 줄에 붙였습니다")
            continue
        if stmt == "BS" and _place_new_line(out["BS"], a, notes):
            continue
        extra[stmt].append(Line("acct", EXTRA_LABEL, code=a["code"], name=a["name"], key=f"{EXTRA_LABEL}({a['section']})", cur=a["cur_disp"], sec=a["section"], area=a.get("area"),
                                cat=("nonop_in" if a["section"] == "수익" else "nonop_out") if stmt == "IS" else None))
        notes.append(f"{a['code']} {a['name']} 은 공시 양식에서 과목을 찾지 못했습니다. 「{EXTRA_LABEL}」 에 두었으니 과목을 정해 주십시오")
    _place_extras(out, extra)
    if equity_free_row:
        host = equity_free_row[1]
        # 시산표는 마감 전이라 당기순손익이 자본에 없다. 손익계산서 Final 당기순손익을 자본에 연동한다
        idx = max(i for i, x in enumerate(out["BS"]) if x.key == host.key)
        out["BS"].insert(idx + 1, Line("acct", host.label, name="당기순손익(손익계산서 연동)", key=host.key, sec="자본", link="NI"))
    else:
        eq_total = next((i for i, x in enumerate(out["BS"]) if x.kind == "total" and "자본총계" in _norm(x.label)), None)
        ni = Line("acct", "당기순손익", name="당기순손익(손익계산서 연동)", key="당기순손익", sec="자본", link="NI")
        if eq_total is not None:
            out["BS"].insert(eq_total, ni)
        else:
            out["BS"].append(ni)
    return {"BS": out["BS"], "IS": out["IS"], "notes": notes, "accounts": accounts, "layout": lay}


def _total_key(n: str) -> str | None:
    """BS 합계 줄 이름 → 자산 · 부채 · 자본 · 합계. 「부채및자본총계」 가 「자본총계」 를 품고 있으니 합계를 먼저 가른다."""
    if "부채및자본" in n or "부채와자본" in n:
        return "합계"
    if "자산총계" in n:
        return "자산"
    if n == "부채총계":
        return "부채"
    if "자본총계" in n:
        return "자본"
    return None


def _place_new_line(bs: list, a: dict, notes: list) -> bool:
    """전기 양식에 없는 당기 표시과목(예: 당기에 처음 생긴 무형자산)을 유동 · 비유동 묶음 끝에 새 줄로 둔다(2026-10-16).
    묶음은 8810 재무제표와 같은 계정코드 범위(report_fs.yaml groups)로 정한다. 양식에 그 묶음 제목이 없으면 놓지 않는다.
    같은 표시과목의 계정이 여럿이면 한 과목 줄 아래에 모인다(공시용 계정과목 열쇠가 같다)."""
    from .report.fs import group_of

    line = a.get("line")
    g = group_of(a["code"])
    if not line or not g or g["side"] not in ("자산", "부채"):
        return False
    gk = _norm(g["key"])
    head = next((i for i, x in enumerate(bs) if x.kind == "head" and _norm(x.label) == gk), None)
    same = [i for i, x in enumerate(bs) if x.kind == "acct" and x.label == line and x.key == line]
    if same:
        pos = same[-1] + 1
    elif head is not None:
        pos = next((i for i in range(head + 1, len(bs)) if bs[i].kind in ("section", "total") or (bs[i].kind == "head" and bs[i].level <= bs[head].level)),
                   len(bs))
    else:
        # 묶음 제목이 없는 양식(과목만 늘어놓은 전기 재무제표): 같은 묶음 계정의 마지막 줄 뒤, 없으면 자산총계 · 부채총계 앞
        mine = [i for i, x in enumerate(bs) if x.kind == "acct" and x.code and (group_of(x.code) or {}).get("key") == g["key"]]
        total = next((i for i, x in enumerate(bs) if x.kind == "total" and _total_key(_norm(x.label)) == g["side"]), None)
        if mine:
            pos = mine[-1] + 1
        elif total is not None:
            pos = total
        else:
            return False
    bs.insert(pos, Line("acct", line, code=a["code"], name=a["name"], key=line, cur=a["cur_disp"], sec=a["section"], area=a.get("area"),
                        prior=0 if a.get("basis") != ADJ_ONLY else None))
    notes.append(f"{a['code']} {a['name']}: 전기 공시 양식에 없는 당기 과목이라 「{line}」 줄을 {g['key']} 아래에 새로 두었습니다. 과목 이름과 위치를 확인해 주십시오")
    return True


def _place_extras(out, extra):
    """양식에 없는 계정을 자기 구분의 합계 앞에 둔다(그래야 정산표 합계와 균형이 시산표와 맞는다).
    재무상태표는 자산 · 부채 · 자본별로, 손익계산서는 수익 · 비용별로 세전이익 계산 앞에 두고 영업외수익 · 영업외비용으로 더한다."""
    for sec in ("자산", "부채", "자본"):
        rows = [x for x in extra["BS"] if x.sec == sec]
        if not rows:
            continue
        idx = next((i for i, x in enumerate(out["BS"]) if x.kind == "total" and _total_key(_norm(x.label)) == sec), len(out["BS"]))
        out["BS"][idx:idx] = [Line("head", f"{EXTRA_LABEL}({sec})", level=1)] + rows
    for sec, cat in (("수익", "nonop_in"), ("비용", "nonop_out")):
        rows = [x for x in extra["IS"] if x.sec == sec]
        if not rows:
            continue
        idx = next((i for i, x in enumerate(out["IS"]) if x.kind == "total" and "차감전" in _norm(x.label)), len(out["IS"]))
        out["IS"][idx:idx] = [Line("head", f"{EXTRA_LABEL}({sec})", level=1)] + rows


def _acct_cat(a) -> str | None:
    """손익 계정의 분류. 분석이 붙인 표시과목(매출액 · 매출원가 · 판매비와관리비 · 금융수익 …)에서 정한다. 없으면 구분(수익 · 비용)으로 영업외로 둔다."""
    c = _is_cat(a.get("line") or "") or _is_cat(a.get("name") or "")
    if c:
        return c
    return "nonop_in" if a.get("section") == "수익" else "nonop_out"


def _is_equity_free(label) -> bool:
    return bool(re.search(r"미처리결손금|미처분이익잉여금|이익잉여금|결손금", str(label)))


def _section_of(stmt, r, accounts):
    return "자본" if _is_equity_free(r["label"]) and stmt == "BS" else None


def _is_cat(label) -> str | None:
    n = _norm(label)
    if n.startswith("매출액") or n in ("수익", "영업수익"):
        return "rev"
    if "매출원가" in n or "서비스공급원가" in n:
        return "cogs"
    if "판매" in n and "관리" in n:
        return "sga"
    if n in ("영업외수익", "금융수익", "기타수익"):
        return "nonop_in"
    if n in ("영업외비용", "금융원가", "금융비용", "기타비용"):
        return "nonop_out"
    if "법인세" in n:
        return "tax"
    return None


# ------------------------------------------------------------------ 시트 쓰기
class Cols:
    """정산표 시트의 열 위치(BS 와 IS 는 계층 열 수가 달라 다르다)."""

    def __init__(self, stmt):
        if stmt == "BS":
            self.code, self.lab, self.name, self.key = 1, (2, 3, 4, 5), 6, 7
            self.mid, self.cur, self.dr, self.cr, self.final, self.p1, self.p2, self.assign = 8, 9, 10, 11, 12, 13, 14, 15
        else:
            self.code, self.lab, self.name, self.key = 1, (2, 3, 4), 5, 6
            self.mid, self.cur, self.dr, self.cr, self.final, self.p1, self.p2, self.assign = 7, 8, 9, 10, 11, 12, 13, 14


def _hier_col(c: Cols, kind, level):
    if kind == "section":
        return c.lab[0]
    if kind == "total":
        return c.lab[0]
    if kind == "head":
        return c.lab[min(level, len(c.lab) - 1)]
    return c.name


def write_statement(ws, stmt, lines: list[Line], adj_range, mat: dict, assigns: dict):
    """한 정산표 시트를 쓴다. 줄 row 번호를 매기고, 계정 줄 · 소계 · 합계를 수식으로 쓴다."""
    c = Cols(stmt)
    ws.sheet_view.showGridLines = False
    widths = {1: 9, 2: 14, 3: 14, 4: 16, 5: 14, 6: 26, 7: 24} if stmt == "BS" else {1: 9, 2: 16, 3: 14, 4: 26, 5: 24, 6: 4}
    for i in range(1, 30):
        ws.column_dimensions[L(i)].width = widths.get(i, 15)
    title = "정산표(재무상태표)" if stmt == "BS" else "정산표(손익계산서)"
    ws.cell(1, 2, title).font = Font(bold=True, size=13)
    # 시점별 증감 판정 기준(홈앤쇼핑 샘플): 비율 기준 · 금액 기준(Materiality × 비율)
    base = c.assign + 2
    ws.cell(2, base, "구분").font = Font(bold=True)
    ws.cell(2, base + 1, "Materiality").font = Font(bold=True)
    ws.cell(2, base + 2, "비율기준").font = Font(bold=True)
    ws.cell(2, base + 3, "금액기준").font = Font(bold=True)
    for i, nm in enumerate(("기말감사(최초)", "기말감사(최종)")):
        r = 3 + i
        ws.cell(r, base, nm)
        ws.cell(r, base + 1, mat.get("overall") if i == 0 else f"={L(base + 1)}3").number_format = WON
        ws.cell(r, base + 2, mat.get("ratio", 0.2)).number_format = PCT
        ws.cell(r, base + 3, f"={L(base + 1)}{r}*{L(base + 2)}{r}").number_format = WON
        if i == 0:
            ws.cell(r, base + 1).fill = YELLOW
            ws.cell(r, base + 2).fill = YELLOW
    hdr = 6
    heads = {c.code: "코드", c.name: "회사 계정과목", c.key: "공시용 계정과목", c.mid: "중간감사", c.cur: "회사 제시", c.dr: "수정 차변",
             c.cr: "수정 대변", c.final: "Final", c.p1: "전기 Final", c.p2: "전전기 Final", c.assign: "담당자"}
    for col, t in heads.items():
        h = ws.cell(hdr, col, t)
        h.font, h.fill = Font(bold=True), HEAD
        h.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, t in enumerate(("최초 Diff %", "최초 Diff 금액", "최종 Diff %", "최종 Diff 금액", "관련 조서(t/m)")):
        h = ws.cell(hdr, base + 4 + i, t)
        h.font, h.fill = Font(bold=True), HEAD
        h.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[hdr].height = 30
    if stmt == "IS":
        h = ws.cell(hdr, c.assign + 1, "분류")
        h.font, h.fill = Font(bold=True, color="707070"), HEAD
    row = hdr + 1
    for ln in lines:
        ln.row = row
        row += 1 if ln.kind != "section" else 1
    A = adj_range                          # 수정사항 시트의 코드 · 차변 · 대변 · 반영 범위(문자열)
    for ln in lines:
        r = ln.row
        if ln.kind == "section":
            ws.cell(r, c.lab[0], ln.label).font = Font(bold=True)
            for col in range(1, c.assign + 1):
                ws.cell(r, col).fill = GRAY
        elif ln.kind == "head":
            ws.cell(r, _hier_col(c, "head", ln.level), ln.label).font = Font(bold=ln.level == 1)
        elif ln.kind == "total":
            ws.cell(r, c.lab[0], ln.label).font = Font(bold=True)
        elif ln.kind == "acct":
            if ln.code:
                ws.cell(r, c.code, ln.code)
            ws.cell(r, c.name, ln.name)
            ws.cell(r, c.key, ln.key).font = Font(color="1F4E79")
            if stmt == "IS" and ln.cat:
                ws.cell(r, c.assign + 1, ln.cat).font = MUTED
            who = assigns.get(ln.area or "") or ""
            ws.cell(r, c.assign, who).fill = YELLOW
            ws.cell(r, c.mid).fill = YELLOW
            ws.cell(r, c.p2).fill = YELLOW
    # 계층 합계: head 줄은 아래 계정 줄과 하위 head 의 합. 스택으로 부모를 정한다.
    # 합계 줄(자산총계 등)은 직전 합계 줄 뒤의 최상위 줄을 더한다(구분 제목 줄이 없는 양식도 있다)
    stack: list[Line] = []
    acc: list[Line] = []
    for ln in lines:
        if ln.kind == "section":
            stack = []
        elif ln.kind == "head":
            while stack and stack[-1].level >= ln.level:
                stack.pop()
            (stack[-1].children if stack else acc).append(ln)
            stack.append(ln)
        elif ln.kind == "acct":
            (stack[-1].children if stack else acc).append(ln)
        elif ln.kind == "total":
            ln.top, acc, stack = acc, [], []
    sums = lambda col, kids: "=SUM(" + ",".join(f"{L(col)}{k.row}" for k in kids) + ")" if kids else None     # noqa: E731
    sign_of = lambda ln: SECTIONS_SIGN.get(ln.sec or "", 1)                                                       # noqa: E731
    for ln in lines:
        r = ln.row
        if ln.kind == "acct":
            sg = sign_of(ln)
            if ln.link == "NI":
                continue                                             # 당기순손익 줄은 손익계산서가 만들어진 뒤 쓴다
            ws.cell(r, c.cur, ln.cur).number_format = WON
            if ln.code:
                dr = f"=SUMIFS({A['dr']},{A['code']},$A{r},{A['use']},\"예\")"
                cr = f"=SUMIFS({A['cr']},{A['code']},$A{r},{A['use']},\"예\")"
                ws.cell(r, c.dr, dr).number_format = WON
                ws.cell(r, c.cr, cr).number_format = WON
            if stmt == "BS":
                f = f"={L(c.cur)}{r}+{L(c.dr)}{r}-{L(c.cr)}{r}" if sg == 1 else f"={L(c.cur)}{r}-{L(c.dr)}{r}+{L(c.cr)}{r}"
            else:
                f = f"={L(c.cur)}{r}+{L(c.dr)}{r}-{L(c.cr)}{r}" if sg == 1 else f"={L(c.cur)}{r}-{L(c.dr)}{r}+{L(c.cr)}{r}"
            ws.cell(r, c.final, f).number_format = WON
            if ln.prior is not None:
                ws.cell(r, c.p1, ln.prior).number_format = WON
        elif ln.kind == "head":
            if ln.children:
                for col in (c.mid, c.cur, c.final, c.p1, c.p2):
                    f = sums(col, ln.children)
                    cell = ws.cell(r, col, f)
                    cell.number_format = WON
                    cell.font = Font(bold=True) if ln.level == 1 else Font()
    # 합계 줄과 손익 계산 줄
    by_label = {}
    for ln in lines:
        if ln.kind == "head":
            by_label.setdefault(_norm(ln.label), ln)
    return c, None, by_label


def finish_bs(ws, c, lines, section_heads, ni_formula_cells):
    """재무상태표 합계 · 검산. 합계 줄은 직전 합계 줄 뒤의 최상위 줄을 더하고, 부채및자본총계는 부채 + 자본이다."""
    total_row = {}
    for ln in lines:
        if ln.kind == "total":
            key = _total_key(_norm(ln.label))
            if key:
                total_row[key] = ln
    for key, ln in total_row.items():
        r = ln.row
        for col in (c.mid, c.cur, c.final, c.p1, c.p2):
            if key == "합계":
                f = f"={L(col)}{total_row['부채'].row}+{L(col)}{total_row['자본'].row}" if "부채" in total_row and "자본" in total_row else None
            else:
                kids = getattr(ln, "top", None) or []
                f = "=SUM(" + ",".join(f"{L(col)}{k.row}" for k in kids) + ")" if kids else None
            if f:
                cell = ws.cell(r, col, f)
                cell.number_format, cell.font = WON, Font(bold=True)
    return total_row


def add_diff_columns(ws, c, lines, base):
    """시점별 증감 판정(홈앤쇼핑 샘플): 전기 대비 증감률이 비율 기준, 증감액이 금액 기준을 넘을 때만 값을 보인다."""
    d1, d2, d3, d4, dt = base + 4, base + 5, base + 6, base + 7, base + 8
    ratio1, amt1, ratio2, amt2 = f"${L(base + 2)}$3", f"${L(base + 3)}$3", f"${L(base + 2)}$4", f"${L(base + 3)}$4"
    for ln in lines:
        if ln.kind not in ("acct", "head") or (ln.kind == "head" and ln.level != 1):
            continue
        r = ln.row
        P, Cu, F = f"{L(c.p1)}{r}", f"{L(c.cur)}{r}", f"{L(c.final)}{r}"
        ws.cell(r, d1, f'=IF({P}=0,0,IF(ABS(({Cu}-{P})/{P})>{ratio1},({Cu}-{P})/{P},0))').number_format = PCT
        ws.cell(r, d2, f"=IF(ABS({Cu}-{P})>{amt1},{Cu}-{P},0)").number_format = WON
        ws.cell(r, d3, f'=IF({P}=0,0,IF(ABS(({F}-{P})/{P})>{ratio2},({F}-{P})/{P},0))').number_format = PCT
        ws.cell(r, d4, f"=IF(ABS({F}-{P})>{amt2},{F}-{P},0)").number_format = WON
        ws.cell(r, dt).fill = YELLOW


SRC_A, SRC_B = "8100A", "8100B 수정 제안"
ADJ_ONLY = "수정분개 계정"            # 시산표에 없고 수정분개에만 나오는 계정(예: 미지급비용). 잔액 0 줄로 둔다(2026-10-16)
RECORD = "WTB양식.json"
ADJ_HEAD_ROW = 4
ADJ_TOTAL = "합계(반영 줄)"


def _prev_adj_rows(prev: Path | None) -> list[dict]:
    """앞서 만든 WTB 파일의 수정사항 줄. 회계사가 직접 적은 줄과 반영 여부를 다시 만들 때 옮기려고 읽는다."""
    if not prev or not Path(prev).exists():
        return []
    from openpyxl import load_workbook

    try:
        wb = load_workbook(prev, read_only=True)
    except Exception:  # noqa: BLE001  깨진 파일이면 옮기지 않는다
        return []
    try:
        if ADJ not in wb.sheetnames:
            return []
        out = []
        for row in wb[ADJ].iter_rows(min_row=ADJ_HEAD_ROW + 1, max_col=14, values_only=True):
            row = list(row) + [None] * (14 - len(row))
            if row[3] == ADJ_TOTAL:
                break
            num = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None  # noqa: E731
            code = str(row[2]).strip() if row[2] not in (None, "") else None
            dr, cr = num(row[4]), num(row[5])
            use = str(row[6]).strip() if row[6] not in (None, "") else None
            if code is None and not dr and not cr and not row[1]:
                continue
            out.append({"no": row[0], "title": row[1], "code": code, "dr": dr, "cr": cr, "use": use, "basis": row[7],
                        "source": str(row[13]).strip() if row[13] else ""})
        return out
    finally:
        wb.close()


def adjustment_rows(d: Path, res: dict, prev: Path | None = None) -> dict:
    """수정사항 시트에 쓸 줄과 그 근거. {"rows": [...], "adjust": [[이슈, 코드, 차변, 대변]], "wtb": [8100B 반영 예 이름표],
    "manual": 직접 적은 줄 수, "notes": [...]}. 8100A 가 정본이고 엔진 후보는 8100A 에서 처리를 정하지 않은 이슈만 싣는다."""
    from . import figures
    from .report import misstatement
    from .report.base import prev_inputs, wp_file

    d = Path(d)
    notes, rows, adjust = [], [], []
    decided: set = set()
    if wp_file(d, "8100A") is not None:
        tb = figures.current_tb(d)
        st = misstatement.read(d)
        decided = {n for n, (dec, _) in st["decisions"].items() if dec} | {e["issue"] for e in st["entries"] if e["issue"] is not None}
        if tb:
            pl = misstatement.plan(d, tb, st)
            for e in pl["adjust"]:
                rows.append({"no": e["issue"], "title": e.get("memo") or f"8100A 이슈 {e['issue']} 수정분개", "code": str(e["code"]),
                             "dr": e["dr"] or None, "cr": e["cr"] or None, "use": "예", "basis": "8100A 처리 「수정」", "source": SRC_A,
                             "name": e.get("name"), "line": e.get("line")})
                adjust.append([e["issue"], str(e["code"]), e["dr"] or 0, e["cr"] or 0])
            if pl["problems"]:
                notes.append(f"8100A 기입에 맞지 않는 곳이 {len(pl['problems'])}건 있어 그 분개는 싣지 않았습니다: " + "; ".join(pl["problems"][:3]))
    prev_rows = _prev_adj_rows(prev)
    # 앞 WTB 에서 회계사가 적은 반영 여부만 앞선다. 빈 칸이면 8100B 값을 쓴다(2026-10-16. 빈 칸이 8100B 「예」 를 지우던 것)
    prev_use = {(str(r["no"]), r["code"]): r["use"] for r in prev_rows if r["source"] == SRC_B and str(r["use"] or "").strip()}
    b_vals = prev_inputs(wp_file(d, "8100B"))
    wtb_yes = sorted(lab[len("반영 여부: "):] for (sh, lab), v in b_vals.items()
                     if sh == "수정 제안" and str(lab).startswith("반영 여부: ") and str(v or "").strip() == "예")
    differ = 0
    for a in res.get("adjustments") or []:
        if a["issue"] is not None and a["issue"] in decided:
            continue
        for ln in a["lines"]:
            b = b_vals.get(("수정 제안", f"반영 여부: 이슈 {a['issue']} {ln['code']}"))
            b = str(b).strip() if b not in (None, "") else None
            use = prev_use.get((str(a["issue"]), str(ln["code"])), b)
            differ += use != b
            rows.append({"no": a["issue"], "title": a["title"], "code": str(ln["code"]), "dr": ln["dr"] or None, "cr": ln["cr"] or None,
                         "use": use, "basis": a["basis"], "source": SRC_B, "name": ln.get("name")})
    if differ:
        notes.append(f"엔진 후보 {differ}줄은 반영 여부가 8100B 와 다릅니다(이 WTB 에서 고친 값을 썼습니다). 계정별 조서는 8100B 값을 봅니다")
    manual = [dict(r, source="") for r in prev_rows if not r["source"]]
    rows += manual
    if manual:
        notes.append(f"수정사항 시트에 직접 적은 줄 {len(manual)}개를 옮겼습니다. 정산표에는 들어가나 재무제표 · 주석 · DART 에는 들어가지 않습니다. "
                     "8100A 에 이슈와 수정분개로 옮기고 처리를 「수정」 으로 적어 주십시오")
    return {"rows": rows, "adjust": sorted(adjust, key=lambda x: json.dumps(x, ensure_ascii=False)), "wtb": wtb_yes, "manual": len(manual), "notes": notes}


def with_adjust_accounts(res: dict, rows: list[dict]) -> tuple[dict, list[str]]:
    """수정사항 줄의 계정 가운데 시산표에 없는 것을 잔액 0 계정으로 더한 분석 결과와 안내 문장.
    이름 · 표시과목은 8100A 의 새 계정명 · 새 표시과목, 엔진 후보의 계정명, 계정 사전(accounts.yaml) 순으로 정한다.
    줄이 없으면 수정사항의 SUMIFS 가 그 금액을 정산표 어디에도 싣지 못해 재무상태표가 맞지 않는다."""
    from .analysis import defs

    have = {a["code"] for a in res["accounts"]}
    extra, notes = [], []
    for x in rows:
        code = x.get("code")
        if not code or code in have:
            continue
        have.add(code)
        m = defs.map_account(code, x.get("name") or "")
        name = x.get("name") or m.get("label") or code
        line = x.get("line") or m.get("label") or name
        extra.append({"code": code, "name": name, "line": line, "compare_line": line, "id": m.get("id"), "area": m.get("area"),
                      "section": m.get("section"), "basis": ADJ_ONLY, "contra": bool(m.get("contra")), "cur": 0, "prior": 0,
                      "cur_disp": 0, "prior_disp": None})
        notes.append(f"{code} {name} 은 시산표에 없고 수정사항에만 있어 잔액 0 줄로 더했습니다(표시과목 {line}, 구분 {m.get('section') or '모름'})")
    if not extra:
        return res, []
    return dict(res, accounts=list(res["accounts"]) + extra), notes


def write_adjustments(wb, rows: list[dict]):
    ws = wb.create_sheet(ADJ)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate((6, 28, 11, 24, 16, 16, 9, 40, 14, 14, 14, 14, 14, 16), 1):
        ws.column_dimensions[L(i)].width = w
    ws.cell(1, 2, "수정사항").font = Font(bold=True, size=13)
    ws.cell(2, 2, "분개 한 줄에 계정코드 하나. 반영 여부가 「예」 인 줄만 정산표의 수정 열에 들어간다. 노란 칸은 감사인이 적는다. "
                  "출처가 8100A 인 줄은 재무제표 · DART 와 같은 분개라 8100A 에서 고친다").font = MUTED
    heads = ["번호", "내역", "계정코드", "계정명", "차변", "대변", "반영 여부", "근거", "자산 영향", "부채 영향", "손익 영향", "자본 영향", "구분", "출처"]
    hr = ADJ_HEAD_ROW
    for i, t in enumerate(heads, 1):
        h = ws.cell(hr, i, t)
        h.font, h.fill = Font(bold=True), HEAD
        h.alignment = Alignment(horizontal="center", wrap_text=True)
    r = hr + 1
    first = r
    for x in rows:
        ws.cell(r, 1, x["no"])
        ws.cell(r, 2, x["title"]).alignment = Alignment(wrap_text=True)
        ws.cell(r, 3, x["code"])
        ws.cell(r, 5, x["dr"] or None).number_format = WON
        ws.cell(r, 6, x["cr"] or None).number_format = WON
        ws.cell(r, 7, x["use"])
        ws.cell(r, 8, x["basis"]).alignment = Alignment(wrap_text=True)
        ws.cell(r, 14, x["source"] or None).font = MUTED
        if x["source"] == SRC_A:
            for col in (1, 2, 3, 5, 6, 7):
                ws.cell(r, col).fill = GRAY
        elif x["source"] == SRC_B:
            ws.cell(r, 7).fill = YELLOW
        else:
            for col in (2, 3, 5, 6, 7, 8):
                ws.cell(r, col).fill = YELLOW
        r += 1
    spare = max(40, r - first + 20)
    for _ in range(spare):
        for col in (2, 3, 5, 6, 7, 8):
            ws.cell(r, col).fill = YELLOW
        r += 1
    last = r - 1
    dv = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G{first}:G{last}")
    # 계정명 · 구분 · 영향 열은 _계정 표에서 찾는다
    for rr in range(first, last + 1):
        ws.cell(rr, 4, f'=IF(C{rr}="","",IFERROR(INDEX(_계정!$B:$B,MATCH(C{rr},_계정!$A:$A,0)),"(코드 없음)"))')
        ws.cell(rr, 13, f'=IF(C{rr}="","",IFERROR(INDEX(_계정!$C:$C,MATCH(C{rr},_계정!$A:$A,0)),""))')
        eff = f'IF($G{rr}="예",N($E{rr})-N($F{rr}),0)'
        ws.cell(rr, 9, f'=IF($M{rr}="자산",{eff},0)').number_format = WON
        ws.cell(rr, 10, f'=IF($M{rr}="부채",-{eff},0)').number_format = WON
        ws.cell(rr, 11, f'=IF(OR($M{rr}="수익",$M{rr}="비용"),-{eff},0)').number_format = WON
        ws.cell(rr, 12, f'=IF($M{rr}="자본",-{eff},0)').number_format = WON
    tot = last + 2
    ws.cell(tot, 4, ADJ_TOTAL).font = Font(bold=True)
    for col, nm in ((5, "E"), (6, "F")):
        ws.cell(tot, col, f'=SUMIFS({nm}{first}:{nm}{last},$G{first}:$G{last},"예")').number_format = WON
    for col in (9, 10, 11, 12):
        ws.cell(tot, col, f"=SUM({L(col)}{first}:{L(col)}{last})").number_format = WON
    ws.cell(tot + 1, 4, "검증: 차변 − 대변 (0)").font = Font(bold=True)
    ws.cell(tot + 1, 5, f"=E{tot}-F{tot}").number_format = WON
    ws.cell(tot + 2, 4, "검증: 자산 − 부채 − 손익 − 자본 영향 (0)").font = Font(bold=True)
    ws.cell(tot + 2, 5, f"=I{tot}-J{tot}-K{tot}-L{tot}").number_format = WON
    ws.cell(tot + 3, 4, "최종 반영 손익").font = Font(bold=True)
    ws.cell(tot + 3, 5, f"=K{tot}").number_format = WON
    ws.cell(tot + 4, 4, "8100A 수정분개 차변 합(반영 줄)").font = Font(bold=True)
    ws.cell(tot + 4, 5, f'=SUMIFS(E{first}:E{last},$G{first}:$G{last},"예",$N{first}:$N{last},"{SRC_A}")').number_format = WON
    ws.freeze_panes = f"A{hr + 1}"
    rng = lambda col: f"{Q(ADJ)}!${col}${first}:${col}${last}"        # noqa: E731
    return {"code": rng("C"), "dr": rng("E"), "cr": rng("F"), "use": rng("G")}, {"dr_cr": f"{Q(ADJ)}!E{tot + 1}", "eff": f"{Q(ADJ)}!E{tot + 2}",
                                                                              "ni": f"{Q(ADJ)}!E{tot + 3}", "a_dr": f"{Q(ADJ)}!E{tot + 4}"}


def write_accounts_table(wb, accounts: dict):
    ws = wb.create_sheet("_계정")
    ws.append(["코드", "계정명", "구분"])
    for a in accounts.values():
        ws.append([a["code"], a["name"], a["section"]])
    ws.sheet_state = "hidden"


def _mat(res) -> dict:
    inp = res.get("inputs") or {}
    ov = (inp.get("materiality.overall") or {}).get("value")
    return {"overall": ov, "ratio": 0.2}


def _assigns(d: Path) -> dict:
    """스태프 배정(assign.of)이 있으면 영역 담당자를 {영역: 이름} 으로 돌려준다. 정산표 계정 줄은 계정의 영역으로 찾는다. 없으면 빈 칸."""
    from . import assign

    try:
        cur = assign.of(Path(d))
    except Exception:  # noqa: BLE001  배정 파일이 없거나 깨져도 WTB 는 만든다
        return {}
    return {no[4:]: a["to"] for no, a in (cur or {}).items() if str(no).startswith("4000") and a.get("to")}


# ------------------------------------------------------------------ 공시용 시트
def write_disclosure(wb, name, lines: list[Line], src_sheet: str, c: Cols, label_note: str):
    """공시용BS · 공시용PL. 계정 줄은 정산표를 공시용 계정과목으로 SUMIF 하고 소계 · 합계는 정산표와 같은 식으로 더한다."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for col, w in {1: 3, 2: 38, 3: 18, 4: 18, 5: 18, 6: 18, 7: 12}.items():
        ws.column_dimensions[L(col)].width = w
    ws.cell(1, 2, name).font = Font(bold=True, size=13)
    ws.cell(2, 2, label_note).font = MUTED
    heads = ["과목", "회사 제시", "Final", "전기", "증감(Final − 전기)", "증감률"]
    for i, t in enumerate(heads, 2):
        h = ws.cell(4, i, t)
        h.font, h.fill = Font(bold=True), HEAD
        h.alignment = Alignment(horizontal="center")
    S = Q(src_sheet)
    keycol, cur, fin, p1 = L(c.key), L(c.cur), L(c.final), L(c.p1)
    row = 5
    done: dict[str, int] = {}
    first_of_key = {}
    for ln in lines:
        if ln.kind == "acct" and ln.key not in first_of_key:
            first_of_key[ln.key] = ln
    rowmap: dict[int, int] = {}
    seq = []
    seen_keys = set()
    for ln in lines:
        if ln.kind == "acct":
            if ln.key in seen_keys:
                continue
            seen_keys.add(ln.key)
        seq.append(ln)
    for ln in seq:
        rowmap[id(ln)] = row
        row += 1
    stack, acc = [], []
    for ln in seq:
        r = rowmap[id(ln)]
        if ln.kind == "section":
            ws.cell(r, 2, ln.label).font = Font(bold=True)
            for col in range(2, 8):
                ws.cell(r, col).fill = GRAY
            stack = []
        elif ln.kind == "head":
            while stack and stack[-1].level >= ln.level:
                stack.pop()
            ln._dkids = []
            (stack[-1]._dkids if stack else acc).append(ln)
            stack.append(ln)
            ws.cell(r, 2, ("  " * (ln.level - 1)) + ln.label).font = Font(bold=ln.level == 1)
        elif ln.kind == "acct":
            ws.cell(r, 2, ("    " if stack else "  ") + ln.key)
            ws.cell(r, 1, ln.key).font = Font(color="FFFFFF", size=1)
            for col, sc in ((3, cur), (4, fin), (5, p1)):
                ws.cell(r, col, f"=SUMIF({S}!${keycol}:${keycol},$A{r},{S}!${sc}:${sc})").number_format = WON
            (stack[-1]._dkids if stack else acc).append(ln)
        elif ln.kind == "total":
            ln._dtop, acc, stack = acc, [], []
            ws.cell(r, 2, ln.label).font = Font(bold=True)
    # 소계 · 증감
    for ln in seq:
        r = rowmap[id(ln)]
        if ln.kind == "head" and getattr(ln, "_dkids", None):
            for col in (3, 4, 5):
                cell = ws.cell(r, col, "=SUM(" + ",".join(f"{L(col)}{rowmap[id(k)]}" for k in ln._dkids) + ")")
                cell.number_format = WON
                cell.font = Font(bold=ln.level == 1)
        if ln.kind in ("acct", "head") or ln.kind == "total":
            ws.cell(r, 6, f"=D{r}-E{r}").number_format = WON
            ws.cell(r, 7, f'=IF(OR(E{r}="",E{r}=0),"",F{r}/ABS(E{r}))').number_format = PCT
    return ws, seq, rowmap, None


def _total_formulas_bs(ws, seq, rowmap, sec_heads):
    tr = {}
    for ln in seq:
        if ln.kind == "total":
            key = _total_key(_norm(ln.label))
            if key:
                tr[key] = ln
    for key, ln in tr.items():
        r = rowmap[id(ln)]
        for col in (3, 4, 5):
            if key == "합계" and "부채" in tr and "자본" in tr:
                f = f"={L(col)}{rowmap[id(tr['부채'])]}+{L(col)}{rowmap[id(tr['자본'])]}"
            else:
                kids = getattr(ln, "_dtop", None) or []
                f = "=SUM(" + ",".join(f"{L(col)}{rowmap[id(k)]}" for k in kids) + ")" if kids else None
            if f:
                cell = ws.cell(r, col, f)
                cell.number_format, cell.font = WON, Font(bold=True)
    return tr


def write_is_totals(ws_stmt, c, lines):
    """손익계산서 정산표의 계산 줄(매출총이익 · 영업이익 · 세전이익 · 당기순이익). 계정 줄의 분류 열을 SUMIF 한다."""
    out = {}
    accts = [ln for ln in lines if ln.kind == "acct"]
    if not accts:
        return out
    h = L(c.assign + 1)
    rng = f"${h}${min(x.row for x in accts)}:${h}${max(x.row for x in accts)}"
    for ln in lines:
        if ln.kind != "total":
            continue
        n = _norm(ln.label)
        for col in (c.mid, c.cur, c.final, c.p1, c.p2):
            vr = f"{L(col)}${min(x.row for x in accts)}:{L(col)}${max(x.row for x in accts)}"
            g = lambda k: f'SUMIF({rng},"{k}",{vr})'               # noqa: E731
            gp = f"({g('rev')}-{g('cogs')})"
            op = f"({gp}-{g('sga')})"
            pbt = f"({op}+{g('nonop_in')}-{g('nonop_out')})"
            if "매출총이익" in n or "매출총손실" in n:
                f, k = f"={gp}", "gp"
            elif "영업이익" in n or "영업손실" in n:
                f, k = f"={op}", "op"
            elif "차감전" in n:
                f, k = f"={pbt}", "pbt"
            elif n in ("당기순이익", "당기순손실", "당기순이익(손실)", "당기순손익"):
                f, k = f"={pbt}-{g('tax')}", "ni"
            else:
                continue
            cell = ws_stmt.cell(ln.row, col, f)
            cell.number_format, cell.font = WON, Font(bold=True)
            out[k] = ln
    return out


def write_equity_statement(wb, bs_seq, bs_rows, sec_heads_bs, pl_ni_ref):
    """공시용자본변동표. 전기말 · 당기말은 공시용BS 의 자본 소계, 변동 줄은 회계사가 적고, 당기순이익은 공시용PL 에서 온다."""
    ws = wb.create_sheet("공시용자본변동표")
    ws.sheet_view.showGridLines = False
    eq_total = next((ln for ln in bs_seq if ln.kind == "total" and _total_key(_norm(ln.label)) == "자본"), None)
    kids = getattr(eq_total, "_dtop", None) or []
    cols = [h for h in kids if h.kind == "head" and h.level == 1] or [h for h in kids if h.kind == "acct"]
    # 마감 전 시산표라 공시용BS 는 당기순손익을 따로 한 줄로 둔다(이익잉여금 줄에 계정이 있을 때). 자본변동표에서는 그 줄을
    # 열로 세우지 않고 이익잉여금 열에 더해 대사한다. 따로 세우면 변동이 없어도 대사가 어긋난다(2026-10-16 시연 검토)
    ni_h = next((h for h in kids if _norm(h.label) in ("당기순손익", "당기순이익", "당기순손실", "당기순이익(손실)")), None)
    if ni_h is not None and len(cols) > 1:
        cols = [h for h in cols if h is not ni_h]
    if not cols:
        ws.cell(1, 2, "공시용자본변동표").font = Font(bold=True, size=13)
        ws.cell(2, 2, "이 회사 공시 양식에 자본 과목이 없어 만들지 못했습니다").font = MUTED
        return {"eq_check_row": 5, "n": 0}
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for i in range(len(cols) + 2):
        ws.column_dimensions[L(3 + i)].width = 17
    ws.cell(1, 2, "공시용자본변동표").font = Font(bold=True, size=13)
    ws.cell(2, 2, "전기말 · 당기말은 공시용BS 자본 소계입니다. 노란 줄에 당기 변동(증자 · 배당 · 기타포괄손익 등)을 적으면 당기말이 공시용BS 와 맞는지 아래 대사 줄이 보입니다").font = MUTED
    heads = ["과목"] + [h.label for h in cols] + ["합계", "검증"]
    for i, t in enumerate(heads, 2):
        h = ws.cell(4, i, t)
        h.font, h.fill = Font(bold=True), HEAD
        h.alignment = Alignment(horizontal="center", wrap_text=True)
    n = len(cols)
    tc, vc = 3 + n, 4 + n
    r = 5
    ws.cell(r, 2, "전기말(= 당기초)").font = Font(bold=True)
    for i, h in enumerate(cols):
        ws.cell(r, 3 + i, f"='공시용BS'!E{bs_rows[id(h)]}").number_format = WON
    ws.cell(r, tc, f"=SUM(C{r}:{L(2 + n)}{r})").number_format = WON
    open_r = r
    r += 1
    ni_r = r
    ws.cell(r, 2, "당기순이익(공시용PL)")
    # 당기순이익은 이익잉여금(마지막 자본 소계)에 들어간다
    ni_col = 3 + max(range(n)) if n else 3
    for i, h in enumerate(cols):
        if re.search(r"이익잉여금|결손금", h.label):
            ni_col = 3 + i
    ws.cell(r, ni_col, f"={pl_ni_ref}").number_format = WON
    ws.cell(r, tc, f"=SUM(C{r}:{L(2 + n)}{r})").number_format = WON
    r += 1
    mv_first = r
    for lab in ("유상증자 · 무상증자", "배당금", "주식보상 · 주식선택권", "기타포괄손익", "기타 변동", "", ""):
        ws.cell(r, 2, lab or None).fill = YELLOW
        for i in range(n):
            ws.cell(r, 3 + i).fill = YELLOW
            ws.cell(r, 3 + i).number_format = WON
        ws.cell(r, tc, f"=SUM(C{r}:{L(2 + n)}{r})").number_format = WON
        r += 1
    mv_last = r - 1
    ws.cell(r, 2, "당기말").font = Font(bold=True)
    for i in range(n):
        ws.cell(r, 3 + i, f"=SUM({L(3 + i)}{open_r}:{L(3 + i)}{mv_last})").number_format = WON
    ws.cell(r, tc, f"=SUM(C{r}:{L(2 + n)}{r})").number_format = WON
    end_r = r
    r += 1
    ws.cell(r, 2, "공시용BS 당기말(Final)").font = MUTED
    for i, h in enumerate(cols):
        extra = f"+'공시용BS'!D{bs_rows[id(ni_h)]}" if ni_h is not None and ni_h not in cols and 3 + i == ni_col and id(ni_h) in bs_rows else ""
        ws.cell(r, 3 + i, f"='공시용BS'!D{bs_rows[id(h)]}{extra}").number_format = WON
    ws.cell(r, tc, f"=SUM(C{r}:{L(2 + n)}{r})").number_format = WON
    bs_r = r
    r += 1
    ws.cell(r, 2, "대사(당기말 − 공시용BS, 0 이어야 함)").font = Font(bold=True)
    for i in range(n + 1):
        ws.cell(r, 3 + i, f"={L(3 + i)}{end_r}-{L(3 + i)}{bs_r}").number_format = WON
    ws.cell(r, vc, f'=IF(SUMPRODUCT(ABS(C{r}:{L(2 + n)}{r}))<1,"OK","변동 줄 확인")')
    return {"eq_check_row": r, "n": n}


# ------------------------------------------------------------------ 본체
def build(d: Path, res: dict, out: Path) -> dict:
    """WTB 양식 파일을 쓴다. {"path", "notes", "bs_rows", "is_rows"}"""
    d = Path(d)
    out = Path(out)
    adj = adjustment_rows(d, res, out)
    res, adj_notes = with_adjust_accounts(res, adj["rows"])
    L_ = build_lines(d, res)
    L_["notes"] = adj_notes + L_["notes"]
    if not L_["BS"] and not L_["IS"]:
        raise ValueError("공시 양식(회사 제시 재무제표)이 없어 WTB 양식을 세울 수 없습니다: " + "; ".join(L_["notes"]))
    wb = Workbook()
    wb.remove(wb.active)
    accounts = L_["accounts"]
    ws_bs, ws_is = wb.create_sheet(BS), wb.create_sheet(IS)
    adj_rng, adj_chk = write_adjustments(wb, adj["rows"])
    write_accounts_table(wb, accounts)
    mat, assigns = _mat(res), _assigns(d)
    cbs, sec_bs, _ = write_statement(ws_bs, "BS", L_["BS"], adj_rng, mat, assigns)
    cis, sec_is, _ = write_statement(ws_is, "IS", L_["IS"], adj_rng, mat, assigns)
    is_tot = write_is_totals(ws_is, cis, L_["IS"])
    # 당기순손익 줄(자본)은 손익계산서 당기순이익 줄과 연결한다
    ni_line = is_tot.get("ni")
    for ln in L_["BS"]:
        if ln.link == "NI":
            r = ln.row
            ws_bs.cell(r, cbs.name, ln.name)
            ws_bs.cell(r, cbs.key, ln.key).font = Font(color="1F4E79")
            if ni_line is not None:
                ws_bs.cell(r, cbs.cur, f"={Q(IS)}!{L(cis.cur)}{ni_line.row}").number_format = WON
                ws_bs.cell(r, cbs.final, f"={Q(IS)}!{L(cis.final)}{ni_line.row}").number_format = WON
                ws_bs.cell(r, cbs.mid, f"={Q(IS)}!{L(cis.mid)}{ni_line.row}").number_format = WON
    total_bs = finish_bs(ws_bs, cbs, L_["BS"], sec_bs, None)
    # 검산 줄
    if "자산" in total_bs and "합계" in total_bs:
        a, g = total_bs["자산"].row, total_bs["합계"].row
        chk = max(ln.row for ln in L_["BS"]) + 2
        ws_bs.cell(chk, cbs.lab[0], "Check").font = Font(bold=True)
        for col in (cbs.cur, cbs.final, cbs.p1):
            ws_bs.cell(chk, col, f'=IF(ROUND({L(col)}{a}-{L(col)}{g},0)=0,"OK","Imbalanced")')
    add_diff_columns(ws_bs, cbs, L_["BS"], cbs.assign + 2)
    add_diff_columns(ws_is, cis, L_["IS"], cis.assign + 2)
    ws_bs.freeze_panes = ws_bs.cell(7, cbs.key + 1)
    ws_is.freeze_panes = ws_is.cell(7, cis.key + 1)
    # 공시용 시트
    wbs, seq_bs, rows_bs, secs_bs = write_disclosure(wb, "공시용BS", L_["BS"], BS, cbs, "정산표(재무상태표)를 공시용 계정과목으로 SUMIF 한 공시 양식 재무상태표")
    tr_bs = _total_formulas_bs(wbs, seq_bs, rows_bs, secs_bs)
    wpl, seq_is, rows_is, secs_is = write_disclosure(wb, "공시용PL", L_["IS"], IS, cis, "정산표(손익계산서)를 공시용 계정과목으로 SUMIF 한 공시 양식 손익계산서")
    # 공시용PL 계산 줄은 정산표 손익계산서의 같은 합계 줄을 그대로 참조한다
    pl_ni = None
    stmt_total = {id(x): x for x in L_["IS"] if x.kind == "total"}
    for ln in seq_is:
        if ln.kind != "total":
            continue
        r = rows_is[id(ln)]
        n = _norm(ln.label)
        if not any(k in n for k in ("매출총", "영업이익", "영업손실", "차감전", "당기순")):
            continue
        for col, sc in ((3, cis.cur), (4, cis.final), (5, cis.p1)):
            cell = wpl.cell(r, col, f"={Q(IS)}!{L(sc)}{ln.row}")
            cell.number_format, cell.font = WON, Font(bold=True)
        if n in ("당기순이익", "당기순손실", "당기순이익(손실)", "당기순손익"):
            pl_ni = f"'공시용PL'!D{r}"
    eq = write_equity_statement(wb, seq_bs, rows_bs, secs_bs, pl_ni or "0")
    # 검산 시트: 엑셀이 계산한 값 ↔ 엔진 기대값(시산표 · 분석 결과의 합계)
    tot = res["totals"]
    checks = []
    if "자산" in total_bs:
        checks += [("자산총계(회사 제시)", f"={Q(BS)}!{L(cbs.cur)}{total_bs['자산'].row}", tot["assets"]),
                   ("부채총계(회사 제시)", f"={Q(BS)}!{L(cbs.cur)}{total_bs['부채'].row}", tot["liabilities"]) if "부채" in total_bs else None,
                   ("자본총계(회사 제시)", f"={Q(BS)}!{L(cbs.cur)}{total_bs['자본'].row}", tot["equity"]) if "자본" in total_bs else None]
    if ni_line is not None:
        checks.append(("당기순이익(회사 제시)", f"={Q(IS)}!{L(cis.cur)}{ni_line.row}", tot["net_income"] if tot.get("net_income") is not None else None))
    checks = [x for x in checks if x]
    ws_ck = wb.create_sheet("_검산")
    ws_ck.append(["검산", "수식(엑셀 계산값)", "엔진 기대값"])
    for label, f, exp in checks:
        ws_ck.append([label, f, exp])
    ws_ck.append(["수정사항 차변 − 대변", f"={adj_chk['dr_cr']}", 0])
    ws_ck.append(["수정사항 자산 − 부채 − 손익 − 자본 영향", f"={adj_chk['eff']}", 0])
    ws_ck.append(["자본변동표 당기말 대사(합계 열)", f"='공시용자본변동표'!{L(3 + eq['n'])}{eq['eq_check_row']}", None])
    ws_ck.append(["8100A 수정분개 차변 합(재무제표 · DART 와 같은 분개)", f"={adj_chk['a_dr']}", sum(x[2] for x in adj["adjust"])])
    ws_ck.sheet_state = "hidden"
    order = [BS, IS, ADJ, "공시용BS", "공시용PL", "공시용자본변동표"]
    wb._sheets = [wb[n] for n in order] + [s for s in wb._sheets if s.title not in order]
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    _record(d, out, adj)
    return {"path": out, "notes": L_["notes"] + adj["notes"], "bs_rows": len(L_["BS"]), "is_rows": len(L_["IS"]),
            "adjust": adj["adjust"], "manual": adj["manual"]}


def default_path(d: Path) -> Path:
    """감사 폴더의 WTB 양식 기본 자리: 8000 감사완결/WTB_<회사>_<결산일>.xlsx (python -m audit wtbform 과 같다)."""
    import json

    d = Path(d)
    case = json.loads((d / "_색인" / "감사건.json").read_text(encoding="utf-8"))
    co = case.get("company")
    co = (co.get("name") if isinstance(co, dict) else co) or d.parent.name
    return d / "8000 감사완결" / f"WTB_{co}_{case.get('fiscal_year_end', '')}.xlsx"


def record(d: Path) -> dict | None:
    """감사 폴더 안에 만든 WTB 양식의 기록(_색인/WTB양식.json). 파일이 없어졌으면 None."""
    import json

    p = Path(d) / "_색인" / RECORD
    if not p.exists():
        return None
    rec = json.loads(p.read_text(encoding="utf-8"))
    return rec if (Path(d) / rec.get("path", "")).is_file() else None


def ensure(d: Path) -> dict:
    """4000 조서보다 WTB 가 먼저다(2026-10-16). WTB 양식이 없으면 분석 결과로 기본 자리에 만든다.
    {"path": 상대 경로 | None, "made": 새로 만들었는지, "why": 만들지 못한 이유}. 분석 결과가 없거나 공시 양식이 없으면 만들지 않는다."""
    import json

    d = Path(d)
    rec = record(d)
    if rec is not None:
        return {"path": rec["path"], "made": False, "why": None}
    p = d / "_색인" / "분석.json"
    if not p.exists():
        return {"path": None, "made": False, "why": "분석 결과가 없어 WTB 를 만들지 못했습니다(python -m audit analyze 먼저)"}
    out = default_path(d)
    try:
        build(d, json.loads(p.read_text(encoding="utf-8")), out)
    except ValueError as e:
        return {"path": None, "made": False, "why": str(e)}
    return {"path": out.relative_to(d).as_posix(), "made": True, "why": None}


def _record(d: Path, out: Path, adj: dict) -> None:
    """감사 폴더 안에 만든 WTB 면 무엇을 담았는지 _색인/WTB양식.json 에 남긴다. 수정 반영이 지금 8100A 와 견준다."""
    idx = Path(d) / "_색인"
    try:
        rel = Path(out).resolve().relative_to(Path(d).resolve()).as_posix()
    except ValueError:
        return                                     # 감사 폴더 밖에 낸 파일은 수정 반영이 다시 만들지 않는다
    if not idx.is_dir():
        return
    from . import folder

    folder._write(idx / RECORD, {"format_version": 1, "path": rel, "at": folder._now(), "adjust": adj["adjust"],
                                 "wtb": adj["wtb"], "manual": adj["manual"]})


# ------------------------------------------------------------------ 엑셀 재계산 검산 (2026-10-16)
def recalc_check(path: Path) -> dict:
    """WTB 파일의 사본을 임시 폴더에서 엑셀로 다시 계산해 `_검산` 시트(엑셀 값 ↔ 엔진 기대값)와 재무상태표 Check,
    공시용자본변동표 대사를 읽는다. 원본은 바꾸지 않는다. 엑셀이 없는 PC 면 {"ok": None, "why": ...}.
    {"ok", "checks": [(항목, 엑셀, 기대, 맞음)], "balance": [OK · Imbalanced], "equity": "OK" | "변동 줄 확인", "errors": [셀]}"""
    import shutil
    import sys
    import tempfile

    from openpyxl import load_workbook

    tools = Path(__file__).resolve().parents[1] / "tools"
    if str(tools) not in sys.path:
        sys.path.insert(0, str(tools))
    try:
        from excel_recalc import recalc
    except ImportError:
        return {"ok": None, "why": "엑셀 재계산 도구(tools/excel_recalc.py)를 찾지 못했습니다"}
    with tempfile.TemporaryDirectory(prefix="audit-wtb-") as tmp:
        cp = Path(tmp) / "WTB_검산.xlsx"
        shutil.copy2(path, cp)
        if not recalc(cp):
            return {"ok": None, "why": "이 PC 에 엑셀(또는 pywin32)이 없어 수식 값을 계산하지 못했습니다"}
        wb = load_workbook(cp, data_only=True)
        checks = []
        for label, got, exp in wb["_검산"].iter_rows(min_row=2, max_col=3, values_only=True):
            if not label:
                continue
            ok = exp is None or (isinstance(got, (int, float)) and abs(got - exp) < 1)
            checks.append((label, got, exp, ok))
        balance = [c.value for row in wb[BS].iter_rows() for c in row if c.value in ("OK", "Imbalanced")]
        eq = None
        if "공시용자본변동표" in wb.sheetnames:
            eq = next((c.value for row in wb["공시용자본변동표"].iter_rows() for c in row if c.value in ("OK", "변동 줄 확인")), None)
        errors = [f"{ws.title}!{c.coordinate}" for ws in wb for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper()]
        wb.close()
    ok = all(x[3] for x in checks) and all(b == "OK" for b in balance) and eq in (None, "OK") and not errors
    return {"ok": ok, "checks": checks, "balance": balance, "equity": eq, "errors": errors[:20]}

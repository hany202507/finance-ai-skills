"""전기조서 내용 읽기 (2026-10-12).

전기조서의 글(계정의 주요 성격 · 메모 · 절차 목록 · 발견사항 · 결론 · 수정사항)을 읽어 `_색인/전기조서.json` 에 두고,
당기 조서가 참고로 싣게 한다. 숫자는 여기서 쓰지 않는다. 전기 잔액은 audit/prior.py 의 원천 순위가 그대로 정한다.

원천은 둘이다.
  손으로 만든 전기조서    받은자료 `전기` 폴더의 엑셀. 라벨(「Audit memo」 · 「감사절차」 · 「발견사항」 · 「Conclusion」 …)로 글 블록을 찾는다.
                         라벨은 definitions/prior_wp.yaml 에 있다. 라벨을 하나도 못 찾은 조서는 AI 판독(확인 전)으로 넘긴다
  작년 감사 폴더의 엔진 조서   `_기입란` 되읽기로 회계사가 적은 값을 그대로 읽는다(증감 사유 · 수행 결과 · 해당 여부 · 결론)

  read_manual(path)            손 조서 한 파일 → {"sheets": [{sheet, role, blocks, procedures, adjustments}], "hidden": n}
  area_of(name)                파일 이름 → (4000 영역 번호 또는 None, 근거)
  looks_like_workpaper(path)   전기 폴더의 엑셀이 전기조서인지(색인 판별)
  collect(d, now)              `_색인/전기조서.json` 을 쓴다. 회계사가 지정한 영역과 AI 판독은 이어받는다
  apply_patch(d, patch)        AI · 사람 판독을 넣는다(확인 전 표시)
  assign(d, path, no)          엔진이 영역을 정하지 못한 파일에 회계사가 영역을 지정한다
  load(d) · for_area(d, no)    당기 조서가 읽는다
  general_inputs(d, no)        작년 감사 폴더의 같은 번호 조서에서 회계사가 적은 값(일반조서 참고용)

숨김 시트(전전기 자료)와 외부 연결은 읽지 않는다. 주민번호 · 계좌번호 꼴의 글자는 옮기지 않는다.
전기에 빠졌던 절차를 한공회 절차와 견줘 짚는 일은 하지 않는다(2026-10-12 대표 결정 6).
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from . import folder, reader
from .errors import DefinitionError

INDEX = "전기조서.json"
PRIOR_WP = "전기조서"
PRIOR_FOLDER = "전기"
MANUAL, ENGINE = "손 조서", "작년 감사 폴더"
KINDS = ("nature", "memo", "purpose", "procedure", "result", "findings", "conclusion", "section")
SKIP_INPUTS = ("검토자", "작성자", "검토 일자")
MASK = "(번호 가림)"
TAIL_COLS = 6
COVER_MARK = "실증절차"
_STRIP = re.compile(r"^[\s\d.()\[\]<>【】「」\-]+|[\s\]>)】」:：]+$")
_HEADING = re.compile(r"^\s*(?:\d+\.\s*\S|\[.+\]\s*$|<.+>\s*$|【.+】\s*$)")
_RRN = re.compile(r"(?<!\d)\d{6}\s*-?\s*[1-8]\d{6}(?!\d)")
_ACCT = re.compile(r"(?<!\d)\d{2,6}(?:-\d{2,8}){2,4}(?!\d)")
_LONG = re.compile(r"(?<![\d,.])\d{10,16}(?![\d,.])")


class PatchRejected(Exception):
    def __init__(self, problems: list[str]):
        super().__init__("; ".join(problems))
        self.problems = problems


def defs() -> dict:
    p = folder.DEFS / "prior_wp.yaml"
    if not p.exists():
        raise DefinitionError(f"정의 파일이 없습니다: {p}")
    return folder.read_yaml(p)


def _label_map(cfg: dict) -> dict[str, str]:
    return {w: kind for kind, words in cfg["labels"].items() for w in words}


def norm_label(v) -> str:
    """라벨로 견줄 글: 앞 번호 · 괄호 · 빈칸을 떼고 소문자로."""
    return re.sub(r"\s+", "", _STRIP.sub("", str(v or ""))).lower()


def mask(text: str) -> str:
    """주민번호 · 계좌번호 꼴(하이픈으로 이은 숫자 묶음, 쉼표 없는 10~16자리 숫자)을 가린다. 사업자번호 · 전화번호도 함께 가려진다."""
    return _LONG.sub(MASK, _ACCT.sub(MASK, _RRN.sub(MASK, text)))


def _text(v, raw: bool = False) -> str | None:
    """칸 값을 글로. 오류 값(#REF! 등)과 빈 칸은 None. raw 면 숫자를 쉼표 없이 둔다."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v != v:
            return None
        if raw:
            return str(v)
        if float(v).is_integer():
            return f"{int(v):,}"
        return f"{v:,.4f}".rstrip("0").rstrip(".")
    s = re.sub(r"\s+", " ", str(v)).strip()
    if not s or s.startswith("#"):
        return None
    return mask(s)


def _line(cells: list, limit: int) -> str | None:
    parts = [t for t in (_text(v) for v in cells) if t]
    if not parts or not any(isinstance(v, str) and _text(v) for v in cells):
        return None                      # 글이 없는 줄(숫자만 있는 표 줄)은 싣지 않는다
    s = " | ".join(parts)
    return s if len(s) <= limit else s[:limit - 1] + "…"


def _col(n: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(n)


# ---------------------------------------------------------------- 손 조서
def _find_label(row: tuple, cfg: dict, lmap: dict) -> tuple[int, str, str, str | None] | None:
    """행에서 라벨을 찾는다. (열 번호, 종류, 라벨 글, 같은 칸의 덧붙은 글). 「1. Audit Procedure : 주석사항을 검토함」 처럼
    라벨 뒤에 쌍점과 글이 이어진 칸도 받는다."""
    sc = cfg["scan"]
    for j, v in enumerate(row[:sc["label_cols"]], 1):
        if not isinstance(v, str) or not v.strip():
            continue
        s = v.strip()
        if len(s) <= sc["label_len"] and norm_label(s) in lmap:
            return j, lmap[norm_label(s)], s, None
        m = re.match(r"^(.{1,30}?)\s*[:：]\s*(\S.*)$", s)
        if m and norm_label(m.group(1)) in lmap:
            return j, lmap[norm_label(m.group(1))], m.group(1).strip(), m.group(2).strip()
        if len(s) <= sc["label_len"] and _HEADING.match(s):
            return j, "section", s, None
        return None                      # 라벨 열의 첫 글이 라벨이 아니면 그 줄은 내용이다
    return None


def _adjust_entry(cells: list) -> dict | None:
    """수정분개 줄: 글 · 금액 · 글 · 금액(차변 계정 · 금액 · 대변 계정 · 금액). 뒤의 글은 참조로 둔다."""
    vals = [v for v in cells if v is not None and not (isinstance(v, str) and not v.strip())]
    if len(vals) < 4:
        return None
    a, x, b, y = vals[:4]
    num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)  # noqa: E731
    if isinstance(a, str) and num(x) and isinstance(b, str) and num(y) and x and y:
        note = " ".join(t for t in (_text(v) for v in vals[4:]) if t)
        return {"dr": mask(a.strip()), "dr_amount": reader._int(float(x)), "cr": mask(b.strip()), "cr_amount": reader._int(float(y)),
                "note": note or None}
    return None


def _read_sheet(ws, cfg: dict, lmap: dict) -> dict:
    sc = cfg["scan"]
    blocks, procedures, adjustments = [], [], []
    cur = None
    for i, row in enumerate(ws.iter_rows(max_row=sc["max_rows"], max_col=sc["max_cols"], values_only=True), 1):
        row = tuple(row)
        if i <= 3 and any(isinstance(v, str) and re.sub(r"\s+", "", v) == COVER_MARK for v in row):
            return {"blocks": [], "procedures": [], "adjustments": []}      # 한공회 표지(실증절차 목록 서식)는 수행한 절차가 아니다
        hit = _find_label(row, cfg, lmap)
        if hit:
            j, kind, label, tail = hit
            cur = {"kind": kind, "title": mask(label), "cell": f"{_col(j)}{i}", "lines": [], "truncated": 0, "_rows": []}
            blocks.append(cur)
            rest = list(row[j:j + TAIL_COLS])          # 같은 줄 오른쪽 끝의 조서 머리(Index · Prepared · 이름)는 읽지 않는다
            if tail:
                rest = [tail] + rest
            if any(v is not None for v in rest):
                cur["_rows"].append(rest)
            continue
        if cur is not None:
            cur["_rows"].append(list(row) if any(v is not None for v in row) else None)      # None 은 빈 줄(표의 끝을 가린다)
    out_blocks = []
    for b in blocks:
        rows = b.pop("_rows")
        if b["kind"] == "lead":
            continue                                   # Lead 표의 숫자는 읽지 않는다
        if b["kind"] == "procedures":
            entries = _procedure_rows(rows)
            if any(e["ref"] for e in entries):
                procedures += [dict(e, cell=b["cell"]) for e in entries]
                continue
            b["kind"] = "procedure"                    # 참조 칸이 없으면 절차 글이다
        elif b["kind"] == "procedure":
            entries = _procedure_rows(rows)
            if len(entries) >= 2 and all(e["ref"] for e in entries):
                procedures += [dict(e, cell=b["cell"]) for e in entries]      # 「[절차]」 아래가 이름 · 참조 표인 Lead
                continue
        rows = [r for r in rows if r is not None]
        if b["kind"] == "findings":
            text_rows = []
            for r in rows:
                e = _adjust_entry(r)
                if e:
                    adjustments.append(dict(e, cell=b["cell"]))      # 분개 줄은 수정사항 목록에만 싣는다
                else:
                    text_rows.append(r)
            rows = text_rows
        for r in rows:
            t = _line(r, sc["line_chars"])
            if not t:
                continue
            if len(b["lines"]) < sc["block_lines"]:
                b["lines"].append(t)
            else:
                b["truncated"] += 1
        if b["lines"] or b["kind"] != "section":
            out_blocks.append(b)
    return {"blocks": out_blocks, "procedures": procedures, "adjustments": adjustments}


def _procedure_rows(rows: list) -> list[dict]:
    """「Audit procedure」 아래 표. 머리줄(Procedure · Reference)은 빼고, 줄마다 첫 글이 절차 이름, 마지막 글이 참조다.
    표는 첫 빈 줄에서 끝난다. 절차 시트의 「Audit procedures」 처럼 아래가 시험 표인 블록을 절차 목록으로 읽지 않게 한다."""
    out = []
    for r in rows:
        if r is None:
            if out:
                break
            continue
        texts = [t for t in (_text(v) for v in r if isinstance(v, str)) if t]
        if not texts:
            continue
        if norm_label(texts[0]) in ("procedure", "procedures", "절차", "감사절차"):
            continue
        ref = texts[-1] if len(texts) > 1 and len(texts[-1]) <= 40 else None
        out.append({"name": texts[0][:120], "ref": ref})
    return out


def read_manual(path: Path, cfg: dict | None = None) -> dict:
    """손으로 만든 전기조서 한 파일. 보이는 시트만 읽는다. 이름 정의가 깨진 파일도 열린다(reader.open_plain)."""
    cfg = cfg or defs()
    lmap = _label_map(cfg)
    wb = reader.open_plain(Path(path))
    try:
        sheets, hidden = [], 0
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                hidden += 1
                continue
            got = _read_sheet(ws, cfg, lmap)
            if got["blocks"] or got["procedures"] or got["adjustments"]:
                role = "lead" if "lead" in ws.title.lower() or len(got["procedures"]) > 1 else "procedure"
                sheets.append({"sheet": ws.title, "role": role, **got})
        return {"sheets": sheets, "hidden": hidden, "visible": len(wb.worksheets) - hidden}
    finally:
        wb.close()


def looks_like_workpaper(path: Path) -> bool:
    """전기 폴더의 엑셀이 전기조서인지: 글 블록 라벨이 정한 수 이상 있다."""
    if Path(path).suffix.lower() not in reader.EXCEL:
        return False
    try:
        cfg = defs()
        got = read_manual(path, cfg)
    except Exception:  # noqa: BLE001  열지 못하는 파일은 전기조서로 보지 않는다
        return False
    n = sum(1 for s in got["sheets"] for b in s["blocks"] if b["kind"] != "section") + sum(len(s["procedures"]) > 0 for s in got["sheets"])
    return n >= cfg["scan"]["min_blocks"]


# ---------------------------------------------------------------- 영역
def area_of(name: str, cfg: dict | None = None) -> tuple[str | None, str]:
    """전기조서 파일 이름 → 4000 영역 번호와 근거. 한공회 번호(4000G), 알파벳 번호(A · AA · BBDD · C · D2 · E3 · R1 · CL …),
    조서 이름(「기타유동부채」) 순으로 본다. 정하지 못하면 (None, 까닭)."""
    cfg = cfg or defs()
    std = folder.load_standard()
    names = {str(k): v for k, v in (std.get("workpaper_names") or {}).items()}
    areas = [str(n)[4:] for f in std["folders"] for n in f.get("workpapers", []) if folder._is_area(str(n))]
    stem = Path(name).stem.strip()
    m = re.match(rf"({folder.AREA_NO})(?=$|[_ .\-])", stem)
    if m:
        return m.group(1), f"파일 이름의 한공회 번호 {m.group(1)}"
    m = re.match(r"([A-Z]+)(\d*)(?=$|[_ .\-(])", stem)
    if m:
        letters, digits = m.group(1), m.group(2)
        code = letters + digits
        over = {str(k): str(v) for k, v in (cfg.get("area_codes") or {}).items()}
        if code in over:
            return "4000" + over[code], f"파일 이름 번호 {code} (prior_wp.yaml area_codes)"
        if not digits and letters in areas:
            return "4000" + letters, f"파일 이름 번호 {letters}"
    body = re.sub(r"\s+", "", stem)
    hits = [(len(v), no) for no, v in names.items() if folder._is_area(no) and v and re.sub(r"\s+", "", v) in body]
    if hits:
        no = max(hits)[1]
        return no, f"조서 이름 「{names[no]}」"
    if m and m.group(1) in areas:
        return "4000" + m.group(1), f"파일 이름 번호 {m.group(1)}{m.group(2)} 의 앞 글자"
    return None, "파일 이름에서 영역을 정하지 못했습니다. 회계사가 지정해 주십시오"


# ---------------------------------------------------------------- 작년 감사 폴더의 엔진 조서
def _engine_inputs(prev: Path, w: dict) -> list[dict]:
    from . import wp

    f = w.get("file")
    if not f or not f.lower().endswith(".xlsx") or not (prev / f).exists():
        return []
    try:
        rows = wp.readback(prev / f)
    except Exception:  # noqa: BLE001  회계사가 다른 꼴로 저장한 파일
        return []
    out = []
    for x in rows:
        lab = str(x["label"] or "")
        if wp._blank(x["value"]) or lab in SKIP_INPUTS or lab.startswith(("수행자", "추가 절차 수행자", "조서번호", "추가 절차 조서번호")):
            continue
        t = _text(x["value"])
        if t:
            out.append({"sheet": x["sheet"], "label": lab, "value": t, "cell": x["ref"]})
    return out


def _engine_area(inputs: list[dict], cfg: dict) -> dict:
    """작년 엔진 4000 조서의 기입값을 블록 · 절차 목록으로 나눈다."""
    blocks: dict[str, dict] = {}
    procs: dict[str, dict] = {}

    def add(kind, title, x, text):
        b = blocks.setdefault(kind, {"kind": kind, "title": title, "sheet": x["sheet"], "cell": x["cell"], "lines": [], "truncated": 0})
        if len(b["lines"]) < cfg["scan"]["block_lines"]:
            b["lines"].append(text)
        else:
            b["truncated"] += 1

    names = cfg["kind_names"]
    for x in inputs:
        lab, v = x["label"], x["value"]
        m = re.match(r"^(?:추가 절차 )?(해당 여부|수행 결과|결과): (.+)$", lab) if x["sheet"] == "표지" else None
        if m:
            p = procs.setdefault(m.group(2), {"name": m.group(2), "ref": None, "applies": None, "result": None, "cell": x["cell"], "sheet": x["sheet"]})
            p["applies" if m.group(1) == "해당 여부" else "result"] = v
        elif lab == "감사목적":
            add("purpose", names["purpose"], x, v)
        elif lab.startswith("증감 사유"):
            add("memo", "증감 사유", x, f"{lab.split(': ', 1)[-1]}: {v}")
        elif lab == "감사절차 및 결과" or lab.startswith(("절차 결과", "결과")):
            add("result", names["result"], x, v if lab == "감사절차 및 결과" else f"{lab}: {v}")
        elif "발견사항" in lab:
            add("findings", names["findings"], x, f"{lab}: {v}")
        elif "결론" in lab:
            add("conclusion", names["conclusion"], x, v if x["sheet"] == "Lead" else f"{x['sheet']}: {v}")
    done = [p for p in procs.values() if p["applies"] or p["result"]]
    return {"blocks": list(blocks.values()), "procedures": done}


# ---------------------------------------------------------------- 모으기
def _idx(d: Path, name: str, default=None):
    return folder._read(Path(d) / "_색인" / name, default)


def _files(d: Path) -> list[str]:
    docs = _idx(d, "문서목록.json", {"files": []}) or {"files": []}
    return sorted(f["path"] for f in docs["files"] if f.get("role") == "received" and f.get("kind") == PRIOR_WP
                  and Path(f["path"]).suffix.lower() in reader.EXCEL)


def text_workpapers(d: Path, paths: list[str]) -> set[str]:
    """숫자 표 판독(start.read_prior_workpapers)에서 뺄 계정 조서: 회계사가 영역을 지정했거나, 파일 이름이 4000 영역으로 이어지고
    글 블록 라벨(Audit memo · 감사절차 · Conclusion …)이 있는 손 조서. 이런 조서의 Lead 표와 표본 표는 정산표 · 수정사항 머리줄과 겹친다.
    라벨이 없는 법인 리드 꼴(계정과목 · Proposed · 수정사항 · Audited 열, R8)은 여기에 들지 않아 앞에서처럼 정산표로 읽힌다."""
    d = Path(d)
    cfg = defs()
    data = load(d) or {}
    assigned = data.get("assigned") or {}
    known = {f["path"]: f for f in data.get("files", []) if f.get("source") == MANUAL and "labels" in f}
    out = set()
    for rel in paths:
        if rel in assigned:
            out.add(rel)
        elif not area_of(Path(rel).name, cfg)[0]:
            continue
        elif rel in known:
            if known[rel]["labels"] >= cfg["scan"]["min_blocks"]:
                out.add(rel)
        elif looks_like_workpaper(d / rel):
            out.add(rel)
    return out


def _src(path: str, sheet: str, cell: str) -> dict:
    return {"path": path, "sheet": sheet, "cell": cell}


def collect(d: Path, now: str | None = None) -> dict:
    """전기조서 내용을 읽어 `_색인/전기조서.json` 에 쓴다. 회계사가 지정한 영역(assigned)과 AI · 사람 판독(basis AI)은 이어받는다."""
    from . import prior

    d = Path(d)
    cfg = defs()
    at = now or folder._now()
    old = _idx(d, INDEX, {}) or {}
    assigned = dict(old.get("assigned") or {})
    out = {"format_version": 1, "read_at": at, "files": [], "areas": {}, "general": {}, "assigned": assigned, "need_ai": [], "warnings": []}

    def area(no: str) -> dict:
        return out["areas"].setdefault(no, {"files": [], "blocks": [], "procedures": [], "adjustments": []})

    trivial = set(cfg.get("trivial") or [])
    for rel in _files(d):
        rec = {"path": rel, "source": MANUAL, "area": None, "area_basis": "", "sheets": 0, "hidden_skipped": 0, "blocks": 0, "need_ai": False}
        out["files"].append(rec)
        try:
            got = read_manual(d / rel, cfg)
        except Exception as e:  # noqa: BLE001  열지 못한 파일은 남기고 다음 파일로 간다
            rec["error"] = f"{type(e).__name__}: {str(e)[:120]}"
            rec["need_ai"] = True
            out["need_ai"].append(rel)
            continue
        no, basis = (assigned[rel], "회계사 지정") if rel in assigned else area_of(Path(rel).name, cfg)
        n = sum(len(s["blocks"]) + len(s["procedures"]) + len(s["adjustments"]) for s in got["sheets"])
        known = sum(1 for s in got["sheets"] for b in s["blocks"] if b["kind"] != "section") + sum(len(s["procedures"]) for s in got["sheets"])
        rec.update(area=no, area_basis=basis, sheets=got["visible"], hidden_skipped=got["hidden"], blocks=n, labels=known)
        if no is None:
            if known >= cfg["scan"]["min_blocks"]:
                out["warnings"].append(f"전기조서의 영역을 정하지 못했습니다. python -m audit prior area 로 지정해 주십시오: {rel}")
            else:
                # 중요성 · 정산표 · 미수정 왜곡 같은 표 조서다. 숫자 표는 시작 자료(start.read_prior_workpapers)가 읽는다
                rec["area_basis"] = "계정 조서가 아님(표는 시작 자료가 읽는다)"
            continue
        if known == 0:
            rec["need_ai"] = True                      # 계정 조서인데 글 블록 라벨을 하나도 찾지 못했다
            out["need_ai"].append(rel)
        a = area(no)
        a["files"].append(rel)
        for s in got["sheets"]:
            for b in s["blocks"]:
                body = re.sub(r"[\s.]", "", " ".join(b["lines"])).lower()
                if s["role"] != "lead" and b["kind"] in ("findings", "conclusion") and (not body or body in trivial):
                    continue                           # 절차 시트의 「특이사항 없음」 은 싣지 않는다
                a["blocks"].append({"kind": b["kind"], "title": b["title"], "lines": b["lines"], "truncated": b["truncated"],
                                    "role": s["role"], "source": _src(rel, s["sheet"], b["cell"]), "basis": "엔진", "confirmed": True})
            a["procedures"] += [{"name": p["name"], "ref": p["ref"], "applies": None, "result": None,
                                 "source": _src(rel, s["sheet"], p["cell"]), "basis": "엔진", "confirmed": True} for p in s["procedures"]]
            a["adjustments"] += [{k: e[k] for k in ("dr", "dr_amount", "cr", "cr_amount", "note")} | {"source": _src(rel, s["sheet"], e["cell"])}
                                 for e in s["adjustments"]]

    prev = prior.prior_engagement(d)
    if prev is not None:
        for w in (folder._read(prev / "_색인" / "조서목록.json", {"workpapers": []}) or {"workpapers": []})["workpapers"]:
            inputs = _engine_inputs(prev, w)
            if not inputs:
                continue
            rel = prior._rel(d, prev / w["file"])
            rec = {"path": rel, "source": ENGINE, "area": w["no"] if folder._is_area(w["no"]) else None,
                   "area_basis": "작년 조서번호", "sheets": len({x["sheet"] for x in inputs}), "hidden_skipped": 0, "blocks": len(inputs),
                   "need_ai": False}
            out["files"].append(rec)
            if folder._is_area(w["no"]):
                a = area(w["no"])
                a["files"].append(rel)
                got = _engine_area(inputs, cfg)
                a["blocks"] += [{"kind": b["kind"], "title": b["title"], "lines": b["lines"], "truncated": b["truncated"], "role": "lead",
                                 "source": _src(rel, b["sheet"], b["cell"]), "basis": "엔진", "confirmed": True} for b in got["blocks"]]
                a["procedures"] += [{"name": p["name"], "ref": p["ref"], "applies": p["applies"], "result": p["result"],
                                     "source": _src(rel, p["sheet"], p["cell"]), "basis": "엔진", "confirmed": True} for p in got["procedures"]]
            else:
                out["general"][w["no"]] = {"source": rel, "inputs": [{k: x[k] for k in ("sheet", "label", "value", "cell")} for x in inputs]}

    # AI · 사람 판독은 다시 읽어도 남긴다
    for no, a in (old.get("areas") or {}).items():
        keep_b = [b for b in a.get("blocks", []) if b.get("basis") == "AI"]
        keep_p = [p for p in a.get("procedures", []) if p.get("basis") == "AI"]
        if keep_b or keep_p:
            cur = area(no)
            cur["blocks"] += keep_b
            cur["procedures"] += keep_p
            for b in keep_b + keep_p:
                p = b["source"]["path"]
                if p not in cur["files"]:
                    cur["files"].append(p)
    have_ai = {b["source"]["path"] for a in out["areas"].values() for b in a["blocks"] + a["procedures"] if b.get("basis") == "AI"}
    out["need_ai"] = [p for p in out["need_ai"] if p not in have_ai]
    for f in out["files"]:
        if f["path"] in have_ai:
            f["need_ai"] = False
    folder._write(d / "_색인" / INDEX, out)
    case = _idx(d, "감사건.json")
    if case is not None:
        n_blocks = sum(len(a["blocks"]) + len(a["procedures"]) for a in out["areas"].values())
        what = (f"전기조서 {len(out['files'])}개에서 영역 {len(out['areas'])}개의 글 블록 · 절차 {n_blocks}개를 읽었습니다"
                + (f". AI 판독이 필요한 파일 {len(out['need_ai'])}개" if out["need_ai"] else "")
                + (f". 영역을 정하지 못한 파일 {len(out['warnings'])}개" if out["warnings"] else ""))
        case.setdefault("log", []).insert(0, {"at": at, "who": "전기조서", "what": what,
                                              "read": [f["path"] for f in out["files"]][:50], "wrote": [f"_색인/{INDEX}"]})
        folder._write(d / "_색인" / "감사건.json", case)
    folder.validate_index(d)
    return out


def assign(d: Path, path: str, no: str, now: str | None = None) -> dict:
    """회계사가 전기조서 파일의 영역을 지정한다. 지정은 전기조서.json assigned 에 남고 다시 읽을 때 쓴다."""
    d = Path(d)
    if not folder._is_area(no):
        raise ValueError(f"4000 영역 번호가 아닙니다: {no}")
    files = _files(d)
    hit = [f for f in files if f == path or Path(f).name == path]
    if len(hit) != 1:
        raise ValueError(f"전기조서 파일을 찾지 못했습니다: {path} (있는 파일: {', '.join(Path(f).name for f in files) or '없음'})")
    data = _idx(d, INDEX, {}) or {"format_version": 1, "read_at": now or folder._now(), "files": [], "areas": {}, "general": {},
                                  "assigned": {}, "need_ai": [], "warnings": []}
    data.setdefault("assigned", {})[hit[0]] = no
    folder._write(d / "_색인" / INDEX, data)
    return collect(d, now=now)


def apply_patch(d: Path, patch: dict, now: str | None = None) -> dict:
    """AI · 사람 판독을 넣는다. 꼴: {"files": [{"path", "area": "4000C", "basis": "AI"|"사람", "blocks": [{kind, title, lines, sheet, cell}],
    "procedures": [{name, ref}]}]}. basis 가 사람이면 그 파일의 AI 판독을 확인한 것으로 바꾼다. 하나라도 어긋나면 아무것도 쓰지 않는다."""
    d = Path(d)
    cfg = defs()
    known = set(_files(d))
    problems = []
    items = patch.get("files") if isinstance(patch, dict) else None
    if not isinstance(items, list):
        raise PatchRejected(["patch 는 {\"files\": [...]} 꼴이어야 합니다"])
    for n, it in enumerate(items, 1):
        p = it.get("path")
        if p not in known:
            problems.append(f"{n}번: 받은자료의 전기조서가 아닙니다: {p}")
        if it.get("basis", "AI") not in ("AI", "사람"):
            problems.append(f"{n}번: basis 는 AI 또는 사람입니다")
        if it.get("basis", "AI") == "AI" and not folder._is_area(str(it.get("area") or "")):
            problems.append(f"{n}번: area 는 4000 영역 번호여야 합니다: {it.get('area')}")
        for b in it.get("blocks") or []:
            if b.get("kind") not in KINDS or not isinstance(b.get("lines"), list) or not all(isinstance(x, str) for x in b["lines"]):
                problems.append(f"{n}번: 블록은 kind({' · '.join(KINDS)})와 글 줄 목록 lines 가 있어야 합니다")
        for pr in it.get("procedures") or []:
            if not isinstance(pr.get("name"), str) or not pr["name"].strip():
                problems.append(f"{n}번: 절차는 name 이 있어야 합니다")
    if problems:
        raise PatchRejected(problems)
    data = _idx(d, INDEX, None) or collect(d, now=now)
    added = confirmed = 0
    for it in items:
        p = it["path"]
        if it.get("basis", "AI") == "사람":
            for a in data["areas"].values():
                for b in a["blocks"] + a["procedures"]:
                    if b["source"]["path"] == p and b.get("basis") == "AI" and not b.get("confirmed"):
                        b["confirmed"] = True
                        confirmed += 1
            continue
        a = data["areas"].setdefault(it["area"], {"files": [], "blocks": [], "procedures": [], "adjustments": []})
        a["blocks"] = [b for b in a["blocks"] if not (b["source"]["path"] == p and b.get("basis") == "AI")]
        a["procedures"] = [x for x in a["procedures"] if not (x["source"]["path"] == p and x.get("basis") == "AI")]
        if p not in a["files"]:
            a["files"].append(p)
        lim = cfg["scan"]
        for b in it.get("blocks") or []:
            lines = [mask(re.sub(r"\s+", " ", x).strip())[:lim["line_chars"]] for x in b["lines"] if x.strip()]
            a["blocks"].append({"kind": b["kind"], "title": mask(str(b.get("title") or cfg["kind_names"][b["kind"]])),
                                "lines": lines[:lim["block_lines"]], "truncated": max(0, len(lines) - lim["block_lines"]), "role": "lead",
                                "source": _src(p, str(b.get("sheet") or ""), str(b.get("cell") or "")), "basis": "AI", "confirmed": False})
            added += 1
        for pr in it.get("procedures") or []:
            a["procedures"].append({"name": mask(pr["name"].strip()), "ref": pr.get("ref"), "applies": None, "result": None,
                                    "source": _src(p, str(pr.get("sheet") or ""), str(pr.get("cell") or "")), "basis": "AI", "confirmed": False})
            added += 1
    data["need_ai"] = [p for p in data.get("need_ai", []) if p not in {it["path"] for it in items}]
    for f in data["files"]:
        if f["path"] in {it["path"] for it in items}:
            f["need_ai"] = False
    folder._write(d / "_색인" / INDEX, data)
    folder.validate_index(d)
    return {"added": added, "confirmed": confirmed}


# ---------------------------------------------------------------- 당기 조서가 읽는 것
def load(d: Path) -> dict | None:
    return _idx(d, INDEX, None)


def for_area(d: Path, no: str, data: dict | None = None) -> dict | None:
    data = data if data is not None else load(d)
    a = ((data or {}).get("areas") or {}).get(no)
    return a if a and (a["blocks"] or a["procedures"] or a["adjustments"]) else None


def general_inputs(d: Path, no: str, data: dict | None = None) -> dict | None:
    """작년 감사 폴더의 같은 번호 조서에서 회계사가 적은 값 {"source", "inputs": [{sheet, label, value}]}."""
    data = data if data is not None else load(d)
    g = ((data or {}).get("general") or {}).get(no)
    return g if g and g.get("inputs") else None


def _where(path: str) -> str:
    """출처 파일을 작년 감사 폴더 이름과 함께(2026-10-16 시연: 파일 이름만 적혀 당기 4000P_매출_v1.1.xlsx 와 구별되지 않았다).
    「…/2024 기말감사/_보관/2025-03-14/4000 계정별/4000P_매출_v1.1.xlsx」 이면 「2024 기말감사 · _보관/2025-03-14 · 4000P_매출_v1.1.xlsx」."""
    parts = Path(path).parts
    i = next((k for k, x in enumerate(parts) if x.endswith(("기말감사", "반기검토"))), None)
    if i is None:
        return Path(path).name
    keep = [parts[i]]
    if len(parts) > i + 3 and parts[i + 1] == "_보관":
        keep.append(f"{parts[i + 1]}/{parts[i + 2]}")
    return " · ".join(keep + [Path(path).name])


def source_text(b: dict) -> str:
    s = b["source"]
    where = _where(s["path"]) + (f" '{s['sheet']}'" if s.get("sheet") else "") + (f"!{s['cell']}" if s.get("cell") else "")
    return where + ("" if b.get("confirmed", True) else " · AI 판독, 확인 전")


def match_procedures(steps: list[str], procedures: list[dict], cfg: dict | None = None) -> tuple[dict[int, list[dict]], list[dict]]:
    """전기 절차를 당기 표지의 절차 줄에 잇는다. ({표지 줄 번호: [전기 절차]}, 잇지 못한 전기 절차).
    작년 엔진 조서의 절차는 이름이 같으면 잇고, 손 조서의 절차는 prior_wp.yaml procedure_match 의 낱말로 잇는다."""
    cfg = cfg or defs()
    rules = cfg.get("procedure_match") or []
    out: dict[int, list[dict]] = {}
    left = []
    flat = lambda t: re.sub(r"\s+", " ", str(t)).strip()  # noqa: E731
    low = [s.lower() for s in steps]
    for p in procedures:
        name = p["name"].lower()
        head = flat(p["name"]).rstrip("…")              # 엔진 조서의 기입란 이름은 절차 글을 30자에서 자른 것이다
        hit = next((i for i, s in enumerate(steps) if head and flat(s).startswith(head)), None)
        if hit is None:
            for r in rules:
                if any(w.lower() in name for w in r["words"]):
                    hit = next((i for i, s in enumerate(low) if any(k.lower() in s for k in r["steps"])), None)
                    if hit is not None:
                        break
        if hit is None:
            left.append(p)
        else:
            out.setdefault(hit, []).append(p)
    return out, left


def procedure_text(p: dict) -> str:
    t = p["name"] + (f" ({p['ref']})" if p.get("ref") else "")
    if p.get("applies"):
        t += f" · {p['applies']}"
    if p.get("result"):
        t += f" · {p['result']}"
    return t


def reference_sheet(b, d: Path, no: str, title: str = "전기 기입(참고)") -> bool:
    """일반조서에 작년 같은 번호 조서의 회계사 기입을 참고 시트로 싣는다. 기입란이 아니다. 실은 것이 있으면 True."""
    from . import wp

    g = general_inputs(d, no)
    if not g:
        return False
    ws, _ = b.sheet(title, [22, 44, 70], head=False)        # 조서 머리(검토자 기입란)를 두지 않는다. 이 시트에는 기입란이 없다
    ws["A1"] = f"{b.no}  {title}"
    ws["A1"].font = wp.font(size=14, bold=True, color=wp.NAVY)
    ws["A2"] = f"작년 조서에서 회계사가 적은 내용입니다. 참고이고 당기 기입란이 아닙니다. 출처: {g['source']}"
    ws["A2"].font = wp.font(size=9, color=wp.MUTED)
    row = b.head(ws, 4, ["시트", "항목", "전기에 적은 내용"])
    for x in g["inputs"]:
        b.put(ws, row, 1, x["sheet"])
        b.put(ws, row, 2, x["label"], wrap=True)
        b.put(ws, row, 3, x["value"], wrap=True)
        row += 1
    return True

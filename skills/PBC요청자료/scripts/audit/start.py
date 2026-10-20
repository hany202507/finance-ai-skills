"""시작 자료: 전기 감사보고서 · 전기조서 · 전기 재무제표 · DART 로 회사 정보와 첫 조서를 채운다(R2, 2026-10-05).

  build(d, dart_target=None, force=False)       읽기 → _색인/시작자료.json → 감사건 company → 조서(PBC · 1100 · 1200 · 2100)
                                               → 요청 메일 초안 → 폴더 훑기 · 지휘. 보고(dict)를 돌려준다
  read_prior_report(path)                      전기 감사보고서 PDF · DART html 판독
  read_prior_workpapers(paths)                 전기조서 엑셀(2700A · 8100B · 8100A) 판독
  apply(d, patch)                              엔진이 못 찾은 글 항목(감사인 · 의견 · 업종 등)을 AI · 사람 판독으로 넣는다
  mark_requested(d, at)                        요청 메일을 보낸 날을 빠진 자료에 적는다(회계사가 보낸 뒤)
  write_mail(d)                                빠진 자료 요청 메일 초안(.txt)을 1000 감사계약/ 에 쓴다

숫자는 엔진이 표 칸 · 엑셀 셀에서 읽은 값만 쓴다. 감사의견 · 감사인 같은 글 항목은 정해진 문형으로 찾고,
못 찾으면 비워 두고 need_ai 에 남긴다. 판단 칸(계약 수락 · 위험 평가 · 결론)은 기입란으로 비운다.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path

from . import dart, director, figures, folder, index, prior, reader, wp
from .errors import IntakeError

START_FILE = "시작자료.json"
PRIOR_TEXT_AI = "전기조서에서 글 블록 라벨을 찾지 못했습니다. 읽어서 prior apply 로 넣어 주십시오(확인 전)"
AREA_WP = "계정 조서(글은 전기조서.json)"
PRIOR_TABLE_AI = "전기조서에서 엔진이 아는 표(중요성 · 수정사항 · 미수정 왜곡 · 정산표)를 찾지 못했습니다"
PRIOR_REPORT = "전기 감사보고서"
PRIOR_WP = "전기조서"
PRIOR_FS = figures.PRIOR_FS
TAX_KINDS = ("부가세 신고서", "법인세 신고서", "원천세 신고서")
KEY_FIGS = ("자산총계", "부채총계", "자본총계", "매출액", "영업이익", "당기순이익")
MAIL_PREFIX = "PBC_요청자료목록_메일_"
AI_FIELDS = {"prior_report.auditor": str, "prior_report.opinion": str, "prior_report.emphasis": bool,
             "prior_report.going_concern": bool, "prior_report.report_date": str, "company.industry": str}

# PBC 분류(한공회 요청자료 목록 순서). 번호는 pbc_match.yaml 의 row 와 같다
PBC_CATEGORY = {1: "결산", 2: "결산", 3: "결산", 4: "결산", 5: "세무", 6: "세무", 7: "세무", 8: "지배구조", 9: "지배구조",
                10: "현금·차입", 11: "현금·차입", 12: "매출·채권", 13: "재고", 14: "유형·무형", 15: "유형·무형",
                16: "인건비", 17: "우발·약정", 18: "기타"}


def _now() -> str:
    return folder._now()


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _won(v) -> str:
    return f"{v:,}원" if isinstance(v, int) else str(v)


def _idx(d: Path, name: str, default=None):
    p = Path(d) / "_색인" / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def _fy_end(case: dict) -> date:
    return date.fromisoformat(case["fiscal_year_end"])


# ---------------------------------------------------------------- 전기 감사보고서
def _pdf_lines(path: Path) -> list[tuple[int, str]]:
    out = []
    for no, text in reader.pdf_text(path):
        for line in text.splitlines():
            if line.strip():
                out.append((no, line.strip()))
    return out


AMOUNT_TAIL = re.compile(r"^(.*?[가-힣A-Za-z\)])\s*(\(?[△\-]?[\d,]+\)?|-)(?:\s+(\(?[△\-]?[\d,]+\)?|-))?\s*$")


PAGE_NOISE = re.compile(r"^(전자공시시스템\s*dart\.fss\.or\.kr|Page\s*\d+|-\s*\d+\s*-)$")
STATEMENT_START = ("재무상태표", "대차대조표")
STATEMENT_STOP = ("자본변동표", "현금흐름표", "이익잉여금처분계산서", "결손금처리계산서", "이익잉여금처분계산서(안)")


def _statement_lines(lines: list[tuple[int, str]]) -> list[tuple[int, str]]:
    """재무상태표 제목부터 자본변동표 · 현금흐름표 제목 전까지(재무상태표 · 손익계산서 쪽). 목차 줄(점선)은 제목으로 보지 않는다.
    현금흐름표 「매출채권의 증감」 · 주석 표가 요약 숫자로 잡히지 않게 한다(2026-10-18 실제 회사 자료). 제목이 없으면 전부."""
    start = next((i for i, (_, t) in enumerate(lines) if dart.norm(t) in STATEMENT_START), None)
    if start is None:
        return lines
    end = next((i for i in range(start + 1, len(lines)) if dart.norm(lines[i][1]) in STATEMENT_STOP), len(lines))
    return lines[start:end]


def _pdf_items(lines: list[tuple[int, str]]) -> dict[str, tuple[int, int]]:
    """'과목 금액' 이 한 줄이거나, 과목 줄 다음 줄이 금액인 꼴을 읽는다. {과목: (금액, 쪽)}
    과목 줄의 숫자는 주석 번호 「(주석4,12)」 를 뗀 뒤에 본다. 떼지 않으면 「매출채권(주석4,12)」 줄을 놓친다."""
    out: dict[str, tuple[int, int]] = {}
    lines = [x for x in lines if not PAGE_NOISE.match(x[1])]      # 쪽 머리글 · 바닥글이 과목 줄과 금액 줄 사이에 끼면 놓친다
    for i, (page, line) in enumerate(lines):
        m = AMOUNT_TAIL.match(line)
        if m and dart.amount(m.group(2)) is not None and not re.search(r"\(주석[\d,\s]*$", m.group(1)):
            name, v = dart.item_name(m.group(1)), dart.amount(m.group(2))
        elif i + 1 < len(lines) and not re.search(r"\d", dart.item_name(line)) and dart.amount(lines[i + 1][1]) is not None:
            name, v = dart.item_name(line), dart.amount(lines[i + 1][1])
        else:
            continue
        if name and len(name) <= 30 and name not in out:
            out[name] = (v, page)
    return out


def _report_date(texts: list[str], fiscal_year_end: str | None) -> str | None:
    """감사보고서일. 「감사보고서일(2024년 3월 6일) 현재로 유효」 문형이 먼저, 없으면 첨부 재무제표 앞의 날짜만 있는 줄 가운데
    결산일이 아닌 마지막 줄(서명 날짜). 결산일을 감사보고서일로 읽던 것을 고쳤다(2026-10-18 실제 회사 자료)."""
    m = re.search(r"감사보고서일\(?(20\d\d년\d{1,2}월\d{1,2}일)", dart.norm(" ".join(texts)))
    if m:
        return dart.date_of(m.group(1))
    stop = next((i for i, t in enumerate(texts) if dart.norm(t) in ("(첨부)재무제표",) + STATEMENT_START), len(texts))
    dates = [dart.date_of(t) for t in texts[:stop] if re.fullmatch(r"20\d\d년\d{1,2}월\d{1,2}일", dart.norm(t))]
    dates = [x for x in dates if x != fiscal_year_end]
    return dates[-1] if dates else None


def corp_name(name: str | None) -> str:
    """회사명 비교용. 「주식회사 · (주) · ㈜」 와 공백을 뗀다(감사 폴더 「가온정밀」 과 보고서 「주식회사 가온정밀」 은 같다)."""
    n = dart.norm(name or "")
    for t in ("주식회사", "(주)", "（주）", "㈜"):
        n = n.replace(t, "")
    return n


def read_prior_report(path: Path) -> dict:
    """전기 감사보고서 한 파일(PDF 또는 DART html). 숫자는 표에서, 의견 · 감사인은 정해진 문형으로."""
    path = Path(path)
    if path.suffix.lower() in (".html", ".htm"):
        r = dart.parse(path.read_bytes().decode("utf-8", errors="replace"))
        figs = {k: v[0] for k, v in {**r["key_bs"], **r["key_is"]}.items() if v[0] is not None}
        return {"format": "DART", "company": r["company"], "auditor": r["auditor"], "opinion": r["opinion"],
                "emphasis": r["emphasis"], "going_concern": r["going_concern"], "report_date": r["report_date"],
                "fiscal_year_end": r["fiscal_year_end"], "framework": r["framework"], "figures": figs,
                "figure_pages": {}, "need_ai": list(r["need_ai"]), "dart": r, "term": r.get("term")}
    lines = _pdf_lines(path)
    if not lines:
        return {"format": "PDF", "company": None, "auditor": None, "opinion": None, "emphasis": None, "going_concern": None,
                "report_date": None, "fiscal_year_end": None, "framework": None, "figures": {}, "figure_pages": {}, "term": None,
                "need_ai": ["PDF 에서 글자를 뽑지 못했습니다(스캔본). 글자 판독이 필요합니다"]}
    texts = [t for _, t in lines]
    whole = " ".join(texts)
    out = {"format": "PDF", "need_ai": []}
    company = None
    for i, t in enumerate(texts):
        n = dart.norm(t)
        if n.endswith("귀중"):
            company = n[: n.find("주주")] if "주주" in n else ""
            if not company and i:
                company = dart.norm(texts[i - 1])
            break
    out["company"] = company or None
    a = whole.find("감사의견")
    b = whole.find("근거", a + 4) if a >= 0 else -1
    op_text = whole[a: b if b > 0 else None] if a >= 0 else whole
    out["opinion"] = dart.opinion_of(op_text)
    out["framework"] = dart.framework_of(op_text) or dart.framework_of(whole)
    out["fiscal_year_end"] = dart.date_of(op_text)
    nw = dart.norm(whole)
    out["emphasis"] = "강조사항" in nw
    out["going_concern"] = "계속기업관련중요한불확실성" in nw
    out["auditor"] = dart.auditor_of(texts)
    out["term"] = dart.term_of(texts)
    out["report_date"] = _report_date(texts, out["fiscal_year_end"])
    # 요약 숫자는 재무상태표 · 손익계산서 쪽에서 먼저 찾고, 거기 없는 과목만 나머지 쪽에서 찾는다
    items = _pdf_items(_statement_lines(lines))
    for k, v in _pdf_items(lines).items():
        items.setdefault(k, v)
    out["figures"], out["figure_pages"] = {}, {}
    for name, (v, page) in items.items():
        for key, pat in dart.KEY_BS + dart.KEY_IS:
            if key not in out["figures"] and re.search(pat, name):
                out["figures"][key] = dart.signed(name, v)          # 영업손실 · 당기순손실 · 결손금은 음수
                out["figure_pages"][key] = f"{page}쪽"
                break
    for k, label in (("company", "회사명"), ("auditor", "감사인"), ("opinion", "감사의견"), ("framework", "회계기준")):
        if not out[k]:
            out["need_ai"].append(f"{label} 문형을 찾지 못했습니다")
    if not out["figures"]:
        out["need_ai"].append("재무제표 요약 숫자를 찾지 못했습니다")
    return out


# ---------------------------------------------------------------- 전기조서
def _rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _header(rows: list[tuple], need: list[tuple[str, ...]]) -> tuple[int, dict[str, int]] | None:
    """첫 15행에서 머리줄을 찾는다. need 는 (열쇠, 후보 낱말...) 목록. 모두 찾으면 (행 번호, {열쇠: 열})."""
    for i, r in enumerate(rows[:15]):
        cells = [dart.norm(str(v)) if v is not None else "" for v in r]
        cols = {}
        for key, *words in need:
            j = next((j for j, c in enumerate(cells) if c and any(w in c for w in words) and j not in cols.values()), None)
            if j is None:
                break
            cols[key] = j
        else:
            return i, cols
    return None


def _cell(r: tuple, j: int | None):
    return r[j] if j is not None and j < len(r) else None


def _materiality(wb) -> dict:
    labels = [("performance", ("수행중요성",)), ("trivial", ("명백히사소한", "명백하게사소한")),
              ("overall", ("재무제표전체중요성", "전체중요성", "재무제표전체에대한중요성"))]
    out, sheet_unit = {}, {}
    for ws in wb.worksheets:
        for r in _rows(ws):
            first = next(((j, v) for j, v in enumerate(r) if isinstance(v, str) and v.strip()), None)
            if not first:
                continue
            j, label = first
            n = dart.norm(label)
            if len(n) > 20:
                continue
            key = next((k for k, words in labels if any(w in n for w in words)), None)
            if not key or key in out:
                continue
            k = next((k for k in range(j + 1, len(r)) if _is_num(r[k])), None)
            if k is None:
                continue
            basis = next((r[x] for x in range(k + 1, len(r)) if isinstance(r[x], str) and r[x].strip()), None)
            # 단위(R8, 진단 #11): 금액 바로 뒤 칸이 「백만원」 · 「천원」 이면 그 단위, 없으면 시트 위쪽 「(단위: 백만원)」
            unit = _unit_of(basis) or sheet_unit.setdefault(ws.title, _sheet_unit(ws))
            mult = UNITS.get(unit, 1)
            row = {"amount": reader._int(float(r[k]) * mult), "basis": None if _unit_of(basis) else (basis.strip() if basis else None),
                   "sheet": ws.title}
            if unit and unit != "원":
                row["unit"] = unit
                row["amount_in_unit"] = reader._int(float(r[k]))
            out[key] = row
    return out


UNITS = {"원": 1, "천원": 1_000, "백만원": 1_000_000, "억원": 100_000_000}


def _unit_of(v) -> str | None:
    """'백만원' · '(단위: 천원)' · '단위:원' 같은 칸이면 단위 이름."""
    if not isinstance(v, str):
        return None
    m = re.fullmatch(r"\(?\s*(?:단위\s*[:：]?\s*)?(원|천원|백만원|억원)\s*\)?", v.strip())
    return m.group(1) if m else None


def _sheet_unit(ws) -> str | None:
    """시트 첫 6행의 단위 표기."""
    for r in ws.iter_rows(values_only=True, max_row=6):
        for v in r:
            if isinstance(v, str) and "단위" in v:
                m = re.search(r"단위\s*[:：]?\s*(원|천원|백만원|억원)", v)
                if m:
                    return m.group(1)
    return None


def _adjustments(wb, unadjusted: bool) -> list[dict]:
    out = []
    for ws in wb.worksheets:
        rows = _rows(ws)
        # 머리줄: 「내용 · 금액」, 또는 회계법인 수정 시트 「구분 · 적요 · 차변 · 금액 · 대변 · 금액」(R8, 진단 #19)
        h = _header(rows, [("text", "내용"), ("amount", "금액")]) or _header(rows, [("text", "적요"), ("amount", "금액")])
        if not h:
            continue
        i, c = h
        heads = [dart.norm(str(v)) if v is not None else "" for v in rows[i]]

        def col(*words):
            return next((j for j, x in enumerate(heads) if x and any(w in x for w in words)), None)

        cols = {"no": col("번호", "#"), "kind": col("구분"), "dr": col("차변"), "cr": col("대변"),
                "account": col("계정"), "booked": col("반영"), "reason": col("사유")}
        for r in rows[i + 1:]:
            text = _cell(r, c["text"])
            amt = _cell(r, c["amount"])
            if not isinstance(text, str) or not text.strip() or not _is_num(amt):
                continue
            if dart.norm(text) in ("합계", "계", "명백히사소한금액") or dart.norm(text).startswith("명백히사소한"):
                continue
            row = {"text": text.strip(), "amount": reader._int(float(amt)), "sheet": ws.title}
            for k, j in cols.items():
                v = _cell(r, j)
                if v is not None and str(v).strip():
                    row[k] = reader._int(float(v)) if k == "no" and _is_num(v) else str(v).strip()
            out.append(row)
    if not out:
        out = _firm_adjustments(wb)
    return out


def _firm_adjustments(wb) -> list[dict]:
    """감사 법인 WTB 의 AJE · RJE 시트(No. | 내역 | 차변 A/C · 금액 | 대변 A/C · 금액, 두 줄 머리). 머리줄에 「금액」 이 다음 줄에 있어
    위 머리줄 규칙으로는 못 읽는다. 읽기는 DSD 정산표 대사(audit/report/dsd_wtb.py)와 같은 함수를 쓴다(2026-10-19 가상 회사 세트)."""
    from .report import dsd_wtb

    out = []
    for ws in wb.worksheets:
        if not dsd_wtb.AJE_SHEET.match(ws.title):
            continue
        try:
            entries = dsd_wtb._firm_aje(ws)
        except Exception:  # noqa: BLE001  읽기 전용 시트 등은 건너뛴다
            continue
        for e in entries:
            dr = [x for x in e["lines"] if x[0] == "차변"]
            cr = [x for x in e["lines"] if x[0] == "대변"]
            if not dr and not cr:
                continue
            row = {"text": e["desc"] or f"수정분개 {e['no']}", "amount": reader._int(float(sum(x[2] for x in dr) or sum(x[2] for x in cr))),
                   "sheet": ws.title, "no": e["no"], "dr": " · ".join(x[1] for x in dr), "cr": " · ".join(x[1] for x in cr)}
            out.append(row)
    return out


FIRM_FS_SHEETS = {"BS": "BS", "IS": "IS", "PL": "IS"}


def _firm_fs_lines(wb) -> dict[str, list[dict]]:
    """감사 법인 WTB 의 BS · IS 시트(과목 | 당기 | 전기, 표시과목 줄). {시트: [{name, kind, final, company}]}.
    이 양식은 계정 단위 WBS · WIS 와 표시과목 단위 BS · IS 를 함께 둔다. 전기 재무제표 줄은 표시과목 단위인 BS · IS 에서 읽는다.
    WBS · WIS 를 리드 꼴로 읽으면 계정 이름과 표시과목이 섞인다(2026-10-19 가상 회사 세트). 읽기는 DSD 정산표 대사와 같은 함수다."""
    from .report import dsd_wtb

    out = {}
    for ws in wb.worksheets:
        kind = FIRM_FS_SHEETS.get(dart.norm(ws.title).upper())
        if not kind:
            continue
        try:
            got = dsd_wtb._firm_fs(ws)
        except Exception:  # noqa: BLE001  양식이 다르면 리드 꼴로 읽는다
            got = None
        # 결손금 · 영업손실처럼 이름으로 손실을 나타낸 줄은 감사보고서 판독과 같이 음수로 둔다(dart.signed)
        lines = [{"name": x["name"], "kind": kind, "final": dart.signed(x["name"], reader._int(float(x["final"]))), "company": None}
                 for x in got or [] if x.get("final") is not None and x["name"]]
        if lines:
            out[ws.title] = lines
    return out


def _wtb(wb) -> dict:
    out = {"accounts": [], "lines": [], "sheets": []}
    firm = _firm_fs_lines(wb)
    for title, lines in firm.items():
        out["lines"] += lines
        out["sheets"].append(title)
    for ws in wb.worksheets:
        if ws.title in firm:
            continue
        rows = _rows(ws)
        h = _header(rows, [("code", "계정코드", "코드"), ("name", "계정명", "계정과목"), ("final", "최종")])
        if h:
            i, c = h
            line_col = _header(rows, [("line", "표시과목")])
            for r in rows[i + 1:]:
                code = reader._code(_cell(r, c["code"]))
                v = _cell(r, c["final"])
                if code is None or not _is_num(v):
                    continue
                out["accounts"].append({"code": str(code), "name": str(_cell(r, c["name"]) or "").strip(),
                                        "line": str(_cell(r, line_col[1]["line"]) or "").strip() if line_col else None,
                                        "final": reader._int(float(v))})
            out["sheets"].append(ws.title)
            continue
        h = _header(rows, [("line", "표시과목"), ("final", "최종")])
        if h:
            i, c = h
            heads = [dart.norm(str(v)) if v is not None else "" for v in rows[i]]
            kind_col = next((j for j, x in enumerate(heads) if x == "구분"), None)
            comp_col = next((j for j, x in enumerate(heads) if "회사제시" in x), None)
            for r in rows[i + 1:]:
                name, v = _cell(r, c["line"]), _cell(r, c["final"])
                if not isinstance(name, str) or not name.strip() or not _is_num(v):
                    continue
                kind = str(_cell(r, kind_col) or "").strip().upper() or None
                comp = _cell(r, comp_col)
                out["lines"].append({"name": name.strip(), "kind": kind, "final": reader._int(float(v)),
                                     "company": reader._int(float(comp)) if _is_num(comp) else None})
            out["sheets"].append(ws.title)
            continue
        got = [] if firm else _lead(rows)
        if got:
            out["lines"] += got
            out["sheets"].append(ws.title)
    return out


LEAD_NAME = ("계정과목", "과목", "BS", "PL", "IS", "NAME", "계정")
LEAD_FINAL = ("AUDITED", "FINAL", "REVIEWED", "최종", "감사후", "수정후")
LEAD_COMPANY = ("PROPOSED", "PRELIMINARY", "회사제시", "감사전", "수정전")


def _lead(rows: list[tuple]) -> list[dict]:
    """회계법인 계정별 리드(R8, 진단 #19): 두 줄 머리 「계정과목 | 기말일 | 수정사항 | 기말일 | 전기말」 /
    「 | Proposed | 차변 | 대변 | Audited | Audited」 처럼 코드 열 없이 수정 전 · 수정 · 수정 후(Audited · Final · Reviewed) 열을 둔 표.
    수정 후 열이 둘 이상이면 왼쪽(당기)이다. [{name, kind: None, final, company}]"""
    for i in range(min(15, len(rows))):
        top = [dart.norm(str(v)).upper() if v is not None else "" for v in rows[i]]
        nxt = [dart.norm(str(v)).upper() if v is not None else "" for v in rows[i + 1]] if i + 1 < len(rows) else []
        name_col = next((j for j, c in enumerate(top) if c in LEAD_NAME), None)
        if name_col is None:
            continue
        both = [(j, (top[j] if j < len(top) else "") + "|" + (nxt[j] if j < len(nxt) else "")) for j in range(name_col + 1, max(len(top), len(nxt)))]
        final = next((j for j, c in both if any(w in c for w in LEAD_FINAL)), None)
        if final is None:
            continue
        comp = next((j for j, c in both if any(w in c for w in LEAD_COMPANY) and j < final), None)
        start = i + 2 if any(any(w in c for w in LEAD_FINAL + LEAD_COMPANY) for c in nxt) else i + 1
        out = []
        for r in rows[start:]:
            name, v = _cell(r, name_col), _cell(r, final)
            if not isinstance(name, str) or not name.strip() or not _is_num(v):
                continue
            n = dart.norm(name)
            if n in ("합계", "계", "총계", "소계") or n.upper() in ("TOTAL", "SUBTOTAL"):
                continue
            c = _cell(r, comp)
            out.append({"name": name.strip(), "kind": None, "final": reader._int(float(v)),
                        "company": reader._int(float(c)) if _is_num(c) else None})
        if out:
            return out
    return []


def read_prior_workpapers(d: Path, paths: list[str]) -> dict:
    """전기조서 엑셀들에서 전기 중요성 · 수정사항 · 미수정 왜곡 · 정산표 최종 금액을 읽는다.
    글 블록 라벨이 있는 계정 조서(파일 이름이 4000 영역으로 이어지고 Audit memo · 감사절차 · Conclusion 같은 라벨이 있는 손 조서)는
    여기서 읽지 않는다. Lead 표와 표본 표가 정산표 · 수정사항 머리줄과 겹쳐 전기 숫자로 읽히기 때문이다(예시 조서 25개 가운데 10개가
    그랬고 한 파일은 수정사항이 1,044줄로 읽혔다, 2026-10-12). 그 조서의 글은 audit/priorwp.py 가 읽어 _색인/전기조서.json 에 둔다.
    라벨이 없는 법인 리드 꼴(계정과목 · Proposed · 수정사항 · Audited 열, R8)은 앞에서처럼 정산표로 읽는다."""
    from . import priorwp

    out = {"materiality": {}, "adjustments": [], "unadjusted": [], "wtb": {"accounts": [], "lines": []}, "files": {}, "need_ai": []}
    skip = priorwp.text_workpapers(Path(d), paths)
    for rel in paths:
        p = Path(d) / rel
        if rel in skip:
            out["files"][rel] = [AREA_WP]
            continue
        if p.suffix.lower() not in reader.EXCEL:
            out["need_ai"].append(f"엑셀이 아닌 전기조서는 엔진이 읽지 않습니다: {rel}")
            continue
        try:
            wb = reader.open_plain(p)                  # 손 조서는 깨진 이름 정의가 있어도 열리게 이름 정의를 빼고 연다(2026-10-12)
        except Exception as e:  # noqa: BLE001  열지 못한 파일은 남기고 다음 파일로 간다
            out["need_ai"].append(f"전기조서를 열지 못했습니다({type(e).__name__}): {rel}")
            continue
        try:
            name = dart.norm(p.stem)
            got = []
            mat = _materiality(wb)
            if mat and not out["materiality"]:
                out["materiality"] = {k: dict(v, path=rel) for k, v in mat.items()}
                got.append("중요성")
            unadj = "미수정" in name or "8100A" in name or any("미수정" in ws.title for ws in wb.worksheets)
            adj = _adjustments(wb, unadj)
            if adj:
                (out["unadjusted"] if unadj else out["adjustments"]).extend(dict(a, path=rel) for a in adj)
                got.append("미수정 왜곡" if unadj else "수정사항")
            w = _wtb(wb)
            if w["accounts"] or w["lines"]:
                out["wtb"]["accounts"] += w["accounts"]
                out["wtb"]["lines"] += w["lines"]
                out["wtb"]["path"] = rel
                got.append("정산표")
            out["files"][rel] = got
            if not got:
                out["need_ai"].append(f"{PRIOR_TABLE_AI}: {rel}")
        finally:
            wb.close()
    return out


# ---------------------------------------------------------------- 다른 받은자료
# 업종 칸 옆에 오는 다른 칸 이름(2026-10-18 실제 회사 자료 부가세 신고서: 「업 종 코 드」 · 「해당국가」 를 업종으로 읽었다)
INDUSTRY_LABEL = re.compile(r"코드|국가|금액|번호|구분|일자|합계|^\(?\d")


def _industry_from_docs(d: Path) -> tuple[str, str] | None:
    """세무신고 PDF 의 '업종' 칸. (업종, '파일 n쪽')."""
    for kind in TAX_KINDS:
        for rel in figures.received(d, kind):
            p = Path(d) / rel
            if p.suffix.lower() != ".pdf":
                continue
            lines = _pdf_lines(p)
            for i, (page, t) in enumerate(lines):
                n = dart.norm(t)
                if n == "업종" and i + 1 < len(lines) and not INDUSTRY_LABEL.search(dart.norm(lines[i + 1][1])):
                    return lines[i + 1][1].strip(), f"{rel} {page}쪽"
                m = re.match(r"^업\s*종\s*[:：]?\s*(\S.*)$", t)
                if m and not INDUSTRY_LABEL.search(dart.norm(m.group(1))):      # 「업 종 코 드」 칸 이름을 업종으로 읽지 않는다
                    return m.group(1).strip(), f"{rel} {page}쪽"
    return None


def _current_journal(d: Path, fy: date) -> str | None:
    docs = _idx(d, "문서목록.json", {"files": []})
    for f in docs["files"]:
        if f.get("role") == "received" and f.get("kind") == reader.JOURNAL and Path(f["path"]).suffix.lower() in reader.EXCEL:
            rng = reader.journal_range(f.get("facts") or [])
            if rng and rng[0].year == rng[1].year == fy.year:
                return f["path"]
    return None


def _code3(code) -> int | None:
    if code is None:
        return None
    return code if code < 1000 else int(str(code)[:3])


def top_parties(d: Path, fy: date, n: int = 10) -> dict:
    """당기 분개장에서 매출처(더존 401~430 대변 순액)와 매입처(251~252 대변) 상위 n 곳."""
    rel = _current_journal(d, fy)
    if not rel:
        return {"path": None, "sales": [], "purchases": [], "sales_total": 0, "purchases_total": 0}
    _, lines = reader._journal(Path(d) / rel)
    sales: dict[str, float] = {}
    buys: dict[str, float] = {}
    for _, _, dr, cr, party, _, code in lines:
        c = _code3(code)
        if c is None or not party:
            continue
        if 401 <= c <= 430:
            sales[party] = sales.get(party, 0) + cr - dr
        elif 251 <= c <= 252:
            buys[party] = buys.get(party, 0) + cr

    def top(m):
        total = sum(v for v in m.values() if v > 0)
        rows = sorted(((k, reader._int(v)) for k, v in m.items() if v > 0), key=lambda x: (-x[1], x[0]))
        return [{"party": k, "amount": v, "share": round(v / total, 4) if total else None} for k, v in rows[:n]], reader._int(total)

    s, st = top(sales)
    b, bt = top(buys)
    return {"path": rel, "sales": s, "purchases": b, "sales_total": st, "purchases_total": bt, "sales_count": len(sales),
            "purchases_count": len(buys)}


# ---------------------------------------------------------------- 당기 회사제시와 전기 비교
def _credit(code: int) -> bool:
    """더존 계정 대역에서 대변 성격(부채 · 자본 · 매출 · 영업외수익)."""
    return 251 <= code <= 430 or 901 <= code <= 950


def current_statement(d: Path) -> dict:
    """회사제시 시산표(마감 전)를 표시 금액으로. 이익잉여금에는 당기순이익을 더해 마감 후로 맞춘다. 합계는 더존 대역으로 센다."""
    rows = figures.current_tb(d)
    if not rows:
        return {"lines": {}, "totals": {}, "kinds": {}}
    lines: dict[str, int] = {}
    kinds: dict[str, str] = {}
    codes: dict[str, int] = {}
    for r in rows:
        c = _code3(int(r["code"]))
        sign = -1 if _credit(c) else 1
        lines[r["line"]] = lines.get(r["line"], 0) + sign * r["amount"]
        kinds.setdefault(r["line"], "BS" if c < 400 else "IS")
        codes.setdefault(r["line"], c)
    pl = [(_code3(int(r["code"])), r["amount"]) for r in rows]
    ni = -sum(a for c, a in pl if c >= 400)
    for name, c in codes.items():
        if 370 <= c <= 379 or "이익잉여금" in name:
            lines[name] += ni
            break
    t = {"자산총계": sum(a for c, a in pl if c < 251), "부채총계": -sum(a for c, a in pl if 251 <= c <= 330),
         "자본총계": -sum(a for c, a in pl if 331 <= c <= 399) + ni, "매출액": -sum(a for c, a in pl if 401 <= c <= 430),
         "매출총이익": -sum(a for c, a in pl if 401 <= c <= 499), "영업이익": -sum(a for c, a in pl if 401 <= c <= 899),
         "법인세비용차감전순이익": -sum(a for c, a in pl if 401 <= c <= 997), "당기순이익": ni}
    t["부채와자본총계"] = t["부채총계"] + t["자본총계"]
    return {"lines": lines, "totals": t, "kinds": kinds}


def big_changes(cur: dict, prior_fs: dict | None, threshold: int | None) -> dict:
    """표시과목 · 합계마다 전기 대비 증감. 큰 변동: |증감| ≥ 기준 금액(있으면) 이고 |증감률| ≥ 10% 또는 전기 없음."""
    if not prior_fs:
        return {"rows": [], "unmatched_current": sorted(cur["lines"]), "unmatched_prior": [], "threshold": threshold}
    prior = {**prior_fs.get("bs", {}), **prior_fs.get("is", {})}
    rows, seen = [], set()
    names = list(cur["lines"]) + [k for k in cur["totals"] if k not in cur["lines"]]
    for name in names:
        c = cur["lines"].get(name, cur["totals"].get(name))
        p = prior.get(name)
        if p is None and name in cur["totals"]:
            continue
        seen.add(name)
        diff = c - (p or 0)
        ratio = None if not p else diff / abs(p)
        # 전기 과목이 없으면 새 과목인지 이름만 다른지 엔진이 모른다. 큰 변동으로 세지 않고 따로 표시한다
        big = p is not None and (threshold is None or abs(diff) >= threshold) and (p == 0 or abs(ratio) >= 0.1) and diff != 0
        rows.append({"name": name, "kind": cur["kinds"].get(name) or ("IS" if name in prior_fs.get("is", {}) else "BS"),
                     "prior": p, "current": c, "diff": diff, "ratio": None if ratio is None else round(ratio, 4),
                     "total": name in cur["totals"] and name not in cur["lines"], "big": big, "prior_missing": p is None})
    return {"rows": rows, "unmatched_current": [r["name"] for r in rows if r["prior_missing"]],
            "unmatched_prior": [k for k in prior if k not in seen], "threshold": threshold}


# ---------------------------------------------------------------- 모으기
def _pick_report(d: Path) -> tuple[str | None, str | None]:
    """(회사가 준 전기 감사보고서 PDF, DART html). 여럿이면 이름 순 마지막(최근 연도)."""
    files = sorted(figures.received(d, PRIOR_REPORT))
    pdf = [p for p in files if p.lower().endswith(".pdf")]
    htm = [p for p in files if p.lower().endswith((".html", ".htm"))]
    return (pdf[-1] if pdf else None), (htm[-1] if htm else None)


def _dart_name(r: dict, src: str) -> str:
    key = r.get("rcp_no") or re.sub(r"[^0-9A-Za-z가-힣]+", "_", Path(src).stem)[:30]
    return f"DART_감사보고서_{key}.html"


def gather(d: Path, dart_target: str | None = None) -> dict:
    """받은자료와 DART 에서 시작 자료를 모은다(아직 쓰지 않는다)."""
    d = Path(d)
    case = _idx(d, "감사건.json")
    _fy_end(case)                    # 결산일이 없으면 여기서 멈춘다
    # 전기말 = 사업연도 시작일 전날(반기 · 짧은 사업연도는 감사건 fiscal_year_start, R9)
    prior_end = folder.periods(case)["prior_end"].isoformat()
    warnings: list[str] = []
    saved_dart = None
    if dart_target:
        html = dart.fetch(dart_target)
        r = dart.parse(html)
        dest = d / "00 받은자료" / "전기" / _dart_name(r, dart_target)
        src = Path(dart_target)
        if not (src.exists() and src.resolve() == dest.resolve()):
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_text(html, encoding="utf-8")
        saved_dart = dest.relative_to(d).as_posix()
        folder.scan(d)
    else:
        # 받은자료 목록을 새로 훑는다(내용 판독은 index read 몫). 요청 단계에서 전기 감사보고서 · 전기조서만 넣고 start 를 돌리면
        # 앞선 색인에 그 파일이 없어 「전기 감사보고서가 받은자료에 없습니다」 가 나왔다(2026-10-19 가상 회사 세트 단계 1)
        folder.scan(d, quiet_log=True)

    pdf, htm = _pick_report(d)
    rep = read_prior_report(d / pdf) if pdf else None
    drep = read_prior_report(d / htm) if htm else None
    prior_report = None
    if rep or drep:
        base = rep or drep
        src = pdf or htm
        prior_report = {k: base.get(k) for k in ("auditor", "opinion", "emphasis", "going_concern", "report_date",
                                                  "fiscal_year_end", "framework", "figures", "term")}
        prior_report.update(path=src, format=base["format"], figure_pages=base.get("figure_pages", {}),
                            need_ai=list(base["need_ai"]), company=base.get("company"))
        prior_report["basis"] = {k: ("엔진" if prior_report.get(k) is not None else "판독 필요")
                                 for k in ("auditor", "opinion", "emphasis", "report_date", "framework")}
        if prior_report["fiscal_year_end"] and prior_report["fiscal_year_end"] != prior_end:
            warnings.append(f"전기 감사보고서의 결산일({prior_report['fiscal_year_end']})이 전기 결산일({prior_end})과 다릅니다: {src}")

    dart_part = None
    if drep:
        r = drep["dart"]
        dart_part = {"path": htm, "source": r.get("source"), "rcp_no": r.get("rcp_no"), "corp_code": r.get("corp_code"),
                     "company": r.get("company"), "industry": r.get("industry"), "auditor": r.get("auditor"),
                     "opinion": r.get("opinion"), "framework": r.get("framework"), "emphasis": r.get("emphasis"),
                     "going_concern": r.get("going_concern"), "report_date": r.get("report_date"),
                     "fiscal_year_end": r.get("fiscal_year_end"), "unit": r.get("unit"), "term": r.get("term"),
                     "key_bs": {k: v[0] for k, v in r["key_bs"].items()}, "key_is": {k: v[0] for k, v in r["key_is"].items()},
                     "bs_items": len(r["bs"]), "is_items": len(r["is"]), "overview": r.get("overview"), "need_ai": r["need_ai"]}
        if rep and prior_report:
            # 회사가 준 감사보고서에서 못 찾은 글 항목은 같은 보고서의 DART 공시에서 엔진이 읽은 값으로 채운다
            filled = []
            for k in ("auditor", "opinion", "report_date", "framework"):
                if prior_report.get(k) is None and drep.get(k):
                    prior_report[k] = drep[k]
                    prior_report["basis"][k] = "엔진"
                    filled.append(k)
            if filled:
                prior_report["filled_from_dart"] = filled
                labels = [NEED_LABEL.get(k, k) for k in filled]
                prior_report["need_ai"] = [n for n in prior_report["need_ai"] if not any(n.startswith(x) for x in labels)]
        if rep:
            diffs = []
            for k in ("opinion", "auditor"):
                if rep.get(k) and drep.get(k) and dart.norm(rep[k]) != dart.norm(drep[k]):
                    diffs.append(f"{k}: 감사보고서 {rep[k]} · DART {drep[k]}")
            for k, v in rep["figures"].items():
                dv = drep["figures"].get(k)
                if dv is not None and dv != v:
                    diffs.append(f"{k}: 감사보고서 {v:,} · DART {dv:,}")
            dart_part["diffs_vs_report"] = diffs
            warnings += [f"전기 감사보고서와 DART 가 다릅니다. {x}" for x in diffs]

    wps_paths = sorted(figures.received(d, PRIOR_WP))
    pw = read_prior_workpapers(d, wps_paths) if wps_paths else None

    # 전기 재무제표: 받은자료 엑셀 → 전기조서 정산표 → DART → 감사보고서 요약
    prior_fs = None
    xl = figures.prior_excel(d)
    if xl and (xl["bs"] or xl["is"]):
        prior_fs = {"as_of": xl.get("as_of"), "bs": xl["bs"], "is": xl["is"], "source": xl["path"], "basis": "전기 재무제표 엑셀"}
    elif pw and pw["wtb"]["lines"]:
        ls = pw["wtb"]["lines"]
        prior_fs = {"as_of": prior_end, "bs": {x["name"]: x["final"] for x in ls if x["kind"] != "IS"},
                    "is": {x["name"]: x["final"] for x in ls if x["kind"] == "IS"}, "source": pw["wtb"]["path"],
                    "basis": "전기조서 8100B 최종 열"}
    elif drep and (drep["dart"]["bs"] or drep["dart"]["is"]):
        prior_fs = {"as_of": drep.get("fiscal_year_end") or prior_end, "bs": drep["dart"]["bs"], "is": drep["dart"]["is"],
                    "source": htm, "basis": "DART 감사보고서 재무제표"}
    elif rep and rep["figures"]:
        f = rep["figures"]
        prior_fs = {"as_of": rep.get("fiscal_year_end") or prior_end,
                    "bs": {k: v for k, v in f.items() if k in dict(dart.KEY_BS)},
                    "is": {k: v for k, v in f.items() if k in dict(dart.KEY_IS)}, "source": pdf, "basis": "전기 감사보고서 요약"}

    # 회사 정보
    name = case["company"]["name"]
    basis = ["폴더 이름"]      # 감사건의 앞선 basis 를 이어 쓰지 않는다(다시 돌릴 때 쌓이지 않게)
    pairs = ([("전기 감사보고서", rep.get("company"))] if rep else []) + ([("DART", drep.get("company"))] if drep else [])
    for label, x in pairs:
        if not x:
            continue
        if corp_name(x) == corp_name(name):          # 「주식회사 · (주) · ㈜」 표기와 공백 차이는 같은 회사로 본다
            basis.append(f"{label} 회사명 일치")
        else:
            warnings.append(f"감사 폴더 회사명({name})과 전기 감사보고서 · DART 회사명({x})이 다릅니다. 같은 회사인지 확인해 주십시오")
    industry, industry_basis = None, None
    if dart_part and dart_part.get("industry"):
        industry, industry_basis = dart_part["industry"], "DART 기업개황 업종명"
    else:
        found = _industry_from_docs(d)
        if found:
            industry, industry_basis = found[0], f"세무신고 업종 칸({found[1]})"
    corp_code = dart_part.get("corp_code") if dart_part else None
    if corp_code:
        basis.append("DART 고유번호")
    if industry_basis:
        basis.append(industry_basis.split("(")[0])
    company = {"name": name, "dart_corp_code": corp_code, "industry": industry, "industry_basis": industry_basis,
               "framework": (prior_report or {}).get("framework") or (dart_part or {}).get("framework"),
               "basis": " · ".join(dict.fromkeys(basis)), "overview": (dart_part or {}).get("overview")}

    need = []
    if prior_report:
        need += [f"{prior_report['path']}: {x}" for x in prior_report["need_ai"]]
    else:
        need.append("전기 감사보고서가 받은자료에 없습니다(초도감사면 DART 로 받습니다)")
    if pw:
        need += pw["need_ai"]
    elif case.get("engagement_kind") == "계속감사":
        need.append("전기조서가 받은자료에 없습니다")
    if not industry:
        need.append("업종을 찾지 못했습니다(DART 기업개황 · 세무신고 업종 칸 없음)")

    return {"prior_end": prior_end, "prior_report": prior_report, "dart": dart_part, "prior_workpapers": pw,
            "prior_fs": prior_fs, "company": company, "need_ai": need, "warnings": warnings, "saved_dart": saved_dart,
            "term": term_of(d, prior_report, dart_part)}


def term_of(d: Path, prior_report: dict | None, dart_part: dict | None) -> dict | None:
    """당기 기수(2026-10-06). 작년 감사 폴더 감사건의 기수 → 전기 감사보고서 · DART 의 「제 N 기」 순서로 전기 기수를 찾고
    당기 = 전기 + 1 로 둔다. 못 찾으면 None 이고, 회계사가 8600 「회사 정보」 시트에 적는다(엔진은 추측하지 않는다)."""
    pe = prior.prior_engagement(d)
    if pe is not None:
        t = (_idx(pe, "감사건.json", {}) or {}).get("term") or {}
        if isinstance(t, dict) and isinstance(t.get("current"), int):
            return {"current": t["current"] + 1, "prior": t["current"], "basis": f"작년 감사 폴더 감사건 「제 {t['current']} 기」"}
    for src, label in ((prior_report, "전기 감사보고서"), (dart_part, "DART 감사보고서")):
        n = (src or {}).get("term")
        if isinstance(n, int):
            return {"current": n + 1, "prior": n, "basis": f"{label}({(src or {}).get('path')}) 「제 {n} 기」"}
    return None


# ---------------------------------------------------------------- AI · 사람 판독
NEED_LABEL = {"auditor": "감사인", "opinion": "감사의견", "industry": "업종", "framework": "회계기준", "report_date": "감사보고서일"}


def _merge_ai(data: dict) -> None:
    """엔진이 비워 둔 글 항목만 AI · 사람 판독으로 채운다. 엔진 값은 덮지 않는다. 근거는 basis 에 남긴다."""
    for x in data.get("ai", []):
        sect, key = x["field"].split(".")
        part = data.get(sect)
        if not isinstance(part, dict):
            continue
        if sect == "prior_report":
            b = part.setdefault("basis", {})
            if part.get(key) is not None and b.get(key) == "엔진":
                continue
            part[key] = x["value"]
            b[key] = x["basis"]
            label = NEED_LABEL.get(key)
            if label:
                part["need_ai"] = [n for n in part.get("need_ai", []) if not n.startswith(label)]
        else:
            if part.get(key) is not None and not str(part.get("industry_basis") or "").startswith(("AI", "사람")):
                continue
            part[key] = x["value"]
            part["industry_basis"] = f"{x['basis']} 판독({x.get('source_page') or '쪽 없음'})"
        label = NEED_LABEL.get(key)
        if label:
            data["need_ai"] = [n for n in data.get("need_ai", []) if label not in n]


def apply(d: Path, patch: dict, now: str | None = None) -> dict:
    """{"fields": [{"field": "prior_report.auditor", "value": "...", "basis": "AI", "source_page": "1쪽"}]}
    글 항목만 받는다. 금액으로 보이는 숫자가 있으면 아무것도 쓰지 않고 거부한다."""
    d = Path(d)
    data = _idx(d, START_FILE)
    if data is None:
        raise IntakeError("시작자료.json 이 없습니다. 먼저 python -m audit start <감사 폴더> 를 돌려 주십시오")
    items = patch.get("fields") if isinstance(patch, dict) else None
    if not isinstance(items, list) or not items:
        raise index.PatchRejected(['patch 는 {"fields": [...]} 꼴이고 항목이 하나 이상 있어야 합니다'])
    probs = []
    for i, x in enumerate(items, 1):
        where = f"{i}번째 항목"
        if not isinstance(x, dict) or x.get("field") not in AI_FIELDS:
            probs.append(f"{where}: 넣을 수 있는 항목은 {', '.join(AI_FIELDS)} 입니다")
            continue
        if x.get("basis") not in ("AI", "사람"):
            probs.append(f"{where}: basis 는 AI 또는 사람 이어야 합니다")
        t = AI_FIELDS[x["field"]]
        v = x.get("value")
        if not isinstance(v, t):
            probs.append(f"{where}: {x['field']} 값의 꼴이 맞지 않습니다")
            continue
        if x["field"] == "prior_report.opinion" and v not in dart.OPINIONS:
            probs.append(f"{where}: 의견은 {' · '.join(dart.OPINIONS)} 가운데 하나입니다")
        if x["field"] == "prior_report.report_date" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            probs.append(f"{where}: 날짜는 2025-03-14 꼴입니다")
        if isinstance(v, str) and x["field"] != "prior_report.report_date":
            amt = index.amount_in_text(v)
            if amt:
                probs.append(f"{where}: 금액으로 보이는 숫자 '{amt}' 가 있습니다. 숫자는 엔진 판독만 씁니다")
    if probs:
        raise index.PatchRejected(probs)
    keep = {x["field"]: x for x in data.get("ai", [])}
    for x in items:
        keep[x["field"]] = {"field": x["field"], "value": x["value"], "basis": x["basis"], "source_page": x.get("source_page")}
    data["ai"] = list(keep.values())
    folder._write(d / "_색인" / START_FILE, data)
    return build(d, now=now, reuse=True)


# ---------------------------------------------------------------- 색인 쓰기
def _doc_facts(d: Path, data: dict) -> list[str]:
    """시작 자료로 읽은 받은자료에 엔진 사실을 붙인다(AI 사실은 둔다)."""
    docs = _idx(d, "문서목록.json")
    by = {f["path"]: f for f in docs["files"]}
    touched = []

    def put(rel, facts):
        f = by.get(rel)
        if not f or not facts:
            return
        f["facts"] = facts + [x for x in f.get("facts") or [] if x.get("basis") != "엔진"]
        touched.append(rel)

    pr = data.get("prior_report")
    if pr and pr.get("basis"):
        fs = []
        for k, label in (("opinion", "전기 감사의견"), ("auditor", "전기 감사인"), ("report_date", "전기 감사보고서일")):
            if pr.get(k) and pr["basis"].get(k) == "엔진":
                fs.append({"text": f"{label} {pr[k]}", "basis": "엔진", "value": None})
        if pr.get("emphasis") is not None and pr["basis"].get("emphasis") == "엔진":
            fs.append({"text": "강조사항 문단 " + ("있음" if pr["emphasis"] else "없음"), "basis": "엔진", "value": None})
        for k, v in (pr.get("figures") or {}).items():
            f = {"text": f"{k} {_won(v)}", "basis": "엔진", "value": v}
            if pr.get("figure_pages", {}).get(k):
                f["source_page"] = pr["figure_pages"][k]
            fs.append(f)
        put(pr["path"], fs)
    dp = data.get("dart")
    if dp and dp.get("path") != (pr or {}).get("path"):
        fs = [{"text": f"DART 감사의견 {dp['opinion']}", "basis": "엔진", "value": None}] if dp.get("opinion") else []
        if dp.get("auditor"):
            fs.append({"text": f"DART 감사인 {dp['auditor']}", "basis": "엔진", "value": None})
        for k in ("자산총계", "매출액", "당기순이익"):
            v = (dp.get("key_bs") or {}).get(k, (dp.get("key_is") or {}).get(k))
            if v is not None:
                fs.append({"text": f"DART {k} {_won(v)}", "basis": "엔진", "value": v})
        put(dp["path"], fs)
    pw = data.get("prior_workpapers")
    if pw:
        mat = pw.get("materiality") or {}
        names = {"overall": "전기 재무제표 전체 중요성", "performance": "전기 수행중요성", "trivial": "전기 명백히 사소한 금액"}
        by_file: dict[str, list] = {}
        for k, v in mat.items():
            by_file.setdefault(v["path"], []).append({"text": f"{names[k]} {_won(v['amount'])}", "basis": "엔진", "value": v["amount"],
                                                      "source_page": f"{v['sheet']} 시트"})
        for key, label in (("adjustments", "전기 수정사항"), ("unadjusted", "전기 미수정 왜곡")):
            for p in {a["path"] for a in pw.get(key, [])}:
                rows = [a for a in pw[key] if a["path"] == p]
                total = sum(a["amount"] for a in rows)
                by_file.setdefault(p, []).append({"text": f"{label} {len(rows)}건, 합계 {_won(total)}", "basis": "엔진", "value": total})
        if pw["wtb"].get("path"):
            n = len(pw["wtb"]["accounts"])
            by_file.setdefault(pw["wtb"]["path"], []).append({"text": f"전기 정산표 계정 {n}개 · 표시과목 {len(pw['wtb']['lines'])}개", "basis": "엔진", "value": n})
        for p, fs in by_file.items():
            put(p, fs)
    folder._write(d / "_색인" / "문서목록.json", docs)
    return touched


def _update_case(d: Path, data: dict, at: str, wrote: list[str], read: list[str]) -> None:
    case = _idx(d, "감사건.json")
    c = data["company"]
    basis = c.get("basis") or "폴더 이름"
    if str(c.get("industry_basis") or "").startswith(("AI", "사람")):
        basis += f" · 업종 {c['industry_basis'].split()[0]} 판독" + (" (확인 전)" if c["industry_basis"].startswith("AI") else "")
    case["company"] = {"name": c["name"], "dart_corp_code": c.get("dart_corp_code"), "industry": c.get("industry"), "basis": basis}
    if data.get("term"):                 # 기수는 찾았을 때만 쓴다. 못 찾았으면 앞선 값(회계사가 고친 값 포함)을 둔다
        case["term"] = data["term"]
    got = [x for x, ok in (("전기 감사보고서", data.get("prior_report")), ("전기조서", data.get("prior_workpapers")),
                           ("DART", data.get("dart")), ("전기 재무제표", data.get("prior_fs"))) if ok]
    what = f"시작 자료를 읽었습니다({' · '.join(got) or '읽은 자료 없음'}). {prior.label(data.get('prior_source'))}. 조서 {len(wrote)}개를 만들었습니다"
    if data.get("need_ai"):
        what += f". 판독이 필요한 것 {len(data['need_ai'])}건"
    case["log"].insert(0, {"at": at, "who": "시작 자료", "what": what, "read": read, "wrote": wrote + ["_색인/시작자료.json", "_색인/감사건.json"]})
    folder._write(d / "_색인" / "감사건.json", case)


# ---------------------------------------------------------------- 조서
def _may_write(path: Path, made: dict, force: bool) -> bool:
    """없거나, 엔진이 만든 뒤 아무도 손대지 않은 파일(해시가 같음)이면 다시 쓴다. force 면 언제나."""
    if force or not path.exists():
        return True
    return made.get(path.name) == _sha(path)


def _rel(d: Path, p: Path) -> str:
    return p.relative_to(d).as_posix()


def _company_rows(b: wp.Book, ws, row: int, data: dict, case: dict) -> int:
    c = data["company"]
    pr = data.get("prior_report") or {}
    row = b.head(ws, row, ["항목", "내용", "출처"])
    rows = [("회사명", c["name"], c.get("basis")),
            ("DART 고유번호", c.get("dart_corp_code") or "없음", "DART" if c.get("dart_corp_code") else "DART 링크를 주면 채웁니다"),
            ("업종", c.get("industry") or "판독 필요", c.get("industry_basis") or ""),
            ("결산일", case.get("fiscal_year_end"), "감사건"),
            ("감사 구분", case.get("engagement_kind"), "감사건"),
            ("전기 회계기준", c.get("framework") or "판독 필요", pr.get("path") or ""),
            ("회사제시 판", case.get("company_version") or "없음", "회사제시 폴더")]
    for k, v, src in rows:
        b.put(ws, row, 1, k)
        b.put(ws, row, 2, v, flag=v in ("판독 필요",), wrap=True)
        b.put(ws, row, 3, src, wrap=True, color=wp.MUTED)
        row += 1
    return row + 1


def _prior_summary_rows(b: wp.Book, ws, row: int, data: dict) -> int:
    fs = data.get("prior_fs")
    row = b.section(ws, row, "전기 재무 요약 (엔진, 표시 금액)")
    row = b.head(ws, row, ["과목", "전기 금액", "출처"])
    allf = {**(fs or {}).get("bs", {}), **(fs or {}).get("is", {})}
    pr_f = (data.get("prior_report") or {}).get("figures") or {}
    for k in KEY_FIGS:
        v = allf.get(k, pr_f.get(k))
        src = (fs or {}).get("source") if k in allf else ((data.get("prior_report") or {}).get("path") if k in pr_f else "없음")
        b.put(ws, row, 1, k)
        b.put(ws, row, 2, v, fmt=wp.WON)
        b.put(ws, row, 3, src, color=wp.MUTED, wrap=True)
        row += 1
    return row + 1


def _checklist(b: wp.Book, ws, row: int, no: str) -> int:
    items = json.loads((Path(__file__).resolve().parent.parent / "tools" / "proto_data" / "kicpa_checklists.json").read_text(encoding="utf-8"))[no]
    row = b.section(ws, row, items["title"] + " (한공회 감사조서 서식 문항)")
    row = b.head(ws, row, ["항목", "조서번호", "수행자", "수행 결과·비고"])
    for n, (text, ref) in enumerate(items["items"], 1):
        numbered = bool(re.match(r"^\d+[.,]", text.strip()))
        b.put(ws, row, 1, text, wrap=True, bold=numbered)
        b.put(ws, row, 2, ref or None)
        if numbered:
            b.ask(ws, row, 3, f"점검 {n} 수행자", rule="optional")
            b.ask(ws, row, 4, f"점검 {n} 수행 결과", rule="required")
        ws.row_dimensions[row].height = 15 * max(1, min(6, len(text) // 60 + 1))
        row += 1
    return row + 1


def make_1100(d: Path, data: dict, case: dict) -> wp.Book:
    b = wp.Book(d, "1100", standard="감사기준서 210 · 220 · 한공회 서식 1100")
    ws, row = b.sheet("1100", [34, 44, 40, 40])
    row = b.section(ws, row, "회사 개황 (엔진)")
    row = _company_rows(b, ws, row, data, case)
    row = _prior_summary_rows(b, ws, row, data)
    pr = data.get("prior_report") or {}
    row = b.section(ws, row, "전기 감사 (전기 감사보고서 판독)")
    row = b.head(ws, row, ["항목", "내용", "근거"])
    for k, label in (("auditor", "전기 감사인"), ("opinion", "전기 감사의견"), ("emphasis", "강조사항 문단"),
                     ("going_concern", "계속기업 관련 중요한 불확실성"), ("report_date", "전기 감사보고서일")):
        v = pr.get(k)
        if isinstance(v, bool):
            v = "있음" if v else "없음"
        basis = (pr.get("basis") or {}).get(k, "엔진" if v is not None else "판독 필요")
        b.put(ws, row, 1, label)
        b.put(ws, row, 2, v if v is not None else "판독 필요", flag=v is None, wrap=True)
        b.put(ws, row, 3, f"{basis} · {pr.get('path') or '전기 감사보고서 없음'}", color=wp.MUTED, wrap=True)
        row += 1
    dp = data.get("dart")
    if dp:
        for k, label in (("auditor", "DART 감사인"), ("opinion", "DART 감사의견")):
            b.put(ws, row, 1, label)
            b.put(ws, row, 2, dp.get(k) or "판독 필요", flag=not dp.get(k))
            b.put(ws, row, 3, f"엔진 · {dp.get('source') or dp['path']}", color=wp.MUTED, wrap=True)
            row += 1
        for x in dp.get("diffs_vs_report") or []:
            b.put(ws, row, 1, "감사보고서와 DART 차이")
            b.put(ws, row, 2, x, flag=True, wrap=True)
            row += 1
    if pr.get("opinion") not in (None, "적정") or pr.get("emphasis") or pr.get("going_concern"):
        ws.cell(row, 1, "전기 의견이 적정이 아니거나 강조사항 · 계속기업 불확실성이 있습니다. 계약 수락 판단에 반영해 주십시오").font = wp.font(size=9, color="B3261E")
        row += 1
    row += 1
    row = b.section(ws, row, "감사인 교체")
    row = b.head(ws, row, ["항목", "내용", "근거"])
    b.put(ws, row, 1, "감사 구분(감사건)")
    b.put(ws, row, 2, case.get("engagement_kind"))
    b.put(ws, row, 3, "초도감사면 전기 감사인과 다른 감사인입니다", color=wp.MUTED, wrap=True)
    row += 1
    b.put(ws, row, 1, "전기 감사인")
    b.put(ws, row, 2, pr.get("auditor") or "판독 필요", flag=not pr.get("auditor"))
    row += 1
    b.put(ws, row, 1, "감사인 교체 여부(예 · 아니오)")
    b.ask(ws, row, 2, "감사인 교체 여부", rule="yesno")
    row += 1
    b.put(ws, row, 1, "교체라면 전임 감사인과 커뮤니케이션 결과")
    b.ask(ws, row, 2, "전임 감사인 커뮤니케이션", rule="optional")
    row += 2
    pw = data.get("prior_workpapers") or {}
    row = b.section(ws, row, "전기 수정사항 · 미수정 왜곡 (전기조서)")
    row = b.head(ws, row, ["구분", "내용", "금액", "출처"])
    for key, label in (("adjustments", "수정사항"), ("unadjusted", "미수정 왜곡")):
        for a in pw.get(key, []):
            b.put(ws, row, 1, label)
            b.put(ws, row, 2, a["text"], wrap=True)
            b.put(ws, row, 3, a["amount"], fmt=wp.WON)
            b.put(ws, row, 4, a["path"], color=wp.MUTED, wrap=True)
            row += 1
    if not pw.get("adjustments") and not pw.get("unadjusted"):
        b.put(ws, row, 1, "없음" if pw else "전기조서 없음")
        row += 1
    row += 1
    row = _checklist(b, ws, row, "1100")
    row = b.section(ws, row, "계약 수락 · 유지 판단 (감사인)")
    b.put(ws, row, 1, "위험 요인 메모")
    b.ask(ws, row, 2, "위험 요인 메모", rule="optional")
    row += 1
    b.conclusion(ws, row, 4, label="계약 수락 · 유지 결론")
    return b


def make_2100(d: Path, data: dict, case: dict) -> wp.Book:
    b = wp.Book(d, "2100", standard="감사기준서 315 · 한공회 서식 2100")
    fy = _fy_end(case)
    ws, row = b.sheet("기업이해", [30, 44, 40, 36])
    row = b.section(ws, row, "회사 개황 (엔진)")
    row = _company_rows(b, ws, row, data, case)
    ov = data["company"].get("overview")
    if ov:
        row = b.section(ws, row, "회사의 개요 (DART 주석 1 원문 앞부분)")
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=4)
        b.put(ws, row, 1, ov, wrap=True)
        row += 4
    row = _prior_summary_rows(b, ws, row, data)
    from . import priorwp

    was = {x["label"]: x["value"] for x in (priorwp.general_inputs(d, "2100") or {}).get("inputs", []) if x["sheet"] == "기업이해"}
    row = b.section(ws, row, "사업 이해 (감사인)" + (". 셋째 칸은 작년 조서에 적은 내용(참고)" if was else ""))
    for label in ("사업 내용 · 산업 환경", "규제 · 외부 요인", "회계정책 선택과 변경"):
        b.put(ws, row, 1, label)
        b.ask(ws, row, 2, label, rule="required")
        if was.get(label):
            b.put(ws, row, 3, was[label], wrap=True, color=wp.MUTED)
        row += 1
    row += 1
    row = _checklist(b, ws, row, "2100")
    b.conclusion(ws, row, 4, label="기업 이해 결론")

    tp = data["parties"]
    ws2, row = b.sheet("주요 거래처", [8, 28, 20, 12])
    ws2.cell(6, 1, f"출처: {tp.get('path') or '당기 분개장 없음'} (매출처는 더존 401~430 대변 순액, 매입처는 251~252 대변 합)").font = wp.font(size=9, color=wp.MUTED)
    for key, label, total in (("sales", "매출처 상위 10", "sales_total"), ("purchases", "매입처 상위 10", "purchases_total")):
        row = b.section(ws2, row, f"{label} (전체 {tp.get(key + '_count', 0)}곳)")
        row = b.head(ws2, row, ["순위", "거래처", "금액", "비중"])
        first = row
        for n, x in enumerate(tp.get(key, []), 1):
            b.put(ws2, row, 1, n)
            b.put(ws2, row, 2, x["party"])
            b.put(ws2, row, 3, x["amount"], fmt=wp.WON)
            b.put(ws2, row, 4, x["share"], fmt=wp.PCT)
            row += 1
        b.put(ws2, row, 2, f"전체 합계({label.split()[0]})", bold=True)
        b.put(ws2, row, 3, tp.get(total) or 0, fmt=wp.WON, bold=True)
        if tp.get(key):
            b.verify(f"{label} 합계 재계산", f"=SUM('주요 거래처'!C{first}:C{row - 1})", sum(x["amount"] for x in tp[key]))
        row += 2
    b.put(ws2, row, 1, "메모")
    b.ask(ws2, row, 2, "주요 거래처 메모(특수관계 · 집중도)", rule="optional")

    _carry_sheet(b, d, data)

    ch = data["changes"]
    ws3, row = b.sheet("전기 대비 변동", [30, 18, 18, 18, 10, 10, 36])
    thr = ch.get("threshold")
    ws3.cell(6, 1, ("큰 변동: |증감| ≥ " + (f"{thr:,}원(전기 수행중요성)" if thr else "기준 금액 없음") +
                    " 이고 |증감률| ≥ 10%. 전기 과목 없음: 새 과목이거나 전기와 이름이 다른 과목(매핑 확인). "
                    "당기는 회사제시 시산표(이익잉여금은 당기순이익을 더한 마감 후)")).font = wp.font(size=9, color=wp.MUTED)
    row = b.head(ws3, row, ["과목", f"전기({data['prior_end']})", f"당기({case['fiscal_year_end']}, 회사제시)", "증감", "증감률", "큰 변동", "회사 설명 · 감사인 메모"])
    for x in ch["rows"]:
        b.put(ws3, row, 1, x["name"], bold=x["total"])
        b.put(ws3, row, 2, x["prior"], fmt=wp.WON)
        b.put(ws3, row, 3, x["current"], fmt=wp.WON)
        b.put(ws3, row, 4, f"=C{row}-N(B{row})", fmt=wp.WON)
        b.put(ws3, row, 5, f'=IF(N(B{row})=0,"",D{row}/ABS(B{row}))', fmt=wp.PCT)
        mark = "큰 변동" if x["big"] else ("전기 과목 없음" if x["prior_missing"] and x["current"] else "")
        b.put(ws3, row, 6, mark, flag=bool(mark))
        if mark:
            b.ask(ws3, row, 7, f"{x['name']} 변동 설명", rule="optional")
        b.verify(f"{x['name']} 증감", f"='전기 대비 변동'!D{row}", x["diff"])
        row += 1
    row += 1
    if ch.get("unmatched_prior"):
        ws3.cell(row, 1, "전기 과목 가운데 당기 표시과목과 이름이 맞지 않은 것(계정 매핑은 분석적 검토 단계에서): " + ", ".join(ch["unmatched_prior"])).font = wp.font(size=9, color="B3261E")
        row += 1
    if ch.get("unmatched_current"):
        ws3.cell(row, 1, "당기 표시과목 가운데 전기 과목에 없는 것: " + ", ".join(ch["unmatched_current"])).font = wp.font(size=9, color="B3261E")
    priorwp.reference_sheet(b, d, "2100")       # 작년 2100 에 회계사가 적은 내용(참고 시트, 2026-10-12)
    return b


CARRY_SHEET = "전기에서 이어진 사항"


def _carry_sheet(b: wp.Book, d: Path, data: dict) -> None:
    """2100 「전기에서 이어진 사항」(R7). 작년 감사 폴더 · 전임 전기조서 · 전기 감사보고서에서 이어받을 사항과
    회계사 판단 기입란(당기 이슈로 볼지). 엔진은 당기 이슈로 정하지 않는다. 앞서 적은 판단은 같은 항목 이름으로 이어받는다."""
    from .report.base import prev_inputs

    keep = prev_inputs(wp.path_for(d, "2100"))
    ps = data.get("prior_source") or {}
    ws, row = b.sheet(CARRY_SHEET, [6, 22, 56, 36, 16, 40])
    ws.cell(6, 1, prior.label(ps) + ". 미수정 왜곡 · 계속기업 · 강조사항 · 의견 변형 · 열린 채 종결된 이슈를 옮긴다(감사기준서 510 문단 9 · 450 문단 11(b))"
            ).font = wp.font(size=9, color=wp.MUTED)
    row = b.head(ws, row, ["번호", "구분", "내용", "출처", "당기 이슈로 볼지", "판단 근거"])
    carry = ps.get("carry") or []
    if not carry:
        b.put(ws, row, 3, "전기 자료 없음" if not ps or ps.get("situation") == prior.SITUATIONS[2] and not ps.get("opinion")
              else "전기에서 이어진 사항 없음", color=wp.MUTED)
        return
    for k, c in enumerate(carry, 1):
        b.put(ws, row, 1, k)
        b.put(ws, row, 2, c["what"], wrap=True)
        b.put(ws, row, 3, c["text"], wrap=True)
        b.put(ws, row, 4, c["source"], wrap=True, color=wp.MUTED)
        for col, lab, rule in ((5, f"이어진 사항 {k} 당기 이슈 여부", "yesno"), (6, f"이어진 사항 {k} 판단 근거", "optional")):
            b.ask(ws, row, col, lab, rule=rule, value=keep.get((CARRY_SHEET, lab)))
        ws.row_dimensions[row].height = 30
        row += 1
    ws.cell(row + 1, 1, "엔진은 이어받은 사항으로 이슈를 만들지 않습니다. 「예」로 판단한 사항은 이슈 대장과 대응 조서에 반영해 주십시오").font = wp.font(size=9, color=wp.MUTED)


def make_pbc(d: Path, data: dict) -> tuple[wp.Book, list[dict]]:
    """요청자료 조서. 시트 네 개는 실제 초도감사 요청자료 양식(audit/pbcform.py)이다. 받음 · 빠짐은 엔진 18건(pbc_match.yaml)만 판정한다."""
    from . import pbcform

    rows = pbc_status(d)
    b = wp.Book(d, "PBC", standard="감사 시작 전 요청자료 목록 · 받음은 폴더에 파일이 있는지로 엔진이 정한다")
    pbcform.write(b, d, rows)
    ws = b.wb.worksheets[0]
    row = ws.max_row + 3
    # 이슈 대응 자료와 분석적 검토가 대사하지 못해 요청하는 자료(분석.json requests, 2026-10-18)를 함께 적는다
    issue = [m for m in data.get("missing", []) if m.get("issue") or m.get("source") == "분석" or m.get("detail")]
    ws.cell(row, 2, "이슈 대응에 필요한 추가 자료 (지휘가 이슈.json · 분석.json 에서 정함)").font = wp.font(size=11, bold=True, color=wp.NAVY)
    row = b.head(ws, row + 1, ["이슈", "구분", "요청자료명", "요청세부내역", "", "요청한 날", "상태"], col=2)
    for m in issue:
        if m.get("issue"):
            b.put(ws, row, 2, f"이슈 {m['issue']}")
            b.put(ws, row, 4, "이슈 대응 자료")
            b.put(ws, row, 5, m["item"], wrap=True)
        else:
            b.put(ws, row, 2, "분석")
            b.put(ws, row, 4, "분석적 검토 대사 자료" if m.get("source") == "분석" else m["item"], wrap=True)
            b.put(ws, row, 5, m["item"] if m.get("source") == "분석" else "; ".join(m.get("detail") or []), wrap=True)
        b.put(ws, row, 3, m.get("workpaper"))
        b.put(ws, row, 7, m.get("requested_at") or "요청 전")
        b.put(ws, row, 8, "빠짐", flag=True)
        row += 1
    if not issue:
        b.put(ws, row, 5, "없음")
    return b, rows


def pbc_status(d: Path) -> list[dict]:
    """요청자료 18건(pbc_match.yaml 번호)의 받음 · 빠짐. 판정은 지휘(director)가 빠진 자료를 셀 때와 같은 규칙이다."""
    c = director._Case(d)
    stage = director._stage_of(c)
    nobal = c.no_balance()
    out = []
    for e in c.match.get("pbc", []):
        hits = _hits(e, c.files)
        ok, lack = c.judge(e)
        due_later = bool(e.get("from_stage")) and director.STAGES.index(stage) < director.STAGES.index(e["from_stage"])
        status = "받음" if ok else (f"{e['from_stage']} 단계에 받음" if due_later else "빠짐")
        if not ok and e.get("area") and str(e["area"]) in nobal:
            status = director.NO_BALANCE      # 시산표 · 재무제표에 그 영역 잔액이 당기 · 전기 모두 없다(2026-10-18)
        note = f"빠진 것: {' · '.join(lack)}" if lack and not ok else None
        out.append({"no": e["row"], "category": PBC_CATEGORY.get(e["row"], "기타"), "item": e["item"],
                    "workpaper": str(e.get("workpaper") or ""), "status": status, "files": hits, "note": note})
    return out


def _hits(spec: dict, files: list[dict]) -> list[str]:
    if spec.get("any_of"):
        return sorted({p for s in spec["any_of"] for p in _hits(s, files)})
    kinds = spec.get("kinds") or []
    kws = spec.get("keywords") or []
    return sorted(f["path"] for f in files if (not kinds or f["kind"] in kinds) and (not kws or any(k in director._hay(f) for k in kws)))


def make_1200(d: Path, data: dict, case: dict, path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt

    c = data["company"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10)
    doc.core_properties.author = "audit 엔진"
    doc.core_properties.comments = "audit 엔진 초안. 노란 칸은 감사인 기입란"
    doc.core_properties.title = "1200 감사업무조건합의"

    def blank(p, text):
        r = p.add_run(f"[기입란: {text}]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return r

    doc.add_heading("1200 감사업무조건합의 (감사계약서 초안)", level=1)
    doc.add_paragraph("감사기준서 210 감사업무 조건의 합의. 엔진이 감사 폴더에서 채운 초안이며, 노란 칸은 감사인이 정해 적는다.")
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in (("회사", c["name"]), ("감사 대상 재무제표", f"{case['fiscal_year_end']} 로 끝나는 보고기간의 재무제표"),
                 ("적용 회계기준(전기 감사보고서 기준)", c.get("framework") or "판독 필요"),
                 ("감사 구분", case.get("engagement_kind")), ("전기 감사인", (data.get("prior_report") or {}).get("auditor") or "판독 필요")):
        row = t.add_row().cells
        row[0].text, row[1].text = k, str(v)
    doc.add_heading("1. 감사의 목적과 범위", level=2)
    doc.add_paragraph("감사인은 회계감사기준에 따라 회사의 재무제표를 감사하고, 재무제표가 중요성의 관점에서 적용 회계기준에 따라 "
                      "공정하게 표시되어 있는지에 대한 의견을 감사보고서로 표명한다.")
    doc.add_heading("2. 경영진의 책임", level=2)
    for s in ("적용 회계기준에 따라 재무제표를 작성하고 공정하게 표시할 책임",
              "부정이나 오류로 인한 중요한 왜곡표시가 없는 재무제표 작성에 필요한 내부통제를 설계하고 운영할 책임",
              "감사에 필요한 기록 · 문서와 그 밖의 자료를 제공하고, 감사인이 필요하다고 정한 사람에게 제한 없이 접근하게 할 책임",
              "감사 종료 때 서면진술서를 제공할 책임"):
        doc.add_paragraph(s, style="List Bullet")
    doc.add_heading("3. 감사인의 책임", level=2)
    doc.add_paragraph("감사인은 직업적 회의주의를 유지하며 감사를 계획하고 수행한다. 감사의 고유한 한계로 중요한 왜곡표시가 "
                      "발견되지 않을 위험이 있다. 감사인은 직무상 알게 된 회사의 비밀을 지킨다.")
    doc.add_heading("4. 보수와 일정", level=2)
    blank(doc.add_paragraph(), "감사 보수, 지급 시기, 현장 감사 일정, 감사보고서 제출 예정일")
    doc.add_heading("5. 자료 처리와 AI 도구 사용 고지", level=2)
    doc.add_paragraph("감사 자료의 처리 장소와 보존, AI 도구를 쓰는지와 그 범위는 감사인이 정해 아래에 적는다. 엔진은 문구를 쓰지 않는다.")
    blank(doc.add_paragraph(), "AI 도구 사용 고지 문구(감사인이 작성)")
    doc.add_heading("6. 그 밖의 조건", level=2)
    blank(doc.add_paragraph(), "감사 범위 제한, 다른 감사인 이용, 내부감사 활용 등")
    doc.add_heading("서명", level=2)
    s = doc.add_table(rows=2, cols=2)
    s.style = "Table Grid"
    s.cell(0, 0).text, s.cell(0, 1).text = f"회사: {c['name']} 대표이사", "감사인"
    blank(s.cell(1, 0).paragraphs[0], "서명 · 날짜")
    blank(s.cell(1, 1).paragraphs[0], "서명 · 날짜")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------- 요청 메일
def write_mail(d: Path, at: date | None = None) -> Path | None:
    """문서목록.missing 으로 요청 메일 초안을 쓴다. 빠진 자료가 없으면 쓰지 않는다. 보내기는 회계사가 한다."""
    d = Path(d)
    docs = _idx(d, "문서목록.json")
    case = _idx(d, "감사건.json")
    miss = docs.get("missing", [])
    if not miss:
        return None
    at = at or date.today()
    fy = case["fiscal_year_end"]
    name = case["company"]["name"]
    new = [m for m in miss if not m.get("requested_at")]
    again = [m for m in miss if m.get("requested_at")]
    # 제목 · 첫 문장은 감사 폴더 이름(「2025 기말감사」 · 「2026 중간감사」)을 따른다. 중간감사 폴더에 「기말감사」 · 「결산」 이 나갔다(2026-10-19)
    eng = d.name if re.fullmatch(r"\d{4} \S*감사|\d{4} \S*검토", d.name) else f"{fy[:4]} 기말감사"
    when = f"{fy[:4]}년 {int(fy[5:7])}월 {int(fy[8:10])}일"
    purpose = f"{when} 기준 중간감사" if "중간" in eng else (f"{when} 기준 반기검토" if "반기" in eng else f"{when} 결산 외부감사")
    lines = ["[초안] 회계사가 확인한 뒤 보냅니다. 받는 사람과 서명은 회계사가 채웁니다.", "",
             f"제목: [{name} {eng}] 감사 자료 요청 ({len(miss)}건)", "",
             f"{name} 담당자님께", "",
             f"{purpose}를 위해 아래 자료를 요청드립니다.",
             "파일로 보내 주시면 감사 폴더에 정리하겠습니다. 이미 보내신 자료가 목록에 있으면 보내신 날을 알려 주시면 확인하겠습니다.", ""]
    n = 0
    numbered = []
    for title, group in (("새로 요청드리는 자료", new), ("앞서 요청드렸으나 아직 받지 못한 자료", again)):
        if not group:
            continue
        lines.append(title)
        for m in group:
            n += 1
            numbered.append((n, m))
            tail = f" (요청한 날 {m['requested_at']})" if m.get("requested_at") else ""
            lines.append(f"{n}. {m['item']}{tail}")
            for x in m.get("detail") or []:          # 분석적 검토가 붙인 구체 문구(2026-10-18)
                lines.append(f"   - {x}")
        lines.append("")
    lines += ["자료 준비에 시간이 걸리는 항목은 예정일을 알려 주시면 일정에 반영하겠습니다.", "감사합니다.", "",
              "감사인 드림", "", f"(감사 폴더 _색인/문서목록.json 의 빠진 자료 {len(miss)}건으로 만든 초안, {at.isoformat()})"]
    # 스태프 배정이 있으면 자료를 받을 담당자별로 번호를 묶어 둔다(2026-10-12). 감사반이 보는 구역이라 보내기 전에 지운다
    from . import assign as _assign

    who = _assign.of(d)
    if who:
        by: dict[str, list] = {}
        for k, m in numbered:
            by.setdefault((who.get(str(m.get("workpaper") or "")) or {}).get("to") or "미배정", []).append((k, m))
        lines += ["", "[감사반 확인용: 담당자별 받을 자료. 보내기 전에 이 아래를 지웁니다]"]
        for name in sorted(by, key=lambda x: (x == "미배정", x)):
            lines.append(f"{name}: " + ", ".join(f"{k}번" + (f"(조서 {m['workpaper']})" if m.get("workpaper") else "") for k, m in by[name]))
    p = d / "1000 감사계약" / f"{MAIL_PREFIX}{at.isoformat()}.txt"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return p


def mark_requested(d: Path, at: str, now: str | None = None) -> int:
    """회계사가 요청 메일을 보낸 뒤 빠진 자료에 요청한 날을 적는다. 이미 적힌 것은 두고, 적은 수를 돌려준다."""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", at):
        raise IntakeError(f"보낸 날은 2026-01-21 꼴로 적어 주십시오: {at}")
    d = Path(d)
    docs = _idx(d, "문서목록.json")
    n = 0
    for m in docs.get("missing", []):
        if not m.get("requested_at"):
            m["requested_at"] = at
            n += 1
    folder._write(d / "_색인" / "문서목록.json", docs)
    case = _idx(d, "감사건.json")
    case["log"].insert(0, {"at": now or _now(), "who": "시작 자료", "what": f"요청 메일을 {at} 에 보냈다고 빠진 자료 {n}건에 적었습니다",
                           "wrote": ["_색인/문서목록.json"]})
    folder._write(d / "_색인" / "감사건.json", case)
    folder.validate_index(d)
    return n


def _carry_people(d: Path, ps: dict) -> str | None:
    """작년 감사 폴더 1000 감사계약/감사반인원.xlsx 를 당기 폴더에 없을 때만 복사한다. 복사한 상대경로 또는 None."""
    from .checklists import people_path

    src = (ps or {}).get("inputs", {}).get("1300 감사반 인원")
    dst = people_path(d)
    if not src or dst.exists():
        return None
    p = (Path(d) / src).resolve()
    if not p.is_file():
        return None
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(p, dst)
    return _rel(d, dst)


# ---------------------------------------------------------------- 한 번에
def build(d: Path, dart_target: str | None = None, force: bool = False, now: str | None = None, reuse: bool = False) -> dict:
    """시작 자료를 읽고 색인 · 조서 · 메일 초안을 쓴다. reuse 면 받은자료를 다시 읽지 않고 시작자료.json 에서 조서만 다시 만든다."""
    d = Path(d)
    at = now or _now()
    if _idx(d, "감사건.json") is None:
        raise IntakeError(f"_색인/감사건.json 이 없습니다. 먼저 감사 폴더를 만들어 주십시오: {d}")
    old = _idx(d, START_FILE, {}) or {}
    # 전기조서의 글(메모 · 절차 · 발견사항 · 결론)을 먼저 읽어 _색인/전기조서.json 에 둔다(2026-10-12). 당기 조서가 참고로 싣고,
    # 아래 표 판독(read_prior_workpapers)이 글 블록 라벨이 있는 계정 조서를 가려내는 데 쓴다
    from . import priorwp

    pwc = priorwp.collect(d, now=at)
    if reuse and old:
        data = copy.deepcopy(old)
    else:
        g = gather(d, dart_target)
        case = _idx(d, "감사건.json")
        data = {"format_version": 1, "made_at": at, "prior_end": g["prior_end"], "company": g["company"],
                "prior_report": g["prior_report"], "dart": g["dart"], "prior_workpapers": g["prior_workpapers"],
                "prior_fs": g["prior_fs"], "need_ai": g["need_ai"], "warnings": g["warnings"], "term": g["term"],
                "ai": old.get("ai", []), "made": old.get("made", {})}
        fy = _fy_end(case)
        data["parties"] = top_parties(d, fy)
        mat = ((g["prior_workpapers"] or {}).get("materiality") or {}).get("performance")
        data["changes"] = big_changes(current_statement(d), g["prior_fs"], mat["amount"] if mat else None)
        if g["saved_dart"]:
            data["dart_saved"] = g["saved_dart"]
    _merge_ai(data)
    data["made_at"] = at
    # 전기 자료(R7): 작년 감사 폴더 · 전임 전기조서 · 전기조서 없음을 판정하고 전기 판단 정보를 모은다
    data["prior_source"] = prior.collect(d, gathered=data)
    # 계정 조서(Lead · 절차 시트)는 중요성 · 정산표 같은 표가 없는 것이 정상이다. 영역으로 이어진 전기조서는 표 판독 필요에서 뺀다
    read_ok = {f["path"] for f in pwc["files"] if f["blocks"] or f["area"]}
    data["need_ai"] = [x for x in data.get("need_ai", []) if not x.startswith(PRIOR_TEXT_AI)
                       and not (x.startswith(PRIOR_TABLE_AI) and x.rsplit(": ", 1)[-1] in read_ok)] + \
        [f"{PRIOR_TEXT_AI}: {p}" for p in pwc["need_ai"]]
    case = _idx(d, "감사건.json")

    # 작년 감사 폴더의 감사반 인원(1300)을 당기 폴더에 없을 때만 옮긴다. 있으면 둔다
    copied = _carry_people(d, data["prior_source"])

    # 빠진 자료를 먼저 지휘 규칙으로 정한다(PBC 와 메일이 같은 목록을 쓴다)
    director.run(d, now=at)
    data["missing"] = _idx(d, "문서목록.json")["missing"]

    made = dict(data.get("made", {}))
    wrote, skipped = [], []
    for no, maker in (("PBC", None), ("1100", make_1100), ("2100", make_2100)):
        path = wp.path_for(d, no)
        if not _may_write(path, made, force):
            skipped.append(_rel(d, path))
            continue
        b = make_pbc(d, data)[0] if no == "PBC" else maker(d, data, case)
        b.save(path)
        made[path.name] = _sha(path)
        wrote.append(_rel(d, path))
    p1200 = wp.path_for(d, "1200", ext=".docx")
    if _may_write(p1200, made, force):
        make_1200(d, data, case, p1200)
        made[p1200.name] = _sha(p1200)
        wrote.append(_rel(d, p1200))
    else:
        skipped.append(_rel(d, p1200))
    data["made"] = made
    if copied:
        wrote.append(copied)
    data["pbc"] = [{k: r[k] for k in ("no", "item", "status")} for r in pbc_status(d)]
    missing = data.pop("missing")       # 빠진 자료의 정본은 문서목록.missing 이다. 여기 두지 않는다
    folder._write(d / "_색인" / START_FILE, data)
    data["missing"] = missing

    touched = _doc_facts(d, data)
    mail = write_mail(d, date.fromisoformat(at[:10]) if re.match(r"\d{4}-\d{2}-\d{2}", at) else None)
    if mail:
        wrote.append(_rel(d, mail))
    read = [x for x in ((data.get("prior_report") or {}).get("path"), (data.get("dart") or {}).get("path")) if x]
    read += list(((data.get("prior_workpapers") or {}).get("files") or {}).keys())
    _update_case(d, data, at, wrote, read)
    folder.scan(d, now=at, quiet_log=True)
    director.run(d, now=at)
    folder.validate_index(d)
    return {"wrote": wrote, "skipped": skipped, "facts_on": touched, "company": data["company"],
            "prior_report": data.get("prior_report"), "dart": data.get("dart"), "prior_workpapers": data.get("prior_workpapers"),
            "prior_fs_basis": (data.get("prior_fs") or {}).get("basis"), "need_ai": data["need_ai"], "warnings": data["warnings"],
            "missing": _idx(d, "문서목록.json")["missing"], "mail": _rel(d, mail) if mail else None,
            "big_changes": [r["name"] for r in data["changes"]["rows"] if r["big"]], "pbc": data["pbc"]}

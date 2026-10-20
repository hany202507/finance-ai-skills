"""한공회 점검표 조서(R5, 2026-10-05): 1300 독립성 · 3200 · 3300 · 3400 · 3500 · 3600 · 3700 · 8200 · 8600.
3500 에는 특수관계자 거래 · 잔액 시트(audit/relparty.py, 2026-10-16)가 붙는다.

문항은 한공회 감사조서서식 2026 원본에서 뽑은 tools/proto_data/kicpa_checklists_r5.json(없으면 kicpa_checklists.json)이다.
문항마다 예 · 아니오 · 해당 없음 기입란(rule yesno)과 근거 기입란을 둔다. 판단은 회계사 몫이라 엔진은 답을 넣지 않는다.
엔진이 색인에서 아는 사실(의사록 판독 사실, 3600-JET 조서 경로 등)은 표 위에 "엔진이 아는 사실" 로 적는다.

  load_items()                     조서번호 → {"title", "items": [[문항, 조서번호], ...]}
  classify(items, narrative)       문항마다 종류(head 제목 · item 확인 · fill 적어 넣기 · note 주석)를 정한다
  people(d, given)                 1300 감사반 인원. 명령줄 입력이 앞서고, 없으면 1000 감사계약/감사반인원.xlsx
  hints(d, no)                     엔진이 아는 사실 문장
  going_concern(d)                 3700 계속기업 재무적 사항(시산표 · 전기 재무제표 · 계정명세서 차입금 표)과 점검 항목의 「해당 후보」 표시
  make(d, no, ...)                 조서 하나를 만든다. refresh 면 앞선 파일의 기입란 값을 이어받는다
  people_template(d)               빈 감사반인원.xlsx 를 만든다(이미 있으면 두고 경로만 돌려준다)
  carry_inputs(book, old_path)     앞선 파일의 기입란 값을 (시트, 항목) 으로 맞춰 새 책에 옮긴다
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from . import folder, params, wp

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "tools" / "proto_data"
YESNO = '"예,아니오,해당 없음"'
EVAL = '"양호,미흡,해당 없음"'


def load_defs() -> dict:
    from . import params

    return params.overlay("close", folder.read_yaml(folder.DEFS / "close.yaml"))   # 2700B 값(2026-10-07)


def load_items() -> dict:
    out = {}
    for name in ("kicpa_checklists.json", "kicpa_checklists_r5.json"):   # 뒤 파일(원본 전체 문항)이 앞선다
        p = DATA / name
        if p.exists():
            out.update({k: v for k, v in json.loads(p.read_text(encoding="utf-8")).items() if not k.startswith("$")})
    return out


# ---------------------------------------------------------------- 문항 종류
_NOTE = re.compile(r"^(\(\*+\)|\*|※|\(주\d*\)|☞|◈|▣|■|□|●|◆|◎)")
_LEVELS = [
    (1, re.compile(r"^\d+\s*[.,](?!\d)|^\d+-\d+\.")),
    (2, re.compile(r"^(\(\d+\)|[⑴-⒇])")),
    (3, re.compile(r"^\d+\)")),
    (4, re.compile(r"^([①-⑳]|\([a-z]+\)|[a-z]\)|[ivx]+\)|[-∙•·])")),
]


def _level(text: str) -> int:
    for n, rx in _LEVELS:
        if rx.match(text):
            return n
    return 5


def classify(items: list, narrative: bool = False) -> list[dict]:
    """문항마다 kind 를 붙인다.

    head: 아래에 더 깊은 문항이 있는 제목(기입란 없음). note: 주석(기입란 없음).
    fill: 값을 적어 넣는 줄(끝이 ':' 이거나 '( )' 가 있음). item: 예 · 아니오 · 해당 없음으로 답하는 문항.
    narrative(3200): 제목 아래 설명 줄과 짝 없는 제목에 내용 기입란을 둔다. 첫 주석부터는 모두 주석.
    """
    rows = [{"text": str(t).strip(), "ref": str(r or "").strip()} for t, r in items if str(t).strip()]
    in_notes = False
    for x in rows:                      # 1) 주석: 3200 처럼 서술형이면 첫 주석부터 끝까지
        t = x["text"]
        x["level"] = _level(t)
        if in_notes or _NOTE.match(t) or (t.startswith("(") and t.endswith(")") and x["level"] == 5):
            x["kind"] = "note"
            in_notes = in_notes or narrative
    for i, x in enumerate(rows):        # 2) 제목 · 적어 넣기 · 확인 문항
        if x.get("kind") == "note":
            continue
        t = x["text"]
        nxt = None
        for y in rows[i + 1:]:
            if y.get("kind") == "note":
                if narrative:
                    break
                continue
            nxt = y
            break
        if "：" in t or "( )" in t:
            x["kind"] = "fill"
        elif nxt is not None and nxt["level"] > x["level"] and x["level"] < 5:
            x["kind"] = "head"
        elif narrative or t.endswith(":"):
            x["kind"] = "fill"
        else:
            x["kind"] = "item"
    return rows


# ---------------------------------------------------------------- 감사반 인원
def people_path(d: Path) -> Path:
    return Path(d) / load_defs()["independence"]["people_file"]


def people_template(d: Path) -> Path:
    p = people_path(d)
    if p.exists():
        return p
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "감사반인원"
    ws.append(["성명", "역할", "등록번호"])
    ws.append([])
    ws["E1"] = "역할: " + " · ".join(load_defs()["independence"]["roles"]) + ". 감사반 소속 공인회계사 전원을 적어 주십시오(외감법 제9조 제3항)"
    for col, w in zip("ABCE", (14, 16, 14, 60)):
        ws.column_dimensions[col].width = w
    wb.save(p)
    return p


def people(d: Path, given: list[str] | None = None) -> tuple[list[dict], str]:
    """(인원 목록, 출처). given 은 '성명:역할[:등록번호]' 목록. 인원을 지어내지 않는다: 없으면 빈 목록."""
    if given:
        out = []
        for g in given:
            parts = [x.strip() for x in str(g).split(":")]
            if parts and parts[0]:
                out.append({"name": parts[0], "role": parts[1] if len(parts) > 1 else "", "reg": parts[2] if len(parts) > 2 else ""})
        return out, "명령줄 입력"
    p = people_path(d)
    if not p.exists():
        return [], ""
    wb = load_workbook(p, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], ""
    head = [str(v or "").strip() for v in rows[0]]
    col = {k: head.index(k) for k in ("성명", "역할", "등록번호") if k in head}
    if "성명" not in col:
        return [], ""
    out = []
    for r in rows[1:]:
        name = str(r[col["성명"]] or "").strip() if col["성명"] < len(r) else ""
        if not name:
            continue
        get = lambda k: str(r[col[k]] or "").strip() if k in col and col[k] < len(r) else ""   # noqa: E731
        out.append({"name": name, "role": get("역할"), "reg": get("등록번호")})
    return out, p.relative_to(d).as_posix()


# ---------------------------------------------------------------- 엔진이 아는 사실
def _read(d: Path, name: str, default):
    p = Path(d) / "_색인" / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def _wp_file(d: Path, no: str) -> str | None:
    w = next((x for x in _read(d, "조서목록.json", {"workpapers": []})["workpapers"] if x["no"] == no), None)
    return w.get("file") if w else None


def hints(d: Path, no: str) -> list[str]:
    """색인에 있는 사실만 옮긴다. 숫자를 새로 만들지 않고, AI 판독 사실은 확인 전이라고 붙인다."""
    d = Path(d)
    docs = _read(d, "문서목록.json", {"files": []})["files"]
    case = _read(d, "감사건.json", {})
    out = []

    def facts_of(pred, limit=12):
        n = 0
        for f in docs:
            if not pred(f):
                continue
            for x in f.get("facts") or []:
                tag = "엔진 판독" if x.get("basis") == "엔진" else ("AI 판독, 회계사 확인 끝" if f.get("confirmed") else "AI 판독, 확인 전")
                page = f" {x['source_page']}" if x.get("source_page") else ""
                out.append(f"{x['text']} ({f['path']}{page}, {tag})")
                n += 1
                if n >= limit:
                    return

    if no == "3500":
        facts_of(lambda f: f.get("role") == "received" and "의사록" in (f.get("kind") or ""))
        perm = [f["path"] for f in docs if f.get("role") == "permanent" and f.get("kind") == "특수관계자목록"]
        out += [f"영구조서 특수관계자목록: {p}" for p in perm]
        if not perm:
            out.append("영구조서에 150 특수관계자목록이 없습니다")
    elif no == "3600":
        jet = _wp_file(d, "3600-JET")
        out.append(f"분개 테스트 조서: {jet}" if jet else "3600-JET 분개테스트 조서가 아직 없습니다")
        issues = [i for i in _read(d, "이슈.json", {"issues": []})["issues"] if i.get("status") == "열림"]
        out += [f"열린 이슈 {i['no']}: {i['title']} (발생 {i['origin_wp']})" for i in issues[:10]]
    elif no == "3300":
        for n in ("2120A", "8110A"):
            p = _wp_file(d, n)
            if p:
                out.append(f"분석적절차 조서: {p}")
        if (d / "_색인" / "분석.json").exists():
            out.append("분석 결과 색인: _색인/분석.json")
    elif no == "3400":
        areas = [w for w in _read(d, "조서목록.json", {"workpapers": []})["workpapers"] if folder._is_area(w["no"]) and w.get("file")]
        out += [f"계정별 조서: {w['file']}" for w in areas]
    elif no == "8200":
        fy = case.get("fiscal_year_end") or ""
        after = [f for f in docs if f.get("role") == "received" and f.get("period") and fy and str(f["period"]) > fy[:7]]
        out += [f"결산일 뒤 기간의 받은자료: {f['path']} ({f['kind']}, {f['period']})" for f in after]
        facts_of(lambda f: f in after)
    elif no == "8600":
        for n in ("8100A", "8820"):
            p = _wp_file(d, n)
            if p:
                out.append(f"관련 조서: {p}")
        opened = [i for i in _read(d, "이슈.json", {"issues": []})["issues"] if i.get("status") == "열림"]
        out.append(f"열린 이슈 {len(opened)}건(이슈.json)")
    elif no == "2110":
        out += _materiality(d)
        issues = [i for i in _read(d, "이슈.json", {"issues": []})["issues"] if i.get("status") == "열림"]
        out += [f"열린 이슈 {i['no']}: {i['title']} → 대응 조서 {i.get('response_wp') or '미정'}" for i in issues[:15]]
        if not issues:
            out.append("열린 이슈가 없습니다(이슈.json)")
        ppl, src = people(d)
        out.append("감사반 인원: " + (", ".join(f"{p['name']}({p['role'] or '역할 미기재'})" for p in ppl) + f" [{src}]"
                                    if ppl else f"감사반 인원 목록이 없습니다({load_defs()['independence']['people_file']})"))
    elif no == "3700":
        out += [f"사실 {i}. {x['text']}" for i, x in enumerate(going_concern(d), 1)]
    elif no == "2301":
        ana = _read(d, "분석.json", None)
        if ana and ana.get("issues"):
            out += [f"분석 상위 {x['rank']}: {x['title']} (영역 {x.get('area') or '-'}, 이슈 {x['issue']})" for x in ana["issues"][:10]]
        else:
            out.append("_색인/분석.json 이 없거나 상위 이슈가 없습니다(분석적 검토 전)")
        p2302 = _wp_file(d, "2302")
        out.append(f"경영진주장 수준 위험 조서: {p2302}" if p2302 else "2302 조서가 아직 없습니다")
    elif no in SIZE_WPS:
        n = 0
        for f in docs:
            if f.get("role") != "received" or f.get("kind") not in SIZE_KINDS:
                continue
            for x in f.get("facts") or []:
                if x.get("basis") == "엔진" and any(k in x["text"] for k in SIZE_KEYS):
                    out.append(f"회사 규모: {x['text']} ({f['path']}, 엔진 판독)")
                    n += 1
        if not n:
            out.append("회사 규모(매출 · 인원)를 엔진이 읽은 받은자료가 없습니다")
        out.append("소규모 기업 간소화 대상인지는 회계사 판단입니다. 엔진은 정하지 않습니다")
    return out


# ---------------------------------------------------------------- 3700 계속기업 (2026-10-18)
def _won(x) -> str:
    return f"{int(round(x)):,}원"


def _prior_pl(pf: dict | None, stem: str):
    """전기 손익계산서의 영업 · 당기순 손익(손실은 음수). 「영업손실」 처럼 손실 이름 줄은 금액 부호와 관계없이 손실로 본다."""
    from . import reader

    for k, v in ((pf or {}).get("is") or {}).items():
        if not isinstance(v, (int, float)):
            continue
        n = reader._norm(k)
        if n in (f"{stem}이익", f"{stem}손익"):
            return v
        if n == f"{stem}손실":
            return -abs(v)
    return None


def _prior_bs(pf: dict | None, name: str):
    from . import reader

    for k, v in ((pf or {}).get("bs") or {}).items():
        if reader._norm(k) == name and isinstance(v, (int, float)):
            return v
    return None


def _pl_text(what: str, cur, prior) -> tuple[str, bool]:
    """(문장, 당기 손실인가). 금액은 시산표 · 전기 재무제표 값 그대로."""
    def one(label, v):
        return f"{what}{'손실' if v < 0 else '이익'} {_won(abs(v))}({label})"
    t = one("당기", cur)
    if prior is None:
        t += f". 전기 {what}손익은 전기 재무제표에서 찾지 못했습니다"
    else:
        t += ", " + one("전기", prior)
        if cur < 0 and prior < 0:
            t += f" (당기 · 전기 연속 {what}손실)"
    return t, cur < 0


def going_concern(d: Path) -> list[dict]:
    """3700 「엔진이 아는 사실」. [{key, text, candidate}]. key 는 close.yaml going_concern.links 의 사실 종류다.
    숫자는 회사제시 시산표(figures.current_tb) · 전기 재무제표(figures.prior) · 계정명세서 차입금 표(recalc.loans)에서만 온다.
    candidate 는 한공회 3700 점검 항목에 「해당 후보」 를 붙일지다. 해당 여부(예 · 아니오)와 계속기업 판단은 회계사 몫이다."""
    from datetime import date as _date

    from . import figures, reader
    from .workpapers import common, recalc

    d = Path(d)
    gc = load_defs().get("going_concern") or {}
    case = _read(d, "감사건.json", {})
    out: list[dict] = []

    def add(key: str, text: str, candidate: bool = False) -> None:
        out.append({"key": key, "text": text, "candidate": candidate})

    rows = [r for r in figures.current_tb(d) if not figures.is_group_row(r.get("name"))]
    pf = figures.prior(d)
    actual: dict = {}
    if not rows:
        add("tb", f"재무적 사항을 계산하지 못했습니다. {figures.tb_absent_reason(d)}")
    else:
        def grp(r) -> str | None:
            if r.get("fs_group"):
                return r["fs_group"]
            k = common.code3(r["code"])
            return next((g for g, (lo, hi) in figures.FS_RANGES.items() if k is not None and lo <= k <= hi), None)

        def tot(g: str) -> int:
            return sum(r["amount"] for r in rows if grp(r) == g)

        ver = next((r.get("version") for r in rows if r.get("version")), None)
        src = f"회사제시 시산표{' ' + ver if ver else ''}"
        ca, cl = tot("유동자산"), -tot("유동부채")
        assets, liab = ca + tot("비유동자산"), cl - tot("비유동부채")
        equity = assets - liab
        capital = -sum(r["amount"] for r in rows if grp(r) == "자본" and reader._norm(r["name"]).endswith("자본금"))
        t = f"당기말 자본총계 {_won(equity)} (자산총계 {_won(assets)} − 부채총계 {_won(liab)}, {src})"
        pe = _prior_bs(pf, "자본총계")
        if pe is not None:
            t += f". 전기말 자본총계 {_won(pe)} (전기 재무제표)"
        if equity < 0:
            add("capital_deficit", t + ". 자본총계가 음수다(완전자본잠식)", True)
        elif capital > 0 and equity < capital:
            add("capital_deficit", t + f". 자본금 {_won(capital)} 보다 작다(부분자본잠식, 자본잠식률 {(capital - equity) / capital:.1%})")
        else:
            add("capital_deficit", t)
        t = f"당기말 유동자산 {_won(ca)}, 유동부채 {_won(cl)} ({src}, 계정코드 범위 유동자산 101~175 · 유동부채 251~290)"
        if cl > ca:
            add("net_current_liabilities", t + f". 유동부채가 유동자산보다 많다(순유동부채 {_won(cl - ca)})", True)
        else:
            add("net_current_liabilities", t)
        # 영업손익은 매출액 · 매출원가 · 판매비와관리비 묶음과 제조 · 도급원가 계정(코드 501~800), 당기순손익은 손익 계정 전부
        op_groups = ("매출액", "매출원가", "판매비와관리비")
        pl_groups = op_groups + ("영업외수익", "영업외비용", "법인세비용")
        cost = lambda r: grp(r) is None and 501 <= (common.code3(r["code"]) or 0) <= 800      # noqa: E731
        op = -sum(r["amount"] for r in rows if grp(r) in op_groups or cost(r))
        ni = -sum(r["amount"] for r in rows if grp(r) in pl_groups or cost(r))
        t, loss = _pl_text("영업", op, _prior_pl(pf, "영업"))
        mfg = " · 제조 · 도급원가(코드 501~800)" if any(cost(r) and r["amount"] for r in rows) else ""
        add("operating_loss", t + f" ({src} 매출액 · 매출원가 · 판매비와관리비{mfg}, 전기 재무제표)", loss)
        actual = {"revenue": -tot("매출액"), "op": op}
        t, loss = _pl_text("당기순", ni, _prior_pl(pf, "당기순"))
        add("net_loss", t + f" ({src} 손익 계정, 전기 재무제표)", loss)
        short = [r for r in rows if grp(r) == "유동부채" and any(k in reader._norm(r["name"]) for k in ("단기차입금", "유동성장기"))]
        if short:
            v = -sum(r["amount"] for r in short)
            names = " · ".join(dict.fromkeys(r["name"] for r in short))
            add("borrowings_due", f"유동부채로 분류된 차입금({names}) 합계 {_won(v)} ({src})", v > 0)
        else:
            add("borrowings_due", f"시산표 유동부채에 단기차입금 · 유동성장기부채 계정이 없습니다 ({src})")
    try:
        ls = recalc.loans(d)
    except Exception:            # 계정명세서를 읽지 못하면 명세가 없는 것으로 본다
        ls = []
    fy = case.get("fiscal_year_end")
    if ls and fy:
        e = _date.fromisoformat(fy)
        try:
            limit = e.replace(year=e.year + 1)
        except ValueError:       # 2월 29일
            limit = _date(e.year + 1, 2, 28)
        due = [(x, reader._date(x.get("만기"))) for x in ls]
        hit = [x for x, m in due if m and m <= limit]
        unknown = sum(1 for _, m in due if m is None)
        t = (f"차입금 표(계정명세서 · 회사 차입금 리스트) {len(ls)}건 가운데 만기가 {limit.isoformat()} 까지인 차입 {len(hit)}건 "
             f"합계 {_won(sum(abs(x['잔액']) for x in hit))}")
        if unknown:
            t += f". 만기를 읽지 못한 차입 {unknown}건(만기 열이 없거나 날짜가 아님)"
        add("borrowings_due", t + ". 갱신 · 상환 가능성은 약정서 · 조회로 회계사가 본다", bool(hit))
    elif not ls:
        add("borrowings_due", "계정명세서 · 회사 차입금 리스트에 차입금 표가 없어 1년 안 만기 차입금 합계를 내지 못했습니다")
    keys = [str(k).lower() for k in gc.get("plan_keys") or []]
    docs = _read(d, "문서목록.json", {"files": []})["files"]
    plans = [f["path"] for f in docs if f.get("role") == "received" and any(k in Path(f["path"]).name.lower() for k in keys)]
    nxt = f"{int(fy[:4]) + 1}년 " if fy else "다음 해 "
    from . import folder as _folder, plans as _plans

    read_plans = [(p, _plans.read(d / p)) for p in plans]
    years = sorted({g["year"] for _, g in read_plans if g and g.get("year")})
    if years:                                  # 파일에서 읽은 계획 연도(중간감사는 결산 연도 계획이다, 2026-10-20)
        nxt = " · ".join(f"{y}년" for y in years) + " "
    if plans:
        add("plan", f"{nxt}사업계획 · 추정 자료로 보이는 받은자료 {len(plans)}개(파일 이름으로 찾음): " + ", ".join(plans[:5])
            + (f" 외 {len(plans) - 5}개" if len(plans) > 5 else "") + ". 내용과 기간은 회계사가 확인한다")
        # 사업계획 ↔ 실적(2026-10-20): 엑셀 계획을 읽어 같은 기간 합계(중간감사 · 예산) 또는 다음 해 연간 계획과 당기 실적을 견준다
        if actual and fy:
            start, end = _folder.fiscal_start(case), _date.fromisoformat(fy[:10])
            for p, got in read_plans:
                c = _plans.compare(got, start, end, actual) if got else None
                if c:
                    short = c["kind"] == "같은 기간" and any(ln["diff"] < 0 for ln in c["lines"])
                    add("plan_actual", _plans.text(c, _won) + ". 계획의 실현 가능성은 회계사가 본다", short)
    else:
        add("plan", f"받은자료에 {nxt}사업계획 · 추정재무제표 · 자금계획으로 보이는 파일이 없습니다(파일 이름으로 찾음)")
    return out


def going_concern_marks(rows: list[dict], facts: list[dict]) -> dict[int, str]:
    """점검 항목 줄 번호(rows 의 0부터) → 「해당 후보: 사실 1 · 2」 같은 표시. 답 칸은 비워 둔다."""
    links = (load_defs().get("going_concern") or {}).get("links") or {}
    out: dict[int, str] = {}
    for i, x in enumerate(rows):
        if x["kind"] != "item":
            continue
        cand, ref = [], []
        for n, f in enumerate(facts, 1):
            if any(k in x["text"] for k in links.get(f["key"], [])):
                (cand if f["candidate"] else ref).append(str(n))
        parts = ([f"해당 후보: 사실 {' · '.join(cand)}"] if cand else []) + ([f"참고: 사실 {' · '.join(ref)}"] if ref else [])
        if parts:
            out[i] = ". ".join(parts)
    return out


SIZE_WPS = ("2511", "2512", "2513", "2520", "2530", "2531", "2532", "2533")
SIZE_KINDS = ("회사제시 재무제표·시산표", "급여대장")
SIZE_KEYS = ("매출액", "자산총계", "인원")


def _materiality(d: Path) -> list[str]:
    """2700A 중요성 기입란 가운데 회계사가 채운 값. 없으면 2700A 미입력."""
    f = _wp_file(d, "2700A")
    if not f or not (Path(d) / f).exists() or not f.endswith(".xlsx"):
        return ["2700A 미입력(중요성 조서 없음)"]
    vals = [x for x in wp.readback(Path(d) / f) if not wp._blank(x["value"]) and x["label"] not in ("작성자", "검토자")]
    if not vals:
        return ["2700A 미입력"]
    return [f"2700A {x['label']}: {x['value']:,}" if isinstance(x["value"], (int, float)) else f"2700A {x['label']}: {x['value']}"
            for x in vals[:8]]


# ---------------------------------------------------------------- 기입란 이어받기
def carry_inputs(book: wp.Book, old_path: Path) -> int:
    """앞선 파일에서 같은 (시트, 항목) 기입란에 적힌 값을 새 책에 옮긴다. 옮긴 칸 수."""
    if not Path(old_path).exists():
        return 0
    old = {(x["sheet"], x["label"]): x["value"] for x in wp.readback(old_path) if not wp._blank(x["value"])}
    n = 0
    for sheet, ref, label, _rule in book.inputs:
        v = old.get((sheet, label))
        if v is not None:
            book.wb[sheet][ref].value = v
            n += 1
    return n


def _yesno_validation(ws) -> DataValidation:
    dv = DataValidation(type="list", formula1=YESNO, allow_blank=True)
    ws.add_data_validation(dv)
    return dv


def preparer_input(b: wp.Book, ws) -> None:
    """머리 작성자 칸을 기입란으로 바꾼다(검사 통과 때 이 이름으로 작성 서명을 단다)."""
    note = b.preparer
    b.ask(ws, 3, 5, "작성자", rule="optional")
    ws.cell(3, 5).value = None          # ws.cell(.., value=None) 은 앞 값을 지우지 않는다
    ws.cell(3, 6, note).font = wp.font(size=9, color=wp.MUTED)


# ---------------------------------------------------------------- 조서 만들기
def _table(b: wp.Book, ws, row: int, rows: list[dict], no: str, key: str, dv, evaluate=None, marks: dict | None = None) -> int:
    """evaluate 가 있으면(통제 이해 2511~2513) 첫째 수준 고려요소마다 평가결과(양호 · 미흡 · 해당 없음) 기입란을 둔다.
    marks 가 있으면(3700) 여섯째 열에 엔진 사실과 이은 표시(해당 후보 · 참고)를 적는다. 기입란이 아니다."""
    third = "예 · 아니오 · 해당 없음" + (" / 평가결과" if evaluate else "")
    row = b.head(ws, row, ["항목", "조서번호", third, "근거 · 수행 결과", "수행자"] + (["엔진 표시(답은 회계사)"] if marks else []))
    if marks:
        ws.column_dimensions["F"].width = 30
    for i, x in enumerate(rows, 1):
        if marks and (i - 1) in marks:
            b.put(ws, row, 6, marks[i - 1], wrap=True, color="B3261E")
        kind = x["kind"]
        b.put(ws, row, 1, x["text"], wrap=True, bold=(kind == "head"), color=wp.MUTED if kind == "note" else None)
        ws.row_dimensions[row].height = 15 * max(1, min(6, len(x["text"]) // 60 + 1))
        b.put(ws, row, 2, x["ref"] or None)
        lab = f"{key}{i}. {x['text'][:40]}"
        if kind == "head" and evaluate is not None and x["level"] == 1 and not any(w in x["text"][:12] for w in ("결론", "기타")):
            evaluate.add(b.ask(ws, row, 3, f"{lab} | 평가결과", rule="required").coordinate)
            b.ask(ws, row, 4, f"{lab} | 평가 근거", rule="optional")
        elif kind == "item":
            cell = b.ask(ws, row, 3, f"{lab} | 답", rule="yesno")
            dv.add(cell.coordinate)
            b.ask(ws, row, 4, f"{lab} | 근거", rule="optional")
            b.ask(ws, row, 5, f"{lab} | 수행자", rule="optional")
        elif kind == "fill":
            b.ask(ws, row, 4, f"{lab} | 내용", rule="required" if no == "3200" else "optional")
        row += 1
    return row


def _hint_block(b: wp.Book, ws, row: int, lines: list[str]) -> int:
    row = b.section(ws, row, "엔진이 아는 사실(색인에서 옮김. 판단은 회계사가 합니다)")
    if not lines:
        ws.cell(row, 1, "색인에 이 조서와 이어진 사실이 없습니다").font = wp.font(size=9, color=wp.MUTED)
        return row + 2
    for t in lines:
        c = ws.cell(row, 1, t)
        c.font = wp.font(size=9, color="B3261E")
        row += 1
    return row + 1


def _independence(b: wp.Book, ws, row: int, ppl: list[dict], source: str, dv) -> int:
    ind = load_defs()["independence"]
    row = b.section(ws, row, "감사반 소속 공인회계사별 독립성 확인(외감법 제9조 제3항 · 공인회계사법 제21조)")
    ws.cell(row, 1, "답: 예(해당하지 않음을 확인) · 아니오(해당함, 근거 칸에 내용) · 해당 없음").font = wp.font(size=9, color=wp.MUTED)
    row += 1
    if not ppl:
        ws.cell(row, 1, f"감사반 인원 목록이 없습니다. {ind['people_file']} 에 성명 · 역할을 적고 "
                        "python -m audit close make <감사 폴더> 1300 --refresh 로 다시 만드십시오").font = wp.font(color="B3261E", bold=True)
        b.ask(ws, row + 1, 1, "감사반 인원 목록", rule="required")
        return row + 3
    ws.cell(row, 1, f"인원 출처: {source}").font = wp.font(size=9, color=wp.MUTED)
    row += 1
    for k, p in enumerate(ppl, 1):
        head = f"{p['name']} ({p['role'] or '역할 미기재'}{', ' + p['reg'] if p.get('reg') else ''})"
        row = b.section(ws, row + 1, f"{k}. {head}")
        row = b.head(ws, row, ["확인 항목", "근거 조문", "예 · 아니오 · 해당 없음", "근거 · 내용", "확인일"])
        for j, it in enumerate(ind["items"], 1):
            b.put(ws, row, 1, it["text"], wrap=True)
            ws.row_dimensions[row].height = 30
            b.put(ws, row, 2, it["law"], wrap=True)
            lab = f"독립 {p['name']} {j}"
            dv.add(b.ask(ws, row, 3, f"{lab} | 답", rule="yesno").coordinate)
            b.ask(ws, row, 4, f"{lab} | 근거", rule="optional")
            b.ask(ws, row, 5, f"{lab} | 확인일", rule="optional", fmt="yyyy-mm-dd")
            row += 1
    return row + 1


def _register(b: wp.Book, ws, row: int, data: dict) -> int:
    """한공회 서식의 빈 기입 표. 템플릿 예시 행은 옮기지 않는다. 칸은 모두 선택 기입(줄 수가 회사마다 다르다)."""
    cols = data["columns"]
    row = b.section(ws, row + 1, data.get("table_title") or "기입 표(한공회 서식 열)")
    for j in range(len(cols)):
        letter = chr(ord("A") + j)
        cur = ws.column_dimensions[letter].width or 10
        ws.column_dimensions[letter].width = max(cur, 18)
    row = b.head(ws, row, cols)
    ws.row_dimensions[row - 1].height = 45
    for i in range(1, int(data.get("register_rows", 10)) + 1):
        for j, c in enumerate(cols, 1):
            b.ask(ws, row, j, f"표 {i}행 {c[:24]}", rule="optional")
        row += 1
    for t, _ in data.get("notes", []):
        ws.cell(row, 1, t).font = wp.font(size=9, color=wp.MUTED)
        row += 1
    return row + 1


@params.engaged
def make(d: Path, no: str, people_given: list[str] | None = None, refresh: bool = False) -> Path:
    d = Path(d)
    defs = load_defs()["checklists"]
    if no not in defs:
        raise ValueError(f"점검표 조서가 아닙니다: {no}")
    spec = defs[no]
    data = load_items().get(no)
    if data is None:
        raise ValueError(f"한공회 점검표 문항이 없습니다: {no}")
    path = wp.path_for(d, no)
    if path.exists() and not refresh:
        return path
    b = wp.Book(d, no, standard=spec["standard"] + " · 한공회 감사조서서식 2026 문항")
    ws, row = b.sheet(no, widths=[70, 14, 14, 40, 12])
    preparer_input(b, ws)
    dv = _yesno_validation(ws)
    title = data["title"].split(" ", 1)[1].strip() if " " in data["title"] else data["title"]
    ws.cell(2, 1, f"{ws.cell(2, 1).value} · 원본 제목: {title}").font = wp.font(size=9, color=wp.MUTED)
    if spec.get("per_person"):
        ppl, source = people(d, people_given)
        row = _independence(b, ws, row, ppl, source, dv)
    else:
        row = _hint_block(b, ws, row, hints(d, no))
    layout = data.get("layout")
    if data["items"]:
        rows = classify(data["items"], narrative=bool(spec.get("narrative")))
        ev = None
        if layout == "control":
            ev = DataValidation(type="list", formula1=EVAL, allow_blank=True)
            ws.add_data_validation(ev)
        marks = going_concern_marks(rows, going_concern(d)) if no == "3700" else None
        row = b.section(ws, row, "한공회 점검 항목")
        row = _table(b, ws, row, rows, no, "", dv, evaluate=ev, marks=marks)
    if layout == "register":
        row = _register(b, ws, row, data)
    b.conclusion(ws, row, 5)
    ws.sheet_view.showGridLines = False
    if no == "3200":                    # 소송 조회 대장(R9 A7). 사건 줄은 회계사 기입
        from .workpapers import recalc

        recalc.litigation_sheet(b, d)
    if no == "8600":                    # 감사의견 선택 시트(2026-10-06). 8820 · DART 문형이 이 값을 따른다
        from .report import opinion

        opinion.add_sheet(b)
    if no == "3500":                    # 특수관계자 거래 · 잔액(2026-10-16). 목록은 영구조서 150, 거래는 분개장
        from . import relparty

        relparty.add_sheet(b, d)
    if no == "3300":                    # 3300A 실증적 분석절차 부속 시트(R9 B2). 기대치는 분석 결과를 쓴다
        from .analysis import substantive

        substantive.add_sheet(b)
    from . import priorwp

    priorwp.reference_sheet(b, d, no)   # 작년 같은 번호 조서에 회계사가 적은 내용(참고 시트, 2026-10-12)
    if refresh:
        carry_inputs(b, path)
    return b.save(path)


@params.engaged
def make_all(d: Path, nos: list[str] | None = None, people_given: list[str] | None = None, refresh: bool = False) -> list[Path]:
    targets = nos or list(load_defs()["checklists"])
    return [make(d, no, people_given, refresh) for no in targets]


def counts(path: Path) -> dict:
    """기입란 채움 현황: {yesno, yesno_filled, required, required_filled}."""
    out = {"yesno": 0, "yesno_filled": 0, "required": 0, "required_filled": 0}
    for x in wp.readback(path):
        if x["rule"] == "yesno":
            out["yesno"] += 1
            out["yesno_filled"] += 0 if wp._blank(x["value"]) else 1
        elif x["rule"] == "required":
            out["required"] += 1
            out["required_filled"] += 0 if wp._blank(x["value"]) else 1
    return out


__all__ = ["load_items", "classify", "people", "people_template", "hints", "going_concern", "make", "make_all", "carry_inputs",
           "preparer_input", "counts"]

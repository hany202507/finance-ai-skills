"""ctax employ: 급여대장 · 사원명부 → 통합고용세액공제 인원 · 공제액 · 검증 워크북. 합성 회사만 쓴다."""
from __future__ import annotations

import json
import shutil
from datetime import date
from fractions import Fraction
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ctax.employ import book, calc, cli
from ctax.employ.model import COMPANY_ROWS, PAY_COLS, STAFF_COLS, InputError, load_rules, read_inputs

D = date
# 사번, 성명, 생년월일, 입사일, 퇴사일, 고용형태, 그 밖의 열
STAFF = [
    ("A01", "김대표", D(1970, 1, 1), D(2019, 1, 1), None, "정규직", {"임원": "Y", "최대주주·친족": "Y"}),
    ("A02", "이정규", D(1985, 3, 10), D(2020, 1, 1), None, "정규직", {}),
    ("A03", "박청년", D(1995, 6, 15), D(2024, 7, 1), None, "정규직", {}),
    ("A04", "최경계", D(1990, 5, 28), D(2025, 4, 1), None, "정규직", {}),
    ("A05", "정기간", D(1998, 1, 1), D(2025, 3, 1), None, "기간제", {"계약 1년 미만": "Y"}),
    ("A06", "한단시", D(1980, 1, 1), D(2023, 1, 1), None, "단시간", {"월 소정근로시간": 80}),
    ("A07", "오퇴사", D(1996, 2, 2), D(2023, 3, 1), D(2025, 8, 15), "정규직", {}),
    ("A08", "윤고령", D(1960, 1, 1), D(2024, 10, 1), None, "정규직", {}),
    ("A09", "강장애", D(1975, 5, 5), D(2025, 7, 1), None, "단시간", {"월 소정근로시간": 100, "0.75 요건": "Y", "장애인 등": "Y"}),
    ("A10", "문병역", D(1989, 9, 1), D(2022, 1, 1), None, "정규직", {"입대일": D(2010, 3, 1), "전역일": D(2012, 2, 29)}),
    ("A11", "서무급", D(1992, 1, 1), D(2025, 1, 1), None, "정규직", {}),
    ("A12", "임외국", D(1990, 1, 1), D(2021, 1, 1), None, "정규직", {"내국인": "N"}),
]
NO_DEDUCTION = {"A11"}


def make_input(path: Path, company: dict | None = None, staff=STAFF, years=(2024, 2025), skip=()) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "회사"
    values = {"과세연도 시작일": D(2025, 1, 1), "과세연도 종료일": D(2025, 12, 31), "기업 구분": "중소기업", "증가 지역 수도권": "Y",
              "직전 1차년도 공제액 청년등": 14500000, "직전 1차년도 공제액 청년등 외": 8500000,
              "직전전 1차년도 공제액 청년등 외": 8500000, "전기 신고 직전전 상시근로자 수": 3.0, "전기 신고 직전전 청년등 수": 1.0}
    values.update(company or {})
    ws.append(["항목", "값"])
    for label, *_ in COMPANY_ROWS:
        ws.append([label, values.get(label)])
    ws = wb.create_sheet("사원명부")
    head = [c[0] for c in STAFF_COLS]
    ws.append(head)
    for sid, name, birth, hired, left, kind, more in staff:
        row = {"사번": sid, "성명": name, "생년월일": birth, "입사일": hired, "퇴사일": left, "고용형태": kind, **more}
        ws.append([row.get(h) for h in head])
    ws = wb.create_sheet("급여대장")
    ws.append([c[0] for c in PAY_COLS])
    for sid, name, _, hired, left, *_ in staff:
        for y in years:
            for m in range(1, 13):
                first, last = D(y, m, 1), calc.month_end(D(y, m, 1))
                if hired <= last and (left is None or left >= first) and (sid, y, m) not in skip:
                    ded = (0, 0, 0) if sid in NO_DEDUCTION else (50000, 135000, 106000)
                    ws.append([f"{y}-{m:02d}", sid, name, 3000000, *ded])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture
def case_dir(tmp_path):
    root = tmp_path / "가상회사" / "2025 법인세"
    root.mkdir(parents=True)
    make_input(root / "00 받은자료" / "급여" / "통합고용_입력.xlsx")
    return root


def run(path: Path):
    inp = read_inputs(path, company="가상회사")
    rules, rates = load_rules(), cli.load_rates()
    return inp, calc.compute(inp, rules, rates), rules, rates


def test_headcount_per_worker_and_month_end(case_dir):
    _, res, _, _ = run(case_dir / "00 받은자료" / "급여" / "통합고용_입력.xlsx")
    cur, prev = res["counts"]["해당"], res["counts"]["직전"]
    assert res["method"] == "근로자별"
    assert (cur.total, cur.youth_deem, cur.youth_age) == (Fraction("6.20"), Fraction("4.70"), Fraction("4.03"))
    assert (prev.total, prev.youth_age) == (Fraction("4.25"), Fraction("2.75"))
    assert (cur.me_total, cur.me_youth) == (Fraction("6.20"), Fraction("4.04"))      # 74.5 / 12, 48.5 / 12
    assert (prev.me_total, prev.me_youth) == (Fraction("4.25"), Fraction("2.75"))
    by = {r.person.id: r for r in cur.rows}
    assert by["A07"].n_work == Fraction("0.58") and sum(by["A07"].worked) == 7          # 8월 15일 퇴사: 8월은 월말 재직이 아니다
    assert by["A09"].n_work == Fraction("0.37") and by["A09"].kind == "장애인"           # 0.75 × 6 / 12 = 0.375
    assert by["A06"].weight == Fraction("0.5") and by["A06"].n_youth_deem == 0
    assert by["A04"].n_youth_deem == Fraction("0.75") and by["A04"].n_youth_age == Fraction("0.08")
    reasons = {r.person.id: r.reason for r in cur.rows if not r.regular}
    assert reasons == {"A01": calc.REASONS["c"], "A05": calc.REASONS["a"], "A11": calc.REASONS["e"], "A12": calc.REASONS["f_res"]}


def test_youth_cutoff_dates():
    rules = load_rules()
    inp_staff = {s[0]: s for s in STAFF}

    def person(sid):
        from ctax.employ.model import Person
        sid, name, birth, hired, left, kind, more = inp_staff[sid]
        return Person(id=sid, name=name, birth=birth, hired=hired, left=left, kind=kind,
                      mil_from=more.get("입대일"), mil_to=more.get("전역일"))
    # 별지 제10호의10서식 작성방법 25: 1990년 5월 28일생은 2025년 5월 28일부터 청년이 아니다
    assert calc.youth_cutoff(person("A04"), "중소기업", rules, deemed=False) == D(2025, 5, 28)
    # 작성방법 18: 2025년 4월 1일 계약이면 3년 또는 2년이 지난 2028-04-01 또는 2027-04-01
    assert calc.youth_cutoff(person("A04"), "중소기업", rules, deemed=True) == D(2028, 4, 1)
    assert calc.youth_cutoff(person("A04"), "일반기업", rules, deemed=True) == D(2027, 4, 1)
    # 병역 731일을 빼면 35세가 되는 날이 2026-09-02 로 늦춰진다
    assert calc.age_limits(person("A10"), rules)[1] == D(2026, 9, 2)
    assert calc.is_senior(person("A08"), rules) and not calc.is_senior(person("A02"), rules)


def test_deemed_only_for_contracts_from_2025(tmp_path):
    """청년 간주는 2025년 1월 1일 이후 체결한 계약만이다(2026-10-12 회계사 확인). 2024년 이전 계약은 현재 연령만 본다."""
    staff = [
        ("B01", "이전계약", D(1990, 9, 15), D(2024, 6, 3), None, "정규직", {}),      # 34세에 계약, 2025-09-15 에 35세
        ("B02", "올해계약", D(1990, 9, 15), D(2025, 6, 2), None, "정규직", {}),      # 34세에 계약, 3년 간주
    ]
    inp, res, rules, _ = run(make_input(tmp_path / "a.xlsx", staff=staff))
    by = {r.person.id: r for r in res["counts"]["해당"].rows}
    assert calc.deemed_until(inp.staff[0], "중소기업", rules) is None
    assert calc.deemed_until(inp.staff[1], "중소기업", rules) == D(2028, 6, 2)
    assert by["B01"].youth_deem == by["B01"].youth_age and sum(by["B01"].youth_deem) == 8      # 1 ~ 8월 말
    assert sum(by["B02"].youth_deem) == 7 and sum(by["B02"].youth_age) == 3                    # 6 ~ 12월 말, 현재 연령은 6 ~ 8월 말


def test_credit_first_second_third(case_dir):
    _, res, _, _ = run(case_dir / "00 받은자료" / "급여" / "통합고용_입력.xlsx")
    assert (res["d_t"], res["d_y"], res["d_o"]) == (Fraction("1.95"), Fraction("1.95"), Fraction(0))
    assert (res["n_y"], res["n_o"]) == (Fraction("1.95"), Fraction(0))
    assert res["first"] == 28_275_000                    # 1.95 × 1,450만원
    assert res["second"] == 23_000_000 and res["second_cut"] == ""
    assert res["third"] == 8_500_000 and res["source"]["직전전"] == "전기 신고서"
    assert res["total"] == 59_775_000


def test_second_year_cut_when_total_falls(tmp_path):
    staff = [s if s[0] != "A02" else (*s[:4], D(2024, 12, 31), *s[5:]) for s in STAFF]       # 2024년 말 퇴사
    staff = [s if s[0] not in ("A03", "A08") else (*s[:4], D(2025, 1, 15), *s[5:]) for s in staff]
    p = make_input(tmp_path / "in.xlsx", staff=staff)
    _, res, _, _ = run(p)
    assert res["legacy"]["해당"][0] < res["legacy"]["직전"][0]
    assert res["second"] == 0 and res["second_cut"] == "전체 감소" and res["first"] == 0


def test_startup_and_missing_prior(tmp_path):
    p = make_input(tmp_path / "a.xlsx", years=(2025,), company={"직전 1차년도 공제액 청년등": None, "직전 1차년도 공제액 청년등 외": None})
    _, res, _, _ = run(p)
    assert res["source"]["직전"] == "없음" and res["first"] == 0
    p = make_input(tmp_path / "b.xlsx", years=(2025,), company={"해당 과세연도 창업": "Y", "직전 1차년도 공제액 청년등": None,
                                                               "직전 1차년도 공제액 청년등 외": None})
    _, res, _, _ = run(p)
    assert res["source"]["직전"] == "창업(0)" and res["n_y"] == Fraction("4.70") and res["n_o"] == Fraction("1.50")


def test_month_end_method_before_2025(tmp_path):
    p = make_input(tmp_path / "in.xlsx", years=(2023, 2024),
                   company={"과세연도 시작일": D(2024, 1, 1), "과세연도 종료일": D(2024, 12, 31)})
    _, res, _, _ = run(p)
    assert res["method"] == "월말합" and not res["periods"]["해당"].deemed
    assert res["used"]["해당"] == (Fraction("4.25"), Fraction("2.75"))


def test_input_errors(tmp_path):
    bad = [("A01", "김대표", D(1970, 1, 1), D(2019, 1, 1), D(2018, 1, 1), "정규직", {})]
    with pytest.raises(InputError, match="퇴사일이 입사일보다"):
        read_inputs(make_input(tmp_path / "a.xlsx", staff=bad))
    bad = [("A01", "김대표", D(1970, 1, 1), D(2019, 1, 1), None, "프리랜서", {})]
    with pytest.raises(InputError, match="고용형태"):
        read_inputs(make_input(tmp_path / "b.xlsx", staff=bad))
    with pytest.raises(InputError, match="입력 파일이 없습니다"):
        read_inputs(tmp_path / "none.xlsx")


def test_workbook_is_formulas(case_dir):
    inp, res, rules, rates = run(case_dir / "00 받은자료" / "급여" / "통합고용_입력.xlsx")
    wb, checks = book.build(inp, res, rules, rates, ["확인"])
    assert wb.sheetnames == [book.S_LIST, book.S_BASE, book.S_STAFF, book.S_PAY, *book.S_JUDGE.values(), book.S_COUNT,
                             book.S_CREDIT, book.S_FORM, book.S_ROSTER, book.S_CHECK]
    # 판정 · 인원 · 공제액 · 서식의 숫자 칸은 값이 아니라 수식이다
    for name in (*book.S_JUDGE.values(), book.S_COUNT, book.S_CREDIT, book.S_FORM):
        ws = wb[name]
        nums = [c for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
        assert not nums, f"{name}: 값으로 박힌 숫자 {[(c.coordinate, c.value) for c in nums[:5]]}"
    form = wb[book.S_FORM]
    text = " ".join(str(c.value) for row in form.iter_rows() for c in row if c.value)
    for word in ("⑥ 상시근로자 수(⑦+⑧)", "⑬ 상시근로자 증가 인원 수(⑪-⑫)", "⑳ 1차년도 세액공제액", "㉖ 2차년도 세액공제액", "㉜ 3차년도 세액공제액",
                 "㊱ 추가공제 세액공제액"):
        assert word in text
    assert len([k for k in checks if not k.startswith("_")]) > 60


def test_cli_writes_outputs_and_info(case_dir, capsys):
    code = cli.main([str(case_dir), "--no-excel"])
    assert code == 0
    data = json.loads((case_dir / "_색인" / "통합고용.json").read_text(encoding="utf-8"))
    assert data["credit"]["합계"] == 59_775_000 and data["headcount"]["해당"] == {"상시": 6.2, "청년등": 4.7, "청년등외": 1.5, "출처": "급여대장"}
    info = json.loads((case_dir / "_색인" / "신고건.json").read_text(encoding="utf-8"))
    block = info["상시근로자"]
    assert block["출처"] == "ctax employ" and block["해당"] == {"청년등": 4.7, "청년등 외": 1.5}
    assert block["직전"] == {"청년등": 2.75, "청년등 외": 1.5} and block["1차년도 공제액"]["직전"]["청년등"] == 14500000
    assert (case_dir / "2 조서" / "통합고용_검증.xlsx").exists()
    # 사람이 넣은 값은 덮어쓰지 않는다
    info["상시근로자"] = {"해당": {"청년등": 1, "청년등 외": 1}}
    (case_dir / "_색인" / "신고건.json").write_text(json.dumps(info, ensure_ascii=False), encoding="utf-8")
    cli.main([str(case_dir), "--no-excel"])
    assert json.loads((case_dir / "_색인" / "신고건.json").read_text(encoding="utf-8"))["상시근로자"]["해당"]["청년등"] == 1
    assert "고치지 않았습니다" in capsys.readouterr().out


def test_excel_formulas_match_engine(case_dir):
    """엑셀이 있는 PC 에서만: 수식을 계산시킨 값이 엔진 값과 모두 같고 검산이 모두 통과한다."""
    pytest.importorskip("win32com.client")
    code = cli.main([str(case_dir)])
    data = json.loads((case_dir / "_색인" / "통합고용.json").read_text(encoding="utf-8"))
    if data["excel"] is None:
        pytest.skip("엑셀이 없습니다")
    assert data["excel"]["bad"] == [] and data["excel"]["errors"] == []
    assert [t for t in data["excel"]["tests"] if t["result"] != "통과"] == []
    assert code == 0
    wb = load_workbook(case_dir / "2 조서" / "통합고용_검증.xlsx", data_only=True)
    form = wb[book.S_FORM]
    values = [c.value for row in form.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert 59_775_000 in values and 28_275_000 in values


def test_month_without_pay_is_not_counted(tmp_path):
    """급여대장이 있는 과세연도는 그 달 급여가 있어야 근무개월이다(2026-10-12 결정). 「급여 없는 달 포함」 Y 면 넣는다."""
    skip = {("A03", 2024, 7)}                    # 2024년 7월 1일 입사, 첫 급여는 8월
    _, res, _, _ = run(make_input(tmp_path / "a.xlsx", skip=skip))
    row = {r.person.id: r for r in res["counts"]["직전"].rows}["A03"]
    assert sum(row.worked) == 5 and row.n_work == Fraction("0.41") and row.no_pay == 1
    assert res["counts"]["직전"].total == Fraction("4.16")
    _, res, _, _ = run(make_input(tmp_path / "b.xlsx", skip=skip, company={"급여 없는 달 포함": "Y"}))
    row = {r.person.id: r for r in res["counts"]["직전"].rows}["A03"]
    assert sum(row.worked) == 6 and row.no_pay == 1 and res["counts"]["직전"].total == Fraction("4.25")


def test_sample_twenty(tmp_path):
    """`--sample` 이 만드는 합성 회사 스무 명. 기대값은 손으로 계산했고 README 0절에 같은 숫자가 있다."""
    root = tmp_path / "샘플회사" / "2025 법인세"
    assert cli.main([str(root), "--sample"]) == 0
    src = root / "00 받은자료" / "급여" / "통합고용_입력.xlsx"
    inp, res, _, _ = run(src)
    assert len(inp.staff) == 20
    assert res["used"]["해당"] == (Fraction("13.31"), Fraction("8.73")) and res["used"]["직전"] == (Fraction("11.32"), Fraction("6.99"))
    assert (res["d_t"], res["n_y"], res["n_o"]) == (Fraction("1.99"), Fraction("1.74"), Fraction("0.25"))
    assert res["legacy"]["해당"] == (Fraction("13.33"), Fraction("8.66")) and res["legacy"]["직전"] == (Fraction("11.33"), Fraction("7.00"))
    assert (res["first"], res["second"], res["third"], res["extra"], res["total"]) == (27_355_000, 37_500_000, 14_500_000, 13_000_000,
                                                                                 92_355_000)
    by = {r.person.id: r for r in res["counts"]["해당"].rows}
    assert (by["S07"].n_work, by["S07"].n_youth_deem) == (Fraction(1), Fraction("0.58"))
    assert (by["S09"].n_youth_deem, by["S09"].n_youth_age) == (Fraction("0.91"), Fraction("0.83"))
    assert by["S12"].n_work == Fraction("0.41") and by["S15"].regular and by["S15"].n_youth_deem == 0
    assert {r.person.id for r in res["counts"]["해당"].rows if not r.regular and any(r.paid)} == {"S01", "S02", "S14", "S17"}
    assert cli.main([str(root), "--sample"]) == 0 and cli.main([str(root), "--no-excel"]) == 0      # 있는 입력은 덮어쓰지 않는다


def _kordoc_or_skip():
    pytest.importorskip("fitz")                                    # 합성 PDF 를 만드는 데만 쓴다
    if not shutil.which("npx"):
        pytest.skip("kordoc(npx) 이 없는 PC")


def test_douzone_parse_kordoc_table():
    """kordoc 표 JSON 에서 머리글 네 줄 · 사람마다 세 줄을 읽는다. 빈 칸은 0, 퇴사일은 셋째 줄, 지급합계는 셋째 줄의 지급합계 열."""
    from ctax.employ import douzone_pdf
    head = [["인적사항", "", "기본급여및제수당", "", "공제및차인지급액", "", "", ""],
            ["사원번호", "성 명", "기본급", "휴일근로수당", "국민연금", "건강보험", "소득세", "지방소득세"],
            ["입사일", "직 급", "", "", "건강보험정산", "", "", "공제합계"],
            ["퇴사일", "부 서", "", "지급합계", "", "", "", "차인지급액"]]
    body = [["113", "가 나", "2,000,000", "50,000", "90,000", "70,900", "", ""],
            ["2023-03-01", "", "", "", "", "", "", "160,900"],
            ["2025-10-20", "", "", "2,250,000", "", "", "", "2,089,100"],
            ["합계(1명)", "", "", "2,250,000", "", "", "", ""]]
    cells = [[{"text": t} for t in r] for r in head + body]
    doc = {"blocks": [{"type": "paragraph", "text": "2025년10월분 급여대장"}, {"type": "paragraph", "text": "[귀속:2025년10월] [지급:2025년11월10일]"},
                      {"type": "table", "table": {"cells": cells}}]}
    rows, total = douzone_pdf.parse(doc)
    assert total == 1
    assert rows == [{"month": "2025-10", "사번": "113", "성명": "가나", "입사일": "2023-03-01", "퇴사일": "2025-10-20",
                     "지급액": 2_250_000, "소득세": 0, "국민연금": 90_000, "건강보험": 70_900}]


def test_douzone_pdfs_round_trip(tmp_path):
    """합성 더존 급여대장 PDF(달마다 하나, 열한 명이 넘는 달은 두 쪽)를 kordoc 으로 읽으면 만든 값 그대로다."""
    _kordoc_or_skip()
    from ctax.employ import douzone_pdf, sample
    pdfs = sample.write_douzone_pdfs(tmp_path / "pdf", 2025)
    assert len(pdfs) == 24
    want = sample.pay_rows(2025, numeric=True)
    got = []
    for _, rows, total in douzone_pdf.read_folder(tmp_path / "pdf"):
        assert total == len(rows)
        got += rows
    key = lambda r: (r["month"], r["id"] if "id" in r else r["사번"])
    assert [(r["month"], r["사번"], r["성명"], r["지급액"], r["소득세"], r["국민연금"], r["건강보험"]) for r in sorted(got, key=key)] ==         [(r["month"], r["id"], r["name"], r["pay"], r["tax"], r["pension"], r["health"]) for r in sorted(want, key=key)]
    by = {r["사번"]: r for r in got}
    assert by["1012"]["퇴사일"] == "2025-06-20" and by["1001"]["입사일"] == "2018-03-01"


def test_assemble_staff_company_pdfs(tmp_path):
    """사원명부 · 회사정보 · 급여대장 PDF 2년치를 assemble 로 묶으면 --sample 과 같은 공제액이 나온다."""
    _kordoc_or_skip()
    from ctax.employ import assemble, sample
    kit = tmp_path / "받은 파일"
    sample.write_staff_file(kit / "사원명부.xlsx")
    sample.write_company_file(kit / "회사정보.xlsx")
    sample.write_douzone_pdfs(kit / "급여대장")
    root = tmp_path / "샘플회사" / "2025 법인세"
    args = [str(root), "--staff", str(kit / "사원명부.xlsx"), "--company", str(kit / "회사정보.xlsx"), "--pdfs", str(kit / "급여대장")]
    assert assemble.main(args) == 0
    assert assemble.main(args) == 2                               # 이미 묶은 입력은 덮어쓰지 않는다
    inp, res, _, _ = run(root / "00 받은자료" / "급여" / "통합고용_입력.xlsx")
    assert len(inp.staff) == 20 and res["total"] == 92_355_000
    assert res["used"]["해당"] == (Fraction("13.31"), Fraction("8.73"))

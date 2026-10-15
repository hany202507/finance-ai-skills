"""받은 파일 여러 개를 통합고용 입력 양식 하나로 묶는다.

`python -m ctax.employ.assemble "<신고 건>" --staff <사원명부.xlsx> --company <회사정보.xlsx> --pdfs <급여대장 PDF 폴더>`

1. 빈 양식(`ctax employ --template` 과 같은 것)을 `00 받은자료/급여/통합고용_입력.xlsx` 에 만든다. 이미 있으면 멈춘다.
2. 회사정보 파일의 「항목 · 값」 을 양식 「회사」 시트의 같은 항목에 옮긴다. 항목 이름은 양식과 같아야 한다.
3. 사원명부 파일의 머리글을 양식 머리글과 맞춰 줄마다 옮긴다. 양식에 없는 머리글은 옮기지 않고 알린다.
4. 급여대장 PDF 폴더를 `douzone_pdf`(kordoc, Node.js 18 이상)로 급여대장 시트에 붙인다(달마다 사람 수를 「합계(N명)」 과 맞춘다).

계산은 하지 않는다. 묶은 뒤 `python -m ctax employ "<신고 건>"` 을 돌린다.
끝 코드  0 다 묶음, 1 넣지 못한 달 · 옮기지 못한 열이 있음, 2 입력 파일 없음 · 양식이 이미 있음
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

from ..case import Case, MissingInput
from . import douzone_pdf
from .cli import write_template
from .model import COMPANY_ROWS, INPUT_DIR, INPUT_NAME, STAFF_COLS, InputError, _header, _norm


def _company(src: Path, ws_to) -> tuple[int, list[str]]:
    """옮긴 항목 수, 모르는 항목 이름."""
    want = {_norm(r[0]) for r in COMPANY_ROWS}
    rows = {_norm(ws_to.cell(row=r, column=1).value): r for r in range(1, ws_to.max_row + 1)}
    ws = load_workbook(src, data_only=True).worksheets[0]
    n, unknown = 0, []
    for r in range(1, ws.max_row + 1):
        key, val = _norm(ws.cell(row=r, column=1).value), ws.cell(row=r, column=2).value
        if not key or key in ("항목",) or val in (None, ""):
            continue
        if key not in want or key not in rows:
            unknown.append(str(ws.cell(row=r, column=1).value))
            continue
        ws_to.cell(row=rows[key], column=2).value = val
        n += 1
    return n, unknown


def _staff(src: Path, ws_to) -> tuple[int, list[str]]:
    """옮긴 사람 수, 양식에 없어 옮기지 않은 머리글."""
    wb = load_workbook(src, data_only=True)
    ws = wb["사원명부"] if "사원명부" in wb.sheetnames else wb.worksheets[0]
    head, at = _header(ws, STAFF_COLS)
    known = {_norm(c[0]) for c in STAFF_COLS}
    extra = [str(ws.cell(row=head, column=c).value) for c in range(1, ws.max_column + 1)
             if ws.cell(row=head, column=c).value not in (None, "") and _norm(ws.cell(row=head, column=c).value) not in known]
    to_col = {c[1]: j for j, c in enumerate(STAFF_COLS, start=1)}
    out_row, n = 2, 0
    for r in range(head + 1, ws.max_row + 1):
        vals = {attr: ws.cell(row=r, column=c).value for attr, c in at.items()}
        if all(v in (None, "") for v in vals.values()):
            continue
        for attr, v in vals.items():
            cell = ws_to.cell(row=out_row, column=to_col[attr])
            cell.value = v
            if hasattr(v, "year"):
                cell.number_format = "yyyy-mm-dd"
        out_row += 1
        n += 1
    return n, extra


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="ctax.employ.assemble", description="사원명부 · 회사정보 · 급여대장 PDF 를 통합고용 입력 양식 하나로 묶는다")
    ap.add_argument("folder", help="신고 건 폴더(<회사>/<연도> 법인세)")
    ap.add_argument("--staff", required=True, help="사원명부 엑셀(머리글은 양식 「사원명부」 와 같은 이름)")
    ap.add_argument("--company", required=True, help="회사정보 엑셀(A 열 항목 · B 열 값, 항목은 양식 「회사」 와 같은 이름)")
    ap.add_argument("--pdfs", required=True, help="더존 급여대장 PDF 폴더(해당 · 직전 과세연도 달마다 하나)")
    args = ap.parse_args(argv)
    for label, p in (("사원명부", args.staff), ("회사정보", args.company), ("급여대장 PDF 폴더", args.pdfs)):
        if not Path(p).exists():
            print(f"{label} 가 없습니다: {p}")
            return 2
    Path(args.folder).mkdir(parents=True, exist_ok=True)
    try:
        case = Case(args.folder)
    except MissingInput as e:
        print(f"신고 건을 열지 못했습니다: {e}")
        return 2
    dst = case.root / INPUT_DIR / INPUT_NAME
    if dst.exists():
        print(f"{dst} 가 이미 있습니다. 다시 묶으려면 이 파일을 지우고 돌리십시오")
        return 2
    write_template(dst, case)
    wb = load_workbook(dst)
    code = 0
    try:
        n_co, unknown = _company(Path(args.company), wb["회사"])
        n_staff, extra = _staff(Path(args.staff), wb["사원명부"])
    except InputError as e:
        dst.unlink()
        print(f"묶지 못했습니다: {e}")
        return 2
    wb.save(dst)
    print(f"회사 항목 {n_co}개, 사원 {n_staff}명을 옮겼습니다")
    if unknown:
        print(f"양식에 없는 회사 항목이라 옮기지 않음: {' · '.join(unknown)}")
        code = 1
    if extra:
        print(f"양식에 없는 사원명부 열이라 옮기지 않음: {' · '.join(extra)}")
        code = 1
    pdf_code = douzone_pdf.main([args.pdfs, str(dst)])
    print(f"묶은 입력: {dst}")
    print(f'이어서 python -m ctax employ "{args.folder}" 를 돌리십시오')
    return max(code, pdf_code)


if __name__ == "__main__":
    sys.exit(main())

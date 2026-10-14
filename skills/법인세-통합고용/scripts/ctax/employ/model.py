"""통합고용세액공제 입력(`00 받은자료/급여/통합고용_입력.xlsx`)을 읽는다. 원천 데이터 정의는 NOTES.md 1절.

시트 셋: 「회사」(항목 · 값), 「사원명부」(한 사람 한 줄), 「급여대장」(한 사람 한 달 한 줄). 머리글 이름으로 열을 찾는다.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
INPUT_DIR = "00 받은자료/급여"
INPUT_NAME = "통합고용_입력.xlsx"
SIZES = ("중소기업", "중견기업", "일반기업")
KINDS = ("정규직", "기간제", "단시간", "파견")
PERIODS = ("해당", "직전", "직전전")

# 사원명부 열: (머리글, 속성, 종류, 필수, 설명)
STAFF_COLS = [
    ("사번", "id", "text", True, "급여대장과 잇는 키. 없으면 성명과 생년월일로 만든다"),
    ("성명", "name", "text", True, ""),
    ("생년월일", "birth", "date", True, "청년(15 ~ 34세) · 고령자(계약일 현재 60세 이상) 판정"),
    ("입사일", "hired", "date", True, "근로계약 체결일. 재입사자는 줄을 따로 쓴다"),
    ("퇴사일", "left", "date", False, "마지막 근무일. 재직 중이면 비운다"),
    ("내국인", "resident", "yn", False, "소득세법상 거주자면 Y(국적 무관). 비우면 Y"),
    ("고용형태", "kind", "kind", True, "정규직 · 기간제 · 단시간 · 파견. 기간제 · 단시간 · 파견은 청년에서 빠진다"),
    ("계약 1년 미만", "short", "yn", False, "근로계약기간이 1년 미만이면 Y. 갱신으로 총 기간이 1년 이상이면 N"),
    ("월 소정근로시간", "hours", "num", False, "단시간근로자만. 60시간 미만이면 상시근로자가 아니다"),
    ("0.75 요건", "supported", "yn", False, "단시간근로자가 시행령 제11조의2 제8항 제2호 나 ~ 라목(기간의 정함 없는 계약 · 차별 없음 · 최저임금 130%, 중소기업 120% 이상)을 채우면 Y. 가목(상시근로자 수 감소 없음)은 엔진이 확인한다"),
    ("임원", "officer", "yn", False, "법인세법 시행령 제40조 제1항의 임원이면 Y"),
    ("최대주주·친족", "owner", "yn", False, "최대주주 · 최대출자자와 그 배우자 · 직계존비속 · 친족이면 Y"),
    ("기타 제외", "other", "yn", False, "그 밖의 까닭으로 상시근로자가 아니면 Y(일용근로자, 근로계약이 아닌 위임 · 도급 등). 까닭은 비고에"),
    ("입대일", "mil_from", "date", False, "병역 이행 기간(6년 한도)만큼 청년 나이에서 뺀다"),
    ("전역일", "mil_to", "date", False, ""),
    ("장애인 등", "disabled", "yn", False, "장애인 · 상이자 · 5 · 18 부상자 · 고엽제후유의증(장애등급)"),
    ("경력단절", "career", "yn", False, "조세특례제한법 제29조의8 제2항의 경력단절 근로자"),
    ("북한이탈주민", "defector", "yn", False, ""),
    ("수도권 근무지", "capital", "yn", False, "근무지가 서울 · 인천 · 경기면 Y. 비우면 회사 값"),
    ("정규직 전환자", "converted", "yn", False, "명세서 표시용"),
    ("육아휴직 복귀자", "returned", "yn", False, "제29조의8 제5항 요건을 채운 복귀자면 Y"),
    ("비고", "memo", "text", False, ""),
]
PAY_COLS = [
    ("귀속연월", "month", "month", True, "급여가 속하는 달(2025-03 또는 날짜)"),
    ("사번", "id", "text", True, "사원명부의 사번"),
    ("성명", "name", "text", False, ""),
    ("지급액", "pay", "int", True, "그 달 지급 합계(과세 + 비과세)"),
    ("소득세", "tax", "int", False, "원천징수한 근로소득세"),
    ("국민연금", "pension", "int", False, "근로자 부담 공제액"),
    ("건강보험", "health", "int", False, "근로자 부담 공제액(장기요양 포함해도 된다)"),
]
# 회사 항목: (항목, 속성, 종류, 설명)
COMPANY_ROWS = [
    ("과세연도 시작일", "fy_start", "date", "비우면 신고건.json 의 fy_start"),
    ("과세연도 종료일", "fy_end", "date", "비우면 신고건.json 의 fy_end"),
    ("기업 구분", "size", "size", "중소기업 · 중견기업 · 일반기업"),
    ("증가 지역 수도권", "capital", "yn", "인원이 늘어난 곳이 수도권이면 Y(중소기업 1인당 금액이 다르다)"),
    ("소비성서비스업", "consumer", "yn", "호텔 · 여관 · 주점 · 그 밖의 오락 · 유흥 서비스업이면 Y(공제 대상이 아니다)"),
    ("해당 과세연도 창업", "startup", "yn", "Y 면 직전 과세연도 상시근로자 수를 0 으로 본다"),
    ("급여 없는 달 포함", "unpaid", "yn", "비우면 N: 그 과세연도 급여대장이 있는데 그 달 급여가 없으면 근무개월 수에 넣지 않는다. 무급휴직처럼 근로관계가 이어진 달까지 넣으려면 Y"),
    ("청년 간주 적용", "deem", "yn", "비우면 2025년 1월 1일 이후 개시 과세연도에 적용(Y). 근로계약 체결일 기준 청년 간주를 쓰지 않으려면 N"),
    ("직전 1차년도 공제액 청년등", "first_prev_youth", "int", "직전 과세연도 신고서 조특 10호의9 의 1차년도 세액공제액(청년등)"),
    ("직전 1차년도 공제액 청년등 외", "first_prev_other", "int", ""),
    ("직전전 1차년도 공제액 청년등", "first_prev2_youth", "int", "직전전 과세연도 증가분의 1차년도 세액공제액(청년등)"),
    ("직전전 1차년도 공제액 청년등 외", "first_prev2_other", "int", ""),
    ("전기 신고 직전 상시근로자 수", "filed_prev_total", "num", "직전 과세연도 신고서에 적은 그해 상시근로자 수"),
    ("전기 신고 직전 청년등 수", "filed_prev_youth", "num", ""),
    ("전기 신고 직전전 상시근로자 수", "filed_prev2_total", "num", ""),
    ("전기 신고 직전전 청년등 수", "filed_prev2_youth", "num", ""),
    ("육아휴직 복귀자 수", "returned", "num", "비우면 사원명부 「육아휴직 복귀자」 Y 인원"),
]


class InputError(Exception):
    """입력을 읽지 못했다. 명령은 끝 코드 2 로 끝낸다."""


@dataclass
class Person:
    id: str
    name: str
    birth: date
    hired: date
    left: date | None = None
    resident: str = "Y"
    kind: str = "정규직"
    short: str = "N"
    hours: float | None = None
    supported: str = "N"
    officer: str = "N"
    owner: str = "N"
    other: str = "N"
    mil_from: date | None = None
    mil_to: date | None = None
    disabled: str = "N"
    career: str = "N"
    defector: str = "N"
    capital: str = ""
    converted: str = "N"
    returned: str = "N"
    memo: str = ""


@dataclass
class Pay:
    month: date          # 그 달 1일
    id: str
    name: str = ""
    pay: int = 0
    tax: int = 0
    pension: int = 0
    health: int = 0


@dataclass
class Company:
    name: str = ""
    fy_start: date | None = None
    fy_end: date | None = None
    size: str = "중소기업"
    capital: str = "Y"
    consumer: str = "N"
    startup: str = "N"
    deem: str = ""
    unpaid: str = "N"
    first_prev_youth: int = 0
    first_prev_other: int = 0
    first_prev2_youth: int = 0
    first_prev2_other: int = 0
    filed_prev_total: float | None = None
    filed_prev_youth: float | None = None
    filed_prev2_total: float | None = None
    filed_prev2_youth: float | None = None
    returned: float | None = None


@dataclass
class Inputs:
    company: Company
    staff: list[Person]
    pays: list[Pay]
    notes: list[str] = field(default_factory=list)     # 읽으면서 고친 것 · 건너뛴 것


def load_data(path: Path) -> dict:
    """yaml 을 읽는다. 같은 이름의 json 이 옆에 있으면 그것을 읽는다(스킬 꾸러미, PyYAML 없이 돌게)."""
    alt = path.with_suffix(".json")
    if alt.exists():
        return json.loads(alt.read_text(encoding="utf-8"))
    import yaml
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def load_rules() -> dict:
    return load_data(HERE / "rules.yaml")


def rule(rules: dict, group: str, key: str):
    return rules[group][key]["value"]


def basis(rules: dict, group: str, key: str) -> str:
    r = rules[group][key]
    return f"{r['basis']} {r['grade']}"


# ---------------------------------------------------------------- 값 읽기

def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def to_date(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = re.sub(r"[^\d]", "", str(v))
    if len(s) == 8:
        return date(int(s[:4]), int(s[4:6]), int(s[6:]))
    raise ValueError(f"날짜가 아닙니다: {v!r}")


def to_month(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return date(v.year, v.month, 1)
    s = re.sub(r"[^\d]", "", str(v))
    if len(s) in (6, 8):
        return date(int(s[:4]), int(s[4:6]), 1)
    raise ValueError(f"연월이 아닙니다: {v!r}")


def to_yn(v, default: str = "N") -> str:
    s = _text(v).upper()
    if not s:
        return default
    if s in ("Y", "YES", "여", "예", "O", "○", "TRUE", "1", "해당"):
        return "Y"
    if s in ("N", "NO", "부", "아니오", "X", "FALSE", "0", "미해당"):
        return "N"
    raise ValueError(f"Y · N 이 아닙니다: {v!r}")


def to_num(v) -> float | None:
    if v in (None, ""):
        return None
    if isinstance(v, bool):
        raise ValueError(f"숫자가 아닙니다: {v!r}")
    return float(str(v).replace(",", ""))


def to_int(v) -> int:
    n = to_num(v)
    return 0 if n is None else int(round(n))


def to_kind(v) -> str:
    s = _text(v).replace(" ", "")
    alias = {"정규": "정규직", "상용": "정규직", "상용직": "정규직", "무기계약": "정규직", "무기계약직": "정규직",
             "계약직": "기간제", "기간제근로자": "기간제", "단시간근로자": "단시간", "파트타임": "단시간", "알바": "단시간",
             "파견근로자": "파견", "파견직": "파견"}
    s = alias.get(s, s)
    if s not in KINDS:
        raise ValueError(f"고용형태는 {' · '.join(KINDS)} 중 하나입니다: {v!r}")
    return s


def to_size(v) -> str:
    s = _text(v).replace(" ", "")
    s = {"중소": "중소기업", "중견": "중견기업", "일반": "일반기업", "대기업": "일반기업"}.get(s, s)
    if s not in SIZES:
        raise ValueError(f"기업 구분은 {' · '.join(SIZES)} 중 하나입니다: {v!r}")
    return s


CONVERT = {"text": _text, "date": to_date, "month": to_month, "yn": to_yn, "num": to_num, "int": to_int, "kind": to_kind,
           "size": to_size}


def _norm(s) -> str:
    return re.sub(r"[\s·ㆍ・/()]", "", _text(s))


def _header(ws, cols: list, limit: int = 15) -> tuple[int, dict]:
    """머리글 줄과 {속성: 열 번호}. 필수 머리글이 다 있는 첫 줄이 머리글 줄이다."""
    want = {_norm(c[0]): c for c in cols}
    for r in range(1, min(ws.max_row, limit) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            key = _norm(ws.cell(row=r, column=c).value)
            if key in want and want[key][1] not in found:
                found[want[key][1]] = c
        if all(c[1] in found for c in cols if c[3]):
            return r, found
    need = " · ".join(c[0] for c in cols if c[3])
    raise InputError(f"「{ws.title}」 시트에서 머리글 줄을 찾지 못했습니다(필수: {need})")


def _rows(ws, cols: list):
    head, at = _header(ws, cols)
    kinds = {c[1]: c[2] for c in cols}
    names = {c[1]: c[0] for c in cols}
    for r in range(head + 1, ws.max_row + 1):
        raw = {attr: ws.cell(row=r, column=c).value for attr, c in at.items()}
        if all(v in (None, "") for v in raw.values()):
            continue
        out = {}
        for attr, v in raw.items():
            try:
                out[attr] = None if v in (None, "") else CONVERT[kinds[attr]](v)      # 빈 칸은 기본값(Person · Pay)
            except ValueError as e:
                raise InputError(f"「{ws.title}」 {r}행 「{names[attr]}」: {e}") from None
        yield r, out


def read_inputs(path: Path, fy_start: str = "", fy_end: str = "", company: str = "") -> Inputs:
    if not path.exists():
        raise InputError(f"입력 파일이 없습니다: {path}. 「ctax employ <신고 건> --template」 으로 빈 양식을 만드십시오")
    wb = load_workbook(path, data_only=True)
    for name in ("사원명부", "급여대장"):
        if name not in wb.sheetnames:
            raise InputError(f"{path.name} 에 「{name}」 시트가 없습니다")
    notes: list[str] = []
    co = Company(name=company)
    if "회사" in wb.sheetnames:
        ws = wb["회사"]
        spec = {_norm(r[0]): r for r in COMPANY_ROWS}
        for r in range(1, ws.max_row + 1):
            key = _norm(ws.cell(row=r, column=1).value)
            if key not in spec:
                continue
            label, attr, kind, _ = spec[key]
            v = ws.cell(row=r, column=2).value
            if v in (None, ""):
                continue
            try:
                setattr(co, attr, to_yn(v) if kind == "yn" else CONVERT[kind](v))
            except ValueError as e:
                raise InputError(f"「회사」 「{label}」: {e}") from None
    co.fy_start = co.fy_start or to_date(fy_start)
    co.fy_end = co.fy_end or to_date(fy_end)
    if not (co.fy_start and co.fy_end):
        raise InputError("과세연도 시작일 · 종료일이 없습니다(「회사」 시트 또는 신고건.json)")
    if co.fy_start.day != 1:
        raise InputError(f"과세연도 시작일이 1일이 아닙니다({co.fy_start}). 달 중간에 시작하는 과세연도는 다루지 않습니다")

    staff, seen = [], {}
    for r, d in _rows(wb["사원명부"], STAFF_COLS):
        miss = [n for n, a, _, need, _ in STAFF_COLS if need and a != "id" and d.get(a) in (None, "")]
        if miss:
            raise InputError(f"「사원명부」 {r}행: {' · '.join(miss)} 이(가) 비었습니다")
        if not d.get("id"):
            d["id"] = f"{d['name']}_{d['birth']:%Y%m%d}"
            notes.append(f"사원명부 {r}행 {d['name']}: 사번이 없어 「{d['id']}」 로 붙였습니다")
        if d["id"] in seen:
            raise InputError(f"「사원명부」 {r}행: 사번 {d['id']} 이(가) {seen[d['id']]}행과 겹칩니다. 재입사자는 사번을 달리 적으십시오")
        seen[d["id"]] = r
        if d.get("left") and d["left"] < d["hired"]:
            raise InputError(f"「사원명부」 {r}행 {d['name']}: 퇴사일이 입사일보다 앞섭니다")
        if bool(d.get("mil_from")) != bool(d.get("mil_to")):
            raise InputError(f"「사원명부」 {r}행 {d['name']}: 입대일과 전역일은 함께 적습니다")
        d = {k: v for k, v in d.items() if v not in (None, "")}
        d.setdefault("resident", "Y")
        staff.append(Person(**d))
    if not staff:
        raise InputError("「사원명부」 에 사람이 없습니다")

    pays = []
    for r, d in _rows(wb["급여대장"], PAY_COLS):
        if not d.get("month") or not d.get("id"):
            raise InputError(f"「급여대장」 {r}행: 귀속연월 또는 사번이 비었습니다")
        pays.append(Pay(**{k: v for k, v in d.items() if v not in (None, "")}))
    return Inputs(company=co, staff=staff, pays=pays, notes=notes)

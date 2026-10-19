"""시험용 합성 회사(스무 명) 입력을 만든다. `python -m ctax employ "<신고 건>" --sample`.

실제 회사 · 사람과 무관한 값이다. 2025 과세연도, 중소기업, 수도권. 급여대장은 2024 · 2025 두 해이고 직전전(2023) 인원은
전기 신고 인원으로 넣었다. 사람마다 한 가지씩 판정 경우를 넣었고 기대값은 README.md 「시험용 예시」 와 시험 `test_sample_twenty` 에 있다.
"""
from __future__ import annotations

from datetime import date as D
from pathlib import Path

from openpyxl import Workbook

from .calc import month_end
from .model import COMPANY_ROWS, PAY_COLS, STAFF_COLS

COMPANY = {"과세연도 시작일": D(2025, 1, 1), "과세연도 종료일": D(2025, 12, 31), "기업 구분": "중소기업", "증가 지역 수도권": "Y",
           "소비성서비스업": "N", "해당 과세연도 창업": "N",
           "직전 1차년도 공제액 청년등": 29_000_000, "직전 1차년도 공제액 청년등 외": 8_500_000,
           "직전전 1차년도 공제액 청년등": 14_500_000, "직전전 1차년도 공제액 청년등 외": 0,
           "전기 신고 직전전 상시근로자 수": 9.5, "전기 신고 직전전 청년등 수": 5.0}

# 사번, 성명, 생년월일, 입사일, 퇴사일, 고용형태, 월 급여, 그 밖의 열, 무엇을 보는 사람인지
STAFF = [
    ("S01", "강도윤", D(1975, 4, 2), D(2018, 3, 1), None, "정규직", 7_000_000, {"임원": "Y", "최대주주·친족": "Y"}, "대표이사(임원 · 최대주주)"),
    ("S02", "문서연", D(1978, 8, 17), D(2019, 1, 2), None, "정규직", 3_800_000, {"최대주주·친족": "Y"}, "최대주주의 배우자"),
    ("S03", "박준호", D(1980, 5, 20), D(2018, 3, 1), None, "정규직", 5_200_000, {}, "청년등 외, 두 해 모두 열두 달"),
    ("S04", "이수민", D(1987, 9, 9), D(2019, 7, 1), None, "정규직", 4_300_000, {"육아휴직 복귀자": "Y"}, "육아휴직 복귀자(추가공제 1명)"),
    ("S05", "김하늘", D(1996, 3, 15), D(2022, 1, 10), None, "정규직", 3_400_000, {}, "청년, 두 해 모두 열두 달"),
    ("S06", "정우진", D(1998, 11, 2), D(2023, 6, 1), None, "정규직", 3_100_000, {}, "청년"),
    ("S07", "최예린", D(1990, 8, 20), D(2021, 4, 1), None, "정규직", 4_000_000, {}, "2025년 8월에 35세. 그 달부터 청년이 아니다"),
    ("S08", "한지호", D(1999, 1, 25), D(2025, 3, 10), None, "정규직", 2_900_000, {}, "2025년 3월 입사 청년(열 달)"),
    ("S09", "오세영", D(1990, 12, 5), D(2025, 2, 3), None, "정규직", 3_900_000, {}, "2025년 계약 때 34세. 12월에 35세지만 청년 간주 3년"),
    ("S10", "윤태경", D(1985, 2, 14), D(2025, 5, 12), None, "정규직", 4_800_000, {}, "2025년 5월 입사 청년등 외(여덟 달)"),
    ("S11", "서다은", D(1997, 7, 7), D(2024, 9, 2), None, "정규직", 3_000_000, {}, "2024년 9월 입사 청년(직전 넉 달)"),
    ("S12", "임현우", D(1995, 10, 10), D(2022, 8, 1), D(2025, 6, 20), "정규직", 3_300_000, {}, "2025년 6월 20일 퇴사(월말 재직은 5월까지)"),
    ("S13", "조미라", D(1982, 4, 4), D(2020, 2, 1), D(2024, 12, 31), "정규직", 4_500_000, {}, "2024년 말 퇴사"),
    ("S14", "백승현", D(2000, 6, 6), D(2025, 7, 1), D(2025, 12, 31), "기간제", 2_600_000, {"계약 1년 미만": "Y"}, "근로계약기간 1년 미만"),
    ("S15", "신가영", D(1999, 9, 9), D(2024, 3, 1), None, "기간제", 2_800_000, {}, "1년 이상 기간제. 상시근로자지만 청년에서 빠진다"),
    ("S16", "권동민", D(1978, 1, 1), D(2023, 5, 1), None, "단시간", 1_400_000, {"월 소정근로시간": 80}, "단시간 80시간(0.5명)"),
    ("S17", "황유나", D(2001, 2, 2), D(2025, 1, 6), None, "단시간", 600_000, {"월 소정근로시간": 40}, "단시간 40시간(상시근로자가 아니다)"),
    ("S18", "안병철", D(1962, 3, 3), D(2024, 5, 1), None, "정규직", 3_200_000, {}, "계약 때 62세(고령자, 청년등)"),
    ("S19", "송지아", D(1984, 12, 12), D(2021, 11, 1), None, "정규직", 3_600_000, {"장애인 등": "Y"}, "장애인(청년등)"),
    ("S20", "류성훈", D(1989, 6, 30), D(2023, 2, 1), None, "정규직", 4_100_000,
     {"입대일": D(2010, 1, 11), "전역일": D(2011, 10, 10)}, "병역 21개월을 빼면 2025년 내내 청년 나이"),
]
YEARS = (2024, 2025)


def _shift(d, years: int):
    return None if d is None else d.replace(year=d.year + years)


def numeric_id(sid: str) -> str:
    """S01 → 1001. 더존 급여대장 PDF 는 숫자 사번만 찍으므로 여러 파일 시연은 이 사번을 쓴다."""
    return "10" + sid[1:]


def _people(year: int, numeric: bool = False):
    delta = year - 2025
    staff = [(numeric_id(sid) if numeric else sid, name, _shift(b, delta), _shift(h, delta), _shift(l, delta), kind, pay,
              {k: (_shift(v, delta) if isinstance(v, D) else v) for k, v in more.items()}, memo)
             for sid, name, b, h, l, kind, pay, more, memo in STAFF]
    company = {k: (_shift(v, delta) if isinstance(v, D) else v) for k, v in COMPANY.items()}
    if year >= 2027:     # 직전(2026) 과세연도에 늘어난 인원은 신규공제라 1차년도 공제액이 없다. 직전전(2025) 1차 공제는 3차 연도로 이어진다
        company["직전 1차년도 공제액 청년등"] = company["직전 1차년도 공제액 청년등 외"] = None
    return staff, company, tuple(y + delta for y in YEARS)


def pay_rows(year: int = 2025, numeric: bool = False) -> list[dict]:
    """급여대장 줄(달 순서, 사번 순서). write · write_douzone_pdfs · write_payroll_excels 가 같은 값을 쓴다."""
    staff, _, years = _people(year, numeric)
    out = []
    for year_ in years:
        for month in range(1, 13):
            first, last = D(year_, month, 1), month_end(D(year_, month, 1))
            for sid, name, _, hired, left, _, pay, _, _ in staff:
                if hired <= last and (left is None or left >= first):
                    out.append({"month": f"{year_}-{month:02d}", "id": sid, "name": name, "hired": hired, "left": left, "pay": pay,
                                "tax": int(pay * 0.03) // 10 * 10, "pension": int(pay * 0.045) // 10 * 10,
                                "health": int(pay * 0.03545) // 10 * 10})
    return out


def write(path: Path, year: int = 2025, payroll: bool = True, numeric: bool = False) -> Path:
    """year 는 해당 과세연도. 2026 이면 모든 날짜를 한 해 뒤로 옮긴 같은 스무 명이다(급여대장은 2025 · 2026, 직전전 2024 는 전기 신고 인원).
    payroll=False 면 급여대장 시트를 머리글만 남긴다(여러 파일 시연에서 급여대장을 따로 붙일 때). numeric=True 면 사번이 1001 ~ 1020 이다."""
    staff, company, years = _people(year, numeric)
    wb = Workbook()
    ws = wb.active
    ws.title = "회사"
    ws.append(["항목", "값", "설명"])
    for label, _, _, text in COMPANY_ROWS:
        ws.append([label, company.get(label), text])
    ws = wb.create_sheet("사원명부")
    head = [c[0] for c in STAFF_COLS]
    ws.append(head)
    for sid, name, birth, hired, left, kind, _, more, memo in staff:
        row = {"사번": sid, "성명": name, "생년월일": birth, "입사일": hired, "퇴사일": left, "고용형태": kind, "비고": memo, **more}
        ws.append([row.get(h) for h in head])
    ws = wb.create_sheet("급여대장")
    ws.append([c[0] for c in PAY_COLS])
    if payroll:
        for r in pay_rows(year, numeric):
            ws.append([r["month"], r["id"], r["name"], r["pay"], r["tax"], r["pension"], r["health"]])
    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, D):
                    c.number_format = "yyyy-mm-dd"
                elif isinstance(c.value, int) and c.value >= 1000:
                    c.number_format = "#,##0"
    wb["회사"].column_dimensions["A"].width = 34
    wb["회사"].column_dimensions["C"].width = 90
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _by_month(rows: list[dict]) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    for r in rows:
        out.setdefault(r["month"], []).append(r)
    return out


DOUZONE_HEAD = (                                   # 더존 급여대장 머리글 네 줄(실제 양식에서 칸 위치만 따왔다)
    ("인적사항", "", "기본급여및제수당", "", "", "공제및차인지급액", "", "", "", ""),
    ("사원번호", "성 명", "기본급", "식대", "", "국민연금", "건강보험", "고용보험", "소득세", "지방소득세"),
    ("입사일", "직 급", "", "", "", "건강보험정산", "장기요양정산", "기타공제", "", "공제합계"),
    ("퇴사일", "부 서", "", "", "지급합계", "", "", "", "", "차인지급액"),
)


def write_douzone_pdfs(folder: Path, year: int = 2025) -> list[Path]:
    """더존 「급여대장」 PDF 를 흉내 낸 달별 PDF. 칸마다 선을 그은 표이고 머리글 네 줄 · 사람마다 세 줄이다(실제 더존 양식과 같은 짜임).
    첫 줄: 사원번호 · 성명 · 기본급 · 식대 · 국민연금 · 건강보험 · 고용보험 · 소득세 · 지방소득세. 둘째 줄: 입사일 · 공제합계.
    셋째 줄: 퇴사일 · 지급합계 · 차인지급액. 마지막 쪽 끝에 「합계(N명)」 줄. 한 쪽에 열한 명씩(더존과 같다)이라 열한 명이 넘는 달은 두 쪽이다."""
    import fitz
    folder.mkdir(parents=True, exist_ok=True)
    x0, y0, w, h, cols = 30, 90, 78, 12, len(DOUZONE_HEAD[0])
    per_page, out = 11, []
    for month, rows in _by_month(pay_rows(year, numeric=True)).items():
        doc = fitz.open()
        for start in range(0, len(rows), per_page):
            pg = doc.new_page(width=842, height=595)
            pg.insert_text((x0, 50), f"{month[:4]}년{month[5:]}월분 급여대장", fontname="korea", fontsize=14)
            pg.insert_text((x0, 68), "회사명: 샘플회사(가상)", fontname="korea", fontsize=8)
            pg.insert_text((x0, 80), f"[귀속:{month[:4]}년{month[5:]}월] [지급:{month[:4]}년{month[5:]}월25일]", fontname="korea", fontsize=8)
            grid = [list(r) for r in DOUZONE_HEAD]
            for r in rows[start:start + per_page]:
                employ = int(r["pay"] * 0.009) // 10 * 10
                local = r["tax"] // 100 * 10
                deduct = r["pension"] + r["health"] + employ + r["tax"] + local
                meal = min(200_000, r["pay"])
                grid.append([r["id"], r["name"], f"{r['pay'] - meal:,}", f"{meal:,}", "", f"{r['pension']:,}", f"{r['health']:,}",
                             f"{employ:,}", f"{r['tax']:,}" if r["tax"] else "", f"{local:,}" if local else ""])
                grid.append([r["hired"].isoformat()] + [""] * (cols - 2) + [f"{deduct:,}"])
                grid.append([r["left"].isoformat() if r["left"] else "", "", "", "", f"{r['pay']:,}", "", "", "", "",
                             f"{r['pay'] - deduct:,}"])
            if start + per_page >= len(rows):
                grid.append([f"합계({len(rows)}명)", "", "", "", f"{sum(r['pay'] for r in rows):,}", "", "", "", "", ""])
            for i, line in enumerate(grid):
                for j, text in enumerate(line):
                    if text:
                        pg.insert_text((x0 + j * w + 2, y0 + i * h + h - 3.5), str(text), fontname="korea", fontsize=7)
            for i in range(len(grid) + 1):
                pg.draw_line((x0, y0 + i * h), (x0 + cols * w, y0 + i * h), width=0.4)
            for j in range(cols + 1):
                pg.draw_line((x0 + j * w, y0), (x0 + j * w, y0 + len(grid) * h), width=0.4)
        path = folder / f"급여대장_{month}.pdf"
        doc.save(path)
        out.append(path)
    return out


def write_payroll_excels(folder: Path, year: int = 2025) -> list[Path]:
    """달별 엑셀 급여대장. 직전 과세연도와 해당 과세연도의 열 이름 · 순서 · 머리 줄 위치를 다르게 했다(급여 프로그램을 바꾼 회사).
    맨 아래에 합계 줄이 있어 옮긴 뒤 원본 합계와 맞춰 볼 수 있다."""
    folder.mkdir(parents=True, exist_ok=True)
    rows_all = pay_rows(year, numeric=True)
    first_year = min(r["month"][:4] for r in rows_all)
    out = []
    for month, rows in _by_month(rows_all).items():
        wb = Workbook()
        ws = wb.active
        ws.title = "급여"
        if month[:4] == first_year:
            ws.append(["사원코드", "사원명", "귀속월", "총지급액", "근로소득세", "국민연금", "건강보험", "고용보험"])
            for r in rows:
                ws.append([r["id"], r["name"], month, r["pay"], r["tax"], r["pension"], r["health"], int(r["pay"] * 0.009) // 10 * 10])
            ws.append(["합계", f"{len(rows)}명", None, sum(r["pay"] for r in rows), sum(r["tax"] for r in rows),
                       sum(r["pension"] for r in rows), sum(r["health"] for r in rows), None])
        else:
            ws.append([f"{month[:4]}년 {int(month[5:])}월 급여대장"])
            ws.append(["샘플회사(가상)", None, None, f"지급일 {month}-25"])
            ws.append(["사번", "성명", "지급합계", "소득세", "지방소득세", "국민연금(근로자)", "건강보험(근로자, 장기요양 포함)", "차인지급액"])
            for r in rows:
                ded = r["tax"] + r["tax"] // 100 * 10 + r["pension"] + r["health"]
                ws.append([r["id"], r["name"], r["pay"], r["tax"], r["tax"] // 100 * 10, r["pension"], r["health"], r["pay"] - ded])
            ws.append(["합계", f"{len(rows)}명", sum(r["pay"] for r in rows), sum(r["tax"] for r in rows), None,
                       sum(r["pension"] for r in rows), sum(r["health"] for r in rows), None])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 16
        path = folder / f"급여대장_{month}.xlsx"
        wb.save(path)
        out.append(path)
    return out


def write_staff_file(path: Path, year: int = 2025) -> Path:
    """인사 담당이 따로 주는 사원명부 엑셀(숫자 사번, 퇴사자 포함). 머리글은 입력 양식 「사원명부」 와 같은 이름이다."""
    staff, _, _ = _people(year, numeric=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "사원명부"
    ws.append([f"샘플회사(가상) 사원명부 · {year - 1} ~ {year} 재직자 전원"])
    head = [c[0] for c in STAFF_COLS]
    ws.append(head)
    for sid, name, birth, hired, left, kind, _, more, memo in staff:
        row = {"사번": sid, "성명": name, "생년월일": birth, "입사일": hired, "퇴사일": left, "고용형태": kind, "비고": memo, **more}
        ws.append([row.get(h) for h in head])
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if isinstance(c.value, D):
                c.number_format = "yyyy-mm-dd"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_company_file(path: Path, year: int = 2025) -> Path:
    """회사정보 엑셀(A 열 항목 · B 열 값). 항목 이름은 입력 양식 「회사」 시트와 같다. 값이 없는 항목은 적지 않았다."""
    _, company, _ = _people(year)
    wb = Workbook()
    ws = wb.active
    ws.title = "회사정보"
    ws.append(["항목", "값"])
    for label, *_ in COMPANY_ROWS:
        if company.get(label) is not None:
            ws.append([label, company[label]])
            if isinstance(company[label], D):
                ws.cell(row=ws.max_row, column=2).number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

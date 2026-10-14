"""통합고용세액공제 검증 워크북(`2 조서/통합고용_검증.xlsx`)을 쓴다. 숫자는 모두 엑셀 수식이다.

시트 순서가 계산 순서다: 0기준 → 1사원명부 → 2급여대장 → 3판정(해당 · 직전 · 직전전) → 4인원 → 5공제액 → 6서식10의9 → 7서식10의10.
「8엔진대조」 는 calc.py 가 따로 계산한 값을 수식 값 옆에 놓는다. 노란 칸만 사람이 고친다.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .. import xlsx as X
from .calc import FAR, REASONS
from .model import PERIODS, STAFF_COLS, Inputs, basis, rule

S_LIST, S_BASE, S_STAFF, S_PAY = "목록", "0기준", "1사원명부", "2급여대장"
S_JUDGE = {"해당": "3판정_해당", "직전": "3판정_직전", "직전전": "3판정_직전전"}
S_COUNT, S_CREDIT, S_FORM, S_ROSTER, S_CHECK = "4인원", "5공제액", "6서식10의9", "7서식10의10", "8엔진대조"

DATE = "yyyy-mm-dd"
YMD = "yyyymmdd"
NUM2 = "0.00"
STAFF_FIRST, PAY_FIRST, JUDGE_FIRST = 5, 5, 11
BLOCK = 12                      # 판정 시트 한 묶음의 달 수
B_PAY, B_WORK, B_Y1, B_Y2 = 24, 36, 48, 60   # 묶음이 시작하는 열
MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
FORM_FONT = Font(name=X.FONT, size=9)
FORM_BOLD = Font(name=X.FONT, size=10, bold=True)
BLACK = Side(style="thin", color="000000")
FORM_BOX = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
GREY = PatternFill("solid", fgColor="F2F2F2")


def ref(sheet: str, addr: str) -> str:
    return f"'{sheet}'!{addr}"


def absolute(col: int, row: int) -> str:
    return f"${L(col)}${row}"


class KV:
    """항목 · 값 · 설명 세 열로 쓰는 시트. key 로 값 칸 주소를 찾는다."""

    def __init__(self, ws, row: int):
        self.ws, self.row, self.at = ws, row, {}

    def section(self, text: str) -> None:
        self.row += 1
        c = self.ws.cell(row=self.row, column=1, value=text)
        c.font = Font(name=X.FONT, bold=True, size=11)
        self.row += 1

    def put(self, key: str, label: str, value, kind: str = "body", fmt: str = "", note: str = "") -> str:
        X.cell(self.ws, self.row, 1, label, wrap=True)
        c = X.cell(self.ws, self.row, 2, value, kind=kind)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="right" if fmt else "center", vertical="center")
        X.cell(self.ws, self.row, 3, note, wrap=True).font = X.NOTE_FONT
        self.at[key] = ref(self.ws.title, absolute(2, self.row))
        self.row += 1
        return self.at[key]

    def __getitem__(self, key: str) -> str:
        return self.at[key]


def build(inp: Inputs, res: dict, rules: dict, rates: dict, memos: list[str]) -> tuple[Workbook, dict]:
    """워크북과 {이름: (시트, 칸)} 대조 목록. `res` 는 calc.compute 결과, `memos` 는 확인사항 글."""
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {n: wb.create_sheet(n) for n in
              [S_LIST, S_BASE, S_STAFF, S_PAY, *S_JUDGE.values(), S_COUNT, S_CREDIT, S_FORM, S_ROSTER, S_CHECK]}
    n = len(inp.staff)
    staff_last = STAFF_FIRST + n - 1
    pays = sorted(inp.pays, key=lambda x: (x.month, x.id))
    pay_last = PAY_FIRST + max(len(pays), 1) - 1
    judge_last = JUDGE_FIRST + n - 1

    base = _base(sheets[S_BASE], inp, res, rules, rates)
    _staff(sheets[S_STAFF], inp, base, staff_last)
    pay = _pay(sheets[S_PAY], pays, pay_last)
    for per in PERIODS:
        _judge(sheets[S_JUDGE[per]], per, base, pay, n, judge_last, len(pays))
    cnt = _count(sheets[S_COUNT], base, judge_last)
    cr = _credit(sheets[S_CREDIT], base, cnt, rules, rates)
    form = _form(sheets[S_FORM], base, cnt, cr)
    rows = _roster(sheets[S_ROSTER], inp, res, base, cnt, judge_last)
    checks = _check(sheets[S_CHECK], inp, res, cnt, cr, form)
    _list(sheets[S_LIST], inp, res, base, cnt, cr, form, rows, memos, staff_last, pay_last, len(pays))
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.active = 0
    return wb, checks


# ---------------------------------------------------------------- 0기준

def _base(ws, inp: Inputs, res: dict, rules: dict, rates: dict) -> KV:
    co = inp.company
    row = X.title(ws, "0. 기준", "회사 정보 · 계산 방식 · 법령 금액입니다. 노란 칸만 고칩니다. 나머지 시트는 이 값을 읽습니다")
    X.header(ws, row, ["항목", "값", "설명 · 근거"])
    k = KV(ws, row + 1)
    r = lambda g, key: rule(rules, g, key)          # noqa: E731
    e = rates["employment_integrated"]

    k.section("회사와 과세연도")
    k.put("company", "상호 또는 법인명", co.name, kind="input")
    k.put("fy_start", "과세연도 시작일", co.fy_start, kind="input", fmt=DATE)
    k.put("fy_end", "과세연도 종료일", co.fy_end, kind="input", fmt=DATE)
    months = lambda s, t: f"=(YEAR({t})-YEAR({s}))*12+MONTH({t})-MONTH({s})+1"   # noqa: E731
    k.put("months", "과세연도 개월 수", months(k["fy_start"], k["fy_end"]), fmt="0")
    k.put("p_start", "직전 과세연도 시작일", f"=EDATE({k['fy_start']},-12)", kind="input", fmt=DATE,
          note="직전 과세연도가 12개월이 아니면 고칩니다")
    k.put("p_end", "직전 과세연도 종료일", f"={k['fy_start']}-1", kind="input", fmt=DATE)
    k.put("p_months", "직전 과세연도 개월 수", months(k["p_start"], k["p_end"]), fmt="0")
    k.put("pp_start", "직전전 과세연도 시작일", f"=EDATE({k['p_start']},-12)", kind="input", fmt=DATE)
    k.put("pp_end", "직전전 과세연도 종료일", f"={k['p_start']}-1", kind="input", fmt=DATE)
    k.put("pp_months", "직전전 과세연도 개월 수", months(k["pp_start"], k["pp_end"]), fmt="0")
    k.put("size", "기업 구분", co.size, kind="input", note="중소기업 · 중견기업 · 일반기업. 중소기업 여부는 51호 기준검토표와 맞춥니다")
    k.put("capital", "증가 지역 수도권", co.capital, kind="input",
          note="Y · N. 중소기업이 수도권 밖에서 늘었으면 N(조세특례제한법 제29조의8 제1항)")
    k.put("consumer", "소비성서비스업", co.consumer, kind="input",
          note="Y 면 공제 대상이 아닙니다(조세특례제한법 시행령 제26조의8 제1항 · 제29조 제3항)")
    k.put("startup", "해당 과세연도 창업", co.startup, kind="input", note=basis(rules, "count", "startup_prior_zero"))
    k.put("unpaid", "급여 없는 달 포함", co.unpaid, kind="input", note=basis(rules, "count", "unpaid_month"))

    k.section("계산 방식")
    k.put("method", "1차 연도 인원 계산 방식", f'=IF({k["fy_start"]}>=DATE(2025,1,1),"근로자별","월말합")',
          note=basis(rules, "count", "per_worker"))
    k.put("in_time", "1인당 공제금액 적용", f'=IF(YEAR({k["fy_end"]})<=2025,"Y","N")', note=e["first_year_until"]["basis"])
    k.put("deem", "해당 과세연도 청년 간주", "N" if co.deem == "N" else f'=IF({k["method"]}="근로자별","Y","N")', kind="input",
          note="쓰지 않으려면 N 으로 고칩니다. " + basis(rules, "youth", "deemed_from"))
    y25, yo = r("youth", "deemed_years_2025"), r("youth", "deemed_years_other")
    pick = lambda d: (f'=IF({k["size"]}="중소기업",{d["중소기업"]},IF({k["size"]}="중견기업",{d["중견기업"]},'   # noqa: E731
                      f'{d["일반기업"]}))')
    k.put("deem_2025", "청년 간주 연수(2025년 중 계약)", pick(y25), fmt="0", note=basis(rules, "youth", "deemed_years_2025"))
    k.put("deem_other", "청년 간주 연수(그 밖의 계약)", pick(yo), fmt="0", note=basis(rules, "youth", "deemed_years_other"))
    k.put("age_from", "청년 나이 아래 끝", r("youth", "age_from"), fmt="0", note=basis(rules, "youth", "age_from"))
    k.put("age_to", "청년 나이 위 끝", r("youth", "age_to"), fmt="0", note=basis(rules, "youth", "age_to"))
    k.put("mil_cap", "병역 기간 한도(일)", r("youth", "military_cap_days"), fmt="0", note=basis(rules, "youth", "military_cap_days"))
    k.put("senior", "고령자 나이", r("youth", "senior_age"), fmt="0", note=basis(rules, "youth", "senior_age"))
    k.put("hours", "단시간근로자 소정근로시간 기준", r("scope", "part_time_hours"), fmt="0", note=basis(rules, "scope", "part_time_hours"))
    k.put("w_half", "단시간근로자 1명", r("scope", "part_time_weight"), fmt=NUM2, note=basis(rules, "scope", "part_time_weight"))
    k.put("w_sup", "단시간근로자 1명(지원요건 충족)", r("scope", "part_time_weight_supported"), fmt=NUM2,
          note=basis(rules, "scope", "part_time_weight_supported"))

    k.section("1인당 공제금액(법령 값, 고치지 않습니다)")
    for key, label in (("youth_sme_capital", "청년등 중소기업 수도권"), ("youth_sme_local", "청년등 중소기업 수도권 밖"),
                       ("youth_mid", "청년등 중견기업"), ("youth_general", "청년등 일반기업"),
                       ("other_sme_capital", "청년등 외 중소기업 수도권"), ("other_sme_local", "청년등 외 중소기업 수도권 밖"),
                       ("other_mid", "청년등 외 중견기업"), ("other_general", "청년등 외 일반기업")):
        k.put(key, label, int(e[key]["value"]), fmt=X.MONEY, note=f"{e[key]['basis']} {e[key]['grade']}")
    k.put("return_sme", "육아휴직 복귀자 중소기업", int(r("extra", "return_sme")), fmt=X.MONEY, note=basis(rules, "extra", "return_sme"))
    k.put("return_mid", "육아휴직 복귀자 중견기업", int(r("extra", "return_mid")), fmt=X.MONEY, note=basis(rules, "extra", "return_mid"))

    k.section("전기 신고서에서 옮기는 값")
    src = "직전 과세연도 신고서 조특 10호의9 「1차년도 세액공제액」. 당기 2차 연도 공제액이 됩니다"
    k.put("f1y", "직전 과세연도 증가분 1차년도 공제액(청년등)", co.first_prev_youth, kind="input", fmt=X.MONEY, note=src)
    k.put("f1o", "직전 과세연도 증가분 1차년도 공제액(청년등 외)", co.first_prev_other, kind="input", fmt=X.MONEY)
    src = "직전전 과세연도 증가분의 1차년도 세액공제액. 당기 3차 연도 공제액이 됩니다(중소 · 중견기업)"
    k.put("f2y", "직전전 과세연도 증가분 1차년도 공제액(청년등)", co.first_prev2_youth, kind="input", fmt=X.MONEY, note=src)
    k.put("f2o", "직전전 과세연도 증가분 1차년도 공제액(청년등 외)", co.first_prev2_other, kind="input", fmt=X.MONEY)
    src = "그 과세연도 급여대장이 없을 때 이 인원을 씁니다. 급여대장이 있으면 다시 센 인원과 맞춰 봅니다"
    k.put("filed_p_t", "전기 신고 직전 과세연도 상시근로자 수", co.filed_prev_total, kind="input", fmt=NUM2, note=src)
    k.put("filed_p_y", "전기 신고 직전 과세연도 청년등 수", co.filed_prev_youth, kind="input", fmt=NUM2)
    k.put("filed_pp_t", "전기 신고 직전전 과세연도 상시근로자 수", co.filed_prev2_total, kind="input", fmt=NUM2)
    k.put("filed_pp_y", "전기 신고 직전전 과세연도 청년등 수", co.filed_prev2_youth, kind="input", fmt=NUM2)
    k.put("returned", "육아휴직 복귀자 수", res["returned"], kind="input", fmt=NUM2,
          note="해당 과세연도에 복직한 사람 중 제29조의8 제5항 요건을 채운 인원(자녀 1명당 한 번)")
    X.widths(ws, [44, 18, 110])
    return k


# ---------------------------------------------------------------- 1사원명부 · 2급여대장

STAFF_CALC = ["제외사유(과세연도와 무관)", "단시간 유형", "병역을 뺀 생년월일", "15세가 되는 날", "35세가 되는 날(청년 제외시점, 현재 연령)",
              "계약 체결일 기준 청년", "청년 간주 종료일", "청년 제외시점(간주 포함)", "고령자", "나이와 무관한 청년등", "정규직"]
C_REASON, C_WEIGHT, C_ADJ, C_15, C_35, C_HIRE_Y, C_DEEM, C_CUT, C_SENIOR, C_FIXED, C_REGULAR = range(
    len(STAFF_COLS) + 1, len(STAFF_COLS) + 12)


def _staff(ws, inp: Inputs, k: KV, last: int) -> None:
    row = X.title(ws, "1. 사원명부", "한 사람 한 줄입니다. A ~ V 열은 입력(노란 칸), W 열부터는 판정 수식입니다. 재입사자는 줄을 따로 씁니다")
    head = [c[0] for c in STAFF_COLS] + STAFF_CALC
    X.header(ws, row, head)
    ws.row_dimensions[row].height = 44
    attrs = [c[1] for c in STAFF_COLS]
    kinds = {c[1]: c[2] for c in STAFF_COLS}
    for i, p in enumerate(inp.staff):
        r = STAFF_FIRST + i
        for j, a in enumerate(attrs, start=1):
            v = getattr(p, a)
            c = X.cell(ws, r, j, v if v not in (None, "") else None, kind="input")
            if kinds[a] == "date":
                c.number_format = DATE
        a = lambda col: f"${L(col)}{r}"      # noqa: E731
        A = {name: a(attrs.index(name) + 1) for name in attrs}
        f = {
            C_REASON: (f'=IF({A["resident"]}="N","{REASONS["f_res"]}",IF({A["kind"]}="파견","{REASONS["f_disp"]}",'
                       f'IF({A["short"]}="Y","{REASONS["a"]}",IF(AND({A["kind"]}="단시간",N({A["hours"]})<{k["hours"]}),'
                       f'"{REASONS["b"]}",IF({A["officer"]}="Y","{REASONS["c"]}",IF({A["owner"]}="Y","{REASONS["d"]}",'
                       f'IF({A["other"]}="Y","{REASONS["f"]}","")))))))'),
            C_WEIGHT: (f'=IF(AND({A["kind"]}="단시간",N({A["hours"]})>={k["hours"]}),IF({A["supported"]}="Y",{k["w_sup"]},'
                       f'{k["w_half"]}),1)'),
            C_ADJ: (f'={A["birth"]}+IF(AND({A["mil_from"]}<>"",{A["mil_to"]}<>""),'
                    f'MAX(0,MIN({A["mil_to"]}-{A["mil_from"]}+1,{k["mil_cap"]})),0)'),
            C_15: f'=EDATE({A["birth"]},12*{k["age_from"]})',
            C_35: f'=EDATE({a(C_ADJ)},12*({k["age_to"]}+1))',
            C_HIRE_Y: f'=IF(AND({A["kind"]}="정규직",{A["hired"]}>={a(C_15)},{A["hired"]}<{a(C_35)}),"Y","N")',
            C_DEEM: (f'=IF(AND({a(C_HIRE_Y)}="Y",{A["hired"]}>=DATE(2025,1,1)),EDATE({A["hired"]},12*IF(YEAR({A["hired"]})=2025,{k["deem_2025"]},'
                     f'{k["deem_other"]})),"")'),
            C_CUT: f'=IF(ISNUMBER({a(C_DEEM)}),MAX({a(C_35)},{a(C_DEEM)}),{a(C_35)})',
            C_SENIOR: f'=IF({A["hired"]}>=EDATE({A["birth"]},12*{k["senior"]}),"Y","N")',
            C_FIXED: (f'=IF(OR({A["disabled"]}="Y",{a(C_SENIOR)}="Y",{A["career"]}="Y",{A["defector"]}="Y"),"Y","N")'),
            C_REGULAR: f'=IF({A["kind"]}="정규직","Y","N")',
        }
        for col, formula in f.items():
            c = X.cell(ws, r, col, formula)
            if col in (C_ADJ, C_15, C_35, C_DEEM, C_CUT):
                c.number_format = DATE
            elif col == C_WEIGHT:
                c.number_format = NUM2
    ws.freeze_panes = ws.cell(row=STAFF_FIRST, column=3)
    X.widths(ws, [10, 10, 12, 12, 12, 7, 9, 8, 9, 8, 6, 9, 8, 12, 12, 8, 8, 9, 8, 8, 8, 16, 26, 8, 12, 12, 12, 9, 12, 12, 7, 9, 7])
    X.note(ws, last + 2, "열 뜻과 근거는 스킬 문서 「원천 데이터 정의」 와 ctax/employ/NOTES.md 에 있습니다")


def _pay(ws, pays: list, last: int) -> dict:
    row = X.title(ws, "2. 급여대장", "한 사람 한 달 한 줄입니다. 귀속연월은 그 달 1일 날짜로 둡니다. 소득세 · 국민연금 · 건강보험은 납부 사실 확인에 씁니다")
    X.header(ws, row, ["귀속연월", "사번", "성명", "지급액", "소득세", "국민연금", "건강보험"])
    for i, x in enumerate(pays):
        r = PAY_FIRST + i
        X.cell(ws, r, 1, x.month, kind="input").number_format = "yyyy-mm"
        X.cell(ws, r, 2, x.id, kind="input")
        X.cell(ws, r, 3, x.name, kind="input")
        for j, v in enumerate((x.pay, x.tax, x.pension, x.health), start=4):
            X.cell(ws, r, j, v, kind="input", money=True)
    ws.freeze_panes = ws.cell(row=PAY_FIRST, column=1)
    X.widths(ws, [11, 12, 12, 14, 12, 12, 12])
    rng = lambda col: ref(S_PAY, f"${L(col)}${PAY_FIRST}:${L(col)}${last}")   # noqa: E731
    return {"month": rng(1), "id": rng(2), "pay": rng(4), "tax": rng(5), "pension": rng(6), "health": rng(7)}


# ---------------------------------------------------------------- 3판정

J_HEAD = ["사번", "성명", "생년월일", "근로계약 체결일", "퇴사일", "퇴사일(계산용)", "납부 확인", "상시근로자", "제외사유", "단시간 유형",
          "월말 재직 달 수", "상시근로자 근무개월수", "상시근로자별 계산한 수", "청년등 달 수(현재 연령)", "청년등 근무개월수(현재 연령)",
          "계산한 수(현재 연령)", "청년등 달 수(서식 기준)", "청년등 상시근로자 근무개월수", "청년등 상시근로자별 계산한 수", "청년등 구분",
          "재직인데 급여 없는 달", "재직 기간이 아닌데 급여 있는 달"]
(J_ID, J_NAME, J_BIRTH, J_HIRED, J_LEFT, J_LEFTX, J_PROOF, J_REG, J_REASON, J_WEIGHT, J_MW, J_M20, J_N21, J_MY1, J_M22A, J_N23A,
 J_MY2, J_M22, J_N23, J_KIND, J_NOPAY, J_ODD) = range(1, 23)
J_START = {"해당": "fy_start", "직전": "p_start", "직전전": "pp_start"}
J_END = {"해당": "fy_end", "직전": "p_end", "직전전": "pp_end"}
J_MONTHS = {"해당": "months", "직전": "p_months", "직전전": "pp_months"}


def _judge(ws, per: str, k: KV, pay: dict, n: int, last: int, n_pay: int) -> None:
    # 단시간근로자 0.75명 지원요건 가목(시행령 제11조의2 제8항 제2호): 단시간근로자를 뺀 상시근로자 수가 직전 과세연도보다 줄지 않았을 것.
    # 비교할 직전 급여대장이 없으면 사람이 입력한 「0.75 요건」 을 그대로 쓴다
    own = lambda col: f"${L(col)}${JUDGE_FIRST}:${L(col)}${last}"       # noqa: E731
    base_w = ref(S_STAFF, f"${L(C_WEIGHT)}${STAFF_FIRST}:${L(C_WEIGHT)}${STAFF_FIRST + n - 1}")
    core = (f'=SUMPRODUCT(({base_w}=1)*({own(J_REG)}="Y")*ROUNDDOWN(ROUND({own(J_MW)}/$F$4,8),2))' if n else 0)
    i = PERIODS.index(per)
    if i + 1 < len(PERIODS):
        older = S_JUDGE[PERIODS[i + 1]]
        allow = f'=IF({ref(older, "$J$4")}>0,IF($P$4>={ref(older, "$P$4")},"Y","N"),"Y")'
    else:
        allow = "Y"
    row = X.title(ws, f"3. 판정({per} 과세연도)",
                  "사람마다 달마다 월말 재직과 청년등 여부를 가립니다. 「서식 기준」 은 해당 과세연도에만 청년 간주를 넣은 값입니다")
    for j, (label, value, fmt) in enumerate((("과세연도 시작일", f"={k[J_START[per]]}", DATE), ("종료일", f"={k[J_END[per]]}", DATE),
                                             ("개월 수", f"={k[J_MONTHS[per]]}", "0"),
                                             ("청년 간주", f"={k['deem']}" if per == "해당" else "N", ""),
                                             ("급여대장 줄 수", None, "0"), ("급여 없는 달 포함", f"={k['unpaid']}", ""),
                                             ("단시간 0.75 허용", allow, ""), ("단시간 제외 상시근로자 수", core, NUM2))):
        X.cell(ws, 4, 1 + 2 * j, label, kind="total")
        X.cell(ws, 4, 2 + 2 * j, value).number_format = fmt or "General"
    START, END, MONTHS, DEEM, PAYN, UNPAID, ALLOW = "$B$4", "$D$4", "$F$4", "$H$4", "$J$4", "$L$4", "$N$4"
    ws[PAYN.replace("$", "")] = (f'=COUNTIFS({pay["month"]},">="&{START},{pay["month"]},"<="&{END})' if n_pay else 0)
    first = JUDGE_FIRST
    head = first - 1
    X.header(ws, head, J_HEAD)
    ws.row_dimensions[head].height = 58
    blocks = ((B_PAY, "그 달 지급액(급여대장)"), (B_WORK, "근무개월로 세는 달(월말 재직이고 급여 있음)"), (B_Y1, "청년등(현재 연령)"), (B_Y2, "청년등(서식 기준)"))
    for start, text in blocks:
        ws.merge_cells(start_row=head - 1, start_column=start, end_row=head - 1, end_column=start + BLOCK - 1)
        c = ws.cell(row=head - 1, column=start, value=text)
        c.fill, c.font, c.alignment = X.HEAD_FILL, X.HEAD_FONT, MID
        for i in range(BLOCK):
            col = start + i
            hc = ws.cell(row=head, column=col, value=(f"=EOMONTH({START},{i})" if start == B_PAY else f"={L(B_PAY + i)}${head}"))
            hc.number_format, hc.fill, hc.font, hc.border, hc.alignment = "yy-mm", X.HEAD_FILL, X.HEAD_FONT, X.BOX, MID
            X.cell(ws, 6, col, f"=IF({L(col)}${head}<={END},1,0)")
    X.cell(ws, 6, 1, "과세연도 안의 달이면 1", kind="total")
    X.cell(ws, 7, 1, "사람 수(급여가 있는 사람 · 월말 재직)", kind="total")
    X.cell(ws, 8, 1, "월말 현재 인원(상시근로자, 단시간 가중)", kind="total")
    rng = lambda col: f"${L(col)}${first}:${L(col)}${last}"       # noqa: E731
    for i in range(BLOCK):
        X.cell(ws, 7, B_PAY + i, f'=COUNTIF({rng(B_PAY + i)},">0")', kind="total")
        X.cell(ws, 7, B_WORK + i, f"=SUM({rng(B_WORK + i)})", kind="total")
        for start in (B_WORK, B_Y1, B_Y2):
            c = X.cell(ws, 8, start + i, f'=SUMPRODUCT({rng(start + i)},{rng(J_WEIGHT)},--({rng(J_REG)}="Y"))', kind="total")
            c.number_format = NUM2
    st = lambda col, r: ref(S_STAFF, f"${L(col)}{r}")             # noqa: E731
    span = lambda start, r: f"${L(start)}{r}:${L(start + BLOCK - 1)}{r}"   # noqa: E731
    heads = f"${L(B_PAY)}${head}:${L(B_PAY + BLOCK - 1)}${head}"
    acts = f"${L(B_PAY)}$6:${L(B_PAY + BLOCK - 1)}$6"
    for i in range(n):
        r, s = first + i, STAFF_FIRST + i
        a = lambda col: f"${L(col)}{r}"   # noqa: E731
        sumifs = lambda x: (f'SUMIFS({pay[x]},{pay["id"]},{a(J_ID)},{pay["month"]},">="&{START},'      # noqa: E731
                            f'{pay["month"]},"<="&{END})')
        proof = (f'=IF({PAYN}=0,"자료없음",IF(IF({st(7, s)}="단시간",' + "+".join(sumifs(x) for x in ("tax", "pension", "health"))
                 + f',{sumifs("pay")})>0,"Y","N"))') if n_pay else "자료없음"
        f = {
            J_ID: f"={st(1, s)}", J_NAME: f"={st(2, s)}", J_BIRTH: f"={st(3, s)}", J_HIRED: f"={st(4, s)}",
            J_LEFT: f'=IF({st(5, s)}="","",{st(5, s)})', J_LEFTX: f'=IF({st(5, s)}="",DATE(9999,12,31),{st(5, s)})',
            J_PROOF: proof,
            J_REG: f'=IF(AND({st(C_REASON, s)}="",{a(J_PROOF)}<>"N"),"Y","N")',
            J_REASON: f'=IF({st(C_REASON, s)}<>"",{st(C_REASON, s)},IF({a(J_PROOF)}="N","{REASONS["e"]}",""))',
            J_WEIGHT: f'=IF(AND({st(C_WEIGHT, s)}={k["w_sup"]},{ALLOW}="N"),{k["w_half"]},{st(C_WEIGHT, s)})',
            J_MW: f"=SUM({span(B_WORK, r)})",
            J_M20: f'=IF({a(J_REG)}="Y",{a(J_WEIGHT)}*{a(J_MW)},0)',
            J_N21: f"=ROUNDDOWN(ROUND({a(J_M20)}/{MONTHS},8),2)",
            J_MY1: f"=SUM({span(B_Y1, r)})",
            J_M22A: f'=IF({a(J_REG)}="Y",{a(J_WEIGHT)}*{a(J_MY1)},0)',
            J_N23A: f"=ROUNDDOWN(ROUND({a(J_M22A)}/{MONTHS},8),2)",
            J_MY2: f"=SUM({span(B_Y2, r)})",
            J_M22: f'=IF({a(J_REG)}="Y",{a(J_WEIGHT)}*{a(J_MY2)},0)',
            J_N23: f"=ROUNDDOWN(ROUND({a(J_M22)}/{MONTHS},8),2)",
            J_KIND: (f'=IF({a(J_M22)}=0,"",IF({st(attr_col("disabled"), s)}="Y","장애인",IF({st(C_SENIOR, s)}="Y","고령자",'
                     f'IF({st(attr_col("career"), s)}="Y","경력단절",IF({st(attr_col("defector"), s)}="Y","북한이탈주민","청년")))))'),
            J_NOPAY: (f"=IF({PAYN}=0,0,SUMPRODUCT(({heads}>={a(J_HIRED)})*({heads}<={a(J_LEFTX)})*({acts}=1)*({span(B_PAY, r)}=0)))"
                      if n_pay else 0),
            J_ODD: (f"=SUMPRODUCT(({span(B_PAY, r)}>0)*((({heads}<{a(J_HIRED)})+(({heads}-DAY({heads})+1)>{a(J_LEFTX)}))>0))"
                    if n_pay else 0),
        }
        for col, formula in f.items():
            c = X.cell(ws, r, col, formula)
            if col in (J_BIRTH, J_HIRED, J_LEFT, J_LEFTX):
                c.number_format = DATE
            elif col in (J_WEIGHT, J_M20, J_N21, J_M22A, J_N23A, J_M22, J_N23):
                c.number_format = NUM2
        for m in range(BLOCK):
            pc, wc, y1, y2 = B_PAY + m, B_WORK + m, B_Y1 + m, B_Y2 + m
            hd = lambda col: f"{L(col)}${head}"       # noqa: E731
            first_day = f"({hd(pc)}-DAY({hd(pc)})+1)"
            X.cell(ws, r, pc, (f'=IF({L(pc)}$6=1,SUMIFS({pay["pay"]},{pay["id"]},{a(J_ID)},{pay["month"]},">="&{first_day},'
                               f'{pay["month"]},"<="&{hd(pc)}),0)') if n_pay else 0, money=True)
            X.cell(ws, r, wc, (f"=IF(AND({L(wc)}$6=1,{a(J_HIRED)}<={hd(wc)},{a(J_LEFTX)}>={hd(wc)},"
                               f'OR({PAYN}=0,{UNPAID}="Y",{L(pc)}{r}>0)),1,0)'))
            young = lambda col, cut: (f'=IF(AND({L(wc)}{r}=1,OR({st(C_FIXED, s)}="Y",AND({st(C_REGULAR, s)}="Y",'      # noqa: E731
                                      f'{hd(col)}>={st(C_15, s)},{hd(col)}<{cut}))),1,0)')
            X.cell(ws, r, y1, young(y1, st(C_35, s)))
            X.cell(ws, r, y2, young(y2, f'IF({DEEM}="Y",{st(C_CUT, s)},{st(C_35, s)})'))
    ws.freeze_panes = ws.cell(row=first, column=3)
    X.widths(ws, [10, 10, 11, 11, 11, 11, 8, 7, 22, 7, 7, 8, 8, 7, 8, 8, 7, 8, 8, 9, 7, 7, 2] + [8] * BLOCK + [4] * (BLOCK * 3))


def attr_col(name: str) -> int:
    return [c[1] for c in STAFF_COLS].index(name) + 1


# ---------------------------------------------------------------- 4인원

def _count(ws, k: KV, judge_last: int) -> dict:
    row = X.title(ws, "4. 인원", "판정 시트를 과세연도별로 더합니다. 근로자별 방식과 월말 합 방식을 함께 보이고, 서식에 쓰는 인원을 고릅니다")
    out: dict = {}
    X.header(ws, row, ["과세연도", "시작일", "종료일", "개월 수", "급여대장 줄 수", "근로자별 상시", "근로자별 청년등(현재 연령)",
                       "근로자별 청년등(서식 기준)", "월말 인원 합(상시)", "월말 합 방식 상시", "월말 인원 합(청년등)", "월말 합 방식 청년등",
                       "재직인데 급여 없는 달", "재직 기간이 아닌데 급여 있는 달"])
    ws.row_dimensions[row].height = 44
    rng = lambda per, col: ref(S_JUDGE[per], f"${L(col)}${JUDGE_FIRST}:${L(col)}${judge_last}")      # noqa: E731
    top = lambda per, start: ref(S_JUDGE[per], f"${L(start)}$8:${L(start + BLOCK - 1)}$8")          # noqa: E731
    a = {}
    for i, per in enumerate(PERIODS):
        r = row + 1 + i
        j = lambda addr: ref(S_JUDGE[per], addr)       # noqa: E731
        vals = [per, f"={j('$B$4')}", f"={j('$D$4')}", f"={j('$F$4')}", f"={j('$J$4')}",
                f"=SUM({rng(per, J_N21)})", f"=SUM({rng(per, J_N23A)})", f"=SUM({rng(per, J_N23)})",
                f"=SUM({top(per, B_WORK)})", None, f"=SUM({top(per, B_Y1)})", None,
                f"=SUM({rng(per, J_NOPAY)})", f"=SUM({rng(per, J_ODD)})"]
        vals[9] = f"=ROUNDDOWN(ROUND(I{r}/D{r},8),2)"
        vals[11] = f"=ROUNDDOWN(ROUND(K{r}/D{r},8),2)"
        for col, v in enumerate(vals, start=1):
            c = X.cell(ws, r, col, v, kind="total" if col == 1 else "body")
            c.number_format = DATE if col in (2, 3) else ("0" if col in (4, 5, 13, 14) else NUM2)
        a[per] = {name: ref(S_COUNT, absolute(col, r)) for name, col in
                  (("payn", 5), ("w_total", 6), ("w_youth_age", 7), ("w_youth", 8), ("me_sum", 9), ("me_total", 10),
                   ("me_youth_sum", 11), ("me_youth", 12), ("nopay", 13), ("odd", 14))}
    out["raw"] = a

    row += 6
    ws.cell(row=row - 1, column=1, value="서식에 쓰는 인원(1차 연도 계산)").font = Font(name=X.FONT, bold=True, size=11)
    X.header(ws, row, ["과세연도", "출처", "⑥ 상시근로자 수", "⑦ 청년등상시근로자 수", "⑧ 청년등 외(⑥-⑦)", "설명"])
    used = {}
    filed = {"직전": ("filed_p_t", "filed_p_y"), "직전전": ("filed_pp_t", "filed_pp_y")}
    f2 = lambda x: f"ROUNDDOWN(ROUND(N({x}),8),2)"     # noqa: E731
    for i, per in enumerate(PERIODS):
        r = row + 1 + i
        p = a[per]
        if per == "해당":
            src = "급여대장"
            total = f'=IF({k["method"]}="근로자별",{p["w_total"]},{p["me_total"]})'
            youth = f'=IF({k["method"]}="근로자별",{p["w_youth"]},{p["me_youth"]})'
            text = "근로자별 방식이면 청년 간주를 넣은 값(별지 제10호의10서식 2번 표)"
        else:
            t, y = k[filed[per][0]], k[filed[per][1]]
            tail = f'IF({k["startup"]}="Y","창업(0)","없음")' if per == "직전" else '"없음"'
            src = f'=IF({p["payn"]}>0,"급여대장",IF({t}<>"","전기 신고서",{tail}))'
            mine_t = f'IF({k["method"]}="근로자별",{p["w_total"]},{p["me_total"]})' if per == "직전" else p["me_total"]
            mine_y = f'IF({k["method"]}="근로자별",{p["w_youth"]},{p["me_youth"]})' if per == "직전" else p["me_youth"]
            total = f'=IF($B{r}="급여대장",{mine_t},IF($B{r}="전기 신고서",{f2(t)},0))'
            youth = f'=IF($B{r}="급여대장",{mine_y},IF($B{r}="전기 신고서",{f2(y)},0))'
            text = ("청년 간주를 넣지 않습니다(별지 제10호의10서식 작성방법 24 · 25)" if per == "직전"
                    else "월말 합 방식(종전 규정)")
        X.cell(ws, r, 1, per, kind="total")
        X.cell(ws, r, 2, src if per == "해당" else src)
        for col, v in ((3, total), (4, youth), (5, f"=ROUND(C{r}-D{r},2)")):
            X.cell(ws, r, col, v).number_format = NUM2
        X.cell(ws, r, 6, text).font = X.NOTE_FONT
        used[per] = {"src": ref(S_COUNT, f"$B${r}"), "t": ref(S_COUNT, f"$C${r}"), "y": ref(S_COUNT, f"$D${r}"),
                     "o": ref(S_COUNT, f"$E${r}")}
    out["used"] = used

    row += 6
    ws.cell(row=row - 1, column=1, value="2 · 3차 연도 판정에 쓰는 인원(월말 합 방식, 종전 규정)").font = Font(name=X.FONT, bold=True, size=11)
    X.header(ws, row, ["과세연도", "출처", "상시근로자 수", "청년등상시근로자 수", "", "설명"])
    legacy = {}
    for i, per in enumerate(PERIODS):
        r = row + 1 + i
        p = a[per]
        X.cell(ws, r, 1, per, kind="total")
        if per == "해당":
            X.cell(ws, r, 2, "급여대장")
            total, youth = f"={p['me_total']}", f"={p['me_youth']}"
        else:
            t, y = k[filed[per][0]], k[filed[per][1]]
            X.cell(ws, r, 2, f"={used[per]['src']}")
            total = f'=IF($B{r}="급여대장",{p["me_total"]},IF($B{r}="전기 신고서",{f2(t)},0))'
            youth = f'=IF($B{r}="급여대장",{p["me_youth"]},IF($B{r}="전기 신고서",{f2(y)},0))'
        X.cell(ws, r, 3, total).number_format = NUM2
        X.cell(ws, r, 4, youth).number_format = NUM2
        X.cell(ws, r, 6, "부칙 <법률 제21223호> 제34조: 2024 과세연도까지의 증가분은 종전 규정").font = X.NOTE_FONT
        legacy[per] = {"t": ref(S_COUNT, f"$C${r}"), "y": ref(S_COUNT, f"$D${r}")}
    out["legacy"] = legacy

    row += 6
    ws.cell(row=row - 1, column=1, value="전기 신고 인원과 다시 센 인원(월말 합 방식)").font = Font(name=X.FONT, bold=True, size=11)
    X.header(ws, row, ["과세연도", "구분", "전기 신고서", "다시 센 인원", "차이", "설명"])
    recon = {}
    r = row + 1
    for per in ("직전", "직전전"):
        for name, key, mine in (("상시", filed[per][0], a[per]["me_total"]), ("청년등", filed[per][1], a[per]["me_youth"])):
            X.cell(ws, r, 1, per, kind="total")
            X.cell(ws, r, 2, name)
            X.cell(ws, r, 3, f'=IF({k[key]}="","",{k[key]})').number_format = NUM2
            X.cell(ws, r, 4, f'=IF({a[per]["payn"]}>0,{mine},"")').number_format = NUM2
            X.cell(ws, r, 5, f'=IF(OR(C{r}="",D{r}=""),"",ROUND(D{r}-C{r},2))').number_format = NUM2
            X.cell(ws, r, 6, "둘 다 있을 때만 봅니다. 차이가 있으면 사원명부의 제외 사유 · 청년 구분을 전기 명세와 맞춰 보십시오").font = X.NOTE_FONT
            recon[(per, name)] = ref(S_COUNT, f"$E${r}")
            r += 1
    out["recon"] = recon
    X.widths(ws, [12, 12, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 12, 12])
    return out


# ---------------------------------------------------------------- 5공제액

def _credit(ws, k: KV, cnt: dict, rules: dict, rates: dict) -> KV:
    row = X.title(ws, "5. 공제액", "조세특례제한법 제29조의8 제1항 · 제3항 · 제5항. 인원은 4인원 시트에서 옵니다")
    X.header(ws, row, ["항목", "값", "설명 · 근거"])
    c = KV(ws, row + 1)
    u, g = cnt["used"], cnt["legacy"]
    e = rates["employment_integrated"]
    n2 = lambda key, label, f, note="": c.put(key, label, f, fmt=NUM2, note=note)   # noqa: E731
    won = lambda key, label, f, note="", kind="body": c.put(key, label, f, kind=kind, fmt=X.MONEY, note=note)   # noqa: E731

    c.section("가. 1차 연도(해당 과세연도 증가분)")
    n2("11", "⑪ 해당 과세연도 상시근로자 수", f"={u['해당']['t']}")
    n2("12", "⑫ 직전 과세연도 상시근로자 수", f"={u['직전']['t']}")
    n2("13", "⑬ 상시근로자 증가 인원 수(⑪-⑫)", f"=ROUND({c['11']}-{c['12']},2)")
    n2("14", "⑭ 해당 과세연도 청년등상시근로자 수", f"={u['해당']['y']}")
    n2("15", "⑮ 직전 과세연도 청년등상시근로자 수", f"={u['직전']['y']}")
    n2("16", "⑯ 청년등상시근로자 증가 인원 수(⑭-⑮)", f"=ROUND({c['14']}-{c['15']},2)")
    n2("17", "⑰ 해당 과세연도 청년등 외 상시근로자 수", f"={u['해당']['o']}")
    n2("18", "⑱ 직전 과세연도 청년등 외 상시근로자 수", f"={u['직전']['o']}")
    n2("19", "⑲ 청년등 외 상시근로자 증가 인원 수(⑰-⑱)", f"=ROUND({c['17']}-{c['18']},2)")
    c.put("ok", "1차 연도 요건",
          f'=IF(AND({k["consumer"]}<>"Y",{k["in_time"]}="Y",{u["직전"]["src"]}<>"없음",{c["13"]}>0),"충족","미충족")',
          note="소비성서비스업이 아니고, 2025년 12월 31일이 속하는 과세연도까지이며, 직전 인원이 있고, ⑬ 이 0 보다 큼")
    n2("n_y", "청년등 공제 인원(⑬ 한도)", f'=IF({c["ok"]}="충족",MAX(0,MIN({c["16"]},{c["13"]})),0)', e["cap"]["basis"])
    n2("n_o", "청년등 외 공제 인원(⑬ 한도)", f'=IF({c["ok"]}="충족",MAX(0,MIN({c["19"]},{c["13"]})),0)')
    pick = lambda part: (f'=IF({k["size"]}="중소기업",IF({k["capital"]}="N",{k[part + "_sme_local"]},{k[part + "_sme_capital"]}),'   # noqa: E731
                         f'IF({k["size"]}="중견기업",{k[part + "_mid"]},{k[part + "_general"]}))')
    won("a_y", "청년등 1인당 공제금액", pick("youth"))
    won("a_o", "청년등 외 1인당 공제금액", pick("other"))
    won("first_y", "1차 연도 청년등 공제액", f"=ROUND({c['n_y']}*{c['a_y']},0)")
    won("first_o", "1차 연도 청년등 외 공제액", f"=ROUND({c['n_o']}*{c['a_o']},0)")
    won("20", "⑳ 1차년도 세액공제액", f"={c['first_y']}+{c['first_o']}", kind="total")

    c.section("나. 2차 연도(직전 과세연도 증가분, 월말 합 방식 인원)")
    n2("21", "㉑ 2차년도(해당 과세연도) 상시근로자 수", f"={g['해당']['t']}")
    n2("22", "㉒ 1차년도(직전 과세연도) 상시근로자 수", f"={g['직전']['t']}")
    n2("23", "㉓ 상시근로자 증가 인원 수(㉑-㉒)", f"=ROUND({c['21']}-{c['22']},2)")
    c.put("cut2_t", "직전 과세연도 대비 상시근로자 감소여부", f'=IF({c["21"]}<{c["22"]},"여","부")', note=e["followup"]["basis"])
    c.put("cut2_y", "직전 과세연도 대비 청년등상시근로자 수 감소여부", f'=IF({g["해당"]["y"]}<{g["직전"]["y"]},"여","부")',
          note="청년등만 줄면 청년등 몫만 끊깁니다(제29조의8 제3항)")
    won("24", "㉔ 직전 과세연도 청년등 증가 세액공제액", f"={k['f1y']}")
    won("25", "㉕ 직전 과세연도 청년등 외 증가 세액공제액", f"={k['f1o']}")
    won("26", "㉖ 2차년도 세액공제액",
        f'=IF(OR({k["consumer"]}="Y",{c["24"]}+{c["25"]}=0,{u["직전"]["src"]}="없음",{u["직전"]["src"]}="창업(0)",'
        f'{c["cut2_t"]}="여"),0,IF({c["cut2_y"]}="여",0,{c["24"]})+{c["25"]})', kind="total")

    c.section("다. 3차 연도(직전전 과세연도 증가분, 중소 · 중견기업)")
    n2("27", "㉗ 3차년도(해당 과세연도) 상시근로자 수", f"={g['해당']['t']}")
    n2("28", "㉘ 1차년도(직전전 과세연도) 상시근로자 수", f"={g['직전전']['t']}")
    n2("29", "㉙ 상시근로자 증가 인원(㉗-㉘)", f"=ROUND({c['27']}-{c['28']},2)")
    c.put("cut3_t", "직전전 과세연도 대비 상시근로자 감소여부",
          f'=IF(OR({g["직전"]["t"]}<{g["직전전"]["t"]},{g["해당"]["t"]}<{g["직전전"]["t"]}),"여","부")',
          note="직전 · 해당 어느 과세연도든 직전전보다 줄었으면 여")
    c.put("cut3_y", "직전전 과세연도 대비 청년등상시근로자 수 감소여부",
          f'=IF(OR({g["직전"]["y"]}<{g["직전전"]["y"]},{g["해당"]["y"]}<{g["직전전"]["y"]}),"여","부")')
    won("30", "㉚ 직전전 과세연도 청년등 증가 세액공제액", f"={k['f2y']}")
    won("31", "㉛ 직전전 과세연도 청년등 외 증가 세액공제액", f"={k['f2o']}")
    won("32", "㉜ 3차년도 세액공제액",
        f'=IF(OR({k["consumer"]}="Y",{c["30"]}+{c["31"]}=0,{k["size"]}="일반기업",{u["직전전"]["src"]}="없음",'
        f'{u["직전"]["src"]}="없음",{u["직전"]["src"]}="창업(0)",{c["cut3_t"]}="여"),0,IF({c["cut3_y"]}="여",0,{c["30"]})+{c["31"]})',
        kind="total")

    c.section("라. 추가공제(육아휴직 복귀자)")
    n2("33", "㉝ 해당 과세연도 상시근로자 수", f"={u['해당']['t']}", "출산전후휴가 대체 인력이 있으면 휴가 중인 사람을 빼고 다시 셉니다(작성방법 6)")
    n2("34", "㉞ 직전 과세연도 상시근로자 수", f"={u['직전']['t']}")
    n2("35", "㉟ 상시근로자 증가 인원 수(㉝-㉞)", f"=ROUND({c['33']}-{c['34']},2)")
    n2("ret", "육아휴직 복귀자 인원 수", f"={k['returned']}")
    won("36", "㊱ 추가공제 세액공제액",
        f'=IF(AND({k["consumer"]}<>"Y",{k["in_time"]}="Y",{u["직전"]["src"]}<>"없음",{c["35"]}>=0,'
        f'OR({k["size"]}="중소기업",{k["size"]}="중견기업")),ROUND({c["ret"]}*IF({k["size"]}="중소기업",{k["return_sme"]},'
        f'{k["return_mid"]}),0),0)', basis(rules, "extra", "return_sme"), kind="total")

    c.section("세액공제액")
    won("total", "⑥ 세액공제액(⑳ + ㉖ + ㉜ + ㊱)", f"={c['20']}+{c['26']}+{c['32']}+{c['36']}",
        "8호(갑) 코드 18S 당기발생액. 최저한세 · 이월 · 농어촌특별세는 조서 20 공제감면에서 봅니다", kind="total")
    X.widths(ws, [48, 18, 100])
    return c


# ---------------------------------------------------------------- 6서식10의9

class Form:
    """법정 서식 모양으로 그리는 시트. 열 여덟 칸짜리 격자에 글 칸과 값 칸을 놓는다."""

    def __init__(self, ws, cols: int = 8):
        self.ws, self.cols, self.row = ws, cols, 1

    def cell(self, col: int, text, span: int = 1, *, value: bool = False, fmt: str = "", bold: bool = False, rows: int = 1,
             fill: bool = False):
        ws, r = self.ws, self.row
        if span > 1 or rows > 1:
            ws.merge_cells(start_row=r, start_column=col, end_row=r + rows - 1, end_column=col + span - 1)
        for rr in range(r, r + rows):
            for cc in range(col, col + span):
                ws.cell(row=rr, column=cc).border = FORM_BOX
        c = ws.cell(row=r, column=col, value=text)
        c.font = FORM_BOLD if bold else FORM_FONT
        c.alignment = Alignment(horizontal="right" if (value and fmt) else "center", vertical="center", wrap_text=True)
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = GREY
        return c

    def line(self, cells: list, height: float = 24) -> int:
        """cells: (글, 칸 수, 옵션) 을 왼쪽부터. 쓴 행 번호를 돌려준다."""
        col = 1
        for text, span, *opt in cells:
            self.cell(col, text, span, **(opt[0] if opt else {}))
            col += span
        self.ws.row_dimensions[self.row].height = height
        self.row += 1
        return self.row - 1

    def head(self, text: str) -> None:
        self.cell(1, text, self.cols, bold=True, fill=True)
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1


V2 = {"value": True, "fmt": NUM2}
VW = {"value": True, "fmt": X.MONEY}
G = {"fill": True}


def _form(ws, k: KV, cnt: dict, c: KV) -> dict:
    f = Form(ws)
    u = cnt["used"]
    ws.cell(row=1, column=1, value="■ 조세특례제한법 시행규칙 [별지 제10호의9서식(1)] <개정 2026. 3. 20.>").font = X.NOTE_FONT
    f.row = 2
    f.cell(1, "통합고용세액공제 공제세액계산서", 8, bold=True)
    ws.row_dimensions[2].height = 30
    ws.cell(row=2, column=1).font = Font(name=X.FONT, bold=True, size=14)
    f.row = 3
    f.head("① 신청인")
    f.line([("① 상호 또는 법인명", 2, G), (f"={k['company']}", 2), ("② 사업자등록번호", 2, G), (None, 2)])
    f.line([("③ 대표자 성명", 2, G), (None, 2), ("④ 생년월일", 2, G), (None, 2)])
    f.line([("⑤ 주소 또는 본점소재지", 2, G), (None, 6)])
    f.head("② 과세연도")
    f.line([("과세연도", 2, G), (f"={k['fy_start']}", 2, {"fmt": DATE}), ("부터", 1), (f"={k['fy_end']}", 2, {"fmt": DATE}),
            ("까지", 1)])
    f.head("③ 상시근로자 현황")
    f.line([("구분", 2, G), ("직전전 과세연도", 2, G), ("직전 과세연도", 2, G), ("해당 과세연도", 2, G)])
    at = {}
    for code, label, key in (("6", "⑥ 상시근로자 수(⑦+⑧)", "t"), ("7", "⑦ 청년등상시근로자 수", "y"),
                             ("8", "⑧ 청년등상시근로자를 제외한 상시근로자 수", "o")):
        cells = [(label, 2, G)]
        for per in ("직전전", "직전", "해당"):
            src = u[per]["src"]
            v = f"={u[per][key]}" if per == "해당" else f'=IF({src}="없음","",{u[per][key]})'
            if code == "6":
                v = f"={L(3 + 2 * ('직전전', '직전', '해당').index(per))}{f.row + 1}+{L(3 + 2 * ('직전전', '직전', '해당').index(per))}{f.row + 2}" \
                    if per == "해당" else (f'=IF({src}="없음","",{L(3 + 2 * ("직전전", "직전", "해당").index(per))}{f.row + 1}'
                                           f'+{L(3 + 2 * ("직전전", "직전", "해당").index(per))}{f.row + 2})')
            cells.append((v, 2, V2))
        r = f.line(cells, 30)
        for i, per in enumerate(("직전전", "직전", "해당")):
            at[f"{code}_{per}"] = ref(S_FORM, absolute(3 + 2 * i, r))
    f.line([("⑨ 정규직 전환 근로자 수", 2, G), ("-", 2), ("-", 2), (None, 2)])
    f.line([("⑩ 육아휴직 복귀자 수", 2, G), ("-", 2), ("-", 2), (f"={c['ret']}", 2, V2)])
    f.head("④ 기본공제 공제세액 계산내용")
    f.head("가. 1차년도 세제지원 요건 : ⑬ 〉0")
    for title, codes, labels in (
            ("1. 상시근로자 증가 인원", ("11", "12", "13"),
             ("⑪ 해당 과세연도 상시근로자 수", "⑫ 직전 과세연도 상시근로자 수", "⑬ 상시근로자 증가 인원 수(⑪-⑫)")),
            ("2. 청년등상시근로자 증가 인원", ("14", "15", "16"),
             ("⑭ 해당 과세연도 청년등상시근로자 수", "⑮ 직전 과세연도 청년등상시근로자 수", "⑯ 청년등상시근로자 증가 인원 수(⑭-⑮)")),
            ("3. 청년등상시근로자를 제외한 상시근로자 증가 인원", ("17", "18", "19"),
             ("⑰ 해당 과세연도 청년등상시근로자를 제외한 상시근로자 수", "⑱ 직전 과세연도 청년등상시근로자를 제외한 상시근로자 수",
              "⑲ 청년등상시근로자를 제외한 상시근로자 증가 인원 수(⑰-⑱)"))):
        f.line([(title, 8, G)], 18)
        f.line([(labels[0], 3, G), (labels[1], 3, G), (labels[2], 2, G)], 40)
        r = f.row
        cur = at[{"11": "6_해당", "14": "7_해당", "17": "8_해당"}[codes[0]]]
        prev = at[{"12": "6_직전", "15": "7_직전", "18": "8_직전"}[codes[1]]]
        f.line([(f"={cur}", 3, V2), (f"=N({prev})", 3, V2), (f"=ROUND(A{r}-D{r},2)", 2, V2)])
        at[codes[0]], at[codes[1]], at[codes[2]] = (ref(S_FORM, absolute(x, r)) for x in (1, 4, 7))
    f.line([("4. 1차년도 세액공제액 계산", 8, G)], 18)
    f.line([("구분", 2, G), ("구분", 1, G), ("직전 과세연도 대비 상시근로자 증가 인원 수(⑬ 상시근로자 증가 인원 수를 한도로 함)", 2, G),
            ("1인당 공제금액", 1, G), ("⑳ 1차년도 세액공제액", 2, G)], 44)
    groups = (("중소기업", "수도권 내", 'AND({s}="중소기업",{cap}<>"N")', "sme_capital"),
              ("중소기업", "수도권 밖", 'AND({s}="중소기업",{cap}="N")', "sme_local"),
              ("중견기업", "", '{s}="중견기업"', "mid"), ("일반기업", "", '{s}="일반기업"', "general"))
    sums: dict = {}
    for size, region, cond, tail in groups:
        cond = cond.format(s=k["size"], cap=k["capital"])
        for part, n_key in (("청년등", "n_y"), ("청년등 외", "n_o")):
            r = f.row
            rate = k[("youth_" if part == "청년등" else "other_") + tail]
            f.line([(f"{size} {region}".strip(), 2, G), (part, 1, G), (f"=IF({cond},{c[n_key]},0)", 2, V2), (f"={rate}", 1, VW),
                    (f"=ROUND(D{r}*F{r},0)", 2, VW)])
            sums.setdefault(size, []).append(r)
        last = size != "중소기업" or region == "수도권 밖"
        if last:
            r = f.row
            rows = sums[size]
            f.line([(size, 2, G), ("계", 1, G), ("=" + "+".join(f"D{x}" for x in rows), 2, V2), (None, 1),
                    ("=" + "+".join(f"G{x}" for x in rows), 2, VW)])
            sums[size + "_계"] = r
    r = f.row
    f.line([("⑳ 1차년도 세액공제액(중소기업 · 중견기업 · 일반기업 계)", 6, G),
            ("=" + "+".join(f"G{sums[s + '_계']}" for s in ("중소기업", "중견기업", "일반기업")), 2, VW)])
    at["20"] = ref(S_FORM, absolute(7, r))

    def follow(title, sub, codes, labels, per, keys):
        f.head(title)
        f.line([("1. 상시근로자 증가 인원", 8, G)], 18)
        f.line([(labels[0], 3, G), (labels[1], 3, G), (labels[2], 2, G)], 40)
        r = f.row
        f.line([(f"={c[codes[0]]}", 3, V2), (f"={c[codes[1]]}", 3, V2), (f"=ROUND(A{r}-D{r},2)", 2, V2)])
        f.line([(sub, 8, G)], 18)
        f.line([(f"1차년도({per} 과세연도) 대비 상시근로자 감소여부", 2, G),
                (f"1차년도({per} 과세연도) 대비 청년등상시근로자 수 감소여부", 2, G),
                (labels[3], 1, G), (labels[4], 1, G), (labels[5], 2, G)], 58)
        r = f.row
        f.line([(f"={c[keys[0]]}", 2), (f"={c[keys[1]]}", 2), (f"={c[codes[3]]}", 1, VW), (f"={c[codes[4]]}", 1, VW),
                (f"={c[codes[5]]}", 2, VW)])
        at[codes[5]] = ref(S_FORM, absolute(7, r))

    follow("나. 2차년도 세제지원 요건 : ㉓ ≥ 0", "2. 2차년도 세액공제액 계산(상시근로자 감소여부)", ("21", "22", "23", "24", "25", "26"),
           ("㉑ 2차년도(해당 과세연도) 상시근로자 수", "㉒ 1차년도(직전 과세연도) 상시근로자 수", "㉓ 상시근로자 증가 인원 수(㉑-㉒)",
            "㉔ 1차년도(직전 과세연도) 청년등상시근로자 증가 세액공제액", "㉕ 1차년도(직전 과세연도) 청년등 외 상시근로자 증가 세액공제액",
            "㉖ 2차년도 세액공제액"), "직전", ("cut2_t", "cut2_y"))
    follow("다. 3차년도 세제지원 요건(중소ㆍ중견기업만 해당) : ㉙ ≥ 0", "2. 3차년도 세액공제액 계산(상시근로자 감소여부)",
           ("27", "28", "29", "30", "31", "32"),
           ("㉗ 3차년도(해당 과세연도) 상시근로자 수", "㉘ 1차년도(직전전 과세연도) 상시근로자 수", "㉙ 상시근로자 증가 인원(㉗-㉘)",
            "㉚ 1차년도(직전전 과세연도) 청년등상시근로자 증가 세액공제액", "㉛ 1차년도(직전전 과세연도) 청년등 외 상시근로자 증가 세액공제액",
            "㉜ 3차년도 세액공제액"), "직전전", ("cut3_t", "cut3_y"))

    f.head("➎ 추가공제 공제세액 계산내용")
    f.line([("가. 세제지원 요건 : ㉟ ≥ 0", 8, G)], 18)
    f.line([("㉝ 해당 과세연도 상시근로자 수", 3, G), ("㉞ 직전 과세연도 상시근로자 수", 3, G), ("㉟ 상시근로자 증가 인원 수(㉝-㉞)", 2, G)], 30)
    r = f.row
    f.line([(f"={c['33']}", 3, V2), (f"={c['34']}", 3, V2), (f"=ROUND(A{r}-D{r},2)", 2, V2)])
    f.line([("나. 세액공제액 계산", 8, G)], 18)
    f.line([("구분", 2, G), ("구분", 1, G), ("인원 수", 2, G), ("1인당 공제금액", 1, G), ("㊱ 추가공제 세액공제액", 2, G)], 30)
    extra = []
    for size, key in (("중소기업", "return_sme"), ("중견기업", "return_mid")):
        r = f.row
        f.line([(size, 2, G), ("육아휴직 복귀자", 1, G), (f'=IF({k["size"]}="{size}",{c["ret"]},0)', 2, V2), (f"={k[key]}", 1, VW),
                (f'=IF({k["size"]}="{size}",{c["36"]},0)', 2, VW)])
        extra.append(r)
    r = f.row
    f.line([("계", 6, G), ("=" + "+".join(f"G{x}" for x in extra), 2, VW)])
    at["36"] = ref(S_FORM, absolute(7, r))
    r = f.row
    f.line([("⑥ 세액공제액 : ⑳ 1차년도 세액공제액 + ㉖ 2차년도 세액공제액 + ㉜ 3차년도 세액공제액 + ㊱ 추가공제 세액공제액", 6, {"bold": True, "fill": True}),
            (f"={at['20']}+{at['26']}+{at['32']}+{at['36']}", 2, VW)], 34)
    at["total"] = ref(S_FORM, absolute(7, r))
    ws.cell(row=f.row + 1, column=1,
            value="값은 모두 5공제액 · 4인원 시트를 읽는 수식입니다. ② 사업자등록번호 · ③ 대표자 · ④ 생년월일 · ⑤ 주소는 사람이 적습니다. "
                  "서식 글은 법제처 별지 서식(별표 001009)에서 옮겼습니다").font = X.NOTE_FONT
    X.widths(ws, [13, 13, 13, 13, 13, 14, 13, 13])
    return at


# ---------------------------------------------------------------- 7서식10의10

R_HEAD = ["생년월일", "성명", "상시근로자(N/Y)", "상시근로자 제외사유", "단시간근로자 유형(0.5/0.75)", "내국인(N/Y)", "수도권 근무지(N/Y)",
          "근로계약 체결일", "근로기간(MM~MM)", "상시근로자 근무개월수", "상시근로자별 계산한 수(÷①)", "청년등 상시근로자 근무개월수",
          "청년등 상시근로자별 계산한 수(÷①)", "과세연도 기준 청년여부(N/Y)", "근로계약 체결일 기준 청년여부(N/Y)", "청년 제외시점",
          "장애인(N/Y)", "고령자(N/Y)", "경력단절근로자(N/Y)", "북한이탈주민(N/Y)", "정규직 전환자(N/Y)", "육아휴직 복귀자(N/Y)"]
R_NO = {"해당": "⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳㉑㉒㉓㉔㉕㉖㉗㉘㉙㉚㉛㉜", "직전": "㉝㉞㉟㊱㊲㊳㊴㊵㊶㊷㊸㊹㊺㊻㊼㊽㊾㊿"}


def _roster(ws, inp: Inputs, res: dict, k: KV, cnt: dict, judge_last: int) -> dict:
    """상시근로자 명세서. 과세연도와 재직 기간이 겹치는 사람만 줄을 만든다(줄을 고르는 것만 엔진이 하고 값은 수식이다)."""
    ws.cell(row=1, column=1, value="[별지 제10호의10서식] <개정 2026. 5. 22.>").font = X.NOTE_FONT
    ws.cell(row=2, column=1, value="통합고용세액공제 상시근로자 명세서").font = Font(name=X.FONT, bold=True, size=14)
    X.cell(ws, 3, 1, "과세연도", kind="total")
    X.cell(ws, 3, 2, f"={k['fy_start']}").number_format = DATE
    X.cell(ws, 3, 3, f"={k['fy_end']}").number_format = DATE
    X.cell(ws, 3, 4, "상호 또는 법인명", kind="total")
    X.cell(ws, 3, 5, f"={k['company']}")
    row = 5
    ws.cell(row=row, column=1, value="1. 상시근로자 수 계산").font = Font(name=X.FONT, bold=True, size=11)
    row += 1
    X.header(ws, row, ["구분", "① 과세연도 개월수", "② 상시근로자 수", "④ 청년 상시근로자 수", "⑤ 장애인 상시근로자 수", "⑥ 고령자 상시근로자 수",
                       "⑦ 경력단절근로자 상시근로자 수", "⑧ 북한이탈주민 상시근로자 수", "⑨ 합계", "⑩ 청년등 외 상시근로자수(=②-⑨)"])
    ws.row_dimensions[row].height = 44
    rng = lambda per, col: ref(S_JUDGE[per], f"${L(col)}${JUDGE_FIRST}:${L(col)}${judge_last}")     # noqa: E731
    at = {}
    for i, per in enumerate(("해당", "직전")):
        r = row + 1 + i
        on = "TRUE" if per == "해당" else f'{cnt["used"][per]["src"]}="급여대장"'
        cells = [f"{per} 과세연도", f"={ref(S_JUDGE[per], '$F$4')}", f'=IF({on},SUM({rng(per, J_N21)}),"")']
        for kind in ("청년", "장애인", "고령자", "경력단절", "북한이탈주민"):
            cells.append(f'=IF({on},SUMIF({rng(per, J_KIND)},"{kind}",{rng(per, J_N23)}),"")')
        cells += [f'=IF({on},SUM(D{r}:H{r}),"")', f'=IF({on},ROUND(C{r}-I{r},2),"")']
        for col, v in enumerate(cells, start=1):
            X.cell(ws, r, col, v, kind="total" if col == 1 else "body").number_format = "0" if col == 2 else NUM2
        at[per] = {"total": ref(S_ROSTER, f"$C${r}"), "youth": ref(S_ROSTER, f"$I${r}"), "other": ref(S_ROSTER, f"$J${r}")}
    row += 4
    st = lambda col, s: ref(S_STAFF, f"${L(col)}{s}")      # noqa: E731
    for per, title in (("해당", "2. 해당 과세연도 상시근로자별 명세"), ("직전", "3. 직전 과세연도 상시근로자별 명세")):
        period = res["periods"][per]
        ws.cell(row=row, column=1, value=title).font = Font(name=X.FONT, bold=True, size=11)
        row += 1
        nos = R_NO[per]
        X.header(ws, row, [f"{nos[i] if i < len(nos) else ''} {h}".strip() for i, h in enumerate(R_HEAD)])
        ws.row_dimensions[row].height = 66
        row += 1
        if per == "직전" and not period.has_payroll:
            X.note(ws, row, "직전 과세연도 급여대장이 없어 명세를 만들지 않았습니다. 전기 신고서의 명세서를 씁니다")
            row += 2
            continue
        start, end, deem = (ref(S_JUDGE[per], x) for x in ("$B$4", "$D$4", "$H$4"))
        shown = 0
        for i, p in enumerate(inp.staff):
            if not (p.hired <= period.end and (p.left or FAR) >= period.start):
                continue
            j, s = JUDGE_FIRST + i, STAFF_FIRST + i
            jg = lambda col: ref(S_JUDGE[per], f"${L(col)}{j}")      # noqa: E731
            young = f'AND({st(C_REGULAR, s)}="Y",{start}<{st(C_35, s)},{end}>={st(C_15, s)})'
            vals = [
                f"={jg(J_BIRTH)}", f"={jg(J_NAME)}", f"={jg(J_REG)}", f"={jg(J_REASON)}",
                f'=IF({jg(J_WEIGHT)}<1,{jg(J_WEIGHT)},"")', f"={st(attr_col('resident'), s)}",
                f'=IF({st(attr_col("capital"), s)}="",{k["capital"]},{st(attr_col("capital"), s)})', f"={jg(J_HIRED)}",
                (f'=IF({jg(J_MW)}=0,"",TEXT(MONTH(MAX({jg(J_HIRED)},{start})),"00")&"~"&TEXT(MONTH(MIN({jg(J_LEFTX)},{end})),"00"))'),
                f"={jg(J_M20)}", f"={jg(J_N21)}", f"={jg(J_M22)}", f"={jg(J_N23)}",
                f'=IF({young},"Y","N")', f'=IF(AND({deem}="Y",ISNUMBER({st(C_DEEM, s)})),"Y","N")',
                f'=IF({st(C_REGULAR, s)}="Y",IF({deem}="Y",{st(C_CUT, s)},{st(C_35, s)}),"")',
                f"={st(attr_col('disabled'), s)}", f"={st(C_SENIOR, s)}", f"={st(attr_col('career'), s)}",
                f"={st(attr_col('defector'), s)}", f"={st(attr_col('converted'), s)}", f"={st(attr_col('returned'), s)}",
            ]
            for col, v in enumerate(vals, start=1):
                c = X.cell(ws, row, col, v)
                c.number_format = YMD if col in (1, 8, 16) else (NUM2 if col in (5, 10, 11, 12, 13) else "General")
                c.alignment = Alignment(horizontal="center")
            row += 1
            shown += 1
        at[f"rows_{per}"] = shown
        row += 2
    X.note(ws, row, "과세연도와 재직 기간이 겹치는 사람만 줄이 있습니다. 사원명부의 입사일 · 퇴사일을 고쳐 대상이 바뀌면 ctax employ 를 다시 돌립니다. "
                    "주민등록번호는 넣지 않았습니다")
    X.widths(ws, [11, 10, 8, 20, 9, 7, 8, 11, 9, 9, 9, 9, 9, 9, 9, 11, 7, 7, 8, 8, 8, 8])
    return at


# ---------------------------------------------------------------- 8엔진대조 · 목록

def _check(ws, inp: Inputs, res: dict, cnt: dict, c: KV, form: dict) -> dict:
    """calc.py 값(왼쪽)과 엑셀 수식 값(오른쪽). {항목: (엔진 값, 칸 주소)} 를 돌려준다."""
    row = X.title(ws, "8. 엔진 대조", "파이썬 엔진이 날짜로 따로 계산한 값과 이 워크북 수식 값을 나란히 놓습니다. 다르면 수식이나 입력이 바뀐 것입니다")
    X.header(ws, row, ["항목", "엔진 값", "엑셀 수식 값", "일치"])
    out = {}
    r = row + 1

    def put(label: str, engine, address: str, fmt: str = NUM2):
        nonlocal r
        X.cell(ws, r, 1, label)
        X.cell(ws, r, 2, engine).number_format = fmt
        X.cell(ws, r, 3, f"={address}").number_format = fmt
        X.cell(ws, r, 4, f'=IF(ABS(N(B{r})-N(C{r}))<0.005,"일치","불일치")')
        out[label] = (engine, f"C{r}")
        r += 1

    raw, used, legacy = cnt["raw"], cnt["used"], cnt["legacy"]
    for per in PERIODS:
        x = res["counts"][per]
        put(f"{per} 근로자별 상시", float(x.total), raw[per]["w_total"])
        put(f"{per} 근로자별 청년등(현재 연령)", float(x.youth_age), raw[per]["w_youth_age"])
        put(f"{per} 근로자별 청년등(서식 기준)", float(x.youth_deem if per == "해당" else x.youth_age), raw[per]["w_youth"])
        put(f"{per} 월말 인원 합(상시)", float(x.me_sum), raw[per]["me_sum"])
        put(f"{per} 월말 합 방식 상시", float(x.me_total), raw[per]["me_total"])
        put(f"{per} 월말 합 방식 청년등", float(x.me_youth), raw[per]["me_youth"])
        put(f"{per} 서식 ⑥ 상시", float(res["used"][per][0]), used[per]["t"])
        put(f"{per} 서식 ⑦ 청년등", float(res["used"][per][1]), used[per]["y"])
        put(f"{per} 2 · 3차 판정 상시", float(res["legacy"][per][0]), legacy[per]["t"])
        put(f"{per} 2 · 3차 판정 청년등", float(res["legacy"][per][1]), legacy[per]["y"])
    put("청년등 공제 인원", float(res["n_y"]), c["n_y"])
    put("청년등 외 공제 인원", float(res["n_o"]), c["n_o"])
    for label, key, cell in (("⑳ 1차년도 세액공제액", "first", "20"), ("㉖ 2차년도 세액공제액", "second", "26"),
                             ("㉜ 3차년도 세액공제액", "third", "32"), ("㊱ 추가공제 세액공제액", "extra", "36"),
                             ("⑥ 세액공제액", "total", "total")):
        put(label, res[key], c[cell], X.MONEY)
    put("서식 10호의9 ⑥ 세액공제액", res["total"], form["total"], X.MONEY)
    for per in ("해당", "직전"):
        x = res["counts"][per]
        if per == "직전" and not x.period.has_payroll:
            continue
        for i, rw in enumerate(x.rows):
            j = JUDGE_FIRST + i
            if not any(rw.worked):
                continue
            put(f"{per} {rw.person.name}({rw.person.id}) 계산한 수", float(rw.n_work), ref(S_JUDGE[per], f"${L(J_N21)}${j}"))
            y = rw.n_youth_deem if per == "해당" else rw.n_youth_age
            put(f"{per} {rw.person.name}({rw.person.id}) 청년등 계산한 수", float(y), ref(S_JUDGE[per], f"${L(J_N23)}${j}"))
    out["_range"] = f"$D${row + 1}:$D${r - 1}"
    X.widths(ws, [44, 16, 16, 10])
    return out


def _list(ws, inp, res, k: KV, cnt, c: KV, form, roster, memos, staff_last, pay_last, n_pay) -> None:
    co = inp.company
    row = X.title(ws, f"통합고용세액공제 검증({co.name} {co.fy_start:%Y-%m-%d} ~ {co.fy_end:%Y-%m-%d})",
                  "급여대장 · 사원명부에서 상시근로자 수를 세어 조세특례제한법 시행규칙 별지 제10호의9서식과 제10호의10서식을 채웁니다")
    X.header(ws, row, ["시트", "하는 일"])
    steps = [(S_BASE, "회사 정보 · 계산 방식 · 1인당 공제금액. 노란 칸 입력"), (S_STAFF, "사원명부 입력과 사람별 판정(제외 사유 · 청년 나이 · 간주 기간)"),
             (S_PAY, "월별 급여대장 입력(지급액 · 소득세 · 국민연금 · 건강보험)"),
             ("3판정_해당 · 직전 · 직전전", "사람마다 달마다 월말 재직과 청년등 여부, 근무개월수와 계산한 수"),
             (S_COUNT, "과세연도별 인원(근로자별 방식 · 월말 합 방식)과 서식에 쓰는 인원"), (S_CREDIT, "1 · 2 · 3차 연도와 추가공제 금액"),
             (S_FORM, "공제세액계산서(법정 서식)"), (S_ROSTER, "상시근로자 명세서(법정 서식)"), (S_CHECK, "파이썬 엔진 값과 수식 값 대조")]
    for i, (name, text) in enumerate(steps):
        X.cell(ws, row + 1 + i, 1, name, kind="total")
        X.cell(ws, row + 1 + i, 2, text)
    row += len(steps) + 2

    ws.cell(row=row, column=1, value="결과").font = Font(name=X.FONT, bold=True, size=11)
    row += 1
    X.header(ws, row, ["항목", "값"])
    u = cnt["used"]
    for label, f, fmt in (("해당 과세연도 상시근로자 수", f"={u['해당']['t']}", NUM2), ("직전 과세연도 상시근로자 수", f"={u['직전']['t']}", NUM2),
                          ("상시근로자 증가 인원 수", f"={c['13']}", NUM2), ("청년등 증가 인원 수", f"={c['16']}", NUM2),
                          ("1차년도 세액공제액", f"={c['20']}", X.MONEY), ("2차년도 세액공제액", f"={c['26']}", X.MONEY),
                          ("3차년도 세액공제액", f"={c['32']}", X.MONEY), ("추가공제 세액공제액", f"={c['36']}", X.MONEY),
                          ("세액공제액", f"={c['total']}", X.MONEY)):
        row += 1
        X.cell(ws, row, 1, label)
        X.cell(ws, row, 2, f).number_format = fmt
    row += 2

    ws.cell(row=row, column=1, value="검산").font = Font(name=X.FONT, bold=True, size=11)
    row += 1
    X.header(ws, row, ["검산", "결과", "왼쪽 값", "오른쪽 값", "어긋나면 볼 곳"])
    raw = cnt["raw"]
    staff_ids = ref(S_STAFF, f"$A${STAFF_FIRST}:$A${staff_last}")
    pay_ids = ref(S_PAY, f"$B${PAY_FIRST}:$B${pay_last}")
    rc = cnt["recon"]
    checks = [
        ("서식 10호의9 ⑥ 세액공제액 = 5공제액 합계", form["total"], c["total"], "6서식10의9 · 5공제액 수식"),
        ("서식 10호의9 ⑪ = 명세서 ② 해당 과세연도 상시근로자 수", f'IF({k["method"]}="근로자별",{form["11"]},{roster["해당"]["total"]})',
         roster["해당"]["total"], "근로자별 방식일 때 공제세액계산서와 명세서가 같아야 합니다"),
        ("서식 10호의9 ⑭ = 명세서 ⑨ 해당 과세연도 청년등 합계", f'IF({k["method"]}="근로자별",{form["14"]},{roster["해당"]["youth"]})',
         roster["해당"]["youth"], "3판정_해당 청년등 구분"),
        ("해당 과세연도: 월말 재직인데 급여가 없는 달 수 = 0", raw["해당"]["nopay"], "0", "3판정_해당 U 열. 그 달은 근무개월 수에서 뺐습니다. 무급휴직 · 급여대장 누락 · 입사일 오류"),
        ("해당 과세연도: 재직 기간이 아닌데 급여가 있는 달 수 = 0", raw["해당"]["odd"], "0", "3판정_해당 V 열. 입사일 · 퇴사일 오류, 퇴직 뒤 정산 지급"),
        ("직전 과세연도: 월말 재직인데 급여가 없는 달 수 = 0", raw["직전"]["nopay"], "0", "3판정_직전 U 열"),
        ("직전 과세연도: 재직 기간이 아닌데 급여가 있는 달 수 = 0", raw["직전"]["odd"], "0", "3판정_직전 V 열"),
        ("급여대장 사번이 모두 사원명부에 있음(없는 줄 수 = 0)",
         f"SUMPRODUCT(--(COUNTIF({staff_ids},{pay_ids})=0))" if n_pay else "0", "0", "2급여대장 사번 · 1사원명부 누락"),
        ("전기 신고 직전 과세연도 상시근로자 수 = 다시 센 인원(차이 0)", f'N({rc[("직전", "상시")]})', "0", "4인원 맨 아래 표"),
        ("전기 신고 직전 과세연도 청년등 수 = 다시 센 인원(차이 0)", f'N({rc[("직전", "청년등")]})', "0", "4인원 맨 아래 표"),
        ("엔진 대조 불일치 0건", f'COUNTIF({ref(S_CHECK, "$D$1:$D$5000")},"불일치")', "0", "8엔진대조. 수식이나 입력을 고친 뒤면 ctax employ 를 다시 돌립니다"),
    ]
    first = row + 1
    for label, left, right, hint in checks:
        row += 1
        X.cell(ws, row, 1, label)
        X.cell(ws, row, 3, f"={left}").number_format = "#,##0.00"
        X.cell(ws, row, 4, f"={right}").number_format = "#,##0.00"
        X.cell(ws, row, 2, f'=IF(ABS(N(C{row})-N(D{row}))<0.005,"통과","확인")', kind="total")
        X.cell(ws, row, 5, hint).font = X.NOTE_FONT
    ws.cell(row=row + 1, column=1, value="검산 중 「확인」 개수").font = X.NOTE_FONT
    X.cell(ws, row + 1, 2, f'=COUNTIF(B{first}:B{row},"확인")', kind="total")
    row += 3

    ws.cell(row=row, column=1, value="확인할 것(엔진이 만든 글, 숫자는 위 시트의 수식이 정본입니다)").font = Font(name=X.FONT, bold=True, size=11)
    for text in memos:
        row += 1
        cell = ws.cell(row=row, column=1, value=f"· {text}")
        cell.font = X.BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.row_dimensions[row].height = max(16, 15 * (1 + len(text) // 95))
    X.widths(ws, [58, 16, 16, 16, 70])

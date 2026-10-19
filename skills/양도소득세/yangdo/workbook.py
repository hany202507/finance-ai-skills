# -*- coding: utf-8 -*-
"""계산근거 워크북. 입력을 바꾸면 세액이 따라 바뀌도록 수식으로 쓴다.
첫 시트 「요약」은 자산마다 계산 흐름을 세로로 보여 주고, 값은 「계산」 시트를 수식으로 가져온다. 「근거」 시트는 근거 조문 목록이다.
recalc 는 formulas 로 워크북을 독립 재계산해 셀 값을 돌려준다(엔진과 대조용)."""
import contextlib
import io
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from yangdo import calc

COL = {"id": "A", "종류": "B", "세율": "C", "단일세율": "D", "가산세율": "E", "양도가액": "F", "취득가액": "G",
       "필요경비": "H", "양도차익": "I", "고가주택": "J", "고가주택기준": "K", "과세양도차익": "L", "장특공률": "M",
       "장특공": "N", "양도소득금액": "O", "기본공제": "P", "과세표준": "Q", "산출세액": "R", "지방소득세": "S",
       "전액비과세": "T", "지방단일세율": "U", "지방가산세율": "V", "전체양도가액": "W", "합산묶음": "X"}
ROW_CHECK = ("양도차익", "과세양도차익", "장특공", "양도소득금액", "과세표준", "산출세액", "지방소득세")

# 서식. 금액은 천 단위 쉼표, 0 은 「-」, 율은 퍼센트로 보인다. 입력값(엔진이 넣은 숫자)은 파란 글씨, 수식은 검은 글씨다
FONT = "맑은 고딕"
MONEY = '#,##0;-#,##0;"-"'
RATE = '0.0%;-0.0%;"-"'
FLAG = '0'
MONEY_COLS = ("양도가액", "취득가액", "필요경비", "양도차익", "고가주택기준", "과세양도차익", "장특공", "양도소득금액",
              "기본공제", "과세표준", "산출세액", "지방소득세", "전체양도가액")
RATE_COLS = ("단일세율", "가산세율", "장특공률", "지방단일세율", "지방가산세율")
INPUT_COLS = ("양도가액", "취득가액", "필요경비", "고가주택", "고가주택기준", "장특공률", "기본공제", "전액비과세",
              "단일세율", "가산세율", "지방단일세율", "지방가산세율", "전체양도가액")
HELPER_COLS = ("단일세율", "가산세율", "고가주택", "고가주택기준", "전액비과세", "지방단일세율", "지방가산세율",
               "전체양도가액", "합산묶음")   # 수식이 참조하는 보조 열. 머리글을 회색으로 구분한다
NAVY, GREY, BLUE, MUTED = "1F3864", "7F7F7F", "0000FF", "595959"
FILL_HEAD, FILL_HELP = PatternFill("solid", fgColor=NAVY), PatternFill("solid", fgColor=GREY)
FILL_SECTION, FILL_KEY = PatternFill("solid", fgColor="D9E1F2"), PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 요약 시트의 줄. (종류, 이름, 계산 시트 열 이름 또는 판정 키, 서식, 설명). 종류 sec 는 구역 머리, link 는 계산 시트 수식,
# key 는 굵게 칠하는 link, info 는 엔진 판정 글자(수식 아님)
SUMMARY = (
    ("sec", "자산과 판정", None, None, None),
    ("link", "자산 종류", "종류", None, None),
    ("info", "취득일", "취득일", None, "취득 잔금일(등기접수일이 먼저면 그 날)"),
    ("info", "양도일", "양도일", None, "양도 잔금일(등기접수일이 먼저면 그 날)"),
    ("info", "보유기간", "보유년", '0"년"', "장기보유특별공제 보유기간"),
    ("link", "적용 세율", "세율", None, "세율 그룹·종류(세율 숫자는 세율표 시트)"),
    ("info", "장기보유특별공제 표", "장특공", None, "표1 일반, 표2 1세대1주택"),
    ("sec", "양도차익", None, None, None),
    ("link", "양도가액", "양도가액", MONEY, "계약서 전체 금액 × 본인 지분"),
    ("link", "취득가액", "취득가액", MONEY, None),
    ("link", "필요경비", "필요경비", MONEY, "취득세·취득부대비용·자본적지출·양도비"),
    ("link", "양도차익", "양도차익", MONEY, "양도가액 − 취득가액 − 필요경비"),
    ("link", "과세 양도차익", "과세양도차익", MONEY, "고가주택이면 12억원 넘는 부분만, 전액 비과세면 0"),
    ("sec", "과세표준", None, None, None),
    ("link", "장기보유특별공제율", "장특공률", RATE, None),
    ("link", "장기보유특별공제", "장특공", MONEY, "과세 양도차익 × 공제율"),
    ("link", "양도소득금액", "양도소득금액", MONEY, "과세 양도차익 − 장기보유특별공제"),
    ("link", "양도소득 기본공제", "기본공제", MONEY, "연 250만원, 먼저 양도한 자산부터"),
    ("key", "과세표준", "과세표준", MONEY, "양도소득금액 − 기본공제"),
    ("sec", "세액", None, None, None),
    ("key", "산출세액", "산출세액", MONEY, "과세표준 × 세율(세율표 시트)"),
    ("key", "지방소득세", "지방소득세", MONEY, "과세표준 × 지방세율(지방세율표 시트)"),
)
SUMMARY_TOP = 5   # 요약 시트 첫 줄(머리글 줄 바로 아래)
# 요약 시트 합계 구역. (이름, result["합계"] 키, 굵게 칠할지, 설명). 값은 모두 B 열 하나에 둔다
TOTALS = (("자산별 세액 합", "자산별세액", False, "자산마다 계산한 세액의 합"),
          ("호별 합산 세액", "호별합산세액", False, "같은 세율끼리 과세표준을 합쳐 계산한 세액의 합"),
          ("합산 비교 세액", "합산비교세액", False, "과세 자산이 둘 이상이면 과세표준을 모두 합쳐 기본세율로 계산"),
          ("산출세액", "산출세액", True, "호별 합산 세액과 합산 비교 세액 중 큰 금액(소득세법 제104조⑤)"),
          ("지방소득세", "지방소득세", True, "같은 방법으로 지방세율표를 적용"))


def summary_rows():
    """요약 시트의 줄 번호. 계산 시트 열 이름 -> 줄(link·key 줄만), 합계 키 -> 줄."""
    rows, r = {}, SUMMARY_TOP
    for kind, _label, key, _fmt, _note in SUMMARY:
        if kind in ("link", "key"):
            rows[key] = r
        r += 1
    r += 2   # 빈 줄 하나, 합계 구역 머리
    totals = {}
    for _label, key, _bold, _note in TOTALS:
        totals[key] = r
        r += 1
    return rows, totals


def summary_cells(n, groups):
    """n 은 자산 행 수, groups 는 합산묶음 수. 묶음 행은 n+7 부터, 그 바로 아래가 호별 합산 행이다."""
    a, b, c, h = n + 3, n + 4, n + 5, n + 7 + groups
    return {"자산별세액": "R%d" % a, "호별합산세액": "R%d" % h, "합산비교세액": "R%d" % b, "산출세액": "R%d" % c,
            "지방_자산별": "S%d" % a, "지방_호별합산": "S%d" % h, "지방_합산비교": "S%d" % b, "지방소득세": "S%d" % c}


def _groups(rows):
    """합산묶음 이름 -> 대표 구성원의 입력 순번(0부터). 대표는 묶음 세율의 기준일이 되는 가장 이른 양도일의 자산이고
    엔진이 쓰는 calc.group_anchor 가 정한다(정의는 한 곳). 과세 자산만, 대표의 (양도일, 묶음 이름) 순서라 입력 순서와 상관없다."""
    members = {}
    for i, row in enumerate(rows):
        key = row["계산"].get("합산묶음")
        if key:
            members.setdefault(key, []).append(i)
    out = {}
    for key, idx in members.items():
        anchor = calc.group_anchor([rows[i] for i in idx])
        out[key] = next(i for i in idx if rows[i] is anchor)
    return dict(sorted(out.items(), key=lambda kv: (rows[kv[1]]["판정"].get("양도일") or "", kv[0])))


def _num(v):
    return float(v) if v is not None else None


def _rate_label(c):
    """세율 칸 글자. 미등기처럼 세율종류가 없는 행은 「미등기·None」 이 아니라 그룹 이름만 쓴다."""
    return "·".join(str(x) for x in (c.get("세율그룹"), c.get("세율종류")) if x)


TEXT_COLS = ("id", "종류", "세율", "합산묶음")   # 글자를 담는 열. 나머지 열은 숫자이거나 이 모듈이 만든 수식이다
TEXT_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _put(ws, addr, value):
    """셀에 값을 쓴다. 사용자가 적은 글자(자산 id 등)가 수식 문자로 시작해도 수식이 되지 않도록 글자로 고정한다.
    수식은 이 함수를 거치지 않고 직접 쓴다."""
    ws[addr] = value
    if isinstance(value, str) and value.startswith(TEXT_LEAD):
        cell = ws[addr]
        cell.data_type = "s"
        cell.quotePrefix = True


def _p(x, sheet):
    rng = "%s!$A$2:$C$9" % sheet
    return "VLOOKUP({x},{r},2,TRUE)+({x}-VLOOKUP({x},{r},1,TRUE))*VLOOKUP({x},{r},3,TRUE)".format(x=x, r=rng)


def _tax(r, d, e, sheet):
    q = "Q%d" % r
    inner = "IF(ISNUMBER({e}{r}),MAX({p}+{q}*{e}{r},IF(ISNUMBER({d}{r}),{q}*{d}{r},0)),{q}*{d}{r})".format(
        e=e, r=r, d=d, q=q, p=_p(q, sheet))
    return "=IF(T{r}=1,0,ROUNDDOWN(ROUND({i},6),0))".format(r=r, i=inner)


def _style(cell, bold=False, color="000000", size=10, fmt=None, fill=None, border=None, align=None):
    cell.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if align:
        cell.alignment = align


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)


def _table(ws, rows):
    ws.append(["하한", "기준액", "세율"])
    for a, b, c in rows:
        ws.append([int(a), int(b), float(c)])
    for cell in ws[1]:
        _style(cell, bold=True, color="FFFFFF", fill=FILL_HEAD, border=BOX, align=CENTER)
    for line in ws.iter_rows(min_row=2):
        for cell, fmt in zip(line, (MONEY, MONEY, RATE)):
            _style(cell, color=BLUE, fmt=fmt, border=BOX)
    for letter, width in (("A", 16), ("B", 16), ("C", 10)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def _style_calc(ws, n, label_rows, key_row):
    """계산 시트 서식. n 은 자산 행 수, label_rows 는 합계·묶음 줄, key_row 는 최종 산출세액 줄."""
    widths = {"id": 7, "종류": 8, "세율": 16, "합산묶음": 18}
    flags = ("고가주택", "전액비과세")
    for name, letter in COL.items():
        _style(ws["%s1" % letter], bold=True, color="FFFFFF", border=BOX, align=CENTER,
               fill=FILL_HELP if name in HELPER_COLS else FILL_HEAD)
        ws.column_dimensions[letter].width = widths.get(name, 10 if name in HELPER_COLS else 15)
        fmt = MONEY if name in MONEY_COLS else RATE if name in RATE_COLS else FLAG if name in flags else None
        for r in range(2, ws.max_row + 1):
            cell = ws["%s%d" % (letter, r)]
            if r <= n + 1:
                _style(cell, color=BLUE if name in INPUT_COLS else "000000", fmt=fmt, border=BOX)
            elif cell.value is not None:
                _style(cell, bold=r in label_rows and letter == "A", fmt=fmt)
    for letter in ("A", "R", "S"):
        _style(ws["%s%d" % (letter, key_row)], bold=True, fill=FILL_KEY, border=BOX, fmt=None if letter == "A" else MONEY)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _summary(wb, result, totals):
    """첫 시트 「요약」. 자산마다 한 열, 계산 흐름을 위에서 아래로 읽는다. 값은 계산 시트를 수식으로 가져온다."""
    ws = wb.create_sheet("요약", 0)
    ws.sheet_view.showGridLines = False
    rows = result["자산"]
    note_col = 2 + max(len(rows), 1)
    last = note_col
    ws["A1"] = "양도소득세 계산근거"
    _style(ws["A1"], bold=True, size=15, color=NAVY)
    ws["A2"] = ("금액은 원 단위입니다. 값은 「계산」 시트의 수식에서 가져오며, 「계산」 시트의 파란 글씨가 엔진이 넣은 입력값입니다. "
                "근거 조문은 「근거」 시트에 있습니다.")
    _style(ws["A2"], size=9, color=MUTED)
    head = SUMMARY_TOP - 1
    ws.cell(head, 1, "항목")
    for i, row in enumerate(rows):
        _put(ws, "%s%d" % (get_column_letter(2 + i), head), "자산 %s" % row["id"])
    ws.cell(head, note_col, "설명")
    for c in range(1, last + 1):
        _style(ws.cell(head, c), bold=True, color="FFFFFF", fill=FILL_HEAD, border=BOX, align=CENTER)
    r = SUMMARY_TOP
    for kind, label, key, fmt, note in SUMMARY:
        ws.cell(r, 1, label)
        if kind == "sec":
            for c in range(1, last + 1):
                _style(ws.cell(r, c), bold=True, color=NAVY, fill=FILL_SECTION, border=BOX)
            r += 1
            continue
        bold, fill = kind == "key", FILL_KEY if kind == "key" else None
        _style(ws.cell(r, 1), bold=bold, fill=fill, border=BOX)
        for i, row in enumerate(rows):
            addr = "%s%d" % (get_column_letter(2 + i), r)
            if kind == "info":
                val = row["판정"].get(key)
                if key == "장특공" and not val:
                    val = "적용 안 함"
                _put(ws, addr, "-" if val is None else val)
            else:
                ws[addr] = "=계산!%s%d" % (COL[key], i + 2)
            _style(ws[addr], bold=bold, fill=fill, border=BOX, fmt=fmt,
                   align=Alignment(horizontal="right" if fmt else "center"))
        ws.cell(r, note_col, note or "")
        _style(ws.cell(r, note_col), size=9, color=MUTED, border=BOX, align=WRAP)
        r += 1
    r += 1
    ws.cell(r, 1, "합계(전체 자산)")
    ws.cell(r, 2, "금액")
    for c in range(1, last + 1):
        _style(ws.cell(r, c), bold=True, color=NAVY, fill=FILL_SECTION, border=BOX,
               align=Alignment(horizontal="center") if c == 2 else None)
    r += 1
    for label, key, bold, note in TOTALS:
        fill = FILL_KEY if bold else None
        ws.cell(r, 1, label)
        _style(ws.cell(r, 1), bold=bold, fill=fill, border=BOX)
        ws.cell(r, 2, "=계산!%s" % totals[key])
        _style(ws.cell(r, 2), bold=bold, fill=fill, border=BOX, fmt=MONEY)
        for c in range(3, note_col):
            _style(ws.cell(r, c), border=BOX)
        ws.cell(r, note_col, note)
        _style(ws.cell(r, note_col), size=9, color=MUTED, border=BOX, align=WRAP)
        r += 1
    ws.column_dimensions["A"].width = 22
    for c in range(2, note_col):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions[get_column_letter(note_col)].width = 52
    ws.freeze_panes = "B%d" % SUMMARY_TOP
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _basis(wb, result):
    """「근거」 시트. 결과의 근거 조문을 (법령, 조항) 하나에 한 줄로 모은다. 근거가 없으면 만들지 않는다."""
    found = {}
    sources = [("전체", result.get("근거") or [])] + [(row["id"], row["판정"].get("근거") or []) for row in result["자산"]]
    for owner, items in sources:
        for g in items:
            k = (g.get("법령"), g.get("조항"))
            if k not in found:
                found[k] = (g, [])
            if owner not in found[k][1]:
                found[k][1].append(owner)
    if not found:
        return
    ws = wb.create_sheet("근거")
    ws.append(["법령", "조항", "시행일", "판정 항목", "원문 발췌", "자산", "법제처 링크"])
    for g, owners in found.values():
        ws.append([g.get("법령") or "", g.get("조항") or "", g.get("시행일") or "", g.get("key") or "",
                   " / ".join(g.get("발췌") or []), ", ".join(owners), "열기" if g.get("URL") else ""])
        if g.get("URL"):
            ws.cell(ws.max_row, 7).hyperlink = g["URL"]
    for cell in ws[1]:
        _style(cell, bold=True, color="FFFFFF", fill=FILL_HEAD, border=BOX, align=CENTER)
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            link = cell.column == 7 and cell.hyperlink is not None
            _style(cell, color=BLUE if link else "000000", border=BOX, align=WRAP)
            if link:
                cell.font = Font(name=FONT, size=10, color=BLUE, underline="single")
    for letter, width in zip("ABCDEFG", (20, 22, 11, 24, 56, 8, 11)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def write(result, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "계산"
    _table(wb.create_sheet("세율표"), result["세율표"]["국세"])
    _table(wb.create_sheet("지방세율표"), result["세율표"]["지방"])
    ws.append(list(COL))
    rows = result["자산"]
    for r, row in enumerate(rows, start=2):
        v, c = row["판정"], row["계산"]
        rate = c["적용세율"]
        values = {
            "id": row["id"], "종류": v["종류"], "세율": _rate_label(c),
            "단일세율": _num(rate["단일"]), "가산세율": _num(rate["가산"]),
            "양도가액": c["양도가액"], "취득가액": c["취득가액"] or 0, "필요경비": c["필요경비"] or 0,
            "양도차익": "=F{r}-G{r}-H{r}".format(r=r), "고가주택": 1 if v["고가주택"] else 0,
            "고가주택기준": c.get("고가주택기준") or 0,
            "과세양도차익": "=IF(T{r}=1,0,IF(J{r}=1,ROUNDDOWN(ROUND(I{r}*(W{r}-K{r})/W{r},6),0),I{r}))".format(r=r),
            "장특공률": float(c["장특공률"]), "장특공": "=ROUNDDOWN(ROUND(L{r}*M{r},6),0)".format(r=r),
            "양도소득금액": "=L{r}-N{r}".format(r=r), "기본공제": c["기본공제"],
            "과세표준": "=O{r}-P{r}".format(r=r),
            "산출세액": _tax(r, "D", "E", "세율표"), "지방소득세": _tax(r, "U", "V", "지방세율표"),
            "전액비과세": 1 if v["전액비과세"] else 0,
            "지방단일세율": _num(rate["지방단일"]), "지방가산세율": _num(rate["지방가산"]),
            "전체양도가액": c["전체양도가액"], "합산묶음": c.get("합산묶음") or "",
        }
        for name, letter in COL.items():
            if name in TEXT_COLS:
                _put(ws, "%s%d" % (letter, r), values[name])
            else:
                ws["%s%d" % (letter, r)] = values[name]
    n = len(rows)
    last = n + 1
    ws["A%d" % (n + 3)] = "자산별 합"
    ws["A%d" % (n + 4)] = "합산 비교(소득세법 제104조⑤)"
    ws["A%d" % (n + 5)] = "산출세액"
    ws["Q%d" % (n + 4)] = "=SUMIF(T2:T{l},0,Q2:Q{l})".format(l=last)
    groups = _groups(rows)
    first_row = n + 7
    hrow = first_row + len(groups)
    ws["A%d" % hrow] = "호별 합산"
    for g, (key, i) in enumerate(groups.items()):
        gr, src = first_row + g, i + 2
        _put(ws, "A%d" % gr, "묶음 %s" % key)
        ws["Q%d" % gr] = '=SUMIF(X2:X{l},"{k}",Q2:Q{l})'.format(l=last, k=key.replace('"', '""'))
        ws["T%d" % gr] = 0
        for letter in ("D", "E", "U", "V"):
            if ws["%s%d" % (letter, src)].value is not None:
                ws["%s%d" % (letter, gr)] = "=%s%d" % (letter, src)
        ws["R%d" % gr] = _tax(gr, "D", "E", "세율표")
        ws["S%d" % gr] = _tax(gr, "U", "V", "지방세율표")
    for col, sheet in (("R", "세율표"), ("S", "지방세율표")):
        ws["%s%d" % (col, n + 3)] = "=SUM({c}2:{c}{l})".format(c=col, l=last)
        ws["%s%d" % (col, n + 4)] = "=IF(COUNTIF(T2:T{l},0)>=2,ROUNDDOWN(ROUND({p},6),0),0)".format(
            l=last, p=_p("Q%d" % (n + 4), sheet))
        ws["%s%d" % (col, hrow)] = ("=SUM({c}{a}:{c}{b})".format(c=col, a=first_row, b=hrow - 1) if groups else 0)
        ws["%s%d" % (col, n + 5)] = "=MAX({c}{h},{c}{b})".format(c=col, h=hrow, b=n + 4)
    _style_calc(ws, n, set(range(n + 3, hrow + 1)), n + 5)
    ws.sheet_view.tabSelected = False
    _summary(wb, result, summary_cells(n, len(groups)))
    _basis(wb, result)
    wb.active = 0
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


_PROGRESS = re.compile(r"[0-9]+%\||it/s|s/it")


@contextlib.contextmanager
def _quiet_progress():
    """formulas 가 계산 중 표준오류에 찍는 진행 막대(tqdm)를 가린다. 진행 막대가 아닌 출력은 그대로 다시 내보낸다."""
    buf = io.StringIO()
    try:
        with contextlib.redirect_stderr(buf):
            yield
    finally:
        for chunk in re.split(r"[\r\n]+", buf.getvalue()):
            if chunk.strip() and not _PROGRESS.search(chunk):
                print(chunk, file=sys.stderr)


def recalc(path):
    import formulas
    base = os.path.basename(path).upper()
    with _quiet_progress():
        sol = formulas.ExcelModel().loads(path).finish().calculate()
    out = {}
    for k, v in sol.items():
        ku = k.upper()
        if not ku.startswith("'[%s]" % base) or "!" not in ku:
            continue
        sheet = ku.split("]", 1)[1].split("'!", 1)[0]
        addr = ku.split("!", 1)[1]
        try:
            val = v.value[0, 0]
        except Exception:
            val = v.value
        out[(sheet, addr)] = val
    return out


def _eq(got, want):
    try:
        return int(round(float(got))) == int(want or 0)
    except (TypeError, ValueError):
        return False


def compare(path, result):
    cells = recalc(path)
    fails = []
    for i, row in enumerate(result["자산"], start=2):
        c = row["계산"]
        for name in ROW_CHECK:
            if c[name] is None:  # 전액비과세 자산의 양도차익처럼 계산하지 않은 값
                continue
            got = cells.get(("계산", "%s%d" % (COL[name], i)))
            if not _eq(got, c[name]):
                fails.append("W %s %s: 워크북 %s, 엔진 %s" % (row["id"], name, got, c[name]))
    groups = len(_groups(result["자산"]))
    for name, addr in summary_cells(len(result["자산"]), groups).items():
        got = cells.get(("계산", addr))
        if not _eq(got, result["합계"][name]):
            fails.append("W 합계 %s: 워크북 %s, 엔진 %s" % (name, got, result["합계"][name]))
    # 요약 시트가 계산 시트의 맞는 칸을 가져오는지도 본다(줄이나 열이 어긋나도 오류 없이 틀린 값이 나올 수 있다)
    srows, stotals = summary_rows()
    for i, row in enumerate(result["자산"]):
        c = row["계산"]
        for name in ROW_CHECK:
            if c[name] is None:
                continue
            got = cells.get(("요약", "%s%d" % (get_column_letter(2 + i), srows[name])))
            if not _eq(got, c[name]):
                fails.append("W 요약 %s %s: 워크북 %s, 엔진 %s" % (row["id"], name, got, c[name]))
    for name, r in stotals.items():
        got = cells.get(("요약", "B%d" % r))
        if not _eq(got, result["합계"][name]):
            fails.append("W 요약 합계 %s: 워크북 %s, 엔진 %s" % (name, got, result["합계"][name]))
    return fails

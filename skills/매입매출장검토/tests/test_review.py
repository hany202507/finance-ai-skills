# -*- coding: utf-8 -*-
"""합성 장부로 자동 검토 10개 항목, 깨끗한 장부, 수정분개 반영 후 재검토를 확인한다."""
import os
import sys
import glob
import subprocess

import openpyxl
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(HERE, "..", "scripts")
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, HERE)

import review  # noqa: E402
import make_fixture  # noqa: E402

AUTO = ["E3", "F1", "C1", "B7", "B5", "B6", "F2", "B10", "A5", "D2"]
ENV = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")


def cli(script, *args):
    return subprocess.run([sys.executable, os.path.join(SCRIPTS, script), *args],
                          capture_output=True, text=True, encoding="utf-8", env=ENV)


def stops(rep):
    return [f for f in rep.findings if f["등급"] == review.STOP]


def codes(rep, grade=None):
    return {f["검사"] for f in rep.findings if grade is None or f["등급"] == grade}


def rows(path):
    ws = openpyxl.load_workbook(path).active
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it]


@pytest.fixture()
def bad(tmp_path):
    return make_fixture.build(str(tmp_path / "오류판"), "오류판")


@pytest.fixture()
def clean(tmp_path):
    return make_fixture.build(str(tmp_path / "깨끗한판"), "깨끗한판")


# ── 오류판 ───────────────────────────────────────────────────
def test_오류판_자동항목_10개가_전부_걸린다(bad):
    rep, period, led, gjs = review.run(bad)
    assert period == "2026-04~06"
    assert set(AUTO) <= codes(rep), f"안 걸린 항목: {set(AUTO) - codes(rep)}"
    assert rep.skipped == {}


@pytest.mark.parametrize("code,grade", [
    ("E3", "정지"), ("E3", "확인"), ("F1", "정지"), ("C1", "정지"), ("B7", "정지"),
    ("B5", "정지"), ("B6", "정지"), ("F2", "정지"), ("B10", "정지"), ("A5", "확인"),
    ("D2", "확인"),
])
def test_오류판_항목별_판정(bad, code, grade):
    rep = review.run(bad)[0]
    assert code in codes(rep, grade)


def test_오류판_금액(bad):
    rep = review.run(bad)[0]
    by = {}
    for f in rep.findings:
        by.setdefault((f["검사"], f["무엇"]), f)
    f1 = next(f for k, f in by.items() if k[0] == "F1")
    assert "PLTF" in f1["대상"] and "2,200,000원" in f1["대상"]
    c1 = next(f for k, f in by.items() if k[0] == "C1")
    assert "D2C" in c1["대상"] and c1["공급가액 영향"] == -500_000
    b6 = [f for k, f in by.items() if k[0] == "B6"]
    assert len(b6) == 2     # 세금계산서·카드 중복, 카드대금 결제
    f2 = [f for k, f in by.items() if k[0] == "F2"]
    assert len(f2) == 2     # 일반전표로 간 세금계산서, 고정자산 계정


def test_B7_수정분개는_외상매출금을_되살리고_선수금을_세운다(bad):
    rep = review.run(bad)[0]
    b7 = [g for g in rep.gj if "[B7]" in g["적요"]]
    ar = next(g for g in b7 if g["계정과목코드"] == 108)
    adv = next(g for g in b7 if g["계정과목코드"] == 259)
    assert ar["차변"] == 1_100_000 and not ar["대변"]
    assert adv["대변"] == 1_100_000 and not adv["차변"]


def test_F2_일반전표_원건을_찾아_되돌린다(bad):
    rep = review.run(bad)[0]
    f2_mm = next(m for m in rep.mm if "[F2] 매입매출전표로 재계상" in m["품명"])
    assert (f2_mm["기본계정"], f2_mm["상대계정"]) == ("824", "251")
    rev = [g for g in rep.gj if "[F2]" in g["적요"]]
    assert sorted((g["계정과목코드"], g["차변"], g["대변"]) for g in rev) == \
        [(251, None, -1_100_000), (824, -1_100_000, None)]


def test_F2_고정자산은_품명키워드로_계정을_고른다(bad):
    rep = review.run(bad)[0]
    re_ = next(m for m in rep.mm if "[F2] 고정자산매입 재계상" in m["품명"])
    assert re_["기본계정"] == "212"


def test_수정분개는_차대가_맞는다(bad):
    rep = review.run(bad)[0]
    dr = sum(g["차변"] or 0 for g in rep.gj)
    cr = sum(g["대변"] or 0 for g in rep.gj)
    assert dr == cr


# ── 깨끗한판 ─────────────────────────────────────────────────
def test_깨끗한판은_정지가_없다(clean):
    rep = review.run(clean)[0]
    assert stops(rep) == []
    assert rep.findings == []
    assert rep.mm == [] and rep.gj == []
    assert rep.skipped == {}, "자동 항목이 미실행으로 빠지면 통과가 아니다"


def test_매출집계가_없으면_orders_를_접어_쓴다(clean):
    assert not glob.glob(os.path.join(clean, "매출집계*"))
    rep = review.run(clean)[0]
    assert "F1" not in rep.skipped and "C1" not in rep.skipped


# ── 반영 후 재검토 ───────────────────────────────────────────
def test_수정분개_반영후_재검토에서_정지가_닫힌다(bad):
    r1 = cli("review.py", bad)
    assert r1.returncode == 1, r1.stdout + r1.stderr        # 정지가 있으면 1
    out = os.path.join(bad, "검토산출")
    for pat in ("검토리포트_*.md", "확인사항_*.md", "수정_매입매출전표_*.xlsx",
                "수정_일반전표_*.xlsx", "수정분개_근거_*.xlsx"):
        assert glob.glob(os.path.join(out, pat)), pat

    r2 = cli("apply_fix.py", bad)
    assert r2.returncode == 0, r2.stdout + r2.stderr
    fixed = os.path.join(bad, "수정후")
    led = rows(glob.glob(os.path.join(fixed, "매입매출장_*.xlsx"))[0])
    assert any(r["거래처명"] == "가나문구" and r["사업자(주민)등록번호"] == "0000000101"
               and r["월"] == 4 and r["일"] == 5 for r in led), "E3 는 원건 칸을 채운다"

    rep = review.run(fixed)[0]
    assert stops(rep) == [], [f["검사"] for f in stops(rep)]
    # 확인 판정은 수정분개로 닫히지 않는다. 증빙을 받거나 마스터에 등록해야 한다
    assert codes(rep) == {"E3", "A5", "D2"}

    r3 = cli("review.py", fixed)
    assert r3.returncode == 0, r3.stdout + r3.stderr


def test_원본_장부는_고치지_않는다(bad):
    before = rows(glob.glob(os.path.join(bad, "매입매출장_*.xlsx"))[0])
    cli("review.py", bad)
    cli("apply_fix.py", bad)
    after = rows(glob.glob(os.path.join(bad, "매입매출장_*.xlsx"))[0])
    assert before == after


def test_리포트_검토범위에_32개가_다_찍힌다(bad):
    rep, period, led, gjs = review.run(bad)
    md = review.report_md(rep, period, led, gjs)
    cat = review.rule("진단항목.csv")
    assert len(cat) == 32
    for r in cat:
        assert f"| {r['번호']} |" in md


def test_B7_일반전표에_입점사_입금·지급이_있으면_그_줄을_선수금으로_되돌린다(bad):
    """직원이 입점사 판매대금 입금을 외상매출금 회수로, 입점사 지급을 지급수수료로 잡은 경우.
    매출 취소 + 입금·지급 재분류를 하면 외상매출금과 지급수수료는 0, 선수금은 수수료만 남는다."""
    p = glob.glob(os.path.join(bad, "일반전표_*.xlsx"))[0]
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    head = [c.value for c in ws[1]]
    for row in (
            {"월": 4, "일": 25, "계정과목코드": 103, "계정과목명": "보통예금",
             "적요": "입점사 판매대금 수취", "차변": 1_100_000},
            {"월": 4, "일": 25, "계정과목코드": 108, "계정과목명": "외상매출금",
             "적요": "입점사 판매대금 수취", "대변": 1_100_000},
            {"월": 5, "일": 10, "계정과목코드": 831, "계정과목명": "지급수수료",
             "적요": "입점사 지급", "차변": 968_000},
            {"월": 5, "일": 10, "계정과목코드": 103, "계정과목명": "보통예금",
             "적요": "입점사 지급", "대변": 968_000}):
        ws.append([row.get(h) for h in head])
    wb.save(p)

    rep = review.run(bad)[0]
    b7 = [g for g in rep.gj if "[B7]" in g["적요"]]
    net = {}
    for g in b7:
        net[g["계정과목코드"]] = net.get(g["계정과목코드"], 0) + (g["차변"] or 0) - (g["대변"] or 0)
    assert net == {108: 1_100_000, 259: -1_100_000 + 968_000, 831: -968_000}
    # 매출 취소(외상매출금 -1,100,000) + 입금 재분류(+1,100,000) + 원건(+1,100,000 -1,100,000) = 0
    b7f = next(f for f in rep.findings if f["검사"] == "B7")
    assert b7f["잠정"] == "아니오"


def test_B7_일반전표에_입점사_줄이_없으면_잠정으로_올린다(bad):
    rep = review.run(bad)[0]
    b7f = next(f for f in rep.findings if f["검사"] == "B7")
    assert b7f["잠정"].startswith("예")


def _append(path, rows):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    head = [c.value for c in ws[1]]
    for row in rows:
        ws.append([row.get(h) for h in head])
    wb.save(path)


def test_C1_다른_채널의_다음기_배송분을_한_채널에_넣었어도_전부_뺀다(bad):
    """플랫폼 주문(6월 주문·7월 배송)을 자사몰 매출로 넣었다. 자사몰 초과분은 두 채널 몫이다."""
    _append(glob.glob(os.path.join(bad, "매출집계_*.xlsx"))[0],
            [{"배송완료일": "2026-07-03", "채널": "PLTF", "과세구분": "TAXABLE", "취소여부": "정상",
              "주문월": "2026-06", "건수": 1, "판매금액": 550_000}])
    _append(glob.glob(os.path.join(bad, "매입매출장_*.xlsx"))[0],
            [make_fixture.mm("2026-06-30", 1, 17, "자사몰(집계)", 500_000, 50_000, "D2C 과세매출",
                             "401", "108", biz="")])
    rep = review.run(bad)[0]
    c1 = [f for f in rep.findings if f["검사"] == "C1"]
    assert len(c1) == 1 and c1[0]["공급가액 영향"] == -1_000_000


def test_B6_일반전표에_카드대금_결제가_이미_있으면_결제분개를_또_만들지_않는다(bad):
    _append(glob.glob(os.path.join(bad, "일반전표_*.xlsx"))[0], [
        {"월": 5, "일": 25, "계정과목코드": 253, "계정과목명": "미지급금", "거래처": "○○카드",
         "적요": "법인카드 대금 결제", "차변": 440_000},
        {"월": 5, "일": 25, "계정과목코드": 103, "계정과목명": "보통예금", "거래처": "○○카드",
         "적요": "법인카드 대금 결제", "대변": 440_000}])
    rep = review.run(bad)[0]
    assert any("[B6] 카드대금 결제 매입계상 취소" in m["품명"] for m in rep.mm)
    assert not [g for g in rep.gj if "[B6]" in g["적요"]]

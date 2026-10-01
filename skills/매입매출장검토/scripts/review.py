# -*- coding: utf-8 -*-
"""매입매출장과 일반전표를 원천과 대조해 검토 리포트와 수정분개를 만든다.

  python review.py <장부폴더> [산출폴더]

산출
  검토리포트_<기간>.md            검출 건마다 대상·기대·실제·볼 곳·근거
  수정_매입매출전표_<기간>.xlsx    더존 업로드 양식. 원건 음수, 정건 양수
  수정_일반전표_<기간>.xlsx        같음
  수정분개_근거_<기간>.xlsx        어느 검사에서 나왔고 무엇을 고쳤나
  확인사항_<기간>.md              회계사가 정해야 할 것

숫자를 세는 일은 전부 이 코드가 한다. AI 가 합계를 내면 다음 달에 다르게 낸다.
수정분개를 만들 뿐 원장에 적용하지 않는다. 회계사가 보고 더존에 올린다.
"""
import os
import io
import csv
import sys
import glob
from collections import defaultdict

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
RULES = os.path.normpath(os.path.join(HERE, "..", "rules"))

MM_COLS = ["년도", "월", "일", "매입매출구분(1-매출/2-매입)", "과세유형", "불공제사유",
           "신용카드거래처코드", "신용카드사명", "신용카드(가맹점)번호", "거래처명",
           "사업자(주민)등록번호", "공급가액", "부가세", "품명", "전자세금(1.전자)",
           "기본계정", "상대계정", "현금영수증 승인번호"]
GJ_COLS = ["월", "일", "구분", "계정과목코드", "계정과목명", "거래처코드", "거래처",
           "적요", "차변", "대변"]
NOTE_COLS = ["검사", "등급", "무엇", "대상", "기대", "실제", "볼 곳", "근거",
             "공급가액 영향", "세액 영향", "잠정"]

STOP, ASK = "정지", "확인"


# ── 읽기 ─────────────────────────────────────────────────────
def _rows_from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    head = next(it, None)
    if not head:
        return []
    head = [str(c).strip() if c is not None else "" for c in head]
    out = []
    for r in it:
        if all(c is None for c in r):
            continue
        out.append({head[k]: r[k] for k in range(min(len(head), len(r)))})
    wb.close()
    return out


def _rows_from_csv(path):
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load(folder):
    """폴더의 엑셀과 CSV 를 표 이름 -> 행목록 으로 올린다. 같은 이름이면 엑셀."""
    tables = {}
    for path in sorted(glob.glob(os.path.join(folder, "*.csv"))) + \
            sorted(glob.glob(os.path.join(folder, "**", "*.csv"), recursive=True)):
        tables[os.path.splitext(os.path.basename(path))[0]] = ("csv", path)
    for path in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
        if os.path.basename(path).startswith("~$"):
            continue
        tables[os.path.splitext(os.path.basename(path))[0]] = ("xlsx", path)
    out = {}
    for name, (kind, path) in tables.items():
        out[name] = (_rows_from_xlsx if kind == "xlsx" else _rows_from_csv)(path)
    return out


def pick(tables, prefix):
    for name in tables:
        if name.startswith(prefix):
            return tables[name]
    return []


def sales_from_orders(orders):
    """매출집계가 없을 때 주문내역(orders)을 매출집계 모양으로 접는다.
    배송완료일 · 채널 · 과세구분 · 취소여부 · 주문월 별 판매금액 합계."""
    agg = defaultdict(int)
    for r in orders:
        k = (s(r.get("배송완료일"))[:10], s(r.get("채널")), s(r.get("과세구분")),
             s(r.get("취소여부")), s(r.get("주문일"))[:7])
        agg[k] += n(r.get("판매금액"))
    return [{"배송완료일": a, "채널": b, "과세구분": c, "취소여부": d, "주문월": e,
             "판매금액": v} for (a, b, c, d, e), v in sorted(agg.items())]


def rule(name):
    p = os.path.join(RULES, name)
    if not os.path.exists(p):
        return []
    return _rows_from_csv(p)


def n(x):
    if x in (None, ""):
        return 0
    if isinstance(x, (int, float)):
        return int(x)
    try:
        return int(float(str(x).replace(",", "")))
    except ValueError:
        return 0


def s(x):
    return "" if x is None else str(x).strip()


def d3(r):
    return f"{n(r.get('년도')):04d}-{n(r.get('월')):02d}-{n(r.get('일')):02d}"


def net_off(rows):
    """음수 행과 짝이 맞는 원건을 상계해 뺀다.

    더존에서 잘못 친 전표를 음수로 지우고 다시 넣으면 장부에 세 행이 남는다.
    검토는 남은 순액을 봐야 한다. 안 그러면 이미 고친 것이 매번 다시 걸린다.
    짝은 일자·거래처·과세유형·기본계정·금액이 같고 부호만 반대인 것이다."""
    def key(r):
        return (d3(r), s(r["거래처명"]), s(r["사업자(주민)등록번호"]),
                n(r["과세유형"]), s(r["기본계정"]), abs(n(r["공급가액"])),
                abs(n(r["부가세"])))

    pos, neg = defaultdict(list), defaultdict(list)
    for r in rows:
        (neg if n(r["공급가액"]) < 0 or n(r["부가세"]) < 0 else pos)[key(r)].append(r)
    drop = set()
    for k, negs in neg.items():
        for r in negs:
            if pos.get(k):                 # 짝이 있을 때만 뺀다
                drop.add(id(r))
                drop.add(id(pos[k].pop()))
    return [r for r in rows if id(r) not in drop]


# ── 검출 ─────────────────────────────────────────────────────
class Report:
    def __init__(self):
        self.findings = []
        self.mm = []
        self.gj = []
        self.skipped = {}      # 검사번호 -> 못 돈 이유. 검토 범위 표에 「미실행」으로 찍힌다

    def add(self, check, grade, what, target, expect, actual, look, ground,
            sup=0, tax=0, tentative="아니오"):
        self.findings.append({"검사": check, "등급": grade, "무엇": what,
                              "대상": target, "기대": expect, "실제": actual,
                              "볼 곳": look, "근거": ground,
                              "공급가액 영향": sup, "세액 영향": tax,
                              "잠정": tentative})

    def fix_mm(self, dt, gubun, vt, name, biz, sup, vat, item, base="", contra="",
               card_co="", elec=""):
        self.mm.append({"년도": int(dt[:4]), "월": int(dt[5:7]), "일": int(dt[8:10]),
                        "매입매출구분(1-매출/2-매입)": gubun, "과세유형": vt,
                        "불공제사유": "", "신용카드거래처코드": "",
                        "신용카드사명": card_co, "신용카드(가맹점)번호": "",
                        "거래처명": name, "사업자(주민)등록번호": biz,
                        "공급가액": sup, "부가세": vat, "품명": item,
                        "전자세금(1.전자)": elec, "기본계정": base,
                        "상대계정": contra, "현금영수증 승인번호": ""})

    def fix_gj(self, dt, code, cname, pname, memo, debit, credit):
        self.gj.append({"월": int(dt[5:7]), "일": int(dt[8:10]),
                        "구분": 3 if debit else 4, "계정과목코드": code,
                        "계정과목명": cname, "거래처코드": "", "거래처": pname,
                        "적요": memo, "차변": debit or None, "대변": credit or None})


def month_end(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    last = [31, 29 if y % 4 == 0 and (y % 100 or y % 400 == 0) else 28, 31, 30, 31,
            30, 31, 31, 30, 31, 30, 31][m - 1]
    return f"{ym}-{last:02d}"


def run(folder):
    t = load(folder)
    if not pick(t, "매출집계") and pick(t, "orders"):
        t["매출집계_orders"] = sales_from_orders(pick(t, "orders"))
    rep = Report()

    # 더존 업로드 파일 이름(매입매출전표_*)으로 낸 장부도 받는다
    led = pick(t, "매입매출장") or pick(t, "매입매출전표")
    gjs = pick(t, "일반전표")
    if not led:
        print("매입매출장을 못 찾았습니다. 파일 이름이 매입매출장 또는 매입매출전표 로 시작해야 합니다.")
        sys.exit(1)

    def need(no, *tabs):
        """원천 표가 없으면 그 검사는 미실행이다. 통과로 찍지 않는다."""
        miss = [x for x in tabs if not pick(t, x)]
        if miss:
            rep.skipped.setdefault(no, "원천 없음: " + " · ".join(miss))
        return not miss

    months = sorted({f"{n(r['년도'])}-{n(r['월']):02d}" for r in led})
    LO, HI = months[0], months[-1]
    period = f"{LO}~{HI[-2:]}"

    def inp(ymd):
        return bool(ymd) and LO <= s(ymd)[:7] <= HI

    live = net_off(led)
    rep.netted = len(led) - len(live)
    sales = [r for r in live if n(r["매입매출구분(1-매출/2-매입)"]) == 1]
    buys = [r for r in live if n(r["매입매출구분(1-매출/2-매입)"]) == 2]

    # 규칙
    ch_map = {s(r["장부거래처명"]): r for r in rule("채널_매핑.csv")}
    deny = {s(r["가맹점명"]): r for r in rule("불공제_가맹점.csv")}
    fixed_rules = rule("고정자산_계정.csv")
    fixed_acct = {s(r["계정코드"]) for r in fixed_rules}
    th = {s(r["항목"]): n(r["값"]) for r in rule("임계값.csv")}
    tol = th.get("일자별 허용오차", 10)

    agg = pick(t, "매출집계")
    ti = pick(t, "tax_invoices")
    card = pick(t, "card_approvals")
    bank = pick(t, "bank_transactions")
    plat = pick(t, "platform_settlements")
    ten = pick(t, "tenant_settlements")
    partners = {s(r["사업자번호"]) for r in pick(t, "partners")}
    biz_of_partner = {s(r.get("상호")): s(r["사업자번호"]).replace("-", "")
                      for r in pick(t, "partners") if s(r.get("상호"))}

    # ── E1 형식 · E3 마스터 ───────────────────────────────────
    biz_by_name = {}
    for src in (pick(t, "cash_receipts"), pick(t, "tax_invoices"), pick(t, "card_approvals")):
        for r in src:
            nm = s(r.get("상호") or r.get("가맹점명"))
            if nm and s(r.get("사업자번호")):
                biz_by_name.setdefault(nm, s(r["사업자번호"]))
    blank = [r for r in buys if not s(r["사업자(주민)등록번호"])]
    for r in blank:
        found = biz_by_name.get(s(r["거래처명"]), "")
        rep.add("E3", STOP, "매입 전표에 사업자번호가 없다",
                f"{d3(r)} {s(r['거래처명'])} {n(r['공급가액']):,}원",
                "매입매출전표 11열은 필수",
                "빈칸. 더존 업로드에서 떨어진다" +
                (f". 원천에 {found} 이 있다" if found else ". 원천에도 없다"),
                "cash_receipts · tax_invoices · card_approvals 의 사업자번호",
                "진단항목 E3 · 더존 매입매출전표 업로드 양식",
                tentative="아니오" if found else "예(거래처 확인 필요)")
        rep.fix_mm(d3(r), 2, n(r["과세유형"]), s(r["거래처명"]),
                   found or "<원천에도 없다>", 0, 0, "[E3] 사업자번호 보정",
                   base=s(r["기본계정"]), contra=s(r["상대계정"]))
    unreg = sorted({(s(r["거래처명"]), s(r["사업자(주민)등록번호"])) for r in buys
                    if s(r["사업자(주민)등록번호"]) and
                    s(r["사업자(주민)등록번호"]) not in partners})
    if not need("E3", "partners"):
        unreg = []
    if unreg:
        rep.add("E3", ASK, "거래처 마스터에 없는 매입처",
                " · ".join(f"{a}({b})" for a, b in unreg[:8]) +
                (f" 외 {len(unreg) - 8}" if len(unreg) > 8 else ""),
                "partners 에 등록돼 있을 것",
                f"{len(unreg)}곳이 없다",
                "partners 표", "진단항목 E3", tentative="예(마스터 등록)")

    # ── F1 원천 누락 ──────────────────────────────────────────
    has_agg = need("F1", "매출집계")
    has_agg = need("C1", "매출집계") and has_agg   # 둘 다 평가한다. and 로 이으면 C1 이 안 찍힌다
    src_day = defaultdict(int)
    for r in agg:
        if s(r["취소여부"]) != "정상" or not inp(r["배송완료일"]):
            continue
        ch = s(r["채널"])
        if ch == "TENANT":
            continue
        src_day[(s(r["배송완료일"]), ch)] += n(r["판매금액"])
    led_day = defaultdict(int)
    for r in sales:
        m = ch_map.get(s(r["거래처명"]))
        if not m or s(m["채널코드"]) == "TENANT":
            continue
        led_day[(d3(r), s(m["채널코드"]))] += n(r["공급가액"]) + n(r["부가세"])
    # 월 단위로 먼저 본다. 일자 단위만 보면 조정분개가 월말에 붙은 것이 매번 걸린다
    src_mon, led_mon = defaultdict(int), defaultdict(int)
    for (day, ch), amt in src_day.items():
        src_mon[(day[:7], ch)] += amt
    for (day, ch), amt in led_day.items():
        led_mon[(day[:7], ch)] += amt
    short_mon = {k for k, v in src_mon.items() if v - led_mon.get(k, 0) > tol} if has_agg else set()
    if short_mon:
        by_ch = defaultdict(lambda: [0, []])
        for (day, ch), amt in sorted(src_day.items()):
            if (day[:7], ch) not in short_mon:
                continue
            gap = amt - led_day.get((day, ch), 0)
            if gap > tol:
                by_ch[ch][0] += gap
                by_ch[ch][1].append(day)
        for ch, (gap, days) in sorted(by_ch.items()):
            sup = round(gap / 1.1)
            rep.add("F1", STOP, "원천에 있는 매출이 장부에 없다",
                    f"{ch} {len(days)}일 {gap:,}원 ({', '.join(days[:5])}"
                    f"{' 외' if len(days) > 5 else ''})",
                    "배송완료일 집계 = 장부 매출 집계",
                    f"장부가 {gap:,}원 적다",
                    "매출집계_배송완료일별 · orders 채널별",
                    "진단항목 F1", sup, gap - sup)
            m = next((k for k, v in ch_map.items() if s(v["채널코드"]) == ch), ch)
            rep.fix_mm(month_end(days[0][:7]), 1, 17, m, "", sup, gap - sup,
                       f"[F1] {ch} 누락분 매출 계상", base="401", contra="108")

    # ── C1 공급시기 경계 ──────────────────────────────────────
    # 초과는 채널별로 본다. 기간 전체로 합치면 다른 채널의 누락(F1)이 이 초과를 덮는다.
    # 다음 기 배송분은 채널을 가리지 않고 모은다. 직원이 다른 채널 주문까지 한 채널 매출로
    # 넣는 일이 흔해서다. 그 채널 초과분만큼, 모은 다음 기 배송분 안에서 뺀다
    by_ym = defaultdict(int)                        # 주문월 -> 판매금액
    for r in agg:
        if s(r["취소여부"]) != "정상" or s(r["채널"]) == "TENANT":
            continue
        if s(r["과세구분"]) != "TAXABLE":
            continue
        if inp(r["주문월"]) and not inp(r["배송완료일"]):
            by_ym[s(r["주문월"])] += n(r["판매금액"])
    pool = sum(by_ym.values())
    chans = sorted({c for (d, c) in led_day} | {c for (d, c) in src_day})
    over_by = {ch: sum(v for (d, c), v in led_day.items() if c == ch)
               - sum(v for (d, c), v in src_day.items() if c == ch) for ch in chans}
    for ch in sorted(chans, key=lambda c: -over_by[c]):
        over = over_by[ch]
        if not has_agg or over <= tol or pool <= tol:
            continue
        gross = min(over, pool)
        pool -= gross
        sup = round(gross / 1.1)
        rep.add("C1", STOP, "공급시기가 다음 기인 매출이 이번 기에 들어갔다",
                f"{ch} 주문월 {' · '.join(sorted(by_ym))} · 배송완료 다음 기 {gross:,}원",
                "공급시기는 배송완료일. 다음 기 과세표준으로 간다",
                f"{ch} 장부 매출이 원천 배송완료일 집계보다 {over:,}원 많다",
                "매출집계_배송완료일별 의 주문월과 배송완료일",
                "부가가치세법 제15조 제1항 제1호", -sup, -(gross - sup))
        name = next((k for k, v in ch_map.items() if s(v["채널코드"]) == ch), ch)
        rep.fix_mm(month_end(max(by_ym)), 1, 17, name, "", -sup, -(gross - sup),
                   f"[C1] {ch} 다음 기 배송완료분 제외", base="401", contra="108")

    # ── B7 총액·순액 ──────────────────────────────────────────
    if ch_map and not any(s(r["거래처명"]) in ch_map for r in sales):
        rep.skipped.setdefault("B7", "규칙 미매칭: 채널_매핑.csv 의 장부거래처명이 매출 행에 하나도 없다. 이 고객사 규칙인지 확인")
    if not ch_map:
        rep.skipped.setdefault("B7", "규칙 없음: 채널_매핑.csv")
    ten_rows = [r for r in sales if s(ch_map.get(s(r["거래처명"]), {}).get("채널코드"))
                == "TENANT"]
    # 조정분개는 월말에 한 줄로 붙어 원건과 행 단위로 상계되지 않는다. 순액으로 본다
    if ten_rows and sum(n(r["공급가액"]) for r in ten_rows) > tol:
        sup = sum(n(r["공급가액"]) for r in ten_rows)
        vat = sum(n(r["부가세"]) for r in ten_rows)
        gmv = sum(n(r["판매대금"]) for r in ten if inp(r["판매일"]))
        fee = sum(n(r["판매수수료_당사수익"]) for r in ten if inp(r["판매일"]))
        rep.add("B7", STOP, "중개 채널의 판매대금이 매출에 들어갔다",
                f"입점사 {len(ten_rows)}행 · 공급가액 {sup:,}원 "
                f"(원천 판매대금 {gmv:,}원 · 수수료 {fee:,}원)",
                "중개라 수수료만 수익. 판매대금은 선수금",
                f"입점사 지급액 {sup + vat:,}원이 매출로 잡혀 있다",
                "tenant_settlements 판매대금·판매수수료_당사수익 · 입점 계약서의 공급자 조항",
                "부가가치세법 제29조 제1항·제3항 · 진단항목 B7", -sup, -vat)
        by_m = defaultdict(lambda: [0, 0])
        for r in ten_rows:
            k = f"{n(r['년도'])}-{n(r['월']):02d}"
            by_m[k][0] += n(r["공급가액"])
            by_m[k][1] += n(r["부가세"])
        # 매출로 잡은 직원은 보통 입점사 판매대금 입금을 외상매출금 회수로, 입점사 지급을
        # 비용으로 처리한다. 매출을 음수로 지우면 외상매출금이 한 번 더 줄어든다.
        # 그래서 일반전표의 그 줄을 찾아 선수금으로 되돌린다. 적요나 거래처에 입점사가 든 줄이다
        words = ("입점사",) + tuple(k.split("(")[0].replace("판매분", "").strip()
                                    for k, v in ch_map.items()
                                    if s(v["채널코드"]) == "TENANT" and k)
        words = tuple(w for w in words if w)

        def tenant_line(g):
            return any(w in s(g.get("적요")) + s(g.get("거래처")) for w in words)
        rcv_m, pay_m = defaultdict(int), defaultdict(lambda: defaultdict(int))
        acct_name = {}
        for g in gjs:
            # 일반전표 업로드 양식에는 연도 칸이 없다. 과세기간은 한 해 안이라 장부 연도를 쓴다
            ym = f"{months[0][:4]}-{n(g.get('월')):02d}"
            if not tenant_line(g) or not inp(ym):
                continue
            code = n(g.get("계정과목코드"))
            if code == 108 and n(g.get("대변")) > 0:
                rcv_m[ym] += n(g.get("대변"))
            elif n(g.get("차변")) > 0 and code not in (101, 102, 103, 259):
                pay_m[ym][code] += n(g.get("차변"))
                acct_name[code] = s(g.get("계정과목명"))
        for ym, (a, b) in sorted(x for x in by_m.items() if x[1][0]):
            rep.fix_mm(month_end(ym), 1, 17, s(ten_rows[0]["거래처명"]), "", -a, -b,
                       f"[B7] 입점사 판매대금 총액계상 취소 {ym}", base="401", contra="108")
        if rcv_m or pay_m:
            for ym in sorted(set(rcv_m) | set(pay_m)):
                me = month_end(ym)
                if rcv_m[ym]:
                    rep.fix_gj(me, 108, "외상매출금", "입점사",
                               f"[B7] 입점사 판매대금 입금 외상매출금 회수 취소 {ym}",
                               rcv_m[ym], 0)
                    rep.fix_gj(me, 259, "선수금", "입점사",
                               f"[B7] 입점사 판매대금 입금 선수금 계상 {ym}", 0, rcv_m[ym])
                for code, amt in sorted(pay_m[ym].items()):
                    rep.fix_gj(me, 259, "선수금", "입점사",
                               f"[B7] 입점사 지급 선수금 차감 {ym}", amt, 0)
                    rep.fix_gj(me, code, acct_name.get(code, ""), "입점사",
                               f"[B7] 입점사 지급 {acct_name.get(code, code)} 취소 {ym}", 0, amt)
        else:
            # 일반전표에서 입점사 입금·지급 줄을 못 찾았다. 입금이 매출 금액만큼 외상매출금을
            # 지웠다고 보고 그만큼 선수금으로 옮긴다. 실제 처리는 회계사가 확인한다
            for ym, (a, b) in sorted(x for x in by_m.items() if x[1][0]):
                me = month_end(ym)
                rep.fix_gj(me, 108, "외상매출금", "입점사", f"[B7] 외상매출금 환원 {ym}",
                           a + b, 0)
                rep.fix_gj(me, 259, "선수금", "입점사", f"[B7] 선수금 계상 {ym}", 0, a + b)
            rep.findings[-1]["잠정"] = "예(일반전표의 입점사 입금·지급 처리 확인)"

    # ── B5 불공제 ─────────────────────────────────────────────
    if deny and not any(s(r["거래처명"]) in deny for r in buys):
        rep.skipped.setdefault("B5", "규칙 미매칭: 불공제_가맹점.csv 의 가맹점명이 매입 행에 하나도 없다. 이 고객사 규칙인지 확인")
    if not deny:
        rep.skipped.setdefault("B5", "규칙 없음: 불공제_가맹점.csv")
    bad = [r for r in buys if s(r["거래처명"]) in deny and n(r["과세유형"]) in (57, 61, 51)]
    if bad:
        sup = sum(n(r["공급가액"]) for r in bad)
        by_reason = defaultdict(lambda: [0, 0])
        for r in bad:
            by_reason[s(deny[s(r["거래처명"])]["사유"])][0] += n(r["공급가액"])
            by_reason[s(deny[s(r["거래처명"])]["사유"])][1] += 1
        for reason, (a, c) in sorted(by_reason.items()):
            g = next(s(v["근거"]) for v in deny.values() if s(v["사유"]) == reason)
            rep.add("B5", STOP, "매입세액 불공제 대상이 공제매입으로 들어갔다",
                    f"{reason} {c}건 {a:,}원",
                    "공제매입세액에 넣지 않는다",
                    "신고서 그 밖의 공제매입세액란에 포함돼 있다",
                    "card_approvals 가맹점명·업종", g, -a, -round(a * .1))
        for r in bad:
            rr = deny[s(r["거래처명"])]
            rep.fix_mm(d3(r), 2, n(r["과세유형"]), s(r["거래처명"]),
                       s(r["사업자(주민)등록번호"]), -n(r["공급가액"]), -n(r["부가세"]),
                       f"[B5] {s(rr['사유'])} 공제취소", base=s(r["기본계정"]),
                       contra=s(r["상대계정"]), card_co=s(r["신용카드사명"]))
            rep.fix_gj(d3(r), n(rr["대체계정"]), s(rr["대체계정명"]), s(r["거래처명"]),
                       f"[B5] {s(r['거래처명'])} 불공제 전액 비용",
                       n(r["공급가액"]) + n(r["부가세"]), 0)
            rep.fix_gj(d3(r), 253, "미지급금", s(r["거래처명"]),
                       "[B5] 법인카드 미지급", 0, n(r["공급가액"]) + n(r["부가세"]))

    # ── B6 매입 이중계상 ──────────────────────────────────────
    for no in ("B6", "F2", "B10"):
        need(no, "tax_invoices")
    ti_key = defaultdict(list)
    for r in ti:
        if s(r["매출매입구분"]) == "매입" and inp(r["작성일자"]):
            ti_key[(s(r["작성일자"]), s(r["사업자번호"]), n(r["공급가액"]))].append(r)
    dup = [r for r in buys if n(r["과세유형"]) in (57, 61)
           and (d3(r), s(r["사업자(주민)등록번호"]), n(r["공급가액"])) in ti_key]
    if dup:
        sup = sum(n(r["공급가액"]) for r in dup)
        rep.add("B6", STOP, "같은 매입이 세금계산서와 카드 양쪽에 있다",
                " · ".join(f"{s(r['거래처명'])} {d3(r)} {n(r['공급가액']):,}원"
                           for r in dup[:5]),
                "증빙 하나로만 계상",
                f"{len(dup)}건이 세금계산서 수취분과 그 밖의 공제매입세액에 두 번 들어갔다",
                "tax_invoices 와 card_approvals 를 일자·사업자번호·금액으로 대사",
                "진단항목 B6", -sup, -round(sup * .1))
        for r in dup:
            rep.fix_mm(d3(r), 2, n(r["과세유형"]), s(r["거래처명"]),
                       s(r["사업자(주민)등록번호"]), -n(r["공급가액"]), -n(r["부가세"]),
                       "[B6] 세금계산서와 중복 계상 취소", base=s(r["기본계정"]),
                       contra=s(r["상대계정"]), card_co=s(r["신용카드사명"]))

    # ── B6 카드대금 결제를 매입으로 ──────────────────────────
    # 적요는 은행마다 다르다. "카드 대금" · "법인카드 정산" · "카드대금 결제" 를 다 받고,
    # "카드가맹 수수료"(PG 수수료 출금)는 카드대금이 아니라 뺀다
    def is_cardpay(memo):
        m = s(memo).replace(" ", "")
        return "카드" in m and any(k in m for k in ("대금", "정산", "결제")) and "수수료" not in m
    if not need("B6", "bank_transactions"):
        pass
    pay_key = {(s(b["거래일자"]), n(b["출금액"])) for b in bank
               if is_cardpay(b["적요"]) and inp(b["거래일자"])}
    cardpay = [r for r in buys
               if (d3(r), n(r["공급가액"]) + n(r["부가세"])) in pay_key]
    if cardpay:
        sup = sum(n(r["공급가액"]) for r in cardpay)
        rep.add("B6", STOP, "법인카드 대금 결제가 매입으로 또 계상됐다",
                " · ".join(f"{d3(r)} {n(r['공급가액']) + n(r['부가세']):,}원"
                           for r in cardpay),
                "카드매입은 승인일에 계상. 결제일은 미지급금 상계",
                f"{len(cardpay)}건이 매입으로 두 번",
                "bank_transactions 법인카드 대금 결제 · 카드 승인분",
                "진단항목 B6", -sup, -round(sup * .1))
        for r in cardpay:
            rep.fix_mm(d3(r), 2, n(r["과세유형"]), s(r["거래처명"]),
                       s(r["사업자(주민)등록번호"]), -n(r["공급가액"]), -n(r["부가세"]),
                       "[B6] 카드대금 결제 매입계상 취소", base=s(r["기본계정"]),
                       contra=s(r["상대계정"]))
            # 일반전표에 결제 분개(차 미지급금)가 이미 있으면 매입만 지우면 된다.
            # 또 넣으면 보통예금이 두 번 빠진다. 결제가 매입으로만 들어간 경우에만 세운다
            gross = n(r["공급가액"]) + n(r["부가세"])
            if any(n(g.get("월")) == n(r["월"]) and n(g.get("일")) == n(r["일"])
                   and n(g.get("계정과목코드")) == 253 and n(g.get("차변")) == gross
                   for g in gjs):
                continue
            rep.fix_gj(d3(r), 253, "미지급금", "법인카드사",
                       "[B6] 카드대금 미지급금 상계", gross, 0)
            rep.fix_gj(d3(r), 103, "보통예금", "법인카드사", "[B6] 카드대금 결제", 0, gross)

    # ── F2 전표 종류 오분류 ───────────────────────────────────
    led_ti = {(d3(r), s(r["사업자(주민)등록번호"]), n(r["공급가액"])) for r in buys
              if n(r["과세유형"]) in (51, 53, 54)}
    orphan = [r for key, rs in ti_key.items() for r in rs if key not in led_ti]

    def gj_original(r):
        """세금계산서 한 장이 일반전표의 어느 줄로 들어갔나. 같은 날 같은 거래처의
        차변(공급가액 또는 합계금액)과 대변(합계금액) 줄을 찾는다. 못 찾으면 None."""
        ymd = s(r["작성일자"])
        mon, day = int(ymd[5:7]), int(ymd[8:10])
        sup0, gross0 = n(r["공급가액"]), n(r["공급가액"]) + n(r["부가세"])
        same = [g for g in gjs if n(g.get("월")) == mon and n(g.get("일")) == day
                and s(g.get("거래처")) == s(r["상호"])]
        dr = next((g for g in same if n(g.get("차변")) in (gross0, sup0)), None)
        cr = next((g for g in same if n(g.get("대변")) == gross0), None)
        return (dr, cr) if dr and cr else None

    if orphan:
        sup = sum(n(r["공급가액"]) for r in orphan)
        names = sorted({s(r["상호"]) for r in orphan})
        found = {id(r): gj_original(r) for r in orphan}
        in_gj = sum(1 for v in found.values() if v)
        rep.add("F2", STOP, "세금계산서를 받은 매입이 매입매출전표에 없다",
                f"{' · '.join(names)} {len(orphan)}건 {sup:,}원",
                "매입매출전표 51(면세 계산서는 53). 매입처별 세금계산서합계표에 올라간다",
                ("일반전표로 들어가" if in_gj == len(orphan) else
                 "장부 어디에도 없어" if in_gj == 0 else
                 f"{in_gj}건은 일반전표로 들어가고 {len(orphan) - in_gj}건은 장부 어디에도 없어")
                + " 세금계산서 수취분 매입세액과 합계표에서 빠졌다",
                "tax_invoices 와 매입매출장을 일자·사업자번호·금액으로 대사",
                "진단항목 F2", sup, round(sup * .1),
                tentative="아니오" if in_gj == len(orphan) else "예(일반전표 원건 확인)")
        for r in orphan:
            gross = n(r["공급가액"]) + n(r["부가세"])
            hit = found[id(r)]
            base = s(hit[0].get("계정과목코드")) if hit else ""
            contra = s(hit[1].get("계정과목코드")) if hit else ""
            rep.fix_mm(s(r["작성일자"]), 2, 53 if n(r["부가세"]) == 0 else 51,
                       s(r["상호"]), s(r["사업자번호"]),
                       n(r["공급가액"]), n(r["부가세"]),
                       "[F2] 매입매출전표로 재계상", base=base, contra=contra,
                       elec=1 if s(r["전자여부"]) == "전자" else "")
            if not hit:
                continue
            dr, cr = hit
            rep.fix_gj(s(r["작성일자"]), n(dr.get("계정과목코드")), s(dr.get("계정과목명")),
                       s(r["상호"]), "[F2] 일반전표 계상 취소", -n(dr.get("차변")), 0)
            if n(dr.get("차변")) != gross:     # 부가세를 따로 차변에 둔 원건
                vat_row = next((g for g in gjs if g is not dr and n(g.get("월")) == n(dr.get("월"))
                                and n(g.get("일")) == n(dr.get("일"))
                                and s(g.get("거래처")) == s(r["상호"])
                                and n(g.get("차변")) == gross - n(dr.get("차변"))), None)
                if vat_row:
                    rep.fix_gj(s(r["작성일자"]), n(vat_row.get("계정과목코드")),
                               s(vat_row.get("계정과목명")), s(r["상호"]),
                               "[F2] 일반전표 계상 취소", -n(vat_row.get("차변")), 0)
            rep.fix_gj(s(r["작성일자"]), n(cr.get("계정과목코드")), s(cr.get("계정과목명")),
                       s(r["상호"]), "[F2] 일반전표 계상 취소", 0, -gross)

    def fixed_account(item):
        """품명 키워드로 고정자산 계정을 고른다. 못 고르면 (첫 계정, 잠정)."""
        for fr in fixed_rules:
            words = [w for w in s(fr.get("품명키워드")).replace(",", "·").split("·") if w]
            if s(fr.get("계정과목")) and s(fr["계정과목"]) in item or any(w in item for w in words):
                return s(fr["계정코드"]), False
        return (s(fixed_rules[0]["계정코드"]) if fixed_rules else "212"), True

    wrong_fixed = [r for r in buys if "고정자산" in s(r["품명"])
                   and s(r["기본계정"]) not in fixed_acct]
    for r in wrong_fixed:
        acct, guess = fixed_account(s(r["품명"]))
        rep.add("F2", STOP, "고정자산 매입이 일반매입 계정으로 들어갔다",
                f"{s(r['거래처명'])} {d3(r)} {n(r['공급가액']):,}원",
                "고정자산매입란. 건물등 감가상각자산 취득명세서에 올라간다",
                f"기본계정 {s(r['기본계정'])} 이라 일반매입란으로 간다",
                "tax_invoices 품목 · 계정과목표 유형자산 계정",
                "부가가치세법 시행규칙 별지 제21호서식 고정자산매입란 · 진단항목 F2",
                tentative="예(자산 계정 확인)" if guess else "아니오")
        rep.fix_mm(d3(r), 2, n(r["과세유형"]), s(r["거래처명"]), s(r["사업자(주민)등록번호"]),
                   -n(r["공급가액"]), -n(r["부가세"]), "[F2] 일반매입 계상 취소",
                   base=s(r["기본계정"]), contra=s(r["상대계정"]), elec=s(r["전자세금(1.전자)"]))
        rep.fix_mm(d3(r), 2, n(r["과세유형"]), s(r["거래처명"]), s(r["사업자(주민)등록번호"]),
                   n(r["공급가액"]), n(r["부가세"]), "[F2] 고정자산매입 재계상",
                   base=acct, contra=s(r["상대계정"]), elec=s(r["전자세금(1.전자)"]))

    # ── B10 면세 매입이 과세로 ────────────────────────────────
    zero_ti = {(s(r["작성일자"]), s(r["사업자번호"]), n(r["공급가액"]))
               for r in ti if s(r["매출매입구분"]) == "매입" and n(r["부가세"]) == 0}
    faux = [r for r in buys if n(r["과세유형"]) == 51
            and (d3(r), s(r["사업자(주민)등록번호"]), n(r["공급가액"])) in zero_ti]
    if faux:
        sup = sum(n(r["공급가액"]) for r in faux)
        rep.add("B10", STOP, "면세 계산서를 과세 세금계산서로 입력했다",
                " · ".join(f"{s(r['거래처명'])} {d3(r)} {n(r['공급가액']):,}원"
                           for r in faux[:5]) +
                (f" 외 {len(faux) - 5}건" if len(faux) > 5 else ""),
                "과세유형 53. 매입처별 계산서합계표로 가고 신고서 매입세액란에 안 올라간다",
                f"과세유형 51 로 들어가 세금계산서 수취분 매입에 섞였다(장부 세액 "
                f"{sum(n(r['부가세']) for r in faux):,}원)",
                "tax_invoices 부가세 0원 행",
                "부가가치세법 제26조 제1항 · 제38조 제1항 · 진단항목 B10",
                -sup, -sum(n(r["부가세"]) for r in faux))
        for r in faux:
            rep.fix_mm(d3(r), 2, 51, s(r["거래처명"]), s(r["사업자(주민)등록번호"]),
                       -n(r["공급가액"]), -n(r["부가세"]), "[B10] 과세매입 계상 취소",
                       base=s(r["기본계정"]), contra=s(r["상대계정"]),
                       elec=s(r["전자세금(1.전자)"]))
            rep.fix_mm(d3(r), 2, 53, s(r["거래처명"]), s(r["사업자(주민)등록번호"]),
                       n(r["공급가액"]), 0, "[B10] 면세 계산서로 재계상",
                       base=s(r["기본계정"]), contra=s(r["상대계정"]),
                       elec=s(r["전자세금(1.전자)"]))

    # ── A5 증빙 누락 ──────────────────────────────────────────
    # 정산서의 플랫폼명(○○마켓)과 세금계산서 상호(○○마켓(주))는 다를 수 있다. 사업자번호로 맞춘다.
    # 정산서에 사업자번호 열이 없으면 partners 의 상호로 사업자번호를 찾고, 그래도 없으면 상호로 맞춘다
    def biz_or_name(r, name_col):
        return (s(r.get("사업자번호")).replace("-", "")
                or biz_of_partner.get(s(r.get(name_col)), "")
                or s(r.get(name_col)))
    fee_settle, plat_name_of = defaultdict(int), {}
    if need("A5", "platform_settlements", "tax_invoices"):
        for r in plat:
            if inp(r["정산일"]):
                k = (biz_or_name(r, "플랫폼"), s(r["정산일"])[:7])
                fee_settle[k] += n(r["판매수수료"])
                plat_name_of[k] = s(r["플랫폼"])
    fee_ti = defaultdict(lambda: [0, 0])
    for r in ti:
        if s(r["매출매입구분"]) == "매입" and "수수료" in s(r["품목"]):
            k = (biz_or_name(r, "상호"), s(r["작성일자"])[:7])
            fee_ti[k][0] += n(r["공급가액"])
            fee_ti[k][1] += n(r["부가세"])
    for k in sorted(fee_settle):
        plat_name, ym, fee = plat_name_of[k], k[1], fee_settle[k]
        if k not in fee_ti:
            rep.add("A5", ASK, "정산 수수료에 대응하는 세금계산서가 없다",
                    f"{plat_name} {ym} 수수료 {fee:,}원",
                    "수수료 정산액마다 매입 세금계산서를 받는다",
                    "tax_invoices 에 같은 사업자번호의 해당 월 수수료 행이 없다",
                    "platform_settlements 판매수수료 대 tax_invoices 수수료 (사업자번호·월)",
                    "부가가치세법 제39조 제1항 제2호", tentative="예(수취 요청)")
            continue
        sup, vat = fee_ti[k]
        # 정산 수수료가 부가세 포함이면 합계금액과, 아니면 공급가액과 맞는다. 둘 다 아니면 차이다
        if min(abs(fee - sup - vat), abs(fee - sup)) > tol:
            rep.add("A5", ASK, "정산 수수료와 세금계산서 금액이 다르다",
                    f"{plat_name} {ym} 정산 {fee:,}원 · 세금계산서 공급가액 {sup:,}원 세액 {vat:,}원",
                    "정산 수수료 = 세금계산서 합계금액(또는 공급가액)",
                    f"차이 {fee - sup - vat:+,}원(합계 기준)",
                    "platform_settlements 회차별 판매수수료 · tax_invoices 수수료 행",
                    "부가가치세법 제39조 제1항 제2호", tentative="예(정산서 확인)")

    # ── D 이상 ────────────────────────────────────────────────
    prior = pick(t, "journal_prior")
    if need("D2", "journal_prior"):
        # 전기 장부에 사업자번호 열이 있으면 사업자번호로, 없으면 거래처명으로 맞춘다
        by_biz = any(s(r.get("사업자번호")) for r in prior)
        if by_biz:
            seen = {s(r.get("사업자번호")).replace("-", "") for r in prior}
            new = sorted({s(r["거래처명"]) for r in buys
                          if s(r["사업자(주민)등록번호"]) and
                          s(r["사업자(주민)등록번호"]).replace("-", "") not in seen})
        else:
            seen = {s(r.get("거래처") or r.get("거래처명")) for r in prior}
            new = sorted({s(r["거래처명"]) for r in buys
                          if s(r["거래처명"]) and s(r["거래처명"]) not in seen})
        if new:
            rep.add("D2", ASK, "전기에 없던 매입처",
                    " · ".join(new[:8]) + (f" 외 {len(new) - 8}" if len(new) > 8 else ""),
                    "전기 분개장에 있던 거래처",
                    f"{len(new)}곳이 journal_prior {'사업자번호' if by_biz else '거래처'}에 없다",
                    "journal_prior 거래처 · 신규 거래처 계약서와 증빙", "진단항목 D2", tentative="예")

    return rep, period, led, gjs


# ── 산출 ─────────────────────────────────────────────────────
def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    head = PatternFill("solid", fgColor="1F3864")
    neg = PatternFill("solid", fgColor="FCE4E4")
    thin = Side(style="thin", color="8EA9DB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    wb.remove(wb.active)
    for title, cols, rows, widths in sheets:
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF", size=9)
            cell.fill = head
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = box
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            has_neg = any(isinstance(c.value, (int, float)) and c.value < 0 for c in row)
            for c in row:
                c.border = box
                c.font = Font(size=9)
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = "#,##0"
                if has_neg:
                    c.fill = neg
    wb.save(path)


def impact(mm):
    """수정분개가 납부세액을 얼마나 움직이나. 매출세액 변동 - 매입세액 변동."""
    out = sum(n(r["부가세"]) for r in mm if n(r["매입매출구분(1-매출/2-매입)"]) == 1)
    inn = sum(n(r["부가세"]) for r in mm if n(r["매입매출구분(1-매출/2-매입)"]) == 2)
    return out, inn, out - inn


def coverage_md(findings, cat, skipped=None):
    """진단항목 전수에 대해 무엇이 돌았고 무엇이 안 돌았나. rules/진단항목.csv 가 없으면 비운다.
    skipped 는 자동 항목인데 원천이나 규칙이 없어 못 돈 것. 통과로 찍지 않는다."""
    if not cat:
        return []
    skipped = skipped or {}
    hit = defaultdict(int)
    for f in findings:
        hit[s(f["검사"])] += 1
    cnt = defaultdict(int)
    L = ["## 검토 범위", "",
         "검토항목 전수입니다. 자동은 이 스크립트가 대조한 것, 검토자는 회계사가 검토 방법대로 "
         "직접 보고 확인사항에 적는 것, 자료 미확보는 원천자료가 없어 지금은 못 보는 것입니다.", "",
         "| 번호 | 검토 항목 | 판정 | 수행 | 결과 |", "|---|---|---|---|---|"]
    for r in cat:
        no, how = s(r["번호"]), s(r["수행"])
        cnt[how] += 1
        if how == "자동" and no in skipped and not hit[no]:
            res = f"미실행 ({skipped[no]})"
        elif how == "자동":
            res = f"검출 {hit[no]}건" if hit[no] else "통과"
        elif how == "검토자":
            res = "검토자 수행 필요"
        else:
            res = "자료 확보 후"
        L.append(f"| {no} | {r['검토 항목']} | {r['판정']} | {how} | {res} |")
    L += ["", f"자동 {cnt['자동']} · 검토자 {cnt['검토자']} · 자료 미확보 {cnt['자료 미확보']}. "
          "검토자 수행 필요와 미실행은 통과가 아닙니다. 검토 방법은 rules/진단항목.csv 에 있습니다.", ""]
    return L


def report_md(rep, period, led, gjs):
    stop = [f for f in rep.findings if f["등급"] == STOP]
    ask = [f for f in rep.findings if f["등급"] == ASK]
    out, inn, pay = impact(rep.mm)
    L = [f"# 매입매출장 검토 리포트 · {period}", "",
         f"- 매입매출장 {len(led):,}행 · 일반전표 {len(gjs):,}행",
         f"- 검출 {len(rep.findings)}건 (정지 {len(stop)} · 확인 {len(ask)})",
         f"- 수정분개 매입매출 {len(rep.mm)}줄 · 일반 {len(rep.gj)}줄",
         f"- 수정하면 매출세액 {out:+,}원 · 매입세액 {inn:+,}원 · "
         f"**납부세액 {pay:+,}원** 움직입니다", ""]
    if stop:
        L += ["## 정지", "",
              "**신고서 산출로 넘어가지 않습니다.** 원천이나 처리 기준이 틀렸습니다.", ""]
        for f in stop:
            L += [f"### [{f['검사']}] {f['무엇']}", "",
                  f"- 대상 : {f['대상']}", f"- 기대 : {f['기대']}",
                  f"- 실제 : {f['실제']}", f"- 볼 곳 : {f['볼 곳']}",
                  f"- 근거 : {f['근거']}",
                  f"- 영향 : 공급가액 {f['공급가액 영향']:+,}원 · 세액 {f['세액 영향']:+,}원", ""]
    if ask:
        L += ["## 확인", "", "회계사가 읽고 판정합니다. 판정 결과를 `rules/` 에 올립니다.", ""]
        for f in ask:
            L += [f"### [{f['검사']}] {f['무엇']}", "",
                  f"- 대상 : {f['대상']}", f"- 기대 : {f['기대']}",
                  f"- 실제 : {f['실제']}", f"- 볼 곳 : {f['볼 곳']}",
                  f"- 근거 : {f['근거']}", f"- 잠정 : {f['잠정']}", ""]
    if not rep.findings:
        L += ["## 통과", "", "자동 검사에서 검출이 없습니다. 검토자 항목은 아래 검토 범위대로 "
              "회계사가 본 뒤에 신고서 산출로 넘어갑니다.", ""]
    L += coverage_md(rep.findings, rule("진단항목.csv"), rep.skipped)
    L += ["---", "",
          "수정분개는 파일로만 냅니다. **원장에 적용하지 않았습니다.**",
          "회계사가 보고 더존에 올립니다.", ""]
    return "\n".join(L)


def confirm_md(rep, period):
    L = [f"# 확인사항 · {period}", "",
         "회계사가 정해야 할 것입니다. 정한 결과를 `rules/` 에 행으로 올리면 "
         "다음 달에도 다음 고객사에도 그대로 갑니다.", "",
         "| 항목 | 상황 | 선택지 | 추천안과 이유 | 재무제표·신고서 영향 |",
         "|---|---|---|---|---|"]
    for f in rep.findings:
        if f["등급"] == STOP and f["잠정"] == "아니오":
            opts = "수정분개 반영 / 현행 유지"
            rec = f"수정분개 반영. {f['기대']}"
        else:
            opts = f["잠정"]
            rec = f["기대"]
        L.append(f"| [{f['검사']}] {f['무엇']} | {f['대상']} | {opts} | {rec} | "
                 f"공급가액 {f['공급가액 영향']:+,} · 세액 {f['세액 영향']:+,} |")
    L += ["", "**표에 한 줄로 올라간 것만 등재입니다.** "
          "리포트 글 속에 적어 둔 것은 등재가 아닙니다.", ""]
    return "\n".join(L)


MM_W = [6, 5, 5, 9, 7, 12, 9, 10, 12, 22, 16, 16, 14, 30, 8, 8, 8, 14]
GJ_W = [5, 5, 5, 10, 14, 9, 16, 38, 16, 16]
NOTE_W = [7, 6, 34, 46, 40, 40, 36, 40, 16, 14, 14]


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    folder = os.path.abspath(sys.argv[1])
    out = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 \
        else os.path.join(folder, "검토산출")
    os.makedirs(out, exist_ok=True)

    rep, period, led, gjs = run(folder)

    with open(os.path.join(out, f"검토리포트_{period}.md"), "w", encoding="utf-8") as f:
        f.write(report_md(rep, period, led, gjs))
    with open(os.path.join(out, f"확인사항_{period}.md"), "w", encoding="utf-8") as f:
        f.write(confirm_md(rep, period))
    write_xlsx(os.path.join(out, f"수정_매입매출전표_{period}.xlsx"),
               [("매출자료 & 매입자료", MM_COLS, rep.mm, MM_W)])
    write_xlsx(os.path.join(out, f"수정_일반전표_{period}.xlsx"),
               [("일반전표", GJ_COLS, rep.gj, GJ_W)])
    write_xlsx(os.path.join(out, f"수정분개_근거_{period}.xlsx"),
               [("근거", NOTE_COLS, rep.findings, NOTE_W)])

    stop = sum(1 for f in rep.findings if f["등급"] == STOP)
    ask = len(rep.findings) - stop
    print("=" * 72)
    print(f"매입매출장 검토 · {period}")
    print("=" * 72)
    print(f"  매입매출장 {len(led):,}행 · 일반전표 {len(gjs):,}행")
    o, i2, pay = impact(rep.mm)
    print(f"  검출 {len(rep.findings)}건   정지 {stop} · 확인 {ask}")
    print(f"  수정분개 매입매출 {len(rep.mm)}줄 · 일반 {len(rep.gj)}줄")
    print(f"  납부세액 영향 {pay:+,}원 (매출세액 {o:+,} · 매입세액 {i2:+,})\n")
    for f in rep.findings:
        print(f"  [{f['검사']:<3}] {f['등급']}  {f['무엇']}")
        print(f"         대상 {f['대상'][:80]}")
        print(f"         세액 {f['세액 영향']:+,}원")
    print(f"\n산출 -> {out}")
    return 1 if stop else 0


if __name__ == "__main__":
    sys.exit(main())

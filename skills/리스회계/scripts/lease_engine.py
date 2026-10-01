# -*- coding: utf-8 -*-
"""리스회계(K-IFRS 1116) 결정적 엔진 + 수식 엑셀 생성기.

사용법:
    python lease_engine.py <조건.json> <출력.xlsx>

입력(조건.json) — 계약서에서 추출한 1116 조건:
    계약명, 개시일(YYYY-MM-DD), 리스기간_개월, 월리스료, 부가세구분(별도|포함|면세),
    연할인율(소수, 예 0.05), 할인율_출처, 지급시점(선급|후급), 상각기간_개월(0=리스기간),
    보증금, 리스개설직접원가, 선급리스료, 리스인센티브, 복구원가추정, 단기소액면제(bool)

원칙:
    - 계산은 결정적: 현재가치 -> 유효이자율 상각 -> 정액 감가상각 -> 전표.
    - 보증금은 리스료가 아님 -> 임차보증금 자산으로 분리(PV 계산 제외).
    - 엑셀 4시트 전부 수식. 할인율/리스료를 바꾸면 자동 재계산.
    - 전표 금액은 ROUND 정수 원, 상각표는 정밀값(잔액 정확히 0).
    - 차변합=대변합, 리스부채/사용권자산 기말잔액 0, 총리스료-리스부채=총이자.
"""
import sys, json
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
from datetime import datetime
from dateutil.relativedelta import relativedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ----- 스타일 -----
TITLE = Font(bold=True, size=13)
HDR_F = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="305496")
LBL_F = Font(bold=True)
CALC_FILL = PatternFill("solid", fgColor="E2EFDA")   # 계산(수식)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # 입력값
CHK_FILL = PatternFill("solid", fgColor="DDEBF7")    # 검증
WON = "#,##0"
WON4 = "#,##0.0000"
PCT = "0.0000%"
DATEF = "yyyy-mm-dd"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def q(sheet):
    """크로스시트 참조용 시트명 인용(한글/특수문자 안전). reference/회계모델.md D.7."""
    return "'%s'" % sheet


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_F
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


# ===================================================================
# 1) 결정적 계산(검산용 정답값) — 엑셀 수식과 삼중 일치 확인에 사용
# ===================================================================
def compute(cond):
    n = int(cond["리스기간_개월"])
    pmt_in = float(cond["월리스료"])
    vat = cond.get("부가세구분", "별도")
    pmt = round(pmt_in / 1.1, 6) if vat == "포함" else float(pmt_in)
    annual = float(cond["연할인율"])
    i = annual / 12.0
    prepaid = cond.get("지급시점", "선급") == "선급"
    m = int(cond.get("상각기간_개월", 0) or 0) or n
    deposit = float(cond.get("보증금", 0) or 0)
    idc = float(cond.get("리스개설직접원가", 0) or 0)
    prepay = float(cond.get("선급리스료", 0) or 0)
    incent = float(cond.get("리스인센티브", 0) or 0)
    restore = float(cond.get("복구원가추정", 0) or 0)

    # 현재가치(리스부채): 선급=기시(annuity-due, t=0..n-1), 후급=기말(t=1..n)
    pv = 0.0
    for k in range(1, n + 1):
        exp = (k - 1) if prepaid else k
        pv += pmt / ((1.0 + i) ** exp)
    rou = pv + prepay + idc + restore - incent

    # 리스부채 상각 시뮬레이션(정밀)
    bal = pv
    liab_rows = []
    for k in range(1, n + 1):
        opening = bal
        if prepaid:
            after = opening - pmt
            interest = after * i
            closing = after + interest
        else:
            interest = opening * i
            closing = opening + interest - pmt
        liab_rows.append((opening, pmt, interest, closing))
        bal = closing
    liab_end = bal
    total_interest = sum(r[2] for r in liab_rows)
    total_pmt = pmt * n

    # 정액 감가상각(마지막 잔차 보정 -> 기말 0)
    dep_rows = []
    cum = 0.0
    base = rou / m
    for k in range(1, m + 1):
        dep = (rou - cum) if k == m else base
        cum += dep
        dep_rows.append((dep, cum, rou - cum))
    dep_end = rou - cum

    # 전표(정수 원) 차/대 합계 — 각 전표는 차=대 동일액이라 항상 일치
    debit = credit = 0
    # 사용권자산 전표금액 = 구성요소를 각각 반올림해 더한 값(구성요소 전표와 1원도 안 어긋나게)
    rou_j = round(pv) + round(prepay) + round(idc) + round(restore) - round(incent)
    debit += rou_j               # 사용권자산
    credit += round(pv)          # 리스부채
    credit += round(prepay)      # 선급리스료 대체
    credit += round(idc)         # 보통예금(리스개설직접원가)
    credit += round(restore)     # 복구충당부채
    debit += round(incent)       # 보통예금(리스인센티브 수령)
    if deposit:
        debit += round(deposit); credit += round(deposit)  # 임차보증금
    for (op, p, intr, cl) in liab_rows:
        debit += round(p);    credit += round(p)     # 지급
        debit += round(intr); credit += round(intr)  # 이자
    for (dep, cum, bk) in dep_rows:
        debit += round(dep);  credit += round(dep)   # 감가

    return {
        "n": n, "m": m, "i": i, "pmt": pmt, "prepaid": prepaid,
        "deposit": deposit, "idc": idc, "prepay": prepay,
        "incent": incent, "restore": restore,
        "pv": pv, "rou": rou,
        "total_pmt": total_pmt, "total_interest": total_interest,
        "liab_end": liab_end, "dep_end": dep_end,
        "debit": debit, "credit": credit,
        "liab_rows": liab_rows, "dep_rows": dep_rows,
    }


# ===================================================================
# 2) 수식 엑셀 4시트 생성
# ===================================================================
def build_excel(cond, eng, out_path):
    wb = openpyxl.Workbook()
    s_in = wb.active
    s_in.title = "입력가정"
    s_liab = wb.create_sheet("리스부채상각표")
    s_dep = wb.create_sheet("사용권자산상각표")
    s_jnl = wb.create_sheet("전표")

    start = datetime.strptime(cond["개시일"], "%Y-%m-%d")
    n = eng["n"]; m = eng["m"]; prepaid = eng["prepaid"]
    NM_LIAB = "리스부채상각표"; NM_DEP = "사용권자산상각표"; NM_IN = "입력가정"

    # ---------- 입력가정 ----------
    ws = s_in
    ws["A1"] = "리스회계(K-IFRS 1116) 입력가정 — 노란칸=입력, 초록칸=수식, 파란칸=검증"
    ws["A1"].font = TITLE
    ws.merge_cells("A1:D1")

    def put(r, label, value, fmt=None, fill=INPUT_FILL, note=""):
        ws.cell(row=r, column=1, value=label).font = LBL_F
        c = ws.cell(row=r, column=2, value=value)
        if fmt:
            c.number_format = fmt
        c.fill = fill
        c.border = BORDER
        if note:
            ws.cell(row=r, column=3, value=note)
        return "$B$%d" % r

    ws["A2"] = "[입력값]"; ws["A2"].font = LBL_F
    c_name = put(3, "계약명/ID", cond.get("계약명", "리스계약"))
    c_start = put(4, "개시일", start, DATEF)
    c_n = put(5, "리스기간(개월)", n)
    c_pay = put(6, "지급시점", cond.get("지급시점", "선급"), note="선급=월초 선지급 / 후급=월말")
    c_pmt_in = put(7, "월리스료(입력)", float(cond["월리스료"]), WON)
    c_vat = put(8, "부가세구분", cond.get("부가세구분", "별도"), note="별도/면세=그대로, 포함=÷1.1 공급가액")
    c_rate = put(9, "연할인율", float(cond["연할인율"]), PCT, note=cond.get("할인율_출처", ""))
    c_src = put(10, "할인율 출처", cond.get("할인율_출처", ""))
    c_mIn = put(11, "상각기간(개월,0=리스기간)", int(cond.get("상각기간_개월", 0) or 0))
    c_dep0 = put(12, "보증금", float(cond.get("보증금", 0) or 0), WON, note="리스료 아님 -> 임차보증금 자산(PV 제외)")
    c_idc = put(13, "리스개설직접원가", float(cond.get("리스개설직접원가", 0) or 0), WON)
    c_pre = put(14, "선급리스료", float(cond.get("선급리스료", 0) or 0), WON)
    c_inc = put(15, "리스인센티브", float(cond.get("리스인센티브", 0) or 0), WON)
    c_res = put(16, "복구원가추정", float(cond.get("복구원가추정", 0) or 0), WON)

    ws["A18"] = "[계산 — 수식]"; ws["A18"].font = LBL_F

    def putf(r, label, formula, fmt=None, fill=CALC_FILL, note=""):
        ws.cell(row=r, column=1, value=label).font = LBL_F
        c = ws.cell(row=r, column=2, value=formula)
        if fmt:
            c.number_format = fmt
        c.fill = fill
        c.border = BORDER
        if note:
            ws.cell(row=r, column=3, value=note)
        return "$B$%d" % r

    c_i = putf(19, "월이자율", "=%s/12" % c_rate, PCT, note="연할인율/12")
    c_pmt = putf(20, "월리스료(적용,공급가액)", '=IF(%s="포함",%s/1.1,%s)' % (c_vat, c_pmt_in, c_pmt_in), WON)
    c_m = putf(21, "상각기간(적용)", "=IF(%s>0,%s,%s)" % (c_mIn, c_mIn, c_n), None, note="0이면 리스기간")
    # PV = SUM(현가열)  ← SUMPRODUCT 와 같은 값
    c_pv = putf(22, "리스부채(최초,현재가치)",
                "=SUM(%s!$H$2:$H$%d)" % (q(NM_LIAB), n + 1), WON, note="현가합(유효이자율 PV)")
    c_rou = putf(23, "사용권자산(최초,ROU)",
                 "=%s+%s+%s+%s-%s" % (c_pv, c_pre, c_idc, c_res, c_inc), WON,
                 note="PV+선급+직접원가+복구-인센티브")
    c_tot = putf(24, "총리스료", "=%s*%s" % (c_pmt, c_n), WON)
    c_tint = putf(25, "총이자", "=%s-%s" % (c_tot, c_pv), WON, note="총리스료-리스부채")

    ws["A27"] = "[검증 — 수식]"; ws["A27"].font = LBL_F

    def putchk(r, label, formula, expr_pass, fmt=None):
        ws.cell(row=r, column=1, value=label).font = LBL_F
        c = ws.cell(row=r, column=2, value=formula)
        if fmt:
            c.number_format = fmt
        c.fill = CHK_FILL; c.border = BORDER
        rc = ws.cell(row=r, column=3, value=expr_pass)
        rc.fill = CHK_FILL; rc.border = BORDER; rc.font = LBL_F
        return r

    putchk(28, "리스부채 기말잔액",
           "=%s!$F$%d" % (q(NM_LIAB), n + 1),       # F=기말잔액
           '=IF(ABS($B$28)<1,"PASS 잔액0","FAIL")', WON)
    putchk(29, "사용권자산 기말장부",
           "=%s!$E$%d" % (q(NM_DEP), m + 1),         # E=장부금액
           '=IF(ABS($B$29)<1,"PASS 잔액0","FAIL")', WON)
    putchk(30, "총이자 일치(Σ이자 vs 총이자)",
           "=SUM(%s!$E$2:$E$%d)" % (q(NM_LIAB), n + 1),  # E=이자
           '=IF(ABS($B$30-%s)<1,"PASS","FAIL")' % c_tint, WON)
    # 전표 차대 일치는 전표 시트에서 산출 -> 여기서 참조
    putchk(31, "전표 차대 일치", 0,  # 자리표시 — build 말미에 전표 합계행 참조로 치환
           '=IF(ABS($B$31)<1,"PASS 차대일치","FAIL")', WON)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 34

    # 절대참조 문자열(다른 시트에서 사용)
    A = lambda cell: "%s!%s" % (q(NM_IN), cell)  # noqa
    PV = A(c_pv); ROU = A(c_rou); RATE = A(c_i); PMT = A(c_pmt)
    MM = A(c_m); TINT = A(c_tint)

    # ---------- 리스부채상각표 ----------
    ws = s_liab
    heads = ["회차", "일자", "기초잔액", "지급액", "이자", "기말잔액", "현가계수", "현가"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(heads))
    for k in range(1, n + 1):
        r = k + 1
        dt = start + relativedelta(months=k - 1)
        ws.cell(row=r, column=1, value=k)
        dc = ws.cell(row=r, column=2, value=dt); dc.number_format = DATEF
        # 기초잔액
        if k == 1:
            ws.cell(row=r, column=3, value="=%s" % PV)
        else:
            ws.cell(row=r, column=3, value="=$F$%d" % (r - 1))
        # 지급액
        ws.cell(row=r, column=4, value="=%s" % PMT)
        if prepaid:
            # 이자 = (기초-지급)*월이자율 ; 기말 = 기초-지급+이자
            ws.cell(row=r, column=5, value="=($C$%d-$D$%d)*%s" % (r, r, RATE))
            ws.cell(row=r, column=6, value="=$C$%d-$D$%d+$E$%d" % (r, r, r))
            ws.cell(row=r, column=7, value="=1/POWER(1+%s,$A$%d-1)" % (RATE, r))
        else:
            # 이자 = 기초*월이자율 ; 기말 = 기초+이자-지급
            ws.cell(row=r, column=5, value="=$C$%d*%s" % (r, RATE))
            ws.cell(row=r, column=6, value="=$C$%d+$E$%d-$D$%d" % (r, r, r))
            ws.cell(row=r, column=7, value="=1/POWER(1+%s,$A$%d)" % (RATE, r))
        ws.cell(row=r, column=8, value="=$D$%d*$G$%d" % (r, r))
        for j in range(3, 9):
            cc = ws.cell(row=r, column=j)
            cc.number_format = WON4 if j in (7,) else WON
            cc.border = BORDER
    for col, w in zip("ABCDEFGH", (6, 12, 15, 13, 13, 15, 12, 15)):
        ws.column_dimensions[col].width = w

    # ---------- 사용권자산상각표 ----------
    ws = s_dep
    heads = ["회차", "일자", "당기상각비", "감가상각누계액", "장부금액"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(heads))
    for k in range(1, m + 1):
        r = k + 1
        dt = start + relativedelta(months=k - 1)
        ws.cell(row=r, column=1, value=k)
        dc = ws.cell(row=r, column=2, value=dt); dc.number_format = DATEF
        if k == m:
            ws.cell(row=r, column=3, value="=%s-$D$%d" % (ROU, r - 1) if m > 1 else "=%s" % ROU)
        else:
            ws.cell(row=r, column=3, value="=%s/%s" % (ROU, MM))
        if k == 1:
            ws.cell(row=r, column=4, value="=$C$%d" % r)
        else:
            ws.cell(row=r, column=4, value="=$D$%d+$C$%d" % (r - 1, r))
        ws.cell(row=r, column=5, value="=%s-$D$%d" % (ROU, r))
        for j in range(3, 6):
            cc = ws.cell(row=r, column=j); cc.number_format = WON; cc.border = BORDER
    for col, w in zip("ABCDE", (6, 12, 15, 16, 15)):
        ws.column_dimensions[col].width = w

    # ---------- 전표 ----------
    ws = s_jnl
    heads = ["일자", "전표번호", "적요", "계정과목", "차변", "대변"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(heads))

    row = 2
    vno = [0]

    def line(dt, no, desc, acct, debit_f=None, credit_f=None):
        nonlocal row
        ws.cell(row=row, column=1, value=dt).number_format = DATEF
        ws.cell(row=row, column=2, value=no)
        ws.cell(row=row, column=3, value=desc)
        ws.cell(row=row, column=4, value=acct)
        if debit_f is not None:
            cc = ws.cell(row=row, column=5, value=debit_f); cc.number_format = WON
        if credit_f is not None:
            cc = ws.cell(row=row, column=6, value=credit_f); cc.number_format = WON
        for j in range(1, 7):
            ws.cell(row=row, column=j).border = BORDER
        row += 1

    # 최초인식
    vno[0] += 1; v = vno[0]
    # 사용권자산 = 리스부채 + 선급리스료 + 리스개설직접원가 + 복구원가 - 리스인센티브.
    # 구성요소마다 상대계정이 다르므로 한 줄로 묶지 않는다. 금액은 구성요소별 ROUND 의 합이라
    # 차대가 반올림 차이 없이 맞는다.
    rou_parts = "ROUND(%s,0)+ROUND(%s,0)+ROUND(%s,0)+ROUND(%s,0)-ROUND(%s,0)" % (
        PV, A(c_pre), A(c_idc), A(c_res), A(c_inc))
    line(start, v, "리스개시-사용권자산 인식", "사용권자산", debit_f="=" + rou_parts)
    line(start, v, "리스개시-리스부채 인식", "리스부채", credit_f="=ROUND(%s,0)" % PV)
    if eng["prepay"]:
        line(start, v, "리스개시-선급리스료 대체", "선급리스료", credit_f="=ROUND(%s,0)" % A(c_pre))
    if eng["idc"]:
        line(start, v, "리스개시-리스개설직접원가", "보통예금", credit_f="=ROUND(%s,0)" % A(c_idc))
    if eng["restore"]:
        line(start, v, "리스개시-복구원가", "복구충당부채", credit_f="=ROUND(%s,0)" % A(c_res))
    if eng["incent"]:
        line(start, v, "리스개시-리스인센티브 수령", "보통예금", debit_f="=ROUND(%s,0)" % A(c_inc))
    if eng["deposit"]:
        vno[0] += 1; v = vno[0]
        line(start, v, "임차보증금 지급", "임차보증금", debit_f="=ROUND(%s,0)" % A(c_dep0))
        line(start, v, "임차보증금 지급", "보통예금", credit_f="=ROUND(%s,0)" % A(c_dep0))

    # 매월
    maxk = max(n, m)
    for k in range(1, maxk + 1):
        dt = start + relativedelta(months=k - 1)
        lr = k + 1  # 상각표 행
        if k <= n:
            # 지급(선급은 월초) : 차)리스부채 / 대)보통예금
            vno[0] += 1; v = vno[0]
            ref = "%s!$D$%d" % (q(NM_LIAB), lr)
            line(dt, v, "%d회차 리스료 지급" % k, "리스부채", debit_f="=ROUND(%s,0)" % ref)
            line(dt, v, "%d회차 리스료 지급" % k, "보통예금", credit_f="=ROUND(%s,0)" % ref)
            # 이자 : 차)이자비용 / 대)리스부채
            vno[0] += 1; v = vno[0]
            ref = "%s!$E$%d" % (q(NM_LIAB), lr)
            line(dt, v, "%d회차 이자비용" % k, "이자비용", debit_f="=ROUND(%s,0)" % ref)
            line(dt, v, "%d회차 이자비용" % k, "리스부채", credit_f="=ROUND(%s,0)" % ref)
        if k <= m:
            # 감가 : 차)감가상각비 / 대)감가상각누계액
            vno[0] += 1; v = vno[0]
            ref = "%s!$C$%d" % (q(NM_DEP), k + 1)
            line(dt, v, "%d회차 감가상각" % k, "감가상각비", debit_f="=ROUND(%s,0)" % ref)
            line(dt, v, "%d회차 감가상각" % k, "감가상각누계액", credit_f="=ROUND(%s,0)" % ref)

    last_data = row - 1
    # 합계검증
    row += 1
    sumrow = row
    ws.cell(row=row, column=4, value="합계").font = LBL_F
    dc = ws.cell(row=row, column=5, value="=SUM($E$2:$E$%d)" % last_data); dc.number_format = WON
    cc = ws.cell(row=row, column=6, value="=SUM($F$2:$F$%d)" % last_data); cc.number_format = WON
    for j in (4, 5, 6):
        ws.cell(row=row, column=j).fill = CHK_FILL; ws.cell(row=row, column=j).border = BORDER
    row += 1
    ws.cell(row=row, column=4, value="차대 검증").font = LBL_F
    rc = ws.cell(row=row, column=5,
                 value='=IF(ABS($E$%d-$F$%d)<1,"PASS 차대일치","FAIL")' % (sumrow, sumrow))
    rc.font = LBL_F; rc.fill = CHK_FILL; rc.border = BORDER
    ws.cell(row=row, column=4).fill = CHK_FILL; ws.cell(row=row, column=4).border = BORDER
    for col, w in zip("ABCDEF", (12, 10, 26, 16, 14, 14)):
        ws.column_dimensions[col].width = w

    # 입력가정 31행의 전표 차대 참조 자리표시 치환
    diff = "='전표'!$E$%d-'전표'!$F$%d" % (sumrow, sumrow)
    s_in["B31"] = diff

    wb.save(out_path)
    return sumrow


def main():
    cond_path, out_path = sys.argv[1], sys.argv[2]
    with open(cond_path, encoding="utf-8") as f:
        cond = json.load(f)

    # 단기·소액 면제 자동 판단(K-IFRS 1116 문단 5)
    n = int(cond["리스기간_개월"])
    exempt = bool(cond.get("단기소액면제", False)) or (n <= 12 and not cond.get("매수선택권", False))

    eng = compute(cond)
    sumrow = build_excel(cond, eng, out_path)

    result = {
        "계약명": cond.get("계약명", ""),
        "단기소액면제_플래그": exempt,
        "n": eng["n"], "m": eng["m"], "월이자율": eng["i"],
        "리스부채_PV": round(eng["pv"], 2),
        "사용권자산_ROU": round(eng["rou"], 2),
        "총리스료": round(eng["total_pmt"], 2),
        "총이자": round(eng["total_interest"], 2),
        "리스부채_기말잔액": round(eng["liab_end"], 6),
        "사용권자산_기말장부": round(eng["dep_end"], 6),
        "전표_차변합": eng["debit"], "전표_대변합": eng["credit"],
        "검증": {
            "차대일치": eng["debit"] == eng["credit"],
            "리스부채잔액0": abs(eng["liab_end"]) < 1e-6,
            "사용권자산잔액0": abs(eng["dep_end"]) < 1e-6,
            "총이자식": abs((eng["total_pmt"] - eng["pv"]) - eng["total_interest"]) < 1e-6,
        },
        "전표_합계행": sumrow,
        "출력파일": out_path,
    }
    result["all_pass"] = all(result["검증"].values())
    with open(out_path.rsplit(".", 1)[0] + "_result.json", "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print("=== 리스회계 엔진 결과 ===")
    print("계약:", result["계약명"], "| 기간 %d개월 상각 %d개월" % (eng["n"], eng["m"]))
    print("단기·소액 면제 플래그:", exempt)
    print("리스부채(PV): %15s" % format(round(eng["pv"]), ","))
    print("사용권자산(ROU): %13s" % format(round(eng["rou"]), ","))
    print("총리스료: %19s" % format(round(eng["total_pmt"]), ","))
    print("총이자: %21s" % format(round(eng["total_interest"]), ","))
    print("전표 차변합=%s 대변합=%s" % (format(eng["debit"], ","), format(eng["credit"], ",")))
    print("리스부채 기말잔액:", round(eng["liab_end"], 6), "| 사용권자산 기말장부:", round(eng["dep_end"], 6))
    print("ALL_PASS:", result["all_pass"])
    print("출력:", out_path)
    sys.exit(0 if result["all_pass"] else 1)


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""계약서(.docx/.pdf/.txt)에서 텍스트를 추출해 표준출력으로 덤프.

사용법: python read_contract.py <계약서.docx|.pdf|.txt>
- .docx: python-docx로 문단·표 텍스트 추출.
- .pdf : pypdf로 텍스트 레이어 추출(디지털 PDF). 텍스트가 거의 없으면(스캔본) 경고 출력
         -> 상위 에이전트(Claude)가 Read 도구로 PDF를 직접 읽어 추출(OCR 대체).
- .txt/.md: 그대로 읽음.
- 추출된 텍스트를 보고 상위 에이전트(Claude)가 1116 조건을 표로 정리해 조건.json을 만든다.
"""
import sys
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def from_docx(path):
    from docx import Document
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)


def from_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for i, page in enumerate(reader.pages, 1):
        txt = (page.extract_text() or "").strip()
        if txt:
            parts.append("[p.%d]\n%s" % (i, txt))
    text = "\n".join(parts)
    if len(text.strip()) < 30:
        # 텍스트 레이어 없음(스캔 이미지 PDF) -> 상위 에이전트가 Read로 직접 읽도록 신호
        print("[안내] 이 PDF는 텍스트가 거의 없습니다(스캔본 가능성). "
              "Read 도구로 PDF를 직접 읽어 1116 조건을 추출하세요: %s" % path)
        return ""
    return text


def main():
    path = sys.argv[1]
    low = path.lower()
    if low.endswith(".docx"):
        print(from_docx(path))
    elif low.endswith(".pdf"):
        out = from_pdf(path)
        if out:
            print(out)
    elif low.endswith((".txt", ".md")):
        with open(path, encoding="utf-8") as f:
            print(f.read())
    else:
        print("[안내] 지원 형식: .docx/.pdf/.txt/.md. 그 외는 Read 도구로 직접 읽으세요:", path)


if __name__ == "__main__":
    main()
